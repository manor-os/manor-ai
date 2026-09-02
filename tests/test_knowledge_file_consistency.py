"""Knowledge file read/write/edit consistency (architecture fix).

One logical file has three representations — FS bytes, the Document projection
(+ metadata.content_text), and pgvector chunks — and two path resolvers.
These tests pin the unified behavior:

  * every byte change invalidates the derived representations (vector_status
    back to pending, content_text fork dropped) so RAG cannot serve stale
    content — the fundraising_ops.md "saved but reads stale" incident;
  * read/edit locate a workspace-scoped file by its logical name, so write and
    read address the SAME file (no write-A-read-B fork);
  * a genuine miss returns same-basename candidates instead of dead-ending.
"""
from __future__ import annotations

import asyncio
import json
from pathlib import Path

import pytest

from packages.core.ai.tools.generate_file.tool import _generate_file_handler as _generate_file
from sqlalchemy import event, select

from packages.core.ai.tools import file_tools
from packages.core.models.base import generate_ulid
from packages.core.models.document import Document, VectorStatus
from packages.core.models.event import EventLog
from packages.core.models.workspace import Workspace


@pytest.fixture
def fs_enabled(monkeypatch, tmp_path):
    from packages.core.config import get_settings

    settings = get_settings()
    old = (settings.MANOR_FS_ENABLED, settings.MANOR_FS_ROOT, settings.DEPLOYMENT_MODE)
    settings.MANOR_FS_ENABLED = True
    settings.MANOR_FS_ROOT = str(tmp_path)
    settings.DEPLOYMENT_MODE = "oss"

    async def allow_file_mutation(**_kwargs):
        return None

    async def allow_file_resource_access(**_kwargs):
        return None

    async def allow_test_paths(*_args, **_kwargs):
        return set()

    monkeypatch.setattr(file_tools, "runtime_guard_file_mutation", allow_file_mutation)
    monkeypatch.setattr("packages.core.ai.runtime.file_actions.runtime_guard_file_mutation", allow_file_mutation)
    monkeypatch.setattr(
        file_tools,
        "runtime_guard_file_resource_access",
        allow_file_resource_access,
    )
    monkeypatch.setattr(file_tools, "_blocked_doc_paths", allow_test_paths)
    try:
        import packages.core.services.knowledge_sync as ks
        monkeypatch.setattr(ks, "_schedule_document_reembed", lambda _id: None)
    except Exception:
        pass
    yield
    settings.MANOR_FS_ENABLED, settings.MANOR_FS_ROOT, settings.DEPLOYMENT_MODE = old


async def _seed_workspace(db_session, entity_id, workspace_id):
    workspace = Workspace(id=workspace_id, entity_id=entity_id, name="Ops WS", operating_model={})
    db_session.add(workspace)
    await db_session.flush()
    from packages.core.services.workspace_artifacts import ensure_workspace_artifact_folder
    await ensure_workspace_artifact_folder(db_session, workspace)
    await db_session.commit()
    from packages.core.services.entity_fs import provision_entity_filesystem
    provision_entity_filesystem(entity_id)


async def _doc_for(db_session, entity_id, name):
    db_session.expire_all()
    return (await db_session.execute(
        select(Document).where(Document.entity_id == entity_id, Document.name == name)
    )).scalar_one_or_none()


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("outcome", ["success", "late_failure", "projection_failure"])
@pytest.mark.parametrize("mutation", ["text", "format", "paragraph", "paragraph_format", "page_layout"])
async def test_precise_office_patch_preserves_knowledge_identity_and_atomicity(
    db_session, fs_enabled, monkeypatch, extension, outcome, mutation,
):
    import copy
    from tests.test_office_table_patch import table_operations, cell_operation, _open
    from tests.test_office_cell_format import format_operation, assert_cell_style
    from tests.test_office_text_set import text_operations, text_operation, open_paragraphs
    from tests.test_office_template_layout import paragraph_operation, page_operation, assert_paragraph_style

    entity_id, workspace_id = generate_ulid(), generate_ulid()
    await _seed_workspace(db_session, entity_id, workspace_id)
    name = f"table-cells.{extension}"
    generation = table_operations(extension) + [cell_operation(extension, value="Generated value")]
    if mutation in {"paragraph", "paragraph_format", "page_layout"}:
        generation = text_operations(extension) + [text_operation(extension, "Generated value")]
    generated = json.loads(await _generate_file(
        entity_id, kind="document", name=name, workspace_id=workspace_id,
        operations=generation,
    ))
    assert generated.get("created") is True, generated
    original = await _doc_for(db_session, entity_id, name)
    original_id, original_metadata = original.id, copy.deepcopy(original.metadata_)
    before_events = set((await db_session.scalars(select(EventLog.id).where(EventLog.entity_id == entity_id))).all())
    path = Path(file_tools._get_entity_root(entity_id), original.fs_path)
    before = path.read_bytes()
    read = json.loads(await file_tools._read_file(entity_id, path=name, workspace_id=workspace_id, include_structure=True))
    assert "Generated value" in read["content"]
    operations = [cell_operation(extension, value="Native patch")]
    if mutation == "format":
        operations = [format_operation(extension)]
    elif mutation == "paragraph":
        operations = [text_operation(extension, "Native patch")]
    elif mutation == "paragraph_format":
        operations = [paragraph_operation(extension)]
    elif mutation == "page_layout":
        operations = [page_operation(extension)]
    if outcome == "late_failure":
        invalid = format_operation(extension, cell="Z99") if mutation == "format" else cell_operation(extension, cell="Z99")
        operations.append(text_operation(extension, index=999) if mutation in {"paragraph", "paragraph_format", "page_layout"} else invalid)
    elif outcome == "projection_failure":
        def reject_metadata(_path):
            raise OSError("injected table response metadata failure")

        monkeypatch.setattr(file_tools, "_file_meta", reject_metadata)
    result = json.loads(await file_tools._patch_file(
        entity_id, path=name, workspace_id=workspace_id, expected_sha256=read["source_sha256"], operations=operations,
    ))
    document = await _doc_for(db_session, entity_id, name)
    assert document.id == original_id
    assert document.metadata_["origin"]["workspace_id"] == workspace_id
    if outcome == "success":
        assert result.get("patched") is True, result
        assert result["document_id"] == original_id
        assert result["source_sha256"] != read["source_sha256"]
        if mutation in {"paragraph", "paragraph_format", "page_layout"}:
            native, paragraphs = open_paragraphs(str(path), extension)
            assert [p.text for p in paragraphs] == ["same", "Native patch" if mutation == "paragraph" else "Generated value", "same"]
            if mutation == "paragraph_format":
                assert_paragraph_style(paragraphs[1], extension)
            elif mutation == "page_layout":
                assert native.sections[0].left_margin.pt == 54 if extension == "docx" else native.slide_width.pt == 960
        else:
            _, table = _open(str(path), extension)
            assert table.cell(1, 1).text == ("Generated value" if mutation == "format" else "Native patch")
            if mutation == "format":
                assert_cell_style(table.cell(1, 1), extension)
            assert table.cell(0, 0).text == table.cell(0, 1).text == table.cell(1, 0).text == "same"
        assert document.vector_status == VectorStatus.PENDING
    else:
        assert result.get("error"), result
        if outcome == "late_failure":
            assert result["operation_index"] == 1
        assert path.read_bytes() == before
        assert document.metadata_ == original_metadata
        assert set((await db_session.scalars(select(EventLog.id).where(EventLog.entity_id == entity_id))).all()) == before_events


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["docx", "pptx", "xlsx"])
@pytest.mark.parametrize("scoped", [False, True])
async def test_office_operations_generate_read_patch_same_knowledge_document(db_session, fs_enabled, extension, scoped):
    import hashlib
    from tests.test_office_operation_generation import office_operations
    from packages.core.services.entity_fs import provision_entity_filesystem

    entity_id, workspace_id = generate_ulid(), generate_ulid() if scoped else None
    if scoped:
        await _seed_workspace(db_session, entity_id, workspace_id)
    else:
        provision_entity_filesystem(entity_id)
    name = f"native.{extension}"
    result = json.loads(await _generate_file(
        entity_id, kind="document", name=name, operations=office_operations(extension), workspace_id=workspace_id,
    ))
    assert result.get("created") is True, result
    assert result["operations_applied"] == len(office_operations(extension))
    document = await _doc_for(db_session, entity_id, name)
    assert document is not None and document.file_type == extension
    document_id = document.id
    assert (document.metadata_ or {}).get("origin", {}).get("workspace_id") == workspace_id
    assert (document.fs_path != name) is scoped
    path = Path(file_tools._get_entity_root(entity_id), document.fs_path)
    read = json.loads(await file_tools._read_file(entity_id, path=name, workspace_id=workspace_id, include_structure=True))
    assert "Proposal" in read["content"]
    assert read["source_sha256"] == hashlib.sha256(path.read_bytes()).hexdigest() == result["document"]["source_sha256"]
    assert read["structure"]
    followup = {"op": "text.replace", "old_text": "Proposal", "new_text": "Revised"}
    if extension == "xlsx":
        followup = {"op": "cell.set", "sheet": "Sheet", "cell": "A1", "value": "Revised"}
    patched = json.loads(await file_tools._patch_file(
        entity_id, path=name, workspace_id=workspace_id, expected_sha256=read["source_sha256"], operations=[followup],
    ))
    assert patched.get("patched") is True, patched
    assert patched["document_id"] == document_id
    reread = json.loads(await file_tools._read_file(entity_id, path=name, workspace_id=workspace_id))
    assert "Revised" in reread["content"]
    assert reread["source_sha256"] == patched["source_sha256"] != read["source_sha256"]
    document = await _doc_for(db_session, entity_id, name)
    assert document.id == document_id
    assert document.vector_status == VectorStatus.PENDING


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["docx", "pptx", "xlsx"])
@pytest.mark.parametrize("outcome", ["late_failure", "projection_failure", "approval_denied", "existing", "concurrent_create"])
async def test_office_operation_generation_never_commits_partial_or_overwrites(
    db_session, fs_enabled, monkeypatch, extension, outcome,
):
    from tests.test_office_operation_generation import office_operations
    from packages.core.ai.runtime import file_actions, generated_files
    from packages.core.contracts.file_engine import normalize_file_patch_operation
    from packages.core.services.entity_fs import provision_entity_filesystem

    entity_id = generate_ulid()
    provision_entity_filesystem(entity_id)
    root = Path(file_tools._get_entity_root(entity_id))
    name = f"atomic.{extension}"
    path = root / name
    operations = office_operations(extension)
    guarded = []

    async def guard(**kwargs):
        guarded.append(kwargs)
        if outcome == "approval_denied":
            return json.dumps({"error": "approval_required"})

    monkeypatch.setattr(file_actions, "runtime_guard_file_mutation", guard)
    if outcome == "late_failure":
        operations.append({"op": "cell.set", "cell": "A0", "value": "bad"} if extension == "xlsx" else {"op": "text.replace", "old_text": "missing", "new_text": "no"})
    elif outcome == "projection_failure":
        async def fail_metadata(_path):
            raise OSError("injected response metadata failure")

        monkeypatch.setattr(generated_files, "runtime_generated_file_metadata", fail_metadata)
    elif outcome == "existing":
        path.write_bytes(b"original user file")
    elif outcome == "concurrent_create":
        real_write = file_actions.RuntimeFileProjectionTransaction.write_bytes

        def raced_write(transaction, *args, **kwargs):
            path.write_bytes(b"concurrent user file")
            assert kwargs["require_missing"] is True
            return real_write(transaction, *args, **kwargs)

        monkeypatch.setattr(file_actions.RuntimeFileProjectionTransaction, "write_bytes", raced_write)

    result = json.loads(await _generate_file(entity_id, kind="document", name=name, operations=operations))
    assert result.get("error"), result
    if outcome == "existing":
        assert result["error"] == "file_already_exists"
        assert guarded == []
        assert path.read_bytes() == b"original user file"
    else:
        assert len(guarded) == 1
        assert guarded[0]["approval_payload"] == {"file_type": extension, "operations": [normalize_file_patch_operation(op) for op in operations]}
        assert guarded[0]["tool_name"] == "generate_file"
        assert json.loads(guarded[0]["content_preview"]) == guarded[0]["approval_payload"]["operations"]
        if outcome == "concurrent_create":
            assert path.read_bytes() == b"concurrent user file"
        else:
            assert not path.exists()
    if outcome == "late_failure":
        assert result["operation_index"] == len(operations) - 1
    assert await _doc_for(db_session, entity_id, name) is None
    assert not (await db_session.scalars(select(EventLog.id).where(EventLog.entity_id == entity_id))).all()


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["json", "diagram.json", "csv", "tsv", "docx", "pptx", "xlsx", "xlsm", "pdf"])
async def test_structural_patch_preserves_real_knowledge_identity(db_session, fs_enabled, extension):
    from tests.test_file_engine_structured_patch import _source
    from packages.core.services import knowledge_sync
    from packages.core.services.entity_fs import provision_entity_filesystem

    entity_id = generate_ulid()
    provision_entity_filesystem(entity_id)
    root = Path(file_tools._get_entity_root(entity_id))
    path = root / f"source.{extension}"
    operations = _source(path, extension)
    synced = await knowledge_sync.sync_file_to_knowledge(
        entity_id=entity_id, abs_path=str(path), entity_root=str(root), source="manual", db=db_session, commit=True,
    )
    assert synced.synced is True
    original_id = synced.document_id
    read = json.loads(await file_tools._read_file(entity_id, path=path.name, include_structure=True))
    result = json.loads(await file_tools._patch_file(
        entity_id, path=path.name, operations=operations, expected_sha256=read["source_sha256"],
    ))
    assert result.get("patched") is True, result
    assert result["document_id"] == original_id
    document = await _doc_for(db_session, entity_id, path.name)
    assert document is not None and document.id == original_id
    assert document.file_type == extension
    assert document.fs_path == path.name
    assert document.vector_status == VectorStatus.PENDING
    assert document.metadata_["file_integrity"]["mtime_ns"] == path.stat().st_mtime_ns
    assert result["source_sha256"] != read["source_sha256"]


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("outcome", ["success", "late_failure", "projection_failure"])
async def test_rich_text_patch_is_atomic_with_real_knowledge(
    db_session, fs_enabled, monkeypatch, extension, outcome,
):
    import copy
    from tests.test_office_text_patch import _document, _reopen, _run
    from packages.core.services import knowledge_sync
    from packages.core.services.entity_fs import provision_entity_filesystem

    entity_id = generate_ulid()
    provision_entity_filesystem(entity_id)
    root = Path(file_tools._get_entity_root(entity_id))
    path = root / f"rich-text.{extension}"
    source, paragraph = _document(extension)
    _run(paragraph, "OL", bold=True)
    _run(paragraph, "D suffix", italic=True)
    source.save(path)
    original = path.read_bytes()
    synced = await knowledge_sync.sync_file_to_knowledge(
        entity_id=entity_id, abs_path=str(path), entity_root=str(root), source="manual", db=db_session, commit=True,
    )
    assert synced.synced is True
    document = await _doc_for(db_session, entity_id, path.name)
    original_metadata = copy.deepcopy(document.metadata_)
    original_updated = document.updated_at
    original_events = list((await db_session.scalars(
        select(EventLog.id).where(EventLog.entity_id == entity_id),
    )).all())
    read = json.loads(await file_tools._read_file(entity_id, path=path.name))
    operations = [
        {"op": "text.replace", "old_text": "OLD", "new_text": "OLD updated", "replace_all": True},
        {"op": "text.replace", "old_text": "suffix", "new_text": "final"},
    ]
    if outcome == "late_failure":
        operations.append({"op": "text.replace", "old_text": "absent", "new_text": "no"})
    elif outcome == "projection_failure":
        def reject_response_metadata(*_args):
            raise OSError("injected response metadata failure")

        monkeypatch.setattr(file_tools, "_file_meta", reject_response_metadata)

    result = json.loads(await file_tools._patch_file(
        entity_id, path=path.name, operations=operations, expected_sha256=read["source_sha256"],
    ))

    document = await _doc_for(db_session, entity_id, path.name)
    assert document.id == synced.document_id
    assert document.fs_path == path.name
    if outcome == "success":
        assert result.get("patched") is True, result
        assert result["replacements"] == 2
        assert result["operations_applied"] == 2
        assert result["document_id"] == synced.document_id
        _, edited = _reopen({**result, "_persisted_bytes": path.read_bytes()}, extension)
        assert edited.text == "OLD updated final"
        assert edited.runs[0].font.bold is True
        assert edited.runs[1].font.italic is True
        reread = json.loads(await file_tools._read_file(entity_id, path=path.name))
        assert "OLD updated final" in reread["content"]
        assert reread["source_sha256"] == result["source_sha256"] != read["source_sha256"]
        assert document.vector_status == VectorStatus.PENDING
    else:
        assert result.get("error"), result
        if outcome == "late_failure":
            assert result["operation_index"] == 2
        assert path.read_bytes() == original
        assert document.metadata_ == original_metadata
        assert document.updated_at == original_updated
        events = list((await db_session.scalars(
            select(EventLog.id).where(EventLog.entity_id == entity_id),
        )).all())
        assert events == original_events


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["xlsx", "xlsm"])
@pytest.mark.parametrize("outcome", ["success", "late_failure", "projection_failure"])
@pytest.mark.parametrize("template", [False, True])
async def test_spreadsheet_patch_preserves_data_and_reports_real_commit(
    db_session, fs_enabled, monkeypatch, extension, outcome, template,
):
    import copy
    from datetime import datetime
    from openpyxl import load_workbook
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from tests.test_file_engine_structured_patch import _source
    from packages.core.services import knowledge_sync
    from packages.core.services.entity_fs import provision_entity_filesystem

    entity_id = generate_ulid()
    provision_entity_filesystem(entity_id)
    root = Path(file_tools._get_entity_root(entity_id))
    path = root / f"spreadsheet.{extension}"
    _source(path, extension)
    workbook = load_workbook(path, keep_vba=extension == "xlsm", rich_text=True)
    workbook.active["A1"] = datetime(2026, 8, 30, 10, 15)
    rich_text = CellRichText("Keep ", TextBlock(InlineFont(b=True), "rich text"))
    workbook.active["B1"] = rich_text
    workbook.active["C1"] = "=2+3"
    workbook.save(path)
    workbook.close()
    if workbook.vba_archive is not None:
        workbook.vba_archive.close()
    original = path.read_bytes()
    synced = await knowledge_sync.sync_file_to_knowledge(
        entity_id=entity_id, abs_path=str(path), entity_root=str(root), source="manual", db=db_session, commit=True,
    )
    assert synced.synced is True
    original_metadata = copy.deepcopy((await _doc_for(db_session, entity_id, path.name)).metadata_)
    before_events = set((await db_session.scalars(select(EventLog.id).where(EventLog.entity_id == entity_id))).all())
    read = json.loads(await file_tools._read_file(entity_id, path=path.name))
    operations = [
        {"op": "sheet.add", "sheet": "New"},
        {"op": "row.append", "sheet": "New", "values": ["Name", "Value"]},
        {"op": "cell.set", "cell": "A1", "value": "Updated"},
    ]
    if template:
        from tests.test_spreadsheet_template_layout import sheet_layout, print_layout
        operations[0:0] = [
            {**sheet_layout(), "sheet": workbook.sheetnames[0]},
            {**print_layout(), "sheet": workbook.sheetnames[0]},
            {"op": "cell.format", "cell": "A1", "format": {"horizontal": "center", "borders": {"bottom": {"style": "double", "color": "174C46"}}}},
        ]
    if outcome == "late_failure":
        operations[-1].pop("value")
    elif outcome == "projection_failure":
        def reject_response_metadata(*_args):
            raise OSError("injected response metadata failure")

        monkeypatch.setattr(file_tools, "_file_meta", reject_response_metadata)

    result = json.loads(await file_tools._patch_file(
        entity_id, path=path.name, operations=operations, expected_sha256=read["source_sha256"],
    ))

    document = await _doc_for(db_session, entity_id, path.name)
    assert document.id == synced.document_id and document.file_type == extension
    if outcome == "success":
        assert result.get("patched") is True, result
        assert result["document_id"] == synced.document_id
        assert result["operations_applied"] == len(operations)
        assert result["operation_results"][-1]["updated_cells"]["A1"]["old"] == "2026-08-30 10:15:00"
        assert result["source_sha256"] != read["source_sha256"]
        edited = load_workbook(path, rich_text=True, keep_vba=extension == "xlsm")
        try:
            assert edited.active["A1"].value == "Updated"
            assert edited.active["B1"].value == rich_text
            assert edited.active["C1"].value == "=2+3"
            assert edited["New"]["A1"].value == "Name"
            assert edited["New"]["B1"].value == "Value"
            if template:
                assert edited.active.column_dimensions["A"].width == 42
                assert edited.active.freeze_panes == "B4"
                assert edited.active.page_setup.fitToWidth == 1
                assert edited.active["A1"].alignment.horizontal == "center"
                assert edited.active["A1"].border.bottom.style == "double"
            if edited.vba_archive is not None:
                assert edited.vba_archive.read("xl/vbaProject.bin") == b"opaque-macro-payload-must-be-preserved"
        finally:
            edited.close()
            if edited.vba_archive is not None:
                edited.vba_archive.close()
        reread = json.loads(await file_tools._read_file(entity_id, path=path.name))
        assert "Updated" in reread["content"]
        assert reread["source_sha256"] == result["source_sha256"]
        assert document.vector_status == VectorStatus.PENDING
    else:
        assert result.get("error"), result
        if outcome == "late_failure":
            assert result["operation_index"] == len(operations) - 1
        assert path.read_bytes() == original
        assert document.metadata_ == original_metadata
        events = set((await db_session.scalars(select(EventLog.id).where(EventLog.entity_id == entity_id))).all())
        assert events == before_events


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["xlsx", "xlsm"])
@pytest.mark.parametrize("outcome", ["success", "late_failure", "projection_failure"])
async def test_table_append_is_atomic_through_real_knowledge(
    db_session, fs_enabled, monkeypatch, extension, outcome,
):
    import copy
    from openpyxl import load_workbook
    from tests.test_spreadsheet_append import _table_source
    from packages.core.services import knowledge_sync
    from packages.core.services.entity_fs import provision_entity_filesystem

    entity_id = generate_ulid()
    provision_entity_filesystem(entity_id)
    root = Path(file_tools._get_entity_root(entity_id))
    path = root / f"table.{extension}"
    _table_source(path)
    original = path.read_bytes()
    synced = await knowledge_sync.sync_file_to_knowledge(
        entity_id=entity_id, abs_path=str(path), entity_root=str(root), source="manual", db=db_session, commit=True,
    )
    assert synced.synced is True
    original_metadata = copy.deepcopy((await _doc_for(db_session, entity_id, path.name)).metadata_)
    before_events = set((await db_session.scalars(select(EventLog.id).where(EventLog.entity_id == entity_id))).all())
    read = json.loads(await file_tools._read_file(entity_id, path=path.name, include_structure=True))
    sheet = read["structure"]["sheets"][0]
    assert sheet["name"] == "Sales"
    table = sheet["tables"][0]
    assert table["ref"] == "C2:F3" and table["name"] == "Orders"
    operations = [
        {"op": "row.append", "sheet": sheet["name"], "table": table["name"], "values": ["After", 5, 2]},
        {"op": "row.append", "sheet": sheet["name"], "table": table["name"], "values": ["Final", 2, 8]},
    ]
    if outcome == "late_failure":
        operations[-1]["table"] = "Missing"
    elif outcome == "projection_failure":
        def reject_response_metadata(*_args):
            raise OSError("injected response metadata failure")

        monkeypatch.setattr(file_tools, "_file_meta", reject_response_metadata)

    result = json.loads(await file_tools._patch_file(
        entity_id, path=path.name, operations=operations, expected_sha256=read["source_sha256"],
    ))
    document = await _doc_for(db_session, entity_id, path.name)
    assert document.id == synced.document_id and document.file_type == extension
    if outcome == "success":
        assert result.get("patched") is True, result
        assert result["document_id"] == synced.document_id
        assert result["operations_applied"] == 2
        edited = load_workbook(path, rich_text=True, keep_vba=extension == "xlsm")
        try:
            worksheet = edited["Sales"]
            assert worksheet.tables["Orders"].ref == worksheet.tables["Orders"].autoFilter.ref == "C2:F5"
            assert worksheet["C4"].value == "After" and worksheet["C5"].value == "Final"
            assert worksheet["F4"].value == "=D4*E4+$A$1" and worksheet["F5"].value == "=D5*E5+$A$1"
            assert worksheet["F5"]._style == worksheet["F3"]._style
            assert worksheet["J20"].value == "Unrelated note below the table"
            if edited.vba_archive is not None:
                assert edited.vba_archive.read("xl/vbaProject.bin") == b"opaque-macro-payload-must-be-preserved"
        finally:
            edited.close()
            if edited.vba_archive is not None:
                edited.vba_archive.close()
        reread = json.loads(await file_tools._read_file(entity_id, path=path.name, include_structure=True))
        assert "After" in reread["content"] and "Final" in reread["content"]
        assert reread["source_sha256"] == result["source_sha256"] != read["source_sha256"]
        assert reread["structure"]["sheets"][0]["tables"][0]["ref"] == "C2:F5"
        assert document.vector_status == VectorStatus.PENDING
    else:
        assert result.get("error"), result
        if outcome == "late_failure":
            assert result["operation_index"] == 1
        assert path.read_bytes() == original
        assert document.metadata_ == original_metadata
        events = set((await db_session.scalars(select(EventLog.id).where(EventLog.entity_id == entity_id))).all())
        assert events == before_events


@pytest.mark.asyncio
async def test_runtime_projection_delivers_document_event_after_commit(
    db_session,
    fs_enabled,
    monkeypatch,
):
    from packages.core.services import event_emitter
    from packages.core.services.entity_fs import provision_entity_filesystem

    entity_id = generate_ulid()
    provision_entity_filesystem(entity_id)
    deliveries: list[tuple[str, dict]] = []
    delivery_complete = asyncio.Event()

    async def capture_webhook(_entity_id, event_type, payload):
        assert _entity_id == entity_id
        deliveries.append((f"webhook:{event_type}", dict(payload)))

    async def capture_external(_entity_id, event_type, payload, **_delivery_context):
        assert _entity_id == entity_id
        deliveries.append((f"external:{event_type}", dict(payload)))
        delivery_complete.set()

    monkeypatch.setattr(event_emitter, "deliver_webhook_event", capture_webhook)
    monkeypatch.setattr(
        event_emitter,
        "deliver_task_external_event",
        capture_external,
    )

    result = json.loads(
        await _generate_file(
            entity_id,
            kind="document",
            name="event-report.md",
            content="# Report\n",
        )
    )

    assert result["created"] is True
    await asyncio.wait_for(delivery_complete.wait(), timeout=1)
    events = list((await db_session.scalars(
        select(EventLog).where(
            EventLog.entity_id == entity_id,
            EventLog.event_type == "document.uploaded",
        ),
    )).all())
    assert len(events) == 1
    assert events[0].payload["document_id"] == result["document"]["document_id"]
    assert deliveries == [
        (
            "webhook:document.uploaded",
            {"document_id": result["document"]["document_id"], "name": "event-report.md"},
        ),
        (
            "external:document.uploaded",
            {"document_id": result["document"]["document_id"], "name": "event-report.md"},
        ),
    ]


@pytest.mark.asyncio
async def test_filesystem_projection_queues_reembed_only_after_commit(
    db_session,
    fs_enabled,
    monkeypatch,
):
    from packages.core.services import knowledge_sync
    from packages.core.services.entity_fs import provision_entity_filesystem, resolve_path

    entity_id = generate_ulid()
    provision_entity_filesystem(entity_id)
    abs_path = resolve_path(entity_id, "commit-ordered.md")
    Path(abs_path).write_text("committed bytes\n", encoding="utf-8")
    order: list[str] = []

    def record_commit(_session):
        order.append("commit")

    event.listen(db_session.sync_session, "after_commit", record_commit)
    monkeypatch.setattr(
        knowledge_sync,
        "_schedule_document_reembed",
        lambda _document_id: order.append("schedule"),
    )
    try:
        result = await knowledge_sync.sync_file_to_knowledge(
            entity_id=entity_id,
            abs_path=abs_path,
            entity_root=str(Path(abs_path).parent),
            source="manual",
            db=db_session,
            commit=True,
        )
    finally:
        event.remove(db_session.sync_session, "after_commit", record_commit)

    assert result.synced is True
    assert order[-1] == "schedule"
    assert order[:-1]
    assert set(order[:-1]) == {"commit"}


@pytest.mark.asyncio
async def test_filesystem_projection_does_not_queue_reembed_when_commit_fails(
    db_session,
    fs_enabled,
    monkeypatch,
):
    from packages.core.services import knowledge_sync
    from packages.core.services.entity_fs import provision_entity_filesystem, resolve_path

    entity_id = generate_ulid()
    provision_entity_filesystem(entity_id)
    abs_path = resolve_path(entity_id, "commit-failed.md")
    Path(abs_path).write_text("uncommitted bytes\n", encoding="utf-8")
    scheduled: list[str] = []

    def reject_commit(_session):
        raise RuntimeError("commit failed")

    event.listen(db_session.sync_session, "before_commit", reject_commit)
    monkeypatch.setattr(
        knowledge_sync,
        "_schedule_document_reembed",
        lambda document_id: scheduled.append(document_id),
    )
    try:
        with pytest.raises(RuntimeError, match="commit failed"):
            await knowledge_sync.sync_file_to_knowledge(
                entity_id=entity_id,
                abs_path=abs_path,
                entity_root=str(Path(abs_path).parent),
                source="manual",
                db=db_session,
                commit=True,
            )
    finally:
        event.remove(db_session.sync_session, "before_commit", reject_commit)
        await db_session.rollback()

    assert scheduled == []


@pytest.mark.asyncio
async def test_edit_invalidates_stale_embedding_and_content_text(db_session, fs_enabled):
    entity_id, workspace_id = generate_ulid(), generate_ulid()
    await _seed_workspace(db_session, entity_id, workspace_id)

    w = json.loads(
        await _generate_file(
            entity_id,
            kind="document",
            name="fundraising_ops.md",
            content="| Investor | URL |\n| Garry Tan | https://www.ycombinator.com/people/garry tan |\n",
            save_to_knowledge=True,
            workspace_id=workspace_id,
        )
    )
    assert w["created"] is True
    await db_session.commit()

    doc = await _doc_for(db_session, entity_id, "fundraising_ops.md")
    assert doc is not None
    # simulate the indexer having embedded it + a legacy content_text fork
    doc.vector_status = VectorStatus.READY
    doc.metadata_ = {**(doc.metadata_ or {}), "content_text": "STALE bad url garry tan"}
    await db_session.commit()

    e = json.loads(
        await file_tools._patch_file(
            entity_id,
            path=w["document"]["fs_path"],
            workspace_id=workspace_id,
            operations=[
                {
                    "op": "text.replace",
                    "old_text": "https://www.ycombinator.com/people/garry tan",
                    "new_text": "https://www.ycombinator.com/people/garry-tan",
                }
            ],
        )
    )
    assert e.get("edited") or e.get("replacements") or "error" not in e, e
    await db_session.commit()

    doc = await _doc_for(db_session, entity_id, "fundraising_ops.md")
    # derived representations invalidated so RAG re-derives from fresh bytes
    assert doc.vector_status == VectorStatus.PENDING
    assert "content_text" not in (doc.metadata_ or {})


@pytest.mark.asyncio
async def test_agent_office_overwrite_invalidates_current_preview_marker(
    db_session,
    fs_enabled,
):
    from packages.core.services.slide_renderer import publish_current_preview_version

    entity_id = generate_ulid()
    from packages.core.services.entity_fs import provision_entity_filesystem

    entity_root = file_tools._get_entity_root(entity_id)
    provision_entity_filesystem(entity_id)
    result = json.loads(
        await _generate_file(
            entity_id,
            kind="document",
            name="deck.pptx",
            content="# First version\n",
        )
    )
    assert result["created"] is True
    await db_session.commit()

    doc = await _doc_for(db_session, entity_id, "deck.pptx")
    assert doc is not None and doc.fs_path
    source_path = Path(entity_root, doc.fs_path)
    cache_dir = Path(entity_root, ".slide-cache", doc.id)
    version_dir = cache_dir / "0123456789abcdef"
    version_dir.mkdir(parents=True)
    (version_dir / "slide-1.png").write_bytes(b"first-preview")
    (version_dir / ".complete").write_text("1", encoding="utf-8")
    publish_current_preview_version(
        str(cache_dir),
        version_dir.name,
        str(source_path),
    )
    assert (cache_dir / ".current").is_file()

    overwritten = json.loads(
        await _generate_file(
            entity_id,
            kind="document",
            name=doc.fs_path,
            content="# Second version with changed bytes\n",
            expected_sha256=result["document"]["source_sha256"],
        )
    )
    assert overwritten["created"] is True
    await db_session.commit()

    assert not (cache_dir / ".current").exists()


@pytest.mark.asyncio
async def test_write_scoped_then_read_and_edit_by_logical_name(db_session, fs_enabled):
    """A workspace write reroutes a new file under the artifact folder; reading
    and editing it by its bare logical name must still find it."""
    entity_id, workspace_id = generate_ulid(), generate_ulid()
    await _seed_workspace(db_session, entity_id, workspace_id)

    w = json.loads(
        await _generate_file(
            entity_id,
            kind="document",
            name="fundraising_ops.md",
            content="row one\n",
            save_to_knowledge=True,
            workspace_id=workspace_id,
        )
    )
    await db_session.commit()
    # write rerouted the new file away from the bare name
    assert w["document"]["fs_path"] != "fundraising_ops.md"

    r = json.loads(await file_tools._read_file(
        entity_id, path="fundraising_ops.md", workspace_id=workspace_id,
    ))
    assert "error" not in r, r
    assert "row one" in r.get("content", "")

    e = json.loads(
        await file_tools._patch_file(
            entity_id,
            path="fundraising_ops.md",
            workspace_id=workspace_id,
            operations=[{"op": "text.replace", "old_text": "row one", "new_text": "row edited"}],
        )
    )
    assert "error" not in e, e

    r2 = json.loads(await file_tools._read_file(
        entity_id, path="fundraising_ops.md", workspace_id=workspace_id,
    ))
    assert "row edited" in r2.get("content", "")


@pytest.mark.asyncio
async def test_workspace_logical_name_prefers_scoped_file_over_entity_collision(
    db_session,
    fs_enabled,
):
    """Workspace reads, edits, and deletes must target the path used by writes."""
    entity_id, workspace_id = generate_ulid(), generate_ulid()
    await _seed_workspace(db_session, entity_id, workspace_id)

    root_write = json.loads(
        await _generate_file(
            entity_id,
            kind="document",
            name="shared_name.md",
            content="entity root\n",
        )
    )
    workspace_write = json.loads(
        await _generate_file(
            entity_id,
            kind="document",
            name="shared_name.md",
            content="workspace copy\n",
            workspace_id=workspace_id,
        )
    )
    assert root_write["document"]["fs_path"] == "shared_name.md"
    assert workspace_write["document"]["fs_path"] != root_write["document"]["fs_path"]

    workspace_read = json.loads(await file_tools._read_file(
        entity_id,
        path="shared_name.md",
        workspace_id=workspace_id,
    ))
    assert workspace_read.get("content") == "workspace copy\n"

    workspace_edit = json.loads(
        await file_tools._patch_file(
            entity_id,
            path="shared_name.md",
            workspace_id=workspace_id,
            operations=[{"op": "text.replace", "old_text": "workspace copy", "new_text": "workspace edited"}],
        )
    )
    assert "error" not in workspace_edit, workspace_edit

    workspace_delete = json.loads(await file_tools._delete_file(
        entity_id,
        path="shared_name.md",
        workspace_id=workspace_id,
    ))
    assert workspace_delete.get("deleted") is True

    entity_read = json.loads(await file_tools._read_file(
        entity_id,
        path="shared_name.md",
    ))
    assert entity_read.get("content") == "entity root\n"


@pytest.mark.asyncio
async def test_not_found_lists_same_basename_candidates(db_session, fs_enabled):
    entity_id, workspace_id = generate_ulid(), generate_ulid()
    await _seed_workspace(db_session, entity_id, workspace_id)
    # a file exists at a nested path but the model guesses the bare name
    await _generate_file(
        entity_id,
        kind="document",
        name="dev/reports/tracker.md",
        content="x\n",
        workspace_id=None,
    )

    r = json.loads(await file_tools._read_file(entity_id, path="tracker.md"))
    assert r["error"].startswith("File not found")
    assert "dev/reports/tracker.md" in r.get("candidates", [])


def test_same_basename_candidates_hide_runtime_internal_paths(tmp_path):
    hidden = tmp_path / "uploads"
    visible = tmp_path / "reports"
    hidden.mkdir()
    visible.mkdir()
    (hidden / "tracker.md").write_text("internal", encoding="utf-8")
    (visible / "tracker.md").write_text("visible", encoding="utf-8")

    assert file_tools._same_basename_candidates(
        str(tmp_path),
        "tracker.md",
    ) == ["reports/tracker.md"]


@pytest.mark.asyncio
async def test_unchanged_resync_does_not_reset_vector_status(db_session, fs_enabled, monkeypatch):
    """A reconcile/resync of a byte-identical file must NOT churn embeddings."""
    from packages.core.services import knowledge_sync

    entity_id = generate_ulid()
    from packages.core.services.entity_fs import provision_entity_filesystem, resolve_path
    provision_entity_filesystem(entity_id)

    await _generate_file(
        entity_id,
        kind="document",
        name="notes.md",
        content="hello\n",
    )
    await db_session.commit()
    doc = await _doc_for(db_session, entity_id, "notes.md")
    doc.vector_status = VectorStatus.READY
    await db_session.commit()

    # re-sync the SAME bytes (no write) → must stay ready
    abs_path = resolve_path(entity_id, "notes.md")
    root = file_tools._get_entity_root(entity_id)
    await knowledge_sync.sync_file_to_knowledge(
        entity_id=entity_id, abs_path=abs_path, entity_root=root,
        source="filesystem_reconcile",
    )
    await db_session.commit()
    doc = await _doc_for(db_session, entity_id, "notes.md")
    assert doc.vector_status == VectorStatus.READY
