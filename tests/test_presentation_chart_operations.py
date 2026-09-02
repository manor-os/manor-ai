"""Native PPT chart generation and patching share one editable package path."""

from __future__ import annotations

import io
import re
from pathlib import Path
from zipfile import ZipFile

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Pt

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools.file_tools import _apply_office_patch_sequence_sync
from packages.core.contracts.file_engine import (
    OfficeChartType,
    file_patch_operations,
    normalize_file_patch_operation,
)
from packages.core.services.file_engine_patches import describe_file_structure
from tests.test_office_template_generation import template_bytes


def normalized(*operations):
    return [normalize_file_patch_operation(operation) for operation in operations]


def chart_insert(*, chart_type="column", chart_format=None):
    operation = {
        "op": "chart.insert",
        "slide": 1,
        "chart_type": chart_type,
        "categories": ["Jan", "Feb", "Mar"],
        "series": [
            {"name": "Revenue", "values": [12, 15, 19]},
            {"name": "Cost", "values": [8, 9, 11]},
        ],
        "transform": {"x": 72.25, "y": 96.5, "width": 620, "height": 310, "rotation": 2},
    }
    if chart_type in {"pie", "doughnut"}:
        operation["series"] = operation["series"][:1]
    if chart_type.startswith("scatter"):
        operation["categories"] = [1, 2.5, 4]
    if chart_type in {"stock_hlc", "stock_vhlc"}:
        operation["series"] = [
            {"name": "High", "values": [19, 22, 25]},
            {"name": "Low", "values": [10, 12, 14]},
            {"name": "Close", "values": [15, 18, 21]},
        ]
    if chart_type in {"stock_ohlc", "stock_vohlc"}:
        operation["series"] = [
            {"name": "Open", "values": [14, 17, 20]},
            {"name": "High", "values": [19, 22, 25]},
            {"name": "Low", "values": [10, 12, 14]},
            {"name": "Close", "values": [15, 18, 21]},
        ]
    if chart_type in {"stock_vhlc", "stock_vohlc"}:
        operation["series"].insert(0, {"name": "Volume", "values": [1200, 1800, 1500]})
    if chart_format is not None:
        operation["format"] = chart_format
    return operation


def generate(*operations, template=None):
    return _generate_office_operations_sync(
        "pptx",
        normalized(*operations),
        template_bytes=template,
    )


def patch(path: Path, *operations):
    return _apply_office_patch_sequence_sync(str(path), normalized(*operations))


def chart_style():
    return {
        "title": "Quarterly performance",
        "style": 10,
        "has_legend": True,
        "legend_position": "bottom",
        "legend_include_in_layout": False,
        "vary_colors": False,
        "show_data_labels": True,
        "show_value": True,
        "data_label_position": "outside_end",
        "category_axis_title": "Month",
        "value_axis_title": "USD millions",
        "value_axis_min": 0,
        "value_axis_max": 30,
        "value_axis_major_unit": 5,
        "show_major_gridlines": True,
        "series_colors": ["336699", "CC6600"],
    }


def test_chart_operations_and_types_are_shared_public_capabilities():
    assert {"chart.insert", "chart.data", "chart.format"} <= set(file_patch_operations("pptx"))
    assert OfficeChartType.values() == [
        "column", "column_stacked", "column_stacked_100",
        "bar", "bar_stacked", "bar_stacked_100",
        "line", "line_markers", "area", "area_stacked", "area_stacked_100",
        "pie", "doughnut",
        "scatter", "scatter_lines", "scatter_lines_markers",
        "scatter_smooth", "scatter_smooth_markers",
        "combo_column_line",
        "stock_hlc", "stock_ohlc", "stock_vhlc", "stock_vohlc",
    ]


def test_generate_creates_native_editable_chart_with_data_and_professional_format(tmp_path):
    result = generate(
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        chart_insert(chart_format=chart_style()),
    )
    assert not result.get("error"), result
    path = tmp_path / "chart.pptx"
    path.write_bytes(result["_persisted_bytes"])
    shape = Presentation(path).slides[0].shapes[0]
    chart = shape.chart
    assert shape.shape_id == 2 and shape.has_chart
    assert (shape.left.pt, shape.top.pt, shape.width.pt, shape.height.pt, shape.rotation) == (
        72.25, 96.5, 620, 310, 2,
    )
    assert chart.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED
    assert [series.name for series in chart.series] == ["Revenue", "Cost"]
    assert [list(series.values) for series in chart.series] == [[12, 15, 19], [8, 9, 11]]
    assert chart.chart_title.text_frame.text == "Quarterly performance"
    assert chart.chart_style == 10
    assert chart.legend.position == XL_LEGEND_POSITION.BOTTOM
    assert chart.legend.include_in_layout is False
    assert chart.category_axis.axis_title.text_frame.text == "Month"
    assert chart.value_axis.axis_title.text_frame.text == "USD millions"
    assert (chart.value_axis.minimum_scale, chart.value_axis.maximum_scale, chart.value_axis.major_unit) == (0, 30, 5)
    assert chart.plots[0].data_labels.show_value is True
    assert [str(series.format.fill.fore_color.rgb) for series in chart.series] == ["336699", "CC6600"]
    structure = describe_file_structure(str(path))["shapes"][0]["chart"]
    assert structure["type"] == "column"
    assert structure["categories"] == ["Jan", "Feb", "Mar"]
    assert structure["series"][1] == {"name": "Cost", "values": [8, 9, 11]}
    assert structure["format"]["title"] == "Quarterly performance"
    assert structure["format"]["series_colors"] == ["336699", "CC6600"]


@pytest.mark.parametrize("chart_type", OfficeChartType.values())
def test_every_exposed_chart_type_creates_a_native_editable_chart(chart_type):
    result = generate(
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        chart_insert(chart_type=chart_type),
    )
    assert not result.get("error"), result
    if chart_type.startswith("stock_"):
        with ZipFile(io.BytesIO(result["_persisted_bytes"])) as package:
            chart_xml = package.read("ppt/charts/chart1.xml").decode()
        assert "<c:stockChart>" in chart_xml
        assert chart_xml.count("<c:ser>") == {
            "stock_hlc": 3, "stock_ohlc": 4, "stock_vhlc": 4, "stock_vohlc": 5,
        }[chart_type]
        assert ("<c:barChart>" in chart_xml) is chart_type.startswith("stock_v")
    else:
        shape = Presentation(io.BytesIO(result["_persisted_bytes"])).slides[0].shapes[0]
        assert shape.has_chart
        assert len(shape.chart.series) == (1 if chart_type in {"pie", "doughnut"} else 2)


@pytest.mark.parametrize("chart_type,series_names", [
    ("stock_hlc", ["High", "Low", "Close"]),
    ("stock_ohlc", ["Open", "High", "Low", "Close"]),
    ("stock_vhlc", ["Volume", "High", "Low", "Close"]),
    ("stock_vohlc", ["Volume", "Open", "High", "Low", "Close"]),
])
def test_stock_chart_generation_discovery_and_patch_remain_native(tmp_path, chart_type, series_names):
    created = generate(
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        chart_insert(chart_type=chart_type),
    )
    assert not created.get("error"), created
    path = tmp_path / f"{chart_type}.pptx"
    path.write_bytes(created["_persisted_bytes"])
    details = describe_file_structure(str(path))["shapes"][0]["chart"]
    assert details["type"] == chart_type
    assert [series["name"] for series in details["series"]] == series_names
    rejected = patch(path, {
        "op": "chart.format", "slide": 1, "shape_id": 2, "format": {"vary_colors": True},
    })
    assert rejected.get("error") and path.read_bytes() == created["_persisted_bytes"]

    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement

    document = Presentation(path)
    stock_plot = document.slides[0].shapes[0].chart._chartSpace.chart.plotArea.find(qn("c:stockChart"))
    high_low_lines = stock_plot.find(qn("c:hiLowLines"))
    properties = OxmlElement("c:spPr")
    line = OxmlElement("a:ln")
    line.set("w", "42000")
    properties.append(line)
    high_low_lines.append(properties)
    if chart_type in {"stock_ohlc", "stock_vohlc"}:
        stock_plot.find(qn("c:upDownBars")).find(qn("c:gapWidth")).set("val", "77")
    document.save(path)

    updated_series = [
        {"name": name, "values": [20 + index, 24 + index]}
        for index, name in enumerate(series_names)
    ]
    stock_format = {"title": "Updated stock", "series_colors": ["336699"] * len(series_names)}
    if chart_type.startswith("stock_v"):
        stock_format.update({
            "value_axis_title": "Volume",
            "secondary_value_axis_title": "Price",
            "secondary_value_axis_min": 10,
            "secondary_value_axis_max": 40,
            "secondary_value_axis_major_unit": 5,
            "show_secondary_major_gridlines": False,
        })
    edited = patch(path, {
        "op": "chart.data", "slide": 1, "shape_id": 2,
        "categories": ["Apr", "May"], "series": updated_series,
    }, {
        "op": "chart.format", "slide": 1, "shape_id": 2,
        "format": stock_format,
    })
    assert not edited.get("error"), edited
    path.write_bytes(edited["_persisted_bytes"])
    details = describe_file_structure(str(path))["shapes"][0]["chart"]
    assert details["type"] == chart_type
    assert details["categories"] == ["Apr", "May"]
    assert details["series"][-1]["values"] == [20 + len(series_names) - 1, 24 + len(series_names) - 1]
    assert details["format"]["title"] == "Updated stock"
    if chart_type.startswith("stock_v"):
        assert details["format"]["value_axis_title"] == "Volume"
        assert details["format"]["secondary_value_axis_title"] == "Price"
        assert details["format"]["secondary_value_axis_min"] == 10
        assert details["format"]["secondary_value_axis_max"] == 40
        assert details["format"]["secondary_value_axis_major_unit"] == 5
        assert details["format"]["show_secondary_major_gridlines"] is False
    with ZipFile(io.BytesIO(edited["_persisted_bytes"])) as package:
        chart_xml = package.read("ppt/charts/chart1.xml").decode()
    assert "<c:stockChart>" in chart_xml and "<c:hiLowLines" in chart_xml
    assert 'w="42000"' in chart_xml
    assert ("<c:upDownBars" in chart_xml) is (chart_type in {"stock_ohlc", "stock_vohlc"})
    assert ("<c:barChart>" in chart_xml) is chart_type.startswith("stock_v")
    assert len(set(re.findall(r'<c:valAx><c:axId val="(-?\d+)"', chart_xml))) == (
        2 if chart_type.startswith("stock_v") else 1
    )
    if chart_type in {"stock_ohlc", "stock_vohlc"}:
        assert '<c:gapWidth val="77"' in chart_xml


def test_imported_volume_stock_uses_plot_and_axis_ids_instead_of_xml_order(tmp_path):
    created = generate(
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        chart_insert(
            chart_type="stock_vhlc",
            chart_format={"value_axis_title": "Volume", "secondary_value_axis_title": "Price"},
        ),
    )
    assert not created.get("error"), created
    path = tmp_path / "reordered-volume-stock.pptx"
    document = Presentation(io.BytesIO(created["_persisted_bytes"]))
    chart_shape = document.slides[0].shapes[0]

    from pptx.oxml.ns import qn

    plot_area = chart_shape.chart._chartSpace.chart.plotArea
    bar_plot = plot_area.find(qn("c:barChart"))
    stock_plot = plot_area.find(qn("c:stockChart"))
    plot_area.remove(stock_plot)
    bar_plot.addprevious(stock_plot)
    value_axes = plot_area.findall(qn("c:valAx"))
    plot_area.remove(value_axes[1])
    value_axes[0].addprevious(value_axes[1])
    shape_id = chart_shape.shape_id
    document.save(path)

    imported = describe_file_structure(str(path))["shapes"][0]["chart"]
    assert imported["type"] == "stock_vhlc"
    assert [series["name"] for series in imported["series"]] == ["Volume", "High", "Low", "Close"]
    assert imported["format"]["value_axis_title"] == "Volume"
    assert imported["format"]["secondary_value_axis_title"] == "Price"

    edited = patch(path, {
        "op": "chart.data", "slide": 1, "shape_id": shape_id,
        "categories": ["Apr", "May"],
        "series": [
            {"name": "Volume", "values": [2000, 2200]},
            {"name": "High", "values": [28, 30]},
            {"name": "Low", "values": [18, 19]},
            {"name": "Close", "values": [24, 27]},
        ],
    }, {
        "op": "chart.format", "slide": 1, "shape_id": shape_id,
        "format": {"value_axis_title": "Updated volume", "secondary_value_axis_title": "Updated price"},
    })
    assert not edited.get("error"), edited
    path.write_bytes(edited["_persisted_bytes"])

    details = describe_file_structure(str(path))["shapes"][0]["chart"]
    assert details["type"] == "stock_vhlc"
    assert [series["name"] for series in details["series"]] == ["Volume", "High", "Low", "Close"]
    assert details["series"][0]["values"] == [2000, 2200]
    assert details["series"][-1]["values"] == [24, 27]
    assert details["format"]["value_axis_title"] == "Updated volume"
    assert details["format"]["secondary_value_axis_title"] == "Updated price"
    with ZipFile(io.BytesIO(edited["_persisted_bytes"])) as package:
        chart_xml = package.read("ppt/charts/chart1.xml").decode()
    assert chart_xml.count("<c:valAx>") == 2


def test_unsupported_multi_plot_chart_remains_inspectable(tmp_path):
    created = generate(
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        chart_insert(chart_type="stock_hlc"),
    )
    assert not created.get("error"), created
    path = tmp_path / "unsupported-stock-combo.pptx"
    document = Presentation(io.BytesIO(created["_persisted_bytes"]))

    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement

    plot_area = document.slides[0].shapes[0].chart._chartSpace.chart.plotArea
    stock_plot = plot_area.find(qn("c:stockChart"))
    line_plot = OxmlElement("c:lineChart")
    grouping = OxmlElement("c:grouping")
    grouping.set("val", "standard")
    line_plot.append(grouping)
    stock_plot.addnext(line_plot)
    document.save(path)

    chart = describe_file_structure(str(path))["shapes"][0]["chart"]
    assert chart["type"] == "unsupported"
    assert chart["series"] == []


def test_combo_chart_generation_discovery_and_patch_keep_last_series_as_line(tmp_path):
    result = generate(
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        chart_insert(chart_type="combo_column_line"),
    )
    assert not result.get("error"), result
    path = tmp_path / "combo-chart.pptx"
    path.write_bytes(result["_persisted_bytes"])
    chart = Presentation(path).slides[0].shapes[0].chart
    assert [type(plot).__name__ for plot in chart.plots] == ["BarPlot", "LinePlot"]
    assert [len(plot.series) for plot in chart.plots] == [1, 1]
    assert describe_file_structure(str(path))["shapes"][0]["chart"]["type"] == "combo_column_line"

    edited = patch(path, {
        "op": "chart.data", "slide": 1, "shape_id": 2,
        "categories": ["Q1", "Q2"],
        "series": [
            {"name": "Revenue", "values": [20, 25]},
            {"name": "Cost", "values": [12, 14]},
            {"name": "Margin", "values": [40, 44]},
        ],
    }, {
        "op": "chart.format", "slide": 1, "shape_id": 2,
        "format": {
            "series_colors": ["336699", "CC6600", "228844"],
            "show_data_labels": True,
            "show_value": True,
        },
    })
    assert not edited.get("error"), edited
    chart = Presentation(io.BytesIO(edited["_persisted_bytes"])).slides[0].shapes[0].chart
    assert [len(plot.series) for plot in chart.plots] == [2, 1]
    assert [series.name for series in chart.series] == ["Revenue", "Cost", "Margin"]
    assert [str(series.format.fill.fore_color.rgb) for series in chart.series] == [
        "336699", "CC6600", "228844",
    ]
    assert all(plot.has_data_labels and plot.data_labels.show_value for plot in chart.plots)


@pytest.mark.parametrize("chart_type", [
    "scatter", "scatter_lines", "scatter_lines_markers", "scatter_smooth", "scatter_smooth_markers",
])
def test_scatter_chart_generation_discovery_and_data_patch_share_numeric_x_values(tmp_path, chart_type):
    created = generate(
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        chart_insert(chart_type=chart_type),
    )
    path = tmp_path / f"{chart_type}.pptx"
    path.write_bytes(created["_persisted_bytes"])
    structure = describe_file_structure(str(path))["shapes"][0]["chart"]
    assert structure["type"] == chart_type
    assert structure["categories"] == [1, 2.5, 4]
    assert structure["series"][1]["x_values"] == [1, 2.5, 4]

    patched = patch(path, {
        "op": "chart.data", "slide": 1, "shape_id": 2,
        "categories": [0.5, 3],
        "series": [
            {"name": "Observed", "values": [2, 8]},
            {"name": "Forecast", "values": [3, 9]},
        ],
    })
    assert not patched.get("error"), patched
    edited = Presentation(io.BytesIO(patched["_persisted_bytes"])).slides[0].shapes[0].chart
    expected_types = {
        "scatter": XL_CHART_TYPE.XY_SCATTER,
        "scatter_lines": XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS,
        "scatter_lines_markers": XL_CHART_TYPE.XY_SCATTER_LINES,
        "scatter_smooth": XL_CHART_TYPE.XY_SCATTER_SMOOTH_NO_MARKERS,
        "scatter_smooth_markers": XL_CHART_TYPE.XY_SCATTER_SMOOTH,
    }
    assert edited.chart_type == expected_types[chart_type]
    assert [list(series.values) for series in edited.series] == [[2, 8], [3, 9]]


def test_scatter_chart_keeps_native_axes_labels_and_series_format(tmp_path):
    style = {
        **chart_style(),
        "category_axis_min": 0,
        "category_axis_max": 5,
        "category_axis_major_unit": 1,
        "trendlines": [{
            "series_index": 0, "type": "linear", "name": "Observed trend",
            "forward": 1, "display_equation": True, "display_r_squared": True,
        }],
        "error_bars": [{
            "series_index": 1, "type": "percentage", "value": 10,
            "side": "plus", "end_style": "no_cap",
        }],
    }
    created = generate(
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        chart_insert(chart_type="scatter_lines_markers", chart_format=style),
    )
    assert not created.get("error"), created
    path = tmp_path / "formatted-scatter.pptx"
    path.write_bytes(created["_persisted_bytes"])
    structure = describe_file_structure(str(path))["shapes"][0]["chart"]
    assert structure["format"]["show_data_labels"] is True
    assert structure["format"]["show_value"] is True
    assert structure["format"]["category_axis_title"] == "Month"
    assert structure["format"]["category_axis_min"] == 0
    assert structure["format"]["category_axis_max"] == 5
    assert structure["format"]["category_axis_major_unit"] == 1
    assert structure["format"]["value_axis_title"] == "USD millions"
    assert structure["format"]["series_colors"] == ["336699", "CC6600"]
    assert structure["series"][0]["trendline"] == {
        "type": "linear", "name": "Observed trend", "forward": 1.0,
        "display_equation": True, "display_r_squared": True,
    }
    assert structure["series"][1]["error_bars"] == {
        "type": "percentage", "direction": "y", "side": "plus", "end_style": "no_cap", "value": 10.0,
    }

    changed = patch(path, {
        "op": "chart.format", "slide": 1, "shape_id": 2,
        "format": {
            "trendlines": [{"series_index": 0, "type": "polynomial", "order": 3}],
            "error_bars": [{"series_index": 1, "type": "standard_error", "direction": "x"}],
        },
    })
    assert not changed.get("error"), changed
    path.write_bytes(changed["_persisted_bytes"])
    assert describe_file_structure(str(path))["shapes"][0]["chart"]["series"][0]["trendline"] == {
        "type": "polynomial", "order": 3,
    }
    assert describe_file_structure(str(path))["shapes"][0]["chart"]["series"][1]["error_bars"] == {
        "type": "standard_error", "direction": "x", "side": "both", "end_style": "cap",
    }


def test_chart_data_retains_geometry_type_and_format_while_replacing_embedded_workbook(tmp_path):
    initial = generate(
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        chart_insert(chart_format=chart_style()),
    )
    path = tmp_path / "data.pptx"
    path.write_bytes(initial["_persisted_bytes"])
    before = Presentation(path).slides[0].shapes[0]
    result = patch(path, {
        "op": "chart.data",
        "slide": 1,
        "shape_id": 2,
        "categories": ["Apr", "May"],
        "series": [
            {"name": "Revenue", "values": [21, 24]},
            {"name": "Cost", "values": [12, 13]},
        ],
    })
    assert not result.get("error"), result
    assert path.read_bytes() == initial["_persisted_bytes"]
    after = Presentation(io.BytesIO(result["_persisted_bytes"])).slides[0].shapes[0]
    assert (after.left, after.top, after.width, after.height, after.rotation) == (
        before.left, before.top, before.width, before.height, before.rotation,
    )
    assert after.chart.chart_type == before.chart.chart_type
    assert after.chart.chart_title.text_frame.text == "Quarterly performance"
    assert after.chart.chart_style == 10
    assert [str(series.format.fill.fore_color.rgb) for series in after.chart.series] == ["336699", "CC6600"]
    assert [list(series.values) for series in after.chart.series] == [[21, 24], [12, 13]]


def test_pie_chart_format_colors_native_points_and_rejects_axes(tmp_path):
    created = generate(
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        chart_insert(chart_type="pie"),
    )
    path = tmp_path / "pie.pptx"
    path.write_bytes(created["_persisted_bytes"])
    result = patch(path, {
        "op": "chart.format", "slide": 1, "shape_id": 2,
        "format": {
            "title": "Revenue mix", "show_data_labels": True,
            "show_category_name": True, "show_percentage": True,
            "data_label_position": "best_fit",
            "category_colors": ["174C46", "D89B45", "7B61A8"],
        },
    })
    assert not result.get("error"), result
    chart = Presentation(io.BytesIO(result["_persisted_bytes"])).slides[0].shapes[0].chart
    assert chart.chart_title.text_frame.text == "Revenue mix"
    assert chart.plots[0].data_labels.show_category_name is True
    assert chart.plots[0].data_labels.show_percentage is True
    assert [str(point.format.fill.fore_color.rgb) for point in chart.series[0].points] == [
        "174C46", "D89B45", "7B61A8",
    ]
    path.write_bytes(result["_persisted_bytes"])
    rejected = patch(path, {
        "op": "chart.format", "slide": 1, "shape_id": 2,
        "format": {"value_axis_title": "Not supported"},
    })
    assert rejected.get("error") == "this chart type has no value axis"
    assert path.read_bytes() == result["_persisted_bytes"]
    rejected_trendline = patch(path, {
        "op": "chart.format", "slide": 1, "shape_id": 2,
        "format": {"trendlines": [{"series_index": 0, "type": "linear"}]},
    })
    assert "trendlines require" in rejected_trendline.get("error", "")


def test_template_chart_patch_changes_only_target_chart_and_retains_other_parts(tmp_path):
    native = Presentation(io.BytesIO(template_bytes("pptx")))
    data = CategoryChartData()
    data.categories = ["East", "West"]
    data.add_series("Pipeline", [7, 9])
    second = native.slides[0].shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS, Pt(450), Pt(320), Pt(240), Pt(120), data,
    )
    source = io.BytesIO()
    native.save(source)
    before = source.getvalue()
    untouched_part = second.chart.part.partname.lstrip("/")
    untouched_related = {
        relationship.target_part.partname.lstrip("/")
        for relationship in second.chart.part.rels.values()
        if not relationship.is_external
    }
    target = next(shape for shape in native.slides[0].shapes if shape.has_chart and shape.shape_id != second.shape_id)
    path = tmp_path / "template.pptx"
    path.write_bytes(before)
    result = patch(path, {
        "op": "chart.data", "slide": 1, "shape_id": target.shape_id,
        "categories": ["Q1", "Q2"],
        "series": [{"name": "Revenue", "values": [20, 25]}],
    })
    assert not result.get("error"), result
    edited = Presentation(io.BytesIO(result["_persisted_bytes"]))
    assert list(next(shape for shape in edited.slides[0].shapes if shape.shape_id == target.shape_id).chart.series[0].values) == [20, 25]
    assert list(next(shape for shape in edited.slides[0].shapes if shape.shape_id == second.shape_id).chart.series[0].values) == [7, 9]
    with ZipFile(io.BytesIO(before)) as original, ZipFile(io.BytesIO(result["_persisted_bytes"])) as output:
        for name in {untouched_part, *untouched_related}:
            assert output.read(name) == original.read(name), name
        for name in original.namelist():
            if name.startswith(("ppt/slideMasters/", "ppt/notesSlides/")):
                assert output.read(name) == original.read(name), name


@pytest.mark.parametrize("operation", [
    {"op": "chart.insert", "slide": 1, "chart_type": "pie", "categories": ["A"],
     "series": [{"name": "One", "values": [1]}, {"name": "Two", "values": [2]}],
     "transform": {"x": 1, "y": 1, "width": 100, "height": 100}},
    {"op": "chart.insert", "slide": 1, "chart_type": "unknown", "categories": ["A"],
     "series": [{"name": "One", "values": [1]}],
     "transform": {"x": 1, "y": 1, "width": 100, "height": 100}},
    {"op": "chart.insert", "slide": 1, "chart_type": "column", "categories": ["A", "B"],
     "series": [{"name": "One", "values": [1]}],
     "transform": {"x": 1, "y": 1, "width": 100, "height": 100}},
    {"op": "chart.insert", "slide": 1, "chart_type": "scatter", "categories": ["A"],
     "series": [{"name": "One", "values": [1]}],
     "transform": {"x": 1, "y": 1, "width": 100, "height": 100}},
    {"op": "chart.insert", "slide": 1, "chart_type": "combo_column_line", "categories": ["A"],
     "series": [{"name": "Only", "values": [1]}],
     "transform": {"x": 1, "y": 1, "width": 100, "height": 100}},
    {"op": "chart.insert", "slide": 1, "chart_type": "stock_hlc", "categories": ["A"],
     "series": [{"name": "High", "values": [2]}, {"name": "Low", "values": [1]}],
     "transform": {"x": 1, "y": 1, "width": 100, "height": 100}},
    {"op": "chart.insert", "slide": 1, "chart_type": "stock_vhlc", "categories": ["A"],
     "series": [{"name": "Volume", "values": [100]}, {"name": "High", "values": [2]},
                {"name": "Low", "values": [1]}],
     "transform": {"x": 1, "y": 1, "width": 100, "height": 100}},
    {"op": "chart.data", "slide": 1, "shape_id": 2, "categories": ["A"], "series": []},
    {"op": "chart.format", "slide": 1, "shape_id": 2, "format": {"style": 49}},
    {"op": "chart.format", "slide": 1, "shape_id": 2,
     "format": {"value_axis_min": 10, "value_axis_max": 5}},
    {"op": "chart.format", "slide": 1, "shape_id": 2,
     "format": {"secondary_value_axis_title": "Not available"}},
    {"op": "chart.format", "slide": 1, "shape_id": 2,
     "format": {"series_colors": ["#112233"]}},
    {"op": "chart.format", "slide": 1, "shape_id": 2,
     "format": {"trendlines": [{"series_index": 5, "type": "linear"}]}},
    {"op": "chart.format", "slide": 1, "shape_id": 2,
     "format": {"error_bars": [{"series_index": 0, "type": "percentage"}]}},
])
def test_invalid_chart_operation_aborts_the_whole_batch(tmp_path, operation):
    created = generate(
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        chart_insert(),
    )
    path = tmp_path / "atomic.pptx"
    path.write_bytes(created["_persisted_bytes"])
    result = patch(
        path,
        {"op": "chart.format", "slide": 1, "shape_id": 2, "format": {"title": "Must roll back"}},
        operation,
    )
    assert result.get("error"), result
    assert result["operation_index"] == 1
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == created["_persisted_bytes"]
