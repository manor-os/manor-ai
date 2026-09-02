"""Shared native cell formatting survives generation, patching and later edits."""
from __future__ import annotations

import io

import pytest

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools.file_tools import _apply_office_patch_sequence_sync
from packages.core.contracts.file_engine import normalize_file_patch_operation
from tests.test_office_table_patch import _open, cell_operation, table_operations, table_source


CELL_STYLE = {"bold": True, "italic": True, "font_size": 20.5, "font_name": "Arial", "font_color": "112233", "fill_color": "F3D9AA"}


def format_operation(extension, style=None, cell="B2"):
    operation = cell_operation(extension, cell)
    operation.pop("value")
    return {**operation, "op": "cell.format", "format": CELL_STYLE.copy() if style is None else style}


def _patch(path, *operations):
    return _apply_office_patch_sequence_sync(str(path), [normalize_file_patch_operation(op) for op in operations])


def assert_cell_style(cell, extension):
    paragraphs = cell.paragraphs if extension == "docx" else cell.text_frame.paragraphs
    for paragraph in paragraphs:
        for run in paragraph.runs:
            if run.text:
                assert run.font.bold is True and run.font.italic is True
                assert run.font.size.pt == 20.5
                assert run.font.name == "Arial"
                assert str(run.font.color.rgb) == "112233"
    if extension == "docx":
        from docx.oxml.ns import qn

        assert cell._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) == "F3D9AA"
    else:
        assert str(cell.fill.fore_color.rgb) == "F3D9AA"


@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("format_first", [True, False])
def test_generate_format_and_patch_use_identical_operations(tmp_path, extension, format_first):
    cell_edits = [format_operation(extension), cell_operation(extension, value="标题\nSecond line")]
    if not format_first:
        cell_edits.reverse()
    generated = _generate_office_operations_sync(extension, [normalize_file_patch_operation(op) for op in table_operations(extension) + cell_edits])
    assert not generated.get("error"), generated
    _, table = _open(io.BytesIO(generated["_persisted_bytes"]), extension)
    assert_cell_style(table.cell(1, 1), extension)
    assert table.cell(1, 1).text == "标题\nSecond line"
    path = tmp_path / f"formatted.{extension}"
    path.write_bytes(generated["_persisted_bytes"])
    before = path.read_bytes()
    other_cells = [table.cell(row, col)._tc.xml for row, col in ((0, 0), (0, 1), (1, 0))]
    patched = _patch(path, cell_operation(extension, value="Later edit"), format_operation(extension, {"bold": False}))
    assert not patched.get("error"), patched
    _, table = _open(io.BytesIO(patched["_persisted_bytes"]), extension)
    cell = table.cell(1, 1)
    assert cell.text == "Later edit"
    paragraph = cell.paragraphs[0] if extension == "docx" else cell.text_frame.paragraphs[0]
    run = next(run for run in paragraph.runs if run.text)
    assert run.font.bold is False and run.font.italic is True
    assert run.font.size.pt == 20.5 and run.font.name == "Arial"
    assert str(run.font.color.rgb) == "112233"
    assert [table.cell(row, col)._tc.xml for row, col in ((0, 0), (0, 1), (1, 0))] == other_cells
    assert path.read_bytes() == before


@pytest.mark.parametrize("extension", ["docx", "pptx"])
def test_fill_only_preserves_text_runs_and_paragraph_layout(tmp_path, extension):
    path = tmp_path / f"fill.{extension}"
    table_source(path)
    document, table = _open(str(path), extension)
    cell = table.cell(1, 1)
    paragraph = cell.paragraphs[0] if extension == "docx" else cell.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = "Linked"
    run.font.underline = True
    document.save(path)
    original = paragraph._p.xml
    result = _patch(path, format_operation(extension, {"fill_color": "abcdef"}))
    assert not result.get("error"), result
    _, table = _open(io.BytesIO(result["_persisted_bytes"]), extension)
    paragraph = table.cell(1, 1).paragraphs[0] if extension == "docx" else table.cell(1, 1).text_frame.paragraphs[0]
    assert paragraph._p.xml == original


@pytest.mark.parametrize("extension", ["docx", "pptx"])
def test_cell_format_preserves_hyperlinks_and_unrequested_font_properties(tmp_path, extension):
    path = tmp_path / f"links.{extension}"
    table_source(path)
    document, table = _open(str(path), extension)
    cell = table.cell(1, 1)
    paragraph = cell.paragraphs[0] if extension == "docx" else cell.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = "Keep link"
    run.font.underline = True
    if extension == "docx":
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.opc.constants import RELATIONSHIP_TYPE as RT

        link = OxmlElement("w:hyperlink")
        link_id = document.part.relate_to("https://example.com", RT.HYPERLINK, is_external=True)
        link.set(qn("r:id"), link_id)
        paragraph._p.append(link)
        link.append(run._r)
    else:
        run.hyperlink.address = "https://example.com"
    document.save(path)
    result = _patch(path, format_operation(extension))
    assert not result.get("error"), result
    document, table = _open(io.BytesIO(result["_persisted_bytes"]), extension)
    if extension == "docx":
        paragraph = table.cell(1, 1).paragraphs[0]
        hyperlink = paragraph.hyperlinks[0]
        assert hyperlink.address == "https://example.com"
        run = hyperlink.runs[0]
        fonts = run._r.rPr.rFonts
        for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
            assert fonts.get(qn("w:" + slot)) == "Arial"
        assert run.font.cs_bold is True and run.font.cs_italic is True
        assert run._r.rPr.find(qn("w:szCs")).get(qn("w:val")) == "41"
    else:
        run = next(run for run in table.cell(1, 1).text_frame.paragraphs[0].runs if run.text)
        assert run.hyperlink.address == "https://example.com"
        for slot in ("latin", "ea", "cs"):
            assert run._r.rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}" + slot).get("typeface") == "Arial"
    assert run.text == "Keep link" and run.font.underline is True
    assert run.font.bold is True and run.font.italic is True


@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("style", [
    {}, None, [], {"unknown": True}, {"bold": 1}, {"italic": "false"},
    {"font_size": True}, {"font_size": float("inf")}, {"font_size": 0}, {"font_size": 410},
    {"font_size": 12.345}, {"font_name": ""}, {"font_name": "bad\x00font"},
    {"font_color": "#112233"}, {"fill_color": "blue"}, {"wrap_text": True}, {"number_format": "0.00"},
])
def test_invalid_cell_format_rolls_back_entire_sequence(tmp_path, extension, style):
    path = tmp_path / f"invalid.{extension}"
    table_source(path)
    before = path.read_bytes()
    invalid = {**format_operation(extension), "format": style}
    result = _patch(path, cell_operation(extension), invalid)
    assert result.get("error"), result
    assert result["operation_index"] == 1
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before


def test_word_explicit_colors_override_theme_shading_and_font(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    path = tmp_path / "theme.docx"
    table_source(path)
    document, table = _open(str(path), "docx")
    cell = table.cell(1, 1)
    run = cell.paragraphs[0].add_run("Themed")
    run.font.name = "Calibri"
    run._r.rPr.rFonts.set(qn("w:asciiTheme"), "minorHAnsi")
    run._r.rPr.rFonts.set(qn("w:eastAsiaTheme"), "minorEastAsia")
    shading = OxmlElement("w:shd")
    shading.set(qn("w:themeFill"), "accent1")
    shading.set(qn("w:themeFillTint"), "80")
    cell._tc.tcPr.append(shading)
    document.save(path)
    result = _patch(path, format_operation("docx"))
    assert not result.get("error"), result
    _, table = _open(io.BytesIO(result["_persisted_bytes"]), "docx")
    assert_cell_style(table.cell(1, 1), "docx")
    for fonts in table.cell(1, 1)._tc.xpath(".//w:rFonts"):
        assert not any(attribute.lower().endswith("theme") for attribute in fonts.attrib)
    assert "themeFill" not in table.cell(1, 1)._tc.xml


def test_ppt_explicit_rgb_clears_old_color_modifiers_only_when_color_changes(tmp_path):
    from pptx.dml.color import RGBColor
    from pptx.oxml.xmlchemy import OxmlElement

    path = tmp_path / "tinted.pptx"
    table_source(path)
    document, table = _open(str(path), "pptx")
    cell = table.cell(1, 1)
    run = cell.text_frame.paragraphs[0].add_run()
    run.text = "Tinted"
    cell.fill.solid()
    for color in (cell.fill.fore_color, run.font.color):
        color.rgb = RGBColor.from_string("445566")
        color.brightness = 0.5
        alpha = OxmlElement("a:alpha")
        alpha.set("val", "50000")
        color._xFill.eg_colorChoice.append(alpha)
    document.save(path)
    result = _patch(path, format_operation("pptx", {"bold": False}))
    assert not result.get("error"), result
    _, unchanged = _open(io.BytesIO(result["_persisted_bytes"]), "pptx")
    assert unchanged.cell(1, 1).fill.fore_color.brightness == 0.5
    result = _patch(path, format_operation("pptx"))
    assert not result.get("error"), result
    _, table = _open(io.BytesIO(result["_persisted_bytes"]), "pptx")
    cell = table.cell(1, 1)
    assert_cell_style(cell, "pptx")
    run = next(run for run in cell.text_frame.paragraphs[0].runs if run.text)
    for color in (cell.fill.fore_color, run.font.color):
        assert len(color._xFill.eg_colorChoice) == 0


@pytest.mark.parametrize("size", [1, 12.5, 20.03, 20.07, 99.99, 409])
def test_ppt_font_size_preserves_hundredth_point_precision(tmp_path, size):
    path = tmp_path / "size.pptx"
    table_source(path)
    result = _patch(path, format_operation("pptx", {"font_size": size}), cell_operation("pptx"))
    assert not result.get("error"), result
    _, table = _open(io.BytesIO(result["_persisted_bytes"]), "pptx")
    run = next(run for run in table.cell(1, 1).text_frame.paragraphs[0].runs if run.text)
    assert run._r.rPr.sz == round(size * 100)


@pytest.mark.parametrize("extension", ["docx", "pptx"])
def test_formatting_merged_anchor_does_not_change_merge_geometry(tmp_path, extension):
    path = tmp_path / f"merged.{extension}"
    table_source(path)
    document, table = _open(str(path), extension)
    table.cell(0, 0).merge(table.cell(1, 1))
    text = table.cell(0, 0).text
    document.save(path)
    before = path.read_bytes()
    result = _patch(path, format_operation(extension, cell="B2"))
    assert "top-left anchor" in result["error"]
    assert path.read_bytes() == before
    result = _patch(path, format_operation(extension, cell="A1"))
    assert not result.get("error"), result
    _, table = _open(io.BytesIO(result["_persisted_bytes"]), extension)
    assert table.cell(0, 0).text == text
    assert_cell_style(table.cell(0, 0), extension)
    if extension == "docx":
        assert table.cell(0, 0).grid_span == 2
        assert table.cell(0, 0)._tc.vMerge == "restart"
    else:
        assert table.cell(0, 0).span_width == table.cell(0, 0).span_height == 2
