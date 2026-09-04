from __future__ import annotations

import json
from types import SimpleNamespace

import pytest

from packages.core.ai.tools import file_tools


def _tool_names() -> set[str]:
    return {schema["function"]["name"] for schema, _handler in file_tools.get_tools()}


def test_file_engine_tools_are_registered_with_generate_and_edit_tools() -> None:
    names = _tool_names()

    assert {"read_file", "write_file", "edit_file", "inspect_file_engine", "patch_file"} <= names

    from packages.core.ai.tool_pool import ToolPool

    pool = ToolPool()
    pool.initialize()
    registered = set(pool.registered_tool_names())
    assert {"generate_file", "inspect_file_engine", "patch_file"} <= registered
    assert {
        "mcp__manor_mcp_file_engine__inspect",
        "mcp__manor_mcp_file_engine__generate",
        "mcp__manor_mcp_file_engine__patch",
    } <= registered


def test_file_engine_mcp_catalog_matches_module_and_skill_config() -> None:
    from packages.core.ai.mcp import get_module
    from packages.core.ai.tools.mcp_builtin import _SERVER_TOOL_SCHEMAS

    module = get_module("manor_mcp_file_engine")
    assert module is not None
    module_tools = {tool["name"] for tool in module.list_tools()}
    assert module_tools == {"inspect", "generate", "patch"}
    assert {tool["name"] for tool in _SERVER_TOOL_SCHEMAS["manor_mcp_file_engine"]} == module_tools

    from pathlib import Path

    config = json.loads(
        Path("packages/core/ai/skills/file_engine/config.json").read_text(encoding="utf-8")
    )
    assert {
        "mcp__manor_mcp_file_engine__inspect",
        "mcp__manor_mcp_file_engine__generate",
        "mcp__manor_mcp_file_engine__patch",
    } <= set(config["tools"])


@pytest.mark.asyncio
async def test_file_engine_mcp_delegates_calls_through_runtime_boundary(monkeypatch) -> None:
    from packages.core.ai.mcp import manor_mcp_file_engine
    from packages.core.ai.runtime import file_engine as runtime_file_engine

    calls: list[tuple[str, dict[str, object]]] = []

    async def fake_inspect(**kwargs):
        calls.append(("inspect", kwargs))
        return json.dumps({"action": "inspect"})

    async def fake_generate(**kwargs):
        calls.append(("generate", kwargs))
        return json.dumps({"action": "generate"})

    async def fake_patch(**kwargs):
        calls.append(("patch", kwargs))
        return json.dumps({"action": "patch"})

    monkeypatch.setattr(runtime_file_engine, "runtime_inspect_file_engine", fake_inspect)
    monkeypatch.setattr(runtime_file_engine, "runtime_generate_file", fake_generate)
    monkeypatch.setattr(runtime_file_engine, "runtime_patch_file", fake_patch)
    manor_mcp_file_engine.set_call_context({
        "entity_id": "ent_1",
        "user_id": "user_1",
        "workspace_id": "ws_1",
        "conversation_id": "conv_1",
        "task_id": "task_1",
    })
    try:
        for name, arguments in (
            ("inspect", {"file_type": "docx"}),
            ("generate", {"kind": "document", "name": "brief.md"}),
            ("patch", {"path": "brief.md", "operations": []}),
        ):
            result = await manor_mcp_file_engine.call_tool(name, arguments, "unused")
            assert result["structuredContent"] == {"action": name}
    finally:
        manor_mcp_file_engine.clear_call_context()

    assert calls == [
        ("inspect", {"entity_id": "ent_1", "file_type": "docx"}),
        ("generate", {
            "entity_id": "ent_1",
            "user_id": "user_1",
            "workspace_id": "ws_1",
            "conversation_id": "conv_1",
            "task_id": "task_1",
            "kind": "document",
            "name": "brief.md",
        }),
        ("patch", {
            "entity_id": "ent_1",
            "user_id": "user_1",
            "workspace_id": "ws_1",
            "conversation_id": "conv_1",
            "task_id": "task_1",
            "path": "brief.md",
            "operations": [],
        }),
    ]


@pytest.mark.parametrize(
    "query",
    [
        "office engine patch ppt",
        "edit uploaded Word docx in knowledge",
        "Excel spreadsheet add sheet",
        "文档引擎 编辑 ppt",
    ],
)
def test_file_engine_mcp_is_discoverable_for_office_and_knowledge_queries(query: str) -> None:
    from packages.core.ai.runtime.tool_discovery import (
        runtime_server_index,
        runtime_server_query_score,
    )
    from packages.core.ai.runtime.tool_search import runtime_search_tool_candidates
    from packages.core.ai.tools.mcp_builtin import _SERVER_TOOL_SCHEMAS

    index = runtime_server_index()
    entry = index["manor_mcp_file_engine"]
    assert runtime_server_query_score(entry, query) > 0

    pool = [
        (
            f"mcp__manor_mcp_file_engine__{tool['name']}",
            {
                "type": "function",
                "function": {
                    "name": f"mcp__manor_mcp_file_engine__{tool['name']}",
                    "description": tool.get("description", ""),
                    "parameters": tool.get("inputSchema", {}),
                },
            },
        )
        for tool in _SERVER_TOOL_SCHEMAS["manor_mcp_file_engine"]
    ]
    pool.append((
        "mcp__manor_mcp_calendar__list_booking_links",
        {
            "type": "function",
            "function": {
                "name": "mcp__manor_mcp_calendar__list_booking_links",
                "description": "List booking links.",
                "parameters": {"type": "object", "properties": {}},
            },
        },
    ))

    results, suppressed = runtime_search_tool_candidates(
        tool_schemas=pool,
        query=query,
        server_index=index,
        usable_providers=frozenset({"manor_mcp_file_engine", "manor_mcp_calendar"}),
    )

    assert suppressed == []
    assert results
    assert results[0]["name"].startswith("mcp__manor_mcp_file_engine__")


def test_inspect_file_engine_reports_all_generate_file_kinds() -> None:
    from packages.core.contracts.audio_generation import GenerateFileKind

    payload = file_tools._file_engine_capabilities()

    assert set(payload["generate_kinds"]) == set(GenerateFileKind.values())
    assert {
        "diagram",
        "code",
        "document",
        "word_document",
        "pdf",
        "presentation",
        "spreadsheet",
        "image",
        "video",
        "audio",
    } <= set(payload["generate_kinds"])


@pytest.mark.parametrize(
    ("file_type", "can_patch", "operations"),
    [
        ("txt", True, {"replace_text", "text.replace"}),
        ("md", True, {"replace_text", "text.replace"}),
        ("json", True, {"replace_text", "text.replace"}),
        ("html", True, {"replace_text", "text.replace"}),
        ("csv", True, {"replace_text", "text.replace"}),
        ("docx", True, {"replace_text", "text.replace"}),
        ("pptx", True, {"replace_text", "text.replace"}),
        ("xlsx", True, {"set_cell", "cell.set", "update_row", "row.update", "append_row", "row.append", "add_sheet", "sheet.add"}),
        ("xlsm", True, {"set_cell", "cell.set", "update_row", "row.update", "append_row", "row.append", "add_sheet", "sheet.add"}),
        ("pdf", False, set()),
        ("doc", False, set()),
        ("ppt", False, set()),
        ("xls", False, set()),
    ],
)
def test_inspect_file_engine_reports_patch_contract_for_every_user_file_type(
    file_type: str,
    can_patch: bool,
    operations: set[str],
) -> None:
    payload = file_tools._file_engine_capabilities(file_type)["selected"]

    assert payload["can_patch"] is can_patch
    assert set(payload["operations"]) == operations


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("op", "expected_operation"),
    [
        ("text.replace", "replace_text"),
        ("replace_text", "replace_text"),
        ("cell.set", "set_cell"),
        ("row.update", "update_row"),
        ("row.append", "append_row"),
        ("sheet.add", "add_sheet"),
    ],
)
async def test_patch_file_maps_unified_operations_to_existing_editor(monkeypatch, op, expected_operation) -> None:
    seen: dict[str, object] = {}

    monkeypatch.setattr(file_tools, "_get_entity_root", lambda _entity_id: "/tmp")

    class _NoopLock:
        async def __aenter__(self):
            return self

        async def __aexit__(self, exc_type, exc, tb):
            return False

    monkeypatch.setattr(file_tools, "runtime_entity_filesystem_mutation_lock", lambda _root: _NoopLock())

    async def fake_edit(entity_id: str, **kwargs):
        seen["entity_id"] = entity_id
        seen.update(kwargs)
        return json.dumps({"patched": True})

    monkeypatch.setattr(file_tools, "_edit_file_locked", fake_edit)

    result = json.loads(await file_tools._patch_file(
        "entity-1",
        path="docs/example.txt",
        operations=[{"op": op, "old_text": "old", "new_text": "new"}],
        expected_sha256="sha",
    ))

    assert result == {"patched": True}
    assert seen["entity_id"] == "entity-1"
    assert seen["_source_tool_name"] == "patch_file"
    assert seen["operation"] == expected_operation
    assert seen["expected_sha256"] == "sha"


@pytest.mark.asyncio
async def test_patch_file_fails_closed_for_batch_operations(monkeypatch) -> None:
    monkeypatch.setattr(file_tools, "_get_entity_root", lambda _entity_id: "/tmp")

    class _NoopLock:
        async def __aenter__(self):
            return self

        async def __aexit__(self, exc_type, exc, tb):
            return False

    monkeypatch.setattr(file_tools, "runtime_entity_filesystem_mutation_lock", lambda _root: _NoopLock())

    payload = json.loads(await file_tools._patch_file(
        "entity-1",
        path="docs/example.md",
        operations=[
            {"op": "text.replace", "old_text": "a", "new_text": "b"},
            {"op": "text.replace", "old_text": "c", "new_text": "d"},
        ],
    ))

    assert payload["error"] == "multi_operation_patch_not_supported_yet"
    assert payload["capabilities"]["patch_tool"] == "patch_file"


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["txt", "md", "json", "html", "csv"])
async def test_patch_file_rewrites_text_like_files_end_to_end(tmp_path, monkeypatch, extension) -> None:
    entity_id = "entity-1"
    root = tmp_path / entity_id
    root.mkdir()
    rel_path = f"docs/example.{extension}"
    path = root / rel_path
    path.parent.mkdir()
    path.write_text("alpha old omega", encoding="utf-8")

    monkeypatch.setattr(file_tools, "_get_entity_root", lambda _entity_id: str(root))
    async def allow_file_access(**_kwargs):
        return None

    monkeypatch.setattr(file_tools, "runtime_guard_file_resource_access", allow_file_access)
    monkeypatch.setattr(file_tools, "runtime_guard_file_mutation", allow_file_access)

    async def fake_guard_expected_source_sha(**_kwargs):
        return file_tools.ExpectedSourceGuardResult()

    async def fake_commit_file_projection(**kwargs):
        target = root / kwargs["rel_path"]
        target.write_bytes(kwargs["data"])
        return (
            str(target),
            SimpleNamespace(synced=True, document_id="doc-1", reason=None),
            {"size": len(kwargs["data"]), "mtime_ns": 1, "source_sha256": "new-sha"},
        )

    monkeypatch.setattr(file_tools, "_guard_expected_source_sha", fake_guard_expected_source_sha)
    monkeypatch.setattr(file_tools, "_commit_file_projection", fake_commit_file_projection)

    payload = json.loads(await file_tools._patch_file(
        entity_id,
        path=rel_path,
        operations=[{"op": "text.replace", "old_text": "old", "new_text": "new"}],
        user_id="user-1",
    ))

    assert payload["edited"] is True
    assert payload["path"] == rel_path
    assert path.read_text(encoding="utf-8") == "alpha new omega"
    assert payload["knowledge_synced"] is True


def test_docx_patch_rewrites_real_document_package(tmp_path) -> None:
    from docx import Document

    path = tmp_path / "proposal.docx"
    doc = Document()
    doc.add_paragraph("Hello Old Title")
    doc.save(path)

    result = file_tools._replace_docx_sync(str(path), "Old Title", "New Title", False)
    assert result["edited"] is True
    assert result["file_type"] == "docx"

    out = tmp_path / "proposal-edited.docx"
    out.write_bytes(result["_persisted_bytes"])
    edited = Document(out)
    assert "Hello New Title" in "\n".join(p.text for p in edited.paragraphs)


def test_pptx_patch_rewrites_real_presentation_package(tmp_path) -> None:
    from pptx import Presentation

    path = tmp_path / "deck.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Old Roadmap"
    deck.save(path)

    result = file_tools._replace_pptx_sync(str(path), "Old Roadmap", "New Roadmap", False)
    assert result["edited"] is True
    assert result["file_type"] == "pptx"

    out = tmp_path / "deck-edited.pptx"
    out.write_bytes(result["_persisted_bytes"])
    edited = Presentation(out)
    assert edited.slides[0].shapes.title.text == "New Roadmap"


def test_xlsx_patch_rewrites_real_workbook_package(tmp_path) -> None:
    from openpyxl import Workbook, load_workbook

    path = tmp_path / "model.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Forecast"
    ws["A1"] = "Metric"
    ws["B1"] = "Value"
    ws["A2"] = "Revenue"
    ws["B2"] = 100
    wb.save(path)
    wb.close()

    result = file_tools._spreadsheet_edit_sync(
        str(path),
        {"operation": "set_cell", "sheet": "Forecast", "cell": "B2", "value": 120},
    )
    assert result["updated"] is True
    assert result["operation"] == "set_cell"

    out = tmp_path / "model-edited.xlsx"
    out.write_bytes(result["_persisted_bytes"])
    edited = load_workbook(out)
    try:
        assert edited["Forecast"]["B2"].value == 120
    finally:
        edited.close()


def test_xlsx_patch_adds_sheet_to_existing_workbook_package(tmp_path) -> None:
    from openpyxl import Workbook, load_workbook

    path = tmp_path / "model.xlsx"
    wb = Workbook()
    wb.active.title = "Forecast"
    wb.save(path)
    wb.close()

    result = file_tools._spreadsheet_edit_sync(
        str(path),
        {"operation": "add_sheet", "sheet": "Actuals"},
    )
    assert result["updated"] is True
    assert result["operation"] == "add_sheet"
    assert result["sheet"] == "Actuals"
    assert "Actuals" in result["sheet_names"]
    assert "range" not in result

    out = tmp_path / "model-edited.xlsx"
    out.write_bytes(result["_persisted_bytes"])
    edited = load_workbook(out)
    try:
        assert edited.sheetnames == ["Forecast", "Actuals"]
    finally:
        edited.close()


def test_xlsx_patch_rejects_case_insensitive_duplicate_sheet_name(tmp_path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "model.xlsx"
    wb = Workbook()
    wb.active.title = "Forecast"
    wb.save(path)
    wb.close()

    result = file_tools._spreadsheet_edit_sync(
        str(path),
        {"operation": "add_sheet", "new_sheet_name": "forecast"},
    )

    assert result == {"error": "sheet_already_exists", "sheet": "forecast"}


@pytest.mark.asyncio
@pytest.mark.parametrize("scoped", [False, True])
async def test_patch_file_adds_sheet_to_xlsx_end_to_end(tmp_path, monkeypatch, scoped) -> None:
    from openpyxl import Workbook, load_workbook

    entity_id = "entity-1"
    root = tmp_path / entity_id
    root.mkdir()
    rel_path = "sheets/model.xlsx"
    path = root / rel_path
    path.parent.mkdir()
    wb = Workbook()
    wb.active.title = "Forecast"
    wb.save(path)
    wb.close()
    if scoped:
        literal_path = root / "model.xlsx"
        literal_path.write_bytes(path.read_bytes())

        async def locate_scoped_file(**kwargs):
            assert kwargs["path"] == "model.xlsx"
            return str(path), []

        monkeypatch.setattr(file_tools, "_locate_entity_file", locate_scoped_file)

    monkeypatch.setattr(file_tools, "_get_entity_root", lambda _entity_id: str(root))

    async def allow_file_access(**_kwargs):
        return None

    async def fake_guard_expected_source_sha(**_kwargs):
        return file_tools.ExpectedSourceGuardResult()

    async def fake_commit_file_projection(**kwargs):
        target = root / kwargs["rel_path"]
        target.write_bytes(kwargs["data"])
        return (
            str(target),
            SimpleNamespace(synced=True, document_id="sheet-doc-1", reason=None),
            {"size": len(kwargs["data"]), "mtime_ns": 1, "source_sha256": "new-sha"},
        )

    monkeypatch.setattr(file_tools, "runtime_guard_file_resource_access", allow_file_access)
    monkeypatch.setattr(file_tools, "runtime_guard_file_mutation", allow_file_access)
    monkeypatch.setattr(file_tools, "_guard_expected_source_sha", fake_guard_expected_source_sha)
    monkeypatch.setattr(file_tools, "_commit_file_projection", fake_commit_file_projection)

    payload = json.loads(await file_tools._patch_file(
        entity_id,
        path="model.xlsx" if scoped else rel_path,
        operations=[{"op": "sheet.add", "sheet": "Actuals"}],
        user_id="user-1",
    ))

    assert payload["updated"] is True
    assert payload["operation"] == "add_sheet"
    assert payload["path"] == rel_path
    assert payload["document_id"] == "sheet-doc-1"
    assert payload["knowledge_synced"] is True
    edited = load_workbook(path)
    try:
        assert edited.sheetnames == ["Forecast", "Actuals"]
    finally:
        edited.close()
    if scoped:
        untouched = load_workbook(literal_path)
        try:
            assert untouched.sheetnames == ["Forecast"]
        finally:
            untouched.close()
