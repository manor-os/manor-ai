"""Native PPT picture generation and patching share one authorized executor."""

from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path
from zipfile import ZipFile

from PIL import Image
import pytest
from pptx import Presentation

from packages.core.ai.runtime import file_actions, generated_files
from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools import file_tools
from packages.core.ai.tools.file_tools import _apply_office_patch_sequence_sync
from packages.core.ai.tools.generate_file.tool import _generate_file_handler
from packages.core.contracts.file_engine import (
    file_patch_operations,
    normalize_file_operation_resources,
    normalize_file_patch_operation,
)
from packages.core.services.file_engine_patches import describe_file_structure
from packages.core.services import office_operation_resources
from packages.core.services.office_operation_resources import _validated_image_bytes
from packages.core.models.base import generate_ulid
from tests.test_knowledge_file_consistency import _doc_for, _seed_workspace, fs_enabled  # noqa: F401


def image_bytes(color: str, size=(200, 100), image_format="PNG") -> bytes:
    output = io.BytesIO()
    Image.new("RGB", size, color).save(output, image_format)
    return output.getvalue()


def source(path: str, data: bytes) -> dict[str, str]:
    return {"path": path, "expected_sha256": hashlib.sha256(data).hexdigest()}


def normalized(*operations):
    return [normalize_file_patch_operation(operation) for operation in operations]


def generate(operations, resources):
    return _generate_office_operations_sync("pptx", normalized(*operations), resources=resources)


def picture_insert(src, *, fit="cover", picture_format=None):
    operation = {
        "op": "picture.insert",
        "slide": 1,
        "source": src,
        "transform": {"x": 72, "y": 90, "width": 300, "height": 300, "rotation": 7},
        "fit": fit,
    }
    if picture_format is not None:
        operation["format"] = picture_format
    return operation


def test_picture_operations_are_shared_public_capabilities():
    operations = file_patch_operations("pptx")
    assert {"picture.insert", "picture.replace", "picture.format"} <= set(operations)


def test_picture_sources_are_canonical_hash_bound_and_deduplicated():
    digest = "A" * 64
    patches, resources = normalize_file_operation_resources(
        normalized(
            picture_insert({"path": "Assets\\hero.png", "expected_sha256": digest}),
            {"op": "picture.replace", "slide": 1, "shape_id": 2,
             "source": {"path": "Assets/hero.png", "expected_sha256": digest}},
        ),
        normalize_path=lambda path: path.replace("\\", "/"),
        visible_path=lambda path: not path.startswith("."),
    )
    assert resources == [{"path": "Assets/hero.png", "expected_sha256": digest.lower()}]
    assert patches[0]["source"] == patches[1]["source"] == resources[0]


@pytest.mark.parametrize("src", [
    None,
    {},
    {"path": "hero.png"},
    {"path": "../hero.png", "expected_sha256": "a" * 64},
    {"path": "/hero.png", "expected_sha256": "a" * 64},
    {"path": ".ai/hero.png", "expected_sha256": "a" * 64},
    {"path": "hero.png", "expected_sha256": "short"},
    {"path": "hero.png", "expected_sha256": "g" * 64},
])
def test_invalid_picture_sources_fail_before_resource_access(src):
    with pytest.raises(ValueError):
        normalize_file_operation_resources(
            normalized(picture_insert(src)),
            normalize_path=lambda path: path,
            visible_path=lambda path: not path.startswith("."),
        )


def test_generate_inserts_native_editable_picture_with_cover_crop_and_format(tmp_path):
    data = image_bytes("red", (200, 100))
    src = source("Assets/hero.png", data)
    operations = [
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        picture_insert(src, picture_format={"opacity": 0.55, "alt_text": "Product dashboard"}),
        {"op": "shape.transform", "slide": 1, "shape_id": 2,
         "transform": {"flip_horizontal": True, "x": 80}},
    ]
    result = generate(operations, {(src["path"], src["expected_sha256"]): data})
    assert not result.get("error"), result
    path = tmp_path / "pictures.pptx"
    path.write_bytes(result["_persisted_bytes"])
    picture = Presentation(path).slides[0].shapes[0]
    assert picture.shape_id == 2
    assert picture.left.pt == 80
    assert picture.rotation == 7
    assert picture._element.xfrm.flipH is True
    assert picture.crop_left == picture.crop_right == 0.25
    assert picture.crop_top == picture.crop_bottom == 0
    assert picture._element.nvPicPr.cNvPr.get("descr") == "Product dashboard"
    assert 'amt="55000"' in picture._element.blipFill.blip.xml
    structure = describe_file_structure(str(path))["shapes"][0]
    assert structure["picture"] == {
        "crop": {"left": 0.25, "top": 0.0, "right": 0.25, "bottom": 0.0},
        "opacity": 0.55,
        "alt_text": "Product dashboard",
    }
    assert structure["flip_horizontal"] is True


def test_contain_preserves_aspect_and_centers_inside_requested_frame():
    data = image_bytes("red", (200, 100))
    src = source("Assets/hero.png", data)
    result = generate([
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        picture_insert(src, fit="contain"),
    ], {(src["path"], src["expected_sha256"]): data})
    picture = Presentation(io.BytesIO(result["_persisted_bytes"])).slides[0].shapes[0]
    assert picture.left.pt == 72
    assert picture.top.pt == 165
    assert picture.width.pt == 300
    assert picture.height.pt == 150
    assert not any((picture.crop_left, picture.crop_top, picture.crop_right, picture.crop_bottom))


def test_picture_replace_preserves_layout_and_format_but_changes_media(tmp_path):
    red, blue = image_bytes("red"), image_bytes("blue", (50, 50), "JPEG")
    first, second = source("Assets/red.png", red), source("Assets/blue.jpg", blue)
    generated = generate([
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        picture_insert(first, picture_format={"crop": {"top": 0.1}, "opacity": 0.4, "alt_text": "Before"}),
    ], {(first["path"], first["expected_sha256"]): red})
    path = tmp_path / "replace.pptx"
    path.write_bytes(generated["_persisted_bytes"])
    before = Presentation(path).slides[0].shapes[0]
    state = (before.left, before.top, before.width, before.height, before.rotation,
             before.crop_left, before.crop_top, before.crop_right, before.crop_bottom,
             before._element.nvPicPr.cNvPr.get("descr"), before._element.blipFill.blip.xml)
    result = _apply_office_patch_sequence_sync(
        str(path),
        normalized({"op": "picture.replace", "slide": 1, "shape_id": 2, "source": second}),
        resources={(second["path"], second["expected_sha256"]): blue},
    )
    assert not result.get("error"), result
    assert path.read_bytes() == generated["_persisted_bytes"]
    after = Presentation(io.BytesIO(result["_persisted_bytes"])).slides[0].shapes[0]
    assert (after.left, after.top, after.width, after.height, after.rotation,
            after.crop_left, after.crop_top, after.crop_right, after.crop_bottom,
            after._element.nvPicPr.cNvPr.get("descr")) == state[:-1]
    assert 'amt="40000"' in after._element.blipFill.blip.xml
    assert after.image.size == (50, 50)


def test_picture_replace_keeps_a_shared_media_relationship_for_other_pictures(tmp_path):
    red, blue = image_bytes("red"), image_bytes("blue")
    first, second = source("Assets/red.png", red), source("Assets/blue.png", blue)
    insertion = picture_insert(first, fit="stretch")
    other = picture_insert(first, fit="stretch")
    other["transform"] = {"x": 400, "y": 90, "width": 200, "height": 100}
    created = generate([
        {"op": "slide.insert", "index": 0, "layout_index": 6}, insertion, other,
    ], {(first["path"], first["expected_sha256"]): red})
    path = tmp_path / "shared-media.pptx"
    path.write_bytes(created["_persisted_bytes"])
    before = Presentation(path).slides[0]
    assert before.shapes[0]._element.blipFill.blip.rEmbed == before.shapes[1]._element.blipFill.blip.rEmbed
    result = _apply_office_patch_sequence_sync(
        str(path),
        normalized({"op": "picture.replace", "slide": 1, "shape_id": 2, "source": second}),
        resources={(second["path"], second["expected_sha256"]): blue},
    )
    assert not result.get("error"), result
    slide = Presentation(io.BytesIO(result["_persisted_bytes"])).slides[0]
    assert slide.shapes[0].image.blob == blue
    assert slide.shapes[1].image.blob == red
    assert slide.shapes[0]._element.blipFill.blip.rEmbed != slide.shapes[1]._element.blipFill.blip.rEmbed


def test_picture_changes_preserve_template_charts_masters_and_notes(tmp_path):
    from tests.test_office_template_generation import template_bytes

    before = template_bytes("pptx")
    data = image_bytes("green")
    src = source("Assets/chart-photo.png", data)
    result = _generate_office_operations_sync(
        "pptx",
        normalized(picture_insert(src, fit="contain")),
        template_bytes=before,
        resources={(src["path"], src["expected_sha256"]): data},
    )
    assert not result.get("error"), result
    with ZipFile(io.BytesIO(before)) as original, ZipFile(io.BytesIO(result["_persisted_bytes"])) as edited:
        for name in original.namelist():
            if name.startswith(("ppt/charts/", "ppt/embeddings/", "ppt/slideMasters/", "ppt/notesSlides/")):
                assert edited.read(name) == original.read(name), name


@pytest.mark.parametrize("operation", [
    {"op": "picture.format", "slide": 1, "shape_id": 2, "format": {}},
    {"op": "picture.format", "slide": 1, "shape_id": 2, "format": {"opacity": 1.1}},
    {"op": "picture.format", "slide": 1, "shape_id": 2,
     "format": {"crop": {"left": 0.6, "right": 0.4}}},
    {"op": "picture.format", "slide": 1, "shape_id": 2,
     "format": {"alt_text": "x" * 1001}},
    {"op": "picture.insert", "slide": 1, "source": {},
     "transform": {"x": 0, "y": 0, "width": 100, "height": 100}, "fit": "tile"},
])
def test_invalid_picture_operation_is_atomic(tmp_path, operation):
    data = image_bytes("red")
    src = source("Assets/hero.png", data)
    initial = generate([
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        picture_insert(src),
    ], {(src["path"], src["expected_sha256"]): data})["_persisted_bytes"]
    path = tmp_path / "atomic.pptx"
    path.write_bytes(initial)
    result = _apply_office_patch_sequence_sync(str(path), normalized(operation))
    assert result.get("error"), result
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == initial


@pytest.mark.parametrize("image_format,extension", [("WEBP", "webp"), ("ICO", "ico")])
def test_non_ooxml_image_formats_are_validated_and_embedded_as_png(image_format, extension):
    data = image_bytes("purple", (32, 24), image_format)
    normalized_data = _validated_image_bytes(f"Assets/source.{extension}", data)
    assert normalized_data.startswith(b"\x89PNG\r\n\x1a\n")


def test_office_image_rejects_decoded_memory_amplification(monkeypatch):
    data = image_bytes("purple", (20, 20), "WEBP")
    monkeypatch.setattr(
        office_operation_resources,
        "OFFICE_OPERATION_IMAGE_MAX_DECODED_BYTES",
        1_000,
    )

    with pytest.raises(ValueError, match="decoded image exceeds"):
        _validated_image_bytes("Assets/source.webp", data)


async def synced_image(entity_id: str, root: Path, name="Assets/hero.png") -> tuple[Path, bytes, dict[str, str]]:
    from packages.core.services.knowledge_sync import sync_file_to_knowledge

    path = root / name
    path.parent.mkdir(parents=True, exist_ok=True)
    data = image_bytes("red")
    path.write_bytes(data)
    synced = await sync_file_to_knowledge(
        entity_id=entity_id,
        abs_path=str(path),
        entity_root=str(root),
        source="upload",
        created_by="test",
    )
    assert synced.document_id
    return path, data, source(name, data)


@pytest.mark.asyncio
@pytest.mark.usefixtures("fs_enabled")
@pytest.mark.parametrize("outcome", [
    "success", "stale", "read_denied", "approval_denied", "late_failure",
    "projection_failure", "source_changes_after_snapshot",
])
async def test_picture_generation_uses_acl_hash_approval_and_atomic_projection(
    db_session, monkeypatch, outcome,
):
    entity_id, workspace_id = generate_ulid(), generate_ulid()
    await _seed_workspace(db_session, entity_id, workspace_id)
    root = Path(file_tools._get_entity_root(entity_id))
    image_path, original_image, src = await synced_image(entity_id, root)
    original_image_document = await _doc_for(db_session, entity_id, Path(src["path"]).name)
    reads, approvals = [], []

    async def read_guard(**kwargs):
        reads.append(kwargs)
        if outcome == "read_denied":
            return json.dumps({"error": "file_permission_denied"})

    async def approve(**kwargs):
        approvals.append(kwargs)
        if outcome == "approval_denied":
            return json.dumps({"error": "approval_required"})
        if outcome == "source_changes_after_snapshot":
            image_path.write_bytes(image_bytes("blue"))

    monkeypatch.setattr(file_actions, "runtime_guard_file_read_access", read_guard)
    monkeypatch.setattr(file_actions, "runtime_guard_file_mutation", approve)
    if outcome == "projection_failure":
        async def fail_metadata(_path):
            raise OSError("injected picture projection failure")

        monkeypatch.setattr(generated_files, "runtime_generated_file_metadata", fail_metadata)
    supplied = dict(src)
    if outcome == "stale":
        supplied["expected_sha256"] = "a" * 64
    operations = [
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        picture_insert(supplied, picture_format={"alt_text": "Approved image"}),
    ]
    if outcome == "late_failure":
        operations.append({"op": "picture.format", "slide": 1, "shape_id": 999,
                           "format": {"opacity": 0.5}})
    result = json.loads(await _generate_file_handler(
        entity_id,
        kind="presentation",
        name="picture-generation.pptx",
        workspace_id=workspace_id,
        operations=operations,
    ))
    assert len(reads) == 1
    assert reads[0]["paths"] == [src["path"]]
    assert reads[0]["workspace_id"] == workspace_id
    if approvals:
        assert len(approvals) == 1
        assert approvals[0]["approval_payload"]["operations"][1]["source"] == supplied
    output_document = await _doc_for(db_session, entity_id, "picture-generation.pptx")
    if outcome in {"success", "source_changes_after_snapshot"}:
        assert result.get("created"), result
        assert output_document is not None
        output = Path(root, output_document.fs_path)
        picture = Presentation(output).slides[0].shapes[0]
        assert picture.image.blob == original_image
        assert picture._element.nvPicPr.cNvPr.get("descr") == "Approved image"
    else:
        assert result.get("error"), result
        assert output_document is None
        if outcome in {"stale", "read_denied"}:
            assert not approvals
    reloaded_image_document = await _doc_for(db_session, entity_id, Path(src["path"]).name)
    assert reloaded_image_document.id == original_image_document.id
    if outcome != "source_changes_after_snapshot":
        assert image_path.read_bytes() == original_image


@pytest.mark.asyncio
@pytest.mark.usefixtures("fs_enabled")
@pytest.mark.parametrize("outcome", [
    "success", "stale", "read_denied", "approval_denied", "late_failure", "projection_failure",
])
async def test_picture_patch_keeps_target_identity_and_rolls_back_every_failure(
    db_session, monkeypatch, outcome,
):
    entity_id, workspace_id = generate_ulid(), generate_ulid()
    await _seed_workspace(db_session, entity_id, workspace_id)
    created = json.loads(await _generate_file_handler(
        entity_id,
        kind="presentation",
        name="picture-patch.pptx",
        workspace_id=workspace_id,
        operations=[{"op": "slide.insert", "index": 0, "layout_index": 6}],
    ))
    assert created.get("created"), created
    target_document = await _doc_for(db_session, entity_id, "picture-patch.pptx")
    target_id = target_document.id
    root = Path(file_tools._get_entity_root(entity_id))
    target = root / target_document.fs_path
    target_before = target.read_bytes()
    image_path, original_image, src = await synced_image(entity_id, root)
    reads, approvals = [], []

    async def read_guard(**kwargs):
        reads.append(kwargs)
        if outcome == "read_denied":
            return json.dumps({"error": "file_permission_denied"})

    async def approve(**kwargs):
        approvals.append(kwargs)
        if outcome == "approval_denied":
            return json.dumps({"error": "approval_required"})

    monkeypatch.setattr(file_tools, "runtime_guard_file_read_access", read_guard)
    monkeypatch.setattr(file_tools, "runtime_guard_file_mutation", approve)
    if outcome == "projection_failure":
        def fail_metadata(_path):
            raise OSError("injected picture metadata failure")

        monkeypatch.setattr(file_tools, "_file_meta", fail_metadata)
    supplied = dict(src)
    if outcome == "stale":
        supplied["expected_sha256"] = "b" * 64
    operations = [picture_insert(supplied)]
    if outcome == "late_failure":
        operations.append({"op": "picture.format", "slide": 1, "shape_id": 999,
                           "format": {"opacity": 0.5}})
    result = json.loads(await file_tools._patch_file(
        entity_id,
        path=target_document.fs_path,
        workspace_id=workspace_id,
        expected_sha256=hashlib.sha256(target_before).hexdigest(),
        operations=operations,
    ))
    assert len(reads) == 1
    assert reads[0]["paths"] == [src["path"]]
    if approvals:
        assert approvals[0]["content_preview"]["operations"][0]["source"] == supplied
    reloaded = await _doc_for(db_session, entity_id, "picture-patch.pptx")
    assert reloaded.id == target_id
    if outcome == "success":
        assert result.get("patched"), result
        assert Presentation(target).slides[0].shapes[0].image.blob == original_image
    else:
        assert result.get("error"), result
        assert target.read_bytes() == target_before
        if outcome in {"stale", "read_denied"}:
            assert not approvals
    assert image_path.read_bytes() == original_image
