"""Spreadsheet patches must not guess targets or silently lose workbook data."""

from __future__ import annotations

import io
import json
from datetime import date, datetime, time, timedelta

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Border, Side

from packages.core.ai.tools import file_tools


def _source(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", "Value"])
    sheet.append(["Alice", 10])
    workbook.save(path)
    workbook.close()


def _patch(path, *operations):
    return file_tools._apply_office_patch_sequence_sync(str(path), list(operations))


def _workbook(result):
    assert result.get("patched") is True, result
    return load_workbook(io.BytesIO(result["_persisted_bytes"]), rich_text=True)


@pytest.mark.parametrize("operation", ["set_cell", "cell.format"])
@pytest.mark.parametrize("reference", ["XFE1", "ZZZ1", "A1048577", "A0", "A1:B2", "1", "A", "A-1"])
def test_cell_coordinates_reject_out_of_grid_or_non_cell_targets(tmp_path, operation, reference):
    path = tmp_path / "data.xlsx"
    _source(path)
    original = path.read_bytes()

    result = _patch(path, {"operation": operation, "cell": reference, "value": "new", "format": {"bold": True}})

    assert result.get("error"), result
    assert result["operation_index"] == 0
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == original


def test_set_cell_requires_explicit_value_to_clear_existing_data(tmp_path):
    path = tmp_path / "data.xlsx"
    _source(path)
    result = _patch(path, {"operation": "set_cell", "cell": "B2"})
    assert result.get("error") == "value is required for set_cell"
    assert result["operation_index"] == 0

    cleared = _workbook(_patch(path, {"operation": "set_cell", "cell": "B2", "value": None}))
    assert cleared.active["B2"].value is None
    assert cleared.active["A2"].value == "Alice"
    cleared.close()


@pytest.mark.parametrize("operation", ["set_cell", "update_row", "append_row"])
@pytest.mark.parametrize(
    "value",
    ["x" * 32768, float("nan"), float("inf"), [1, 2], {"nested": "value"}, "bad\x01text"],
    ids=["too-long", "nan", "infinity", "array", "object", "control-character"],
)
def test_invalid_cell_values_fail_before_serialization(tmp_path, operation, value):
    path = tmp_path / "data.xlsx"
    _source(path)
    original = path.read_bytes()
    patch = {
        "set_cell": {"cell": "B2", "value": value},
        "update_row": {"match_column": "Name", "match_value": "Alice", "updates": {"Value": value}},
        "append_row": {"values": ["Bob", value]},
    }[operation]

    result = _patch(path, {"operation": operation, **patch})

    assert result.get("error"), result
    assert result["operation_index"] == 0
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == original


@pytest.mark.parametrize("operation", ["update_row", "append_row"])
@pytest.mark.parametrize("header_row", [0, -1, True, 1.5, "1", None])
def test_invalid_header_row_is_not_coerced_to_a_different_row(tmp_path, operation, header_row):
    path = tmp_path / "data.xlsx"
    _source(path)
    result = _patch(
        path,
        {
            "operation": operation,
            "header_row": header_row,
            "match_column": "Name",
            "match_value": "Alice",
            "updates": {"Value": 99},
            "row": {"Name": "Bob", "Value": 20},
        },
    )
    assert result.get("error"), result
    assert result["operation_index"] == 0


@pytest.mark.parametrize("operation", ["set_cell", "cell.format"])
def test_merged_follower_is_rejected_with_anchor_hint(tmp_path, operation):
    path = tmp_path / "merged.xlsx"
    workbook = Workbook()
    workbook.active.merge_cells("A1:B2")
    workbook.active["A1"] = "Keep merged content"
    workbook.save(path)
    workbook.close()

    result = _patch(path, {"operation": operation, "cell": "B2", "value": "new", "format": {"bold": True}})

    assert "merged" in result.get("error", "").lower(), result
    assert "A1" in result["error"]
    assert result["operation_index"] == 0


@pytest.mark.parametrize("operation", ["update_row", "append_row"])
def test_duplicate_header_does_not_pick_the_first_column(tmp_path, operation):
    path = tmp_path / "duplicate.xlsx"
    workbook = Workbook()
    workbook.active.append(["Name", "Value", " value "])
    workbook.active.append(["Alice", 10, 20])
    workbook.save(path)
    workbook.close()

    result = _patch(
        path,
        {
            "operation": operation,
            "match_column": "Name",
            "match_value": "Alice",
            "updates": {"Value": 99},
            "row": {"Name": "Bob", "Value": 99},
        },
    )
    assert "ambiguous" in result.get("error", "").lower(), result
    assert result["operation_index"] == 0


def test_update_row_requires_a_match_value_even_when_blank_rows_exist(tmp_path):
    path = tmp_path / "blank-key.xlsx"
    workbook = Workbook()
    workbook.active.append(["Name", "Value"])
    workbook.active.append([None, 10])
    workbook.save(path)
    workbook.close()

    result = _patch(path, {"operation": "update_row", "match_column": "Name", "updates": {"Value": 99}})
    assert result.get("error") == "match_value is required for update_row"


@pytest.mark.parametrize("values", [[], "not-a-row"])
def test_append_empty_or_invalid_values_does_not_report_success(tmp_path, values):
    path = tmp_path / "data.xlsx"
    _source(path)
    result = _patch(path, {"operation": "append_row", "values": values})
    assert result.get("error"), result


def test_append_rejects_conflicting_row_and_values(tmp_path):
    path = tmp_path / "data.xlsx"
    _source(path)
    result = _patch(path, {"operation": "append_row", "row": {"Name": "Alice"}, "values": ["Bob", 20]})
    assert result.get("error"), result


def test_append_to_a_new_sheet_starts_at_first_row(tmp_path):
    path = tmp_path / "data.xlsx"
    _source(path)
    result = _patch(
        path,
        {"operation": "add_sheet", "sheet": "New"},
        {"operation": "append_row", "sheet": "New", "values": ["Name", "Value"]},
    )
    workbook = _workbook(result)
    assert result["operation_results"][1]["row_number"] == 1
    assert [workbook["New"].cell(1, col).value for col in (1, 2)] == ["Name", "Value"]
    workbook.close()


@pytest.mark.parametrize("name", ["x" * 32, 12, "bad/name", "'quoted", "quoted'"])
def test_add_sheet_rejects_invalid_names_before_save(tmp_path, name):
    path = tmp_path / "data.xlsx"
    _source(path)
    result = _patch(path, {"operation": "add_sheet", "sheet": name})
    assert result.get("error"), result
    assert result["operation_index"] == 0


@pytest.mark.parametrize("sheet", [False, 0, [], {}])
def test_invalid_sheet_selector_does_not_fall_back_to_active_sheet(tmp_path, sheet):
    path = tmp_path / "data.xlsx"
    _source(path)
    result = _patch(path, {"operation": "set_cell", "sheet": sheet, "cell": "B2", "value": 99})
    assert result.get("error"), result


@pytest.mark.parametrize("operation", ["set_cell", "cell.format", "append_row", "add_sheet"])
def test_unedited_rich_text_and_other_workbook_features_survive(tmp_path, operation):
    from openpyxl.comments import Comment
    from openpyxl.chart import BarChart, Reference
    from openpyxl.worksheet.datavalidation import DataValidation

    path = tmp_path / "rich.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Edit here"
    rich = CellRichText("Plain ", TextBlock(InlineFont(b=True, color="FF0000"), "Bold"))
    sheet["B1"] = rich
    sheet["C1"] = "=1+2"
    sheet["D1"] = datetime(2026, 8, 30, 10, 15)
    sheet["A1"].border = Border(bottom=Side(style="thin"))
    sheet["B1"].comment = Comment("Keep comment", "Author")
    sheet["B1"].hyperlink = "https://example.com/source"
    sheet.freeze_panes = "A2"
    validation = DataValidation(type="list", formula1='"One,Two"')
    sheet.add_data_validation(validation)
    validation.add("A2:A10")
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=3, min_row=1, max_row=1))
    sheet.add_chart(chart, "F1")
    workbook.save(path)
    workbook.close()
    patch = {
        "set_cell": {"cell": "A1", "value": "Edited"},
        "cell.format": {"cell": "A1", "format": {"bold": True}},
        "append_row": {"values": ["New row", 20]},
        "add_sheet": {"sheet": "New"},
    }[operation]

    edited = _workbook(_patch(path, {"operation": operation, **patch}))
    sheet = edited["Sheet"]
    assert isinstance(sheet["B1"].value, CellRichText)
    assert sheet["B1"].value == rich
    assert sheet["B1"].comment.text == "Keep comment"
    assert sheet["B1"].hyperlink.target == "https://example.com/source"
    assert sheet["C1"].value == "=1+2"
    assert sheet["D1"].value == datetime(2026, 8, 30, 10, 15)
    assert sheet["A1"].border.bottom.style == "thin"
    assert sheet.freeze_panes == "A2"
    assert len(sheet._charts) == 1
    assert str(sheet.data_validations.dataValidation[0].sqref) == "A2:A10"
    edited.close()


@pytest.mark.parametrize(
    "previous", [date(2026, 8, 30), datetime(2026, 8, 30, 10, 15), time(10, 15), timedelta(hours=4)]
)
def test_replacement_result_is_json_serializable_for_date_cells(tmp_path, previous):
    path = tmp_path / "date.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = previous
    workbook.save(path)
    workbook.close()

    result = _patch(path, {"operation": "set_cell", "cell": "A1", "value": "Replaced"})
    assert result.get("patched") is True, result
    result.pop("_persisted_bytes")
    encoded = json.loads(json.dumps(result, allow_nan=False))
    assert isinstance(encoded["operation_results"][0]["updated_cells"]["A1"]["old"], str)


def test_valid_boundary_values_and_formatting_roundtrip(tmp_path):
    path = tmp_path / "data.xlsx"
    _source(path)
    workbook = _workbook(
        _patch(
            path,
            {"operation": "set_cell", "cell": "XFD1048576", "value": "x" * 32767},
            {"operation": "cell.format", "cell": "XFD1048576", "format": {"bold": True}},
        )
    )
    assert workbook.active["XFD1048576"].value == "x" * 32767
    assert workbook.active["XFD1048576"].font.bold is True
    assert workbook.active["XFD1048576"].font.name == "Calibri"
    workbook.close()


@pytest.mark.parametrize("header", [0, False])
def test_falsy_header_values_remain_addressable(tmp_path, header):
    path = tmp_path / "numeric-header.xlsx"
    workbook = Workbook()
    workbook.active.append(["Name", header])
    workbook.active.append(["Alice", 10])
    workbook.save(path)
    workbook.close()

    edited = _workbook(
        _patch(
            path,
            {
                "operation": "update_row",
                "match_column": "Name",
                "match_value": "Alice",
                "updates": {str(header): 20},
            },
        )
    )
    assert edited.active["B2"].value == 20
    edited.close()


def test_unique_column_can_be_used_when_unrelated_headers_are_duplicated(tmp_path):
    path = tmp_path / "duplicates.xlsx"
    workbook = Workbook()
    workbook.active.append(["Name", "Value", "Other", " other "])
    workbook.active.append(["Alice", 10, 30, 40])
    workbook.save(path)
    workbook.close()

    edited = _workbook(
        _patch(
            path,
            {
                "operation": "update_row",
                "match_column": "Name",
                "match_value": "Alice",
                "updates": {"Value": 20},
            },
        )
    )
    assert [edited.active.cell(2, col).value for col in (2, 3, 4)] == [20, 30, 40]
    edited.close()


@pytest.mark.parametrize("operation", ["update_row", "append_row"])
def test_duplicate_normalized_update_keys_fail_atomically(tmp_path, operation):
    path = tmp_path / "data.xlsx"
    _source(path)
    original = path.read_bytes()
    result = _patch(
        path,
        {
            "operation": operation,
            "match_column": "Name",
            "match_value": "Alice",
            "updates": {"Value": 1, " value ": 2},
            "row": {"Value": 1, " value ": 2},
        },
    )
    assert "duplicate" in result.get("error", "").lower(), result
    assert result["operation_index"] == 0
    assert path.read_bytes() == original


def test_named_sheet_with_surrounding_spaces_is_not_redirected(tmp_path):
    path = tmp_path / "spaces.xlsx"
    workbook = Workbook()
    workbook.active.title = "Data"
    workbook.active["A1"] = "Leave alone"
    workbook.create_sheet(" Data ")["A1"] = "Edit this"
    workbook.save(path)
    workbook.close()

    edited = _workbook(_patch(path, {"operation": "set_cell", "sheet": " Data ", "cell": "A1", "value": "Edited"}))
    assert edited["Data"]["A1"].value == "Leave alone"
    assert edited[" Data "]["A1"].value == "Edited"
    edited.close()


def test_header_row_after_1000_is_used_without_clamping(tmp_path):
    path = tmp_path / "late-header.xlsx"
    workbook = Workbook()
    workbook.active["A1005"], workbook.active["B1005"] = "Name", "Value"
    workbook.active["A1006"], workbook.active["B1006"] = "Alice", 10
    workbook.save(path)
    workbook.close()

    edited = _workbook(
        _patch(
            path,
            {
                "operation": "update_row",
                "header_row": 1005,
                "match_column": "Name",
                "match_value": "Alice",
                "updates": {"Value": 20},
            },
        )
    )
    assert edited.active["B1006"].value == 20
    edited.close()


@pytest.mark.parametrize("target", ["last-row", "too-many-columns"])
def test_append_cannot_exceed_workbook_grid(tmp_path, target):
    path = tmp_path / "full.xlsx"
    workbook = Workbook()
    workbook.active["A1048576" if target == "last-row" else "A1"] = "Last"
    workbook.save(path)
    workbook.close()
    values = ["new"] if target == "last-row" else [None] * 16385
    result = _patch(path, {"operation": "append_row", "values": values})
    assert result.get("error"), result
    assert result["operation_index"] == 0


def test_merged_anchor_value_and_format_remain_editable(tmp_path):
    path = tmp_path / "merged.xlsx"
    workbook = Workbook()
    workbook.active.merge_cells("A1:B2")
    workbook.active["A1"] = "Before"
    workbook.save(path)
    workbook.close()

    edited = _workbook(
        _patch(
            path,
            {"operation": "set_cell", "cell": "A1", "value": "After"},
            {"operation": "cell.format", "cell": "A1", "format": {"bold": True}},
        )
    )
    assert edited.active["A1"].value == "After"
    assert edited.active["A1"].font.bold is True
    assert str(edited.active.merged_cells) == "A1:B2"
    edited.close()


def test_rich_text_old_value_is_plain_json_text(tmp_path):
    path = tmp_path / "rich-old.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = CellRichText("Plain ", TextBlock(InlineFont(b=True), "Bold"))
    workbook.save(path)
    workbook.close()

    result = _patch(path, {"operation": "set_cell", "cell": "A1", "value": "Replaced"})
    assert result.get("patched") is True, result
    result.pop("_persisted_bytes")
    encoded = json.loads(json.dumps(result, allow_nan=False))
    assert encoded["operation_results"][0]["updated_cells"]["A1"]["old"] == "Plain Bold"
