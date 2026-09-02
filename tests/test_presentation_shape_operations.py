"""Native PPT shape creation and patch formatting share one byte-preserving engine."""
from __future__ import annotations

import copy
import hashlib
import io
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from pptx import Presentation
from sqlalchemy import select

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools.file_tools import _apply_office_patch_sequence_sync
from packages.core.ai.tools import file_tools
from packages.core.ai.tools.generate_file.tool import _generate_file_handler
from packages.core.contracts.file_engine import PresentationShapePreset, normalize_file_patch_operation
from packages.core.models.base import generate_ulid
from packages.core.models.event import EventLog
from packages.core.services.file_engine_patches import describe_file_structure
from tests.test_knowledge_file_consistency import _doc_for, _seed_workspace, fs_enabled  # noqa: F401


def card_style():
    return {"fill_color": "174C46", "fill_opacity": 0.85, "line_color": "D89B45", "line_width": 2.5,
            "line_opacity": 0.7, "line_dash": "dash", "corner_radius": 18,
            "margin_left": 18, "margin_right": 18, "margin_top": 12, "margin_bottom": 12,
            "vertical_alignment": "middle", "word_wrap": True,
            "shadow": {"color": "000000", "opacity": 0.25, "blur": 6, "distance": 3, "angle": 45}}


def shape_operations():
    return [
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        {"op": "page.setup", "format": {"width": 960, "height": 540}},
        {"op": "shape.insert", "slide": 1, "preset": "roundRect", "transform": {"x": 72, "y": 90, "width": 480, "height": 180}, "text": "Native card", "format": card_style()},
        {"op": "paragraph.format", "slide": 1, "shape_id": 2, "index": 0, "format": {"font_size": 28, "font_color": "FFFFFF", "font_name": "Arial", "bold": True, "alignment": "center"}},
        {"op": "shape.insert", "slide": 1, "preset": "ellipse", "transform": {"x": 630, "y": 100, "width": 160, "height": 160}, "format": {"gradient_fill": {"angle": 0, "stops": [{"position": 0, "color": "174C46"}, {"position": 1, "color": "A2DBBE", "opacity": 0.6}]}, "line_color": None}},
        {"op": "shape.insert", "slide": 1, "preset": "line", "transform": {"x": 72, "y": 330, "width": 720, "height": 0}, "format": {"line_color": "174C46", "line_width": 3, "line_dash": "dashDot"}},
    ]


def generated(operations=None):
    return _generate_office_operations_sync("pptx", [normalize_file_patch_operation(op) for op in (operations or shape_operations())])


def patch(path, *operations):
    return _apply_office_patch_sequence_sync(str(path), [normalize_file_patch_operation(op) for op in operations])


def shape_format(**style):
    return {"op": "shape.format", "slide": 1, "shape_id": 2, "format": style}


def test_generated_shapes_are_native_and_styles_are_discoverable(tmp_path):
    result = generated()
    assert not result.get("error"), result
    path = tmp_path / "card.pptx"
    path.write_bytes(result["_persisted_bytes"])
    office = Presentation(path)
    assert len(office.slides[0].shapes) == 3
    card, ellipse, line = office.slides[0].shapes
    assert card.text == "Native card"
    assert card.left.pt == 72 and card.width.pt == 480
    assert card.adjustments[0] == 0.1
    assert card.line.width.pt == 2.5
    assert card.text_frame.margin_left.pt == 18
    assert line.height == 0
    structure = describe_file_structure(str(path))
    assert [s["preset"] for s in structure["shapes"]] == ["roundRect", "ellipse", "line"]
    assert structure["shapes"][0]["format"] == card_style()
    assert structure["shapes"][1]["format"]["gradient_fill"]["stops"][1]["opacity"] == 0.6
    assert structure["shapes"][2]["format"]["line_dash"] == "dashDot"
    assert "gradFill" in ellipse._element.xml


@pytest.mark.parametrize("preset", PresentationShapePreset.values())
def test_every_exposed_preset_creates_an_editable_native_object(preset):
    insertion = {"op": "shape.insert", "slide": 1, "preset": preset, "transform": {"x": 30.05, "y": 40.25, "width": 150, "height": 75, "rotation": 15}}
    result = generated([shape_operations()[0], insertion])
    assert not result.get("error"), result
    shape = Presentation(io.BytesIO(result["_persisted_bytes"])).slides[0].shapes[0]
    assert shape.shape_id == 2
    assert shape.rotation == 15
    assert shape.left == round(30.05 * 12700)
    assert shape.top == round(40.25 * 12700)
    assert shape.has_text_frame is (preset != "line")


def test_selective_style_changes_preserve_text_geometry_other_shapes_and_source(tmp_path):
    path = tmp_path / "card.pptx"
    before = generated()["_persisted_bytes"]
    path.write_bytes(before)
    original = Presentation(io.BytesIO(before))
    original_shape = original.slides[0].shapes[0]
    original_text = original_shape.text_frame._txBody.xml
    original_line = original_shape._element.spPr.ln.xml
    other_shapes = [shape._element.xml for shape in original.slides[0].shapes][1:]
    result = patch(path, shape_format(fill_color="336699", fill_opacity=0.4, corner_radius=0, shadow=None))
    assert not result.get("error"), result
    assert path.read_bytes() == before
    office = Presentation(io.BytesIO(result["_persisted_bytes"]))
    card = office.slides[0].shapes[0]
    assert card.text_frame._txBody.xml == original_text
    assert card._element.spPr.ln.xml == original_line
    assert (card.left, card.top, card.width, card.height) == (original_shape.left, original_shape.top, original_shape.width, original_shape.height)
    assert card.adjustments[0] == 0
    assert card._element.spPr.effectLst is not None and len(card._element.spPr.effectLst) == 0
    assert [shape._element.xml for shape in office.slides[0].shapes][1:] == other_shapes
    path.write_bytes(result["_persisted_bytes"])
    followup = patch(path, {"op": "text.set", "slide": 1, "shape_id": 2, "index": 0, "text": "Client card"})
    card = Presentation(io.BytesIO(followup["_persisted_bytes"])).slides[0].shapes[0]
    assert card.text == "Client card"
    assert card.text_frame.paragraphs[0].runs[0].font.bold
    assert str(card.fill.fore_color.rgb) == "336699"


def test_formatting_does_not_create_or_override_an_inherited_line(tmp_path):
    path = tmp_path / "theme.pptx"
    operations = [shape_operations()[0], {"op": "shape.insert", "slide": 1, "preset": "rect", "transform": {"x": 20, "y": 20, "width": 200, "height": 100}}]
    path.write_bytes(generated(operations)["_persisted_bytes"])
    assert Presentation(path).slides[0].shapes[0]._element.spPr.ln is None
    result = patch(path, shape_format(fill_color="336699"))
    assert Presentation(io.BytesIO(result["_persisted_bytes"])).slides[0].shapes[0]._element.spPr.ln is None


def test_template_shapes_do_not_rewrite_charts_masters_or_notes(tmp_path):
    from tests.test_office_template_generation import template_bytes

    before = template_bytes("pptx")
    path = tmp_path / "template.pptx"
    path.write_bytes(before)
    result = patch(path, shape_operations()[2])
    assert not result.get("error"), result
    native = Presentation(io.BytesIO(result["_persisted_bytes"]))
    assert len(native.slides[0].shapes) == 3
    assert list(native.slides[0].shapes[1].chart.series[0].values) == [12, 15]
    with ZipFile(io.BytesIO(before)) as source, ZipFile(io.BytesIO(result["_persisted_bytes"])) as output:
        for name in source.namelist():
            if name.startswith(("ppt/charts/", "ppt/embeddings/", "ppt/slideMasters/", "ppt/notesSlides/")):
                assert output.read(name) == source.read(name), name
    assert path.read_bytes() == before


def test_shadow_edit_rejects_complex_effect_graph_but_fill_keeps_it(tmp_path):
    from pptx.oxml.xmlchemy import OxmlElement

    path = tmp_path / "effects.pptx"
    native = Presentation(io.BytesIO(generated()["_persisted_bytes"]))
    properties = native.slides[0].shapes[0]._element.spPr
    properties.remove(properties.effectLst)
    graph = OxmlElement("a:effectDag")
    properties.append(graph)
    native.save(path)
    before = path.read_bytes()
    result = patch(path, shape_format(shadow=None))
    assert result.get("error"), result
    assert path.read_bytes() == before
    assert "_persisted_bytes" not in result
    result = patch(path, shape_format(fill_color="336699"))
    assert not result.get("error"), result
    assert "effectDag" in Presentation(io.BytesIO(result["_persisted_bytes"])).slides[0].shapes[0]._element.xml


@pytest.mark.parametrize("operation", [
    {"op": "shape.format", "slide": 0, "shape_id": 2, "format": {"fill_color": "112233"}},
    {"op": "shape.format", "slide": 1, "shape_id": 999, "format": {"fill_color": "112233"}},
    {"op": "shape.format", "slide": 1, "shape_id": 2, "text": "ignored", "format": {"fill_color": "112233"}},
    {"op": "shape.format", "slide": 1, "shape_id": 3, "format": {"corner_radius": 10}},
    {"op": "shape.format", "slide": 1, "shape_id": 4, "format": {"fill_color": "112233"}},
])
def test_shape_format_rejects_wrong_selectors_and_inapplicable_fields(tmp_path, operation):
    path = tmp_path / "card.pptx"
    before = generated()["_persisted_bytes"]
    path.write_bytes(before)
    result = patch(path, operation)
    assert result.get("error"), result
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before


@pytest.mark.parametrize("style", [
    {}, {"unknown": True}, {"fill_color": "#112233"}, {"fill_opacity": True}, {"fill_opacity": 1.1},
    {"fill_color": None, "fill_opacity": 0.4}, {"line_width": -1}, {"line_width": 73}, {"line_width": 1.005},
    {"line_color": None, "line_opacity": 0.5}, {"line_dash": "invalid"}, {"corner_radius": 91},
    {"margin_left": 480}, {"vertical_alignment": "center"}, {"word_wrap": 1}, {"shadow": {}},
    {"shadow": {"color": "000000", "opacity": 1.1, "blur": 6, "distance": 3, "angle": 45}},
    {"gradient_fill": {"angle": 0, "stops": []}},
    {"gradient_fill": {"angle": -1, "stops": [{"position": 0, "color": "112233"}, {"position": 1, "color": "445566"}]}},
    {"gradient_fill": {"angle": 0, "stops": [{"position": 1, "color": "112233"}, {"position": 0, "color": "445566"}]}},
    {"gradient_fill": {"angle": 0, "stops": [{"position": 0, "color": "112233"}, {"position": 1, "color": "445566"}]}, "fill_color": "FFFFFF"},
])
def test_invalid_shape_style_is_atomic(tmp_path, style):
    path = tmp_path / "card.pptx"
    before = generated()["_persisted_bytes"]
    path.write_bytes(before)
    result = patch(path, shape_format(fill_color="112233"), shape_format(**copy.deepcopy(style)))
    assert result.get("error"), result
    assert result["operation_index"] == 1
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before


@pytest.mark.parametrize("extra", [
    {"preset": "unknown"}, {"text": None}, {"transform": {"x": 0}}, {"shape_id": 2}, {"style": "ignored"},
    {"preset": "line", "text": "cannot edit line text"},
    {"preset": "line", "transform": {"x": 0, "y": 0, "width": 0, "height": 0}},
    {"preset": "line", "format": {"fill_color": "112233"}},
])
def test_invalid_shape_insert_rejects_without_partial_output(extra):
    insertion = {"op": "shape.insert", "slide": 1, "preset": "rect", "transform": {"x": 20, "y": 20, "width": 200, "height": 100}, **extra}
    result = generated([shape_operations()[0], insertion])
    assert result.get("error"), result
    assert "_persisted_bytes" not in result


@pytest.mark.asyncio
@pytest.mark.usefixtures("fs_enabled")
@pytest.mark.parametrize("mutation", ["insert", "format"])
@pytest.mark.parametrize("outcome", ["success", "late_failure", "projection_failure", "stale", "approval_denied"])
async def test_native_shape_tools_keep_knowledge_identity_and_atomicity(db_session, monkeypatch, mutation, outcome):
    entity_id, workspace_id = generate_ulid(), generate_ulid()
    await _seed_workspace(db_session, entity_id, workspace_id)
    generated_result = json.loads(await _generate_file_handler(
        entity_id, kind="presentation", name="shapes.pptx", workspace_id=workspace_id, operations=shape_operations(),
    ))
    assert generated_result.get("created"), generated_result
    original = await _doc_for(db_session, entity_id, "shapes.pptx")
    original_id, metadata = original.id, copy.deepcopy(original.metadata_)
    path = Path(file_tools._get_entity_root(entity_id), original.fs_path)
    before = path.read_bytes()
    read = json.loads(await file_tools._read_file(entity_id, path="shapes.pptx", workspace_id=workspace_id, include_structure=True))
    assert read["source_sha256"] == hashlib.sha256(before).hexdigest()
    assert read["structure"]["shapes"][0]["format"] == card_style()
    events = set((await db_session.scalars(select(EventLog.id).where(EventLog.entity_id == entity_id))).all())
    operations = [shape_format(fill_color="336699", line_dash="dot")]
    if mutation == "insert":
        operations = [shape_operations()[2]]
    if outcome == "late_failure":
        operations.append(shape_format(line_width=-1))
    if outcome == "projection_failure":
        def reject_metadata(_path):
            raise OSError("injected shape metadata failure")
        monkeypatch.setattr(file_tools, "_file_meta", reject_metadata)
    approvals = []

    async def approve(**kwargs):
        approvals.append(kwargs)
        if outcome == "approval_denied":
            return json.dumps({"error": "approval_required"})

    monkeypatch.setattr(file_tools, "runtime_guard_file_mutation", approve)
    result = json.loads(await file_tools._patch_file(
        entity_id, path="shapes.pptx", workspace_id=workspace_id,
        expected_sha256="a" * 64 if outcome == "stale" else read["source_sha256"], operations=operations,
    ))
    document = await _doc_for(db_session, entity_id, "shapes.pptx")
    assert document.id == original_id
    if approvals:
        assert len(approvals) == 1
        assert approvals[0]["content_preview"]["operations"] == [normalize_file_patch_operation(op) for op in operations]
    if outcome == "success":
        assert result.get("patched"), result
        assert result["document_id"] == original_id
        assert result["source_sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
        native = Presentation(path)
        assert len(native.slides[0].shapes) == (4 if mutation == "insert" else 3)
        assert native.slides[0].shapes[0].text == "Native card"
        assert describe_file_structure(str(path))["shapes"][0]["format"]["line_dash"] == ("dot" if mutation == "format" else "dash")
    else:
        assert result.get("error"), result
        assert path.read_bytes() == before
        assert document.metadata_ == metadata
        assert set((await db_session.scalars(select(EventLog.id).where(EventLog.entity_id == entity_id))).all()) == events
