"""PowerPoint template slides duplicate through the shared native executor."""
from __future__ import annotations

import io
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools.file_tools import _apply_office_patch_sequence_sync
from packages.core.contracts.file_engine import file_patch_operations, normalize_file_patch_operation


def normalized(*operations):
    return [normalize_file_patch_operation(operation) for operation in operations]


def apply(path, *operations):
    return _apply_office_patch_sequence_sync(str(path), normalized(*operations))


def image_bytes() -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (80, 50), "#174c46").save(output, "PNG")
    return output.getvalue()


def template(path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0xF2, 0xED, 0xE4)
    textbox = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(4), Inches(0.8))
    paragraph = textbox.text_frame.paragraphs[0]
    paragraph.text = "Template title"
    paragraph.runs[0].font.bold = True
    paragraph.runs[0].hyperlink.address = "https://example.com/template"
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.5), Inches(2), Inches(0.8))
    shape.text = "Styled shape"
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0x17, 0x4C, 0x46)
    table = slide.shapes.add_table(2, 2, Inches(0.7), Inches(2.6), Inches(3), Inches(1.2))
    table.table.cell(0, 0).text, table.table.cell(0, 1).text = "Item", "Amount"
    table.table.cell(1, 0).text, table.table.cell(1, 1).text = "Design", "1200"
    slide.shapes.add_picture(io.BytesIO(image_bytes()), Inches(4.9), Inches(0.5), width=Inches(1.6))
    chart_data = CategoryChartData()
    chart_data.categories = ["Q1", "Q2"]
    chart_data.add_series("Revenue", [10, 20])
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(4.1), Inches(2.0), Inches(4.5), Inches(3.0), chart_data,
    )
    slide.notes_slide.notes_text_frame.text = "Speaker notes stay with the duplicated slide"
    presentation.save(path)


def chart_shape(slide):
    return next(shape for shape in slide.shapes if getattr(shape, "has_chart", False))


def test_slide_duplicate_is_a_public_pptx_capability():
    assert "slide.duplicate" in file_patch_operations("pptx")


def test_template_slide_duplicate_preserves_layout_objects_notes_and_independent_chart(tmp_path):
    path = tmp_path / "template.pptx"
    template(path)
    result = apply(path, {"op": "slide.duplicate", "slide": 1})
    assert not result.get("error"), result
    duplicated_path = tmp_path / "duplicated.pptx"
    duplicated_path.write_bytes(result["_persisted_bytes"])
    presentation = Presentation(duplicated_path)
    assert len(presentation.slides) == 2
    original, duplicated = presentation.slides
    assert original.slide_layout.part is duplicated.slide_layout.part
    assert [shape.shape_type for shape in original.shapes] == [shape.shape_type for shape in duplicated.shapes]
    assert [shape.text for shape in original.shapes if getattr(shape, "has_text_frame", False)] == [
        shape.text for shape in duplicated.shapes if getattr(shape, "has_text_frame", False)
    ]
    assert original.follow_master_background is duplicated.follow_master_background
    assert original._element.cSld.bg.xml == duplicated._element.cSld.bg.xml
    assert duplicated.shapes[0].text_frame.paragraphs[0].runs[0].hyperlink.address == "https://example.com/template"
    assert original.notes_slide.notes_text_frame.text == duplicated.notes_slide.notes_text_frame.text
    assert original.notes_slide.part is not duplicated.notes_slide.part
    original_chart, duplicated_chart = chart_shape(original).chart, chart_shape(duplicated).chart
    assert original_chart.part is not duplicated_chart.part
    assert original_chart.part.chart_workbook.xlsx_part is not duplicated_chart.part.chart_workbook.xlsx_part
    assert original_chart.part.chart_workbook.xlsx_part.blob == duplicated_chart.part.chart_workbook.xlsx_part.blob
    chart_id = chart_shape(duplicated).shape_id

    edited = apply(
        duplicated_path,
        {
            "op": "chart.data",
            "slide": 2,
            "shape_id": chart_id,
            "categories": ["Q1", "Q2"],
            "series": [{"name": "Revenue", "values": [99, 101]}],
        },
    )
    assert not edited.get("error"), edited
    final = Presentation(io.BytesIO(edited["_persisted_bytes"]))
    assert b">99<" not in chart_shape(final.slides[0]).chart.part.blob
    assert b">99<" in chart_shape(final.slides[1]).chart.part.blob
    with ZipFile(io.BytesIO(edited["_persisted_bytes"])) as archive:
        names = archive.namelist()
        assert len([name for name in names if name.startswith("ppt/charts/chart") and name.endswith(".xml")]) == 2
        assert len([name for name in names if name.startswith("ppt/embeddings/") and name.endswith(".xlsx")]) == 2
        assert len([name for name in names if name.startswith("ppt/notesSlides/notesSlide") and name.endswith(".xml")]) == 2


def test_generation_can_duplicate_then_adjust_template_slide_content():
    result = _generate_office_operations_sync(
        "pptx",
        normalized(
            {"op": "slide.insert", "index": 0, "layout_index": 6},
            {
                "op": "textbox.insert",
                "slide": 1,
                "text": "Template content",
                "transform": {"x": 72, "y": 72, "width": 400, "height": 80},
            },
            {"op": "slide.duplicate", "slide": 1},
            {"op": "text.set", "slide": 2, "shape_id": 2, "index": 0, "text": "Adjusted copy"},
        ),
    )
    assert not result.get("error"), result
    presentation = Presentation(io.BytesIO(result["_persisted_bytes"]))
    assert [slide.shapes[0].text for slide in presentation.slides] == ["Template content", "Adjusted copy"]


def test_slide_duplicate_selector_failure_is_atomic(tmp_path):
    path = tmp_path / "atomic.pptx"
    template(path)
    before = path.read_bytes()
    result = apply(
        path,
        {"op": "text.set", "slide": 1, "shape_id": 2, "index": 0, "text": "Rollback"},
        {"op": "slide.duplicate", "slide": 9},
    )
    assert result.get("operation_index") == 1
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before
