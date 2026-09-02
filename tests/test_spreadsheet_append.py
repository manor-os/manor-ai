"""Named-table append and explicit template inheritance use native patch_file."""

from __future__ import annotations

import io
import json

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableFormula, TableStyleInfo

from packages.core.ai.tools import file_tools
from packages.core.services.file_engine_patches import describe_file_structure


def _table_source(path, *, declared_formula=False, header_only=False):
    if path.suffix == ".xlsm":
        from tests.test_file_engine_structured_patch import _source

        _source(path, "xlsm")
        workbook = load_workbook(path, keep_vba=True)
    else:
        workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    for column, name in enumerate(["Item", "Qty", "Price", "Amount"], 3):
        sheet.cell(2, column, name)
    if not header_only:
        for column, value in enumerate(["Before", 2, 10, "=D3*E3+$A$1"], 3):
            sheet.cell(3, column, value)
        sheet["F3"].number_format = '#,##0.00'
        sheet["F3"].font = Font(bold=True, color="FF123456")
        sheet["F3"].fill = PatternFill("solid", fgColor="FFABCDEF")
        sheet["F3"].border = Border(bottom=Side(style="thin"))
        sheet["F3"].alignment = Alignment(horizontal="right")
    sheet["A1"] = 5
    sheet["J20"] = "Unrelated note below the table"
    table = Table(displayName="Orders", ref="C2:F2" if header_only else "C2:F3")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    table._initialise_columns()
    for column, name in zip(table.tableColumns, ["Item", "Qty", "Price", "Amount"]):
        column.name = name
    if declared_formula:
        table.tableColumns[-1].calculatedColumnFormula = TableFormula(attr_text="[@Qty]*[@Price]")
    sheet.add_table(table)
    workbook.save(path)
    workbook.close()
    if workbook.vba_archive is not None:
        workbook.vba_archive.close()


def _patch(path, *operations):
    return file_tools._apply_office_patch_sequence_sync(str(path), list(operations))


def _open(result):
    assert result.get("patched") is True, result
    return load_workbook(io.BytesIO(result["_persisted_bytes"]), rich_text=True)


@pytest.mark.parametrize("input_kind", ["row", "values"])
@pytest.mark.parametrize("declared_formula", [False, True])
def test_table_append_extends_only_named_table_and_inherits_formula_style(tmp_path, input_kind, declared_formula):
    path = tmp_path / "table.xlsx"
    _table_source(path, declared_formula=declared_formula)
    original = path.read_bytes()
    arguments = {"row": {"Item": "After", "Qty": 3, "Price": 12}} if input_kind == "row" else {"values": ["After", 3, 12]}

    result = _patch(path, {"operation": "append_row", "sheet": "Sales", "table": "Orders", **arguments})

    workbook = _open(result)
    sheet = workbook["Sales"]
    assert sheet.tables["Orders"].ref == "C2:F4"
    assert sheet.tables["Orders"].autoFilter.ref == "C2:F4"
    assert [sheet.cell(4, col).value for col in (3, 4, 5)] == ["After", 3, 12]
    assert sheet["F4"].value == ("=[@Qty]*[@Price]" if declared_formula else "=D4*E4+$A$1")
    assert sheet["F4"]._style == sheet["F3"]._style
    assert sheet["C3"].value == "Before"
    assert sheet["J20"].value == "Unrelated note below the table"
    assert sheet.tables["Orders"].tableStyleInfo.name == "TableStyleMedium2"
    applied = result["operation_results"][0]
    assert applied["row_number"] == 4 and applied["table"] == "Orders"
    assert applied["table_ref"] == "C2:F4"
    assert applied["updated_cells"]["F4"]["new"] == sheet["F4"].value
    assert path.read_bytes() == original
    workbook.close()


@pytest.mark.parametrize("override", [None, 0, "=D4*E4*2"])
def test_explicit_table_formula_override_including_null_is_not_overwritten(tmp_path, override):
    path = tmp_path / "table.xlsx"
    _table_source(path, declared_formula=True)
    workbook = _open(_patch(path, {
        "operation": "append_row", "table": "Orders", "row": {"Item": "After", "Amount": override},
    }))
    assert workbook.active["F4"].value == override
    assert workbook.active["D4"].value is None  # Do not copy old literal values.
    assert workbook.active["E4"].value is None
    workbook.close()


def test_header_only_table_uses_declared_calculated_column_without_copying_headers(tmp_path):
    path = tmp_path / "empty-table.xlsx"
    _table_source(path, declared_formula=True, header_only=True)
    workbook = _open(_patch(path, {"operation": "append_row", "table": "Orders", "values": ["First", 2, 5]}))
    assert workbook.active.tables["Orders"].ref == "C2:F3"
    assert workbook.active["C3"].value == "First"
    assert workbook.active["F3"].value == "=[@Qty]*[@Price]"
    workbook.close()


@pytest.mark.parametrize("obstacle", ["value", "comment", "merged", "other-table", "totals", "external"])
def test_table_append_rejects_unsafe_expansion_without_mutating_source(tmp_path, obstacle):
    from openpyxl.comments import Comment

    path = tmp_path / "blocked.xlsx"
    _table_source(path)
    workbook = load_workbook(path)
    sheet = workbook.active
    if obstacle == "value":
        sheet["E4"] = "Keep this note"
    elif obstacle == "comment":
        sheet["E4"].comment = Comment("Keep this comment", "User")
    elif obstacle == "merged":
        sheet.merge_cells("C4:D4")
    elif obstacle == "other-table":
        other = Table(displayName="Other", ref="C4:F5", headerRowCount=0)
        sheet.add_table(other)
    elif obstacle == "totals":
        sheet.tables["Orders"].totalsRowCount = 1
    else:
        sheet.tables["Orders"].connectionId = 1
    workbook.save(path)
    workbook.close()
    original = path.read_bytes()

    result = _patch(path, {"operation": "append_row", "table": "Orders", "values": ["After", 3, 12]})

    assert result.get("error"), result
    assert result["operation_index"] == 0
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == original


@pytest.mark.parametrize("arguments", [
    {"table": "Missing"}, {"table": ""}, {"table": None}, {"table": 12},
    {"table": "Orders", "inherit_from_row": 3}, {"table": "Orders", "header_row": 1},
])
def test_invalid_table_selector_is_not_ignored(tmp_path, arguments):
    path = tmp_path / "table.xlsx"
    _table_source(path)
    result = _patch(path, {"operation": "append_row", "values": ["After", 3, 12], **arguments})
    assert result.get("error"), result
    assert result["operation_index"] == 0


def test_table_values_must_fit_table_width(tmp_path):
    path = tmp_path / "table.xlsx"
    _table_source(path)
    result = _patch(path, {"operation": "append_row", "table": "Orders", "values": [1, 2, 3, 4, 5]})
    assert result.get("error"), result


def test_two_appends_in_one_batch_use_updated_table_range(tmp_path):
    path = tmp_path / "table.xlsx"
    _table_source(path)
    result = _patch(path,
        {"operation": "append_row", "table": "Orders", "values": ["Second", 3, 4]},
        {"operation": "append_row", "table": "orders", "values": ["Third", 4, 5]},
    )
    workbook = _open(result)
    assert workbook.active.tables["Orders"].ref == "C2:F5"
    assert workbook.active["F4"].value == "=D4*E4+$A$1"
    assert workbook.active["F5"].value == "=D5*E5+$A$1"
    assert result["operations_applied"] == 2
    workbook.close()


def test_table_selector_is_not_silently_ignored_on_another_operation(tmp_path):
    path = tmp_path / "table.xlsx"
    _table_source(path)
    result = _patch(path, {"operation": "set_cell", "table": "Orders", "cell": "A1", "value": 99})
    assert result.get("error"), result


def test_plain_append_inherits_styles_and_translates_only_formulas_when_requested(tmp_path):
    path = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Item", "Qty", "Amount", "Note"])
    sheet.append(["Old", 3, '=B2*$F$1+$B2+B$2+SUM(B2:B2)+\'Rates 2026\'!B2', "Do not copy"])
    sheet["F1"] = 4
    sheet["B2"].number_format = "0.00"
    sheet["C2"].font = Font(bold=True)
    workbook.save(path)
    workbook.close()

    result = _patch(path, {"operation": "append_row", "inherit_from_row": 2, "row": {"Item": "New", "Qty": 5}})

    edited = _open(result)
    assert edited.active["A3"].value == "New" and edited.active["B3"].value == 5
    assert edited.active["C3"].value == '=B3*$F$1+$B3+B$2+SUM(B3:B3)+\'Rates 2026\'!B3'
    assert edited.active["C3"].font.bold is True
    assert edited.active["B3"].number_format == "0.00"
    assert edited.active["D3"].value is None
    assert result["operation_results"][0]["inherited_from_row"] == 2
    edited.close()


@pytest.mark.parametrize("source_row", [0, -1, True, "3", 3.5, None, 100])
def test_invalid_template_row_is_not_ignored(tmp_path, source_row):
    path = tmp_path / "table.xlsx"
    _table_source(path)
    result = _patch(path, {"operation": "append_row", "inherit_from_row": source_row, "values": ["New"]})
    assert result.get("error"), result


def test_append_without_new_selectors_retains_plain_worksheet_semantics(tmp_path):
    path = tmp_path / "table.xlsx"
    _table_source(path)
    workbook = _open(_patch(path, {"operation": "append_row", "values": ["After", 3, 12]}))
    assert workbook.active.tables["Orders"].ref == "C2:F3"
    assert workbook.active["A21"].value == "After"
    assert workbook.active["F21"].value is None
    workbook.close()


def test_spreadsheet_structure_exposes_real_sheet_table_and_column_selectors(tmp_path):
    path = tmp_path / "table.xlsx"
    _table_source(path)
    structure = describe_file_structure(str(path))
    assert structure["sheet_count"] == 1
    assert structure["sheets"][0]["name"] == "Sales"
    table = structure["sheets"][0]["tables"][0]
    assert table["name"] == "Orders" and table["ref"] == "C2:F3"
    assert table["columns"] == ["Item", "Qty", "Price", "Amount"]
    assert table["has_totals"] is False
    assert structure["truncated"] is False
    json.dumps(structure)


@pytest.mark.parametrize("nested", [False, True])
def test_table_append_extends_sort_state_and_preserves_sort_options(tmp_path, nested):
    from openpyxl.worksheet.filters import SortCondition, SortState

    path = tmp_path / "sorted.xlsx"
    _table_source(path)
    workbook = load_workbook(path)
    table = workbook.active.tables["Orders"]
    owner = table.autoFilter if nested else table
    owner.sortState = SortState(ref="C3:F3", caseSensitive=True, sortCondition=[
        SortCondition(ref="D3:D3", descending=True),
    ])
    workbook.save(path)
    workbook.close()

    edited = _open(_patch(path, {"operation": "append_row", "table": "Orders", "values": ["Next", 5, 2]}))
    table = edited.active.tables["Orders"]
    state = (table.autoFilter if nested else table).sortState
    assert state.ref == "C3:F4" and state.caseSensitive is True
    assert state.sortCondition[0].ref == "D3:D4" and state.sortCondition[0].descending is True
    edited.close()


@pytest.mark.parametrize(("formula", "expected"), [
    ("='Sales'!TaxRate*D3", "='Sales'!TaxRate*D4"),
    ("='Sales!FY'!D3", "='Sales!FY'!D4"),
    ("=ZZZ+XFE1+A1048577+D3", "=ZZZ+XFE1+A1048577+D4"),
    ("=SUM(D3:E3)+SUM(D:D)+SUM(3:3)", "=SUM(D4:E4)+SUM(D:D)+SUM(4:4)"),
    ("=SUM(Orders[[#This Row],[Qty]:[Price]])", "=SUM(Orders[[#This Row],[Qty]:[Price]])"),
    ("=@D3*E3", "=@D4*E4"),
])
def test_inherited_formula_preserves_names_and_translates_references(tmp_path, formula, expected):
    path = tmp_path / "formula.xlsx"
    _table_source(path)
    workbook = load_workbook(path)
    workbook.active["F3"] = formula
    workbook.save(path)
    workbook.close()
    edited = _open(_patch(path, {"operation": "append_row", "table": "Orders", "values": ["Next", 5, 2]}))
    assert edited.active["F4"].value == expected
    edited.close()


@pytest.mark.parametrize("formula", [
    "=D1048576", "=SUM(1048576:1048576)", "=D3#", "=SUM('Sales'!D3:'Sales'!E3)", "array", "data-table",
])
def test_unsafe_formula_inheritance_rejects_whole_batch(tmp_path, formula):
    from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

    path = tmp_path / "unsafe-formula.xlsx"
    _table_source(path)
    workbook = load_workbook(path)
    if formula == "array":
        formula = ArrayFormula(ref="F3:F3", text="=D3*E3")
    elif formula == "data-table":
        formula = DataTableFormula(ref="F3:F3", r1="D3")
    workbook.active["F3"] = formula
    workbook.save(path)
    workbook.close()
    original = path.read_bytes()

    result = _patch(path,
        {"operation": "set_cell", "cell": "A1", "value": 9},
        {"operation": "append_row", "table": "Orders", "values": ["Next", 5, 2]},
    )
    assert result.get("error"), result
    assert result["operation_index"] == 1
    assert "_persisted_bytes" not in result and path.read_bytes() == original


@pytest.mark.parametrize("override", [None, 0, "=D21*E21"])
def test_explicit_template_values_override_inheritance_without_copying_literals(tmp_path, override):
    path = tmp_path / "template.xlsx"
    _table_source(path)
    result = _patch(path, {
        "operation": "append_row", "inherit_from_row": 3,
        "values": [None, None, "New", None, None, override],
    })
    edited = _open(result)
    assert edited.active["F21"].value == override
    assert edited.active["D21"].value is None and edited.active["E21"].value is None
    assert edited.active["F21"]._style == edited.active["F3"]._style
    edited.close()


def test_table_selector_never_falls_back_to_a_different_worksheet(tmp_path):
    path = tmp_path / "multiple-sheets.xlsx"
    _table_source(path)
    workbook = load_workbook(path)
    workbook.create_sheet("Other")
    workbook.save(path)
    workbook.close()
    result = _patch(path, {"operation": "append_row", "sheet": "Other", "table": "Orders", "values": ["Next"]})
    assert result.get("error") == "Table not found in worksheet Other: Orders"


@pytest.mark.parametrize("dimension", ["sheets", "tables", "columns"])
def test_spreadsheet_structure_reports_bounded_results(tmp_path, dimension):
    path = tmp_path / "bounded.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    if dimension == "sheets":
        for index in range(200):
            workbook.create_sheet(f"Other{index}")
    elif dimension == "tables":
        for index in range(201):
            sheet.cell(index + 1, 1, f"Header{index}")
            sheet.add_table(Table(displayName=f"Table{index}", ref=f"A{index + 1}:A{index + 1}"))
    else:
        from openpyxl.utils.cell import get_column_letter

        sheet.append([f"Header{index}" for index in range(201)])
        sheet.add_table(Table(displayName="Wide", ref=f"A1:{get_column_letter(201)}1"))
    workbook.save(path)
    workbook.close()
    structure = describe_file_structure(str(path))
    assert structure["truncated"] is True
    if dimension == "sheets":
        assert structure["sheet_count"] == 201 and len(structure["sheets"]) == 200
    elif dimension == "tables":
        assert len(structure["sheets"][0]["tables"]) == 200
    else:
        table = structure["sheets"][0]["tables"][0]
        assert table["column_count"] == 201 and len(table["columns"]) == 200
