"""Word body/header/footer text and tables share generation and patch operations."""
from __future__ import annotations

import io

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools.file_tools import _apply_office_patch_sequence_sync
from packages.core.contracts.file_engine import normalize_file_patch_operation
from packages.core.services.file_engine_patches import describe_file_structure


def normalized(*operations):
    return [normalize_file_patch_operation(operation) for operation in operations]


def test_generation_and_patch_share_header_footer_text_table_and_cell_operations(tmp_path):
    created = _generate_office_operations_sync("docx", normalized(
        {"op": "paragraph.insert", "story": "header", "section_index": 0, "index": 0, "text": "Company"},
        {"op": "paragraph.format", "story": "header", "section_index": 0, "index": 0,
         "format": {"bold": True, "font_color": "235A45", "alignment": "right"}},
        {"op": "text.set", "story": "header", "section_index": 0, "index": 1, "text": "Confidential"},
        {"op": "table.insert", "story": "footer", "section_index": 0, "index": 0,
         "rows": [["Page", "1"], ["Status", "Draft"]], "style": "Table Grid"},
        {"op": "cell.set", "story": "footer", "section_index": 0,
         "table_index": 0, "cell": "B2", "value": "Final"},
        {"op": "cell.format", "story": "footer", "section_index": 0,
         "table_index": 0, "cell": "B2", "format": {"bold": True, "fill_color": "FFF2CC"}},
    ))
    assert not created.get("error"), created
    path = tmp_path / "stories.docx"
    path.write_bytes(created["_persisted_bytes"])
    document = Document(path)
    header = document.sections[0].header
    footer = document.sections[0].footer
    assert [paragraph.text for paragraph in header.paragraphs] == ["Company", "Confidential"]
    assert header.paragraphs[0].alignment == 2
    assert header.paragraphs[0].runs[0].bold is True
    assert header.paragraphs[0].runs[0].font.color.rgb == (35, 90, 69)
    assert len(footer.tables) == 1
    assert footer.tables[0].cell(1, 1).text == "Final"
    assert footer.tables[0].cell(1, 1).paragraphs[0].runs[0].bold is True

    details = describe_file_structure(str(path))
    structures = {item["story"]: item for item in details["story_structures"]}
    assert structures["header"]["section_indices"] == [0]
    assert [item["text"] for item in structures["header"]["paragraphs"]] == ["Company", "Confidential"]
    assert structures["footer"]["table_count"] == 1
    assert {
        key: structures["footer"]["tables"][0][key]
        for key in ("index", "row_count", "column_count")
    } == {"index": 0, "row_count": 2, "column_count": 2}

    patched = _apply_office_patch_sequence_sync(str(path), normalized(
        {"op": "text.set", "story": "header", "section_index": 0, "index": 0, "text": "Updated Company"},
        {"op": "paragraph.delete", "story": "header", "section_index": 0, "index": 1},
        {"op": "paragraph.insert", "story": "footer", "section_index": 0, "index": 1, "text": "Approved"},
        {"op": "table.delete", "story": "footer", "section_index": 0, "table_index": 0},
    ))
    assert not patched.get("error"), patched
    edited = Document(io.BytesIO(patched["_persisted_bytes"]))
    assert [paragraph.text for paragraph in edited.sections[0].header.paragraphs] == ["Updated Company"]
    assert [paragraph.text for paragraph in edited.sections[0].footer.paragraphs] == ["", "Approved"]
    assert len(edited.sections[0].footer.tables) == 0


def test_linked_later_section_story_edit_is_rejected_atomically(tmp_path):
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Shared"
    document.add_section(WD_SECTION.NEW_PAGE)
    assert document.sections[1].header.is_linked_to_previous
    path = tmp_path / "linked.docx"
    document.save(path)
    before = path.read_bytes()
    result = _apply_office_patch_sequence_sync(str(path), normalized(
        {"op": "paragraph.insert", "index": 0, "text": "Must roll back"},
        {"op": "paragraph.insert", "story": "header", "section_index": 1, "index": 0, "text": "Unsafe"},
    ))
    assert result.get("error") and result["operation_index"] == 1
    assert "linked to the previous section" in result["error"]
    assert "_persisted_bytes" not in result and path.read_bytes() == before


def test_header_footer_must_keep_one_paragraph(tmp_path):
    document = Document()
    document.sections[0].footer.paragraphs[0].text = "Keep or clear"
    path = tmp_path / "required-paragraph.docx"
    document.save(path)
    result = _apply_office_patch_sequence_sync(str(path), normalized(
        {"op": "paragraph.delete", "story": "footer", "section_index": 0, "index": 0},
    ))
    assert result.get("error")
    assert "must retain at least one paragraph" in result["error"]


def test_section_insert_is_shared_by_generation_and_patch_with_independent_page_furniture(tmp_path):
    created = _generate_office_operations_sync("docx", normalized(
        {"op": "paragraph.insert", "index": 0, "text": "Portrait section"},
        {
            "op": "section.insert", "index": 1, "start_type": "new_page",
            "inherit_headers_footers": False,
            "format": {
                "width": 792, "height": 612,
                "margin_top": 36, "margin_bottom": 36, "margin_left": 72, "margin_right": 72,
            },
        },
        {"op": "paragraph.insert", "index": 2, "text": "Landscape section"},
        {"op": "text.set", "story": "header", "section_index": 1, "index": 0, "text": "Independent header"},
    ))
    assert not created.get("error"), created
    path = tmp_path / "generated-sections.docx"
    path.write_bytes(created["_persisted_bytes"])
    document = Document(path)
    assert len(document.sections) == 2
    assert document.sections[1].start_type == WD_SECTION.NEW_PAGE
    assert document.sections[1].orientation == WD_ORIENT.LANDSCAPE
    assert document.sections[1].page_width.pt == 792
    assert document.sections[1].page_height.pt == 612
    assert document.sections[1].header.is_linked_to_previous is False
    assert document.sections[1].header.paragraphs[0].text == "Independent header"
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "Portrait section", "", "Landscape section",
    ]

    template = Document()
    template.add_paragraph("Existing")
    template_path = tmp_path / "template.docx"
    template.save(template_path)
    patched = _apply_office_patch_sequence_sync(str(template_path), normalized(
        {"op": "section.insert", "index": 1, "start_type": "continuous"},
        {"op": "paragraph.insert", "index": 2, "text": "Added section"},
    ))
    assert not patched.get("error"), patched
    edited = Document(io.BytesIO(patched["_persisted_bytes"]))
    assert len(edited.sections) == 2
    assert edited.sections[1].start_type == WD_SECTION.CONTINUOUS
    assert edited.sections[1].header.is_linked_to_previous
    assert edited.paragraphs[-1].text == "Added section"


def test_section_insert_rejects_non_append_index_atomically(tmp_path):
    document = Document()
    document.add_paragraph("Keep")
    path = tmp_path / "atomic-section.docx"
    document.save(path)
    before = path.read_bytes()
    result = _apply_office_patch_sequence_sync(str(path), normalized(
        {"op": "paragraph.insert", "index": 1, "text": "Must roll back"},
        {"op": "section.insert", "index": 0, "start_type": "new_page"},
    ))
    assert result.get("error") and result["operation_index"] == 1
    assert "index must equal section_count" in result["error"]
    assert "_persisted_bytes" not in result and path.read_bytes() == before
