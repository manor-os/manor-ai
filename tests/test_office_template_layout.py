"""Template typography and page geometry must be both generatable and patchable."""
from __future__ import annotations

import io
import json

import pytest

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools import file_tools
from packages.core.contracts.file_engine import normalize_file_patch_operation
from packages.core.services.file_engine_patches import describe_file_structure
from tests.test_office_text_set import apply_text, open_paragraphs, text_operation, text_source
from tests.test_file_engine_structured_patch import file_runtime as file_runtime


PARAGRAPH_STYLE = {
    "font_name": "Arial", "font_size": 24, "font_color": "174C46", "bold": True,
    "italic": True, "underline": True, "strike": False, "alignment": "center",
    "space_before": 6, "space_after": 12, "line_spacing": 1.25,
    "indent_left": 18, "indent_right": 9, "first_line_indent": -9,
}


def paragraph_operation(extension, style=None, index=1):
    operation = text_operation(extension, index=index)
    operation.pop("text")
    return {**operation, "op": "paragraph.format", "format": dict(PARAGRAPH_STYLE if style is None else style)}


def page_operation(extension):
    if extension == "docx":
        return {"op": "page.setup", "section_index": 0, "format": {
            "width": 612, "height": 792, "margin_top": 54, "margin_bottom": 54,
            "margin_left": 54, "margin_right": 54, "header_distance": 24, "footer_distance": 24,
        }}
    return {"op": "page.setup", "format": {"width": 960, "height": 540}}


def template_operations(extension):
    """A deliberate title/body hierarchy, not unstyled fixture-only text."""
    if extension == "docx":
        content = [
            {"op": "paragraph.insert", "index": 0, "text": "Service proposal", "style": "Title"},
            {"op": "paragraph.insert", "index": 1, "text": "Scope and delivery"},
            {"op": "paragraph.insert", "index": 2, "text": "A focused engagement with clear deliverables and weekly progress reviews."},
        ]
    else:
        content = [
            {"op": "slide.insert", "index": 0, "layout_index": 6},
            {"op": "textbox.insert", "slide": 1, "text": "Service proposal\nScope and delivery\nA focused engagement with clear deliverables.",
             "transform": {"x": 72, "y": 90, "width": 816, "height": 340}},
        ]
    title = {"font_name": "Arial", "font_size": 36 if extension == "docx" else 44, "font_color": "174C46", "bold": True, "space_after": 28}
    subtitle = {"font_name": "Arial", "font_size": 20 if extension == "docx" else 28, "bold": True, "space_after": 12}
    body = {"font_name": "Arial", "font_size": 12 if extension == "docx" else 22, "line_spacing": 1.25, "space_after": 12}
    return [page_operation(extension), *content,
            paragraph_operation(extension, title, 0), paragraph_operation(extension, subtitle, 1), paragraph_operation(extension, body, 2)]


def assert_paragraph_style(paragraph, extension):
    run = next(run for run in paragraph.runs if run.text)
    assert run.font.name == "Arial"
    assert run.font.size.pt == 24
    assert str(run.font.color.rgb) == "174C46"
    assert run.font.bold is True and run.font.italic is True and run.font.underline is True
    assert int(paragraph.alignment) == 1 if extension == "docx" else int(paragraph.alignment) == 2
    layout = paragraph.paragraph_format if extension == "docx" else paragraph
    assert layout.space_before.pt == 6 and layout.space_after.pt == 12
    assert layout.line_spacing == 1.25
    if extension == "docx":
        assert layout.left_indent.pt == 18 and layout.right_indent.pt == 9 and layout.first_line_indent.pt == -9
        assert run.font.strike is False
    else:
        properties = paragraph._p.pPr
        assert properties.get("marL") == str(18 * 12700)
        assert properties.get("marR") == str(9 * 12700)
        assert properties.get("indent") == str(-9 * 12700)
        assert run._r.rPr.get("strike") == "noStrike"


@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("empty", [False, True])
def test_template_paragraph_format_roundtrips_and_preserves_neighbors(tmp_path, extension, empty):
    path = tmp_path / f"format.{extension}"
    text_source(path)
    original, paragraphs = open_paragraphs(str(path), extension)
    index = 1 if empty else 0
    untouched = [p._p.xml for i, p in enumerate(paragraphs) if i != index]
    result = apply_text(path, [paragraph_operation(extension, index=index), text_operation(extension, "Formatted content", index)])
    assert not result.get("error"), result
    document, paragraphs = open_paragraphs(io.BytesIO(result["_persisted_bytes"]), extension)
    assert_paragraph_style(paragraphs[index], extension)
    assert [p._p.xml for i, p in enumerate(paragraphs) if i != index] == untouched
    assert paragraphs[index].text == "Formatted content"
    if extension == "pptx":
        assert document.slides[0].shapes[1]._element.xml == original.slides[0].shapes[1]._element.xml


@pytest.mark.parametrize("extension", ["docx", "pptx"])
def test_generate_template_then_restyle_through_same_operations(tmp_path, extension):
    operations = template_operations(extension)
    result = _generate_office_operations_sync(extension, [normalize_file_patch_operation(op) for op in operations])
    assert not result.get("error"), result
    path = tmp_path / f"proposal.{extension}"
    path.write_bytes(result["_persisted_bytes"])
    before = path.read_bytes()
    structure = describe_file_structure(str(path))
    if extension == "docx":
        assert "Title" in structure["paragraph_styles"]
        assert structure["sections"][0]["margin_left"] == 54
        assert structure["sections"][0]["width"] == 612
    else:
        assert structure["page_size"] == {"width": 960, "height": 540}
    patched = apply_text(path, [paragraph_operation(extension), text_operation(extension, "Revised scope")])
    assert not patched.get("error"), patched
    document, paragraphs = open_paragraphs(io.BytesIO(patched["_persisted_bytes"]), extension)
    assert paragraphs[0].text == "Service proposal"
    assert paragraphs[1].text == "Revised scope"
    assert_paragraph_style(paragraphs[1], extension)
    assert path.read_bytes() == before


@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("invalid", [{}, {"fill_color": "112233"}, {"alignment": "diagonal"},
    {"space_before": -1}, {"space_after": float("inf")}, {"space_after": 0.001},
    {"line_spacing": True}, {"line_spacing": 0}, {"font_size": 12.345}, {"underline": "yes"},
    {"strike": None}, {"indent_left": -1}, {"first_line_indent": -10001}, {"font_color": "#abc"}])
def test_invalid_template_format_returns_no_partial_bytes(tmp_path, extension, invalid):
    path = tmp_path / f"invalid.{extension}"
    text_source(path)
    before = path.read_bytes()
    result = apply_text(path, [text_operation(extension), paragraph_operation(extension, invalid)])
    assert result.get("error"), result
    assert result["operation_index"] == 1
    assert "_persisted_bytes" not in result and path.read_bytes() == before


def test_word_named_style_and_pagination_flags_do_not_change_runs(tmp_path):
    path = tmp_path / "style.docx"
    text_source(path)
    _, paragraphs = open_paragraphs(str(path), "docx")
    runs = [r._r.xml for r in paragraphs[0].runs]
    operation = paragraph_operation("docx", {"keep_with_next": True, "keep_together": True, "page_break_before": False, "widow_control": True}, 0)
    operation["style"] = "Heading 1"
    result = apply_text(path, [operation])
    assert not result.get("error"), result
    _, paragraphs = open_paragraphs(io.BytesIO(result["_persisted_bytes"]), "docx")
    paragraph = paragraphs[0]
    assert paragraph.style.name == "Heading 1"
    assert paragraph.paragraph_format.keep_with_next is True
    assert paragraph.paragraph_format.keep_together is True
    assert paragraph.paragraph_format.page_break_before is False
    assert paragraph.paragraph_format.widow_control is True
    assert [r._r.xml for r in paragraph.runs] == runs
    for name in ["does not exist", "Default Paragraph Font"]:
        assert apply_text(path, [{**operation, "style": name}]).get("error")


def test_ppt_paragraph_spacing_keeps_hundredth_point_precision(tmp_path):
    path = tmp_path / "precise.pptx"
    text_source(path)
    result = apply_text(path, [paragraph_operation("pptx", {"space_before": 20.03, "space_after": 20.07, "level": 2})])
    assert not result.get("error"), result
    _, paragraphs = open_paragraphs(io.BytesIO(result["_persisted_bytes"]), "pptx")
    assert paragraphs[1].space_before.pt == 20.03
    assert paragraphs[1].space_after.pt == 20.07
    assert paragraphs[1].level == 2


def test_word_page_setup_changes_only_selected_section(tmp_path):
    from docx.enum.section import WD_SECTION, WD_ORIENT
    path = tmp_path / "sections.docx"
    text_source(path)
    document, _ = open_paragraphs(str(path), "docx")
    document.add_section(WD_SECTION.NEW_PAGE)
    document.sections[0].header.paragraphs[0].text = "Keep header"
    document.save(path)
    original = document.sections[0]._sectPr.xml
    operation = {"op": "page.setup", "section_index": 1, "format": {"width": 792, "height": 612, "margin_left": 36}}
    result = apply_text(path, [operation])
    assert not result.get("error"), result
    document, _ = open_paragraphs(io.BytesIO(result["_persisted_bytes"]), "docx")
    assert document.sections[0]._sectPr.xml == original
    assert document.sections[1].orientation == WD_ORIENT.LANDSCAPE
    assert document.sections[1].left_margin.pt == 36
    assert document.sections[1].header.paragraphs[0].text == "Keep header"


@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("invalid", [{"width": 0}, {"height": True}, {"width": float("nan")},
    {"width": 612.001}, {"margin_left": 700}, {"orientation": "landscape"}, {}])
def test_invalid_page_setup_is_atomic(tmp_path, extension, invalid):
    path = tmp_path / f"page.{extension}"
    text_source(path)
    before = path.read_bytes()
    result = apply_text(path, [text_operation(extension), {**page_operation(extension), "format": invalid}])
    assert result.get("error"), result
    assert result["operation_index"] == 1
    assert "_persisted_bytes" not in result and path.read_bytes() == before


@pytest.mark.parametrize("extension", ["docx", "pptx"])
def test_page_setup_preserves_text_and_ppt_shape_geometry(tmp_path, extension):
    path = tmp_path / f"page.{extension}"
    text_source(path)
    original, paragraphs = open_paragraphs(str(path), extension)
    source_xml = [p._p.xml for p in paragraphs]
    result = apply_text(path, [page_operation(extension)])
    assert not result.get("error"), result
    document, paragraphs = open_paragraphs(io.BytesIO(result["_persisted_bytes"]), extension)
    assert [p._p.xml for p in paragraphs] == source_xml
    if extension == "pptx":
        assert document.slides[0]._element.xml == original.slides[0]._element.xml
        assert document.slide_width.pt == 960 and document.slide_height.pt == 540


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("kind", ["paragraph", "page"])
@pytest.mark.parametrize("guard", ["resource", "approval", "stale"])
async def test_template_layout_keeps_native_guards(file_runtime, monkeypatch, extension, kind, guard):
    root, calls = file_runtime
    path = root / f"guarded.{extension}"
    text_source(path)
    before = path.read_bytes()
    async def reject(**_kwargs):
        return json.dumps({"error": f"{guard}_denied"})
    if guard != "stale":
        monkeypatch.setattr(file_tools, "runtime_guard_file_resource_access" if guard == "resource" else "runtime_guard_file_mutation", reject)
    result = json.loads(await file_tools._patch_file(
        "entity", path=path.name, operations=[paragraph_operation(extension) if kind == "paragraph" else page_operation(extension)],
        expected_sha256="0" * 64 if guard == "stale" else file_tools._file_meta(str(path))["source_sha256"],
    ))
    assert result.get("error"), result
    assert path.read_bytes() == before
    assert not any("commit" in call for call in calls)
