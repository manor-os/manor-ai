"""Word text-box contents use the shared generation and patch operation engine."""
from __future__ import annotations

import io

from docx import Document
from docx.oxml import parse_xml

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools.file_tools import _apply_office_patch_sequence_sync
from packages.core.contracts.file_engine import file_patch_operations, normalize_file_patch_operation
from packages.core.services.file_engine_patches import describe_file_structure


def normalized(*operations):
    return [normalize_file_patch_operation(operation) for operation in operations]


def text_box_template() -> bytes:
    document = Document()
    paragraph = document.add_paragraph("Body before")
    paragraph._p.append(parse_xml("""
      <w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
           xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
           xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
           xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
        <w:drawing>
          <wp:anchor distT="0" distB="0" distL="114300" distR="114300" simplePos="0"
                     relativeHeight="251659264" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">
            <wp:simplePos x="0" y="0"/>
            <wp:positionH relativeFrom="column"><wp:posOffset>914400</wp:posOffset></wp:positionH>
            <wp:positionV relativeFrom="paragraph"><wp:posOffset>457200</wp:posOffset></wp:positionV>
            <wp:extent cx="2743200" cy="1371600"/>
            <wp:effectExtent l="0" t="0" r="0" b="0"/>
            <wp:wrapSquare wrapText="bothSides"/>
            <wp:docPr id="42" name="Template text box" descr="Editable callout"/>
            <wp:cNvGraphicFramePr/>
            <a:graphic>
              <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
                <wps:wsp>
                  <wps:cNvSpPr txBox="1"/>
                  <wps:spPr>
                    <a:xfrm><a:off x="0" y="0"/><a:ext cx="2743200" cy="1371600"/></a:xfrm>
                    <a:prstGeom prst="roundRect"><a:avLst/></a:prstGeom>
                    <a:solidFill><a:srgbClr val="EAF5F0"/></a:solidFill>
                    <a:ln w="25400"><a:solidFill><a:srgbClr val="174C46"/></a:solidFill></a:ln>
                  </wps:spPr>
                  <wps:txbx>
                    <w:txbxContent>
                      <w:p><w:r><w:rPr><w:b/><w:color w:val="174C46"/></w:rPr><w:t>Template title</w:t></w:r></w:p>
                      <w:p><w:r><w:t>Template body</w:t></w:r></w:p>
                    </w:txbxContent>
                  </wps:txbx>
                  <wps:bodyPr wrap="square" lIns="91440" tIns="45720" rIns="91440" bIns="45720" anchor="t"/>
                </wps:wsp>
              </a:graphicData>
            </a:graphic>
          </wp:anchor>
        </w:drawing>
      </w:r>
    """))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def grouped_text_box_template() -> bytes:
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    document = Document()
    paragraph = document.add_paragraph("Body before group")
    group = parse_xml("""
      <w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
           xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
           xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
           xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup"
           xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
        <w:drawing>
          <wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0"
                     relativeHeight="251659264" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">
            <wp:simplePos x="0" y="0"/>
            <wp:positionH relativeFrom="column"><wp:posOffset>914400</wp:posOffset></wp:positionH>
            <wp:positionV relativeFrom="paragraph"><wp:posOffset>457200</wp:posOffset></wp:positionV>
            <wp:extent cx="3657600" cy="1828800"/>
            <wp:effectExtent l="0" t="0" r="0" b="0"/>
            <wp:wrapSquare wrapText="bothSides"/>
            <wp:docPr id="52" name="Grouped text boxes" descr="Editable grouped callouts"/>
            <wp:cNvGraphicFramePr/>
            <a:graphic>
              <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup">
                <wpg:wgp>
                  <wpg:cNvGrpSpPr/>
                  <wpg:grpSpPr>
                    <a:xfrm>
                      <a:off x="0" y="0"/><a:ext cx="3657600" cy="1828800"/>
                      <a:chOff x="0" y="0"/><a:chExt cx="3657600" cy="1828800"/>
                    </a:xfrm>
                  </wpg:grpSpPr>
                  <wps:wsp>
                    <wps:cNvSpPr txBox="1"/>
                    <wps:spPr>
                      <a:xfrm><a:off x="0" y="0"/><a:ext cx="1600200" cy="914400"/></a:xfrm>
                      <a:prstGeom prst="roundRect"><a:avLst/></a:prstGeom>
                    </wps:spPr>
                    <wps:txbx><w:txbxContent><w:p><w:hyperlink><w:r><w:t>Grouped first</w:t></w:r></w:hyperlink></w:p></w:txbxContent></wps:txbx>
                    <wps:bodyPr wrap="square"/>
                  </wps:wsp>
                  <wps:wsp>
                    <wps:cNvSpPr txBox="1"/>
                    <wps:spPr>
                      <a:xfrm><a:off x="1828800" y="0"/><a:ext cx="1600200" cy="914400"/></a:xfrm>
                      <a:prstGeom prst="ellipse"><a:avLst/></a:prstGeom>
                    </wps:spPr>
                    <wps:txbx><w:txbxContent><w:p><w:hyperlink><w:r><w:t>Grouped second</w:t></w:r></w:hyperlink></w:p></w:txbxContent></wps:txbx>
                    <wps:bodyPr wrap="square"/>
                  </wps:wsp>
                </wpg:wgp>
              </a:graphicData>
            </a:graphic>
          </wp:anchor>
        </w:drawing>
      </w:r>
    """)
    links = list(group.iter(qn("w:hyperlink")))
    links[0].set(qn("r:id"), document.part.relate_to(
        "https://example.com/removed", RT.HYPERLINK, is_external=True,
    ))
    links[1].set(qn("r:id"), document.part.relate_to(
        "https://example.com/retained", RT.HYPERLINK, is_external=True,
    ))
    paragraph._p.append(group)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def vml_text_box_template() -> bytes:
    document = Document()
    paragraph = document.add_paragraph("Legacy body")
    paragraph._p.append(parse_xml("""
      <w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
           xmlns:v="urn:schemas-microsoft-com:vml">
        <w:pict>
          <v:shape id="LegacyTextBox" title="Legacy note" type="#_x0000_t202"
                   style="position:absolute;margin-left:36pt;margin-top:18pt;width:144pt;height:72pt"
                   filled="t" fillcolor="FFFFCC" stroked="t" strokecolor="333333" strokeweight="1pt">
            <v:textbox inset="5pt,5pt,5pt,5pt">
              <w:txbxContent><w:p><w:r><w:t>Legacy text</w:t></w:r></w:p></w:txbxContent>
            </v:textbox>
          </v:shape>
        </w:pict>
      </w:r>
    """))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def grouped_vml_text_box_template() -> bytes:
    document = Document()
    paragraph = document.add_paragraph("Legacy group")
    paragraph._p.append(parse_xml("""
      <w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
           xmlns:v="urn:schemas-microsoft-com:vml">
        <w:pict>
          <v:group id="LegacyGroup" style="position:absolute;width:300pt;height:100pt" coordsize="300,100">
            <v:shape id="LegacyFirst" type="#_x0000_t202"
                     style="position:absolute;left:0;top:0;width:140pt;height:80pt">
              <v:textbox><w:txbxContent><w:p><w:r><w:t>Legacy grouped first</w:t></w:r></w:p></w:txbxContent></v:textbox>
            </v:shape>
            <v:shape id="LegacySecond" type="#_x0000_t202"
                     style="position:absolute;left:150pt;top:0;width:140pt;height:80pt">
              <v:textbox><w:txbxContent><w:p><w:r><w:t>Legacy grouped second</w:t></w:r></w:p></w:txbxContent></v:textbox>
            </v:shape>
          </v:group>
        </w:pict>
      </w:r>
    """))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def edits():
    return normalized(
        {"op": "text.set", "text_box_index": 0, "index": 0, "text": "Updated title"},
        {"op": "paragraph.format", "text_box_index": 0, "index": 0,
         "format": {"font_size": 24, "font_color": "235A45", "alignment": "center"}},
        {"op": "paragraph.insert", "text_box_index": 0, "index": 2, "text": "Added detail"},
        {"op": "table.insert", "text_box_index": 0, "index": 2,
         "rows": [["Metric", "Value"], ["Status", "Draft"]], "style": "Table Grid"},
        {"op": "cell.set", "text_box_index": 0, "table_index": 0, "cell": "B2", "value": "Final"},
        {"op": "cell.format", "text_box_index": 0, "table_index": 0, "cell": "B2",
         "format": {"bold": True, "fill_color": "FFF2CC"}},
    )


def test_template_generation_and_patch_share_precise_text_box_content_operations(tmp_path):
    template = text_box_template()
    generated = _generate_office_operations_sync("docx", edits(), template_bytes=template)
    assert not generated.get("error"), generated
    assert generated["_persisted_bytes"] != template

    path = tmp_path / "text-box.docx"
    path.write_bytes(template)
    patched = _apply_office_patch_sequence_sync(str(path), edits())
    assert not patched.get("error"), patched
    assert path.read_bytes() == template

    for result in (generated, patched):
        output_path = tmp_path / f"result-{id(result)}.docx"
        output_path.write_bytes(result["_persisted_bytes"])
        structure = describe_file_structure(str(output_path))
        assert structure["text_box_count"] == 1
        text_box = structure["text_boxes"][0]
        assert text_box["story"] == "body" and text_box["preset"] == "roundRect"
        assert text_box["width"] == 216 and text_box["height"] == 108
        assert text_box["position_x"] == 72 and text_box["position_y"] == 36
        assert [paragraph["text"] for paragraph in text_box["paragraphs"]] == [
            "Updated title", "Template body", "Added detail",
        ]
        assert {
            key: text_box["tables"][0][key]
            for key in ("index", "row_count", "column_count")
        } == {"index": 0, "row_count": 2, "column_count": 2}
        xml = Document(output_path).element.xml
        assert "Updated title" in xml and "Final" in xml and "FFF2CC" in xml
        assert "Editable callout" in xml and "roundRect" in xml


def test_text_box_selector_errors_are_atomic(tmp_path):
    path = tmp_path / "atomic-text-box.docx"
    path.write_bytes(text_box_template())
    before = path.read_bytes()
    result = _apply_office_patch_sequence_sync(str(path), normalized(
        {"op": "text.set", "text_box_index": 0, "index": 1, "text": "Must roll back"},
        {"op": "text.set", "text_box_index": 9, "index": 0, "text": "Invalid"},
    ))
    assert result.get("error") and result["operation_index"] == 1
    assert "text_box_index" in result["error"]
    assert "_persisted_bytes" not in result and path.read_bytes() == before


def test_text_box_must_retain_one_paragraph(tmp_path):
    path = tmp_path / "required-text-box-paragraph.docx"
    path.write_bytes(text_box_template())
    result = _apply_office_patch_sequence_sync(str(path), normalized(
        {"op": "paragraph.delete", "text_box_index": 0, "index": 1},
        {"op": "paragraph.delete", "text_box_index": 0, "index": 0},
    ))
    assert result.get("error") and result["operation_index"] == 1
    assert "must retain at least one paragraph" in result["error"]


def test_word_text_box_generation_has_full_shared_create_transform_format_and_delete(tmp_path):
    expected = {"textbox.insert", "shape.transform", "shape.format", "shape.delete"}
    assert expected <= set(file_patch_operations("docx"))
    operations = normalized(
        {
            "op": "textbox.insert", "index": 0, "text": "Generated callout", "preset": "roundRect",
            "transform": {"x": 72, "y": 36, "width": 240, "height": 96},
            "format": {
                "fill_color": "EAF5F0", "fill_opacity": 0.8,
                "line_color": "174C46", "line_opacity": 0.7, "line_width": 2,
                "margin_left": 10, "margin_right": 11, "margin_top": 6, "margin_bottom": 7,
                "vertical_alignment": "middle", "word_wrap": True, "alt_text": "Editable callout",
                "wrap": "square", "behind_text": False, "allow_overlap": True,
                "layout_in_cell": True, "distance_left": 8, "distance_right": 9,
            },
        },
        {"op": "shape.transform", "text_box_index": 0,
         "transform": {"x": 90, "y": 54, "width": 252, "height": 108, "rotation": 7}},
        {"op": "text.set", "text_box_index": 0, "index": 0, "text": "Edited callout"},
        {"op": "paragraph.insert", "text_box_index": 0, "index": 1, "text": "Editable details"},
        {"op": "textbox.insert", "index": 1, "text": "Temporary", "preset": "ellipse",
         "transform": {"x": 360, "y": 54, "width": 120, "height": 72}},
        {"op": "shape.delete", "text_box_index": 1},
    )
    result = _generate_office_operations_sync("docx", operations)
    assert not result.get("error"), result
    path = tmp_path / "generated-text-box.docx"
    path.write_bytes(result["_persisted_bytes"])
    structure = describe_file_structure(str(path))
    assert structure["text_box_count"] == 1
    text_box = structure["text_boxes"][0]
    assert text_box["grouped"] is False
    assert text_box["preset"] == "roundRect"
    assert (text_box["position_x"], text_box["position_y"]) == (90, 54)
    assert (text_box["width"], text_box["height"]) == (252, 108)
    assert text_box["rotation"] == 7
    assert text_box["format"]["fill_color"] == "EAF5F0"
    assert text_box["format"]["line_color"] == "174C46"
    assert text_box["format"]["vertical_alignment"] == "middle"
    assert text_box["format"]["alt_text"] == "Editable callout"
    assert [item["text"] for item in text_box["paragraphs"]] == ["Edited callout", "Editable details"]
    xml = Document(path).element.xml
    assert "Editable callout" in xml and 'rot="420000"' in xml
    assert 'val="EAF5F0"' in xml and 'val="174C46"' in xml
    assert "Temporary" not in xml


def test_grouped_word_text_box_member_delete_is_shared_by_generation_and_patch(tmp_path):
    template = grouped_text_box_template()
    operation = normalized(
        {"op": "shape.delete", "text_box_index": 0},
        {"op": "shape.transform", "text_box_index": 0,
         "transform": {"x": 150, "y": 12, "width": 132, "height": 84, "rotation": 9}},
    )
    generated = _generate_office_operations_sync("docx", operation, template_bytes=template)
    assert not generated.get("error"), generated

    source = tmp_path / "grouped-template.docx"
    source.write_bytes(template)
    patched = _apply_office_patch_sequence_sync(str(source), operation)
    assert not patched.get("error"), patched
    assert source.read_bytes() == template

    for index, result in enumerate((generated, patched)):
        output = tmp_path / f"grouped-member-{index}.docx"
        output.write_bytes(result["_persisted_bytes"])
        structure = describe_file_structure(str(output))
        assert structure["text_box_count"] == 1
        assert structure["text_boxes"][0]["grouped"] is True
        assert structure["text_boxes"][0]["group_depth"] == 1
        assert structure["text_boxes"][0]["coordinate_space"] == "group_local"
        assert (
            structure["text_boxes"][0]["position_x"],
            structure["text_boxes"][0]["position_y"],
            structure["text_boxes"][0]["width"],
            structure["text_boxes"][0]["height"],
            structure["text_boxes"][0]["rotation"],
        ) == (150, 12, 132, 84, 9)
        assert structure["text_boxes"][0]["paragraphs"] == [
            {"index": 0, "text": "Grouped second"},
        ]
        edited = Document(output)
        xml = edited.element.xml
        assert "Grouped first" not in xml and "Grouped second" in xml
        assert "wordprocessingGroup" in xml and "<wpg:wgp" in xml
        relationship_targets = {relationship.target_ref for relationship in edited.part.rels.values()}
        assert "https://example.com/removed" not in relationship_targets
        assert "https://example.com/retained" in relationship_targets

    remaining = tmp_path / "remaining-group-member.docx"
    remaining.write_bytes(generated["_persisted_bytes"])
    shared_frame_edit = _apply_office_patch_sequence_sync(
        str(remaining), normalized(
            {"op": "shape.format", "text_box_index": 0, "format": {"wrap": "none"}},
        ),
    )
    assert shared_frame_edit.get("error")
    assert "Grouped Word text boxes" in shared_frame_edit["error"]
    deleted_last = _apply_office_patch_sequence_sync(
        str(remaining), normalized({"op": "shape.delete", "text_box_index": 0}),
    )
    assert not deleted_last.get("error"), deleted_last
    final_document = Document(io.BytesIO(deleted_last["_persisted_bytes"]))
    final_path = tmp_path / "deleted-empty-group.docx"
    final_path.write_bytes(deleted_last["_persisted_bytes"])
    assert describe_file_structure(str(final_path))["text_box_count"] == 0
    assert "Grouped second" not in final_document.element.xml
    assert "<wpg:wgp" not in final_document.element.xml
    assert "https://example.com/retained" not in {
        relationship.target_ref for relationship in final_document.part.rels.values()
    }


def test_template_patch_changes_text_box_geometry_and_style_without_replacing_outer_shape(tmp_path):
    path = tmp_path / "template-shape.docx"
    template = text_box_template()
    path.write_bytes(template)
    result = _apply_office_patch_sequence_sync(str(path), normalized(
        {"op": "shape.transform", "text_box_index": 0,
         "transform": {"x": 80, "y": 44, "width": 230, "height": 120, "rotation": -12}},
        {"op": "shape.format", "text_box_index": 0,
         "format": {"preset": "wedgeRoundRectCallout", "fill_color": "DDEBF7",
                    "line_color": None, "margin_top": 9, "vertical_alignment": "bottom",
                    "word_wrap": False, "alt_text": "Updated native shape", "wrap": "top_bottom"}},
    ))
    assert not result.get("error"), result
    assert path.read_bytes() == template
    edited = Document(io.BytesIO(result["_persisted_bytes"]))
    xml = edited.element.xml
    assert "Template title" in xml and "Template body" in xml
    assert "Updated native shape" in xml and "wedgeRoundRectCallout" in xml
    assert 'rot="-720000"' in xml and "wrapTopAndBottom" in xml
    assert "noFill" in xml and 'val="DDEBF7"' in xml


def test_legacy_vml_text_box_content_geometry_and_supported_style_are_editable(tmp_path):
    path = tmp_path / "legacy-vml.docx"
    path.write_bytes(vml_text_box_template())
    result = _apply_office_patch_sequence_sync(str(path), normalized(
        {"op": "text.set", "text_box_index": 0, "index": 0, "text": "Edited legacy text"},
        {"op": "shape.transform", "text_box_index": 0,
         "transform": {"x": 48, "y": 24, "width": 180, "height": 90, "rotation": 5}},
        {"op": "shape.format", "text_box_index": 0,
         "format": {"fill_color": "D9EAD3", "line_color": "38761D", "line_width": 1.5,
                    "margin_left": 6, "margin_top": 7, "margin_right": 8, "margin_bottom": 9,
                    "alt_text": "Updated legacy note"}},
    ))
    assert not result.get("error"), result
    output = tmp_path / "legacy-vml-output.docx"
    output.write_bytes(result["_persisted_bytes"])
    structure = describe_file_structure(str(output))
    assert structure["text_boxes"][0]["paragraphs"][0]["text"] == "Edited legacy text"
    assert structure["text_boxes"][0]["position_x"] == 48
    assert structure["text_boxes"][0]["width"] == 180
    assert structure["text_boxes"][0]["rotation"] == 5
    assert structure["text_boxes"][0]["format"]["alt_text"] == "Updated legacy note"
    xml = Document(output).element.xml
    assert "Updated legacy note" in xml and "D9EAD3" in xml and "38761D" in xml
    assert "strokeweight=\"1.5pt\"" in xml and "rotation:5" in xml


def test_legacy_vml_group_member_delete_preserves_its_sibling(tmp_path):
    path = tmp_path / "legacy-vml-group.docx"
    path.write_bytes(grouped_vml_text_box_template())
    result = _apply_office_patch_sequence_sync(str(path), normalized(
        {"op": "shape.delete", "text_box_index": 0},
        {"op": "shape.transform", "text_box_index": 0,
         "transform": {"x": 160, "y": 5, "width": 130, "height": 70, "rotation": 3}},
    ))
    assert not result.get("error"), result
    output = tmp_path / "legacy-vml-group-output.docx"
    output.write_bytes(result["_persisted_bytes"])
    structure = describe_file_structure(str(output))
    assert structure["text_box_count"] == 1
    assert structure["text_boxes"][0]["grouped"] is True
    assert structure["text_boxes"][0]["coordinate_space"] == "group_local"
    assert (
        structure["text_boxes"][0]["position_x"],
        structure["text_boxes"][0]["position_y"],
        structure["text_boxes"][0]["width"],
        structure["text_boxes"][0]["height"],
        structure["text_boxes"][0]["rotation"],
    ) == (160, 5, 130, 70, 3)
    assert structure["text_boxes"][0]["paragraphs"] == [
        {"index": 0, "text": "Legacy grouped second"},
    ]
    xml = Document(output).element.xml
    assert "Legacy grouped first" not in xml and "Legacy grouped second" in xml
    assert "<v:group" in xml


def test_textbox_insert_targets_header_story_and_returns_package_wide_index(tmp_path):
    result = _generate_office_operations_sync("docx", normalized(
        {"op": "paragraph.insert", "index": 0, "text": "Body remains independent"},
        {"op": "textbox.insert", "story": "header", "section_index": 0, "index": 0,
         "text": "Header callout", "transform": {"x": 360, "y": 12, "width": 144, "height": 36}},
        {"op": "text.set", "text_box_index": 0, "index": 0, "text": "Updated header callout"},
    ))
    assert not result.get("error"), result
    path = tmp_path / "header-text-box.docx"
    path.write_bytes(result["_persisted_bytes"])
    structure = describe_file_structure(str(path))
    assert structure["text_box_count"] == 1
    assert structure["text_boxes"][0]["story"] == "header"
    assert structure["text_boxes"][0]["section_indices"] == [0]
    assert structure["text_boxes"][0]["paragraphs"] == [
        {"index": 0, "text": "Updated header callout"},
    ]
    assert Document(path).paragraphs[0].text == "Body remains independent"
