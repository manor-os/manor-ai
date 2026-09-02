"""Word and PowerPoint tables share native layout operations across generation and patching."""
from __future__ import annotations

import io
import json

import pytest

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools import file_tools
from packages.core.contracts.file_engine import normalize_file_patch_operation
from packages.core.services.file_engine_patches import describe_file_structure
from tests.test_office_table_patch import table_operations
from tests.test_file_engine_structured_patch import file_runtime as file_runtime


def table_format_operation(extension: str, *, alternate: bool = False) -> dict:
    if extension == "docx":
        return {
            "op": "table.format",
            "table_index": 0,
            "format": {
                "style_name": "Table Grid",
                "alignment": "left" if alternate else "center",
                "autofit": alternate,
                "width": 340 if alternate else 360,
                "indent": 12 if alternate else 18,
                "column_widths": {"A": 100 if alternate else 120, "B": 240},
                "row_heights": {"1": 50 if alternate else 24, "2": 100 if alternate else 36},
                "row_height_rules": {"1": "at_least" if alternate else "exact", "2": "at_least"},
                "repeat_header_rows": 0 if alternate else 1,
            },
        }
    return {
        "op": "table.format",
        "slide": 1,
        "shape_id": 2,
        "format": {
            "column_widths": {"A": 100 if alternate else 120, "B": 240},
            "row_heights": {"1": 50 if alternate else 24, "2": 100 if alternate else 36},
            "first_row": not alternate,
            "last_row": alternate,
            "first_column": alternate,
            "last_column": not alternate,
            "banded_rows": alternate,
            "banded_columns": not alternate,
        },
    }


def open_table(data: bytes, extension: str):
    if extension == "docx":
        from docx import Document

        document = Document(io.BytesIO(data))
        return document, document.tables[0]
    from pptx import Presentation

    document = Presentation(io.BytesIO(data))
    return document, document.slides[0].shapes[0].table


def assert_layout(data: bytes, extension: str, *, alternate: bool = False) -> None:
    document, table = open_table(data, extension)
    expected_columns = (100, 240) if alternate else (120, 240)
    expected_rows = (50, 100) if alternate else (24, 36)
    assert tuple(round(column.width.pt) for column in table.columns) == expected_columns
    assert tuple(round(row.height.pt) for row in table.rows) == expected_rows
    if extension == "docx":
        from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
        from docx.oxml.ns import qn

        assert table.style.name == "Table Grid"
        assert table.alignment == (WD_TABLE_ALIGNMENT.LEFT if alternate else WD_TABLE_ALIGNMENT.CENTER)
        assert table.autofit is alternate
        assert table.rows[0].height_rule == (
            WD_ROW_HEIGHT_RULE.AT_LEAST if alternate else WD_ROW_HEIGHT_RULE.EXACTLY
        )
        assert table.rows[1].height_rule == WD_ROW_HEIGHT_RULE.AT_LEAST
        header = table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader"))
        assert (header is not None) is (not alternate)
    else:
        assert table.first_row is (not alternate)
        assert table.last_row is alternate
        assert table.first_col is alternate
        assert table.last_col is (not alternate)
        assert table.horz_banding is alternate
        assert table.vert_banding is (not alternate)
        assert document.slides[0].shapes[0].shape_id == 2


@pytest.mark.parametrize("extension", ["docx", "pptx"])
def test_table_format_is_shared_by_blank_generation_template_generation_and_patch(tmp_path, extension):
    base_operations = [normalize_file_patch_operation(operation) for operation in table_operations(extension)]
    generated = _generate_office_operations_sync(
        extension,
        [*base_operations, normalize_file_patch_operation(table_format_operation(extension))],
    )
    assert not generated.get("error"), generated
    assert_layout(generated["_persisted_bytes"], extension)

    template = _generate_office_operations_sync(extension, base_operations)["_persisted_bytes"]
    templated = _generate_office_operations_sync(
        extension,
        [normalize_file_patch_operation(table_format_operation(extension))],
        template_bytes=template,
    )
    assert not templated.get("error"), templated
    assert_layout(templated["_persisted_bytes"], extension)

    path = tmp_path / f"layout.{extension}"
    path.write_bytes(templated["_persisted_bytes"])
    before = path.read_bytes()
    patched = file_tools._apply_office_patch_sequence_sync(
        str(path),
        [normalize_file_patch_operation(table_format_operation(extension, alternate=True))],
    )
    assert not patched.get("error"), patched
    assert_layout(patched["_persisted_bytes"], extension, alternate=True)
    assert path.read_bytes() == before

    patched_path = tmp_path / f"patched.{extension}"
    patched_path.write_bytes(patched["_persisted_bytes"])
    structure = describe_file_structure(str(patched_path))
    table_structure = structure["tables"][0] if extension == "docx" else structure["shapes"][0]["table"]
    assert tuple(round(value) for value in table_structure["column_widths"]) == (100, 240)
    assert tuple(round(value) for value in table_structure["row_heights"]) == (50, 100)
    if extension == "docx":
        assert table_structure["style_name"] == "Table Grid"
        assert table_structure["alignment"] == "left"
        assert table_structure["repeat_header_rows"] == 0
    else:
        assert table_structure["first_row"] is False
        assert table_structure["banded_rows"] is True


@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize(
    "invalid_format",
    [
        {},
        {"column_widths": {"Z": 12}},
        {"column_widths": {"A": 12, "a": 24}},
        {"row_heights": {"0": 12}},
        {"row_heights": {"1": 12.345}},
        {"unknown": True},
    ],
)
def test_invalid_table_format_is_atomic(tmp_path, extension, invalid_format):
    source = _generate_office_operations_sync(
        extension,
        [normalize_file_patch_operation(operation) for operation in table_operations(extension)],
    )["_persisted_bytes"]
    path = tmp_path / f"invalid.{extension}"
    path.write_bytes(source)
    before = path.read_bytes()
    operation = table_format_operation(extension)
    operation["format"] = invalid_format
    result = file_tools._apply_office_patch_sequence_sync(
        str(path), [normalize_file_patch_operation(operation)],
    )
    assert result.get("error"), result
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before


def test_powerpoint_table_format_requires_a_table_shape(tmp_path):
    from pptx import Presentation
    from pptx.util import Pt

    path = tmp_path / "not-table.pptx"
    document = Presentation()
    slide = document.slides.add_slide(document.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Pt(10), Pt(10), Pt(100), Pt(40))
    document.save(path)
    operation = table_format_operation("pptx")
    operation["shape_id"] = textbox.shape_id
    result = file_tools._apply_office_patch_sequence_sync(
        str(path), [normalize_file_patch_operation(operation)],
    )
    assert "PowerPoint table" in result["error"]
    assert "_persisted_bytes" not in result


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("guard", ["resource", "approval", "stale"])
async def test_table_format_keeps_runtime_guards(file_runtime, monkeypatch, extension, guard):
    root, calls = file_runtime
    source = _generate_office_operations_sync(
        extension,
        [normalize_file_patch_operation(operation) for operation in table_operations(extension)],
    )["_persisted_bytes"]
    path = root / f"guarded-table-format.{extension}"
    path.write_bytes(source)
    before = path.read_bytes()

    async def reject(**_kwargs):
        return json.dumps({"error": f"{guard}_denied"})

    if guard != "stale":
        monkeypatch.setattr(
            file_tools,
            "runtime_guard_file_resource_access" if guard == "resource" else "runtime_guard_file_mutation",
            reject,
        )
    expected_sha256 = "0" * 64 if guard == "stale" else file_tools._file_meta(str(path))["source_sha256"]
    result = json.loads(await file_tools._patch_file(
        "entity",
        path=path.name,
        operations=[table_format_operation(extension)],
        expected_sha256=expected_sha256,
    ))
    assert result.get("error"), result
    assert path.read_bytes() == before
    assert not any("commit" in call for call in calls)
