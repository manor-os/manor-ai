"""Excel merged ranges share the native generation and patch executor."""
from __future__ import annotations

import io
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from openpyxl.worksheet.table import Table
import pytest

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools.file_tools import _apply_office_patch_sequence_sync
from packages.core.contracts.file_engine import file_patch_operations, normalize_file_patch_operation
from packages.core.services.file_engine_patches import describe_file_structure


def normalized(*operations):
    return [normalize_file_patch_operation(operation) for operation in operations]


def apply(path, *operations):
    return _apply_office_patch_sequence_sync(str(path), normalized(*operations))


def test_merge_operations_are_shared_xlsx_xlsm_capabilities():
    assert {"merge.set", "merge.clear"} <= set(file_patch_operations("xlsx"))
    assert {"merge.set", "merge.clear"} <= set(file_patch_operations("xlsm"))


def test_generation_creates_native_merge_with_editable_anchor_and_descriptor(tmp_path):
    result = _generate_office_operations_sync(
        "xlsx",
        normalized(
            {"op": "cell.set", "cell": "A1", "value": "Quarterly report"},
            {"op": "cell.format", "cell": "A1", "format": {"bold": True, "fill_color": "174C46"}},
            {"op": "merge.set", "range": "a1:c2"},
            {"op": "cell.set", "cell": "A1", "value": "Updated report"},
        ),
    )
    assert not result.get("error"), result
    path = tmp_path / "merge.xlsx"
    path.write_bytes(result["_persisted_bytes"])
    workbook = load_workbook(path)
    assert str(workbook.active.merged_cells) == "A1:C2"
    assert workbook.active["A1"].value == "Updated report"
    assert workbook.active["A1"].font.bold is True
    assert workbook.active["A1"].fill.fgColor.rgb == "FF174C46"
    workbook.close()
    structure = describe_file_structure(str(path))
    assert structure["merge_count"] == 1
    assert structure["sheets"][0]["merged_ranges"] == ["A1:C2"]


def test_merge_set_is_idempotent_and_clear_requires_exact_range(tmp_path):
    workbook = Workbook()
    workbook.active.merge_cells("B2:D3")
    path = tmp_path / "clear.xlsx"
    workbook.save(path)
    workbook.close()
    same = apply(path, {"op": "merge.set", "range": "B2:D3"})
    assert not same.get("error"), same
    assert same["operation_results"][0]["updated"] is False
    cleared = apply(path, {"op": "merge.clear", "range": "B2:D3"})
    assert not cleared.get("error"), cleared
    edited = load_workbook(io.BytesIO(cleared["_persisted_bytes"]))
    assert not edited.active.merged_cells.ranges
    edited.active["C3"] = "Now editable"
    edited.close()
    rejected = apply(path, {"op": "merge.clear", "range": "B2:C3"})
    assert "exact existing" in rejected.get("error", "")


@pytest.mark.parametrize("obstacle", ["value", "comment", "hyperlink", "style", "merge", "table"])
def test_merge_set_rejects_data_loss_and_overlap_atomically(tmp_path, obstacle):
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Keep anchor"
    if obstacle == "value":
        sheet["B2"] = "Keep value"
    elif obstacle == "comment":
        sheet["B2"].comment = Comment("Keep comment", "Manor")
    elif obstacle == "hyperlink":
        sheet["B2"].hyperlink = "https://example.com"
    elif obstacle == "style":
        sheet["B2"].fill = PatternFill("solid", fgColor="FF0000")
    elif obstacle == "merge":
        sheet.merge_cells("B2:D2")
    else:
        sheet["A1"], sheet["B1"] = "One", "Two"
        sheet["A2"], sheet["B2"] = 1, 2
        sheet.add_table(Table(displayName="KeepTable", ref="A1:B2"))
    path = tmp_path / f"{obstacle}.xlsx"
    workbook.save(path)
    workbook.close()
    before = path.read_bytes()
    result = apply(
        path,
        {"op": "cell.set", "cell": "A1", "value": "Rollback"},
        {"op": "merge.set", "range": "A1:C3"},
    )
    assert result.get("operation_index") == 1
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before


def test_xlsm_merge_patch_preserves_macro_payload(tmp_path):
    created = _generate_office_operations_sync("xlsx", normalized({"op": "cell.set", "cell": "A1", "value": "Keep"}))
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(created["_persisted_bytes"])) as archive, zipfile.ZipFile(output, "w") as target:
        for item in archive.infolist():
            payload = archive.read(item.filename)
            if item.filename == "[Content_Types].xml":
                payload = payload.replace(
                    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                )
            target.writestr(item, payload)
        target.writestr("xl/vbaProject.bin", b"opaque-macro-payload")
    path = tmp_path / "macro.xlsm"
    path.write_bytes(output.getvalue())
    result = apply(path, {"op": "merge.set", "range": "A1:C2"})
    assert not result.get("error"), result
    with zipfile.ZipFile(io.BytesIO(result["_persisted_bytes"])) as archive:
        assert archive.read("xl/vbaProject.bin") == b"opaque-macro-payload"
