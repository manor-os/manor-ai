"""Native Excel charts use the same generation and patch operation contract."""
from __future__ import annotations

import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.worksheet.datavalidation import DataValidation

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools import file_tools
from packages.core.contracts.file_engine import OfficeChartType, file_patch_operations, normalize_file_patch_operation
from packages.core.services.file_engine_patches import describe_file_structure, hydrate_spreadsheet_chart_styles


def chart_insert(*, chart_type: str = "column", chart_format=None):
    operation = {
        "op": "chart.insert",
        "sheet": "Sheet",
        "chart_type": chart_type,
        "categories": ["Jan", "Feb", "Mar"],
        "series": [
            {"name": "Revenue", "values": [12, 15, 19]},
            {"name": "Cost", "values": [8, 9, 11]},
        ],
        "anchor": "F5",
        "transform": {"width": 520, "height": 280},
    }
    if chart_type in {"pie", "doughnut"}:
        operation["series"] = operation["series"][:1]
    if chart_type.startswith("scatter"):
        operation["categories"] = [1, 2.5, 4]
    if chart_type in {"stock_hlc", "stock_vhlc"}:
        operation["series"] = [
            {"name": "High", "values": [19, 22, 25]},
            {"name": "Low", "values": [10, 12, 14]},
            {"name": "Close", "values": [15, 18, 21]},
        ]
    if chart_type in {"stock_ohlc", "stock_vohlc"}:
        operation["series"] = [
            {"name": "Open", "values": [14, 17, 20]},
            {"name": "High", "values": [19, 22, 25]},
            {"name": "Low", "values": [10, 12, 14]},
            {"name": "Close", "values": [15, 18, 21]},
        ]
    if chart_type in {"stock_vhlc", "stock_vohlc"}:
        operation["series"].insert(0, {"name": "Volume", "values": [1200, 1800, 1500]})
    if chart_format is not None:
        operation["format"] = chart_format
    return operation


def chart_style():
    return {
        "title": "Quarterly performance",
        "style": 10,
        "legend_position": "bottom",
        "legend_include_in_layout": False,
        "show_data_labels": True,
        "show_value": True,
        "category_axis_title": "Month",
        "value_axis_title": "USD millions",
        "value_axis_min": 0,
        "value_axis_max": 30,
        "value_axis_major_unit": 5,
        "show_major_gridlines": True,
        "series_colors": ["336699", "CC6600"],
    }


def apply(path, operations):
    return file_tools._apply_office_patch_sequence_sync(
        str(path), [normalize_file_patch_operation(operation) for operation in operations],
    )


def generate(*operations):
    result = _generate_office_operations_sync(
        "xlsx", [normalize_file_patch_operation(operation) for operation in operations],
    )
    assert not result.get("error"), result
    return result


def workbook_from(result):
    assert result.get("patched") or result.get("generated"), result
    return load_workbook(io.BytesIO(result["_persisted_bytes"]))


def test_chart_operations_and_types_are_one_shared_public_contract():
    expected = {
        "column", "column_stacked", "column_stacked_100",
        "bar", "bar_stacked", "bar_stacked_100",
        "line", "line_markers", "area", "area_stacked", "area_stacked_100",
        "pie", "doughnut",
        "scatter", "scatter_lines", "scatter_lines_markers",
        "scatter_smooth", "scatter_smooth_markers",
        "combo_column_line",
        "stock_hlc", "stock_ohlc", "stock_vhlc", "stock_vohlc",
    }
    assert set(OfficeChartType.values()) == expected
    assert {"chart.insert", "chart.data", "chart.format"} <= set(file_patch_operations("xlsx"))
    assert {"chart.insert", "chart.data", "chart.format"} <= set(file_patch_operations("xlsm"))
    capabilities = file_tools._file_engine_capabilities("xlsx")
    assert set(capabilities["limits"]["office_chart_types"]) == expected
    assert "very-hidden" in capabilities["limits"]["spreadsheet_charts"]


def test_generate_creates_linked_editable_chart_and_discoverable_data(tmp_path):
    result = generate(chart_insert(chart_format=chart_style()))
    path = tmp_path / "chart.xlsx"
    path.write_bytes(result["_persisted_bytes"])
    workbook = load_workbook(path)
    hydrate_spreadsheet_chart_styles(str(path), workbook)
    chart = workbook["Sheet"]._charts[0]
    assert len(chart.ser) == 2
    assert chart.title.tx.rich.p[0].r[0].t == "Quarterly performance"
    assert chart.style == 10 and chart.legend.position == "b" and chart.legend.overlay is True
    assert chart.y_axis.scaling.min == 0 and chart.y_axis.scaling.max == 30
    assert chart.y_axis.majorUnit == 5 and chart.dLbls.showVal is True
    assert [series.graphicalProperties.solidFill.srgbClr for series in chart.ser] == ["336699", "CC6600"]
    data_sheet = workbook["_manor_chart_data"]
    assert data_sheet.sheet_state == "veryHidden"
    assert [data_sheet.cell(row, 1).value for row in range(4, 7)] == ["Jan", "Feb", "Mar"]
    workbook.close()

    structure = describe_file_structure(str(path))
    details = structure["sheets"][0]["charts"][0]
    assert structure["chart_count"] == 1
    assert details["index"] == 0 and details["type"] == "column" and details["anchor"] == "F5"
    assert details["width"] == pytest.approx(520, abs=0.01)
    assert details["height"] == pytest.approx(280, abs=0.01)
    assert details["categories"] == ["Jan", "Feb", "Mar"]
    assert [(series["name"], series["values"]) for series in details["series"]] == [
        ("Revenue", [12, 15, 19]), ("Cost", [8, 9, 11]),
    ]
    assert structure["sheets"][1]["managed_chart_data"] is True


@pytest.mark.parametrize("chart_type", OfficeChartType.values())
def test_every_exposed_chart_type_generates_a_native_excel_chart(chart_type):
    workbook = workbook_from(generate(chart_insert(chart_type=chart_type)))
    chart = workbook.active._charts[0]
    series_count = sum(len(component.ser) for component in chart._charts)
    expected_series = {
        "pie": 1, "doughnut": 1, "stock_hlc": 3, "stock_ohlc": 4,
        "stock_vhlc": 4, "stock_vohlc": 5,
    }.get(chart_type, 2)
    assert series_count == expected_series
    if chart_type == "line_markers":
        assert all(series.marker.symbol == "circle" for series in chart.ser)
    workbook.close()


@pytest.mark.parametrize("chart_type,series_names", [
    ("stock_hlc", ["High", "Low", "Close"]),
    ("stock_ohlc", ["Open", "High", "Low", "Close"]),
    ("stock_vhlc", ["Volume", "High", "Low", "Close"]),
    ("stock_vohlc", ["Volume", "Open", "High", "Low", "Close"]),
])
def test_stock_chart_generation_discovery_and_patch_remain_native(tmp_path, chart_type, series_names):
    created = generate(chart_insert(chart_type=chart_type))
    path = tmp_path / f"{chart_type}.xlsx"
    path.write_bytes(created["_persisted_bytes"])
    details = describe_file_structure(str(path))["sheets"][0]["charts"][0]
    assert details["type"] == chart_type
    assert [series["name"] for series in details["series"]] == series_names
    rejected = apply(path, [{
        "op": "chart.format", "sheet": "Sheet", "chart_index": 0, "format": {"vary_colors": True},
    }])
    assert rejected.get("error") and path.read_bytes() == created["_persisted_bytes"]

    updated_series = [
        {"name": name, "values": [20 + index, 24 + index]}
        for index, name in enumerate(series_names)
    ]
    stock_format = {"title": "Updated stock", "series_colors": ["336699"] * len(series_names)}
    if chart_type.startswith("stock_v"):
        stock_format.update({
            "value_axis_title": "Volume",
            "secondary_value_axis_title": "Price",
            "secondary_value_axis_min": 10,
            "secondary_value_axis_max": 40,
            "secondary_value_axis_major_unit": 5,
            "show_secondary_major_gridlines": False,
        })
    edited = apply(path, [{
        "op": "chart.data", "sheet": "Sheet", "chart_index": 0,
        "categories": ["Apr", "May"], "series": updated_series,
    }, {
        "op": "chart.format", "sheet": "Sheet", "chart_index": 0,
        "format": stock_format,
    }])
    assert not edited.get("error"), edited
    path.write_bytes(edited["_persisted_bytes"])
    details = describe_file_structure(str(path))["sheets"][0]["charts"][0]
    assert details["type"] == chart_type
    assert details["categories"] == ["Apr", "May"]
    assert details["series"][-1]["values"] == [20 + len(series_names) - 1, 24 + len(series_names) - 1]
    assert details["format"]["title"] == "Updated stock"
    if chart_type.startswith("stock_v"):
        assert details["format"]["value_axis_title"] == "Volume"
        assert details["format"]["secondary_value_axis_title"] == "Price"
        assert details["format"]["secondary_value_axis_min"] == 10
        assert details["format"]["secondary_value_axis_max"] == 40
        assert details["format"]["secondary_value_axis_major_unit"] == 5
        assert details["format"]["show_secondary_major_gridlines"] is False
    with zipfile.ZipFile(io.BytesIO(edited["_persisted_bytes"])) as package:
        chart_xml = package.read("xl/charts/chart1.xml").decode()
    assert "<stockChart>" in chart_xml and "<hiLowLines" in chart_xml
    assert ("<upDownBars" in chart_xml) is (chart_type in {"stock_ohlc", "stock_vohlc"})
    assert ("<barChart>" in chart_xml) is chart_type.startswith("stock_v")
    assert chart_xml.count("<valAx>") == (2 if chart_type.startswith("stock_v") else 1)


def test_combo_chart_generation_discovery_and_patch_keep_last_series_as_line(tmp_path):
    created = generate(chart_insert(chart_type="combo_column_line"))
    path = tmp_path / "combo-chart.xlsx"
    path.write_bytes(created["_persisted_bytes"])
    workbook = load_workbook(path)
    chart = workbook["Sheet"]._charts[0]
    assert [type(component).__name__ for component in chart._charts] == ["BarChart", "LineChart"]
    assert [len(component.ser) for component in chart._charts] == [1, 1]
    workbook.close()
    details = describe_file_structure(str(path))["sheets"][0]["charts"][0]
    assert details["type"] == "combo_column_line"
    assert [series["name"] for series in details["series"]] == ["Revenue", "Cost"]

    edited = apply(path, [{
        "op": "chart.data", "sheet": "Sheet", "chart_index": 0,
        "categories": ["Q1", "Q2"],
        "series": [
            {"name": "Revenue", "values": [20, 25]},
            {"name": "Cost", "values": [12, 14]},
            {"name": "Margin", "values": [40, 44]},
        ],
    }, {
        "op": "chart.format", "sheet": "Sheet", "chart_index": 0,
        "format": {
            "series_colors": ["336699", "CC6600", "228844"],
            "show_data_labels": True,
            "show_value": True,
        },
    }])
    assert not edited.get("error"), edited
    workbook = workbook_from(edited)
    chart = workbook["Sheet"]._charts[0]
    assert [len(component.ser) for component in chart._charts] == [2, 1]
    assert [
        series.graphicalProperties.solidFill.srgbClr
        for component in chart._charts for series in component.ser
    ] == ["336699", "CC6600", "228844"]
    assert all(component.dLbls is not None and component.dLbls.showVal for component in chart._charts)
    workbook.close()


@pytest.mark.parametrize("chart_type", [
    "scatter", "scatter_lines", "scatter_lines_markers", "scatter_smooth", "scatter_smooth_markers",
])
def test_scatter_chart_generation_discovery_and_data_patch_share_numeric_x_values(tmp_path, chart_type):
    created = generate(chart_insert(chart_type=chart_type))
    path = tmp_path / f"{chart_type}.xlsx"
    path.write_bytes(created["_persisted_bytes"])
    structure = describe_file_structure(str(path))["sheets"][0]["charts"][0]
    assert structure["type"] == chart_type
    assert structure["categories"] == [1, 2.5, 4]
    assert structure["series"][1]["x_values"] == [1, 2.5, 4]

    patched = apply(path, [{
        "op": "chart.data", "sheet": "Sheet", "chart_index": 0,
        "categories": [0.5, 3],
        "series": [
            {"name": "Observed", "values": [2, 8]},
            {"name": "Forecast", "values": [3, 9]},
        ],
    }])
    assert not patched.get("error"), patched
    path.write_bytes(patched["_persisted_bytes"])
    edited = load_workbook(path)
    chart = edited["Sheet"]._charts[0]
    assert chart.scatterStyle == {
        "scatter": "marker",
        "scatter_lines": "line",
        "scatter_lines_markers": "lineMarker",
        "scatter_smooth": "smooth",
        "scatter_smooth_markers": "smoothMarker",
    }[chart_type]
    assert len(chart.ser) == 2
    edited.close()


def test_scatter_chart_uses_shared_numeric_x_axis_format_in_generation_and_patch(tmp_path):
    created = generate(chart_insert(chart_type="scatter_smooth_markers", chart_format={
        "category_axis_title": "Elapsed seconds",
        "category_axis_min": 0,
        "category_axis_max": 5,
        "category_axis_major_unit": 1,
        "value_axis_title": "Response",
        "trendlines": [{
            "series_index": 0, "type": "linear", "name": "Observed trend",
            "forward": 1, "display_equation": True, "display_r_squared": True,
        }],
        "error_bars": [{
            "series_index": 1, "type": "percentage", "value": 10,
            "side": "plus", "end_style": "no_cap",
        }],
    }))
    path = tmp_path / "formatted-scatter.xlsx"
    path.write_bytes(created["_persisted_bytes"])
    structure = describe_file_structure(str(path))["sheets"][0]["charts"][0]
    assert structure["format"]["category_axis_title"] == "Elapsed seconds"
    assert structure["format"]["category_axis_min"] == 0
    assert structure["format"]["category_axis_max"] == 5
    assert structure["format"]["category_axis_major_unit"] == 1
    assert structure["series"][0]["trendline"] == {
        "type": "linear", "name": "Observed trend", "forward": 1.0,
        "display_equation": True, "display_r_squared": True,
    }
    assert structure["series"][1]["error_bars"] == {
        "type": "percentage", "direction": "y", "side": "plus", "end_style": "no_cap", "value": 10.0,
    }

    patched = apply(path, [{
        "op": "chart.format", "sheet": "Sheet", "chart_index": 0,
        "format": {
            "category_axis_max": 10, "category_axis_major_unit": 2,
            "trendlines": [{"series_index": 0, "type": "polynomial", "order": 3}],
            "error_bars": [{"series_index": 1, "type": "standard_error", "direction": "x"}],
        },
    }])
    assert not patched.get("error"), patched
    workbook = workbook_from(patched)
    chart = workbook["Sheet"]._charts[0]
    assert chart.x_axis.scaling.min == 0 and chart.x_axis.scaling.max == 10
    assert chart.x_axis.majorUnit == 2
    assert chart.ser[0].trendline.trendlineType == "poly" and chart.ser[0].trendline.order == 3
    assert chart.ser[1].errBars.errValType == "stdErr" and chart.ser[1].errBars.errDir == "x"
    workbook.close()


def test_chart_data_retains_native_style_and_chart_move_retains_exact_size(tmp_path):
    initial = generate(chart_insert(chart_format=chart_style()))
    path = tmp_path / "chart.xlsx"
    path.write_bytes(initial["_persisted_bytes"])
    patched = apply(path, [
        {"op": "chart.data", "sheet": "Sheet", "chart_index": 0,
         "categories": ["Q1", "Q2"],
         "series": [{"name": "Revenue", "values": [21, 24]}, {"name": "Cost", "values": [12, 13]}]},
        {"op": "chart.format", "sheet": "Sheet", "chart_index": 0,
         "format": {"anchor": "H7", "title": "Updated performance"}},
    ])
    assert not patched.get("error"), patched
    path.write_bytes(patched["_persisted_bytes"])
    workbook = load_workbook(path)
    hydrate_spreadsheet_chart_styles(str(path), workbook)
    chart = workbook.active._charts[0]
    assert chart.title.tx.rich.p[0].r[0].t == "Updated performance"
    assert chart.style == 10 and chart.legend.position == "b"
    assert [series.graphicalProperties.solidFill.srgbClr for series in chart.ser] == ["336699", "CC6600"]
    assert chart.anchor._from.col == 7 and chart.anchor._from.row == 6
    assert chart.anchor.ext.width == 520 * 12700
    assert chart.anchor.ext.height == 280 * 12700
    data = workbook["_manor_chart_data"]
    assert [data.cell(data.max_row - 1 + offset, 1).value for offset in range(2)] == ["Q1", "Q2"]
    assert data.max_row == 5
    workbook.close()

    path.write_bytes(patched["_persisted_bytes"])
    second_patch = apply(path, [{
        "op": "chart.data", "chart_index": 0,
        "categories": ["H1"],
        "series": [{"name": "Revenue", "values": [30]}, {"name": "Cost", "values": [14]}],
    }])
    workbook = workbook_from(second_patch)
    assert workbook["_manor_chart_data"].max_row == 4
    workbook.close()


def test_pie_category_colors_use_reference_count_and_survive_save(tmp_path):
    initial = generate(chart_insert(chart_type="pie"))
    path = tmp_path / "pie.xlsx"
    path.write_bytes(initial["_persisted_bytes"])
    patched = apply(path, [{"op": "chart.format", "chart_index": 0, "format": {
        "title": "Revenue mix", "show_data_labels": True,
        "show_category_name": True, "show_percentage": True,
        "category_colors": ["336699", "CC6600", "009966"],
    }}])
    assert not patched.get("error"), patched
    workbook = workbook_from(patched)
    chart = workbook.active._charts[0]
    assert [point.graphicalProperties.solidFill.srgbClr for point in chart.ser[0].dPt] == [
        "336699", "CC6600", "009966",
    ]
    workbook.close()


def test_template_chart_patch_preserves_unrelated_chart_validation_formula_and_sheet(tmp_path):
    initial = generate(chart_insert(chart_format=chart_style()))
    path = tmp_path / "template.xlsx"
    path.write_bytes(initial["_persisted_bytes"])
    workbook = load_workbook(path)
    sheet = workbook.active
    sheet["A1"] = "Keep"
    sheet["B1"] = "=1+2"
    validation = DataValidation(type="list", formula1='"Yes,No"')
    sheet.add_data_validation(validation)
    validation.add("C1")
    values = workbook.create_sheet("Source")
    values.append(["Name", "Value"])
    values.append(["Untouched", 7])
    second = BarChart()
    second.style = 12
    second.add_data(Reference(values, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    second.set_categories(Reference(values, min_col=1, min_row=2, max_row=2))
    sheet.add_chart(second, "N5")
    workbook.save(path)
    workbook.close()
    with zipfile.ZipFile(path) as archive:
        before_chart = archive.read("xl/charts/chart2.xml")

    patched = apply(path, [{"op": "chart.data", "sheet": "Sheet", "chart_index": 0,
        "categories": ["Q1", "Q2"], "series": [{"name": "Revenue", "values": [20, 25]}]}])
    assert not patched.get("error"), patched
    with zipfile.ZipFile(io.BytesIO(patched["_persisted_bytes"])) as archive:
        assert archive.read("xl/charts/chart2.xml") == before_chart
    workbook = workbook_from(patched)
    sheet = workbook["Sheet"]
    assert sheet["A1"].value == "Keep" and sheet["B1"].value == "=1+2"
    assert list(sheet.data_validations.dataValidation[0].sqref.ranges)[0].coord == "C1"
    assert len(sheet._charts) == 2 and workbook["Source"]["B2"].value == 7
    assert len(sheet._charts[0].ser) == 1 and len(sheet._charts[1].ser) == 1
    workbook.close()


def test_two_cell_template_chart_moves_without_geometry_conversion_and_resize_is_explicit(tmp_path):
    path = tmp_path / "two-cell.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Name", "Value"])
    sheet.append(["A", 1])
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    chart.anchor = TwoCellAnchor(
        _from=AnchorMarker(col=4, row=3, colOff=1234, rowOff=5678),
        to=AnchorMarker(col=10, row=20, colOff=4321, rowOff=8765),
    )
    sheet.add_chart(chart)
    workbook.save(path)
    workbook.close()

    moved = apply(path, [{"op": "chart.format", "chart_index": 0, "format": {"anchor": "G8"}}])
    workbook = workbook_from(moved)
    anchor = workbook.active._charts[0].anchor
    assert isinstance(anchor, TwoCellAnchor)
    assert (anchor._from.col, anchor._from.row, anchor._from.colOff, anchor._from.rowOff) == (6, 7, 1234, 5678)
    assert (anchor.to.col, anchor.to.row, anchor.to.colOff, anchor.to.rowOff) == (12, 24, 4321, 8765)
    workbook.close()

    rejected = apply(path, [{"op": "chart.format", "chart_index": 0, "format": {"width": 500}}])
    assert "requires both width and height" in rejected.get("error", "")
    converted = apply(path, [{"op": "chart.format", "chart_index": 0,
                              "format": {"width": 500, "height": 260}}])
    workbook = workbook_from(converted)
    assert workbook.active._charts[0].anchor.ext.width == 500 * 12700
    assert workbook.active._charts[0].anchor.ext.height == 260 * 12700
    workbook.close()


@pytest.mark.parametrize("operation", [
    {"op": "chart.insert", "chart_type": "pie", "categories": ["A"],
     "series": [{"name": "One", "values": [1]}, {"name": "Two", "values": [2]}], "anchor": "A1"},
    {"op": "chart.insert", "chart_type": "unknown", "categories": ["A"],
     "series": [{"name": "One", "values": [1]}], "anchor": "A1"},
    {"op": "chart.insert", "chart_type": "column", "categories": ["A", "B"],
     "series": [{"name": "One", "values": [1]}], "anchor": "A1"},
    {"op": "chart.insert", "chart_type": "scatter", "categories": ["A"],
     "series": [{"name": "One", "values": [1]}], "anchor": "A1"},
    {"op": "chart.insert", "chart_type": "combo_column_line", "categories": ["A"],
     "series": [{"name": "Only", "values": [1]}], "anchor": "A1"},
    {"op": "chart.insert", "chart_type": "stock_hlc", "categories": ["A"],
     "series": [{"name": "High", "values": [2]}, {"name": "Low", "values": [1]}], "anchor": "A1"},
    {"op": "chart.insert", "chart_type": "stock_vhlc", "categories": ["A"],
     "series": [{"name": "Volume", "values": [100]}, {"name": "High", "values": [2]},
                {"name": "Low", "values": [1]}], "anchor": "A1"},
    {"op": "chart.data", "chart_index": 0, "categories": ["A"], "series": []},
    {"op": "chart.format", "chart_index": 0, "format": {"style": 49}},
    {"op": "chart.format", "chart_index": 0, "format": {"category_colors": ["112233"]}},
    {"op": "chart.format", "chart_index": 0,
     "format": {"secondary_value_axis_title": "Not available"}},
    {"op": "chart.format", "chart_index": 0,
     "format": {"trendlines": [{"series_index": 5, "type": "linear"}]}},
    {"op": "chart.format", "chart_index": 0,
     "format": {"error_bars": [{"series_index": 0, "type": "percentage"}]}},
])
def test_invalid_chart_operation_aborts_the_whole_batch(tmp_path, operation):
    initial = generate(chart_insert())
    path = tmp_path / "atomic.xlsx"
    path.write_bytes(initial["_persisted_bytes"])
    before = path.read_bytes()
    result = apply(path, [
        {"op": "cell.set", "cell": "A1", "value": "Must roll back"},
        operation,
    ])
    assert result.get("error"), result
    assert result["operation_index"] == 1 and "_persisted_bytes" not in result
    assert path.read_bytes() == before
