from __future__ import annotations

import csv
import hashlib
import io
import json
from types import SimpleNamespace

import pytest

from packages.core.ai.tools import file_tools
from packages.core.services.file_engine_patches import apply_pdf_patch_sequence, apply_text_patch_sequence
from packages.core.contracts.file_engine import KNOWN_FILE_TYPES, file_type_capability


@pytest.fixture
def file_runtime(tmp_path, monkeypatch):
    root = tmp_path / "entity"
    root.mkdir()
    calls = []

    async def allow(**kwargs):
        calls.append(kwargs)

    async def visible(*_args, **_kwargs):
        return set()

    async def commit(**kwargs):
        path = root / kwargs["rel_path"]
        version = kwargs["expected_source_version"]
        if version is not None:
            assert version.matches(path.stat())
        path.write_bytes(kwargs["data"])
        calls.append({"commit": kwargs})
        return (
            str(path),
            SimpleNamespace(synced=True, document_id="same-document", reason=None),
            file_tools._file_meta(str(path)),
        )

    monkeypatch.setattr(file_tools, "_get_entity_root", lambda _entity: str(root))
    monkeypatch.setattr(file_tools, "runtime_guard_file_resource_access", allow)
    monkeypatch.setattr(file_tools, "runtime_guard_file_mutation", allow)
    monkeypatch.setattr(file_tools, "_blocked_doc_paths", visible)
    monkeypatch.setattr(file_tools, "_commit_file_projection", commit)
    return root, calls


def _source(path, extension):
    if extension == "diagram.json":
        from packages.core.ai.tools.generate_file.diagram import _diagram_document_from_prompt

        document = _diagram_document_from_prompt("flowchart LR; A[Start] --> B[Finish]", name="before.diagram.json")
        path.write_text(json.dumps(document))
        return [{"op": "json.replace", "pointer": "/title", "value": "after"}]
    if extension == "json":
        path.write_text('{"nodes":[{"label":"before"}],"keep":true}\n')
        return [{"op": "json.replace", "pointer": "/nodes/0/label", "value": "after"}]
    if extension in {"csv", "tsv"}:
        separator = "\t" if extension == "tsv" else ","
        path.write_bytes(f"name{separator}value\r\nitem{separator}before\r\n".encode())
        return [
            {"op": "cell.set", "cell": "B2", "value": 'after,"quoted"\nline'},
            {"op": "row.append", "values": ["second", 2]},
        ]
    if extension == "docx":
        from docx import Document

        document = Document()
        document.add_paragraph("before").runs[0].bold = True
        document.save(path)
        return [{"op": "paragraph.insert", "index": 0, "text": "after", "style": "Heading 1"}]
    if extension == "pptx":
        from pptx import Presentation
        from pptx.util import Inches

        document = Presentation()
        shape = document.slides.add_slide(document.slide_layouts[6]).shapes.add_textbox(
            Inches(1), Inches(1), Inches(2), Inches(1)
        )
        shape.text = "before"
        shape.text_frame.paragraphs[0].runs[0].font.bold = True
        document.save(path)
        return [
            {
                "op": "shape.transform",
                "slide": 1,
                "shape_id": shape.shape_id,
                "transform": {"x": 144, "width": 216, "rotation": 30},
            }
        ]
    if extension in {"xlsx", "xlsm"}:
        from openpyxl import Workbook
        from openpyxl.styles import Border, Side

        workbook = Workbook()
        workbook.active["A1"] = "=1+2"
        workbook.active["B1"] = "visible workbook content"
        workbook.active["A1"].border = Border(bottom=Side(style="thin"))
        workbook.save(path)
        workbook.close()
        if extension == "xlsm":
            import zipfile

            output = io.BytesIO()
            with zipfile.ZipFile(path) as source, zipfile.ZipFile(output, "w") as target:
                for item in source.infolist():
                    data = source.read(item.filename)
                    if item.filename == "[Content_Types].xml":
                        data = data.replace(
                            b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                            b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                        )
                    target.writestr(item, data)
                target.writestr("xl/vbaProject.bin", b"opaque-macro-payload-must-be-preserved")
            path.write_bytes(output.getvalue())
        return [
            {
                "op": "cell.format",
                "cell": "A1",
                "format": {"bold": True, "font_size": 18, "fill_color": "123abc", "number_format": "0.00"},
            }
        ]
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=300)
    writer.add_metadata({"/Title": "Keep metadata"})
    writer.write(path)
    writer.close()
    return [{"op": "page.rotate", "page": 1, "degrees": 90}]


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["json", "diagram.json", "csv", "tsv", "docx", "pptx", "xlsx", "xlsm", "pdf"])
async def test_native_read_patch_roundtrip_for_new_operations(file_runtime, extension):
    root, calls = file_runtime
    path = root / f"source.{extension}"
    operations = _source(path, extension)
    before = path.read_bytes()
    read = json.loads(await file_tools._read_file("entity", path=path.name, include_structure=True))
    assert read["source_sha256"] == hashlib.sha256(before).hexdigest()
    if extension in {"xlsx", "xlsm"}:
        assert "visible workbook content" in read["content"]
    if extension == "docx":
        assert read["structure"]["paragraphs"][0]["text"] == "before"
    elif extension == "pptx":
        assert read["structure"]["shapes"][0]["shape_id"] == operations[0]["shape_id"]

    result = json.loads(
        await file_tools._patch_file(
            "entity",
            path=path.name,
            operations=operations,
            expected_sha256=read["source_sha256"],
        )
    )
    assert result.get("patched") is True, result
    assert result["document_id"] == "same-document"
    assert result["source_sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert result["source_sha256"] != read["source_sha256"]
    assert len([call for call in calls if "commit" in call]) == 1
    assert any(call.get("tool_name") == "patch_file" and "content_preview" in call for call in calls)
    if extension == "diagram.json":
        document = json.loads(path.read_text())
        assert document["title"] == "after"
        assert document["elements"] == json.loads(before)["elements"]
        assert document["version"] == "editable_diagram_v1"
    elif extension == "json":
        assert json.loads(path.read_text()) == {"nodes": [{"label": "after"}], "keep": True}
    elif extension in {"csv", "tsv"}:
        rows = list(
            csv.reader(
                io.StringIO(path.read_bytes().decode(), newline=""), delimiter="\t" if extension == "tsv" else ","
            )
        )
        assert rows == [["name", "value"], ["item", 'after,"quoted"\nline'], ["second", "2"]]
        assert path.read_bytes().endswith(b"\r\n")
    elif extension == "docx":
        from docx import Document

        document = Document(path)
        assert [p.text for p in document.paragraphs] == ["after", "before"]
        assert document.paragraphs[0].style.name == "Heading 1"
        assert document.paragraphs[1].runs[0].bold is True
    elif extension == "pptx":
        from pptx import Presentation

        shape = Presentation(path).slides[0].shapes[0]
        assert shape.left.pt == 144 and shape.width.pt == 216 and shape.rotation == 30
        assert shape.text == "before" and shape.text_frame.paragraphs[0].runs[0].font.bold is True
    elif extension in {"xlsx", "xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, keep_vba=extension == "xlsm")
        cell = workbook.active["A1"]
        assert cell.value == "=1+2" and cell.font.bold is True and cell.font.sz == 18
        assert cell.fill.fgColor.rgb == "FF123ABC" and cell.number_format == "0.00"
        assert cell.border.bottom.style == "thin"
        workbook.close()
        if workbook.vba_archive is not None:
            assert workbook.vba_archive.read("xl/vbaProject.bin") == b"opaque-macro-payload-must-be-preserved"
            workbook.vba_archive.close()
    else:
        from pypdf import PdfReader

        document = PdfReader(path)
        assert document.pages[0].rotation == 90
        assert document.metadata.title == "Keep metadata"


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["docx", "pptx", "xlsx", "pdf"])
async def test_non_text_changes_invalidate_expected_sha(file_runtime, extension):
    root, _calls = file_runtime
    path = root / f"source.{extension}"
    operations = _source(path, extension)
    source_hash = file_tools._file_meta(str(path))["source_sha256"]
    result = json.loads(
        await file_tools._patch_file("entity", path=path.name, operations=operations, expected_sha256=source_hash)
    )
    assert result.get("patched") is True, result
    current = path.read_bytes()
    stale = json.loads(
        await file_tools._patch_file("entity", path=path.name, operations=operations, expected_sha256=source_hash)
    )
    assert stale["error"] == "source_changed"
    assert path.read_bytes() == current


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["json", "csv", "docx", "pptx", "xlsx", "pdf"])
async def test_new_patch_late_failure_never_commits(file_runtime, extension):
    root, calls = file_runtime
    path = root / f"source.{extension}"
    operations = _source(path, extension)
    original = path.read_bytes()
    invalid = {
        "json": {"op": "json.remove", "pointer": "/missing"},
        "csv": {"op": "cell.set", "cell": "B900", "value": "no"},
        "docx": {"op": "paragraph.insert", "index": 900, "text": "no"},
        "pptx": {"op": "shape.transform", "slide": 900, "shape_id": 1, "transform": {"x": 1}},
        "xlsx": {"op": "cell.format", "cell": "A1", "format": {"bold": "yes"}},
        "pdf": {"op": "page.rotate", "page": 1, "degrees": 45},
    }[extension]
    result = json.loads(await file_tools._patch_file("entity", path=path.name, operations=operations + [invalid]))
    assert "error" in result
    assert result["operation_index"] == len(operations), result
    assert path.read_bytes() == original
    assert not any("commit" in call for call in calls)


@pytest.mark.parametrize("extension", ["yaml", "py", "svg", "drawio", "srt", "scss", "toml", "dockerfile"])
def test_text_patch_preserves_bom_crlf_and_unicode(tmp_path, extension):
    path = tmp_path / f"source.{extension}"
    original = b"\xef\xbb\xbf" + "旧内容\r\nkeep\r\n".encode()
    path.write_bytes(original)
    result = apply_text_patch_sequence(
        str(path), [{"operation": "replace_text", "old_text": "旧内容", "new_text": "新内容"}]
    )
    assert result["_persisted_bytes"] == b"\xef\xbb\xbf" + "新内容\r\nkeep\r\n".encode()
    assert path.read_bytes() == original
    assert file_type_capability(extension)["can_patch"] is True


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["json", "diagram.json", "csv", "tsv", "docx", "pptx", "xlsx", "xlsm", "pdf"])
async def test_new_operations_cannot_bypass_mutation_approval(file_runtime, monkeypatch, extension):
    root, calls = file_runtime
    path = root / f"source.{extension}"
    operations = _source(path, extension)
    original = path.read_bytes()

    async def deny(**kwargs):
        assert kwargs["tool_name"] == "patch_file"
        assert kwargs["content_preview"]["operations"]
        return json.dumps({"status": "needs_approval"})

    monkeypatch.setattr(file_tools, "runtime_guard_file_mutation", deny)
    result = json.loads(await file_tools._patch_file("entity", path=path.name, operations=operations))
    assert result["status"] == "needs_approval"
    assert path.read_bytes() == original
    assert not any("commit" in call for call in calls)


@pytest.mark.asyncio
@pytest.mark.parametrize("selector", [{"table": "Orders"}, {"inherit_from_row": 3}])
@pytest.mark.parametrize("guard", ["resource", "approval", "stale"])
async def test_append_selectors_keep_native_guards(file_runtime, monkeypatch, selector, guard):
    from tests.test_spreadsheet_append import _table_source

    root, calls = file_runtime
    path = root / "table.xlsx"
    _table_source(path)
    original = path.read_bytes()

    async def deny(**_kwargs):
        return json.dumps({"error": "denied", "status": "needs_approval"})

    if guard == "resource":
        monkeypatch.setattr(file_tools, "runtime_guard_file_resource_access", deny)
    elif guard == "approval":
        monkeypatch.setattr(file_tools, "runtime_guard_file_mutation", deny)
    result = json.loads(await file_tools._patch_file(
        "entity", path=path.name,
        operations=[{"op": "row.append", "values": ["Next", 5, 2], **selector}],
        expected_sha256="0" * 64 if guard == "stale" else hashlib.sha256(original).hexdigest(),
    ))
    assert not result.get("patched"), result
    assert result.get("error") or result.get("status") == "needs_approval", result
    assert path.read_bytes() == original
    assert not any("commit" in call for call in calls)


@pytest.mark.asyncio
@pytest.mark.parametrize("selector", [{"table": "Orders"}, {"inherit_from_row": 3}])
@pytest.mark.parametrize("extension", ["csv", "tsv", "xlsx"])
async def test_append_selectors_are_rejected_on_wrong_type_or_operation(file_runtime, selector, extension):
    root, calls = file_runtime
    path = root / f"source.{extension}"
    _source(path, extension)
    original = path.read_bytes()
    operation = {"op": "cell.set", "cell": "A1", "value": "Next"} if extension == "xlsx" else {
        "op": "row.append", "values": ["Next", "Value"],
    }
    result = json.loads(await file_tools._patch_file(
        "entity", path=path.name, operations=[{**operation, **selector}],
    ))
    assert result.get("error"), result
    assert result["operation_index"] == 0
    assert path.read_bytes() == original
    assert not any("content_preview" in call or "commit" in call for call in calls)


@pytest.mark.parametrize("raw", [b"old\xff", b"old\x00binary", b"%PDF-1.3\nold\n%%EOF"])
def test_text_patch_never_lossily_decodes_binary(tmp_path, raw):
    path = tmp_path / "fake.txt"
    path.write_bytes(raw)
    result = apply_text_patch_sequence(str(path), [{"operation": "replace_text", "old_text": "old", "new_text": "new"}])
    assert "error" in result and "_persisted_bytes" not in result
    assert path.read_bytes() == raw


@pytest.mark.parametrize(
    "pointer", ["/missing/label", "/nodes/99", "/nodes/01", "/nodes/-", "/bad~2key", "not/a/pointer"]
)
def test_json_invalid_pointer_fails_without_output(tmp_path, pointer):
    path = tmp_path / "source.json"
    path.write_text('{"nodes":[1]}')
    result = apply_text_patch_sequence(str(path), [{"operation": "json.replace", "pointer": pointer, "value": None}])
    assert "error" in result and "_persisted_bytes" not in result


def test_json_add_remove_escaped_pointer_and_nested_values(tmp_path):
    path = tmp_path / "source.json"
    path.write_text('{"a/b":{"~key":[]},"obsolete":true}')
    result = apply_text_patch_sequence(
        str(path),
        [
            {"operation": "json.add", "pointer": "/a~1b/~0key/-", "value": {"ok": [1, None, True]}},
            {"operation": "json.remove", "pointer": "/obsolete"},
        ],
    )
    assert json.loads(result["_persisted_bytes"]) == {"a/b": {"~key": [{"ok": [1, None, True]}]}}


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["png", "gif", "mp4", "mp3", "zip", "doc", "unknown"])
async def test_unsupported_binary_patch_is_rejected_before_approval(file_runtime, extension):
    root, calls = file_runtime
    path = root / f"source.{extension}"
    original = b"old\x00\xffdata"
    path.write_bytes(original)
    result = json.loads(
        await file_tools._patch_file(
            "entity", path=path.name, operations=[{"op": "text.replace", "old_text": "old", "new_text": "new"}]
        )
    )
    assert result["error"] == "unsupported_binary_edit"
    assert path.read_bytes() == original
    assert not any("content_preview" in call or "commit" in call for call in calls)


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["mp4", "wav", "zip", "unknown"])
async def test_binary_read_returns_metadata_not_replacement_characters(file_runtime, extension):
    root, _calls = file_runtime
    path = root / f"source.{extension}"
    path.write_bytes(b"old\x00\xffdata")
    result = json.loads(await file_tools._read_file("entity", path=path.name))
    assert result["read_mode"] == "metadata" and "content" not in result
    assert result["source_sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()


def test_unknown_types_are_not_advertised_as_generatable():
    for extension in ("unknown", "exe", "arbitrary", "zip"):
        assert file_type_capability(extension)["can_generate"] is False
    macro = file_type_capability("xlsm")
    assert macro["can_generate"] is True
    assert macro["generate_kind"] == "spreadsheet"
    assert macro["can_generate_from_operations"] is False
    assert macro["can_generate_from_template"] is True
    for extension in KNOWN_FILE_TYPES:
        selected = file_type_capability(extension)
        assert selected["known_type"] is True
        assert selected["can_patch"] == bool(selected["operations"])


@pytest.mark.parametrize("protection", ["encrypted", "signed"])
def test_pdf_rotation_rejects_protected_documents(tmp_path, protection):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject

    path = tmp_path / "protected.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=300)
    if protection == "encrypted":
        writer.encrypt("password")
    else:
        writer.root_object[NameObject("/Perms")] = DictionaryObject({NameObject("/DocMDP"): DictionaryObject()})
    writer.write(path)
    writer.close()
    original = path.read_bytes()
    result = apply_pdf_patch_sequence(str(path), [{"operation": "page.rotate", "page": 1, "degrees": 90}])
    assert result["error"] == f"{protection}_pdf_not_editable"
    assert path.read_bytes() == original


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["xlsx", "png", "zip", "mp4", "unknown"])
async def test_document_generator_rejects_text_disguised_as_binary(tmp_path, monkeypatch, extension):
    from packages.core.ai.runtime import file_actions, generated_files

    monkeypatch.setattr(file_actions, "runtime_entity_file_root", lambda _entity: str(tmp_path))
    result = json.loads(
        await generated_files.runtime_generate_document_file(
            entity_id="entity",
            user_id="user",
            conversation_id="",
            name=f"invalid.{extension}",
            content="not a binary file",
            file_type=extension,
        )
    )
    assert result["error"] == "unsupported_document_generation_type", result
    assert not (tmp_path / f"invalid.{extension}").exists()


@pytest.mark.asyncio
async def test_code_generator_rejects_binary_bundle_entry(tmp_path, monkeypatch):
    from packages.core.ai.tools.generate_file import code

    monkeypatch.setattr(code, "runtime_entity_file_root", lambda _entity: str(tmp_path))
    result = json.loads(
        await code.handle_code(
            entity_id="entity",
            user_id="user",
            conversation_id="",
            prompt="",
            name="bundle",
            agent_id=None,
            params={"files": [{"path": "index.html", "content": "valid"}, {"path": "fake.png", "content": "invalid"}]},
            kwargs={},
        )
    )
    assert result["error"] == "unsupported_code_file_type", result
    assert not list(tmp_path.rglob("index.html"))
