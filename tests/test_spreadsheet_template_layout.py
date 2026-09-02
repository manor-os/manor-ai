"""The same native operations create and restyle workbook templates."""
from __future__ import annotations

import io
import json

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools import file_tools
from packages.core.contracts.file_engine import normalize_file_patch_operation
from packages.core.services.file_engine_patches import describe_file_structure
from tests.test_file_engine_structured_patch import file_runtime as file_runtime


def sheet_layout():
    return {"op": "sheet.format", "sheet": "Sheet", "format": {
        "column_widths": {"A": 42, "B": 18, "C": 22}, "row_heights": {"1": 40, "3": 28},
        "freeze_panes": "B4", "show_gridlines": False,
    }}


def print_layout():
    return {"op": "page.setup", "sheet": "Sheet", "format": {
        "orientation": "landscape", "paper_size": "A4", "print_area": "A1:C8",
        "repeat_rows": "1:3", "repeat_columns": "A:A", "fit_width": 1, "fit_height": 0,
        "margin_left": 36, "margin_right": 36, "margin_top": 54, "margin_bottom": 54,
        "header_distance": 18, "footer_distance": 18,
    }}


def invoice_operations():
    return [
        {"op": "row.append", "values": ["Service proposal"]},
        {"op": "row.append", "values": ["Prepared for your team"]},
        {"op": "row.append", "values": ["Description", "Quantity", "Amount"]},
        {"op": "row.append", "values": ["Implementation", 2, "=B4*125"]},
        {"op": "row.append", "values": ["Training", 1, "=B5*100"]},
        {"op": "cell.set", "cell": "C7", "value": "=SUM(C4:C5)"},
        {"op": "cell.format", "cell": "A1", "format": {"font_size": 26, "font_color": "174C46", "bold": True}},
        *[{"op": "cell.format", "cell": cell, "format": {
            "font_color": "FFFFFF", "fill_color": "174C46", "bold": True,
            "horizontal": "left", "vertical": "center", "wrap_text": True,
            "borders": {"bottom": {"style": "medium", "color": "174C46"}},
        }} for cell in ["A3", "B3", "C3"]],
        *[{"op": "cell.format", "cell": cell, "format": {
            "number_format": '"$"#,##0.00', "horizontal": "right",
        }} for cell in ["C4", "C5", "C7"]],
        sheet_layout(), print_layout(), {"op": "sheet.add", "sheet": "Notes"},
        {"op": "cell.set", "sheet": "Notes", "cell": "A1", "value": "Keep notes"},
    ]


def apply_layout(path, operations):
    return file_tools._apply_office_patch_sequence_sync(str(path), [normalize_file_patch_operation(op) for op in operations])


def source(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Keep title"
    sheet["B4"] = "=1+2"
    sheet["A1"].font = Font(name="Arial", italic=True)
    sheet["A1"].alignment = Alignment(wrap_text=True, textRotation=15)
    sheet["A1"].border = Border(left=Side(style="thin", color="112233"))
    workbook.create_sheet("Notes")["A1"] = "Keep notes"
    workbook.save(path)
    workbook.close()


def assert_layout(workbook):
    sheet = workbook["Sheet"]
    assert sheet.column_dimensions["A"].width == 42
    assert sheet.row_dimensions[1].height == 40
    assert sheet.freeze_panes == "B4"
    assert sheet.sheet_view.showGridLines is False
    assert sheet.page_setup.orientation == "landscape"
    assert sheet.page_setup.paperSize == int(sheet.PAPERSIZE_A4)
    assert sheet.page_setup.fitToWidth == 1 and sheet.page_setup.fitToHeight == 0
    assert sheet.sheet_properties.pageSetUpPr.fitToPage is True
    assert sheet.page_margins.left == 0.5 and sheet.page_margins.top == 0.75
    assert sheet.page_margins.header == 0.25
    assert sheet.print_area == "'Sheet'!$A$1:$C$8"
    assert sheet.print_title_rows == "$1:$3" and sheet.print_title_cols == "$A:$A"
    assert workbook["Notes"]["A1"].value == "Keep notes"


@pytest.mark.parametrize("extension", ["xlsx", "xlsm"])
def test_existing_template_layout_preserves_values_and_other_sheets(tmp_path, extension):
    path = tmp_path / f"template.{extension}"
    source(path)
    before = path.read_bytes()
    result = apply_layout(path, [sheet_layout(), print_layout()])
    assert result.get("patched"), result
    workbook = load_workbook(io.BytesIO(result["_persisted_bytes"]))
    assert_layout(workbook)
    assert workbook["Sheet"]["A1"].value == "Keep title" and workbook["Sheet"]["B4"].value == "=1+2"
    assert path.read_bytes() == before
    workbook.close()


def test_invoice_generation_and_patch_use_identical_layout_operations(tmp_path):
    result = _generate_office_operations_sync("xlsx", [normalize_file_patch_operation(op) for op in invoice_operations()])
    assert not result.get("error"), result
    path = tmp_path / "invoice.xlsx"
    path.write_bytes(result["_persisted_bytes"])
    workbook = load_workbook(path)
    assert_layout(workbook)
    assert workbook["Sheet"]["C7"].value == "=SUM(C4:C5)"
    assert workbook["Sheet"]["A3"].border.bottom.style == "medium"
    workbook.close()
    changed = apply_layout(path, [{"op": "sheet.format", "format": {"column_widths": {"A": 48}, "row_heights": {"1": 48}}}])
    assert not changed.get("error"), changed
    path.write_bytes(changed["_persisted_bytes"])
    layout = describe_file_structure(str(path))["sheets"][0]["layout"]
    assert layout["column_widths"]["A"] == 48 and layout["row_heights"]["1"] == 48
    assert layout["freeze_panes"] == "B4" and layout["show_gridlines"] is False


def test_cell_layout_changes_only_requested_style_properties(tmp_path):
    path = tmp_path / "style.xlsx"
    source(path)
    result = apply_layout(path, [{"op": "cell.format", "cell": "A1", "format": {
        "horizontal": "center", "vertical": "center", "indent": 2, "rotation": 30,
        "shrink_to_fit": True, "underline": True, "strike": True,
        "borders": {"bottom": {"style": "double", "color": "ABCDEF"}},
    }}])
    assert not result.get("error"), result
    workbook = load_workbook(io.BytesIO(result["_persisted_bytes"]))
    cell = workbook.active["A1"]
    assert cell.value == "Keep title" and cell.font.name == "Arial" and cell.font.italic is True
    assert cell.font.underline == "single" and cell.font.strike is True
    assert cell.alignment.horizontal == "center" and cell.alignment.vertical == "center"
    assert cell.alignment.indent == 2 and cell.alignment.textRotation == 30 and cell.alignment.wrap_text is True
    assert cell.alignment.shrinkToFit is True
    assert cell.border.left.style == "thin" and cell.border.bottom.style == "double"
    assert cell.border.bottom.color.rgb == "FFABCDEF"
    workbook.close()


def test_grouped_column_width_change_splits_only_target_column(tmp_path):
    path = tmp_path / "grouped.xlsx"
    source(path)
    workbook = load_workbook(path)
    workbook.active.column_dimensions.group("A", "D", outline_level=2, hidden=True)
    workbook.active.column_dimensions["A"].width = 20
    workbook.save(path)
    workbook.close()
    result = apply_layout(path, [{"op": "sheet.format", "format": {"column_widths": {"B": 35}}}])
    assert not result.get("error"), result
    workbook = load_workbook(io.BytesIO(result["_persisted_bytes"]))
    dimensions = sorted(workbook.active.column_dimensions.values(), key=lambda item: item.min)
    assert [(d.min, d.max, d.width) for d in dimensions] == [(1, 1, 20), (2, 2, 35), (3, 4, 20)]
    assert all(d.hidden and d.outlineLevel == 2 for d in dimensions)
    workbook.close()


@pytest.mark.parametrize("operation,style", [
    ("sheet.format", {}), ("sheet.format", {"column_widths": {"XFE": 20}}),
    ("sheet.format", {"column_widths": {"A1": 20}}), ("sheet.format", {"column_widths": {"A": True}}),
    ("sheet.format", {"column_widths": {"A": 256}}), ("sheet.format", {"row_heights": {"1048577": 20}}),
    ("sheet.format", {"row_heights": {"1": float("nan")}}), ("sheet.format", {"row_heights": {"1": 0}}),
    ("sheet.format", {"freeze_panes": "A0"}), ("sheet.format", {"show_gridlines": 0}),
    ("page.setup", {"print_area": "Other!A1:C8"}), ("page.setup", {"repeat_rows": "3:1"}),
    ("page.setup", {"repeat_columns": "A:XFE"}), ("page.setup", {"margin_left": -1}),
    ("page.setup", {"fit_width": 1.5}), ("page.setup", {"orientation": "diagonal"}),
    ("page.setup", {"paper_size": "not-paper"}), ("page.setup", {"width": 123}),
    ("cell.format", {"horizontal": "diagonal"}), ("cell.format", {"vertical": "left"}),
    ("cell.format", {"rotation": 181}), ("cell.format", {"indent": True}),
    ("cell.format", {"borders": {"unknown": {"style": "thin"}}}),
    ("cell.format", {"borders": {"bottom": {"style": "thinn", "color": "123456"}}}),
    ("cell.format", {"borders": {"bottom": {"style": "thin", "color": "#fff"}}}),
])
def test_invalid_template_layout_is_atomic(tmp_path, operation, style):
    path = tmp_path / "invalid.xlsx"
    source(path)
    before = path.read_bytes()
    patch = {"op": operation, "format": style, **({"cell": "A1"} if operation == "cell.format" else {})}
    result = apply_layout(path, [{"op": "cell.set", "cell": "A1", "value": "Do not commit"}, patch])
    assert result.get("error"), result
    assert result["operation_index"] == 1 and "_persisted_bytes" not in result
    assert path.read_bytes() == before


def test_reset_freeze_print_ranges_and_selected_border(tmp_path):
    path = tmp_path / "reset.xlsx"
    source(path)
    result = apply_layout(path, [sheet_layout(), print_layout(),
        {"op": "sheet.format", "format": {"freeze_panes": None, "show_gridlines": True}},
        {"op": "page.setup", "format": {"print_area": None, "repeat_rows": None, "repeat_columns": None}},
        {"op": "cell.format", "cell": "A1", "format": {"borders": {"left": None}}},
    ])
    assert not result.get("error"), result
    workbook = load_workbook(io.BytesIO(result["_persisted_bytes"]))
    sheet = workbook.active
    assert sheet.freeze_panes is None and sheet.sheet_view.showGridLines is True
    assert not sheet.print_area and sheet.print_title_rows is None and sheet.print_title_cols is None
    assert sheet["A1"].border.left.style is None
    workbook.close()


@pytest.mark.parametrize("target", ["C5", None])
def test_repeated_freeze_changes_do_not_duplicate_pane_selections(tmp_path, target):
    path = tmp_path / "panes.xlsx"
    source(path)
    workbook = load_workbook(path)
    workbook.active.sheet_view.selection[0].activeCell = "D6"
    workbook.active.sheet_view.selection[0].sqref = "D6"
    workbook.save(path)
    workbook.close()
    result = apply_layout(path, [sheet_layout(), {"op": "sheet.format", "format": {"freeze_panes": target}}])
    assert not result.get("error"), result
    workbook = load_workbook(io.BytesIO(result["_persisted_bytes"]))
    selections = workbook.active.sheet_view.selection
    assert [selection.pane for selection in selections] == (["topRight", "bottomLeft", "bottomRight"] if target else [None])
    assert selections[-1].activeCell == "D6" and selections[-1].sqref == "D6"
    workbook.close()


@pytest.mark.asyncio
@pytest.mark.parametrize("operation", [sheet_layout(), print_layout(), {"op": "cell.format", "cell": "A1", "format": {"horizontal": "center"}}])
@pytest.mark.parametrize("guard", ["resource", "approval", "stale"])
async def test_workbook_layout_uses_native_guards(file_runtime, monkeypatch, operation, guard):
    root, calls = file_runtime
    path = root / "guarded.xlsx"
    source(path)
    before = path.read_bytes()
    async def reject(**_kwargs):
        return json.dumps({"error": f"{guard}_denied"})
    if guard != "stale":
        monkeypatch.setattr(file_tools, "runtime_guard_file_resource_access" if guard == "resource" else "runtime_guard_file_mutation", reject)
    result = json.loads(await file_tools._patch_file("entity", path=path.name, operations=[operation],
        expected_sha256="0" * 64 if guard == "stale" else file_tools._file_meta(str(path))["source_sha256"]))
    assert result.get("error"), result
    assert path.read_bytes() == before and not any("commit" in call for call in calls)
