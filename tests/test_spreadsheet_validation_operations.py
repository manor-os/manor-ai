"""Native Excel data validation shares generation and patch execution."""
from __future__ import annotations

import io
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
import pytest

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools.file_tools import _apply_office_patch_sequence_sync
from packages.core.contracts.file_engine import file_patch_operations, normalize_file_patch_operation
from packages.core.services.file_engine_patches import describe_file_structure


def normalized(*operations):
    return [normalize_file_patch_operation(operation) for operation in operations]


def apply(path, *operations):
    return _apply_office_patch_sequence_sync(str(path), normalized(*operations))


def test_validation_operations_are_shared_xlsx_xlsm_capabilities():
    expected = {"validation.insert", "validation.format", "validation.delete"}
    assert expected <= set(file_patch_operations("xlsx"))
    assert expected <= set(file_patch_operations("xlsm"))


def test_generation_creates_and_formats_native_list_validation(tmp_path):
    result = _generate_office_operations_sync(
        "xlsx",
        normalized(
            {"op": "cell.set", "cell": "B1", "value": "Status"},
            {
                "op": "validation.insert",
                "range": "B2:B10",
                "validation": {
                    "type": "list",
                    "values": ["Planned", "Active", "Done"],
                    "allow_blank": True,
                    "show_dropdown": True,
                    "show_error_message": True,
                    "error_style": "stop",
                    "error_title": "Choose a status",
                    "error": "Use one of the listed statuses.",
                },
            },
            {
                "op": "validation.format",
                "validation_index": 0,
                "range": "B2:B20",
                "validation": {
                    "show_input_message": True,
                    "prompt_title": "Status",
                    "prompt": "Select the current status.",
                },
            },
        ),
    )
    assert not result.get("error"), result
    path = tmp_path / "validation.xlsx"
    path.write_bytes(result["_persisted_bytes"])
    workbook = load_workbook(path)
    validation = workbook.active.data_validations.dataValidation[0]
    assert str(validation.sqref) == "B2:B20"
    assert validation.type == "list" and validation.formula1 == '"Planned,Active,Done"'
    assert validation.allowBlank is True and validation.showDropDown is False
    assert validation.showErrorMessage is True and validation.errorStyle == "stop"
    assert validation.showInputMessage is True and validation.promptTitle == "Status"
    workbook.close()
    structure = describe_file_structure(str(path))
    assert structure["validation_count"] == 1
    assert structure["sheets"][0]["validations"] == [{
        "index": 0,
        "ranges": ["B2:B20"],
        "type": "list",
        "operator": None,
        "formula1": '"Planned,Active,Done"',
        "formula2": None,
        "allow_blank": True,
        "show_dropdown": True,
        "show_error_message": True,
        "error_style": "stop",
        "error_title": "Choose a status",
        "error": "Use one of the listed statuses.",
        "show_input_message": True,
        "prompt_title": "Status",
        "prompt": "Select the current status.",
    }]


def test_template_validation_format_preserves_unmodified_native_properties(tmp_path):
    workbook = Workbook()
    validation = DataValidation(
        type="whole",
        operator="between",
        formula1=1,
        formula2=10,
        allow_blank=True,
        showErrorMessage=True,
        errorStyle="warning",
        errorTitle="Original",
        error="Keep this message",
        imeMode="halfAlpha",
    )
    validation.add("C2:C9")
    workbook.active.add_data_validation(validation)
    workbook.active["E1"] = "Neighbor"
    path = tmp_path / "template.xlsx"
    workbook.save(path)
    workbook.close()
    result = apply(
        path,
        {
            "op": "validation.format",
            "validation_index": 0,
            "validation": {"operator": "equal", "formula1": 5, "error_title": "Exactly five"},
        },
    )
    assert not result.get("error"), result
    edited = load_workbook(io.BytesIO(result["_persisted_bytes"]))
    updated = edited.active.data_validations.dataValidation[0]
    assert updated.type == "whole" and updated.operator == "equal"
    assert updated.formula1 == "5" and updated.formula2 is None
    assert str(updated.sqref) == "C2:C9"
    assert updated.allowBlank is True and updated.showErrorMessage is True
    assert updated.errorStyle == "warning" and updated.error == "Keep this message"
    assert updated.errorTitle == "Exactly five" and updated.imeMode == "halfAlpha"
    assert edited.active["E1"].value == "Neighbor"
    edited.close()


@pytest.mark.parametrize(
    "validation",
    [
        {"type": "unknown", "formula1": 1},
        {"type": "whole", "operator": "between", "formula1": 1},
        {"type": "whole", "operator": "equal", "formula1": 1, "formula2": 2},
        {"type": "list", "values": []},
        {"type": "list", "values": ["Has,comma"]},
        {"type": "list", "values": ["One"], "operator": "equal"},
        {"type": "custom", "formula1": "=A1>0", "show_dropdown": True},
        {"type": "list", "values": ["One"], "error_title": "x" * 33},
    ],
)
def test_validation_insert_rejects_invalid_rules_atomically(tmp_path, validation):
    workbook = Workbook()
    workbook.active["A1"] = "Keep"
    path = tmp_path / "invalid.xlsx"
    workbook.save(path)
    workbook.close()
    before = path.read_bytes()
    result = apply(
        path,
        {"op": "cell.set", "cell": "A1", "value": "Rollback"},
        {"op": "validation.insert", "range": "B2:B10", "validation": validation},
    )
    assert result.get("operation_index") == 1
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before


def test_validation_overlap_rejected_and_delete_removes_only_selected_rule(tmp_path):
    workbook = Workbook()
    first = DataValidation(type="list", formula1='"One,Two"')
    first.add("A2:A10")
    second = DataValidation(type="custom", formula1="=B2>0")
    second.add("B2:B10")
    workbook.active.add_data_validation(first)
    workbook.active.add_data_validation(second)
    path = tmp_path / "delete.xlsx"
    workbook.save(path)
    workbook.close()
    rejected = apply(
        path,
        {"op": "validation.insert", "range": "A9:C12", "validation": {"type": "whole", "formula1": 1, "formula2": 10}},
    )
    assert "overlaps existing" in rejected.get("error", "")
    result = apply(path, {"op": "validation.delete", "validation_index": 0})
    assert not result.get("error"), result
    edited = load_workbook(io.BytesIO(result["_persisted_bytes"]))
    remaining = edited.active.data_validations.dataValidation
    assert len(remaining) == 1 and remaining[0].type == "custom" and str(remaining[0].sqref) == "B2:B10"
    edited.close()


def test_template_whole_column_validation_can_be_adjusted_and_blocks_overlap(tmp_path):
    workbook = Workbook()
    validation = DataValidation(type="list", formula1='"Yes,No"')
    validation.add("A1:A1048576")
    workbook.active.add_data_validation(validation)
    path = tmp_path / "whole-column.xlsx"
    workbook.save(path)
    workbook.close()
    result = apply(
        path,
        {"op": "validation.format", "validation_index": 0, "validation": {"allow_blank": True}},
    )
    assert not result.get("error"), result
    assert result["operation_results"][0]["ranges"] == ["A1:A1048576"]
    rejected = apply(
        path,
        {"op": "validation.insert", "range": "A20:A30", "validation": {"type": "custom", "formula1": "=A20>0"}},
    )
    assert "overlaps existing" in rejected.get("error", "")


def test_xlsm_validation_patch_preserves_macro_payload(tmp_path):
    created = _generate_office_operations_sync(
        "xlsx",
        normalized({"op": "cell.set", "cell": "A1", "value": "Keep"}),
    )
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(created["_persisted_bytes"])) as archive, zipfile.ZipFile(output, "w") as target:
        for item in archive.infolist():
            payload = archive.read(item.filename)
            if item.filename == "[Content_Types].xml":
                payload = payload.replace(
                    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                )
            target.writestr(item, payload)
        target.writestr("xl/vbaProject.bin", b"opaque-macro-payload")
    path = tmp_path / "macro.xlsm"
    path.write_bytes(output.getvalue())
    result = apply(
        path,
        {"op": "validation.insert", "range": "A2:A9", "validation": {"type": "custom", "formula1": "=A2<>\"\""}},
    )
    assert not result.get("error"), result
    with zipfile.ZipFile(io.BytesIO(result["_persisted_bytes"])) as archive:
        assert archive.read("xl/vbaProject.bin") == b"opaque-macro-payload"
