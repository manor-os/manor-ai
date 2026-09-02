"""Word inline pictures share the authorized generation and patch executor."""
from __future__ import annotations

import hashlib
import io
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from PIL import Image
import pytest

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools.file_tools import _apply_office_patch_sequence_sync
from packages.core.contracts.file_engine import file_patch_operations, normalize_file_patch_operation
from packages.core.services.file_engine_patches import describe_file_structure


def image_bytes(color: str, size=(200, 100), image_format="PNG") -> bytes:
    output = io.BytesIO()
    Image.new("RGB", size, color).save(output, image_format)
    return output.getvalue()


def source(path: str, data: bytes) -> dict[str, str]:
    return {"path": path, "expected_sha256": hashlib.sha256(data).hexdigest()}


def normalized(*operations):
    return [normalize_file_patch_operation(operation) for operation in operations]


def generate(operations, resources):
    result = _generate_office_operations_sync("docx", normalized(*operations), resources=resources)
    assert not result.get("error"), result
    return result


def picture_insert(src, **overrides):
    return {
        "op": "picture.insert", "index": 0, "source": src,
        "transform": {"width": 240},
        "format": {"alt_text": "Company logo", "alignment": "center"},
        **overrides,
    }


def test_word_picture_operations_are_shared_public_capabilities():
    assert {"picture.insert", "picture.replace", "picture.format"} <= set(file_patch_operations("docx"))


def test_generate_inserts_editable_inline_picture_with_aspect_alt_text_and_alignment(tmp_path):
    data = image_bytes("red", (200, 100))
    src = source("Assets/logo.png", data)
    result = generate([
        {"op": "paragraph.insert", "index": 0, "text": "Proposal"},
        picture_insert(src),
    ], {(src["path"], src["expected_sha256"]): data})
    path = tmp_path / "picture.docx"
    path.write_bytes(result["_persisted_bytes"])
    document = Document(path)
    assert [paragraph.text for paragraph in document.paragraphs] == ["", "Proposal"]
    picture = document.inline_shapes[0]
    assert picture.width / 12700 == 240 and picture.height / 12700 == 120
    assert picture._inline.docPr.get("descr") == "Company logo"
    assert document.paragraphs[0].alignment == 1
    assert document.part.related_parts[picture._inline.graphic.graphicData.pic.blipFill.blip.embed].blob == data
    details = describe_file_structure(str(path))
    assert details["picture_count"] == 1
    assert details["pictures"] == [{
        "picture_index": 0, "story": "body", "stories": ["body"], "section_indices": [],
        "layout": "inline", "paragraph_index": 0,
        "width": 240.0, "height": 120.0,
        "name": "Picture 1", "alt_text": "Company logo", "alignment": "center",
    }]


def test_picture_replace_preserves_layout_and_accessibility_but_changes_media(tmp_path):
    red, blue = image_bytes("red"), image_bytes("blue", (80, 160), "JPEG")
    first, second = source("Assets/red.png", red), source("Assets/blue.jpg", blue)
    initial = generate([picture_insert(first)], {(first["path"], first["expected_sha256"]): red})
    path = tmp_path / "replace.docx"
    path.write_bytes(initial["_persisted_bytes"])
    before = Document(path)
    state = (before.inline_shapes[0].width, before.inline_shapes[0].height,
             before.inline_shapes[0]._inline.docPr.get("descr"), before.paragraphs[0].alignment)
    result = _apply_office_patch_sequence_sync(
        str(path), normalized({"op": "picture.replace", "picture_index": 0, "source": second}),
        resources={(second["path"], second["expected_sha256"]): blue},
    )
    assert not result.get("error"), result
    document = Document(io.BytesIO(result["_persisted_bytes"]))
    picture = document.inline_shapes[0]
    assert (picture.width, picture.height, picture._inline.docPr.get("descr"), document.paragraphs[0].alignment) == state
    assert document.part.related_parts[picture._inline.graphic.graphicData.pic.blipFill.blip.embed].blob == blue


def test_replace_one_of_two_shared_pictures_keeps_the_other_relationship(tmp_path):
    red, blue = image_bytes("red"), image_bytes("blue")
    first, second = source("Assets/red.png", red), source("Assets/blue.png", blue)
    created = generate([
        picture_insert(first),
        picture_insert(first, index=1, transform={"width": 120}, format={"alt_text": "Second"}),
    ], {(first["path"], first["expected_sha256"]): red})
    path = tmp_path / "shared.docx"
    path.write_bytes(created["_persisted_bytes"])
    document = Document(path)
    first_rel = document.inline_shapes[0]._inline.graphic.graphicData.pic.blipFill.blip.embed
    second_rel = document.inline_shapes[1]._inline.graphic.graphicData.pic.blipFill.blip.embed
    assert first_rel == second_rel
    result = _apply_office_patch_sequence_sync(
        str(path), normalized({"op": "picture.replace", "picture_index": 0, "source": second}),
        resources={(second["path"], second["expected_sha256"]): blue},
    )
    document = Document(io.BytesIO(result["_persisted_bytes"]))
    relationships = [
        picture._inline.graphic.graphicData.pic.blipFill.blip.embed
        for picture in document.inline_shapes
    ]
    assert relationships[0] != relationships[1]
    assert document.part.related_parts[relationships[0]].blob == blue
    assert document.part.related_parts[relationships[1]].blob == red


def test_picture_format_resizes_proportionally_and_changes_only_requested_metadata(tmp_path):
    data = image_bytes("green", (300, 100))
    src = source("Assets/banner.png", data)
    initial = generate([picture_insert(src, transform={"width": 300}, format={"alt_text": "Before"})], {
        (src["path"], src["expected_sha256"]): data,
    })
    path = tmp_path / "format.docx"
    path.write_bytes(initial["_persisted_bytes"])
    result = _apply_office_patch_sequence_sync(str(path), normalized(
        {"op": "picture.format", "picture_index": 0,
         "format": {"height": 80, "alt_text": "After", "alignment": "right"}},
    ))
    document = Document(io.BytesIO(result["_persisted_bytes"]))
    picture = document.inline_shapes[0]
    assert picture.width / 12700 == 240 and picture.height / 12700 == 80
    assert picture._inline.docPr.get("descr") == "After"
    assert document.paragraphs[0].alignment == 2
    assert document.part.related_parts[picture._inline.graphic.graphicData.pic.blipFill.blip.embed].blob == data


def test_template_picture_patch_preserves_unrelated_header_footer_and_table(tmp_path):
    red, blue = image_bytes("red"), image_bytes("blue")
    second = source("Assets/blue.png", blue)
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Keep header"
    document.sections[0].footer.paragraphs[0].text = "Keep footer"
    document.add_paragraph().add_run().add_picture(io.BytesIO(red))
    table = document.add_table(1, 1)
    table.cell(0, 0).text = "Keep table"
    path = tmp_path / "template.docx"
    document.save(path)
    with ZipFile(path) as archive:
        before_header = archive.read("word/header1.xml")
        before_footer = archive.read("word/footer1.xml")
    result = _apply_office_patch_sequence_sync(
        str(path), normalized({"op": "picture.replace", "picture_index": 0, "source": second}),
        resources={(second["path"], second["expected_sha256"]): blue},
    )
    with ZipFile(io.BytesIO(result["_persisted_bytes"])) as archive:
        assert archive.read("word/header1.xml") == before_header
        assert archive.read("word/footer1.xml") == before_footer
    edited = Document(io.BytesIO(result["_persisted_bytes"]))
    assert edited.tables[0].cell(0, 0).text == "Keep table"


def test_generate_and_patch_header_footer_pictures_through_shared_operation_family(tmp_path):
    red, blue, green = image_bytes("red"), image_bytes("blue"), image_bytes("green")
    first, second, third = (
        source("Assets/header.png", red), source("Assets/replacement.png", blue), source("Assets/footer.png", green)
    )
    created = generate([
        {
            "op": "picture.insert", "story": "header", "section_index": 0, "index": 0, "source": first,
            "transform": {"width": 120}, "format": {"alt_text": "Header logo", "alignment": "right"},
        },
        {
            "op": "picture.insert", "story": "footer", "section_index": 0, "index": 0, "source": third,
            "transform": {"height": 36}, "format": {"alt_text": "Footer mark", "alignment": "center"},
        },
    ], {
        (first["path"], first["expected_sha256"]): red,
        (third["path"], third["expected_sha256"]): green,
    })
    path = tmp_path / "stories.docx"
    path.write_bytes(created["_persisted_bytes"])
    details = describe_file_structure(str(path))
    assert [(item["picture_index"], item["story"], item["section_indices"], item["layout"])
            for item in details["pictures"]] == [
        (0, "header", [0], "inline"),
        (1, "footer", [0], "inline"),
    ]
    assert details["pictures"][0]["paragraph_index"] == 0
    assert details["pictures"][0]["alt_text"] == "Header logo"
    assert details["pictures"][1]["alt_text"] == "Footer mark"

    replaced = _apply_office_patch_sequence_sync(
        str(path), normalized(
            {"op": "picture.replace", "picture_index": 0, "source": second},
            {"op": "picture.delete", "picture_index": 1},
        ),
        resources={(second["path"], second["expected_sha256"]): blue},
    )
    assert not replaced.get("error"), replaced
    edited = Document(io.BytesIO(replaced["_persisted_bytes"]))
    header = edited.sections[0].header
    footer = edited.sections[0].footer
    header_blip = next(header._element.iter(qn("a:blip")))
    assert header.part.related_parts[header_blip.get(qn("r:embed"))].blob == blue
    assert next(footer._element.iter(qn("a:blip")), None) is None


def test_floating_picture_generation_discovery_and_partial_patch_preserve_anchor_state(tmp_path):
    red, blue = image_bytes("red", (200, 100)), image_bytes("blue", (100, 200))
    first, second = source("Assets/floating.png", red), source("Assets/new.png", blue)
    created = generate([{
        "op": "picture.insert", "index": 0, "source": first, "transform": {"width": 200},
        "format": {
            "layout": "floating", "position_x": 72, "position_y": 36,
            "relative_from_horizontal": "page", "relative_from_vertical": "page",
            "wrap": "square", "behind_text": False, "allow_overlap": False,
            "layout_in_cell": True, "distance_left": 9, "distance_right": 9,
            "alt_text": "Floating hero",
        },
    }], {(first["path"], first["expected_sha256"]): red})
    path = tmp_path / "floating.docx"
    path.write_bytes(created["_persisted_bytes"])
    details = describe_file_structure(str(path))
    picture = details["pictures"][0]
    assert picture["layout"] == "floating" and picture["width"] == 200 and picture["height"] == 100
    assert picture["floating"] == {
        "position_x": 72.0, "position_y": 36.0,
        "relative_from_horizontal": "page", "relative_from_vertical": "page",
        "wrap": "square", "behind_text": False, "allow_overlap": False, "layout_in_cell": True,
        "distance_top": 0.0, "distance_bottom": 0.0, "distance_left": 9.0, "distance_right": 9.0,
    }
    with ZipFile(path) as archive:
        before_xml = archive.read("word/document.xml")
        assert b"<wp:anchor" in before_xml and b'relativeFrom="page"' in before_xml

    result = _apply_office_patch_sequence_sync(
        str(path), normalized(
            {"op": "picture.replace", "picture_index": 0, "source": second},
            {"op": "picture.format", "picture_index": 0,
             "format": {"height": 50, "position_y": 48, "behind_text": True}},
        ),
        resources={(second["path"], second["expected_sha256"]): blue},
    )
    assert not result.get("error"), result
    output = tmp_path / "floating-edited.docx"
    output.write_bytes(result["_persisted_bytes"])
    edited = describe_file_structure(str(output))["pictures"][0]
    assert edited["layout"] == "floating"
    assert edited["width"] == 100 and edited["height"] == 50
    assert edited["floating"]["position_x"] == 72 and edited["floating"]["position_y"] == 48
    assert edited["floating"]["wrap"] == "square" and edited["floating"]["behind_text"] is True
    document = Document(output)
    anchor = next(document.element.iter(qn("wp:anchor")))
    relationship_id = next(anchor.iter(qn("a:blip"))).get(qn("r:embed"))
    assert document.part.related_parts[relationship_id].blob == blue


def test_picture_layout_can_convert_back_to_inline_without_losing_media_or_metadata(tmp_path):
    data = image_bytes("purple")
    src = source("Assets/purple.png", data)
    created = generate([picture_insert(src, format={
        "layout": "floating", "position_x": 12, "position_y": 24, "alt_text": "Convertible",
    })], {(src["path"], src["expected_sha256"]): data})
    path = tmp_path / "convert.docx"
    path.write_bytes(created["_persisted_bytes"])
    result = _apply_office_patch_sequence_sync(str(path), normalized({
        "op": "picture.format", "picture_index": 0, "format": {"layout": "inline", "alignment": "left"},
    }))
    assert not result.get("error"), result
    document = Document(io.BytesIO(result["_persisted_bytes"]))
    assert len(document.inline_shapes) == 1
    assert document.inline_shapes[0]._inline.docPr.get("descr") == "Convertible"
    assert document.paragraphs[0].alignment == 0


@pytest.mark.parametrize("operation", [
    {"op": "picture.insert", "index": 0, "source": {}, "fit": "cover"},
    {"op": "picture.insert", "index": 0, "source": {}, "transform": {"x": 10}},
    {"op": "picture.format", "picture_index": 0, "format": {}},
    {"op": "picture.format", "picture_index": 0, "format": {"width": 0}},
    {"op": "picture.format", "picture_index": 0, "format": {"alignment": "justify"}},
    {"op": "picture.format", "picture_index": 0, "format": {"layout": "floating", "alignment": "left"}},
    {"op": "picture.format", "picture_index": 0, "format": {"position_x": 10}},
    {"op": "picture.insert", "story": "sidebar", "source": {}},
    {"op": "picture.format", "picture_index": 0, "format": {"alt_text": "x" * 1001}},
    {"op": "picture.replace", "picture_index": 9,
     "source": {"path": "Assets/other.png", "expected_sha256": "a" * 64}},
])
def test_invalid_picture_operation_aborts_the_whole_batch(tmp_path, operation):
    data = image_bytes("red")
    src = source("Assets/red.png", data)
    initial = generate([picture_insert(src)], {(src["path"], src["expected_sha256"]): data})
    path = tmp_path / "atomic.docx"
    path.write_bytes(initial["_persisted_bytes"])
    before = path.read_bytes()
    result = _apply_office_patch_sequence_sync(str(path), normalized(
        {"op": "paragraph.insert", "index": 1, "text": "Must roll back"}, operation,
    ))
    assert result.get("error"), result
    assert result["operation_index"] == 1 and "_persisted_bytes" not in result
    assert path.read_bytes() == before
