"""Native Office deletion uses the same generation and patch executor."""
from __future__ import annotations

import hashlib
import io
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.drawing.image import Image as SpreadsheetImage
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools.file_tools import _apply_office_patch_sequence_sync
from packages.core.contracts.file_engine import file_patch_operations, normalize_file_patch_operation


def normalized(*operations):
    return [normalize_file_patch_operation(operation) for operation in operations]


def apply(path, *operations):
    return _apply_office_patch_sequence_sync(str(path), normalized(*operations))


def image_bytes(color: str) -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (100, 60), color).save(output, "PNG")
    return output.getvalue()


def source(path: str, data: bytes) -> dict[str, str]:
    return {"path": path, "expected_sha256": hashlib.sha256(data).hexdigest()}


def test_delete_operations_are_public_for_each_native_office_type():
    assert {"paragraph.delete", "table.delete", "picture.delete"} <= set(file_patch_operations("docx"))
    assert {
        "paragraph.delete", "table.delete", "slide.delete", "shape.delete",
        "picture.delete", "chart.delete",
    } <= set(file_patch_operations("pptx"))
    assert {"sheet.delete", "table.delete", "picture.delete", "chart.delete"} <= set(
        file_patch_operations("xlsx")
    )
    assert set(file_patch_operations("xlsx")) == set(file_patch_operations("xlsm"))


def test_word_delete_removes_only_selected_native_content_and_media(tmp_path):
    red = image_bytes("red")
    src = source("Assets/logo.png", red)
    result = _generate_office_operations_sync(
        "docx",
        normalized(
            {"op": "paragraph.insert", "index": 0, "text": "Keep"},
            {"op": "paragraph.insert", "index": 1, "text": "Remove"},
            {"op": "table.insert", "index": 2, "rows": [["Remove table"]]},
            {"op": "picture.insert", "index": 2, "source": src},
            {"op": "picture.insert", "index": 3, "source": src},
            {"op": "paragraph.delete", "index": 1},
            {"op": "table.delete", "table_index": 0},
            {"op": "picture.delete", "picture_index": 0},
        ),
        resources={(src["path"], src["expected_sha256"]): red},
    )
    assert not result.get("error"), result
    path = tmp_path / "deleted.docx"
    path.write_bytes(result["_persisted_bytes"])
    document = Document(path)
    assert "Remove" not in [paragraph.text for paragraph in document.paragraphs]
    assert "Keep" in [paragraph.text for paragraph in document.paragraphs]
    assert len(document.tables) == 0
    assert len(document.inline_shapes) == 1
    remaining = document.inline_shapes[0]
    assert document.part.related_parts[remaining._inline.graphic.graphicData.pic.blipFill.blip.embed].blob == red


def test_word_rejects_section_break_deletion_atomically(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("Section owner")
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:sectPr"))
    path = tmp_path / "section.docx"
    document.save(path)
    before = path.read_bytes()
    result = apply(path, {"op": "paragraph.delete", "index": 0})
    assert "section break" in result.get("error", "")
    assert path.read_bytes() == before and "_persisted_bytes" not in result


def test_powerpoint_deletes_typed_shapes_paragraphs_and_slides(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    textbox.text_frame.text = "Keep"
    textbox.text_frame.add_paragraph().text = "Remove"
    picture = slide.shapes.add_picture(io.BytesIO(image_bytes("blue")), Inches(1), Inches(2))
    table = slide.shapes.add_table(1, 1, Inches(3), Inches(2), Inches(2), Inches(1))
    table.table.cell(0, 0).text = "Delete"
    shape = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(2), Inches(1))
    shape.text = "Delete shape"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    path = tmp_path / "delete.pptx"
    presentation.save(path)

    result = apply(
        path,
        {"op": "paragraph.delete", "slide": 1, "shape_id": textbox.shape_id, "index": 1},
        {"op": "picture.delete", "slide": 1, "shape_id": picture.shape_id},
        {"op": "table.delete", "slide": 1, "shape_id": table.shape_id},
        {"op": "shape.delete", "slide": 1, "shape_id": shape.shape_id},
        {"op": "slide.delete", "index": 1},
    )
    assert not result.get("error"), result
    edited = Presentation(io.BytesIO(result["_persisted_bytes"]))
    assert len(edited.slides) == 1
    assert [item.shape_id for item in edited.slides[0].shapes] == [textbox.shape_id]
    assert [paragraph.text for paragraph in edited.slides[0].shapes[0].text_frame.paragraphs] == ["Keep"]
    with ZipFile(io.BytesIO(result["_persisted_bytes"])) as archive:
        assert not [name for name in archive.namelist() if name.startswith("ppt/media/")]


def test_powerpoint_chart_delete_prunes_native_chart_parts():
    created = _generate_office_operations_sync(
        "pptx",
        normalized(
            {"op": "slide.insert", "index": 0, "layout_index": 6},
            {
                "op": "chart.insert", "slide": 1, "chart_type": "column",
                "categories": ["A", "B"], "series": [{"name": "Sales", "values": [1, 2]}],
                "transform": {"x": 72, "y": 72, "width": 360, "height": 240},
            },
            {"op": "chart.delete", "slide": 1, "shape_id": 2},
        ),
    )
    assert not created.get("error"), created
    presentation = Presentation(io.BytesIO(created["_persisted_bytes"]))
    assert not presentation.slides[0].shapes
    with ZipFile(io.BytesIO(created["_persisted_bytes"])) as archive:
        assert not [name for name in archive.namelist() if name.startswith("ppt/charts/")]
        assert not [name for name in archive.namelist() if name.startswith("ppt/embeddings/")]


def test_powerpoint_rejects_wrong_typed_and_only_paragraph_deletes(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    textbox.text = "Only"
    path = tmp_path / "atomic.pptx"
    presentation.save(path)
    before = path.read_bytes()
    wrong_type = apply(path, {"op": "picture.delete", "slide": 1, "shape_id": textbox.shape_id})
    only_paragraph = apply(path, {"op": "paragraph.delete", "slide": 1, "shape_id": textbox.shape_id, "index": 0})
    assert "picture" in wrong_type.get("error", "")
    assert "only paragraph" in only_paragraph.get("error", "")
    assert path.read_bytes() == before


def test_excel_deletes_native_objects_but_preserves_cells_and_table_range(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.append(["Name", "Value"])
    sheet.append(["A", 1])
    sheet.add_table(Table(displayName="ReportTable", ref="A1:B2"))
    sheet.add_image(SpreadsheetImage(io.BytesIO(image_bytes("green"))), "D4")
    workbook.create_sheet("Spare")
    path = tmp_path / "delete.xlsx"
    workbook.save(path)
    workbook.close()

    result = apply(
        path,
        {"op": "table.delete", "sheet": "Report", "table": "reporttable"},
        {"op": "picture.delete", "sheet": "Report", "picture_index": 0},
        {"op": "sheet.delete", "sheet": "Spare"},
    )
    assert not result.get("error"), result
    edited = load_workbook(io.BytesIO(result["_persisted_bytes"]))
    assert edited.sheetnames == ["Report"]
    assert edited["Report"]["A2"].value == "A" and edited["Report"]["B2"].value == 1
    assert not edited["Report"].tables and not edited["Report"]._images
    edited.close()


def test_excel_chart_delete_removes_managed_data_sheet():
    result = _generate_office_operations_sync(
        "xlsx",
        normalized(
            {
                "op": "chart.insert", "sheet": "Sheet", "chart_type": "line",
                "categories": ["Jan", "Feb"], "series": [{"name": "Sales", "values": [1, 2]}],
                "anchor": "E3",
            },
            {"op": "chart.delete", "sheet": "Sheet", "chart_index": 0},
        ),
    )
    assert not result.get("error"), result
    workbook = load_workbook(io.BytesIO(result["_persisted_bytes"]))
    assert not workbook.active._charts
    assert workbook.sheetnames == ["Sheet"]
    workbook.close()


def test_excel_sheet_delete_removes_its_managed_chart_data_sheet():
    result = _generate_office_operations_sync(
        "xlsx",
        normalized(
            {"op": "sheet.add", "sheet": "Keep"},
            {
                "op": "chart.insert", "sheet": "Sheet", "chart_type": "column",
                "categories": ["Jan", "Feb"], "series": [{"name": "Sales", "values": [1, 2]}],
                "anchor": "E3",
            },
            {"op": "sheet.delete", "sheet": "Sheet"},
        ),
    )
    assert not result.get("error"), result
    workbook = load_workbook(io.BytesIO(result["_persisted_bytes"]))
    assert workbook.sheetnames == ["Keep"]
    workbook.close()


def test_excel_rejects_last_visible_managed_and_referenced_sheet_deletes(tmp_path):
    workbook = Workbook()
    source_sheet = workbook.active
    source_sheet.title = "Inputs"
    source_sheet["A1"] = 10
    report = workbook.create_sheet("Report")
    report["A1"] = "=Inputs!A1"
    path = tmp_path / "refs.xlsx"
    workbook.save(path)
    workbook.close()
    before = path.read_bytes()
    referenced = apply(path, {"op": "sheet.delete", "sheet": "Inputs"})
    assert "referenced by formula Report!A1" in referenced.get("error", "")
    assert path.read_bytes() == before

    workbook = Workbook()
    workbook.save(path)
    workbook.close()
    last = apply(path, {"op": "sheet.delete", "sheet": "Sheet"})
    assert "last visible" in last.get("error", "")


def test_excel_sheet_delete_rejects_validation_formatting_and_hyperlink_references(tmp_path):
    for kind in ("validation", "conditional formatting", "hyperlink"):
        workbook = Workbook()
        inputs = workbook.active
        inputs.title = "Inputs"
        report = workbook.create_sheet("Report")
        if kind == "validation":
            validation = DataValidation(type="list", formula1="'Inputs'!$A$1:$A$2")
            report.add_data_validation(validation)
            validation.add("B1")
        elif kind == "conditional formatting":
            report.conditional_formatting.add("B1", FormulaRule(formula=["Inputs!$A$1>0"]))
        else:
            report["B1"].hyperlink = "#Inputs!A1"
        path = tmp_path / f"{kind.replace(' ', '-')}.xlsx"
        workbook.save(path)
        workbook.close()
        result = apply(path, {"op": "sheet.delete", "sheet": "Inputs"})
        assert kind in result.get("error", ""), result
