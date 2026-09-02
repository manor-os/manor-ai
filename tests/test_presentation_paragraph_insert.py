"""PowerPoint paragraphs can be inserted into an existing native text frame."""
from __future__ import annotations

import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools.file_tools import _apply_office_patch_sequence_sync
from packages.core.contracts.file_engine import file_patch_operations, normalize_file_patch_operation


def normalized(*operations):
    return [normalize_file_patch_operation(operation) for operation in operations]


def test_presentation_paragraph_insert_is_a_shared_public_operation():
    assert "paragraph.insert" in file_patch_operations("pptx")


def test_generation_inserts_and_formats_paragraph_in_existing_textbox():
    result = _generate_office_operations_sync(
        "pptx",
        normalized(
            {"op": "slide.insert", "index": 0, "layout_index": 6},
            {
                "op": "textbox.insert", "slide": 1, "text": "First",
                "transform": {"x": 72, "y": 72, "width": 360, "height": 144},
            },
            {
                "op": "paragraph.format", "slide": 1, "shape_id": 2, "index": 0,
                "format": {"level": 1, "font_size": 22, "font_color": "123ABC"},
            },
            {
                "op": "paragraph.insert", "slide": 1, "shape_id": 2, "index": 1,
                "text": "Second", "format": {"bold": True, "space_before": 6},
            },
        ),
    )
    assert not result.get("error"), result
    shape = Presentation(io.BytesIO(result["_persisted_bytes"])).slides[0].shapes[0]
    paragraphs = shape.text_frame.paragraphs
    assert [paragraph.text for paragraph in paragraphs] == ["First", "Second"]
    assert paragraphs[1].level == 1
    assert paragraphs[1].runs[0].font.size.pt == 22
    assert paragraphs[1].runs[0].font.color.rgb == RGBColor(0x12, 0x3A, 0xBC)
    assert paragraphs[1].runs[0].font.bold is True
    assert paragraphs[1].space_before.pt == 6


def test_template_insert_at_start_and_middle_inherits_nearest_paragraph_style(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    first = shape.text_frame.paragraphs[0]
    first.text = "Original first"
    first.level = 2
    first.runs[0].font.name = "Aptos Display"
    first.runs[0].font.size = Pt(19)
    first.runs[0].font.italic = True
    last = shape.text_frame.add_paragraph()
    last.text = "Original last"
    last.level = 0
    path = tmp_path / "paragraphs.pptx"
    presentation.save(path)
    geometry = (shape.left, shape.top, shape.width, shape.height)

    result = _apply_office_patch_sequence_sync(
        str(path),
        normalized(
            {"op": "paragraph.insert", "slide": 1, "shape_id": shape.shape_id, "index": 1, "text": "Middle"},
            {"op": "paragraph.insert", "slide": 1, "shape_id": shape.shape_id, "index": 0, "text": "New first"},
        ),
    )
    assert not result.get("error"), result
    edited_shape = Presentation(io.BytesIO(result["_persisted_bytes"])).slides[0].shapes[0]
    assert (edited_shape.left, edited_shape.top, edited_shape.width, edited_shape.height) == geometry
    paragraphs = edited_shape.text_frame.paragraphs
    assert [paragraph.text for paragraph in paragraphs] == [
        "New first", "Original first", "Middle", "Original last",
    ]
    for paragraph in (paragraphs[0], paragraphs[2]):
        assert paragraph.level == 2
        assert paragraph.runs[0].font.name == "Aptos Display"
        assert paragraph.runs[0].font.size.pt == 19
        assert paragraph.runs[0].font.italic is True


def test_invalid_paragraph_insert_aborts_batch_without_touching_source(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    shape.text = "Keep"
    path = tmp_path / "atomic.pptx"
    presentation.save(path)
    before = path.read_bytes()
    result = _apply_office_patch_sequence_sync(
        str(path),
        normalized(
            {"op": "text.set", "slide": 1, "shape_id": shape.shape_id, "index": 0, "text": "Rollback"},
            {"op": "paragraph.insert", "slide": 1, "shape_id": shape.shape_id, "index": 9, "text": "Invalid"},
        ),
    )
    assert result.get("operation_index") == 1
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before
