"""Precise paragraph edits share generation, patching and the native save path."""
from __future__ import annotations

import io
import json

import pytest

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools import file_tools
from packages.core.contracts.file_engine import file_type_capability, normalize_file_patch_operation
from packages.core.services.file_engine_patches import describe_file_structure
from tests.test_file_engine_structured_patch import file_runtime as file_runtime


def text_operations(extension):
    if extension == "docx":
        return [{"op": "paragraph.insert", "index": i, "text": text} for i, text in enumerate(["same", "", "same"])]
    return [
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        {"op": "textbox.insert", "slide": 1, "text": "same\n\nsame",
         "transform": {"x": 72, "y": 72, "width": 480, "height": 144}},
        {"op": "textbox.insert", "slide": 1, "text": "same",
         "transform": {"x": 72, "y": 250, "width": 480, "height": 72}},
    ]


def text_operation(extension, text="Target only", index=1):
    return {"op": "text.set", "index": index, "text": text,
            **({"slide": 1, "shape_id": 2} if extension == "pptx" else {})}


def open_paragraphs(source, extension):
    if extension == "docx":
        from docx import Document
        document = Document(source)
        return document, document.paragraphs
    from pptx import Presentation
    document = Presentation(source)
    return document, document.slides[0].shapes[0].text_frame.paragraphs


def text_source(path):
    extension = path.suffix[1:]
    result = _generate_office_operations_sync(extension, [normalize_file_patch_operation(op) for op in text_operations(extension)])
    assert not result.get("error"), result
    path.write_bytes(result["_persisted_bytes"])


def apply_text(path, operations):
    return file_tools._apply_office_patch_sequence_sync(str(path), [normalize_file_patch_operation(op) for op in operations])


@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("index", [0, 1, 2])
@pytest.mark.parametrize("text", ["Changed", "", "中文\tEnglish\r\nSecond\rThird"])
def test_generate_and_patch_exact_paragraph(tmp_path, extension, index, text):
    operations = text_operations(extension) + [text_operation(extension, text, index)]
    result = _generate_office_operations_sync(extension, [normalize_file_patch_operation(op) for op in operations])
    assert not result.get("error"), result
    assert "text.set" in file_type_capability(extension)["operations"]
    document, paragraphs = open_paragraphs(io.BytesIO(result["_persisted_bytes"]), extension)
    expected = ["same", "", "same"]
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    expected[index] = normalized if extension == "docx" else normalized.replace("\n", "\v")
    assert [p.text for p in paragraphs] == expected
    untouched = [p._p.xml for i, p in enumerate(paragraphs) if i != index]
    path = tmp_path / f"precise.{extension}"
    document.save(path)
    before = path.read_bytes()
    result = apply_text(path, [text_operation(extension, "Followup\nline", index)])
    assert not result.get("error"), result
    document, paragraphs = open_paragraphs(io.BytesIO(result["_persisted_bytes"]), extension)
    expected[index] = "Followup\nline" if extension == "docx" else "Followup\vline"
    assert [p.text for p in paragraphs] == expected
    assert [p._p.xml for i, p in enumerate(paragraphs) if i != index] == untouched
    assert path.read_bytes() == before
    if extension == "pptx":
        shape = document.slides[0].shapes[0]
        assert (shape.left.pt, shape.top.pt, shape.width.pt, shape.height.pt) == (72, 72, 480, 144)
        assert document.slides[0].shapes[1].text == "same"


@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("empty", [False, True])
def test_text_set_keeps_layout_links_and_typing_style(tmp_path, extension, empty):
    path = tmp_path / f"styled.{extension}"
    text_source(path)
    document, paragraphs = open_paragraphs(str(path), extension)
    index = 1 if empty else 0
    paragraph = paragraphs[index]
    if extension == "docx":
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt
        paragraph.paragraph_format.space_after = Pt(14)
        paragraph.style = "List Bullet"
        section = OxmlElement("w:sectPr")
        paragraph._p.get_or_add_pPr().append(section)
        if empty:
            properties = OxmlElement("w:rPr")
            properties.append(OxmlElement("w:b"))
            paragraph._p.get_or_add_pPr().append(properties)
        else:
            from docx.opc.constants import RELATIONSHIP_TYPE as RT
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), paragraph.part.relate_to("https://example.com", RT.HYPERLINK, is_external=True))
            run = paragraph.runs[0]
            run.bold = True
            paragraph._p.append(link)
            link.append(run._r)
    else:
        from pptx.oxml.xmlchemy import OxmlElement
        from pptx.util import Pt
        paragraph.space_after = Pt(14)
        paragraph.level = 2
        if empty:
            properties = OxmlElement("a:endParaRPr")
            properties.set("b", "1")
            paragraph._p.append(properties)
        else:
            paragraph.runs[0].font.bold = True
            paragraph.runs[0].hyperlink.address = "https://example.com"
    properties = paragraph._p.pPr.xml
    document.save(path)
    result = apply_text(path, [text_operation(extension, "New\nline", index)])
    assert not result.get("error"), result
    document, paragraphs = open_paragraphs(io.BytesIO(result["_persisted_bytes"]), extension)
    paragraph = paragraphs[index]
    assert paragraph._p.pPr.xml == properties
    if extension == "docx":
        from docx.text.run import Run
        runs = [Run(r, paragraph) for r in paragraph._p.xpath(".//w:r") if r.text]
        if not empty:
            assert len(paragraph.hyperlinks) == 1
            assert paragraph.hyperlinks[0].address == "https://example.com"
    else:
        runs = [run for run in paragraph.runs if run.text]
        if not empty:
            assert all(run.hyperlink.address == "https://example.com" for run in runs)
    assert runs and all(run.font.bold is True for run in runs)


@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("invalid", [
    {"index": True}, {"index": -1}, {"index": 3}, {"index": None},
    {"text": None}, {"text": 42}, {"text": "bad\x00value"}, {"text": "x" * 32768},
    {"cell": "A1"}, {"table_index": 0}, {"sheet": "Sheet"},
])
def test_invalid_late_text_set_is_atomic(tmp_path, extension, invalid):
    path = tmp_path / f"invalid.{extension}"
    text_source(path)
    before = path.read_bytes()
    result = apply_text(path, [text_operation(extension), {**text_operation(extension), **invalid}])
    assert result.get("error"), result
    assert result["operation_index"] == 1
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before


@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("kind", ["field", "drawing", "annotation"])
@pytest.mark.parametrize("operation", ["text.set", "paragraph.format"])
def test_text_set_does_not_destroy_structured_content(tmp_path, extension, kind, operation):
    path = tmp_path / f"protected.{extension}"
    text_source(path)
    document, paragraphs = open_paragraphs(str(path), extension)
    if extension == "docx":
        from docx.oxml import OxmlElement
        element = OxmlElement({"field": "w:fldSimple", "drawing": "w:drawing", "annotation": "w:bookmarkStart"}[kind])
        paragraphs[1]._p.append(element)
    else:
        from pptx.oxml.xmlchemy import OxmlElement
        paragraphs[1]._p.append(OxmlElement({"field": "a:fld", "drawing": "a:unknown", "annotation": "a:extLst"}[kind]))
    document.save(path)
    before = path.read_bytes()
    target = {**text_operation(extension), "op": operation, "format": {"bold": True}}
    result = apply_text(path, [text_operation(extension, "Earlier", 0), target])
    assert "cannot replace" in result["error"]
    assert result["operation_index"] == 1
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before


def test_word_text_set_rejects_cross_paragraph_field_and_page_break(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    path = tmp_path / "fields.docx"
    text_source(path)
    document, paragraphs = open_paragraphs(str(path), "docx")
    field = OxmlElement("w:fldChar")
    field.set(qn("w:fldCharType"), "begin")
    paragraphs[0].add_run()._r.append(field)
    paragraphs[1].add_run("cached field text")
    document.save(path)
    result = apply_text(path, [text_operation("docx")])
    assert "cannot replace" in result["error"]
    field.getparent().remove(field)
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    paragraphs[1].add_run()._r.append(br)
    document.save(path)
    assert "page or column break" in apply_text(path, [text_operation("docx")])["error"]


@pytest.mark.parametrize("target", ["table", "missing", "missing_index", "bool_slide"])
def test_powerpoint_text_set_requires_exact_text_frame(tmp_path, target):
    from pptx.util import Pt
    path = tmp_path / "target.pptx"
    text_source(path)
    document, _ = open_paragraphs(str(path), "pptx")
    shapes = document.slides[0].shapes
    table = shapes.add_table(1, 1, Pt(0), Pt(0), Pt(50), Pt(50))
    group = shapes.add_group_shape()
    group.shapes.add_textbox(Pt(0), Pt(0), Pt(50), Pt(50))
    document.save(path)
    op = text_operation("pptx", index=0)
    if target == "missing_index":
        op.pop("index")
    elif target == "bool_slide":
        op["slide"] = True
    else:
        op["shape_id"] = {"table": table.shape_id, "missing": 999}[target]
    result = apply_text(path, [op])
    assert result.get("error"), result
    assert "_persisted_bytes" not in result


def test_powerpoint_text_set_targets_a_group_child_by_slide_unique_shape_id(tmp_path):
    from pptx import Presentation
    from pptx.util import Pt

    path = tmp_path / "group-child.pptx"
    text_source(path)
    document, _ = open_paragraphs(str(path), "pptx")
    group = document.slides[0].shapes.add_group_shape()
    child = group.shapes.add_textbox(Pt(0), Pt(0), Pt(50), Pt(50))
    child.text = "Before"
    document.save(path)
    before = path.read_bytes()
    result = apply_text(path, [{**text_operation("pptx", "After", 0), "shape_id": child.shape_id}])
    assert not result.get("error"), result
    assert path.read_bytes() == before
    patched = Presentation(io.BytesIO(result["_persisted_bytes"]))
    assert patched.slides[0].shapes[-1].shapes[0].text == "After"
    assert result["operation_results"][0]["group_path"] == [group.shape_id, child.shape_id]


def test_ppt_structure_discovers_text_paragraphs_with_global_bound(tmp_path):
    path = tmp_path / "structure.pptx"
    text_source(path)
    structure = describe_file_structure(str(path))
    frame = structure["shapes"][0]["text_frame"]
    assert frame == {"paragraph_count": 3, "paragraphs": [{"index": i, "text": text} for i, text in enumerate(["same", "", "same"])], "truncated": False}
    assert structure["shapes"][0]["has_text_frame"] is True
    document, _ = open_paragraphs(str(path), "pptx")
    document.slides[0].shapes[0].text = "\n".join(str(i) for i in range(201))
    document.save(path)
    structure = describe_file_structure(str(path))
    assert structure["truncated"] is True
    assert sum(len(shape["text_frame"]["paragraphs"]) for shape in structure["shapes"]) == 200
    assert structure["shapes"][1]["text_frame"]["paragraphs"] == []
    assert structure["shapes"][1]["text_frame"]["paragraph_count"] == 1
    assert structure["shapes"][1]["text_frame"]["truncated"] is True


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("guard", ["resource", "approval", "stale"])
async def test_text_set_keeps_native_guards(file_runtime, monkeypatch, extension, guard):
    root, calls = file_runtime
    path = root / f"guarded.{extension}"
    text_source(path)
    before = path.read_bytes()
    async def reject(**_kwargs):
        return json.dumps({"error": f"{guard}_denied"})
    if guard != "stale":
        monkeypatch.setattr(file_tools, "runtime_guard_file_resource_access" if guard == "resource" else "runtime_guard_file_mutation", reject)
    result = json.loads(await file_tools._patch_file(
        "entity", path=path.name, operations=[text_operation(extension)],
        expected_sha256="0" * 64 if guard == "stale" else file_tools._file_meta(str(path))["source_sha256"],
    ))
    assert result.get("error"), result
    assert path.read_bytes() == before
    assert not any("commit" in call for call in calls)


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["docx", "pptx"])
async def test_text_set_approval_binds_full_target_and_commits_once(file_runtime, extension):
    root, calls = file_runtime
    path = root / f"approved.{extension}"
    text_source(path)
    read = json.loads(await file_tools._read_file("entity", path=path.name, include_structure=True))
    operations = [text_operation(extension, "First", 0), text_operation(extension, "Second", 1)]
    result = json.loads(await file_tools._patch_file(
        "entity", path=path.name, operations=operations, expected_sha256=read["source_sha256"],
    ))
    assert result.get("patched") is True, result
    approvals = [call for call in calls if "content_preview" in call]
    assert len(approvals) == 1
    assert approvals[0]["content_preview"]["operations"] == [normalize_file_patch_operation(op) for op in operations]
    assert len([call for call in calls if "commit" in call]) == 1
    assert result["document_id"] == "same-document"
    _, paragraphs = open_paragraphs(str(path), extension)
    assert [p.text for p in paragraphs] == ["First", "Second", "same"]
