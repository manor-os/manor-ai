"""Native Excel pictures share generation and patch operations with Word and PowerPoint."""
from __future__ import annotations

import hashlib
import io
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as SpreadsheetImage
from openpyxl.drawing.spreadsheet_drawing import AbsoluteAnchor, AnchorMarker, TwoCellAnchor
from openpyxl.drawing.xdr import XDRPoint2D, XDRPositiveSize2D
from PIL import Image
import pytest

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools import file_tools
from packages.core.contracts.file_engine import file_patch_operations, normalize_file_patch_operation
from packages.core.services.file_engine_patches import describe_file_structure


def image_bytes(color: str, size=(200, 100), image_format="PNG") -> bytes:
    output = io.BytesIO()
    Image.new("RGB", size, color).save(output, image_format)
    return output.getvalue()


def source(path: str, data: bytes) -> dict[str, str]:
    return {"path": path, "expected_sha256": hashlib.sha256(data).hexdigest()}


def normalized(*operations):
    return [normalize_file_patch_operation(operation) for operation in operations]


def generate(operations, resources):
    result = _generate_office_operations_sync("xlsx", normalized(*operations), resources=resources)
    assert not result.get("error"), result
    return result


def apply(path, operations, resources=None):
    return file_tools._apply_office_patch_sequence_sync(
        str(path), normalized(*operations), resources=resources,
    )


def picture_insert(src, **overrides):
    return {
        "op": "picture.insert", "sheet": "Sheet", "source": src, "anchor": "C4",
        "transform": {"width": 240}, "format": {"alt_text": "Company logo"},
        **overrides,
    }


def picture_blob(workbook, index=0):
    return workbook.active._images[index]._data()


def test_spreadsheet_picture_operations_are_shared_public_capabilities():
    expected = {"picture.insert", "picture.replace", "picture.format"}
    assert expected <= set(file_patch_operations("xlsx"))
    assert expected <= set(file_patch_operations("xlsm"))
    assert "native Excel pictures" in file_tools._file_engine_capabilities("xlsx")["limits"]["spreadsheet_pictures"]


def test_generate_inserts_native_picture_with_anchor_size_alt_text_and_descriptor(tmp_path):
    data = image_bytes("red", (200, 100))
    src = source("Assets/logo.png", data)
    result = generate([picture_insert(src)], {(src["path"], src["expected_sha256"]): data})
    path = tmp_path / "picture.xlsx"
    path.write_bytes(result["_persisted_bytes"])
    workbook = load_workbook(path)
    picture = workbook.active._images[0]
    assert (picture.anchor._from.col, picture.anchor._from.row) == (2, 3)
    assert (picture.anchor.ext.width, picture.anchor.ext.height) == (240 * 12700, 120 * 12700)
    assert picture.anchor.pic.nvPicPr.cNvPr.descr == "Company logo"
    assert picture_blob(workbook) == data
    workbook.close()

    structure = describe_file_structure(str(path))
    details = structure["sheets"][0]["pictures"][0]
    assert structure["picture_count"] == 1
    assert details == {
        "picture_index": 0, "anchor": "C4", "to_anchor": None,
        "anchor_type": "OneCellAnchor", "width": 240.0, "height": 120.0,
        "intrinsic_width_px": 200, "intrinsic_height_px": 100,
        "name": "logo.png", "alt_text": "Company logo", "format": "png",
    }


def test_picture_replace_preserves_native_geometry_name_and_alt_text(tmp_path):
    red, blue = image_bytes("red"), image_bytes("blue", (80, 160), "JPEG")
    first, second = source("Assets/red.png", red), source("Assets/blue.jpg", blue)
    created = generate([picture_insert(first)], {(first["path"], first["expected_sha256"]): red})
    path = tmp_path / "replace.xlsx"
    path.write_bytes(created["_persisted_bytes"])
    result = apply(path, [{"op": "picture.replace", "sheet": "Sheet", "picture_index": 0, "source": second}], {
        (second["path"], second["expected_sha256"]): blue,
    })
    assert not result.get("error"), result
    workbook = load_workbook(io.BytesIO(result["_persisted_bytes"]))
    picture = workbook.active._images[0]
    assert (picture.anchor._from.col, picture.anchor._from.row) == (2, 3)
    assert (picture.anchor.ext.width, picture.anchor.ext.height) == (240 * 12700, 120 * 12700)
    assert picture.anchor.pic.nvPicPr.cNvPr.name == "red.png"
    assert picture.anchor.pic.nvPicPr.cNvPr.descr == "Company logo"
    assert picture_blob(workbook) == blue
    workbook.close()


def test_picture_format_moves_and_resizes_using_displayed_aspect_ratio(tmp_path):
    data = image_bytes("green", (300, 100))
    src = source("Assets/banner.png", data)
    created = generate([picture_insert(src, transform={"width": 300})], {
        (src["path"], src["expected_sha256"]): data,
    })
    path = tmp_path / "format.xlsx"
    path.write_bytes(created["_persisted_bytes"])
    result = apply(path, [{
        "op": "picture.format", "sheet": "Sheet", "picture_index": 0,
        "format": {"anchor": "F8", "height": 80, "alt_text": "Updated banner"},
    }])
    assert not result.get("error"), result
    workbook = load_workbook(io.BytesIO(result["_persisted_bytes"]))
    picture = workbook.active._images[0]
    assert (picture.anchor._from.col, picture.anchor._from.row) == (5, 7)
    assert (picture.anchor.ext.width, picture.anchor.ext.height) == (240 * 12700, 80 * 12700)
    assert picture.anchor.pic.nvPicPr.cNvPr.descr == "Updated banner"
    assert picture_blob(workbook) == data
    workbook.close()


def test_two_cell_template_picture_moves_exactly_and_requires_explicit_size(tmp_path):
    data = image_bytes("purple")
    path = tmp_path / "two-cell.xlsx"
    workbook = Workbook()
    picture = SpreadsheetImage(io.BytesIO(data))
    picture.anchor = TwoCellAnchor(
        _from=AnchorMarker(col=3, row=2, colOff=1234, rowOff=5678),
        to=AnchorMarker(col=8, row=13, colOff=4321, rowOff=8765),
    )
    workbook.active.add_image(picture)
    workbook.save(path)
    workbook.close()

    moved = apply(path, [{"op": "picture.format", "picture_index": 0, "format": {"anchor": "G8"}}])
    workbook = load_workbook(io.BytesIO(moved["_persisted_bytes"]))
    anchor = workbook.active._images[0].anchor
    assert isinstance(anchor, TwoCellAnchor)
    assert (anchor._from.col, anchor._from.row, anchor._from.colOff, anchor._from.rowOff) == (6, 7, 1234, 5678)
    assert (anchor.to.col, anchor.to.row, anchor.to.colOff, anchor.to.rowOff) == (11, 18, 4321, 8765)
    workbook.close()

    rejected = apply(path, [{"op": "picture.format", "picture_index": 0, "format": {"width": 300}}])
    assert "requires both width and height" in rejected.get("error", "")
    converted = apply(path, [{
        "op": "picture.format", "picture_index": 0, "format": {"width": 300, "height": 180},
    }])
    workbook = load_workbook(io.BytesIO(converted["_persisted_bytes"]))
    assert workbook.active._images[0].anchor.ext.width == 300 * 12700
    assert workbook.active._images[0].anchor.ext.height == 180 * 12700
    workbook.close()


def test_absolute_template_picture_requires_size_when_converting_to_a_cell_anchor(tmp_path):
    path = tmp_path / "absolute.xlsx"
    workbook = Workbook()
    picture = SpreadsheetImage(io.BytesIO(image_bytes("teal")))
    picture.anchor = AbsoluteAnchor(
        pos=XDRPoint2D(x=1000, y=2000),
        ext=XDRPositiveSize2D(cx=180 * 12700, cy=90 * 12700),
    )
    workbook.active.add_image(picture)
    workbook.save(path)
    workbook.close()

    rejected = apply(path, [{"op": "picture.format", "picture_index": 0, "format": {"anchor": "C3"}}])
    assert "requires both width and height" in rejected.get("error", "")
    converted = apply(path, [{
        "op": "picture.format", "picture_index": 0,
        "format": {"anchor": "C3", "width": 180, "height": 90},
    }])
    workbook = load_workbook(io.BytesIO(converted["_persisted_bytes"]))
    anchor = workbook.active._images[0].anchor
    assert (anchor._from.col, anchor._from.row) == (2, 2)
    assert (anchor.ext.width, anchor.ext.height) == (180 * 12700, 90 * 12700)
    workbook.close()


def test_xlsm_picture_patch_preserves_macro_payload(tmp_path):
    data = image_bytes("orange")
    src = source("Assets/logo.png", data)
    created = generate([picture_insert(src)], {(src["path"], src["expected_sha256"]): data})
    path = tmp_path / "macro.xlsm"
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(created["_persisted_bytes"])) as archive, zipfile.ZipFile(output, "w") as target:
        for item in archive.infolist():
            payload = archive.read(item.filename)
            if item.filename == "[Content_Types].xml":
                payload = payload.replace(
                    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                )
            target.writestr(item, payload)
        target.writestr("xl/vbaProject.bin", b"opaque-macro-payload")
    path.write_bytes(output.getvalue())
    result = apply(path, [{"op": "picture.format", "picture_index": 0, "format": {"alt_text": "Macro logo"}}])
    assert not result.get("error"), result
    with zipfile.ZipFile(io.BytesIO(result["_persisted_bytes"])) as archive:
        assert archive.read("xl/vbaProject.bin") == b"opaque-macro-payload"


@pytest.mark.parametrize("operation", [
    {"op": "picture.insert", "source": {}, "anchor": "A0"},
    {"op": "picture.insert", "source": {}, "transform": {"x": 10}},
    {"op": "picture.format", "picture_index": 0, "format": {}},
    {"op": "picture.format", "picture_index": 0, "format": {"width": 0}},
    {"op": "picture.format", "picture_index": 0, "format": {"alt_text": "x" * 1001}},
    {"op": "picture.replace", "picture_index": 9,
     "source": {"path": "Assets/other.png", "expected_sha256": "a" * 64}},
])
def test_invalid_picture_operation_aborts_the_whole_batch(tmp_path, operation):
    data = image_bytes("red")
    src = source("Assets/red.png", data)
    created = generate([picture_insert(src)], {(src["path"], src["expected_sha256"]): data})
    path = tmp_path / "atomic.xlsx"
    path.write_bytes(created["_persisted_bytes"])
    before = path.read_bytes()
    result = apply(path, [
        {"op": "cell.set", "cell": "A1", "value": "Must roll back"}, operation,
    ])
    assert result.get("error"), result
    assert result["operation_index"] == 1 and "_persisted_bytes" not in result
    assert path.read_bytes() == before
