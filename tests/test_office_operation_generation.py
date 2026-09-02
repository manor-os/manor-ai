"""Office creation and patching share operations, validation and editable OOXML."""
from __future__ import annotations

import io
import json

import pytest

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools import file_tools
from packages.core.ai.tools.generate_file import document, tool
from packages.core.contracts.file_engine import normalize_file_patch_operation
from packages.core.services.file_engine_patches import describe_file_structure


def office_operations(extension):
    return {
        "docx": [
            {"op": "paragraph.insert", "index": 0, "text": "Proposal", "style": "Heading 1"},
            {"op": "table.insert", "index": 1, "rows": [["Item", "Price"], ["Service", 12]], "style": "Table Grid"},
        ],
        "pptx": [
            {"op": "slide.insert", "index": 0, "layout_index": 6},
            {"op": "textbox.insert", "slide": 1, "text": "Proposal", "transform": {"x": 72, "y": 72, "width": 360, "height": 72}},
        ],
        "xlsx": [
            {"op": "cell.set", "cell": "A1", "value": "Proposal"},
            {"op": "cell.set", "cell": "B1", "value": "=2*6"},
            {"op": "cell.format", "cell": "A1", "format": {"bold": True, "fill_color": "123ABC"}},
            {"op": "sheet.add", "sheet": "Summary"},
        ],
    }[extension]


@pytest.mark.parametrize("extension", ["docx", "pptx", "xlsx"])
def test_generate_uses_patch_executor_and_remains_editable(tmp_path, monkeypatch, extension):
    real_apply = file_tools._apply_office_patch_sequence_sync
    calls = []

    def apply(path, operations):
        calls.append(operations)
        return real_apply(path, operations)

    monkeypatch.setattr(file_tools, "_apply_office_patch_sequence_sync", apply)
    operations = [normalize_file_patch_operation(op) for op in office_operations(extension)]
    generated = _generate_office_operations_sync(extension, operations)
    assert calls == [operations]
    assert generated["operations_applied"] == len(operations)
    path = tmp_path / f"generated.{extension}"
    path.write_bytes(generated["_persisted_bytes"])
    before = path.read_bytes()
    followup = {"op": "text.replace", "old_text": "Proposal", "new_text": "Revised"}
    if extension == "xlsx":
        followup = {"op": "cell.set", "sheet": "Sheet", "cell": "A1", "value": "Revised"}
    patched = real_apply(str(path), [normalize_file_patch_operation(followup)])
    assert not patched.get("error"), patched
    assert path.read_bytes() == before  # executor never commits the source
    data = io.BytesIO(patched["_persisted_bytes"])
    if extension == "docx":
        from docx import Document

        reopened = Document(data)
        assert reopened.paragraphs[0].text == "Revised"
        assert reopened.paragraphs[0].style.name == "Heading 1"
        assert reopened.tables[0].cell(1, 1).text == "12"
        assert reopened.tables[0].style.name == "Table Grid"
        assert describe_file_structure(str(path))["table_count"] == 1
    elif extension == "pptx":
        from pptx import Presentation

        reopened = Presentation(data)
        assert len(reopened.slides) == 1
        shape = reopened.slides[0].shapes[0]
        assert shape.text == "Revised"
        assert (shape.left.pt, shape.top.pt, shape.width.pt, shape.height.pt) == (72, 72, 360, 72)
        assert describe_file_structure(str(path))["layouts"][6]["name"] == "Blank"
    else:
        from openpyxl import load_workbook

        reopened = load_workbook(data)
        try:
            assert reopened.sheetnames == ["Sheet", "Summary"]
            assert reopened["Sheet"]["A1"].value == "Revised"
            assert reopened["Sheet"]["B1"].value == "=2*6"
            assert reopened["Sheet"]["A1"].font.bold is True
            assert reopened["Sheet"]["A1"].fill.fgColor.rgb == "FF123ABC"
        finally:
            reopened.close()


def test_new_structures_preserve_existing_order_and_style(tmp_path):
    from docx import Document
    from pptx import Presentation

    word = Document()
    word.add_paragraph("Before").runs[0].bold = True
    word.add_table(1, 1).cell(0, 0).text = "Original table"
    word.add_paragraph("After")
    word.sections[0].header.paragraphs[0].text = "Keep header"
    path = tmp_path / "existing.docx"
    word.save(path)
    result = file_tools._apply_office_patch_sequence_sync(str(path), [
        normalize_file_patch_operation({"op": "table.insert", "index": 1, "rows": [["New", None], [True, 2.5]]}),
    ])
    edited = Document(io.BytesIO(result["_persisted_bytes"]))
    assert result["operation_results"][0]["table_index"] == 1
    assert [element.tag.rsplit("}", 1)[-1] for element in edited.element.body] == ["p", "tbl", "tbl", "p", "sectPr"]
    assert edited.paragraphs[0].runs[0].bold is True
    assert edited.tables[0].cell(0, 0).text == "Original table"
    assert edited.tables[1].cell(0, 1).text == ""
    assert edited.sections[0].header.paragraphs[0].text == "Keep header"

    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[0])
    first.shapes.title.text = "Before"
    first.shapes.title.text_frame.paragraphs[0].runs[0].font.bold = True
    first.notes_slide.notes_text_frame.text = "Keep notes"
    second = deck.slides.add_slide(deck.slide_layouts[0])
    second.shapes.title.text = "After"
    ids = [slide.slide_id for slide in deck.slides]
    path = tmp_path / "existing.pptx"
    deck.save(path)
    result = file_tools._apply_office_patch_sequence_sync(str(path), [
        normalize_file_patch_operation({"op": "slide.insert", "index": 1, "layout_index": 6}),
        normalize_file_patch_operation({"op": "textbox.insert", "slide": 2, "text": "Inserted", "transform": {"x": 20, "y": 30, "width": 200, "height": 60, "rotation": -90}}),
    ])
    edited = Presentation(io.BytesIO(result["_persisted_bytes"]))
    assert [edited.slides[index].slide_id for index in (0, 2)] == ids
    assert edited.slides[0].shapes.title.text == "Before"
    assert edited.slides[0].shapes.title.text_frame.paragraphs[0].runs[0].font.bold is True
    assert edited.slides[0].notes_slide.notes_text_frame.text == "Keep notes"
    assert edited.slides[2].shapes.title.text == "After"
    assert edited.slides[1].shapes[0].text == "Inserted"
    assert edited.slides[1].shapes[0].rotation == 270


@pytest.mark.parametrize("extension,invalid", [
    ("docx", {"op": "table.insert", "index": 1, "rows": []}),
    ("docx", {"op": "table.insert", "index": 1, "rows": [[1, 2], [3]]}),
    ("docx", {"op": "table.insert", "index": 1, "rows": [[{}]]}),
    ("docx", {"op": "table.insert", "index": 1, "rows": [["x"]] * 201}),
    ("docx", {"op": "table.insert", "index": 1, "rows": [["x"] * 51]}),
    ("docx", {"op": "table.insert", "index": True, "rows": [[1]]}),
    ("pptx", {"op": "slide.insert", "index": 0}),
    ("pptx", {"op": "slide.insert", "index": 2, "layout_index": 6}),
    ("pptx", {"op": "slide.insert", "index": 0, "layout_index": 999}),
    ("pptx", {"op": "textbox.insert", "slide": 1, "text": "bad", "transform": {"x": 0}}),
    ("pptx", {"op": "textbox.insert", "slide": 1, "text": "bad", "transform": {"x": 0, "y": 0, "width": 0, "height": 1}}),
    ("xlsx", {"op": "cell.set", "cell": "A0", "value": "bad"}),
])
def test_late_operation_failure_does_not_produce_partial_file(extension, invalid):
    operations = office_operations(extension) + [invalid]
    result = _generate_office_operations_sync(extension, [normalize_file_patch_operation(op) for op in operations])
    assert result.get("error"), result
    assert result["operation_index"] == len(operations) - 1
    assert "_persisted_bytes" not in result


@pytest.mark.asyncio
@pytest.mark.parametrize("kind,extension", [("document", "docx"), ("document", "pptx"), ("document", "xlsx"), ("word_document", "docx"), ("presentation", "pptx"), ("spreadsheet", "xlsx")])
@pytest.mark.parametrize("nested", [False, True])
async def test_operations_route_without_specialist_or_model(monkeypatch, kind, extension, nested):
    captured = []

    async def runtime(**kwargs):
        captured.append(kwargs)
        return json.dumps({"created": True})

    async def unexpected(**_kwargs):
        pytest.fail("Operation generation must not call a model or specialist skill")

    monkeypatch.setattr(document, "runtime_generate_document_file", runtime)
    for name in ("handle_word_document", "handle_spreadsheet", "handle_presentation"):
        monkeypatch.setattr(tool, name, unexpected)
    arguments = {"operations": office_operations(extension)}
    if nested:
        arguments = {"params": arguments}
    result = json.loads(await tool._generate_file_handler("entity", kind=kind, name=f"new.{extension}", **arguments))
    assert result["created"] is True
    assert captured[0]["file_type"] == extension
    assert captured[0]["operations"] == office_operations(extension)
    assert captured[0]["content"] is None


@pytest.mark.asyncio
@pytest.mark.parametrize("arguments", [
    {"kind": "image"}, {"kind": "pdf"}, {"kind": "code"},
    {"content": "ignored"}, {"prompt": "ignored"}, {"params": {"prompt": "ignored"}},
    {"params": {"content": "ignored"}}, {"options": {}}, {"files": []},
    {"kind": "presentation", "file_type": "docx"},
    {"file_type": "pdf"}, {"file_type": "docx", "params": {"file_type": "pptx"}},
    {"params": {"operations": [{"op": "paragraph.insert", "index": 0, "text": "different"}]}},
    {"operations": None}, {"operations": []},
])
async def test_invalid_generation_contract_never_dispatches(monkeypatch, arguments):
    async def unexpected(**_kwargs):
        pytest.fail("Invalid operation generation reached the runtime")

    monkeypatch.setattr(tool, "handle_document", unexpected)
    monkeypatch.setattr(tool, "handle_image", unexpected)
    monkeypatch.setattr(tool, "handle_pdf", unexpected)
    monkeypatch.setattr(tool, "handle_code", unexpected)
    kwargs = {"kind": "document", "name": "new.docx", "operations": office_operations("docx"), **arguments}
    result = json.loads(await tool._generate_file_handler("entity", **kwargs))
    assert result.get("error"), result


@pytest.mark.asyncio
@pytest.mark.parametrize("arguments", [
    {"file_type": "pptx"},
    {"name": "new.xlsm", "file_type": "xlsm"},
    {"name": "../escape.docx"},
    {"name": ".private.docx"},
    {"expected_sha256": "old-source-hash"},
    {"content": "ignored"},
    {"options": {}},
    {"operations": []},
    {"operations": [{"op": "sheet.add", "sheet": "Wrong format"}]},
    {"operations": [{"op": "table.insert", "index": 0, "rows": [[float("nan")]]}]},
])
async def test_runtime_rejects_invalid_operations_before_approval(tmp_path, monkeypatch, arguments):
    from packages.core.ai.runtime import file_actions, generated_files

    async def no_approval(**_kwargs):
        pytest.fail("Invalid generation reached mutation approval")

    monkeypatch.setattr(file_actions, "runtime_entity_file_root", lambda _entity: str(tmp_path))
    monkeypatch.setattr(file_actions, "runtime_guard_file_mutation", no_approval)
    kwargs = {"entity_id": "entity", "user_id": "", "conversation_id": "", "name": "new.docx", "file_type": "docx", "operations": office_operations("docx"), **arguments}
    result = json.loads(await generated_files.runtime_generate_document_file(**kwargs))
    assert result.get("error"), result
    assert not list(tmp_path.glob("*.docx"))


@pytest.mark.asyncio
async def test_runtime_rejects_empty_content_before_approval(tmp_path, monkeypatch):
    from packages.core.ai.runtime import file_actions, generated_files

    async def no_approval(**_kwargs):
        pytest.fail("Empty content reached mutation approval")

    monkeypatch.setattr(file_actions, "runtime_entity_file_root", lambda _entity: str(tmp_path))
    monkeypatch.setattr(file_actions, "runtime_guard_file_mutation", no_approval)
    result = json.loads(await generated_files.runtime_generate_document_file(
        entity_id="entity",
        user_id="",
        conversation_id="",
        name="empty.txt",
        file_type="txt",
        content="",
    ))

    assert result == {"error": "name and non-empty content or operations are required"}


@pytest.mark.asyncio
async def test_content_generation_never_overwrites_without_expected_hash(tmp_path, monkeypatch):
    from packages.core.ai.runtime import file_actions, generated_files

    target = tmp_path / "existing.md"
    target.write_text("original", encoding="utf-8")

    async def no_approval(**_kwargs):
        pytest.fail("Unversioned overwrite reached mutation approval")

    monkeypatch.setattr(file_actions, "runtime_entity_file_root", lambda _entity: str(tmp_path))
    monkeypatch.setattr(file_actions, "runtime_guard_file_mutation", no_approval)
    result = json.loads(await generated_files.runtime_generate_document_file(
        entity_id="entity",
        user_id="",
        conversation_id="",
        name=target.name,
        file_type="md",
        content="replacement",
    ))

    assert result["error"] == "file_already_exists"
    assert target.read_text(encoding="utf-8") == "original"
