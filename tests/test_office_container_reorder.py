"""Slide and worksheet order are native shared operations, not generator-only state."""
from __future__ import annotations

import io
import zipfile

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from pptx import Presentation

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools.file_tools import _apply_office_patch_sequence_sync
from packages.core.contracts.file_engine import file_patch_operations, normalize_file_patch_operation


def _operations(*operations: dict) -> list[dict]:
    return [normalize_file_patch_operation(operation) for operation in operations]


def _presentation_seed() -> list[dict]:
    operations = []
    for index, title in enumerate(("First", "Second", "Third")):
        operations.extend((
            {"op": "slide.insert", "index": index, "layout_index": 6},
            {
                "op": "textbox.insert", "slide": index + 1, "text": title,
                "transform": {"x": 72, "y": 72, "width": 320, "height": 60},
            },
        ))
    return _operations(*operations)


def _presentation_titles(data: bytes) -> list[str]:
    document = Presentation(io.BytesIO(data))
    return [slide.shapes[0].text for slide in document.slides]


def _spreadsheet_seed() -> list[dict]:
    return _operations(
        {"op": "cell.set", "sheet": "Sheet", "cell": "A1", "value": "First"},
        {"op": "sheet.add", "sheet": "Second"},
        {"op": "cell.set", "sheet": "Second", "cell": "A1", "value": "Second"},
        {"op": "sheet.add", "sheet": "Third"},
        {"op": "cell.set", "sheet": "Third", "cell": "A1", "value": "Third"},
    )


def _spreadsheet_order(data: bytes) -> tuple[list[str], list[str], str]:
    workbook = load_workbook(io.BytesIO(data))
    try:
        return (
            list(workbook.sheetnames),
            [workbook[name]["A1"].value for name in workbook.sheetnames],
            workbook.active.title,
        )
    finally:
        workbook.close()


def test_sheet_add_accepts_a_shared_zero_based_insertion_index(tmp_path):
    operation = normalize_file_patch_operation({
        "op": "sheet.add", "sheet": "Cover", "index": 0,
    })
    blank = _generate_office_operations_sync("xlsx", [operation])
    template = _generate_office_operations_sync(
        "xlsx", _operations({"op": "cell.set", "cell": "A1", "value": "Keep"}),
    )["_persisted_bytes"]
    templated = _generate_office_operations_sync(
        "xlsx", [operation], template_bytes=template,
    )
    path = tmp_path / "insert-sheet.xlsx"
    path.write_bytes(template)
    patched = _apply_office_patch_sequence_sync(str(path), [operation])
    for result in (blank, templated, patched):
        assert not result.get("error"), result
        assert result["operation_results"][-1]["index"] == 0
        workbook = load_workbook(io.BytesIO(result["_persisted_bytes"]))
        try:
            assert workbook.sheetnames == ["Cover", "Sheet"]
        finally:
            workbook.close()

    invalid = _apply_office_patch_sequence_sync(str(path), _operations({
        "op": "sheet.add", "sheet": "Bad", "index": 2,
    }))
    assert invalid.get("error"), invalid
    assert "_persisted_bytes" not in invalid
    assert path.read_bytes() == template


def test_slide_reorder_is_shared_by_blank_generation_template_generation_and_patch(tmp_path):
    assert "slide.reorder" in file_patch_operations("pptx")
    reorder = normalize_file_patch_operation(
        {"op": "slide.reorder", "slide": 1, "index": 2},
    )
    blank = _generate_office_operations_sync(
        "pptx", [*_presentation_seed(), reorder],
    )
    template = _generate_office_operations_sync(
        "pptx", _presentation_seed(),
    )["_persisted_bytes"]
    templated = _generate_office_operations_sync(
        "pptx", [reorder], template_bytes=template,
    )
    path = tmp_path / "slides.pptx"
    path.write_bytes(template)
    patched = _apply_office_patch_sequence_sync(str(path), [reorder])
    for result in (blank, templated, patched):
        assert not result.get("error"), result
        assert result["operation_results"][-1] == {
            "operation": "slide.reorder",
            "source_slide": 1,
            "slide": 3,
            "previous_index": 0,
            "index": 2,
        }
        assert _presentation_titles(result["_persisted_bytes"]) == [
            "Second", "Third", "First",
        ]
    assert path.read_bytes() == template


def test_sheet_reorder_is_shared_by_blank_generation_template_generation_and_patch(tmp_path):
    assert "sheet.reorder" in file_patch_operations("xlsx")
    reorder = normalize_file_patch_operation(
        {"op": "sheet.reorder", "sheet": "Sheet", "index": 2},
    )
    blank = _generate_office_operations_sync(
        "xlsx", [*_spreadsheet_seed(), reorder],
    )
    template = _generate_office_operations_sync(
        "xlsx", _spreadsheet_seed(),
    )["_persisted_bytes"]
    templated = _generate_office_operations_sync(
        "xlsx", [reorder], template_bytes=template,
    )
    path = tmp_path / "sheets.xlsx"
    path.write_bytes(template)
    patched = _apply_office_patch_sequence_sync(str(path), [reorder])
    for result in (blank, templated, patched):
        assert not result.get("error"), result
        assert result["operation_results"][-1] == {
            "operation": "sheet.reorder",
            "updated": True,
            "sheet": "Sheet",
            "previous_index": 0,
            "index": 2,
            "sheet_names": ["Second", "Third", "Sheet"],
        }
        assert _spreadsheet_order(result["_persisted_bytes"]) == (
            ["Second", "Third", "Sheet"],
            ["Second", "Third", "First"],
            "Sheet",
        )
    assert path.read_bytes() == template


def test_invalid_slide_reorder_is_atomic(tmp_path):
    source = _generate_office_operations_sync("pptx", _presentation_seed())["_persisted_bytes"]
    path = tmp_path / "invalid-slides.pptx"
    path.write_bytes(source)
    before = path.read_bytes()
    result = _apply_office_patch_sequence_sync(str(path), _operations(
        {"op": "slide.reorder", "slide": 1, "index": 2},
        {"op": "slide.reorder", "slide": 9, "index": 0},
    ))
    assert result.get("error"), result
    assert result["operation_index"] == 1
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before


def test_invalid_sheet_reorder_is_atomic(tmp_path):
    source = _generate_office_operations_sync("xlsx", _spreadsheet_seed())["_persisted_bytes"]
    path = tmp_path / "invalid-sheets.xlsx"
    path.write_bytes(source)
    before = path.read_bytes()
    result = _apply_office_patch_sequence_sync(str(path), _operations(
        {"op": "sheet.reorder", "sheet": "Sheet", "index": 2},
        {"op": "sheet.reorder", "sheet": "Missing", "index": 0},
    ))
    assert result.get("error"), result
    assert result["operation_index"] == 1
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before


def _spreadsheet_rename_seed() -> list[dict]:
    return _operations(
        {"op": "sheet.add", "sheet": "Source Data"},
        {"op": "cell.set", "sheet": "Source Data", "cell": "A1", "value": 2},
        {"op": "cell.set", "sheet": "Sheet", "cell": "A1", "value": "='Source Data'!A1*2"},
    )


def _assert_renamed_sheet(data: bytes) -> None:
    workbook = load_workbook(io.BytesIO(data), data_only=False)
    try:
        assert workbook.sheetnames == ["Sheet", "Rates 2026"]
        assert workbook["Sheet"]["A1"].value == "='Rates 2026'!A1*2"
        assert workbook["Rates 2026"]["A1"].value == 2
        assert workbook.calculation.fullCalcOnLoad is True
        assert workbook.calculation.forceFullCalc is True
    finally:
        workbook.close()


def test_sheet_rename_is_shared_by_blank_generation_template_generation_and_patch(tmp_path):
    assert "sheet.rename" in file_patch_operations("xlsx")
    rename = normalize_file_patch_operation({
        "op": "sheet.rename", "sheet": "Source Data", "new_sheet": "Rates 2026",
    })
    blank = _generate_office_operations_sync(
        "xlsx", [*_spreadsheet_rename_seed(), rename],
    )
    template = _generate_office_operations_sync(
        "xlsx", _spreadsheet_rename_seed(),
    )["_persisted_bytes"]
    templated = _generate_office_operations_sync(
        "xlsx", [rename], template_bytes=template,
    )
    path = tmp_path / "rename.xlsx"
    path.write_bytes(template)
    patched = _apply_office_patch_sequence_sync(str(path), [rename])
    for result in (blank, templated, patched):
        assert not result.get("error"), result
        assert result["operation_results"][-1] == {
            "operation": "sheet.rename",
            "updated": True,
            "sheet": "Source Data",
            "new_sheet": "Rates 2026",
            "sheet_names": ["Sheet", "Rates 2026"],
        }
        _assert_renamed_sheet(result["_persisted_bytes"])
    assert path.read_bytes() == template


def test_sheet_rename_updates_native_references_without_touching_external_or_string_literals(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "reference-rename.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    source = workbook.create_sheet("Source Data")
    other = workbook.create_sheet("Other")
    source.append([2, 3])
    other["A1"] = 4
    summary["A1"] = "='Source Data'!A1*2"
    summary["A2"] = '=IF("Source Data!A1"="Source Data!A1",\'Source Data\'!A1,0)'
    summary["A3"] = "='[1]Source Data'!A1"
    summary["A4"] = "=SUM('Source Data:Other'!A1)"
    summary["D1"] = "Jump"
    summary["D1"]._hyperlink = Hyperlink(
        ref="D1", location="#'Source Data'!A1", display="Jump",
    )
    summary["D2"] = "External"
    summary["D2"]._hyperlink = Hyperlink(
        ref="D2", target="https://example.com/book.xlsx", location="#'Source Data'!A1",
    )
    workbook.defined_names.add(DefinedName("SourceValue", attr_text="'Source Data'!$A$1"))
    validation = DataValidation(type="whole", formula1="'Source Data'!$A$1", formula2="'Source Data'!$B$1")
    summary.add_data_validation(validation)
    validation.add("B1:B3")
    summary.conditional_formatting.add("C1:C3", FormulaRule(formula=["'Source Data'!A1>0"]))
    chart = BarChart()
    chart.add_data(Reference(source, min_col=1, min_row=1, max_col=2, max_row=1))
    summary.add_chart(chart, "F1")
    workbook.save(path)
    workbook.close()

    result = _apply_office_patch_sequence_sync(str(path), _operations({
        "op": "sheet.rename", "sheet": "Source Data", "new_sheet": "Rates 2026",
    }))
    assert not result.get("error"), result
    renamed = load_workbook(io.BytesIO(result["_persisted_bytes"]), data_only=False)
    try:
        summary = renamed["Summary"]
        assert summary["A1"].value == "='Rates 2026'!A1*2"
        assert summary["A2"].value == '=IF("Source Data!A1"="Source Data!A1",\'Rates 2026\'!A1,0)'
        assert summary["A3"].value == "='[1]Source Data'!A1"
        assert summary["A4"].value == "=SUM('Rates 2026:Other'!A1)"
        assert renamed.defined_names["SourceValue"].attr_text == "'Rates 2026'!$A$1"
        assert summary.data_validations.dataValidation[0].formula1 == "'Rates 2026'!$A$1"
        assert summary.data_validations.dataValidation[0].formula2 == "'Rates 2026'!$B$1"
        rule = next(iter(summary.conditional_formatting._cf_rules.values()))[0]
        assert rule.formula == ["'Rates 2026'!A1>0"]
        assert [series.val.numRef.f for series in summary._charts[0].series] == [
            "'Rates 2026'!$A$1", "'Rates 2026'!$B$1",
        ]
        assert summary["D1"].hyperlink.location == "#'Rates 2026'!A1"
        assert summary["D2"].hyperlink.location == "#'Source Data'!A1"
    finally:
        renamed.close()


def test_sheet_rename_preserves_xlsm_macro_payload(tmp_path):
    source = _generate_office_operations_sync("xlsx", _spreadsheet_rename_seed())["_persisted_bytes"]
    path = tmp_path / "rename.xlsm"
    with zipfile.ZipFile(io.BytesIO(source), "r") as archive, zipfile.ZipFile(path, "w") as target:
        for info in archive.infolist():
            payload = archive.read(info.filename)
            if info.filename == "[Content_Types].xml":
                payload = payload.replace(
                    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                )
            target.writestr(info, payload)
        target.writestr("xl/vbaProject.bin", b"opaque-macro-payload")
    result = _apply_office_patch_sequence_sync(str(path), _operations({
        "op": "sheet.rename", "sheet": "Source Data", "new_sheet": "Rates 2026",
    }))
    assert not result.get("error"), result
    with zipfile.ZipFile(io.BytesIO(result["_persisted_bytes"]), "r") as archive:
        assert archive.read("xl/vbaProject.bin") == b"opaque-macro-payload"


def test_invalid_sheet_rename_is_atomic(tmp_path):
    source = _generate_office_operations_sync("xlsx", _spreadsheet_rename_seed())["_persisted_bytes"]
    path = tmp_path / "invalid-rename.xlsx"
    path.write_bytes(source)
    before = path.read_bytes()
    result = _apply_office_patch_sequence_sync(str(path), _operations(
        {"op": "sheet.rename", "sheet": "Source Data", "new_sheet": "Rates 2026"},
        {"op": "sheet.rename", "sheet": "Rates 2026", "new_sheet": "Sheet"},
    ))
    assert result.get("error"), result
    assert result["operation_index"] == 1
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before
