import hashlib
import json
import sys
import zipfile
from pathlib import Path

import pytest


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "packages/core/ai/skills/pptx/scripts"
sys.path.insert(0, str(SCRIPTS))

import pptx_quality_gate  # noqa: E402


pptx = pytest.importorskip("pptx")
PIL = pytest.importorskip("PIL")
from PIL import Image, ImageDraw  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import PP_ALIGN  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402


def _add_text(slide, text, *, x, y, width, height, size):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    box.text_frame.clear()
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor(30, 36, 44)
    return box


def _editable_deck(path: Path) -> None:
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    blank = deck.slide_layouts[6]

    slide = deck.slides.add_slide(blank)
    _add_text(slide, "A decision-ready launch plan", x=0.9, y=1.0, width=11.5, height=1.0, size=54)
    _add_text(slide, "Evidence, choices, and the next move", x=0.9, y=3.0, width=8.0, height=0.6, size=20)
    slide.notes_slide.notes_text_frame.text = "Opening context."

    slide = deck.slides.add_slide(blank)
    _add_text(slide, "Demand is concentrated in one urgent workflow", x=0.8, y=0.45, width=11.8, height=0.6, size=38)
    _add_text(slide, "Operators lose a day each week reconciling fragmented requests.", x=0.8, y=2.0, width=5.8, height=1.0, size=18)
    slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(1.8), Inches(3.8), Inches(3.2))
    slide.notes_slide.notes_text_frame.text = "Evidence framing."

    slide = deck.slides.add_slide(blank)
    _add_text(slide, "Two options separate speed from control", x=0.8, y=0.45, width=11.8, height=0.6, size=38)
    _add_text(slide, "FAST\nLaunch in two weeks\nAccept manual review", x=0.8, y=2.0, width=5.2, height=2.2, size=18)
    _add_text(slide, "CONTROLLED\nLaunch in four weeks\nAutomate policy checks", x=7.2, y=2.0, width=5.2, height=2.2, size=18)
    slide.notes_slide.notes_text_frame.text = "[Sources]\nhttps://example.com/evidence"

    slide = deck.slides.add_slide(blank)
    _add_text(slide, "Approve a controlled pilot with one accountable owner", x=0.8, y=0.45, width=11.8, height=0.6, size=38)
    _add_text(slide, "Start Monday. Review evidence after 30 days.", x=2.0, y=2.8, width=9.3, height=0.8, size=24)
    slide.notes_slide.notes_text_frame.text = "Closing action."

    deck.save(path)


def _project_evidence(project: Path, slides: int) -> None:
    (project / "svg_output").mkdir(parents=True)
    (project / "notes").mkdir()
    (project / "design_spec.md").write_text("# Design\n", encoding="utf-8")
    (project / "spec_lock.md").write_text("# Lock\n", encoding="utf-8")
    plan = {
        "version": 1,
        "canvas": {"width": 1280, "height": 720},
        "slides": {},
    }
    families = ["split-hero", "claim-evidence", "comparison", "decision-split"]
    visuals = ["image", "diagram", "table", "metric"]
    for index in range(1, slides + 1):
        family = families[(index - 1) % len(families)]
        visual = visuals[(index - 1) % len(visuals)]
        background = "dark" if index % 2 else "light"
        density = "sparse" if index in {1, slides} else "dense"
        text_roles = (
            ["deck-title", "dek"]
            if index == 1
            else ["decision", "body"]
            if index == slides
            else ["slide-title", "body"]
        )
        plan["slides"][f"P{index:02d}"] = {
            "role": "cover" if index == 1 else "evidence",
            "layout_family": family,
            "background": background,
            "dominant_visual": visual,
            "density": density,
            "text_roles": text_roles,
            "content_budget": {"title_lines": 1, "max_words": 60, "max_items": 5},
            "regions": [{"id": "main", "x": 64, "y": 48, "width": 1152, "height": 610}],
        }
        (project / "svg_output" / f"{index:02d}_page.svg").write_text(
            f'''<svg viewBox="0 0 1280 720" xmlns="http://www.w3.org/2000/svg"
 data-layout-family="{family}" data-background="{background}"
 data-dominant-visual="{visual}" data-density="{density}"
 data-text-roles="{','.join(text_roles)}">
  <text x="80" y="110" font-size="24" data-text-role="{text_roles[0]}">Measured sample</text>
  <text x="80" y="150" font-size="24" data-text-role="{text_roles[1]}">Supporting detail</text>
</svg>''',
            encoding="utf-8",
        )
        (project / "notes" / f"{index:02d}_page.md").write_text("Notes\n", encoding="utf-8")
    (project / "layout_plan.json").write_text(json.dumps(plan), encoding="utf-8")


def _render_evidence(render_dir: Path, slides: int) -> None:
    render_dir.mkdir(parents=True)
    for index in range(1, slides + 1):
        image = Image.new("RGB", (1280, 720), (245, 245 - index * 5, 235))
        draw = ImageDraw.Draw(image)
        draw.rectangle((60 * index, 80, 400 + 40 * index, 620), fill=(20 * index, 70, 120))
        image.save(render_dir / f"slide-{index}.png")


def _visual_inspection_receipt(
    project: Path,
    render_dir: Path,
    deck_path: Path,
    slides: int,
) -> None:
    checks = {name: "pass" for name in pptx_quality_gate.VISUAL_RECEIPT_CHECKS}
    render_hashes = {
        path.name: hashlib.sha256(path.read_bytes()).hexdigest()
        for path in sorted(render_dir.glob("slide-*.png"))
    }
    payload = {
        "version": 1,
        "pptx_sha256": hashlib.sha256(deck_path.read_bytes()).hexdigest(),
        "inspected_slides": list(range(1, slides + 1)),
        "full_size_reads": list(range(1, slides + 1)),
        "contact_sheet_read": True,
        "checks": checks,
        "defects": [],
        "render_sha256": render_hashes,
        "text_containment": {
            "status": "pass",
            "evidence_path": "qa/text-containment.json",
        },
    }
    (project / "qa").mkdir(parents=True, exist_ok=True)
    containment_slides = []
    for index in range(1, slides + 1):
        render_path = render_dir / f"slide-{index}.png"
        containment_slides.append(
            {
                "slide": index,
                "status": "pass",
                "method": "pixel-bbox" if index == 1 else "full-size-review",
                "source_pixel_dimensions": [2560, 1440] if index == 1 else None,
                "render_sha256": hashlib.sha256(render_path.read_bytes()).hexdigest(),
                "records": (
                    [
                        {
                            "text": "Measured sample",
                            "bbox": [80, 80, 300, 120],
                            "containing_region_bbox": [60, 60, 320, 140],
                            "parent_region_bbox": [0, 0, 1280, 720],
                            "minimum_required_padding": 16,
                            "actual_minimum_padding": 20,
                            "font_size": 24,
                            "minimum_font_size": 18,
                            "contained": True,
                        }
                    ]
                    if index == 1
                    else []
                ),
            }
        )
    (project / "qa" / "text-containment.json").write_text(
        json.dumps(
            {
                "version": 2,
                "pptx_sha256": hashlib.sha256(deck_path.read_bytes()).hexdigest(),
                "inspected_slides": list(range(1, slides + 1)),
                "slides": containment_slides,
            }
        ),
        encoding="utf-8",
    )
    (project / "qa" / "visual-inspection.json").write_text(
        json.dumps(payload),
        encoding="utf-8",
    )


def test_editable_deck_passes_final_quality_gate(tmp_path: Path) -> None:
    deck_path = tmp_path / "quality-deck.pptx"
    project = tmp_path / "project"
    render_dir = project / "qa" / "final-render"
    _editable_deck(deck_path)
    _project_evidence(project, 4)
    _render_evidence(render_dir, 4)

    result = pptx_quality_gate.run_gate(
        deck_path,
        mode="editable",
        project=project,
        render_dir=render_dir,
        min_score=90,
    )

    assert result["status"] == "pass", json.dumps(result, indent=2)
    assert result["quality_score"] >= 90
    assert result["slide_count"] == 4
    assert result["metrics"]["rendered_slide_count"] == 4


def _add_portability_table(slide, *, explicit_style: bool) -> None:
    table = slide.shapes.add_table(
        2,
        2,
        Inches(0.9),
        Inches(4.5),
        Inches(7.0),
        Inches(1.4),
    ).table
    for row, values in enumerate((("Dimension", "Effect"), ("Ownership", "Cleaner handoff"))):
        for column, value in enumerate(values):
            cell = table.cell(row, column)
            cell.text = value
            paragraph = cell.text_frame.paragraphs[0]
            if explicit_style:
                paragraph.alignment = PP_ALIGN.LEFT
                paragraph.runs[0].font.name = "Carlito"
                paragraph.runs[0].font.size = Pt(18)


def test_gate_rejects_table_text_that_depends_on_theme_font_defaults(tmp_path: Path) -> None:
    deck_path = tmp_path / "implicit-table-font.pptx"
    _editable_deck(deck_path)
    deck = Presentation(deck_path)
    _add_portability_table(deck.slides[1], explicit_style=False)
    deck.save(deck_path)

    result = pptx_quality_gate.run_gate(deck_path, mode="editable", min_score=0)

    assert any("Table text lacks an explicit typeface" in error for error in result["errors"])
    assert any("Table text lacks explicit paragraph alignment" in error for error in result["errors"])
    assert any(action["code"] == "repair-font-portability" for action in result["repair_actions"])


def test_gate_accepts_explicit_portable_table_typography(tmp_path: Path) -> None:
    deck_path = tmp_path / "portable-table-font.pptx"
    _editable_deck(deck_path)
    deck = Presentation(deck_path)
    _add_portability_table(deck.slides[1], explicit_style=True)
    deck.save(deck_path)

    result = pptx_quality_gate.run_gate(deck_path, mode="editable", min_score=0)

    assert not any("Table text lacks" in error for error in result["errors"])
    assert result["metrics"]["table_text_runs"] == 4
    assert result["metrics"]["table_text_runs_with_explicit_typeface"] == 4
    assert result["metrics"]["table_paragraphs_with_explicit_alignment"] == 4


def test_project_gate_requires_declared_native_chart_objects(tmp_path: Path) -> None:
    deck_path = tmp_path / "missing-native-chart.pptx"
    project = tmp_path / "project"
    render_dir = project / "qa" / "final-render"
    _editable_deck(deck_path)
    _project_evidence(project, 4)
    _render_evidence(render_dir, 4)
    (project / "native_charts.json").write_text(
        json.dumps(
            {
                "version": 1,
                "slides": {
                    "P02": [
                        {
                            "id": "evidence-chart",
                            "slot": "evidence-chart",
                            "type": "column",
                            "position": {"x": 100, "y": 180, "width": 700, "height": 400},
                            "categories": ["Before", "After"],
                            "series": [{"name": "Hours", "values": [46, 18]}],
                        }
                    ]
                },
            }
        ),
        encoding="utf-8",
    )

    result = pptx_quality_gate.run_gate(
        deck_path,
        mode="editable",
        project=project,
        render_dir=render_dir,
        min_score=0,
    )

    assert result["status"] == "fail"
    assert any("native chart manifest requires" in error for error in result["errors"])
    assert any(action["code"] == "repair-native-chart" for action in result["repair_actions"])


def test_project_gate_requires_declared_embedded_video_objects(tmp_path: Path) -> None:
    deck_path = tmp_path / "missing-embedded-video.pptx"
    project = tmp_path / "project"
    render_dir = project / "qa" / "final-render"
    _editable_deck(deck_path)
    _project_evidence(project, 4)
    _render_evidence(render_dir, 4)
    (project / "media").mkdir()
    (project / "media" / "demo.mp4").write_bytes(b"manifest validation checks packaging inputs")
    (project / "embedded_media.json").write_text(
        json.dumps(
            {
                "version": 1,
                "slides": {
                    "P02": [
                        {
                            "id": "demo-video",
                            "type": "video",
                            "source": "media/demo.mp4",
                            "position": {
                                "x": 100,
                                "y": 180,
                                "width": 960,
                                "height": 400,
                            },
                        }
                    ]
                },
            }
        ),
        encoding="utf-8",
    )

    result = pptx_quality_gate.run_gate(
        deck_path,
        mode="editable",
        project=project,
        render_dir=render_dir,
        min_score=0,
    )

    assert result["status"] == "fail"
    assert any("embedded video manifest requires" in error for error in result["errors"])
    assert any(action["code"] == "repair-embedded-media" for action in result["repair_actions"])


def test_gate_rejects_off_canvas_shape_and_placeholder(tmp_path: Path) -> None:
    deck_path = tmp_path / "broken.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _add_text(slide, "{{TITLE}}", x=0.8, y=0.5, width=7.0, height=0.8, size=54)
    _add_text(slide, "Invisible", x=14.0, y=2.0, width=2.0, height=1.0, size=18)
    deck.save(deck_path)

    result = pptx_quality_gate.run_gate(deck_path, min_score=0)

    assert result["status"] == "fail"
    assert any("unresolved placeholder" in error for error in result["errors"])
    assert any("entirely outside" in error for error in result["errors"])


def test_gate_rejects_text_that_crosses_slide_boundary(tmp_path: Path) -> None:
    deck_path = tmp_path / "clipped-title.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _add_text(
        slide,
        "This title is visibly clipped at the right edge",
        x=0.8,
        y=0.5,
        width=13.0,
        height=0.8,
        size=54,
    )
    deck.save(deck_path)

    result = pptx_quality_gate.run_gate(deck_path, min_score=0)

    assert result["status"] == "fail"
    assert any("Text crosses the slide boundary" in error for error in result["errors"])


def test_gate_ignores_text_fully_contained_by_visual_panel(tmp_path: Path) -> None:
    deck_path = tmp_path / "contained-label.pptx"
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _add_text(slide, "Contained label", x=1.2, y=1.1, width=3.2, height=0.6, size=38)
    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.9),
        Inches(0.85),
        Inches(3.8),
        Inches(1.1),
    )
    # Move the panel behind the label without changing the geometric relation.
    sp_tree = slide.shapes._spTree
    sp_tree.remove(panel._element)
    sp_tree.insert(2, panel._element)
    deck.save(deck_path)

    result = pptx_quality_gate.run_gate(deck_path, min_score=0)

    assert not any("Unintended text" in error for error in result["errors"])


def test_gate_rejects_text_text_overlap(tmp_path: Path) -> None:
    deck_path = tmp_path / "overlapping-labels.pptx"
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _add_text(slide, "Primary statement", x=1.0, y=1.0, width=5.5, height=1.0, size=38)
    _add_text(slide, "Secondary statement", x=2.0, y=1.25, width=5.5, height=1.0, size=28)
    deck.save(deck_path)

    result = pptx_quality_gate.run_gate(deck_path, min_score=0)

    assert any("Unintended text or partial-object overlap" in error for error in result["errors"])
    defect = next(
        item for item in result["metrics"]["layout_defects"]
        if item["kind"] == "unintended-overlap"
    )
    assert defect["slide"] == 1
    assert {shape["text"] for shape in defect["shapes"]} == {
        "Primary statement",
        "Secondary statement",
    }


def test_gate_rejects_text_that_cannot_fit_its_text_frame(tmp_path: Path) -> None:
    deck_path = tmp_path / "text-frame-overflow.pptx"
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _add_text(
        slide,
        "Text frame capacity validation",
        x=0.8,
        y=0.45,
        width=9.5,
        height=0.7,
        size=38,
    )
    overflow = _add_text(
        slide,
        "This sentence cannot possibly fit inside such a narrow fixed frame.",
        x=1.0,
        y=2.0,
        width=1.5,
        height=0.45,
        size=24,
    )
    overflow.text_frame.word_wrap = False
    deck.save(deck_path)

    result = pptx_quality_gate.run_gate(deck_path, min_score=0)

    assert any("does not fit inside" in error for error in result["errors"])
    defect = next(
        item
        for item in result["metrics"]["layout_defects"]
        if item["kind"] == "text-frame-overflow"
    )
    assert defect["fit"]["axes"] == ["horizontal"]
    assert defect["fit"]["required_width_pt"] > defect["fit"]["available_width_pt"]
    assert any(
        action["code"] == "repair-text-frame-fit"
        for action in result["repair_actions"]
    )


def test_gate_rejects_distributed_alignment_and_excessive_tracking(
    tmp_path: Path,
) -> None:
    deck_path = tmp_path / "distributed-text.pptx"
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = _add_text(
        slide,
        "Character spacing stays controlled",
        x=0.8,
        y=0.45,
        width=10.5,
        height=0.7,
        size=38,
    )
    paragraph = title.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.DISTRIBUTE
    paragraph.runs[0]._r.get_or_add_rPr().set("spc", "240")
    deck.save(deck_path)

    result = pptx_quality_gate.run_gate(deck_path, min_score=0)

    assert any("Distributed paragraph alignment" in error for error in result["errors"])
    assert any("Character spacing exceeds" in error for error in result["errors"])
    assert result["metrics"]["distributed_alignment_slides"] == [1]
    assert result["metrics"]["excessive_character_spacing_slides"] == [1]
    assert any(
        action["code"] == "repair-character-spacing"
        for action in result["repair_actions"]
    )


def test_gate_rejects_visual_object_layered_over_text(tmp_path: Path) -> None:
    deck_path = tmp_path / "foreground-object-over-text.pptx"
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _add_text(slide, "Foreground collision", x=1.0, y=1.0, width=5.5, height=1.0, size=38)
    slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(2.0),
        Inches(1.2),
        Inches(2.5),
        Inches(0.7),
    )
    deck.save(deck_path)

    result = pptx_quality_gate.run_gate(deck_path, min_score=0)

    assert any("Unintended text or partial-object overlap" in error for error in result["errors"])


def test_gate_rejects_text_outside_safe_boundary(tmp_path: Path) -> None:
    deck_path = tmp_path / "unsafe-margin.pptx"
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _add_text(slide, "Too close to the edge", x=0.05, y=0.5, width=5.0, height=0.8, size=38)
    deck.save(deck_path)

    result = pptx_quality_gate.run_gate(deck_path, min_score=0)

    assert any("safe boundary" in error for error in result["errors"])
    defect = next(
        item for item in result["metrics"]["layout_defects"]
        if item["kind"] == "text-safe-boundary"
    )
    assert defect["shapes"][0]["text"] == "Too close to the edge"


def test_gate_rejects_non_text_shape_crossing_boundary(tmp_path: Path) -> None:
    deck_path = tmp_path / "shape-overflow.pptx"
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _add_text(slide, "Boundary validation", x=0.8, y=0.5, width=8.0, height=0.8, size=38)
    slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(12.8),
        Inches(2.0),
        Inches(1.0),
        Inches(2.0),
    )
    deck.save(deck_path)

    result = pptx_quality_gate.run_gate(deck_path, min_score=0)

    assert any("Non-text shapes cross the slide boundary" in error for error in result["errors"])


def test_full_page_image_mode_requires_one_image_and_no_text(tmp_path: Path) -> None:
    image_path = tmp_path / "page.png"
    Image.new("RGB", (1280, 720), (18, 42, 70)).save(image_path)
    deck_path = tmp_path / "image-deck.pptx"
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(image_path), 0, 0, width=deck.slide_width, height=deck.slide_height)
    deck.save(deck_path)

    passing = pptx_quality_gate.run_gate(deck_path, mode="full_page_image")
    assert passing["status"] == "pass"

    deck = Presentation(str(deck_path))
    _add_text(deck.slides[0], "Extra editable text", x=1, y=1, width=4, height=1, size=20)
    deck.save(deck_path)
    failing = pptx_quality_gate.run_gate(deck_path, mode="full_page_image", min_score=0)
    assert failing["status"] == "fail"
    assert any("editable text or extra shapes" in error for error in failing["errors"])


def test_full_page_image_project_requires_layout_and_svg_evidence(tmp_path: Path) -> None:
    image_path = tmp_path / "page.png"
    Image.new("RGB", (1280, 720), (18, 42, 70)).save(image_path)
    deck_path = tmp_path / "image-deck.pptx"
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(image_path), 0, 0, width=deck.slide_width, height=deck.slide_height)
    deck.save(deck_path)
    project = tmp_path / "project"
    project.mkdir()

    result = pptx_quality_gate.run_gate(
        deck_path,
        mode="full_page_image",
        project=project,
        min_score=0,
    )

    assert result["status"] == "fail"
    assert any("layout_plan.json" in error for error in result["errors"])
    assert any("svg_output contains 0 page(s), expected 1" in error for error in result["errors"])


def test_full_page_image_project_requires_pixel_bound_visual_receipt(tmp_path: Path) -> None:
    image_path = tmp_path / "page.png"
    Image.new("RGB", (2560, 1440), (18, 42, 70)).save(image_path)
    deck_path = tmp_path / "image-deck.pptx"
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(image_path), 0, 0, width=deck.slide_width, height=deck.slide_height)
    deck.save(deck_path)
    project = tmp_path / "project"
    render_dir = project / "qa" / "final-render"
    _project_evidence(project, 1)
    _render_evidence(render_dir, 1)

    missing = pptx_quality_gate.run_gate(
        deck_path,
        mode="full_page_image",
        project=project,
        render_dir=render_dir,
        min_score=0,
    )
    assert missing["status"] == "fail"
    assert any("Visual inspection receipt missing" in error for error in missing["errors"])
    assert any(action["code"] == "repair-visual-inspection" for action in missing["repair_actions"])

    _visual_inspection_receipt(project, render_dir, deck_path, 1)
    passing = pptx_quality_gate.run_gate(
        deck_path,
        mode="full_page_image",
        project=project,
        render_dir=render_dir,
        min_score=90,
    )
    assert passing["status"] == "pass", json.dumps(passing, indent=2)
    assert passing["metrics"]["visual_inspection_receipt"]["inspected_slide_count"] == 1

    svg_path = project / "svg_output" / "01_page.svg"
    original_svg = svg_path.read_text(encoding="utf-8")
    svg_path.write_text(
        original_svg.replace(
            '<text x="80" y="110" font-size="24" data-text-role="deck-title">Measured sample</text>',
            '<image href="page.png" x="0" y="0" width="1280" height="720"/>',
        ).replace(
            '<text x="80" y="150" font-size="24" data-text-role="dek">Supporting detail</text>',
            "",
        ),
        encoding="utf-8",
    )
    flattened_source = pptx_quality_gate.run_gate(
        deck_path,
        mode="full_page_image",
        project=project,
        render_dir=render_dir,
        min_score=0,
    )
    assert flattened_source["status"] == "fail"
    assert any("source SVG has no text primitives" in error for error in flattened_source["errors"])
    svg_path.write_text(original_svg, encoding="utf-8")

    Image.new("RGB", (1280, 720), (200, 20, 20)).save(render_dir / "slide-1.png")
    stale = pptx_quality_gate.run_gate(
        deck_path,
        mode="full_page_image",
        project=project,
        render_dir=render_dir,
        min_score=0,
    )
    assert stale["status"] == "fail"
    assert any("render hashes are missing or stale" in error for error in stale["errors"])


def test_full_page_image_receipt_requires_connector_routing_review(tmp_path: Path) -> None:
    image_path = tmp_path / "page.png"
    Image.new("RGB", (2560, 1440), (18, 42, 70)).save(image_path)
    deck_path = tmp_path / "image-deck.pptx"
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(image_path), 0, 0, width=deck.slide_width, height=deck.slide_height)
    deck.save(deck_path)
    project = tmp_path / "project"
    render_dir = project / "qa" / "final-render"
    _project_evidence(project, 1)
    _render_evidence(render_dir, 1)
    _visual_inspection_receipt(project, render_dir, deck_path, 1)

    receipt_path = project / "qa" / "visual-inspection.json"
    receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
    receipt["checks"]["connector_routing"] = "fail"
    receipt_path.write_text(json.dumps(receipt), encoding="utf-8")

    result = pptx_quality_gate.run_gate(
        deck_path,
        mode="full_page_image",
        project=project,
        render_dir=render_dir,
        min_score=0,
    )

    assert result["status"] == "fail"
    assert any("connector_routing" in error for error in result["errors"])


def test_full_page_image_receipt_requires_text_containment_review(tmp_path: Path) -> None:
    image_path = tmp_path / "page.png"
    Image.new("RGB", (2560, 1440), (18, 42, 70)).save(image_path)
    deck_path = tmp_path / "image-deck.pptx"
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(image_path), 0, 0, width=deck.slide_width, height=deck.slide_height)
    deck.save(deck_path)
    project = tmp_path / "project"
    render_dir = project / "qa" / "final-render"
    _project_evidence(project, 1)
    _render_evidence(render_dir, 1)
    _visual_inspection_receipt(project, render_dir, deck_path, 1)

    receipt_path = project / "qa" / "visual-inspection.json"
    receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
    receipt["checks"]["text_containment"] = "fail"
    receipt_path.write_text(json.dumps(receipt), encoding="utf-8")

    result = pptx_quality_gate.run_gate(
        deck_path,
        mode="full_page_image",
        project=project,
        render_dir=render_dir,
        min_score=0,
    )

    assert result["status"] == "fail"
    assert any("text_containment" in error for error in result["errors"])


def test_full_page_image_receipt_rejects_partial_text_containment_evidence(
    tmp_path: Path,
) -> None:
    image_path = tmp_path / "page.png"
    Image.new("RGB", (2560, 1440), (18, 42, 70)).save(image_path)
    deck_path = tmp_path / "image-deck.pptx"
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    for _index in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_picture(
            str(image_path), 0, 0, width=deck.slide_width, height=deck.slide_height
        )
    deck.save(deck_path)
    project = tmp_path / "project"
    render_dir = project / "qa" / "final-render"
    _project_evidence(project, 2)
    _render_evidence(render_dir, 2)
    _visual_inspection_receipt(project, render_dir, deck_path, 2)

    evidence_path = project / "qa" / "text-containment.json"
    evidence = json.loads(evidence_path.read_text(encoding="utf-8"))
    evidence["inspected_slides"] = [1]
    evidence["slides"] = evidence["slides"][:1]
    evidence_path.write_text(json.dumps(evidence), encoding="utf-8")

    result = pptx_quality_gate.run_gate(
        deck_path,
        mode="full_page_image",
        project=project,
        render_dir=render_dir,
        min_score=0,
    )

    assert result["status"] == "fail"
    assert any("every final slide" in error for error in result["errors"])
    assert any("no slide record for: 2" in error for error in result["errors"])


def test_full_page_image_receipt_rejects_text_container_outside_parent(
    tmp_path: Path,
) -> None:
    image_path = tmp_path / "page.png"
    Image.new("RGB", (1280, 720), (18, 42, 70)).save(image_path)
    deck_path = tmp_path / "image-deck.pptx"
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(
        str(image_path), 0, 0, width=deck.slide_width, height=deck.slide_height
    )
    deck.save(deck_path)
    project = tmp_path / "project"
    render_dir = project / "qa" / "final-render"
    _project_evidence(project, 1)
    _render_evidence(render_dir, 1)
    _visual_inspection_receipt(project, render_dir, deck_path, 1)

    evidence_path = project / "qa" / "text-containment.json"
    evidence = json.loads(evidence_path.read_text(encoding="utf-8"))
    record = evidence["slides"][0]["records"][0]
    record.update(
        {
            "bbox": [940, 634, 1100, 654],
            "containing_region_bbox": [932, 614, 1192, 668],
            "parent_region_bbox": [884, 180, 1224, 650],
            "minimum_required_padding": 6,
            "actual_minimum_padding": 14,
            "font_size": 20,
            "minimum_font_size": 22,
        }
    )
    evidence_path.write_text(json.dumps(evidence), encoding="utf-8")

    result = pptx_quality_gate.run_gate(
        deck_path,
        mode="full_page_image",
        project=project,
        render_dir=render_dir,
        min_score=0,
    )

    assert result["status"] == "fail"
    assert any("containing region escapes its parent" in error for error in result["errors"])
    assert any("misses minimum font size" in error for error in result["errors"])
    assert any("requires at least 192 DPI" in error for error in result["errors"])


def test_project_gate_rejects_svg_text_that_escapes_its_card(tmp_path: Path) -> None:
    deck_path = tmp_path / "svg-overflow.pptx"
    project = tmp_path / "project"
    render_dir = project / "qa" / "final-render"
    _editable_deck(deck_path)
    _project_evidence(project, 4)
    _render_evidence(render_dir, 4)
    (project / "svg_output" / "02_page.svg").write_text(
        '''<svg viewBox="0 0 1280 720" xmlns="http://www.w3.org/2000/svg"
 data-layout-family="claim-evidence" data-background="light"
 data-dominant-visual="diagram" data-density="dense">
  <rect x="0" y="0" width="1280" height="720" fill="#FFFFFF"/>
  <text x="72" y="90" fill="#111827" font-size="48">Readable slide title</text>
  <rect x="72" y="170" width="300" height="260" fill="#F4F6F8"/>
  <text x="96" y="230" fill="#111827" font-size="24">This body sentence visibly escapes the narrow card.</text>
</svg>''',
        encoding="utf-8",
    )

    result = pptx_quality_gate.run_gate(
        deck_path,
        mode="editable",
        project=project,
        render_dir=render_dir,
        min_score=0,
    )

    assert result["status"] == "fail"
    assert result["metrics"]["source_svg_error_count"] >= 1
    assert any("SVG quality" in error and "Text fit failure" in error for error in result["errors"])


def test_project_gate_rejects_blank_rendered_svg_media_region(tmp_path: Path) -> None:
    deck_path = tmp_path / "blank-media.pptx"
    project = tmp_path / "project"
    render_dir = project / "qa" / "final-render"
    _editable_deck(deck_path)
    _project_evidence(project, 4)
    _render_evidence(render_dir, 4)
    (project / "images").mkdir()
    Image.new("RGB", (640, 360), (40, 110, 190)).save(project / "images" / "evidence.png")
    (project / "svg_output" / "02_page.svg").write_text(
        '''<svg viewBox="0 0 1280 720" xmlns="http://www.w3.org/2000/svg"
 data-layout-family="claim-evidence" data-background="light"
 data-dominant-visual="image" data-density="dense">
  <rect x="0" y="0" width="1280" height="720" fill="#FFFFFF"/>
  <image href="../images/evidence.png" x="800" y="100" width="350" height="250"/>
</svg>''',
        encoding="utf-8",
    )

    result = pptx_quality_gate.run_gate(
        deck_path,
        mode="editable",
        project=project,
        render_dir=render_dir,
        min_score=0,
    )

    assert result["status"] == "fail"
    assert result["metrics"]["rendered_svg_media_slots"] == 1
    assert any("media region appears blank" in error for error in result["errors"])
    assert any(action["code"] == "repair-blank-media" for action in result["repair_actions"])


def test_gate_rejects_non_footer_text_below_16_points(tmp_path: Path) -> None:
    deck_path = tmp_path / "undersized-body.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _add_text(slide, "Readable decisions need readable type", x=0.8, y=0.5, width=11.4, height=0.8, size=54)
    _add_text(slide, "This body copy is too small for presentation use.", x=0.8, y=2.0, width=8, height=0.8, size=14)
    deck.save(deck_path)

    result = pptx_quality_gate.run_gate(deck_path, min_score=0)

    assert result["status"] == "fail"
    assert any("Body or label text below 16 pt" in error for error in result["errors"])


def test_gate_rejects_title_that_cannot_fit_its_text_box(tmp_path: Path) -> None:
    deck_path = tmp_path / "wrapped-title.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _add_text(
        slide,
        "A strategically important title that cannot remain on one line",
        x=0.8,
        y=0.5,
        width=5.0,
        height=0.8,
        size=54,
    )
    deck.save(deck_path)

    result = pptx_quality_gate.run_gate(deck_path, min_score=0)

    assert result["status"] == "fail"
    assert any("title needs approximately" in error for error in result["errors"])


def test_auto_mode_detects_full_page_images_and_records_exact_hash(tmp_path: Path) -> None:
    image_path = tmp_path / "page.png"
    Image.new("RGB", (1280, 720), (18, 42, 70)).save(image_path)
    deck_path = tmp_path / "image-deck.pptx"
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(image_path), 0, 0, width=deck.slide_width, height=deck.slide_height)
    deck.save(deck_path)

    result = pptx_quality_gate.run_gate(deck_path, mode="auto")

    assert result["status"] == "pass"
    assert result["mode"] == "full_page_image"
    assert result["metrics"]["pptx_sha256"]
    assert result["metrics"]["pptx_size_bytes"] == deck_path.stat().st_size


def test_relationship_validator_rejects_missing_internal_target(tmp_path: Path) -> None:
    package = tmp_path / "broken.pptx"
    rels = """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="example" Target="ppt/missing.xml"/>
</Relationships>"""
    with zipfile.ZipFile(package, "w") as archive:
        archive.writestr("_rels/.rels", rels)

    errors = pptx_quality_gate.validate_relationships(package)

    assert errors == ["_rels/.rels: missing target ppt/missing.xml"]
