"""Native Excel conditional formatting shares generation, template and patch execution."""
from __future__ import annotations

import io
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule, Rule
from openpyxl.styles import Alignment, PatternFill, Protection
from openpyxl.styles.differential import DifferentialStyle
import pytest

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools.file_tools import _apply_office_patch_sequence_sync
from packages.core.contracts.file_engine import file_patch_operations, normalize_file_patch_operation
from packages.core.services.file_engine_patches import describe_file_structure


def normalized(*operations):
    return [normalize_file_patch_operation(operation) for operation in operations]


def apply(path, *operations):
    return _apply_office_patch_sequence_sync(str(path), normalized(*operations))


def threshold(kind, value=None, **extra):
    result = {"type": kind, **extra}
    if value is not None:
        result["value"] = value
    return result


def test_conditional_format_operations_are_shared_xlsx_xlsm_capabilities():
    expected = {
        "conditional_format.insert",
        "conditional_format.format",
        "conditional_format.delete",
    }
    assert expected <= set(file_patch_operations("xlsx"))
    assert expected <= set(file_patch_operations("xlsm"))


def test_generation_creates_all_supported_native_conditional_format_types(tmp_path):
    result = _generate_office_operations_sync(
        "xlsx",
        normalized(
            {
                "op": "conditional_format.insert",
                "range": "B2:B10",
                "conditional_format": {
                    "type": "cell",
                    "operator": "greaterThan",
                    "formulas": [0],
                    "stop_if_true": True,
                    "format": {
                        "bold": True,
                        "font_color": "FFFFFF",
                        "fill_color": "C00000",
                        "borders": {"bottom": {"style": "thin", "color": "800000"}},
                    },
                },
            },
            {
                "op": "conditional_format.insert",
                "range": "C2:C10",
                "conditional_format": {
                    "type": "formula",
                    "formulas": ["=MOD(C2,2)=0"],
                    "format": {"italic": True, "font_color": "006100"},
                },
            },
            {
                "op": "conditional_format.insert",
                "range": "D2:D10",
                "conditional_format": {
                    "type": "color_scale",
                    "thresholds": [
                        threshold("min"),
                        threshold("percentile", 50),
                        threshold("max"),
                    ],
                    "colors": ["F8696B", "FFEB84", "63BE7B"],
                },
            },
            {
                "op": "conditional_format.insert",
                "range": "E2:E10",
                "conditional_format": {
                    "type": "data_bar",
                    "thresholds": [threshold("num", 0), threshold("num", 100)],
                    "color": "638EC6",
                    "show_value": False,
                    "min_length": 5,
                    "max_length": 95,
                },
            },
            {
                "op": "conditional_format.insert",
                "range": "F2:F10",
                "conditional_format": {
                    "type": "icon_set",
                    "icon_style": "3Arrows",
                    "thresholds": [
                        threshold("percent", 0),
                        threshold("percent", 33),
                        threshold("percent", 67),
                    ],
                    "show_value": False,
                    "reverse": True,
                    "percent": True,
                },
            },
            {
                "op": "conditional_format.format",
                "conditional_format_index": 3,
                "conditional_format": {"show_value": True, "max_length": 90},
            },
        ),
    )
    assert not result.get("error"), result
    path = tmp_path / "conditional-formats.xlsx"
    path.write_bytes(result["_persisted_bytes"])
    workbook = load_workbook(path)
    rules = [
        rule
        for rules in workbook.active.conditional_formatting._cf_rules.values()
        for rule in rules
    ]
    assert [rule.type for rule in rules] == [
        "cellIs",
        "expression",
        "colorScale",
        "dataBar",
        "iconSet",
    ]
    assert rules[0].operator == "greaterThan" and rules[0].formula == ["0"]
    assert rules[0].dxf.font.b is True and rules[0].dxf.fill.fgColor.rgb == "FFC00000"
    assert rules[1].formula == ["MOD(C2,2)=0"] and rules[1].dxf.font.i is True
    assert len(rules[2].colorScale.cfvo) == 3 and len(rules[2].colorScale.color) == 3
    assert rules[3].dataBar.showValue is True and rules[3].dataBar.maxLength == 90
    assert rules[4].iconSet.iconSet == "3Arrows" and rules[4].iconSet.reverse is True
    workbook.close()

    structure = describe_file_structure(str(path))
    assert structure["conditional_format_count"] == 5
    descriptions = structure["sheets"][0]["conditional_formats"]
    assert [item["index"] for item in descriptions] == list(range(5))
    assert descriptions[0]["format"]["fill_color"] == "C00000"
    assert descriptions[2]["thresholds"][1] == {"type": "percentile", "value": 50.0}
    assert descriptions[4]["icon_style"] == "3Arrows"


def test_template_partial_format_and_range_move_preserve_sibling_and_opaque_style(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    dxf = DifferentialStyle(
        alignment=Alignment(horizontal="center"),
        protection=Protection(locked=False),
    )
    first = Rule(
        type="expression",
        formula=["B2>0"],
        stopIfTrue=True,
        text="opaque-native-text",
        dxf=dxf,
    )
    second = FormulaRule(formula=["B2<0"], fill=PatternFill("solid", fgColor="FFFF0000"))
    sheet.conditional_formatting.add("B2:B10", first)
    sheet.conditional_formatting.add("B2:B10", second)
    path = tmp_path / "template.xlsx"
    workbook.save(path)
    workbook.close()

    result = apply(
        path,
        {
            "op": "conditional_format.format",
            "conditional_format_index": 0,
            "range": "C2:C10",
            "conditional_format": {
                "stop_if_true": False,
                "format": {"fill_color": "FFF2CC", "bold": True},
            },
        },
    )
    assert not result.get("error"), result
    edited = load_workbook(io.BytesIO(result["_persisted_bytes"]))
    containers = list(edited.active.conditional_formatting._cf_rules.items())
    assert len(containers) == 2
    assert str(containers[0][0].sqref) == "B2:B10"
    assert len(containers[0][1]) == 1 and containers[0][1][0].formula == ["B2<0"]
    assert str(containers[1][0].sqref) == "C2:C10"
    moved = containers[1][1][0]
    assert moved.formula == ["B2>0"] and moved.stopIfTrue is False
    assert moved.priority == 1 and moved.text == "opaque-native-text"
    assert moved.dxf.alignment.horizontal == "center" and moved.dxf.protection.locked is False
    assert moved.dxf.font.b is True and moved.dxf.fill.fgColor.rgb == "FFFFF2CC"
    edited.close()


def test_range_only_move_preserves_unsupported_template_rule(tmp_path):
    workbook = Workbook()
    rule = Rule(type="top10", rank=10, percent=True, bottom=True)
    workbook.active.conditional_formatting.add("A2:A20", rule)
    path = tmp_path / "unsupported-template-rule.xlsx"
    workbook.save(path)
    workbook.close()

    result = apply(
        path,
        {
            "op": "conditional_format.format",
            "conditional_format_index": 0,
            "range": "D2:D20",
        },
    )
    assert not result.get("error"), result
    edited = load_workbook(io.BytesIO(result["_persisted_bytes"]))
    container, rules = next(iter(edited.active.conditional_formatting._cf_rules.items()))
    assert str(container.sqref) == "D2:D20"
    assert rules[0].type == "top10" and rules[0].rank == 10
    assert rules[0].percent is True and rules[0].bottom is True
    edited.close()


def test_delete_removes_only_selected_rule_from_shared_range(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.conditional_formatting.add("A2:A10", FormulaRule(formula=["A2>0"]))
    sheet.conditional_formatting.add("A2:A10", FormulaRule(formula=["A2<0"]))
    path = tmp_path / "delete.xlsx"
    workbook.save(path)
    workbook.close()

    result = apply(path, {"op": "conditional_format.delete", "conditional_format_index": 0})
    assert not result.get("error"), result
    edited = load_workbook(io.BytesIO(result["_persisted_bytes"]))
    containers = list(edited.active.conditional_formatting._cf_rules.items())
    assert len(containers) == 1 and str(containers[0][0].sqref) == "A2:A10"
    assert len(containers[0][1]) == 1 and containers[0][1][0].formula == ["A2<0"]
    edited.close()


@pytest.mark.parametrize(
    "conditional_format",
    [
        {"type": "unknown"},
        {"type": "cell", "operator": "between", "formulas": [1]},
        {"type": "formula", "formulas": []},
        {
            "type": "color_scale",
            "thresholds": [threshold("min"), threshold("max")],
            "colors": ["FF0000"],
        },
        {
            "type": "color_scale",
            "thresholds": [threshold("max"), threshold("min")],
            "colors": ["FF0000", "00FF00"],
        },
        {
            "type": "data_bar",
            "thresholds": [threshold("num", 0), threshold("num", 100)],
            "color": "not-rgb",
        },
        {
            "type": "icon_set",
            "icon_style": "3Arrows",
            "thresholds": [threshold("percent", 0), threshold("percent", 50)],
        },
        {
            "type": "formula",
            "formulas": ["A1>0"],
            "format": {"borders": {"left": {"style": "thin"}}},
        },
    ],
)
def test_invalid_conditional_format_batch_is_atomic(tmp_path, conditional_format):
    workbook = Workbook()
    workbook.active["A1"] = "Keep"
    path = tmp_path / "invalid.xlsx"
    workbook.save(path)
    workbook.close()
    before = path.read_bytes()

    result = apply(
        path,
        {"op": "cell.set", "cell": "A1", "value": "Rollback"},
        {
            "op": "conditional_format.insert",
            "range": "B2:B10",
            "conditional_format": conditional_format,
        },
    )
    assert result.get("operation_index") == 1
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before


def test_xlsm_conditional_format_patch_preserves_macro_payload(tmp_path):
    created = _generate_office_operations_sync(
        "xlsx",
        normalized({"op": "cell.set", "cell": "A1", "value": "Keep"}),
    )
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(created["_persisted_bytes"])) as archive, zipfile.ZipFile(output, "w") as target:
        for item in archive.infolist():
            payload = archive.read(item.filename)
            if item.filename == "[Content_Types].xml":
                payload = payload.replace(
                    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                )
            target.writestr(item, payload)
        target.writestr("xl/vbaProject.bin", b"opaque-macro-payload")
    path = tmp_path / "macro.xlsm"
    path.write_bytes(output.getvalue())

    result = apply(
        path,
        {
            "op": "conditional_format.insert",
            "range": "A2:A9",
            "conditional_format": {
                "type": "formula",
                "formulas": ["=A2<>\"\""],
                "format": {"fill_color": "D9EAD3"},
            },
        },
    )
    assert not result.get("error"), result
    with zipfile.ZipFile(io.BytesIO(result["_persisted_bytes"])) as archive:
        assert archive.read("xl/vbaProject.bin") == b"opaque-macro-payload"
        sheet_xml = archive.read("xl/worksheets/sheet1.xml")
        assert b"conditionalFormatting" in sheet_xml and b"A2:A9" in sheet_xml
