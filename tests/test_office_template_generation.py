"""Template-based generation is patching a snapshot into a new Knowledge identity."""
from __future__ import annotations

import copy
import hashlib
import io
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from sqlalchemy import select

from packages.core.ai.runtime import file_actions, generated_files
from packages.core.ai.runtime.file_contracts import FileApprovalOperationFactory
from packages.core.ai.tools import file_tools
from packages.core.ai.tools.generate_file import document, tool
from packages.core.config import get_settings
from packages.core.contracts.file_engine import normalize_file_patch_operation
from packages.core.models.base import generate_ulid
from packages.core.models.event import EventLog
from tests.test_knowledge_file_consistency import _doc_for, _seed_workspace, fs_enabled  # noqa: F401
from tests.test_office_operation_generation import office_operations


def template_patch(extension, text="Template client"):
    if extension in {"xlsx", "xlsm"}:
        return {"op": "cell.set", "sheet": "Sheet", "cell": "A1", "value": text}
    return {"op": "text.set", "index": 0, "text": text, **({"slide": 1, "shape_id": 2} if extension == "pptx" else {})}


def template_bytes(extension):
    if extension == "xlsm":
        return macro_template_bytes()
    data = generated_files._generate_office_operations_sync(
        extension, [normalize_file_patch_operation(op) for op in office_operations(extension)],
    )["_persisted_bytes"]
    output = io.BytesIO()
    if extension == "docx":
        from docx import Document
        from docx.shared import Pt

        office = Document(io.BytesIO(data))
        office.sections[0].header.paragraphs[0].text = "Keep letterhead"
        office.sections[0].footer.paragraphs[0].text = "Keep footer"
        office.sections[0].left_margin = Pt(45)
        office.paragraphs[0].runs[0].font.size = Pt(28)
    elif extension == "pptx":
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Pt

        office = Presentation(io.BytesIO(data))
        office.slides[0].notes_slide.notes_text_frame.text = "Keep speaker notes"
        office.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.size = Pt(28)
        chart_data = CategoryChartData()
        chart_data.categories = ["Jan", "Feb"]
        chart_data.add_series("Revenue", [12, 15])
        office.slides[0].shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(450), Pt(100), Pt(240), Pt(200), chart_data)
    else:
        from openpyxl import load_workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.worksheet.datavalidation import DataValidation

        office = load_workbook(io.BytesIO(data))
        sheet = office["Sheet"]
        sheet["A3"], sheet["B3"], sheet["A4"], sheet["B4"] = "Jan", 12, "Feb", 15
        sheet.freeze_panes = "B3"
        sheet.column_dimensions["A"].width = 38
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=2, min_row=3, max_row=4))
        chart.set_categories(Reference(sheet, min_col=1, min_row=3, max_row=4))
        sheet.add_chart(chart, "D3")
        validation = DataValidation(type="whole", operator="between", formula1=0, formula2=100)
        sheet.add_data_validation(validation)
        validation.add("B3:B4")
    office.save(output)
    if extension == "xlsx":
        office.close()
    return output.getvalue()


def macro_template_bytes() -> bytes:
    source = template_bytes("xlsx")
    output = io.BytesIO()
    with ZipFile(io.BytesIO(source), "r") as archive, ZipFile(output, "w") as target:
        for info in archive.infolist():
            payload = archive.read(info.filename)
            if info.filename == "[Content_Types].xml":
                payload = payload.replace(
                    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                )
            target.writestr(info, payload)
        target.writestr("xl/vbaProject.bin", b"opaque-template-macro-payload")
    return output.getvalue()


def assert_template_result(data, extension, text):
    if extension == "docx":
        from docx import Document

        office = Document(io.BytesIO(data))
        assert office.paragraphs[0].text == text
        assert office.paragraphs[0].style.name == "Heading 1"
        assert office.paragraphs[0].runs[0].font.size.pt == 28
        assert office.sections[0].header.paragraphs[0].text == "Keep letterhead"
        assert office.sections[0].footer.paragraphs[0].text == "Keep footer"
        assert office.sections[0].left_margin.pt == 45
        assert office.tables[0].cell(1, 1).text == "12"
    elif extension == "pptx":
        from pptx import Presentation

        office = Presentation(io.BytesIO(data))
        assert office.slides[0].shapes[0].text == text
        assert office.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.size.pt == 28
        assert office.slides[0].shapes[0].left.pt == 72
        assert office.slides[0].notes_slide.notes_text_frame.text == "Keep speaker notes"
        assert list(office.slides[0].shapes[1].chart.series[0].values) == [12, 15]
    else:
        from openpyxl import load_workbook

        office = load_workbook(io.BytesIO(data))
        try:
            sheet = office["Sheet"]
            assert office.sheetnames == ["Sheet", "Summary"]
            assert sheet["A1"].value == text
            assert sheet["B1"].value == "=2*6"
            assert sheet["A1"].font.bold
            assert sheet["A1"].fill.fgColor.rgb == "FF123ABC"
            assert sheet.column_dimensions["A"].width == 38
            assert sheet.freeze_panes == "B3"
            assert len(sheet._charts) == 1
            assert len(sheet.data_validations.dataValidation) == 1
        finally:
            office.close()
        if extension == "xlsm":
            with ZipFile(io.BytesIO(data), "r") as archive:
                assert archive.read("xl/vbaProject.bin") == b"opaque-template-macro-payload"


@pytest.mark.parametrize("extension", ["docx", "pptx", "xlsx", "xlsm"])
def test_template_executor_preserves_existing_objects_and_supports_followup(tmp_path, monkeypatch, extension):
    original = template_bytes(extension)
    operations = [normalize_file_patch_operation(template_patch(extension))]
    applied = []
    real_apply = file_tools._apply_office_patch_sequence_sync

    def apply(path, patches):
        assert Path(path).read_bytes() == original
        applied.append(patches)
        return real_apply(path, patches)

    monkeypatch.setattr(file_tools, "_apply_office_patch_sequence_sync", apply)
    result = generated_files._generate_office_operations_sync(extension, operations, template_bytes=original)
    assert not result.get("error"), result
    assert applied == [operations]
    assert_template_result(result["_persisted_bytes"], extension, "Template client")
    # Untouched native objects/masters remain in the actual package, not a preview.
    with ZipFile(io.BytesIO(original)) as before, ZipFile(io.BytesIO(result["_persisted_bytes"])) as after:
        prefix = {"docx": ("word/header", "word/footer", "word/theme/"), "pptx": ("ppt/slideMasters/", "ppt/charts/", "ppt/embeddings/", "ppt/notesSlides/"), "xlsx": ("xl/charts/",), "xlsm": ("xl/charts/", "xl/vbaProject.bin")}[extension]
        retained = [name for name in before.namelist() if name.startswith(prefix)]
        assert retained
        for name in retained:
            assert after.read(name) == before.read(name), name
    path = tmp_path / f"output.{extension}"
    path.write_bytes(result["_persisted_bytes"])
    followup = real_apply(str(path), [normalize_file_patch_operation(template_patch(extension, "Followup edit"))])
    assert_template_result(followup["_persisted_bytes"], extension, "Followup edit")


def test_xlsm_template_generation_uses_patch_executor_and_preserves_vba():
    original = macro_template_bytes()
    operations = [normalize_file_patch_operation({
        "op": "cell.set",
        "sheet": "Sheet",
        "cell": "A1",
        "value": "Macro template",
    })]

    result = generated_files._generate_office_operations_sync(
        "xlsm", operations, template_bytes=original,
    )

    assert not result.get("error"), result
    with ZipFile(io.BytesIO(result["_persisted_bytes"]), "r") as archive:
        assert archive.read("xl/vbaProject.bin") == b"opaque-template-macro-payload"
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(result["_persisted_bytes"]))
    try:
        assert workbook["Sheet"]["A1"].value == "Macro template"
    finally:
        workbook.close()


def test_blank_xlsm_operation_generation_is_rejected():
    operations = [normalize_file_patch_operation({
        "op": "cell.set", "cell": "A1", "value": 1,
    })]

    with pytest.raises(ValueError, match="unsupported operation generation type"):
        generated_files._generate_office_operations_sync("xlsm", operations)


@pytest.mark.asyncio
@pytest.mark.parametrize("extension,kind", [("docx", "document"), ("docx", "word_document"), ("pptx", "presentation"), ("xlsx", "spreadsheet"), ("xlsm", "spreadsheet")])
@pytest.mark.parametrize("nested", [False, True])
async def test_template_routes_through_document_runtime(monkeypatch, extension, kind, nested):
    captured = []

    async def runtime(**kwargs):
        captured.append(kwargs)
        return json.dumps({"created": True})

    monkeypatch.setattr(document, "runtime_generate_document_file", runtime)
    operation_type = "xlsx" if extension == "xlsm" else extension
    arguments = {"template": {"path": f"source.{extension}", "expected_sha256": "a" * 64}, "operations": [template_patch(operation_type)]}
    result = json.loads(await tool._generate_file_handler("entity", kind=kind, name=f"copy.{extension}", **({"params": arguments} if nested else arguments)))
    assert result["created"] is True
    assert captured[0]["template"] == arguments["template"]
    assert captured[0]["operations"] == arguments["operations"]
    assert captured[0]["content"] is None


@pytest.mark.asyncio
@pytest.mark.parametrize("extra", [
    {"operations": None}, {"operations": []}, {"content": "ignored"}, {"prompt": "ignored"},
    {"kind": "image"}, {"kind": "pdf"}, {"kind": "code"}, {"kind": "video"}, {"kind": "audio"},
    {"template": None}, {"template": "source.docx"},
    {"params": {"template": {"path": "different.docx", "expected_sha256": "b" * 64}}},
])
async def test_invalid_template_never_reaches_a_generator(monkeypatch, extra):
    async def unexpected(**_kwargs):
        pytest.fail("Invalid template request reached a generator")

    for name in ("handle_document", "handle_image", "handle_pdf", "handle_code", "handle_video", "handle_audio"):
        monkeypatch.setattr(tool, name, unexpected)
    result = json.loads(await tool._generate_file_handler("entity", **{
        "kind": "document", "name": "copy.docx", "operations": [template_patch("docx")],
        "template": {"path": "source.docx", "expected_sha256": "a" * 64}, **extra,
    }))
    assert result.get("error"), result


@pytest.mark.asyncio
@pytest.mark.parametrize("template", [
    {}, {"path": "source.docx"}, {"path": "source.docx", "expected_sha256": "short"},
    {"path": "source.docx", "expected_sha256": "g" * 64},
    {"path": "source.docx", "expected_sha256": "a" * 64, "extra": True},
    *[{"path": path, "expected_sha256": "a" * 64} for path in [None, "", "../source.docx", "/other/source.docx", " /other/source.docx", "..\\source.docx", ".ai/source.docx", "nested/../source.docx", "source\x00.docx", "source.xlsm", "source.pptx"]],
])
async def test_runtime_rejects_invalid_template_before_read_or_approval(tmp_path, monkeypatch, template):
    async def unexpected(**_kwargs):
        pytest.fail("Invalid template reached authorization")

    monkeypatch.setattr(file_actions, "runtime_entity_file_root", lambda _entity: str(tmp_path))
    monkeypatch.setattr(file_actions, "runtime_guard_file_mutation", unexpected)
    monkeypatch.setattr(file_actions, "runtime_guard_file_read_access", unexpected)
    result = json.loads(await generated_files.runtime_generate_document_file(
        entity_id="entity", user_id="", conversation_id="", name="copy.docx", file_type="docx",
        operations=[template_patch("docx")], template=template,
    ))
    assert result.get("error"), result


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["docx", "pptx", "xlsx", "xlsm"])
@pytest.mark.parametrize("outcome", ["success", "late_failure", "projection_failure", "approval_denied", "stale", "read_denied", "snapshot_changed_after_read", "existing"])
@pytest.mark.usefixtures("fs_enabled")
async def test_template_generation_knowledge_transaction(db_session, monkeypatch, extension, outcome):
    from packages.core.services.knowledge_sync import sync_file_to_knowledge

    entity_id, workspace_id = generate_ulid(), generate_ulid()
    await _seed_workspace(db_session, entity_id, workspace_id)
    root = Path(file_tools._get_entity_root(entity_id))
    source = root / f"template.{extension}"
    before = template_bytes(extension)
    source.write_bytes(before)
    sync = await sync_file_to_knowledge(entity_id=entity_id, abs_path=str(source), entity_root=str(root), source="upload", created_by="test")
    assert sync.document_id
    original = await _doc_for(db_session, entity_id, source.name)
    original_id, original_meta = original.id, copy.deepcopy(original.metadata_)
    original_events = set((await db_session.scalars(select(EventLog.id).where(EventLog.entity_id == entity_id))).all())
    reads, approvals = [], []

    async def read_guard(**kwargs):
        reads.append(kwargs)
        if outcome == "read_denied":
            return json.dumps({"error": "file_permission_denied"})

    async def approve(**kwargs):
        approvals.append(kwargs)
        if outcome == "approval_denied":
            return json.dumps({"error": "approval_required"})
        if outcome == "snapshot_changed_after_read":
            source.write_bytes(b"external writer changed the template after snapshot")

    monkeypatch.setattr(file_actions, "runtime_guard_file_read_access", read_guard)
    monkeypatch.setattr(file_actions, "runtime_guard_file_mutation", approve)
    if outcome == "projection_failure":
        async def fail_metadata(_path):
            raise OSError("injected template projection response failure")

        monkeypatch.setattr(generated_files, "runtime_generated_file_metadata", fail_metadata)
    elif outcome == "read_denied":
        def no_read(*_args):
            pytest.fail("Template bytes/hash read before source ACL")

        monkeypatch.setattr(generated_files, "_read_office_template_sync", no_read)
    name = f"copy.{extension}"
    operations = [template_patch(extension)]
    if outcome == "late_failure":
        operations.append({"op": "cell.set", "cell": "A0", "value": 1} if extension in {"xlsx", "xlsm"} else {"op": "text.set", "index": 999, "text": "invalid", **({"slide": 1, "shape_id": 2} if extension == "pptx" else {})})
    target_rel = await file_tools._workspace_scoped_new_file_path(entity_id=entity_id, entity_root=str(root), workspace_id=workspace_id, path=name)
    target = root / target_rel
    if outcome == "existing":
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(b"Keep existing output")
    template = {"path": source.name, "expected_sha256": "a" * 64 if outcome == "stale" else hashlib.sha256(before).hexdigest()}
    result = json.loads(await tool._generate_file_handler(
        entity_id, kind="document", name=name, workspace_id=workspace_id,
        operations=operations, template=template,
    ))
    reloaded = await _doc_for(db_session, entity_id, source.name)
    assert reloaded.id == original_id
    assert reloaded.metadata_ == original_meta
    if outcome != "snapshot_changed_after_read":
        assert source.read_bytes() == before
    if outcome != "existing":
        assert len(reads) == 1
        assert reads[0]["paths"] == [source.name]
        assert reads[0]["workspace_id"] == workspace_id
    if approvals:
        assert len(approvals) == 1
        assert approvals[0]["approval_payload"]["template"] == template
        assert approvals[0]["approval_payload"]["operations"] == [normalize_file_patch_operation(op) for op in operations]
        assert json.loads(approvals[0]["content_preview"]) == approvals[0]["approval_payload"]
        assert approvals[0]["paths"] == [target_rel]
    else:
        assert outcome in {"stale", "read_denied", "existing"}
    if outcome in {"success", "snapshot_changed_after_read"}:
        assert result.get("created") is True, result
        output = await _doc_for(db_session, entity_id, name)
        assert output.id != original_id
        assert output.metadata_["origin"]["workspace_id"] == workspace_id
        assert output.fs_path == target_rel
        assert result["document"]["document_id"] == output.id
        assert result["document"]["source_sha256"] == hashlib.sha256(target.read_bytes()).hexdigest()
        assert_template_result(target.read_bytes(), extension, "Template client")
        # The input template must not be advertised as a newly produced artifact.
        assert "template" not in result
        patched = json.loads(await file_tools._patch_file(entity_id, path=target_rel, workspace_id=workspace_id, expected_sha256=result["document"]["source_sha256"], operations=[template_patch(extension, "Followup edit")]))
        assert patched["document_id"] == output.id
        assert_template_result(target.read_bytes(), extension, "Followup edit")
    else:
        assert result.get("error"), result
        if outcome == "stale":
            assert result["error"] == "source_changed"
        if outcome == "existing":
            assert target.read_bytes() == b"Keep existing output"
        else:
            assert not target.exists()
        assert await _doc_for(db_session, entity_id, name) is None
        assert set((await db_session.scalars(select(EventLog.id).where(EventLog.entity_id == entity_id))).all()) == original_events


@pytest.mark.parametrize("change", ["path", "hash", "operations"])
def test_approval_binds_template_identity_version_and_patches(change):
    payload = {"file_type": "docx", "template": {"path": "source.docx", "expected_sha256": "a" * 64}, "operations": [template_patch("docx")]}
    approved = FileApprovalOperationFactory.create(tool_name="generate_file", action="create_document", paths=["copy.docx"], payload=payload)
    changed = copy.deepcopy(payload)
    if change == "operations":
        changed["operations"] = [template_patch("docx", "different")]
    else:
        changed["template"]["path" if change == "path" else "expected_sha256"] = "other.docx" if change == "path" else "b" * 64
    retry = FileApprovalOperationFactory.create(tool_name="generate_file", action="create_document", paths=["copy.docx"], payload=changed)
    assert not retry.matches(approved.to_record())


@pytest.mark.parametrize("alias", ["file", "parent"])
def test_template_snapshot_rejects_alias_even_after_acl(tmp_path, monkeypatch, alias):
    from packages.core.services.entity_fs import EntityFilesystemError

    monkeypatch.setattr(get_settings(), "MANOR_FS_ROOT", str(tmp_path))
    monkeypatch.setattr(get_settings(), "MANOR_FS_ENABLED", True)
    root = tmp_path / "entity"
    original = root / "actual" / "template.docx"
    original.parent.mkdir(parents=True)
    original.write_bytes(b"original")
    path = root / "template.docx" if alias == "file" else root / "alias"
    path.symlink_to(original if alias == "file" else original.parent, target_is_directory=alias == "parent")
    assert generated_files._read_office_template_sync("entity", {"path": "actual/template.docx", "expected_sha256": hashlib.sha256(b"original").hexdigest()}) == b"original"
    with pytest.raises(EntityFilesystemError):
        generated_files._read_office_template_sync("entity", {"path": "template.docx" if alias == "file" else "alias/template.docx", "expected_sha256": hashlib.sha256(b"original").hexdigest()})


@pytest.mark.asyncio
@pytest.mark.usefixtures("fs_enabled")
@pytest.mark.parametrize("restriction", ["private", "restricted", "quarantined", "other_workspace", "deleted_workspace", "userless"])
async def test_real_template_read_acl_blocks_before_bytes_or_creation(db_session, monkeypatch, restriction):
    from datetime import datetime, timezone
    from packages.core.models.user import User
    from packages.core.models.workspace import Workspace
    from packages.core.services.document_service import create_document
    from packages.core.services.entity_fs import provision_entity_filesystem

    entity_id = generate_ulid()
    owner = User(entity_id=entity_id, email=f"{entity_id}-owner@test.com", display_name="Owner", password_hash="unused", role="owner", status="active")
    member = User(entity_id=entity_id, email=f"{entity_id}-member@test.com", display_name="Member", password_hash="unused", role="member", status="active")
    db_session.add_all([owner, member])
    await db_session.flush()
    owner_id, member_id = owner.id, member.id
    metadata = {}
    workspace_id = None
    if restriction in {"other_workspace", "deleted_workspace"}:
        source_workspace = Workspace(entity_id=entity_id, name="Source", operating_model={})
        db_session.add(source_workspace)
        await db_session.flush()
        metadata = {"origin": {"workspace_id": source_workspace.id}}
        if restriction == "deleted_workspace":
            source_workspace.deleted_at = datetime.now(timezone.utc)
        else:
            destination = Workspace(entity_id=entity_id, name="Destination", operating_model={})
            db_session.add(destination)
            await db_session.flush()
            workspace_id = destination.id
    source_doc = await create_document(db_session, entity_id, name="template.docx", fs_path="template.docx", file_type="docx", source="upload", visibility="private" if restriction == "private" else "entity", classification="restricted" if restriction == "restricted" else None, owner_id=owner_id, metadata=metadata)
    if restriction == "quarantined":
        source_doc.quarantine_status = "quarantined"
    await db_session.commit()
    provision_entity_filesystem(entity_id)
    source = Path(file_tools._get_entity_root(entity_id)) / "template.docx"
    source.write_bytes(b"protected template must not be parsed")

    def no_read(*_args):
        pytest.fail("Protected source was read before its ACL passed")

    async def no_create(**_kwargs):
        pytest.fail("Protected source reached create approval")

    monkeypatch.setattr(generated_files, "_read_office_template_sync", no_read)
    monkeypatch.setattr(file_actions, "runtime_guard_file_mutation", no_create)
    user_id = owner_id if restriction in {"other_workspace", "deleted_workspace"} else member_id
    result = json.loads(await tool._generate_file_handler(
        entity_id, user_id="" if restriction == "userless" else user_id, kind="document", name="copy.docx", workspace_id=workspace_id,
        operations=[template_patch("docx")], template={"path": source.name, "expected_sha256": "a" * 64},
    ))
    assert result.get("error"), result
    assert result["mode"] == "entity_file_acl_denied"
    assert await _doc_for(db_session, entity_id, "copy.docx") is None


def test_template_snapshot_detects_in_place_changes(tmp_path, monkeypatch):
    import builtins
    from packages.core.services.entity_fs import EntityFilesystemStaleWriteError

    monkeypatch.setattr(get_settings(), "MANOR_FS_ROOT", str(tmp_path))
    monkeypatch.setattr(get_settings(), "MANOR_FS_ENABLED", True)
    path = tmp_path / "entity" / "source.docx"
    path.parent.mkdir()
    path.write_bytes(b"original")
    real_open = builtins.open

    def changed_open(file, *args, **kwargs):
        if str(file).startswith(("/dev/fd/", "/proc/self/fd/")):
            path.write_bytes(b"external changed bytes")
        return real_open(file, *args, **kwargs)

    monkeypatch.setattr(builtins, "open", changed_open)
    with pytest.raises(EntityFilesystemStaleWriteError):
        generated_files._read_office_template_sync("entity", {"path": path.name, "expected_sha256": hashlib.sha256(b"original").hexdigest()})
