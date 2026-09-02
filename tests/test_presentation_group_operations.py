"""Grouped PPT objects use the same slide + shape_id patch contract as top-level objects."""

from __future__ import annotations

import hashlib
import io

from PIL import Image
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Pt

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools.file_tools import _apply_office_patch_sequence_sync
from packages.core.contracts.file_engine import file_patch_operations, normalize_file_patch_operation
from packages.core.services.file_engine_patches import describe_file_structure


def _image_bytes(color: str) -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (160, 90), color).save(output, "PNG")
    return output.getvalue()


def _grouped_template() -> tuple[bytes, dict[str, int], bytes]:
    image = _image_bytes("red")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()

    textbox = group.shapes.add_textbox(Pt(30), Pt(35), Pt(210), Pt(65))
    textbox.text = "Grouped title"
    card = group.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.from_xml("roundRect"), Pt(30), Pt(120), Pt(210), Pt(90),
    )
    card.text = "Grouped card"
    picture = group.shapes.add_picture(io.BytesIO(image), Pt(270), Pt(35), Pt(180), Pt(105))

    table = slide.shapes.add_table(2, 2, Pt(270), Pt(165), Pt(220), Pt(100))
    table.table.cell(0, 0).text = "Metric"
    table.table.cell(0, 1).text = "Value"
    table.table.cell(1, 0).text = "Revenue"
    table.table.cell(1, 1).text = "12"
    group._element.append(table._element)

    chart_data = ChartData()
    chart_data.categories = ["Jan", "Feb"]
    chart_data.add_series("Revenue", [12, 15])
    chart = group.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(520), Pt(35), Pt(300), Pt(210), chart_data,
    )
    nested_group = group.shapes.add_group_shape()
    nested_textbox = nested_group.shapes.add_textbox(Pt(850), Pt(35), Pt(180), Pt(65))
    nested_textbox.text = "Nested note"
    nested_shape = nested_group.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.from_xml("ellipse"), Pt(850), Pt(120), Pt(90), Pt(90),
    )
    group.shapes._recalculate_extents()

    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue(), {
        "group": group.shape_id,
        "textbox": textbox.shape_id,
        "card": card.shape_id,
        "picture": picture.shape_id,
        "table": table.shape_id,
        "chart": chart.shape_id,
        "nested_group": nested_group.shape_id,
        "nested_textbox": nested_textbox.shape_id,
        "nested_shape": nested_shape.shape_id,
    }, image


def _normalized(operations):
    return [normalize_file_patch_operation(operation) for operation in operations]


def _all_shapes(presentation: Presentation):
    def walk(shapes):
        for shape in shapes:
            yield shape
            if shape._element.tag.rsplit("}", 1)[-1] == "grpSp":
                yield from walk(shape.shapes)

    return list(walk(presentation.slides[0].shapes))


def test_group_and_ungroup_are_shared_public_operations_with_stable_z_order():
    assert {"shape.group", "shape.ungroup"} <= set(file_patch_operations("pptx"))
    operations = _normalized([
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        {"op": "shape.insert", "slide": 1, "preset": "rect",
         "transform": {"x": 20, "y": 30, "width": 120, "height": 60}, "text": "First"},
        {"op": "shape.insert", "slide": 1, "preset": "ellipse",
         "transform": {"x": 160, "y": 30, "width": 80, "height": 80}},
        {"op": "shape.insert", "slide": 1, "preset": "rect",
         "transform": {"x": 270, "y": 30, "width": 120, "height": 60}, "text": "Last"},
        {"op": "shape.group", "slide": 1, "shape_ids": [2, 3]},
        {"op": "text.set", "slide": 1, "shape_id": 2, "index": 0, "text": "Grouped first"},
        {"op": "shape.ungroup", "slide": 1, "shape_id": 5},
    ])
    result = _generate_office_operations_sync("pptx", operations)
    assert not result.get("error"), result
    assert result["operation_results"][4] == {
        "operation": "shape.group", "slide": 1, "shape_id": 5,
        "grouped_shape_ids": [2, 3],
    }
    assert result["operation_results"][5]["group_path"] == [5, 2]
    assert result["operation_results"][6]["ungrouped_shape_ids"] == [2, 3]
    presentation = Presentation(io.BytesIO(result["_persisted_bytes"]))
    assert [shape.shape_id for shape in presentation.slides[0].shapes] == [2, 3, 4]
    assert presentation.slides[0].shapes[0].text == "Grouped first"


def test_shape_reorder_is_shared_by_blank_generation_template_generation_and_patch(tmp_path):
    assert "shape.reorder" in file_patch_operations("pptx")
    seed_operations = _normalized([
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        {"op": "shape.insert", "slide": 1, "preset": "rect",
         "transform": {"x": 20, "y": 30, "width": 180, "height": 90}, "text": "Back"},
        {"op": "shape.insert", "slide": 1, "preset": "ellipse",
         "transform": {"x": 20, "y": 30, "width": 180, "height": 90}, "text": "Middle"},
        {"op": "shape.insert", "slide": 1, "preset": "rect",
         "transform": {"x": 20, "y": 30, "width": 180, "height": 90}, "text": "Front"},
    ])
    reorder = normalize_file_patch_operation(
        {"op": "shape.reorder", "slide": 1, "shape_id": 2, "z_index": 2},
    )
    blank = _generate_office_operations_sync("pptx", [*seed_operations, reorder])
    assert not blank.get("error"), blank
    assert blank["operation_results"][-1] == {
        "operation": "shape.reorder",
        "slide": 1,
        "shape_id": 2,
        "previous_z_index": 0,
        "z_index": 2,
        "sibling_count": 3,
    }

    template = _generate_office_operations_sync("pptx", seed_operations)["_persisted_bytes"]
    generated = _generate_office_operations_sync("pptx", [reorder], template_bytes=template)
    path = tmp_path / "z-order-template.pptx"
    path.write_bytes(template)
    patched = _apply_office_patch_sequence_sync(str(path), [reorder])
    for result in (blank, generated, patched):
        assert not result.get("error"), result
        presentation = Presentation(io.BytesIO(result["_persisted_bytes"]))
        assert [shape.shape_id for shape in presentation.slides[0].shapes] == [3, 4, 2]
        assert [shape.text for shape in presentation.slides[0].shapes] == ["Middle", "Front", "Back"]
    assert path.read_bytes() == template

    output = tmp_path / "z-order-output.pptx"
    output.write_bytes(patched["_persisted_bytes"])
    by_id = {shape["shape_id"]: shape for shape in describe_file_structure(str(output))["shapes"]}
    assert (by_id[3]["z_index"], by_id[4]["z_index"], by_id[2]["z_index"]) == (0, 1, 2)
    assert all(shape["sibling_count"] == 3 for shape in by_id.values())


def test_shape_reorder_supports_nested_group_siblings_without_changing_membership(tmp_path):
    template, ids, _image = _grouped_template()
    path = tmp_path / "nested-z-order.pptx"
    path.write_bytes(template)
    result = _apply_office_patch_sequence_sync(str(path), _normalized([
        {"op": "shape.reorder", "slide": 1, "shape_id": ids["textbox"], "z_index": 5},
    ]))
    assert not result.get("error"), result
    assert result["operation_results"][0]["group_path"] == [ids["group"], ids["textbox"]]
    presentation = Presentation(io.BytesIO(result["_persisted_bytes"]))
    outer = presentation.slides[0].shapes[0]
    assert [shape.shape_id for shape in outer.shapes] == [
        ids["card"], ids["picture"], ids["table"], ids["chart"], ids["nested_group"], ids["textbox"],
    ]
    assert presentation.slides[0].shapes[0].shape_id == ids["group"]


def test_shape_reorder_invalid_destination_is_atomic(tmp_path):
    generated = _generate_office_operations_sync("pptx", _normalized([
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        {"op": "shape.insert", "slide": 1, "preset": "rect",
         "transform": {"x": 20, "y": 30, "width": 120, "height": 60}, "text": "One"},
        {"op": "shape.insert", "slide": 1, "preset": "rect",
         "transform": {"x": 160, "y": 30, "width": 120, "height": 60}, "text": "Two"},
    ]))
    path = tmp_path / "atomic-z-order.pptx"
    path.write_bytes(generated["_persisted_bytes"])
    before = path.read_bytes()
    result = _apply_office_patch_sequence_sync(str(path), _normalized([
        {"op": "shape.reorder", "slide": 1, "shape_id": 2, "z_index": 1},
        {"op": "shape.reorder", "slide": 1, "shape_id": 3, "z_index": 2},
    ]))
    assert result.get("error"), result
    assert result["operation_index"] == 1
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before


def test_ungroup_bakes_translation_and_scale_into_member_geometry():
    result = _generate_office_operations_sync("pptx", _normalized([
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        {"op": "shape.insert", "slide": 1, "preset": "rect",
         "transform": {"x": 20, "y": 30, "width": 120, "height": 60}},
        {"op": "shape.insert", "slide": 1, "preset": "ellipse",
         "transform": {"x": 160, "y": 30, "width": 80, "height": 80}},
        {"op": "shape.group", "slide": 1, "shape_ids": [2, 3]},
        {"op": "shape.transform", "slide": 1, "shape_id": 4,
         "transform": {"x": 100, "y": 80, "width": 440, "height": 160}},
        {"op": "shape.ungroup", "slide": 1, "shape_id": 4},
    ]))
    assert not result.get("error"), result
    first, second = Presentation(io.BytesIO(result["_persisted_bytes"])).slides[0].shapes
    assert (first.left.pt, first.top.pt, first.width.pt, first.height.pt) == (100, 80, 240, 120)
    assert (second.left.pt, second.top.pt, second.width.pt, second.height.pt) == (380, 80, 160, 160)


def test_nested_sibling_grouping_uses_a_slide_unique_id_and_round_trips(tmp_path):
    template, ids, _image = _grouped_template()
    path = tmp_path / "nested-grouping.pptx"
    path.write_bytes(template)
    group_id = max(ids.values()) + 1
    grouped = _apply_office_patch_sequence_sync(str(path), _normalized([
        {"op": "shape.group", "slide": 1, "shape_ids": [ids["textbox"], ids["card"]]},
    ]))
    assert not grouped.get("error"), grouped
    assert grouped["operation_results"][0]["shape_id"] == group_id
    assert grouped["operation_results"][0]["group_path"] == [ids["group"], group_id]
    nested = Presentation(io.BytesIO(grouped["_persisted_bytes"]))
    shape_ids = [shape.shape_id for shape in _all_shapes(nested)]
    assert len(shape_ids) == len(set(shape_ids))
    grouped_path = tmp_path / "nested-grouped.pptx"
    grouped_path.write_bytes(grouped["_persisted_bytes"])
    ungrouped = _apply_office_patch_sequence_sync(str(grouped_path), _normalized([
        {"op": "shape.ungroup", "slide": 1, "shape_id": group_id},
    ]))
    assert not ungrouped.get("error"), ungrouped
    restored = Presentation(io.BytesIO(ungrouped["_persisted_bytes"]))
    outer = restored.slides[0].shapes[0]
    assert [shape.shape_id for shape in list(outer.shapes)[:2]] == [ids["textbox"], ids["card"]]


def test_grouping_rejects_noncontiguous_z_order_and_complex_ungroup_is_atomic(tmp_path):
    template, ids, _image = _grouped_template()
    path = tmp_path / "group-guards.pptx"
    path.write_bytes(template)
    noncontiguous = _apply_office_patch_sequence_sync(str(path), _normalized([
        {"op": "shape.group", "slide": 1, "shape_ids": [ids["textbox"], ids["picture"]]},
    ]))
    assert noncontiguous.get("error")
    assert "contiguous sibling" in noncontiguous["error"]
    assert path.read_bytes() == template

    transformed = _apply_office_patch_sequence_sync(str(path), _normalized([
        {"op": "shape.transform", "slide": 1, "shape_id": ids["group"],
         "transform": {"rotation": 15}},
        {"op": "shape.ungroup", "slide": 1, "shape_id": ids["group"]},
    ]))
    assert transformed.get("error")
    assert transformed["operation_index"] == 1
    assert "without changing appearance" in transformed["error"]
    assert "_persisted_bytes" not in transformed
    assert path.read_bytes() == template


def test_group_members_are_discovered_and_editable_through_existing_operations(tmp_path):
    template, ids, _image = _grouped_template()
    path = tmp_path / "grouped-template.pptx"
    path.write_bytes(template)
    structure = describe_file_structure(str(path))
    by_id = {shape["shape_id"]: shape for shape in structure["shapes"]}
    assert structure["shape_count"] == len(ids)
    assert by_id[ids["textbox"]]["group_path"] == [ids["group"], ids["textbox"]]
    assert by_id[ids["nested_textbox"]]["group_path"] == [
        ids["group"], ids["nested_group"], ids["nested_textbox"],
    ]
    assert by_id[ids["nested_textbox"]]["coordinate_space"] == f"group:{ids['nested_group']}"
    original_group_transform = Presentation(io.BytesIO(template)).slides[0].shapes[0]._element.xfrm.xml

    operations = _normalized([
        {"op": "text.set", "slide": 1, "shape_id": ids["textbox"], "index": 0, "text": "Updated title"},
        {"op": "paragraph.insert", "slide": 1, "shape_id": ids["textbox"], "index": 1, "text": "Second line"},
        {"op": "paragraph.format", "slide": 1, "shape_id": ids["textbox"], "index": 1,
         "format": {"bold": True, "font_size": 18, "font_color": "174C46"}},
        {"op": "shape.format", "slide": 1, "shape_id": ids["card"],
         "format": {"fill_color": "DDEEFF", "line_color": "174C46", "corner_radius": 12}},
        {"op": "shape.transform", "slide": 1, "shape_id": ids["nested_textbox"],
         "transform": {"x": 875, "rotation": 7}},
        {"op": "cell.set", "slide": 1, "shape_id": ids["table"], "cell": "B2", "value": "19"},
        {"op": "cell.format", "slide": 1, "shape_id": ids["table"], "cell": "B2",
         "format": {"bold": True, "fill_color": "EAF5F0"}},
        {"op": "picture.format", "slide": 1, "shape_id": ids["picture"],
         "format": {"opacity": 0.6, "alt_text": "Grouped dashboard"}},
        {"op": "chart.data", "slide": 1, "shape_id": ids["chart"],
         "categories": ["Mar", "Apr"], "series": [{"name": "Revenue", "values": [18, 21]}]},
        {"op": "chart.format", "slide": 1, "shape_id": ids["chart"],
         "format": {"title": "Grouped revenue", "has_legend": False}},
    ])
    result = _apply_office_patch_sequence_sync(str(path), operations)
    assert not result.get("error"), result
    assert path.read_bytes() == template
    assert all(item["group_path"][0] == ids["group"] for item in result["operation_results"])

    patched = Presentation(io.BytesIO(result["_persisted_bytes"]))
    shapes = {shape.shape_id: shape for shape in _all_shapes(patched)}
    assert shapes[ids["textbox"]].text == "Updated title\nSecond line"
    assert shapes[ids["textbox"]].text_frame.paragraphs[1].runs[0].font.bold
    assert str(shapes[ids["card"]].fill.fore_color.rgb) == "DDEEFF"
    assert shapes[ids["nested_textbox"]].left.pt == 875
    assert shapes[ids["nested_textbox"]].rotation == 7
    assert shapes[ids["table"]].table.cell(1, 1).text == "19"
    assert shapes[ids["table"]].table.cell(1, 1).fill.type is not None
    assert shapes[ids["picture"]]._element.nvPicPr.cNvPr.get("descr") == "Grouped dashboard"
    assert list(shapes[ids["chart"]].chart.series[0].values) == [18, 21]
    assert shapes[ids["chart"]].chart.chart_title.text_frame.text == "Grouped revenue"
    assert patched.slides[0].shapes[0]._element.xfrm.xml == original_group_transform


def test_template_generation_and_patch_share_group_member_executor(tmp_path):
    template, ids, _image = _grouped_template()
    path = tmp_path / "grouped-template.pptx"
    path.write_bytes(template)
    operations = _normalized([
        {"op": "text.set", "slide": 1, "shape_id": ids["nested_textbox"], "index": 0,
         "text": "Generated from grouped template"},
        {"op": "shape.transform", "slide": 1, "shape_id": ids["nested_shape"],
         "transform": {"width": 110, "flip_horizontal": True}},
    ])
    patched = _apply_office_patch_sequence_sync(str(path), operations)
    generated = _generate_office_operations_sync("pptx", operations, template_bytes=template)
    assert not patched.get("error") and not generated.get("error")
    assert patched["operation_results"] == generated["operation_results"]
    for result in (patched, generated):
        shapes = {shape.shape_id: shape for shape in _all_shapes(Presentation(io.BytesIO(result["_persisted_bytes"])))}
        assert shapes[ids["nested_textbox"]].text == "Generated from grouped template"
        assert shapes[ids["nested_shape"]].width.pt == 110
        assert shapes[ids["nested_shape"]]._element.xfrm.flipH is True


def test_grouped_picture_replace_and_type_safe_deletes(tmp_path):
    template, ids, original_image = _grouped_template()
    replacement = _image_bytes("blue")
    digest = hashlib.sha256(replacement).hexdigest()
    source = {"path": "Assets/replacement.png", "expected_sha256": digest}
    path = tmp_path / "grouped-template.pptx"
    path.write_bytes(template)
    operations = _normalized([
        {"op": "picture.replace", "slide": 1, "shape_id": ids["picture"], "source": source},
        {"op": "picture.delete", "slide": 1, "shape_id": ids["picture"]},
        {"op": "table.delete", "slide": 1, "shape_id": ids["table"]},
        {"op": "chart.delete", "slide": 1, "shape_id": ids["chart"]},
        {"op": "shape.delete", "slide": 1, "shape_id": ids["nested_shape"]},
    ])
    result = _apply_office_patch_sequence_sync(
        str(path), operations, resources={(source["path"], digest): replacement},
    )
    assert not result.get("error"), result
    assert original_image not in result["_persisted_bytes"]
    remaining = {shape.shape_id for shape in _all_shapes(Presentation(io.BytesIO(result["_persisted_bytes"])))}
    assert not {ids["picture"], ids["table"], ids["chart"], ids["nested_shape"]} & remaining
    assert {ids["group"], ids["nested_group"], ids["nested_textbox"]} <= remaining


def test_group_member_failures_are_atomic_and_the_only_member_requires_explicit_group_delete(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    member = group.shapes.add_textbox(Pt(30), Pt(30), Pt(200), Pt(60))
    member.text = "Keep me"
    path = tmp_path / "one-member-group.pptx"
    presentation.save(path)
    before = path.read_bytes()
    result = _apply_office_patch_sequence_sync(str(path), _normalized([
        {"op": "text.set", "slide": 1, "shape_id": member.shape_id, "index": 0, "text": "Temporary"},
        {"op": "shape.delete", "slide": 1, "shape_id": member.shape_id},
    ]))
    assert result.get("error")
    assert result["operation_index"] == 1
    assert "only member" in result["error"]
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before

    deleted = _apply_office_patch_sequence_sync(str(path), _normalized([
        {"op": "shape.delete", "slide": 1, "shape_id": group.shape_id},
    ]))
    assert not deleted.get("error"), deleted
    assert len(Presentation(io.BytesIO(deleted["_persisted_bytes"])).slides[0].shapes) == 0
