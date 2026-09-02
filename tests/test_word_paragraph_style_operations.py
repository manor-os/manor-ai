"""Native Word paragraph-style definitions share generation and patch execution."""

from __future__ import annotations

import io

import pytest
from docx import Document

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools.file_tools import _apply_office_patch_sequence_sync
from packages.core.contracts.file_engine import file_patch_operations, normalize_file_patch_operation
from packages.core.services.file_engine_patches import describe_file_structure


def _normalized(operations):
    return [normalize_file_patch_operation(operation) for operation in operations]


def _generate(operations, *, template_bytes=None):
    return _generate_office_operations_sync(
        "docx", _normalized(operations), template_bytes=template_bytes,
    )


def _style_operations():
    return [
        {"op": "paragraph_style.insert", "style": "Client Body",
         "format": {"font_name": "Arial", "font_size": 11, "font_color": "263238",
                    "space_after": 6, "line_spacing": 1.15, "widow_control": True}},
        {"op": "paragraph_style.insert", "style": "Client Heading", "base_style": "Title",
         "next_style": "Client Body",
         "format": {"font_name": "Arial", "font_size": 24, "font_color": "174C46", "bold": True,
                    "space_before": 12, "space_after": 8, "keep_with_next": True}},
        {"op": "paragraph.insert", "index": 0, "text": "Service proposal", "style": "Client Heading"},
        {"op": "paragraph.insert", "index": 1, "text": "Prepared for the client", "style": "Client Body"},
    ]


def test_paragraph_style_operations_are_public_and_generate_native_editable_styles(tmp_path):
    assert {
        "paragraph_style.insert", "paragraph_style.format", "paragraph_style.delete",
    } <= set(file_patch_operations("docx"))
    result = _generate(_style_operations())
    assert not result.get("error"), result
    document = Document(io.BytesIO(result["_persisted_bytes"]))
    heading = document.styles["Client Heading"]
    body = document.styles["Client Body"]
    assert not heading.builtin and not body.builtin
    assert heading.base_style.name == "Title"
    assert heading.next_paragraph_style.name == "Client Body"
    assert heading.font.name == "Arial" and heading.font.size.pt == 24
    assert str(heading.font.color.rgb) == "174C46"
    assert heading.font.bold and heading.paragraph_format.keep_with_next
    assert [paragraph.style.name for paragraph in document.paragraphs] == ["Client Heading", "Client Body"]
    assert 'w:eastAsia="Arial"' in heading._element.xml
    assert 'w:cs="Arial"' in heading._element.xml
    assert '<w:szCs w:val="48"' in heading._element.xml

    path = tmp_path / "styled.docx"
    path.write_bytes(result["_persisted_bytes"])
    structure = describe_file_structure(str(path))
    definitions = {item["name"]: item for item in structure["paragraph_style_definitions"]}
    assert definitions["Client Heading"] == {
        "name": "Client Heading",
        "style_id": heading.style_id,
        "builtin": False,
        "base_style": "Title",
        "next_style": "Client Body",
        "format": {
            "bold": True,
            "font_name": "Arial",
            "font_size": 24.0,
            "font_color": "174C46",
            "space_before": 12.0,
            "space_after": 8.0,
            "keep_with_next": True,
        },
    }


def test_style_format_is_selective_and_template_generation_matches_patch(tmp_path):
    template = _generate(_style_operations())["_persisted_bytes"]
    path = tmp_path / "style-template.docx"
    path.write_bytes(template)
    operations = _normalized([
        {"op": "paragraph_style.format", "style": "Client Heading", "base_style": "Heading 1",
         "next_style": None, "format": {"font_color": "336699", "space_after": 14}},
    ])
    patched = _apply_office_patch_sequence_sync(str(path), operations)
    generated = _generate_office_operations_sync("docx", operations, template_bytes=template)
    assert not patched.get("error") and not generated.get("error")
    assert path.read_bytes() == template
    for result in (patched, generated):
        document = Document(io.BytesIO(result["_persisted_bytes"]))
        style = document.styles["Client Heading"]
        assert style.base_style.name == "Heading 1"
        assert style.next_paragraph_style.name == "Client Heading"
        assert str(style.font.color.rgb) == "336699"
        assert style.font.name == "Arial" and style.font.size.pt == 24 and style.font.bold
        assert style.paragraph_format.space_before.pt == 12
        assert style.paragraph_format.space_after.pt == 14
        assert document.paragraphs[0].style.name == "Client Heading"


def test_style_delete_guards_builtins_content_and_style_references_then_deletes_unused(tmp_path):
    template = _generate([
        {"op": "paragraph_style.insert", "style": "Base Custom", "format": {"font_size": 12}},
        {"op": "paragraph_style.insert", "style": "Derived Custom", "base_style": "Base Custom"},
        {"op": "paragraph_style.insert", "style": "Unused Custom", "format": {"italic": True}},
        {"op": "paragraph.insert", "index": 0, "text": "Uses derived", "style": "Derived Custom"},
    ])["_persisted_bytes"]
    path = tmp_path / "delete-styles.docx"
    path.write_bytes(template)
    for style, message in (
        ("Normal", "Built-in"),
        ("Derived Custom", "used by document content"),
        ("Base Custom", "referenced by another style"),
    ):
        result = _apply_office_patch_sequence_sync(str(path), _normalized([
            {"op": "paragraph_style.delete", "style": style},
        ]))
        assert result.get("error") and message in result["error"]
        assert "_persisted_bytes" not in result
        assert path.read_bytes() == template

    deleted = _apply_office_patch_sequence_sync(str(path), _normalized([
        {"op": "paragraph_style.delete", "style": "Unused Custom"},
    ]))
    assert not deleted.get("error"), deleted
    document = Document(io.BytesIO(deleted["_persisted_bytes"]))
    with pytest.raises(KeyError):
        document.styles["Unused Custom"]
    assert document.styles["Base Custom"] is not None


@pytest.mark.parametrize("operation", [
    {"op": "paragraph_style.insert", "style": "Normal", "format": {"bold": True}},
    {"op": "paragraph_style.insert", "style": " Bad", "format": {"bold": True}},
    {"op": "paragraph_style.insert", "style": "New", "base_style": "New"},
    {"op": "paragraph_style.insert", "style": "New", "base_style": "Missing"},
    {"op": "paragraph_style.insert", "style": "New", "format": {}},
    {"op": "paragraph_style.format", "style": "Normal"},
    {"op": "paragraph_style.format", "style": "Missing", "format": {"bold": True}},
    {"op": "paragraph_style.format", "style": "Normal", "format": {"font_size": 10.25}},
    {"op": "paragraph_style.delete", "style": "Missing"},
])
def test_invalid_paragraph_style_operations_are_atomic(tmp_path, operation):
    document = Document()
    document.add_paragraph("Keep")
    path = tmp_path / "invalid-style.docx"
    document.save(path)
    before = path.read_bytes()
    result = _apply_office_patch_sequence_sync(str(path), _normalized([operation]))
    assert result.get("error"), result
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before


def test_style_insert_rejects_a_name_that_collides_after_word_style_id_normalization(tmp_path):
    document = Document()
    path = tmp_path / "style-id-collision.docx"
    document.save(path)
    before = path.read_bytes()
    result = _apply_office_patch_sequence_sync(str(path), _normalized([
        {"op": "paragraph_style.insert", "style": "Client Body", "format": {"font_size": 11}},
        {"op": "paragraph_style.insert", "style": "ClientBody", "format": {"font_size": 12}},
    ]))
    assert result.get("error")
    assert result["operation_index"] == 1
    assert "existing style ID" in result["error"]
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before
