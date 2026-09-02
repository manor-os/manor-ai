"""Native Office text patches must retain the surrounding editable OOXML."""

from __future__ import annotations

import base64
import io
import zipfile

import pytest
from lxml import etree

from packages.core.ai.tools import file_tools


def _document(extension):
    if extension == "docx":
        from docx import Document

        document = Document()
        return document, document.add_paragraph()
    from pptx import Presentation
    from pptx.util import Inches

    document = Presentation()
    shape = document.slides.add_slide(document.slide_layouts[6]).shapes.add_textbox(
        Inches(1), Inches(2), Inches(3), Inches(1)
    )
    return document, shape.text_frame.paragraphs[0]


def _run(paragraph, text, **styles):
    run = paragraph.add_run()
    run.text = text
    for key, value in styles.items():
        setattr(run.font, key, value)
    return run


def _patch(path, old, new, replace_all=False):
    return file_tools._apply_office_patch_sequence_sync(
        str(path),
        [{"operation": "replace_text", "old_text": old, "new_text": new, "replace_all": replace_all}],
    )


def _reopen(result, extension):
    assert result.get("patched") is True, result
    source = io.BytesIO(result["_persisted_bytes"])
    if extension == "docx":
        from docx import Document

        document = Document(source)
        return document, document.paragraphs[0]
    from pptx import Presentation

    document = Presentation(source)
    return document, document.slides[0].shapes[0].text_frame.paragraphs[0]


@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("replace_all", [False, True])
def test_replacement_does_not_search_inserted_text(tmp_path, extension, replace_all):
    document, paragraph = _document(extension)
    _run(paragraph, "Old Old", bold=True)
    path = tmp_path / f"source.{extension}"
    document.save(path)
    original = path.read_bytes()

    result = _patch(path, "Old", "Old updated", replace_all)

    _, edited = _reopen(result, extension)
    assert edited.text == ("Old updated Old updated" if replace_all else "Old updated Old")
    assert result["replacements"] == (2 if replace_all else 1)
    assert edited.runs[0].font.bold is True
    assert path.read_bytes() == original


@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("replacement", ["New", "", "OLD updated", "替换🙂"])
def test_cross_run_replace_keeps_original_run_properties(tmp_path, extension, replacement):
    document, paragraph = _document(extension)
    _run(paragraph, "prefix OL", bold=True)
    _run(paragraph, "D", italic=True)
    _run(paragraph, " suffix", underline=True)
    properties = [etree.tostring(run._r.rPr) for run in paragraph.runs]
    path = tmp_path / f"source.{extension}"
    document.save(path)

    result = _patch(path, "OLD", replacement, True)

    _, edited = _reopen(result, extension)
    assert edited.text == f"prefix {replacement} suffix"
    assert result["replacements"] == 1
    assert [etree.tostring(run._r.rPr) for run in edited.runs] == properties
    assert edited.runs[0].text == f"prefix {replacement}"
    assert edited.runs[1].text == ""
    assert edited.runs[2].text == " suffix"


@pytest.mark.parametrize("extension", ["docx", "pptx"])
def test_replace_first_uses_paragraph_order_not_run_match_order(tmp_path, extension):
    document, paragraph = _document(extension)
    _run(paragraph, "OL", bold=True)
    _run(paragraph, "D then OLD", italic=True)
    path = tmp_path / f"source.{extension}"
    document.save(path)

    result = _patch(path, "OLD", "NEW")

    _, edited = _reopen(result, extension)
    assert edited.text == "NEW then OLD"
    assert result["replacements"] == 1
    assert [run.text for run in edited.runs] == ["NEW", " then OLD"]


@pytest.mark.parametrize("extension", ["docx", "pptx"])
def test_cross_run_matches_are_non_overlapping_and_counted_once(tmp_path, extension):
    document, paragraph = _document(extension)
    for text in ["a", "aa", "a", "a"]:
        _run(paragraph, text, bold=True)
    path = tmp_path / f"source.{extension}"
    document.save(path)

    result = _patch(path, "aa", "aaa", True)

    _, edited = _reopen(result, extension)
    assert edited.text == "aaaaaaa"
    assert result["replacements"] == 2
    assert len(edited.runs) == 4


@pytest.mark.parametrize("extension", ["docx", "pptx"])
def test_line_break_replacement_and_insertion_preserve_styles(tmp_path, extension):
    document, paragraph = _document(extension)
    _run(paragraph, "before OLD", bold=True)
    if extension == "docx":
        paragraph.runs[0].add_break()
    else:
        paragraph.add_line_break()
    _run(paragraph, "TAIL after", italic=True)
    path = tmp_path / f"source.{extension}"
    document.save(path)
    separator = "\n" if extension == "docx" else "\v"

    result = _patch(path, f"OLD{separator}TAIL", "New\nLine\tTab")

    _, edited = _reopen(result, extension)
    assert edited.text == f"before New{separator}Line\tTab after"
    assert edited.runs[0].font.bold is True
    assert edited.runs[-1].font.italic is True
    assert edited.runs[-1].text == " after"


def _word_hyperlink(paragraph, text):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE

    link = OxmlElement("w:hyperlink")
    rid = paragraph.part.relate_to("https://example.com/proposal", RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    properties.append(OxmlElement("w:i"))
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)
    return link, rid


@pytest.mark.parametrize("match_inside_link", [False, True])
def test_word_hyperlink_survives_cross_run_replacement(tmp_path, match_inside_link):
    document, paragraph = _document("docx")
    if match_inside_link:
        link, rid = _word_hyperlink(paragraph, "OL")
        _run(paragraph, "D suffix", bold=True)
    else:
        _run(paragraph, "OL", bold=True)
        _run(paragraph, "D ", italic=True)
        link, rid = _word_hyperlink(paragraph, "Visit")
    path = tmp_path / "linked.docx"
    document.save(path)

    result = _patch(path, "OLD", "NEW")

    edited, paragraph = _reopen(result, "docx")
    assert paragraph.text == ("NEW suffix" if match_inside_link else "NEW Visit")
    assert len(paragraph.hyperlinks) == 1
    assert paragraph.hyperlinks[0].text == ("NEW" if match_inside_link else "Visit")
    assert edited.part.rels[rid].target_ref == "https://example.com/proposal"
    assert paragraph.hyperlinks[0].runs[0].italic is True


def test_word_inline_picture_in_edited_run_is_not_removed(tmp_path):
    document, paragraph = _document("docx")
    run = _run(paragraph, "", bold=True)
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII="
    )
    run.add_picture(io.BytesIO(png))
    run.add_text("Old")
    path = tmp_path / "picture.docx"
    document.save(path)
    drawing = etree.tostring(run._r.xpath("./w:drawing")[0])

    result = _patch(path, "Old", "New")

    edited, paragraph = _reopen(result, "docx")
    assert paragraph.text == "New"
    assert len(edited.inline_shapes) == 1
    assert etree.tostring(paragraph.runs[0]._r.xpath("./w:drawing")[0]) == drawing
    with zipfile.ZipFile(io.BytesIO(result["_persisted_bytes"])) as package:
        assert package.read("word/media/image1.png") == png


def test_pptx_hyperlink_survives_cross_run_replacement(tmp_path):
    document, paragraph = _document("pptx")
    run = _run(paragraph, "OL", bold=True)
    run.hyperlink.address = "https://example.com/proposal"
    _run(paragraph, "D suffix", italic=True)
    path = tmp_path / "linked.pptx"
    document.save(path)

    result = _patch(path, "OLD", "NEW")

    _, edited = _reopen(result, "pptx")
    assert edited.text == "NEW suffix"
    assert edited.runs[0].hyperlink.address == "https://example.com/proposal"
    assert edited.runs[0].font.bold is True
    assert edited.runs[1].font.italic is True


def test_word_merged_cells_and_linked_headers_are_edited_once(tmp_path):
    from docx import Document

    document = Document()
    document.add_table(rows=2, cols=2).cell(0, 0).merge(document.tables[0].cell(1, 1)).text = "Old"
    document.sections[0].header.paragraphs[0].text = "Old"
    document.add_section()
    assert document.sections[1].header.is_linked_to_previous
    path = tmp_path / "merged.docx"
    document.save(path)

    result = _patch(path, "Old", "Old updated", True)

    edited, _ = _reopen(result, "docx")
    assert result["replacements"] == 2
    assert edited.tables[0].cell(0, 0).text == "Old updated"
    assert edited.sections[0].header.paragraphs[0].text == "Old updated"
    assert edited.sections[1].header.is_linked_to_previous


def test_word_text_patch_does_not_create_undefined_headers_and_footers(tmp_path):
    document, paragraph = _document("docx")
    _run(paragraph, "Old")
    path = tmp_path / "no-headers.docx"
    document.save(path)
    with zipfile.ZipFile(path) as package:
        original_parts = set(package.namelist())

    result = _patch(path, "Old", "New", True)

    _reopen(result, "docx")
    with zipfile.ZipFile(io.BytesIO(result["_persisted_bytes"])) as package:
        assert set(package.namelist()) == original_parts


def test_word_replace_first_respects_interleaved_table_order(tmp_path):
    from docx import Document

    document = Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Old in table"
    document.add_paragraph("Old in later paragraph")
    path = tmp_path / "table-first.docx"
    document.save(path)

    result = _patch(path, "Old", "New")

    edited, paragraph = _reopen(result, "docx")
    assert edited.tables[0].cell(0, 0).text == "New in table"
    assert paragraph.text == "Old in later paragraph"


def test_word_explicitly_shared_header_part_is_edited_once(tmp_path):
    from copy import deepcopy
    from docx import Document

    document = Document()
    document.add_paragraph("Body")
    document.sections[0].header.paragraphs[0].text = "Old"
    reference = deepcopy(document.sections[0]._sectPr.xpath("./w:headerReference")[0])
    document.add_section()._sectPr.insert(0, reference)
    path = tmp_path / "shared-part.docx"
    document.save(path)

    result = _patch(path, "Old", "Old updated", True)

    edited, _ = _reopen(result, "docx")
    assert result["replacements"] == 1
    assert all(section.header.paragraphs[0].text == "Old updated" for section in edited.sections)


@pytest.mark.parametrize("field_type", ["word-simple", "word-complex", "powerpoint"])
def test_field_cache_is_preserved_when_surrounding_text_changes(tmp_path, field_type):
    from docx.oxml import OxmlElement as WordElement
    from docx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement as SlideElement

    extension = "pptx" if field_type == "powerpoint" else "docx"
    document, paragraph = _document(extension)
    _run(paragraph, "OL", bold=True)
    _run(paragraph, "D ", italic=True)
    if field_type == "word-simple":
        field = WordElement("w:fldSimple")
        field.set(qn("w:instr"), " DATE ")
        run = WordElement("w:r")
        text = WordElement("w:t")
        text.text = "OLD"
        run.append(text)
        field.append(run)
        paragraph._p.append(field)
        field_nodes = [field]
    elif field_type == "word-complex":
        field_nodes = []
        for kind in ["begin", "instruction", "separate", "cached", "end"]:
            run = paragraph.add_run()
            if kind == "cached":
                run.text = "OLD"
            elif kind == "instruction":
                instruction = WordElement("w:instrText")
                instruction.text = " DATE "
                run._r.append(instruction)
            else:
                marker = WordElement("w:fldChar")
                marker.set(qn("w:fldCharType"), kind)
                run._r.append(marker)
            field_nodes.append(run._r)
    else:
        field = SlideElement("a:fld")
        field.set("id", "{11111111-1111-1111-1111-111111111111}")
        field.set("type", "slidenum")
        text = SlideElement("a:t")
        text.text = "OLD"
        field.append(text)
        paragraph._p.append(field)
        field_nodes = [field]
    field_xml = [etree.tostring(node) for node in field_nodes]
    path = tmp_path / f"field.{extension}"
    document.save(path)

    result = _patch(path, "OLD", "NEW", True)

    _, edited = _reopen(result, extension)
    assert result["replacements"] == 1
    assert edited.runs[0].text == "NEW"
    # Exact field instructions/cache XML survives, not merely its visible text.
    assert [etree.tostring(node) for node in list(edited._p)[-len(field_nodes) :]] == field_xml


def test_word_multiparagraph_field_cache_is_not_edited(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    paragraph = document.add_paragraph("Old outside field")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    paragraph.add_run()._r.append(begin)
    paragraph.add_run()._r.append(OxmlElement("w:instrText"))
    separator = OxmlElement("w:fldChar")
    separator.set(qn("w:fldCharType"), "separate")
    paragraph.add_run()._r.append(separator)
    cached = document.add_paragraph("Old cached table of contents")
    cached.runs[0].italic = True
    field_xml = etree.tostring(cached._p)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    document.add_paragraph().add_run()._r.append(end)
    path = tmp_path / "multiparagraph-field.docx"
    document.save(path)

    result = _patch(path, "Old", "New", True)

    edited, paragraph = _reopen(result, "docx")
    assert result["replacements"] == 1
    assert paragraph.text == "New outside field"
    assert etree.tostring(edited.paragraphs[1]._p) == field_xml


@pytest.mark.parametrize("extension", ["docx", "pptx"])
def test_unmatched_paragraph_and_empty_run_are_not_rewritten(extension):
    from packages.core.services.file_engine_patches import replace_office_paragraph_text

    _, paragraph = _document(extension)
    _run(paragraph, "Before", bold=True)
    _run(paragraph, "", italic=True)
    before = etree.tostring(paragraph._p)

    assert replace_office_paragraph_text(paragraph, "missing", "New", None) == (0, None)
    assert etree.tostring(paragraph._p) == before


@pytest.mark.parametrize("extension", ["docx", "pptx"])
def test_identical_replacement_does_not_move_text_between_styles(extension):
    from packages.core.services.file_engine_patches import replace_office_paragraph_text

    _, paragraph = _document(extension)
    _run(paragraph, "OL", bold=True)
    _run(paragraph, "D", italic=True)
    before = etree.tostring(paragraph._p)

    assert replace_office_paragraph_text(paragraph, "OLD", "OLD", None) == (1, None)
    assert etree.tostring(paragraph._p) == before


@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("replace_all", [False, True])
def test_replacement_is_independent_of_run_splitting(extension, replace_all):
    from packages.core.services.file_engine_patches import replace_office_paragraph_text

    source = "🙂aba ababa end"
    for split in range(len(source) + 1):
        _, paragraph = _document(extension)
        for value in [source[:split], "", source[split:]]:
            _run(paragraph, value, bold=True)
        count, remaining = replace_office_paragraph_text(
            paragraph,
            "aba",
            "aba🙂",
            None if replace_all else 1,
        )
        assert paragraph.text == source.replace("aba", "aba🙂", -1 if replace_all else 1)
        assert count == (source.count("aba") if replace_all else 1)
        assert remaining == (None if replace_all else 0)
        assert len(paragraph.runs) == 3
        assert all(run.font.bold is True for run in paragraph.runs)


@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("invalid", [None, 12, False, {"text": "new"}])
def test_office_replace_rejects_non_string_new_text(tmp_path, extension, invalid):
    document, paragraph = _document(extension)
    _run(paragraph, "Old")
    path = tmp_path / f"source.{extension}"
    document.save(path)
    original = path.read_bytes()

    result = _patch(path, "Old", invalid)

    assert result.get("error") == "new_text must be a string"
    assert result["operation_index"] == 0
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == original
