"""Native tables use the same create/patch operations and precise grid selectors."""
from __future__ import annotations

import io
import json

import pytest

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools import file_tools
from packages.core.contracts.file_engine import normalize_file_patch_operation
from packages.core.services.file_engine_patches import describe_file_structure
from tests.test_file_engine_structured_patch import file_runtime as file_runtime


def table_operations(extension):
    table = {"op": "table.insert", "rows": [["same", "same"], ["same", None]]}
    if extension == "docx":
        return [{**table, "index": 0, "style": "Table Grid"}]
    return [
        {"op": "slide.insert", "index": 0, "layout_index": 6},
        {**table, "slide": 1, "transform": {"x": 72, "y": 144, "width": 360, "height": 144}},
    ]


def cell_operation(extension, cell="B2", value="Changed"):
    selector = {"table_index": 0} if extension == "docx" else {"slide": 1, "shape_id": 2}
    return {"op": "cell.set", **selector, "cell": cell, "value": value}


def _open(data, extension):
    if extension == "docx":
        from docx import Document

        document = Document(data)
        return document, document.tables[0]
    from pptx import Presentation

    document = Presentation(data)
    return document, document.slides[0].shapes[0].table


def table_source(path):
    extension = path.suffix[1:]
    result = _generate_office_operations_sync(extension, [normalize_file_patch_operation(op) for op in table_operations(extension)])
    assert not result.get("error"), result
    path.write_bytes(result["_persisted_bytes"])
    document, table = _open(str(path), extension)
    cell = table.cell(1, 1)
    paragraph = cell.paragraphs[0] if extension == "docx" else cell.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.font.bold = True
    run.font.italic = True
    document.save(path)


@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("value", ["New value", "Line one\nLine two", None, True, 12.5])
def test_generate_and_patch_table_cells_use_same_operations(tmp_path, extension, value):
    operations = table_operations(extension) + [cell_operation(extension, value=value)]
    result = _generate_office_operations_sync(extension, [normalize_file_patch_operation(op) for op in operations])
    assert not result.get("error"), result
    assert result["operations_applied"] == len(operations)
    document, table = _open(io.BytesIO(result["_persisted_bytes"]), extension)
    assert table.cell(1, 1).text == ("" if value is None else str(value))
    assert table.cell(0, 0).text == table.cell(0, 1).text == table.cell(1, 0).text == "same"
    path = tmp_path / f"table.{extension}"
    document.save(path)
    before = path.read_bytes()
    result = file_tools._apply_office_patch_sequence_sync(str(path), [normalize_file_patch_operation(cell_operation(extension, "A1", "One only"))])
    assert not result.get("error"), result
    _, table = _open(io.BytesIO(result["_persisted_bytes"]), extension)
    assert table.cell(0, 0).text == "One only"
    assert table.cell(0, 1).text == table.cell(1, 0).text == "same"
    assert path.read_bytes() == before
    structure = describe_file_structure(str(path))
    if extension == "docx":
        assert {
            key: structure["tables"][0][key] for key in ("index", "row_count", "column_count")
        } == {"index": 0, "row_count": 2, "column_count": 2}
    else:
        assert structure["shapes"][0]["shape_id"] == 2
        assert structure["shapes"][0]["has_table"] is True
        assert {
            key: structure["shapes"][0]["table"][key] for key in ("row_count", "column_count")
        } == {"row_count": 2, "column_count": 2}
        shape = document.slides[0].shapes[0]
        assert (shape.left.pt, shape.top.pt, shape.width.pt, shape.height.pt) == (72, 144, 360, 144)


@pytest.mark.parametrize("extension", ["docx", "pptx"])
def test_blank_and_multiline_cell_edits_preserve_properties(tmp_path, extension):
    path = tmp_path / f"styled.{extension}"
    table_source(path)
    original, table = _open(str(path), extension)
    properties = table.cell(1, 1)._tc.tcPr.xml
    operations = [cell_operation(extension, value="One\nTwo\nThree"), cell_operation(extension, value="Final\nSecond")]
    result = file_tools._apply_office_patch_sequence_sync(str(path), [normalize_file_patch_operation(op) for op in operations])
    assert not result.get("error"), result
    document, table = _open(io.BytesIO(result["_persisted_bytes"]), extension)
    cell = table.cell(1, 1)
    assert cell.text == "Final\nSecond"
    assert cell._tc.tcPr.xml == properties
    paragraphs = cell.paragraphs if extension == "docx" else cell.text_frame.paragraphs
    assert len(paragraphs) == 2
    for paragraph in paragraphs:
        run = next(run for run in paragraph.runs if run.text)
        assert run.font.bold is True and run.font.italic is True
    if extension == "docx":
        assert table.style.name == "Table Grid"
    else:
        assert document.slides[0].shapes[0].shape_id == original.slides[0].shapes[0].shape_id


@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("target", ["B1", "A2", "B2"])
def test_merged_continuations_reject_without_redirecting(tmp_path, extension, target):
    path = tmp_path / f"merged.{extension}"
    table_source(path)
    document, table = _open(str(path), extension)
    table.cell(0, 0).merge(table.cell(1, 1))
    document.save(path)
    before = path.read_bytes()
    result = file_tools._apply_office_patch_sequence_sync(str(path), [normalize_file_patch_operation(cell_operation(extension, target))])
    assert "top-left anchor" in result["error"]
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before
    result = file_tools._apply_office_patch_sequence_sync(str(path), [normalize_file_patch_operation(cell_operation(extension, "A1"))])
    assert not result.get("error"), result
    _, edited = _open(io.BytesIO(result["_persisted_bytes"]), extension)
    assert edited.cell(0, 0).text == "Changed"
    if extension == "docx":
        assert edited.cell(0, 0).grid_span == 2
        assert edited.cell(0, 0)._tc.vMerge == "restart"
    else:
        assert edited.cell(0, 0).is_merge_origin
        assert edited.cell(0, 0).span_width == edited.cell(0, 0).span_height == 2


def test_word_omitted_grid_positions_do_not_shift_addresses(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    path = tmp_path / "omitted.docx"
    table_source(path)
    document, table = _open(str(path), "docx")
    row = table.rows[1]
    row._tr.remove(row._tr.tc_lst[0])
    before = OxmlElement("w:gridBefore")
    before.set(qn("w:val"), "1")
    row._tr.get_or_add_trPr().append(before)
    document.save(path)
    result = file_tools._apply_office_patch_sequence_sync(str(path), [normalize_file_patch_operation(cell_operation("docx", "B2"))])
    assert not result.get("error"), result
    _, edited = _open(io.BytesIO(result["_persisted_bytes"]), "docx")
    assert edited.rows[1].grid_cols_before == 1
    assert edited.rows[1].cells[0].text == "Changed"
    result = file_tools._apply_office_patch_sequence_sync(str(path), [normalize_file_patch_operation(cell_operation("docx", "A2"))])
    assert "omitted" in result["error"]


@pytest.mark.parametrize("kind", ["field", "spanning_field", "nested_table", "drawing", "bookmark", "revision", "page_break", "ppt_field"])
@pytest.mark.parametrize("operation", ["cell.set", "cell.format"])
def test_cell_set_rejects_non_plain_content_without_loss(tmp_path, kind, operation):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    extension = "pptx" if kind == "ppt_field" else "docx"
    path = tmp_path / f"protected.{extension}"
    table_source(path)
    document, table = _open(str(path), extension)
    cell = table.cell(1, 1)
    if kind == "ppt_field":
        from pptx.oxml.xmlchemy import OxmlElement as PptElement

        field = PptElement("a:fld")
        text = PptElement("a:t")
        text.text = "Cached date"
        field.append(text)
        cell.text_frame.paragraphs[0]._p.append(field)
    elif kind == "nested_table":
        cell.add_table(1, 1).cell(0, 0).text = "Keep nested table"
    elif kind in {"field", "spanning_field"}:
        if kind == "spanning_field":
            start = document.add_paragraph()
            table._tbl.addprevious(start._p)
            end = document.add_paragraph()
        else:
            start = end = cell.paragraphs[0]
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        start.add_run()._r.append(begin)
        cell.paragraphs[0].add_run("Cached date")
        finish = OxmlElement("w:fldChar")
        finish.set(qn("w:fldCharType"), "end")
        end.add_run()._r.append(finish)
    else:
        element = OxmlElement({"drawing": "w:drawing", "bookmark": "w:bookmarkStart", "revision": "w:ins", "page_break": "w:br"}[kind])
        if kind == "page_break":
            element.set(qn("w:type"), "page")
        cell.paragraphs[0].add_run()._r.append(element)
    document.save(path)
    before = path.read_bytes()
    patch = {**cell_operation(extension), "op": operation, "format": {"bold": True}}
    result = file_tools._apply_office_patch_sequence_sync(str(path), [normalize_file_patch_operation(patch)])
    assert "cannot replace" in result["error"]
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before


@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("invalid", [{"cell": "A0"}, {"cell": "A1:B2"}, {"cell": "Z99"}, {"value": {}}, {"value": float("inf")}, {"sheet": "wrong"}, {"value": "bad\x00text"}])
def test_late_invalid_cell_update_returns_no_partial_bytes(tmp_path, extension, invalid):
    path = tmp_path / f"invalid.{extension}"
    table_source(path)
    operations = [cell_operation(extension), {**cell_operation(extension), **invalid}]
    result = file_tools._apply_office_patch_sequence_sync(str(path), [normalize_file_patch_operation(op) for op in operations])
    assert result.get("error"), result
    assert result["operation_index"] == 1
    assert "_persisted_bytes" not in result


@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("invalid", ["missing_value", "missing_selector", "bool_selector", "unknown_selector"])
def test_cell_set_requires_explicit_valid_target_and_value(tmp_path, extension, invalid):
    path = tmp_path / f"selectors.{extension}"
    table_source(path)
    before = path.read_bytes()
    operation = cell_operation(extension)
    selector = "table_index" if extension == "docx" else "shape_id"
    if invalid == "missing_value":
        operation.pop("value")
    elif invalid == "missing_selector":
        operation.pop(selector)
    else:
        operation[selector] = True if invalid == "bool_selector" else 999
    result = file_tools._apply_office_patch_sequence_sync(str(path), [normalize_file_patch_operation(operation)])
    assert result.get("error"), result
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before


def test_powerpoint_cell_set_cannot_target_a_textbox(tmp_path):
    path = tmp_path / "textbox.pptx"
    table_source(path)
    document, _ = _open(str(path), "pptx")
    from pptx.util import Pt

    textbox = document.slides[0].shapes.add_textbox(Pt(0), Pt(0), Pt(100), Pt(100))
    textbox.text = "Keep textbox"
    document.save(path)
    before = path.read_bytes()
    operation = {**cell_operation("pptx"), "shape_id": textbox.shape_id}
    result = file_tools._apply_office_patch_sequence_sync(str(path), [normalize_file_patch_operation(operation)])
    assert "PowerPoint table" in result["error"]
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before


@pytest.mark.parametrize("invalid", [
    {"rows": []}, {"rows": [["A"], ["B", "C"]]}, {"rows": [[{}]]},
    {"style": "Table Grid"}, {"transform": {"x": 72, "y": 144, "width": 360}},
    {"transform": {"x": 72, "y": 144, "width": 0, "height": 144}},
])
def test_powerpoint_table_generation_rejects_invalid_content_and_geometry(invalid):
    operations = table_operations("pptx")
    operations[-1].update(invalid)
    result = _generate_office_operations_sync("pptx", [normalize_file_patch_operation(op) for op in operations])
    assert result.get("error"), result
    assert "_persisted_bytes" not in result


@pytest.mark.asyncio
@pytest.mark.parametrize("extension", ["docx", "pptx"])
@pytest.mark.parametrize("guard", ["resource", "approval", "stale"])
@pytest.mark.parametrize("operation", ["cell.set", "cell.format"])
async def test_table_cell_patch_keeps_native_guards(file_runtime, monkeypatch, extension, guard, operation):
    root, calls = file_runtime
    path = root / f"guarded.{extension}"
    table_source(path)
    before = path.read_bytes()

    async def reject(**_kwargs):
        return json.dumps({"error": f"{guard}_denied"})

    if guard != "stale":
        monkeypatch.setattr(file_tools, "runtime_guard_file_resource_access" if guard == "resource" else "runtime_guard_file_mutation", reject)
    result = json.loads(await file_tools._patch_file(
        "entity", path=path.name, operations=[{**cell_operation(extension), "op": operation, "format": {"bold": True}}], expected_sha256="0" * 64 if guard == "stale" else file_tools._file_meta(str(path))["source_sha256"],
    ))
    assert result.get("error"), result
    assert path.read_bytes() == before
    assert not any("commit" in call for call in calls)
