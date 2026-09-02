"""Native Excel tables share generation, template and patch operations."""
from __future__ import annotations

import io
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo
import pytest

from packages.core.ai.runtime.generated_files import _generate_office_operations_sync
from packages.core.ai.tools.file_tools import _apply_office_patch_sequence_sync
from packages.core.contracts.file_engine import file_patch_operations, normalize_file_patch_operation
from packages.core.services.file_engine_patches import describe_file_structure


def normalized(*operations):
    return [normalize_file_patch_operation(operation) for operation in operations]


def apply(path, *operations):
    return _apply_office_patch_sequence_sync(str(path), normalized(*operations))


def test_table_operations_are_shared_xlsx_xlsm_capabilities():
    expected = {"table.insert", "table.format", "table.delete"}
    assert expected <= set(file_patch_operations("xlsx"))
    assert expected <= set(file_patch_operations("xlsm"))


def test_generation_creates_formats_and_appends_to_native_table(tmp_path):
    result = _generate_office_operations_sync(
        "xlsx",
        normalized(
            {"op": "cell.set", "cell": "A1", "value": "Item"},
            {"op": "cell.set", "cell": "B1", "value": "Amount"},
            {"op": "cell.set", "cell": "A2", "value": "Design"},
            {"op": "cell.set", "cell": "B2", "value": 1200},
            {
                "op": "table.insert",
                "range": "A1:B2",
                "table": "Services",
                "format": {"style_name": "TableStyleDark3", "show_row_stripes": True},
            },
            {"op": "table.format", "table": "Services", "format": {"show_last_column": True}},
            {"op": "row.append", "table": "Services", "row": {"Item": "Support", "Amount": 400}},
        ),
    )
    assert not result.get("error"), result
    path = tmp_path / "tables.xlsx"
    path.write_bytes(result["_persisted_bytes"])
    workbook = load_workbook(path)
    table = workbook.active.tables["Services"]
    assert table.ref == "A1:B3"
    assert table.autoFilter.ref == "A1:B3"
    assert table.tableStyleInfo.name == "TableStyleDark3"
    assert table.tableStyleInfo.showRowStripes is True
    assert table.tableStyleInfo.showLastColumn is True
    assert [workbook.active.cell(3, column).value for column in (1, 2)] == ["Support", 400]
    workbook.close()
    descriptor = describe_file_structure(str(path))["sheets"][0]["tables"][0]
    assert descriptor["style"] == {
        "style_name": "TableStyleDark3",
        "show_first_column": False,
        "show_last_column": True,
        "show_row_stripes": True,
        "show_column_stripes": False,
    }


def test_template_table_format_changes_only_selected_style(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["Item", "Amount", "Formula"])
    sheet.append(["Design", 1200, "=B2*2"])
    target = Table(displayName="Services", ref="A1:C2")
    target.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(target)
    sheet["E1"], sheet["E2"] = "Keep", "Neighbor"
    other = workbook.create_sheet("Other")
    other.append(["Key", "Value"])
    other.append(["A", 1])
    other.add_table(Table(displayName="OtherTable", ref="A1:B2"))
    path = tmp_path / "template.xlsx"
    workbook.save(path)
    workbook.close()

    result = apply(
        path,
        {
            "op": "table.format",
            "sheet": "Sales",
            "table": "services",
            "format": {"style_name": "TableStyleLight8", "show_column_stripes": True},
        },
    )
    assert not result.get("error"), result
    edited = load_workbook(io.BytesIO(result["_persisted_bytes"]), data_only=False)
    table = edited["Sales"].tables["Services"]
    assert table.ref == "A1:C2" and table.tableStyleInfo.name == "TableStyleLight8"
    assert table.tableStyleInfo.showRowStripes is True
    assert table.tableStyleInfo.showColumnStripes is True
    assert edited["Sales"]["C2"].value == "=B2*2"
    assert edited["Sales"]["E2"].value == "Neighbor"
    assert edited["Other"].tables["OtherTable"].ref == "A1:B2"
    edited.close()


@pytest.mark.parametrize(
    ("obstacle", "operation"),
    [
        ("empty-header", {"op": "table.insert", "range": "A1:B2", "table": "NewTable"}),
        ("duplicate-header", {"op": "table.insert", "range": "A1:B2", "table": "NewTable"}),
        ("merged", {"op": "table.insert", "range": "A1:B2", "table": "NewTable"}),
        ("overlap", {"op": "table.insert", "range": "A1:B3", "table": "NewTable"}),
        ("duplicate-name", {"op": "table.insert", "range": "D1:E2", "table": "Existing"}),
        ("invalid-name", {"op": "table.insert", "range": "D1:E2", "table": "A1"}),
        ("invalid-style", {"op": "table.format", "table": "Existing", "format": {"style_name": "Unknown"}}),
    ],
)
def test_table_operations_reject_invalid_or_lossy_changes_atomically(tmp_path, obstacle, operation):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["One", "Two", None, "Four", "Five"])
    sheet.append([1, 2, None, 4, 5])
    existing = Table(displayName="Existing", ref="A1:B2")
    existing.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(existing)
    if obstacle == "empty-header":
        sheet["B1"] = None
        del sheet.tables["Existing"]
    elif obstacle == "duplicate-header":
        sheet["B1"] = "one"
        del sheet.tables["Existing"]
    elif obstacle == "merged":
        del sheet.tables["Existing"]
        sheet.merge_cells("A1:B1")
    path = tmp_path / f"{obstacle}.xlsx"
    workbook.save(path)
    workbook.close()
    before = path.read_bytes()
    result = apply(
        path,
        {"op": "cell.set", "cell": "G1", "value": "Rollback"},
        operation,
    )
    assert result.get("operation_index") == 1
    assert "_persisted_bytes" not in result
    assert path.read_bytes() == before


def test_table_format_can_clear_style_without_changing_cells(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["One", "Two"])
    sheet.append([1, 2])
    table = Table(displayName="Plain", ref="A1:B2")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    path = tmp_path / "clear-style.xlsx"
    workbook.save(path)
    workbook.close()
    result = apply(path, {"op": "table.format", "table": "Plain", "format": {"style_name": None}})
    assert not result.get("error"), result
    edited = load_workbook(io.BytesIO(result["_persisted_bytes"]))
    assert edited.active.tables["Plain"].tableStyleInfo is None
    assert edited.active["B2"].value == 2
    edited.close()


def test_table_format_preserves_template_custom_style_when_only_flags_change(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["One", "Two"])
    sheet.append([1, 2])
    table = Table(displayName="Custom", ref="A1:B2")
    table.tableStyleInfo = TableStyleInfo(name="CompanyTableStyle", showRowStripes=True)
    sheet.add_table(table)
    path = tmp_path / "custom-style.xlsx"
    workbook.save(path)
    workbook.close()
    result = apply(
        path,
        {"op": "table.format", "table": "Custom", "format": {"show_first_column": True}},
    )
    assert not result.get("error"), result
    edited = load_workbook(io.BytesIO(result["_persisted_bytes"]))
    style = edited.active.tables["Custom"].tableStyleInfo
    assert style.name == "CompanyTableStyle"
    assert style.showRowStripes is True and style.showFirstColumn is True
    edited.close()


def test_table_insert_rejects_workbook_defined_name_collision(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["One", "Two"])
    sheet.append([1, 2])
    workbook.defined_names.add(DefinedName("ReservedName", attr_text="Sheet!$A$1"))
    path = tmp_path / "defined-name.xlsx"
    workbook.save(path)
    workbook.close()
    result = apply(path, {"op": "table.insert", "range": "A1:B2", "table": "reservedname"})
    assert "already exists" in result.get("error", "")


def test_xlsm_table_patch_preserves_macro_payload(tmp_path):
    created = _generate_office_operations_sync(
        "xlsx",
        normalized(
            {"op": "cell.set", "cell": "A1", "value": "One"},
            {"op": "cell.set", "cell": "A2", "value": 1},
        ),
    )
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(created["_persisted_bytes"])) as archive, zipfile.ZipFile(output, "w") as target:
        for item in archive.infolist():
            payload = archive.read(item.filename)
            if item.filename == "[Content_Types].xml":
                payload = payload.replace(
                    b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                    b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
                )
            target.writestr(item, payload)
        target.writestr("xl/vbaProject.bin", b"opaque-macro-payload")
    path = tmp_path / "macro.xlsm"
    path.write_bytes(output.getvalue())
    result = apply(path, {"op": "table.insert", "range": "A1:A2", "table": "MacroTable"})
    assert not result.get("error"), result
    with zipfile.ZipFile(io.BytesIO(result["_persisted_bytes"])) as archive:
        assert archive.read("xl/vbaProject.bin") == b"opaque-macro-payload"
