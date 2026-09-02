from __future__ import annotations

import asyncio
import hashlib
import json
import os
import tempfile
from functools import wraps
from typing import Any

from packages.core.ai.runtime.file_contracts import FileMutationAction
from packages.core.contracts.file_engine import (
    OPERATION_GENERATION_TYPES,
    TEMPLATE_OPERATION_GENERATION_TYPES,
    TEXT_CONTENT_TYPES,
    file_patch_operations,
    file_type_capability,
    file_type_from_path,
    normalize_file_operation_resources,
    normalize_file_patch_operation,
)
from packages.core.services.office_operation_resources import read_office_operation_resources


def _generate_office_operations_sync(
    file_type: str,
    operations: list[dict],
    *,
    template_bytes: bytes | None = None,
    resources: dict[tuple[str, str], bytes] | None = None,
) -> dict:
    """Use the patch executor on a blank package or an authorized template snapshot."""
    from packages.core.ai.tools.file_tools import _apply_office_patch_sequence_sync

    with tempfile.TemporaryDirectory(prefix="manor-office-generate-") as directory:
        seed = os.path.join(directory, f"blank.{file_type}")
        allowed_types = (
            TEMPLATE_OPERATION_GENERATION_TYPES
            if template_bytes is not None
            else OPERATION_GENERATION_TYPES
        )
        if file_type not in allowed_types:
            raise ValueError("unsupported operation generation type")
        if template_bytes is not None:
            with open(seed, "wb") as target:
                target.write(template_bytes)
        elif file_type == "docx":
            from docx import Document

            Document().save(seed)
        elif file_type == "pptx":
            from pptx import Presentation

            Presentation().save(seed)
        elif file_type == "xlsx":
            from openpyxl import Workbook

            workbook = Workbook()
            try:
                workbook.save(seed)
            finally:
                workbook.close()
        if resources:
            return _apply_office_patch_sequence_sync(seed, operations, resources=resources)
        return _apply_office_patch_sequence_sync(seed, operations)


def _read_office_template_sync(entity_id: str, template: dict[str, str]) -> bytes:
    """Read the approved version without following final or parent symlinks."""
    from packages.core.services.entity_fs import (
        EntityFilesystemStaleWriteError,
        open_entity_file_snapshot,
    )

    with open_entity_file_snapshot(
        entity_id, template["path"], expected_content_sha256=template["expected_sha256"],
    ) as snapshot:
        with open(snapshot.descriptor_path, "rb") as source:
            data = source.read()
        # The descriptor pins identity, not in-place writes by an external process.
        if hashlib.sha256(data).hexdigest() != template["expected_sha256"]:
            raise EntityFilesystemStaleWriteError("Template changed while being read")
        return data


def _serialized_entity_file_mutation(handler):
    """Keep source validation, approval consumption, commit and sync atomic."""

    @wraps(handler)
    async def wrapped(*args, **kwargs):
        entity_id = str(kwargs.get("entity_id") or "")
        from packages.core.ai.runtime.file_actions import (
            runtime_entity_file_root,
            runtime_entity_filesystem_mutation_lock,
            runtime_normalize_entity_file_path,
        )

        clean_name = runtime_normalize_entity_file_path(
            str(kwargs.get("name") or "")
        )
        if not clean_name or (
            kwargs.get("operations") is None and not kwargs.get("content")
        ):
            return json.dumps({
                "error": "name and non-empty content or operations are required"
            })
        entity_root = runtime_entity_file_root(entity_id)
        if not entity_root:
            return json.dumps({"error": "Entity filesystem is not enabled"})
        async with runtime_entity_filesystem_mutation_lock(entity_root):
            return await handler(*args, **kwargs)

    return wrapped


@_serialized_entity_file_mutation
async def runtime_generate_document_file(
    *,
    entity_id: str,
    user_id: str,
    conversation_id: str,
    name: str,
    content: str | None = None,
    file_type: str,
    operations: list[dict] | None = None,
    template: dict[str, str] | None = None,
    approval_token: str | None = None,
    expected_sha256: str | None = None,
    workspace_id: str | None = None,
    task_id: str | None = None,
    agent_id: str | None = None,
    runtime_envelope: Any | None = None,
    options: dict[str, Any] | None = None,
    storage_scope: str = "task",
) -> str:
    """Generate a user-visible document file through the Runtime boundary."""

    from packages.core.ai.runtime.document_actions import (
        runtime_document_markdown_link,
        runtime_document_to_dict,
        runtime_document_viewer_url,
    )
    from packages.core.ai.runtime.file_actions import (
        RuntimeFileCommitError,
        RuntimeFileProjectionError,
        RuntimeFileProjectionTransactionFactory,
        runtime_entity_file_root,
        runtime_get_document_for_entity,
        runtime_guard_file_mutation,
        runtime_guard_file_read_access,
        runtime_normalize_entity_file_path,
        runtime_user_visible_file_path,
    )
    from packages.core.ai.tools.file_tools import _workspace_scoped_new_file_path

    clean_name = runtime_normalize_entity_file_path(name or "")
    if not clean_name or (operations is None and not content):
        return json.dumps({"error": "name and non-empty content or operations are required"})

    patches = None
    if template is not None:
        if operations is None:
            return json.dumps({"error": "template requires Office operations"})
        if not isinstance(template, dict) or set(template) != {"path", "expected_sha256"}:
            return json.dumps({"error": "template requires path and expected_sha256 from read_file"})
        raw_path, digest = template["path"], template["expected_sha256"]
        if (
            not isinstance(raw_path, str)
            or raw_path.strip().startswith(("/", "\\"))
            or any(char in raw_path for char in (":", "\x00"))
            or ".." in raw_path.replace("\\", "/").split("/")
        ):
            return json.dumps({"error": "template.path must be an entity-relative Knowledge path"})
        template_path = runtime_normalize_entity_file_path(raw_path)
        if not runtime_user_visible_file_path(template_path):
            return json.dumps({"error": "template.path must be a user-visible Knowledge path"})
        if (
            not isinstance(digest, str)
            or len(digest) != 64
            or any(char not in "0123456789abcdefABCDEF" for char in digest)
        ):
            return json.dumps({"error": "template.expected_sha256 must be the SHA-256 from read_file"})
        template = {"path": template_path, "expected_sha256": digest.lower()}
    if operations is not None:
        if content is not None or options is not None or expected_sha256:
            return json.dumps({"error": "operations create a new file; omit content/options/expected_sha256 and use patch_file for existing files"})
        if not isinstance(operations, list) or not operations:
            return json.dumps({"error": "operations must contain at least one patch operation"})
        try:
            patches = [normalize_file_patch_operation(operation) for operation in operations]
            patches, operation_resources = normalize_file_operation_resources(
                patches,
                normalize_path=runtime_normalize_entity_file_path,
                visible_path=runtime_user_visible_file_path,
            )
            # Reject non-JSON values before approval hashing or rendering.
            json.dumps(patches, allow_nan=False)
        except (ValueError, TypeError) as exc:
            return json.dumps({"error": str(exc)})

    if storage_scope not in {"task", "workspace"}:
        return json.dumps({"error": "storage_scope must be task or workspace"})
    if storage_scope == "workspace" and not workspace_id:
        return json.dumps({"error": "storage_scope=workspace requires an active Workspace context"})

    entity_dir = runtime_entity_file_root(entity_id)
    if not entity_dir:
        return json.dumps({"error": "Entity filesystem is not enabled"})
    clean_name = await _workspace_scoped_new_file_path(
        entity_id=entity_id,
        entity_root=entity_dir,
        workspace_id=workspace_id,
        task_id=None if storage_scope == "workspace" else task_id,
        path=clean_name,
        expected_sha256=expected_sha256,
    )

    clean_file_type = str(file_type or "txt").lstrip(".").lower()
    if "." not in os.path.basename(clean_name):
        clean_name = f"{clean_name}.{clean_file_type}"
    if not runtime_user_visible_file_path(clean_name):
        return json.dumps({"error": "Cannot upload hidden/system document path"})

    ext = os.path.splitext(clean_name)[1].lower()
    detected_type = file_type_from_path(clean_name)
    if patches is not None:
        allowed_types = (
            TEMPLATE_OPERATION_GENERATION_TYPES
            if template is not None
            else OPERATION_GENERATION_TYPES
        )
        if detected_type not in allowed_types or detected_type != clean_file_type:
            return json.dumps({
                "error": (
                    "operations require matching file_type and extension: docx/pptx/xlsx; "
                    "xlsm requires a matching template"
                ),
            })
        if template is not None and file_type_from_path(template["path"]) != detected_type:
            return json.dumps({"error": "template and output must have the same Office file type"})
        for index, patch in enumerate(patches):
            if patch["operation"] not in file_patch_operations(detected_type):
                return json.dumps({"error": "unsupported_operation_for_file_type", "operation_index": index, "capabilities": file_type_capability(detected_type)})
    elif detected_type not in TEXT_CONTENT_TYPES | {"docx", "pptx", "pdf"}:
        return json.dumps({
            "error": "unsupported_document_generation_type",
            "capabilities": file_type_capability(detected_type),
            "hint": "kind=document writes text, DOCX, PPTX or PDF. Use the reported generate_kind for other formats; never write text under a binary extension.",
        })
    target = None
    if entity_dir:
        target = os.path.normpath(os.path.join(entity_dir, clean_name))
        if os.path.commonpath([os.path.normpath(entity_dir), target]) != os.path.normpath(entity_dir):
            return json.dumps({"error": "Path traversal detected"})
    if not target or not entity_dir:
        return json.dumps({"error": "Entity filesystem is not enabled"})

    clean_expected_sha256 = str(expected_sha256 or "").strip()
    if os.path.lexists(target) and not clean_expected_sha256:
        return json.dumps({"error": "file_already_exists", "hint": "Read the existing file and use patch_file with expected_sha256."})

    if clean_expected_sha256:
        stale = await _runtime_guard_generated_file_expected_source(
            abs_path=target,
            path=clean_name,
            expected_sha256=clean_expected_sha256,
        )
        if stale:
            if isinstance(stale, str):
                return stale
            return stale.error or json.dumps({"error": "source_changed"})
    else:
        from packages.core.ai.tools.file_tools import ExpectedSourceGuardResult

        stale = ExpectedSourceGuardResult()

    content_text = str(content) if content is not None else ""
    approval_payload = {
        "content_sha256": hashlib.sha256(content_text.encode("utf-8")).hexdigest(),
        "file_type": clean_file_type,
        "options": options or {},
    }
    if patches is not None:
        approval_payload = {"file_type": clean_file_type, "operations": patches}
    template_bytes = None
    if template is not None:
        blocked = await runtime_guard_file_read_access(
            entity_id=entity_id, user_id=user_id or None, workspace_id=workspace_id,
            runtime_envelope=runtime_envelope, tool_name="generate_file", paths=[template["path"]],
        )
        if blocked:
            return blocked
        from packages.core.services.entity_fs import EntityFilesystemError, EntityFilesystemStaleWriteError

        try:
            template_bytes = await asyncio.to_thread(_read_office_template_sync, entity_id, template)
        except EntityFilesystemStaleWriteError:
            return json.dumps({"error": "source_changed", "path": template["path"], "hint": "Read the template again before generating a new file."})
        except (EntityFilesystemError, OSError):
            return json.dumps({"error": "template_unavailable", "path": template["path"], "hint": "Read an existing, non-symlink Knowledge file as the template."})
        approval_payload["template"] = template
    resource_bytes: dict[tuple[str, str], bytes] = {}
    if patches is not None and operation_resources:
        blocked = await runtime_guard_file_read_access(
            entity_id=entity_id,
            user_id=user_id or None,
            workspace_id=workspace_id,
            runtime_envelope=runtime_envelope,
            tool_name="generate_file",
            paths=[source["path"] for source in operation_resources],
        )
        if blocked:
            return blocked
        from packages.core.services.entity_fs import EntityFilesystemError, EntityFilesystemStaleWriteError

        try:
            resource_bytes = await asyncio.to_thread(
                read_office_operation_resources, entity_id, operation_resources,
            )
        except EntityFilesystemStaleWriteError:
            return json.dumps({
                "error": "source_changed",
                "paths": [source["path"] for source in operation_resources],
                "hint": "Read the picture source again before generating the presentation.",
            })
        except (EntityFilesystemError, OSError, ValueError) as exc:
            return json.dumps({
                "error": "picture_source_unavailable",
                "detail": str(exc),
                "paths": [source["path"] for source in operation_resources],
            })
    blocked = await runtime_guard_file_mutation(
        entity_id=entity_id,
        user_id=user_id or None,
        conversation_id=conversation_id or None,
        workspace_id=workspace_id,
        task_id=task_id,
        runtime_envelope=runtime_envelope,
        tool_name="generate_file",
        action=FileMutationAction.CREATE_DOCUMENT,
        paths=[clean_name],
        approval_token=approval_token,
        content_preview=(
            json.dumps(approval_payload if template is not None else patches, ensure_ascii=False)
            if patches is not None else content_text
        ),
        approval_payload=approval_payload,
    )
    if blocked:
        return blocked

    binary_content: bytes | None = None
    operation_summary = {}
    if patches is not None:
        try:
            template_args = {"template_bytes": template_bytes} if template is not None else {}
            if resource_bytes:
                template_args["resources"] = resource_bytes
            generated = await asyncio.to_thread(_generate_office_operations_sync, detected_type, patches, **template_args)
            if generated.get("error"):
                return json.dumps(generated)
            binary_content = generated.pop("_persisted_bytes")
            operation_summary = {key: generated[key] for key in ("operations_applied", "operation_results")}
        except Exception as exc:  # noqa: BLE001
            return json.dumps({"error": f"Office operation generation failed: {exc}"})
    elif ext == ".pptx":
        try:
            from packages.core.services.docgen_service import generate_pptx

            title = os.path.splitext(clean_name)[0]
            for line in content_text.split("\n"):
                if line.startswith("# ") and not line.startswith("## "):
                    title = line[2:].strip()
                    break
            binary_content = await generate_pptx(title, content_text)
        except Exception as exc:  # noqa: BLE001
            # Never persist UTF-8 text with a .pptx extension. That fallback
            # produces a corrupt download while claiming the PowerPoint MIME.
            return json.dumps({"error": f"PowerPoint generation failed: {exc}"})
    elif ext == ".docx":
        try:
            from packages.core.services.docgen_service import generate_docx

            docx_options = dict(options or {})
            raw_cover = str(docx_options.get("cover_image") or "").strip()
            if raw_cover:
                candidate = raw_cover[7:] if raw_cover.startswith("file://") else raw_cover
                if not os.path.isabs(candidate):
                    candidate = os.path.join(entity_dir, candidate)
                candidate = os.path.normpath(candidate)
                try:
                    inside_entity = os.path.commonpath([os.path.normpath(entity_dir), candidate]) == os.path.normpath(entity_dir)
                except ValueError:
                    inside_entity = False
                if inside_entity and os.path.isfile(candidate):
                    docx_options["cover_image"] = candidate
                else:
                    docx_options.pop("cover_image", None)
            title = os.path.splitext(os.path.basename(clean_name))[0]
            binary_content = await generate_docx(title, content_text, docx_options)
        except Exception as exc:  # noqa: BLE001
            return json.dumps({"error": f"Word document generation failed: {exc}"})
    elif ext == ".pdf":
        try:
            from packages.core.services.docgen_service import generate_pdf

            binary_content = await generate_pdf(os.path.splitext(os.path.basename(clean_name))[0], content_text)
        except Exception as exc:  # noqa: BLE001
            return json.dumps({"error": f"PDF generation failed: {exc}"})

    persisted_bytes = binary_content if binary_content is not None else content_text.encode("utf-8")
    file_size = len(persisted_bytes)
    mime_by_ext = {
        "txt": "text/plain",
        "md": "text/markdown",
        "csv": "text/csv",
        "json": "application/json",
        "html": "text/html",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
        "pdf": "application/pdf",
    }
    mime_type = mime_by_ext.get(ext.lstrip("."), "text/plain")

    content_sha256 = hashlib.sha256(persisted_bytes).hexdigest()
    try:
        async with RuntimeFileProjectionTransactionFactory.create(entity_id) as transaction:
            target = transaction.write_bytes(
                clean_name,
                persisted_bytes,
                expected_content_sha256=content_sha256,
                expected_size=file_size,
                allow_empty=False,
                require_missing=not clean_expected_sha256,
                expected_source_version=getattr(stale, "version", None),
            )
            sync = await transaction.project_file(
                abs_path=target,
                entity_root=entity_dir,
                source="ai_generated",
                created_by=user_id or "ai-agent",
                force=True,
                workspace_id=workspace_id,
                task_id=task_id,
                agent_id=agent_id,
                conversation_id=conversation_id or None,
                user_id=user_id or None,
                tool_name="generate_file",
                expected_content_sha256=content_sha256,
            )
            written_meta = await runtime_generated_file_metadata(target)
            doc = await runtime_get_document_for_entity(
                transaction.db,
                entity_id=entity_id,
                document_id=sync.document_id,
            )
            if doc:
                data = runtime_document_to_dict(doc, detail="details")
                data.update({
                    "source_sha256": written_meta["source_sha256"],
                    "mtime_ns": written_meta["mtime_ns"],
                })
                result = json.dumps({"created": True, "document": data, **operation_summary})
            else:
                result = json.dumps({
                    "created": True,
                    **operation_summary,
                    "document": {
                        "document_id": sync.document_id,
                        "name": os.path.basename(clean_name),
                        "viewer_url": runtime_document_viewer_url(sync.document_id),
                        "markdown_link": runtime_document_markdown_link(
                            os.path.basename(clean_name),
                            sync.document_id,
                        ),
                        "fs_path": clean_name,
                        "file_size": file_size,
                        "file_type": clean_file_type,
                        "mime_type": mime_type,
                        "source_sha256": written_meta["source_sha256"],
                        "mtime_ns": written_meta["mtime_ns"],
                    },
                })
            await transaction.commit()
    except RuntimeFileCommitError as exc:
        return json.dumps({"error": f"Entity filesystem is not available: {exc}"})
    except RuntimeFileProjectionError as exc:
        if exc.reason == "storage_limit":
            return json.dumps({"error": (
                "Knowledge base storage limit reached for this plan — the file "
                "was rolled back. Ask the user to free up space or upgrade their "
                "plan, then try again."
            )})
        return json.dumps({
            "error": (
                "Document was not synced to Knowledge and its file write was "
                f"rolled back: {exc.reason}"
            ),
        })
    except Exception as exc:  # noqa: BLE001
        return json.dumps({"error": f"Document was not committed: {exc}"})
    return result


async def runtime_generated_file_metadata(abs_path: str) -> dict[str, Any]:
    """Return stable metadata for a generated user-visible file."""

    from packages.core.ai.tools.file_tools import _file_meta

    return _file_meta(abs_path)


async def runtime_guard_generated_file_expected_source(
    *,
    abs_path: str,
    path: str,
    expected_sha256: str,
) -> str | None:
    """Reject generated-file writes when the target source changed."""

    result = await _runtime_guard_generated_file_expected_source(
        abs_path=abs_path,
        path=path,
        expected_sha256=expected_sha256,
    )
    return result if isinstance(result, str) else result.error


async def _runtime_guard_generated_file_expected_source(
    *,
    abs_path: str,
    path: str,
    expected_sha256: str,
):
    from packages.core.ai.tools.file_tools import _guard_expected_source_sha

    return await _guard_expected_source_sha(
        abs_path=abs_path,
        path=path,
        expected_sha256=expected_sha256,
    )
