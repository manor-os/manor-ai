"""Shared deterministic file transforms; return bytes, never persist them."""

from __future__ import annotations

from bisect import bisect_left, bisect_right
import copy
from contextlib import contextmanager
import csv
import html
import io
import json
import math
import re
import zipfile
from itertools import islice
from pathlib import Path
from typing import Any

from packages.core.contracts.file_engine import (
    EXCEL_MAX_CELL_CHARS,
    EXCEL_MAX_COLUMN,
    EXCEL_MAX_ROW,
    OfficeChartType,
    PresentationShapePreset,
    file_type_from_path,
)
from packages.core.services.office_operation_resources import operation_resource_key


def word_complex_field_nodes(root: Any) -> set[Any]:
    """Locate cached text across paragraphs in one Word story, e.g. a TOC."""
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    text_tags = {ns + name for name in ("t", "tab", "br", "cr")}
    protected, depth = set(), 0
    for node in root.iter():
        if node.tag == ns + "fldChar":
            kind = node.get(ns + "fldCharType")
            if kind == "begin":
                depth += 1
            elif kind == "end":
                depth = max(0, depth - 1)
        elif depth and node.tag in text_tags:
            protected.add(node)
    return protected


def replace_office_paragraph_text(
    paragraph: Any, old: str, new: str, remaining: int | None, *, protected_nodes: set[Any] | None = None
) -> tuple[int, int | None]:
    """Match original text once, changing text nodes without clearing runs/objects.

    Inserted text inherits the first matched run. Unmatched text retains its
    original runs, links and styles. Fields and opaque objects are boundaries,
    not plain text that can safely be flattened or have its cached value edited.
    """
    if not old or remaining == 0:
        return 0, remaining
    if "\x00" in old or "\x00" in new:
        raise ValueError("Office text cannot contain NUL characters")
    paragraph_xml = paragraph._p
    ns = paragraph_xml.tag.rsplit("}", 1)[0] + "}"
    word = ns == "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    containers = {ns + name for name in ("p", "r", "hyperlink")}
    ignored = {
        ns + name
        for name in (
            "pPr",
            "rPr",
            "endParaRPr",
            "bookmarkStart",
            "bookmarkEnd",
            "commentRangeStart",
            "commentRangeEnd",
            "proofErr",
            "lastRenderedPageBreak",
        )
    }
    nodes: list[tuple[Any, str]] = []
    field_depth = 0
    protected_nodes = protected_nodes or set()

    def visit(element):
        nonlocal field_depth
        for child in element:
            tag = child.tag
            if child in protected_nodes:
                nodes.append((None, "\x00"))
            elif word and tag == ns + "fldChar":
                kind = child.get(ns + "fldCharType")
                if kind == "begin":
                    field_depth += 1
                elif kind == "end":
                    field_depth = max(0, field_depth - 1)
                nodes.append((None, "\x00"))
            elif tag in containers:
                visit(child)
            elif field_depth or tag in ignored:
                continue
            elif tag == ns + "t":
                nodes.append((child, child.text or ""))
            elif tag == ns + "br" and (not word or child.get(ns + "type", "textWrapping") == "textWrapping"):
                nodes.append((child, "\n" if word else "\v"))
            elif word and tag in {ns + "tab", ns + "cr"}:
                nodes.append((child, "\t" if tag == ns + "tab" else "\n"))
            else:
                nodes.append((None, "\x00"))

    visit(paragraph_xml)
    # Empty runs must survive, but cannot own a character in a match.
    nodes = [(node, text) for node, text in nodes if text]
    source = "".join(text for _, text in nodes)
    matches = []
    for match in re.finditer(re.escape(old), source):
        matches.append((match.start(), match.end()))
        if remaining is not None and len(matches) >= remaining:
            break
    if not matches:
        return 0, remaining
    if old == new:
        return len(matches), None if remaining is None else remaining - len(matches)
    starts, ends, offset = [], [], 0
    updated = [text for _, text in nodes]
    for text in updated:
        starts.append(offset)
        offset += len(text)
        ends.append(offset)
    # Reverse original ranges so an insertion cannot shift an earlier match.
    for start, end in reversed(matches):
        first, last = bisect_right(ends, start), bisect_left(starts, end) - 1
        prefix = updated[first][: start - starts[first]]
        suffix = updated[last][end - starts[last] :]
        updated[first] = prefix + new + (suffix if first == last else "")
        if first != last:
            for index in range(first + 1, last):
                updated[index] = ""
            updated[last] = suffix
    for (node, original), text in zip(nodes, updated):
        if node is not None and text != original:
            _set_office_text_node(node, text, word=word, ns=ns)
    return len(matches), None if remaining is None else remaining - len(matches)


def _set_office_text_node(node: Any, text: str, *, word: bool, ns: str) -> None:
    if word:
        if node.tag == ns + "t" and not any(char in text for char in "\t\n\r"):
            node.text = text
            if text != text.strip():
                node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            return
        from docx.oxml import OxmlElement

        # Use the library's tab/break encoding on a *new* run, then replace only
        # this text token. Assigning the original run.text would delete images.
        temporary = OxmlElement("w:r")
        temporary.text = text
        for child in list(temporary):
            node.addprevious(child)
        node.getparent().remove(node)
        return

    from pptx.oxml.xmlchemy import OxmlElement

    chunks = re.split("\n|\v", text)
    if node.tag == ns + "t":
        node.text = chunks[0]
        anchor = node.getparent()
        properties = anchor.find(ns + "rPr")
        chunks = chunks[1:]
        needs_break = True
    else:
        anchor = node
        properties = node.find(ns + "rPr")
        needs_break = False
    for chunk in chunks:
        if needs_break:
            line_break = OxmlElement("a:br")
            anchor.addnext(line_break)
            anchor = line_break
        run = OxmlElement("a:r")
        if properties is not None:
            run.append(copy.deepcopy(properties))
        run.append(OxmlElement("a:t"))
        run.find(ns + "t").text = chunk
        anchor.addnext(run)
        anchor = run
        needs_break = True
    if node.tag == ns + "br":
        node.getparent().remove(node)


def _integer(value: Any, name: str, minimum: int, maximum: int) -> int:
    if type(value) is not int or not minimum <= value <= maximum:
        raise ValueError(f"{name} must be an integer between {minimum} and {maximum}")
    return value


def _number(value: Any, name: str, minimum: float, maximum: float) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{name} must be a number")
    if not math.isfinite(value) or not minimum <= value <= maximum:
        raise ValueError(f"{name} must be between {minimum} and {maximum}")
    return float(value)


def _json_object(pairs: list[tuple[str, Any]]) -> dict:
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f"Duplicate JSON key: {key}")
        result[key] = value
    return result


def _json_patch(content: str, patch: dict[str, Any]) -> str:
    document = json.loads(content, object_pairs_hook=_json_object)
    pointer = patch.get("pointer")
    if not isinstance(pointer, str) or (pointer and not pointer.startswith("/")):
        raise ValueError("pointer must be a JSON Pointer, for example /nodes/0/label")
    if re.search(r"~(?:[^01]|$)", pointer):
        raise ValueError("Invalid JSON Pointer escape; use ~0 for ~ and ~1 for /")
    parts = [part.replace("~1", "/").replace("~0", "~") for part in pointer.split("/")[1:]]
    operation = patch["operation"]
    if operation != "json.remove" and "value" not in patch:
        raise ValueError("value is required")
    if not parts:
        if operation == "json.remove":
            raise ValueError("Cannot remove the JSON document root")
        document = patch["value"]
    else:
        parent = document
        for part in parts[:-1]:
            if isinstance(parent, list):
                parent = parent[_json_index(part, len(parent) - 1)]
            elif isinstance(parent, dict) and part in parent:
                parent = parent[part]
            else:
                raise ValueError("JSON Pointer parent does not exist")
        key = parts[-1]
        if isinstance(parent, list):
            index = (
                len(parent)
                if key == "-" and operation == "json.add"
                else _json_index(
                    key,
                    len(parent) if operation == "json.add" else len(parent) - 1,
                )
            )
            if operation == "json.add":
                parent.insert(index, patch["value"])
            elif operation == "json.remove":
                parent.pop(index)
            else:
                parent[index] = patch["value"]
        elif isinstance(parent, dict):
            if operation != "json.add" and key not in parent:
                raise ValueError("JSON Pointer target does not exist")
            if operation == "json.remove":
                del parent[key]
            else:
                parent[key] = patch["value"]
        else:
            raise ValueError("JSON Pointer parent is not an object or array")
    newline = "\r\n" if "\r\n" in content else "\n"
    return json.dumps(document, ensure_ascii=False, indent=2, allow_nan=False).replace("\n", newline) + (
        newline if content.endswith("\n") else ""
    )


def _json_index(value: str, maximum: int) -> int:
    if not re.fullmatch(r"0|[1-9][0-9]*", value) or len(value) > 9:
        raise ValueError("Invalid JSON array index")
    return _integer(int(value), "JSON array index", 0, maximum)


def _csv_patch(content: str, extension: str, patch: dict[str, Any]) -> str:
    from openpyxl.utils.cell import coordinate_to_tuple

    delimiter = "\t" if extension == "tsv" else ","
    rows = list(csv.reader(io.StringIO(content, newline=""), delimiter=delimiter, strict=True))
    if patch["operation"] == "set_cell":
        reference = str(patch.get("cell") or "")
        if not re.fullmatch(r"[A-Za-z]{1,3}[1-9][0-9]{0,6}", reference):
            raise ValueError("cell must be an A1 coordinate")
        row, column = coordinate_to_tuple(reference)
        if row > len(rows) or column > len(rows[row - 1]):
            raise ValueError("cell is outside the existing CSV/TSV table")
        if "value" not in patch:
            raise ValueError("value is required")
        rows[row - 1][column - 1] = _csv_value(patch["value"])
    else:
        values = patch.get("values")
        if not isinstance(values, list) or not values:
            raise ValueError("values must be a non-empty array for CSV/TSV row.append")
        if rows and len(values) != len(rows[0]):
            raise ValueError("values must match the number of CSV/TSV columns")
        rows.append([_csv_value(value) for value in values])
    newline = "\r\n" if "\r\n" in content else "\n"
    output = io.StringIO(newline="")
    csv.writer(output, delimiter=delimiter, lineterminator=newline).writerows(rows)
    result = output.getvalue()
    return result if content.endswith("\n") else result.removesuffix(newline)


def _csv_value(value: Any) -> str:
    if value is None:
        return ""
    if type(value) not in {str, int, float, bool}:
        raise ValueError("CSV/TSV cell values must be scalar")
    return str(value)


def apply_text_patch_sequence(abs_path: str, patches: list[dict[str, Any]]) -> dict[str, Any]:
    raw = Path(abs_path).read_bytes()
    if raw.startswith((b"%PDF-", b"PK\x03\x04", b"\xd0\xcf\x11\xe0")) or (
        raw.startswith(b"RIFF") and raw[8:12] in {b"WEBP", b"WAVE", b"AVI "}
    ):
        return {"error": "unsupported_binary_edit"}
    bom = b"\xef\xbb\xbf" if raw.startswith(b"\xef\xbb\xbf") else b""
    try:
        content = raw[len(bom) :].decode("utf-8")
    except UnicodeDecodeError:
        return {"error": "unsupported_text_encoding", "hint": "Convert this file to UTF-8 before patching."}
    if "\x00" in content:
        return {"error": "unsupported_binary_edit"}
    results = []
    replacements = 0
    for index, patch in enumerate(patches):
        try:
            operation = patch["operation"]
            result: dict[str, Any] = {"operation": operation}
            if operation == "replace_text":
                old, new = patch.get("old_text"), patch.get("new_text", "")
                if not isinstance(old, str) or not old:
                    raise ValueError("old_text is required")
                if not isinstance(new, str):
                    raise ValueError("new_text must be a string")
                if old not in content:
                    raise ValueError("old_text not found in file. Ensure exact match including whitespace.")
                replace_all = bool(patch.get("replace_all", False))
                count = content.count(old) if replace_all else 1
                content = content.replace(old, new, -1 if replace_all else 1)
                replacements += count
                result.update(replacements=count, replace_all=replace_all)
            elif operation.startswith("json."):
                content = _json_patch(content, patch)
                result["pointer"] = patch.get("pointer")
            else:
                content = _csv_patch(content, file_type_from_path(abs_path), patch)
            results.append(result)
        except (ValueError, TypeError, KeyError, IndexError, csv.Error) as exc:
            return {"error": str(exc), "operation_index": index}
    return {
        "patched": True,
        "edited": True,
        "operation": "patch",
        "file_type": file_type_from_path(abs_path),
        "operations_applied": len(results),
        "operation_results": results,
        "replacements": replacements,
        "_persisted_bytes": bom + content.encode("utf-8"),
    }


def _office_table_rows(rows: Any) -> list[list[Any]]:
    if not isinstance(rows, list) or not 1 <= len(rows) <= 200:
        raise ValueError("rows must contain 1 to 200 table rows")
    if not isinstance(rows[0], list) or not 1 <= len(rows[0]) <= 50:
        raise ValueError("each table row must contain 1 to 50 cells")
    if any(not isinstance(row, list) or len(row) != len(rows[0]) for row in rows):
        raise ValueError("table rows must have equal column counts")
    return [[spreadsheet_cell_value(value) for value in row] for row in rows]


def _office_table_cell(table: Any, reference: Any, *, word: bool) -> Any:
    from openpyxl.utils.cell import coordinate_to_tuple

    if not isinstance(reference, str) or not re.fullmatch(r"[A-Za-z]{1,3}[1-9][0-9]{0,6}", reference):
        raise ValueError("cell must be a single A1 coordinate within the table")
    row, column = coordinate_to_tuple(reference.upper())
    _integer(row, "cell row", 1, len(table.rows))
    _integer(column, "cell column", 1, len(table.columns))
    if word:
        from docx.table import _Cell

        # Word rows can omit leading/trailing grid positions. table.cell()
        # flattens the package and can select the wrong cell in those rows.
        selected_row = table.rows[row - 1]
        cursor = selected_row.grid_cols_before
        for tc in selected_row._tr.tc_lst:
            if cursor <= column - 1 < cursor + tc.grid_span:
                if column - 1 != cursor or tc.vMerge == "continue":
                    raise ValueError("cell targets a merged continuation; use its top-left anchor")
                return _Cell(tc, table)
            cursor += tc.grid_span
        raise ValueError("cell targets an omitted Word table grid position")
    cell = table.cell(row - 1, column - 1)
    if cell.is_spanned:
        raise ValueError("cell targets a merged continuation; use its top-left anchor")
    return cell


def _validate_plain_office_paragraphs(paragraphs: list[Any], *, word: bool, operation: str, protected_nodes: set[Any] | None = None) -> None:
    """Preflight all targeted content before changing text or run formatting."""
    ns = paragraphs[0]._p.tag.rsplit("}", 1)[0] + "}"
    containers = {ns + name for name in ("p", "r", "hyperlink")}
    properties = {ns + name for name in ("pPr", "rPr", "endParaRPr")}
    tokens = {ns + name for name in (("t", "tab", "br", "cr") if word else ("t", "br"))}

    def validate(element):
        for child in element:
            if protected_nodes and child in protected_nodes:
                raise ValueError(f"{operation} cannot replace fields or embedded/annotated content")
            if child.tag in containers:
                validate(child)
            elif child.tag not in properties | tokens:
                raise ValueError(f"{operation} cannot replace fields or embedded/annotated content")
            elif word and child.tag == ns + "br" and child.get(ns + "type", "textWrapping") != "textWrapping":
                raise ValueError(f"{operation} cannot replace a page or column break")

    for paragraph in paragraphs:
        validate(paragraph._p)


def _plain_office_cell_paragraphs(cell: Any, *, word: bool, operation: str, protected_nodes: set[Any] | None = None) -> list[Any]:
    paragraphs = list(cell.paragraphs if word else cell.text_frame.paragraphs)
    ns = paragraphs[0]._p.tag.rsplit("}", 1)[0] + "}"
    if word and any(child.tag not in {ns + "tcPr", ns + "p"} for child in cell._tc):
        raise ValueError(f"{operation} cannot replace nested tables or structured cell content")
    _validate_plain_office_paragraphs(paragraphs, word=word, operation=operation, protected_nodes=protected_nodes)
    return paragraphs


def _first_office_run_properties(paragraph_xml: Any) -> Any:
    ns = paragraph_xml.tag.rsplit("}", 1)[0] + "}"
    runs = list(paragraph_xml.iter(ns + "r"))
    first = next((run for run in runs if any(node.text for node in run.iter(ns + "t"))), None)
    if first is None:
        # Empty paragraphs may start with an unformatted placeholder run.
        first = next((run for run in runs if run.find(ns + "rPr") is not None), None)
    if first is not None:
        return first.find(ns + "rPr")
    # Imported blank paragraphs can store typing style only on the paragraph
    # mark (Word) or endParaRPr (PowerPoint), with no run to inherit from.
    properties = paragraph_xml.find(ns + "pPr/" + ns + "rPr") if ns.endswith("wordprocessingml/2006/main}") else paragraph_xml.find(ns + "endParaRPr")
    if properties is not None:
        properties = copy.deepcopy(properties)
        properties.tag = ns + "rPr"
    return properties


def _set_plain_office_paragraph_text(paragraph: Any, text: str, *, word: bool) -> None:
    """Replace one preflighted paragraph without changing its layout or identity."""
    if paragraph.text:
        count, _ = replace_office_paragraph_text(paragraph, paragraph.text, text, 1)
        if count != 1:
            raise ValueError("paragraph text cannot be replaced without changing its structure")
    elif text:
        properties = _first_office_run_properties(paragraph._p)
        run = paragraph.add_run()
        if properties is not None:
            run._r.insert(0, copy.deepcopy(properties))
        if word:
            run.text = text
        else:
            # Run.text escapes controls; encode newlines as native soft breaks.
            run.text = ""
            ns = paragraph._p.tag.rsplit("}", 1)[0] + "}"
            _set_office_text_node(run._r.find(ns + "t"), text, word=False, ns=ns)


def _set_office_paragraph_text(paragraph: Any, text: Any, *, word: bool, protected_nodes: set[Any] | None = None) -> None:
    if not isinstance(text, str):
        raise ValueError("text must be a string; use an empty string to clear the paragraph")
    spreadsheet_cell_value(text)
    _validate_plain_office_paragraphs([paragraph], word=word, operation="text.set", protected_nodes=protected_nodes)
    _set_plain_office_paragraph_text(paragraph, text.replace("\r\n", "\n").replace("\r", "\n"), word=word)


def _set_office_table_cell_text(cell: Any, text: str, *, word: bool, protected_nodes: set[Any] | None = None) -> None:
    """Map replacement paragraphs by position, retaining native cell/run styles."""
    paragraphs = _plain_office_cell_paragraphs(cell, word=word, operation="cell.set", protected_nodes=protected_nodes)
    ns = paragraphs[0]._p.tag.rsplit("}", 1)[0] + "}"

    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = text.split("\n")
    template = copy.deepcopy(paragraphs[-1]._p)
    for index, line in enumerate(lines):
        if index < len(paragraphs):
            paragraph = paragraphs[index]
        else:
            paragraph = cell.add_paragraph() if word else cell.text_frame.add_paragraph()
            # Clone paragraph/run properties, not text or paragraph identifiers.
            for child in template:
                if child.tag in {ns + "pPr", ns + "endParaRPr"}:
                    paragraph._p.append(copy.deepcopy(child))
            run_properties = _first_office_run_properties(template)
            if run_properties is not None:
                run = paragraph.add_run()
                run._r.insert(0, copy.deepcopy(run_properties))
        _set_plain_office_paragraph_text(paragraph, line, word=word)
    for paragraph in paragraphs[len(lines):]:
        paragraph._p.getparent().remove(paragraph._p)


def _set_office_rgb(color: Any, value: str, *, word: bool) -> None:
    if word:
        from docx.shared import RGBColor
    else:
        from pptx.dml.color import RGBColor
    color.rgb = RGBColor.from_string(value)
    if not word:
        # Explicit RGB must not retain old brightness/alpha modifiers.
        rgb = color._xFill.eg_colorChoice
        for modifier in list(rgb):
            rgb.remove(modifier)


def _format_office_font(font: Any, style: dict[str, Any], *, word: bool) -> None:
    size_units = 2 if word else 100
    if "font_size" in style and round(style["font_size"] * size_units) / size_units != style["font_size"]:
        raise ValueError("font_size requires half-point Word or hundredth-point PowerPoint precision")
    if word:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt
    else:
        from pptx.oxml.xmlchemy import OxmlElement

    for key in ("bold", "italic", "underline", "strike", "font_size", "font_color", "font_name"):
        if key not in style:
            continue
        value = style[key]
        if key == "strike" and not word:
            font._rPr.set("strike", "sngStrike" if value else "noStrike")
        elif key in {"bold", "italic", "underline", "strike"}:
            setattr(font, key, value)
            if word and key in {"bold", "italic"}:
                setattr(font, "cs_" + key, value)
        elif key == "font_size":
            if word:
                font.size = Pt(value)
                properties = font._element.get_or_add_rPr()
                complex_size = properties.find(qn("w:szCs"))
                if complex_size is None:
                    complex_size = OxmlElement("w:szCs")
                    properties.sz.addnext(complex_size)
                complex_size.set(qn("w:val"), str(int(value * 2)))
            else:
                # Avoid truncating a hundredth-point via a floating EMU.
                font._rPr.sz = round(value * 100)
        elif key == "font_color":
            _set_office_rgb(font.color, value, word=word)
        else:
            font.name = value
            if word:
                fonts = font._element.get_or_add_rPr().rFonts
                for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                    fonts.attrib.pop(qn("w:" + attribute), None)
                for slot in ("eastAsia", "cs"):
                    fonts.set(qn("w:" + slot), value)
            else:
                anchor = font._rPr.latin
                for slot in ("ea", "cs"):
                    element = font._rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}" + slot)
                    if element is None:
                        element = OxmlElement("a:" + slot)
                        anchor.addnext(element)
                    element.set("typeface", value)
                    anchor = element


def _format_office_paragraph_runs(paragraphs: list[Any], style: dict[str, Any], *, word: bool) -> None:
    if word:
        from docx.text.run import Run
    else:
        from pptx.text.text import Font

    if set(style) - {"fill_color"}:
        for paragraph in paragraphs:
            runs = ([Run(element, paragraph) for element in paragraph._p.xpath("./w:r | ./w:hyperlink/w:r")]
                    if word else list(paragraph.runs))
            for run in runs or [paragraph.add_run()]:
                _format_office_font(run.font, style, word=word)
            if not word:
                _format_office_font(paragraph.font, style, word=word)
                end_properties = paragraph._p.find("{http://schemas.openxmlformats.org/drawingml/2006/main}endParaRPr")
                if end_properties is not None:
                    _format_office_font(Font(end_properties), style, word=word)


def _format_office_table_cell(cell: Any, style: Any, *, word: bool, protected_nodes: set[Any] | None = None) -> None:
    style = _validated_cell_format(style, spreadsheet=False)
    paragraphs = _plain_office_cell_paragraphs(cell, word=word, operation="cell.format", protected_nodes=protected_nodes)
    _format_office_paragraph_runs(paragraphs, style, word=word)
    if "fill_color" in style:
        if word:
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn

            properties = cell._tc.get_or_add_tcPr()
            for previous in properties.findall(qn("w:shd")):
                properties.remove(previous)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:val"), "clear")
            shading.set(qn("w:fill"), style["fill_color"])
            properties.insert_element_before(shading, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark", "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange")
        else:
            cell.fill.solid()
            _set_office_rgb(cell.fill.fore_color, style["fill_color"], word=False)


def _office_table_dimension_updates(
    value: Any,
    key: str,
    *,
    count: int,
    columns: bool,
    word: bool,
) -> dict[int, float]:
    if not isinstance(value, dict) or not value or len(value) > count:
        raise ValueError(f"{key} must be a non-empty object with at most {count} entries")
    updates: dict[int, float] = {}
    for raw_index, raw_value in value.items():
        if columns:
            if not isinstance(raw_index, str) or not re.fullmatch(r"[A-Za-z]{1,2}", raw_index):
                raise ValueError(f"{key} keys must be table column letters such as A or B")
            index = 0
            for character in raw_index.upper():
                index = index * 26 + ord(character) - ord("A") + 1
        else:
            if not isinstance(raw_index, str) or not re.fullmatch(r"[1-9][0-9]*", raw_index):
                raise ValueError(f"{key} keys must be 1-based table row numbers")
            index = int(raw_index)
        if index > count:
            raise ValueError(f"{key} target {raw_index} is outside this table")
        if index - 1 in updates:
            raise ValueError(f"{key} contains duplicate target {raw_index}")
        updates[index - 1] = _office_points(
            raw_value,
            f"{key}.{raw_index}",
            word=word,
            minimum=1,
            maximum=10000,
        )
    return updates


def _word_table_property(properties: Any, name: str) -> Any:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    element = properties.find(qn(f"w:{name}"))
    if element is None:
        element = OxmlElement(f"w:{name}")
        properties.insert_element_before(
            element,
            "w:tblBorders",
            "w:shd",
            "w:tblLayout",
            "w:tblCellMar",
            "w:tblLook",
            "w:tblPrChange",
        )
    return element


def _format_word_table(table: Any, value: Any) -> None:
    from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    allowed = {
        "style_name",
        "alignment",
        "autofit",
        "width",
        "indent",
        "column_widths",
        "row_heights",
        "row_height_rules",
        "repeat_header_rows",
    }
    if not isinstance(value, dict) or not value or set(value) - allowed:
        raise ValueError(f"Word table.format supports: {', '.join(sorted(allowed))}")
    if "style_name" in value:
        style_name = value["style_name"]
        if style_name is not None and (not isinstance(style_name, str) or not style_name.strip()):
            raise ValueError("style_name must be an existing Word table style or null")
        try:
            table.style = style_name
        except KeyError:
            raise ValueError(f"Word table style not found: {style_name}") from None
    if "alignment" in value:
        alignment = value["alignment"]
        alignments = {
            "left": WD_TABLE_ALIGNMENT.LEFT,
            "center": WD_TABLE_ALIGNMENT.CENTER,
            "right": WD_TABLE_ALIGNMENT.RIGHT,
        }
        if alignment not in alignments:
            raise ValueError("alignment must be left, center or right")
        table.alignment = alignments[alignment]
    if "autofit" in value:
        if type(value["autofit"]) is not bool:
            raise ValueError("autofit must be boolean")
        table.autofit = value["autofit"]
    properties = table._tbl.tblPr
    for key, property_name in (("width", "tblW"), ("indent", "tblInd")):
        if key not in value:
            continue
        points = _office_points(value[key], key, word=True, minimum=0 if key == "indent" else 1)
        element = _word_table_property(properties, property_name)
        element.set(qn("w:type"), "dxa")
        element.set(qn("w:w"), str(round(points * 20)))
    if "column_widths" in value:
        updates = _office_table_dimension_updates(
            value["column_widths"], "column_widths", count=len(table.columns), columns=True, word=True,
        )
        for index, points in updates.items():
            table.columns[index].width = Pt(points)
        # Word stores preferred widths both in tblGrid and on each cell. Keep
        # merged spans and omitted leading grid positions aligned with the new
        # grid instead of leaving stale tcW values that override the columns.
        grid_points = [column.width.pt for column in table.columns]
        for row in table.rows:
            cursor = row.grid_cols_before
            for cell_element in row._tr.tc_lst:
                span = cell_element.grid_span
                width = sum(grid_points[cursor:cursor + span])
                cell_width = cell_element.get_or_add_tcPr().get_or_add_tcW()
                cell_width.set(qn("w:type"), "dxa")
                cell_width.set(qn("w:w"), str(round(width * 20)))
                cursor += span
    if "row_heights" in value:
        updates = _office_table_dimension_updates(
            value["row_heights"], "row_heights", count=len(table.rows), columns=False, word=True,
        )
        for index, points in updates.items():
            row = table.rows[index]
            row.height = Pt(points)
            if row.height_rule is None:
                row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    if "row_height_rules" in value:
        rules = value["row_height_rules"]
        if not isinstance(rules, dict) or not rules or len(rules) > len(table.rows):
            raise ValueError(f"row_height_rules must be a non-empty object with at most {len(table.rows)} entries")
        rule_values = {
            "auto": WD_ROW_HEIGHT_RULE.AUTO,
            "at_least": WD_ROW_HEIGHT_RULE.AT_LEAST,
            "exact": WD_ROW_HEIGHT_RULE.EXACTLY,
        }
        for raw_index, rule in rules.items():
            if not isinstance(raw_index, str) or not re.fullmatch(r"[1-9][0-9]*", raw_index):
                raise ValueError("row_height_rules keys must be 1-based table row numbers")
            index = int(raw_index)
            if index > len(table.rows):
                raise ValueError(f"row_height_rules target {raw_index} is outside this table")
            if rule not in rule_values:
                raise ValueError("row height rule must be auto, at_least or exact")
            row = table.rows[index - 1]
            row.height_rule = rule_values[rule]
            if rule == "auto":
                row.height = None
    if "repeat_header_rows" in value:
        count = _integer(value["repeat_header_rows"], "repeat_header_rows", 0, len(table.rows))
        for index, row in enumerate(table.rows):
            properties = row._tr.get_or_add_trPr()
            for previous in properties.findall(qn("w:tblHeader")):
                properties.remove(previous)
            if index < count:
                properties.append(OxmlElement("w:tblHeader"))


def _format_presentation_table(table: Any, value: Any) -> None:
    from pptx.util import Emu

    flag_attributes = {
        "first_row": "first_row",
        "last_row": "last_row",
        "first_column": "first_col",
        "last_column": "last_col",
        "banded_rows": "horz_banding",
        "banded_columns": "vert_banding",
    }
    allowed = set(flag_attributes) | {"column_widths", "row_heights"}
    if not isinstance(value, dict) or not value or set(value) - allowed:
        raise ValueError(f"PowerPoint table.format supports: {', '.join(sorted(allowed))}")
    for key, attribute in flag_attributes.items():
        if key not in value:
            continue
        if type(value[key]) is not bool:
            raise ValueError(f"{key} must be boolean")
        setattr(table, attribute, value[key])
    if "column_widths" in value:
        updates = _office_table_dimension_updates(
            value["column_widths"], "column_widths", count=len(table.columns), columns=True, word=False,
        )
        for index, points in updates.items():
            table.columns[index].width = Emu(round(points * 12700))
    if "row_heights" in value:
        updates = _office_table_dimension_updates(
            value["row_heights"], "row_heights", count=len(table.rows), columns=False, word=False,
        )
        for index, points in updates.items():
            table.rows[index].height = Emu(round(points * 12700))


def _office_points(value: Any, key: str, *, word: bool, minimum: float = 0, maximum: float = 10000) -> float:
    value = _number(value, key, minimum, maximum)
    units = 20 if word else 100
    if round(value * units) / units != value:
        raise ValueError(f"{key} requires {'twentieth' if word else 'hundredth'}-point precision")
    return value


def _validated_office_paragraph_format(
    style: Any,
    *,
    word: bool,
    operation: str,
    allow_empty: bool = False,
) -> tuple[dict[str, Any], dict[str, Any]]:
    font_keys = {"bold", "italic", "font_size", "font_name", "font_color", "underline", "strike"}
    spacing_keys = {"space_before", "space_after", "indent_left", "indent_right", "first_line_indent"}
    word_flags = {"keep_with_next", "keep_together", "page_break_before", "widow_control"}
    allowed = font_keys | spacing_keys | {"alignment", "line_spacing"} | (word_flags if word else {"level"})
    if not isinstance(style, dict) or (not style and not allow_empty) or set(style) - allowed:
        raise ValueError(f"{operation} requires supported format keys: {', '.join(sorted(allowed))}")
    font = {key: value for key, value in style.items() if key in font_keys - {"underline", "strike"}}
    if font:
        font = _validated_cell_format(font, spreadsheet=False)
    for key in {"underline", "strike"} & style.keys():
        if type(style[key]) is not bool:
            raise ValueError(f"{key} must be boolean")
        font[key] = style[key]
    layout = {key: value for key, value in style.items() if key not in font_keys}
    for key, value in layout.items():
        if key in spacing_keys:
            maximum = 10000 if word else 4032
            layout[key] = _office_points(
                value, key, word=word, minimum=-maximum if key == "first_line_indent" else 0,
                maximum=maximum,
            )
        elif key == "line_spacing":
            layout[key] = _number(value, key, 0.1, 20)
        elif key == "alignment":
            if value not in ("left", "center", "right", "justify"):
                raise ValueError("alignment must be left, center, right or justify")
        elif key == "level":
            _integer(value, "level", 0, 8)
        elif type(value) is not bool:
            raise ValueError(f"{key} must be boolean")
    return font, layout


def _format_office_paragraph(paragraph: Any, patch: dict[str, Any], *, word: bool, protected_nodes: set[Any] | None = None) -> None:
    style_name = patch.get("style")
    style = patch.get("format", {})
    if style_name is not None and (not word or not isinstance(style_name, str) or not style_name):
        raise ValueError("style must be an existing Word paragraph style")
    if not style and not style_name:
        raise ValueError("paragraph.format requires style or a supported format key")
    font, layout = _validated_office_paragraph_format(
        style, word=word, operation="paragraph.format", allow_empty=bool(style_name),
    )
    _validate_plain_office_paragraphs([paragraph], word=word, operation="paragraph.format", protected_nodes=protected_nodes)
    if word:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
        if style_name is not None:
            paragraph.style = style_name
        target = paragraph.paragraph_format
        alignment_type = WD_ALIGN_PARAGRAPH
    else:
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Emu
        target = paragraph
        alignment_type = PP_ALIGN
    if font:
        _format_office_paragraph_runs([paragraph], font, word=word)
    for key, value in layout.items():
        if key == "alignment":
            paragraph.alignment = getattr(alignment_type, value.upper())
        elif key in {"indent_left", "indent_right", "first_line_indent"}:
            if word:
                setattr(target, {"indent_left": "left_indent", "indent_right": "right_indent"}.get(key, key), Pt(value))
            else:
                paragraph._p.get_or_add_pPr().set({"indent_left": "marL", "indent_right": "marR", "first_line_indent": "indent"}[key], str(round(value * 12700)))
        elif key in {"space_before", "space_after"}:
            setattr(target, key, Pt(value) if word else Emu(round(value * 12700)))
        else:
            setattr(target, key, value)


def _validated_word_paragraph_style_name(name: Any) -> str:
    if (
        not isinstance(name, str)
        or not name
        or name != name.strip()
        or len(name) > 255
        or any(ord(character) < 32 for character in name)
    ):
        raise ValueError("style must be a 1–255 character Word paragraph style name")
    return name


def _word_paragraph_style(document: Any, name: Any) -> Any:
    from docx.enum.style import WD_STYLE_TYPE

    name = _validated_word_paragraph_style_name(name)
    matches = [style for style in document.styles if style.name == name]
    if not matches:
        raise ValueError(f"Word paragraph style not found: {name}") from None
    if len(matches) != 1:
        raise ValueError(f"Word paragraph style name is ambiguous: {name}")
    style = matches[0]
    if style.type != WD_STYLE_TYPE.PARAGRAPH:
        raise ValueError(f"Word style is not a paragraph style: {name}")
    return style


def _format_word_paragraph_style(style: Any, value: Any) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    font, layout = _validated_office_paragraph_format(
        value, word=True, operation="paragraph_style.format",
    )
    if font:
        _format_office_font(style.font, font, word=True)
    target = style.paragraph_format
    for key, item in layout.items():
        if key == "alignment":
            target.alignment = getattr(WD_ALIGN_PARAGRAPH, item.upper())
        elif key in {"indent_left", "indent_right", "first_line_indent"}:
            setattr(target, {"indent_left": "left_indent", "indent_right": "right_indent"}.get(key, key), Pt(item))
        elif key in {"space_before", "space_after"}:
            setattr(target, key, Pt(item))
        else:
            setattr(target, key, item)


def _word_paragraph_style_is_used(document: Any, style_id: str) -> bool:
    from docx.oxml.ns import qn

    for part in document.part.package.parts:
        root = getattr(part, "element", None)
        if root is None:
            root = getattr(part, "_element", None)
        if root is None:
            continue
        if any(node.get(qn("w:val")) == style_id for node in root.iter(qn("w:pStyle"))):
            return True
    return False


def _word_paragraph_style_is_referenced(document: Any, target: Any) -> bool:
    from docx.oxml.ns import qn

    for style in document.styles:
        if style.style_id == target.style_id:
            continue
        for tag in ("basedOn", "next", "link"):
            reference = style._element.find(qn("w:" + tag))
            if reference is not None and reference.get(qn("w:val")) == target.style_id:
                return True
    return False


def _word_paragraph_style_details(style: Any) -> dict[str, Any]:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    direct: dict[str, Any] = {}
    font = style.font
    for key in ("bold", "italic", "underline", "strike"):
        value = getattr(font, key)
        if value is not None:
            direct[key] = bool(value)
    if font.name is not None:
        direct["font_name"] = font.name
    if font.size is not None:
        direct["font_size"] = font.size.pt
    if getattr(font.color, "rgb", None) is not None:
        direct["font_color"] = str(font.color.rgb)
    paragraph = style.paragraph_format
    alignment = {
        WD_ALIGN_PARAGRAPH.LEFT: "left",
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
    }.get(paragraph.alignment)
    if alignment is not None:
        direct["alignment"] = alignment
    for key, attribute in (
        ("space_before", "space_before"),
        ("space_after", "space_after"),
        ("indent_left", "left_indent"),
        ("indent_right", "right_indent"),
        ("first_line_indent", "first_line_indent"),
    ):
        value = getattr(paragraph, attribute)
        if value is not None:
            direct[key] = value.pt
    if isinstance(paragraph.line_spacing, (int, float)):
        direct["line_spacing"] = paragraph.line_spacing
    for key in ("keep_with_next", "keep_together", "page_break_before", "widow_control"):
        value = getattr(paragraph, key)
        if value is not None:
            direct[key] = bool(value)
    return {
        "name": style.name,
        "style_id": style.style_id,
        "builtin": style.builtin,
        "base_style": style.base_style.name if style.base_style is not None else None,
        "next_style": style.next_paragraph_style.name,
        "format": direct,
    }


def _setup_office_page(document: Any, patch: dict[str, Any], *, word: bool) -> dict[str, Any]:
    allowed = {"width", "height"}
    if word:
        allowed |= {"margin_top", "margin_bottom", "margin_left", "margin_right", "header_distance", "footer_distance"}
    layout = patch.get("format")
    if not isinstance(layout, dict) or not layout or set(layout) - allowed:
        raise ValueError(f"page.setup format must contain only {', '.join(sorted(allowed))} in points")
    if any(key in patch for key in ("index", "slide", "shape_id", "cell", "table_index", "sheet", "style")):
        raise ValueError("page.setup targets Word section_index or the whole PPT canvas, not content selectors")
    layout = {key: _office_points(value, key, word=word, minimum=1 if key in {"width", "height"} else 0)
              for key, value in layout.items()}
    if word:
        from docx.enum.section import WD_ORIENT
        from docx.shared import Pt
        index = _integer(patch.get("section_index"), "section_index", 0, len(document.sections) - 1)
        section = document.sections[index]
        width = layout.get("width", section.page_width.pt if section.page_width is not None else None)
        height = layout.get("height", section.page_height.pt if section.page_height is not None else None)
        def margin(side):
            current = getattr(section, side + "_margin")
            return layout.get("margin_" + side, current.pt if current is not None else 0)
        if width is None or height is None or width <= margin("left") + margin("right") or height <= margin("top") + margin("bottom"):
            raise ValueError("page dimensions must leave positive content space after margins")
        attributes = {"width": "page_width", "height": "page_height"}
        attributes.update({"margin_" + side: side + "_margin" for side in ("top", "bottom", "left", "right")})
        for key, value in layout.items():
            setattr(section, attributes.get(key, key), Pt(value))
        if {"width", "height"} & layout.keys():
            section.orientation = WD_ORIENT.LANDSCAPE if width > height else WD_ORIENT.PORTRAIT
        return {"section_index": index}
    if "section_index" in patch:
        raise ValueError("PowerPoint page.setup does not accept section_index")
    from pptx.util import Emu
    for key, value in layout.items():
        # PowerPoint itself supports slide dimensions from one to 56 inches.
        if not 72 <= value <= 4032:
            raise ValueError("PowerPoint width/height must be between 72 and 4032 points")
        setattr(document, "slide_" + key, Emu(round(value * 12700)))
    return {}


def _format_presentation_shape(shape: Any, style: Any) -> None:
    """Change only requested native DrawingML properties; never flatten objects."""
    from pptx.enum.dml import MSO_FILL_TYPE
    from pptx.enum.text import MSO_ANCHOR
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Emu

    margins = {"margin_left", "margin_right", "margin_top", "margin_bottom"}
    allowed = margins | {"fill_color", "fill_opacity", "gradient_fill", "line_color", "line_width", "line_opacity", "line_dash", "corner_radius", "vertical_alignment", "word_wrap", "shadow"}
    if not isinstance(style, dict) or not style or set(style) - allowed:
        raise ValueError(f"shape.format requires supported format keys: {', '.join(sorted(allowed))}")
    if shape._element.tag.rsplit("}", 1)[-1] not in {"sp", "cxnSp"}:
        raise ValueError("shape.format requires a PowerPoint shape, textbox or connector")
    if shape._element.tag.rsplit("}", 1)[-1] == "cxnSp" and set(style) - {"line_color", "line_width", "line_opacity", "line_dash", "shadow"}:
        raise ValueError("connectors accept only line and shadow formatting")
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    properties = shape._element.spPr

    def color(value):
        return _validated_cell_format({"fill_color": value}, spreadsheet=False)["fill_color"]

    def color_node(value, opacity=1):
        element = OxmlElement("a:srgbClr")
        element.set("val", color(value))
        alpha = OxmlElement("a:alpha")
        alpha.set("val", str(round(_number(opacity, "opacity", 0, 1) * 100000)))
        element.append(alpha)
        return element

    if "gradient_fill" in style and {"fill_color", "fill_opacity"} & style.keys():
        raise ValueError("gradient_fill cannot be combined with fill_color/fill_opacity")
    if "fill_color" in style:
        if style["fill_color"] is None:
            shape.fill.background()
        else:
            shape.fill.solid()
            _set_office_rgb(shape.fill.fore_color, color(style["fill_color"]), word=False)
    if "gradient_fill" in style:
        gradient = style["gradient_fill"]
        if not isinstance(gradient, dict) or set(gradient) != {"angle", "stops"}:
            raise ValueError("gradient_fill requires angle and stops")
        angle = _number(gradient["angle"], "gradient angle", 0, 360)
        stops = gradient["stops"]
        if not isinstance(stops, list) or not 2 <= len(stops) <= 16:
            raise ValueError("gradient_fill requires 2–16 ordered stops")
        fill = properties.get_or_change_to_gradFill()
        fill.clear()
        fill.set("rotWithShape", "1")
        stop_list = OxmlElement("a:gsLst")
        previous = -1
        for stop in stops:
            if not isinstance(stop, dict) or not {"position", "color"} <= stop.keys() or set(stop) - {"position", "color", "opacity"}:
                raise ValueError("gradient stops require position, color and optional opacity")
            position = _number(stop["position"], "gradient position", 0, 1)
            if position < previous:
                raise ValueError("gradient stops must be ordered by position")
            previous = position
            node = OxmlElement("a:gs")
            node.set("pos", str(round(position * 100000)))
            node.append(color_node(stop["color"], stop.get("opacity", 1)))
            stop_list.append(node)
        fill.append(stop_list)
        linear = OxmlElement("a:lin")
        linear.set("ang", str(round((angle % 360) * 60000)))
        linear.set("scaled", "1")
        fill.append(linear)
    if "line_color" in style:
        if style["line_color"] is None:
            shape.line.fill.background()
        else:
            shape.line.fill.solid()
            _set_office_rgb(shape.line.color, color(style["line_color"]), word=False)
    for key in ("fill_opacity", "line_opacity"):
        if key not in style:
            continue
        fill = shape.fill if key == "fill_opacity" else shape.line.fill
        opacity = _number(style[key], key, 0, 1)
        if fill.type != MSO_FILL_TYPE.SOLID:
            raise ValueError(f"{key} requires an existing or explicit solid color")
        solid = fill._xPr.solidFill
        color_element = next(iter(solid), None)
        if color_element is None:
            raise ValueError(f"{key} requires a solid color")
        for child in list(color_element):
            if child.tag in {ns + "alpha", ns + "alphaMod", ns + "alphaOff"}:
                color_element.remove(child)
        alpha = OxmlElement("a:alpha")
        alpha.set("val", str(round(opacity * 100000)))
        color_element.append(alpha)
    if "line_width" in style:
        shape.line.width = Emu(round(_office_points(style["line_width"], "line_width", word=False, maximum=72) * 12700))
    if "line_dash" in style:
        dash = style["line_dash"]
        if dash not in ("solid", "dash", "dot", "dashDot", "lgDash", "lgDashDot", "lgDashDotDot", "sysDash", "sysDot", "sysDashDot", "sysDashDotDot"):
            raise ValueError("line_dash must be a DrawingML preset dash name")
        line = properties.get_or_add_ln()
        for child in list(line):
            if child.tag in {ns + "prstDash", ns + "custDash"}:
                line.remove(child)
        node = OxmlElement("a:prstDash")
        node.set("val", dash)
        line.insert_element_before(node, "a:round", "a:bevel", "a:miter", "a:headEnd", "a:tailEnd", "a:extLst")
    if "corner_radius" in style:
        if properties.prstGeom is None or properties.prstGeom.get("prst") != "roundRect":
            raise ValueError("corner_radius requires a roundRect shape")
        side = min(shape.width, shape.height) / 12700
        if side <= 0:
            raise ValueError("corner_radius requires nonzero shape dimensions")
        radius = _office_points(style["corner_radius"], "corner_radius", word=False, maximum=side / 2)
        shape.adjustments[0] = radius / side
    if margins & style.keys() or {"vertical_alignment", "word_wrap"} & style.keys():
        if not shape.has_text_frame:
            raise ValueError("text layout requires a shape with a text frame")
        frame = shape.text_frame
        for key in margins & style.keys():
            setattr(frame, key, Emu(round(_office_points(style[key], key, word=False, maximum=4032) * 12700)))
        if frame.margin_left + frame.margin_right >= shape.width or frame.margin_top + frame.margin_bottom >= shape.height:
            raise ValueError("text margins must leave positive content space")
        if "vertical_alignment" in style:
            value = style["vertical_alignment"]
            if value not in ("top", "middle", "bottom"):
                raise ValueError("vertical_alignment must be top, middle or bottom")
            frame.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[value]
        if "word_wrap" in style:
            if type(style["word_wrap"]) is not bool:
                raise ValueError("word_wrap must be boolean")
            frame.word_wrap = style["word_wrap"]
    if "shadow" in style:
        if properties.find(ns + "effectDag") is not None:
            raise ValueError("shadow formatting cannot replace a complex effectDag")
        effect_list = properties.get_or_add_effectLst()
        for child in list(effect_list):
            if child.tag == ns + "outerShdw":
                effect_list.remove(child)
        shadow = style["shadow"]
        if shadow is not None:
            if not isinstance(shadow, dict) or set(shadow) != {"color", "opacity", "blur", "distance", "angle"}:
                raise ValueError("shadow requires color, opacity, blur, distance and angle; null clears")
            node = OxmlElement("a:outerShdw")
            for key, xml_key in (("blur", "blurRad"), ("distance", "dist")):
                node.set(xml_key, str(round(_office_points(shadow[key], key, word=False, maximum=4032) * 12700)))
            node.set("dir", str(round((_number(shadow["angle"], "shadow angle", 0, 360) % 360) * 60000)))
            node.append(color_node(shadow["color"], shadow["opacity"]))
            following = next((child for child in effect_list if child.tag in {ns + "prstShdw", ns + "reflection", ns + "softEdge"}), None)
            if following is not None:
                following.addprevious(node)
            else:
                effect_list.append(node)


def _format_presentation_slide(slide: Any, style: Any) -> dict[str, Any]:
    """Set an editable native slide background or restore master inheritance."""
    from pptx.oxml.xmlchemy import OxmlElement

    allowed = {"background_color", "background_gradient"}
    if not isinstance(style, dict) or not style or set(style) - allowed:
        raise ValueError(
            "slide.format requires background_color or background_gradient",
        )
    if set(style) == allowed:
        raise ValueError("background_color and background_gradient are mutually exclusive")

    def color(value: Any) -> str:
        return _validated_cell_format(
            {"fill_color": value}, spreadsheet=False,
        )["fill_color"]

    def color_node(value: Any, opacity: Any = 1) -> Any:
        element = OxmlElement("a:srgbClr")
        element.set("val", color(value))
        alpha = OxmlElement("a:alpha")
        alpha.set(
            "val",
            str(round(_number(opacity, "background opacity", 0, 1) * 100000)),
        )
        element.append(alpha)
        return element

    background = None
    if "background_color" in style and style["background_color"] is not None:
        fill = OxmlElement("a:solidFill")
        fill.append(color_node(style["background_color"]))
        background = OxmlElement("p:bg")
        properties = OxmlElement("p:bgPr")
        properties.append(fill)
        properties.append(OxmlElement("a:effectLst"))
        background.append(properties)
    elif "background_gradient" in style:
        gradient = style["background_gradient"]
        if not isinstance(gradient, dict) or set(gradient) != {"angle", "stops"}:
            raise ValueError("background_gradient requires angle and stops")
        angle = _number(gradient["angle"], "background gradient angle", 0, 360)
        stops = gradient["stops"]
        if not isinstance(stops, list) or not 2 <= len(stops) <= 16:
            raise ValueError("background_gradient requires 2–16 ordered stops")
        fill = OxmlElement("a:gradFill")
        fill.set("rotWithShape", "1")
        stop_list = OxmlElement("a:gsLst")
        previous = -1.0
        for stop in stops:
            if (
                not isinstance(stop, dict)
                or not {"position", "color"} <= stop.keys()
                or set(stop) - {"position", "color", "opacity"}
            ):
                raise ValueError(
                    "background gradient stops require position, color and optional opacity",
                )
            position = _number(stop["position"], "background gradient position", 0, 1)
            if position < previous:
                raise ValueError("background gradient stops must be ordered by position")
            previous = position
            node = OxmlElement("a:gs")
            node.set("pos", str(round(position * 100000)))
            node.append(color_node(stop["color"], stop.get("opacity", 1)))
            stop_list.append(node)
        fill.append(stop_list)
        linear = OxmlElement("a:lin")
        linear.set("ang", str(round((angle % 360) * 60000)))
        linear.set("scaled", "1")
        fill.append(linear)
        background = OxmlElement("p:bg")
        properties = OxmlElement("p:bgPr")
        properties.append(fill)
        properties.append(OxmlElement("a:effectLst"))
        background.append(properties)
    elif style["background_color"] is not None:
        raise ValueError("background_color must be null or a six-digit RGB hex string without #")

    common = slide._element.cSld
    namespace = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    existing = common.find(namespace + "bg")
    if existing is not None:
        common.remove(existing)
    if background is not None:
        common.insert(0, background)
    return {"follow_master_background": background is None}


def _office_picture_resource(
    patch: dict[str, Any], resources: dict[tuple[str, str], bytes] | None,
) -> bytes:
    source = patch.get("source")
    if not isinstance(source, dict) or set(source) != {"path", "expected_sha256"}:
        raise ValueError("picture source requires path and expected_sha256 from read_file")
    try:
        data = (resources or {})[operation_resource_key(source)]
    except (KeyError, TypeError):
        raise ValueError("approved picture source snapshot is unavailable") from None
    if not isinstance(data, bytes) or not data:
        raise ValueError("approved picture source snapshot is empty")
    return data


def _iter_presentation_shapes(
    shapes: Any,
    parent_shape_ids: tuple[int, ...] = (),
):
    """Yield every slide shape, including nested group members, in source order."""
    for shape in shapes:
        yield shape, parent_shape_ids
        if shape._element.tag.rsplit("}", 1)[-1] == "grpSp":
            yield from _iter_presentation_shapes(shape.shapes, (*parent_shape_ids, shape.shape_id))


def _presentation_shape_target(slide: Any, shape_id: Any) -> tuple[Any, tuple[int, ...]]:
    """Resolve the slide-unique OOXML shape id across the complete group tree."""
    target_id = _integer(shape_id, "shape_id", 1, 2**31 - 1)
    matches = [item for item in _iter_presentation_shapes(slide.shapes) if item[0].shape_id == target_id]
    if not matches:
        raise ValueError("shape_id not found on the slide")
    if len(matches) != 1:
        raise ValueError("shape_id is ambiguous in this malformed PowerPoint slide")
    return matches[0]


def _presentation_shape_selector(
    slide_number: int,
    shape: Any,
    parent_shape_ids: tuple[int, ...],
    **extra: Any,
) -> dict[str, Any]:
    selector: dict[str, Any] = {"slide": slide_number, "shape_id": shape.shape_id}
    if parent_shape_ids:
        selector["group_path"] = [*parent_shape_ids, shape.shape_id]
    selector.update(extra)
    return selector


def _presentation_shape_tree(slide: Any, parent_shape_ids: tuple[int, ...]) -> Any:
    if not parent_shape_ids:
        return slide.shapes
    parent, _ = _presentation_shape_target(slide, parent_shape_ids[-1])
    if parent._element.tag.rsplit("}", 1)[-1] != "grpSp":
        raise ValueError("PowerPoint shape parent is not a group")
    return parent.shapes


def _flatten_presentation_group_members(group: Any) -> list[Any]:
    """Bake a non-rotated group's translation/scale into its members before lifting them."""
    from pptx.util import Emu

    members = list(group.shapes)
    if not members:
        return []
    transform = group._element.xfrm
    if transform.flipH or transform.flipV or float(transform.rot or 0) % 360:
        raise ValueError("shape.ungroup cannot flatten a rotated or flipped group without changing appearance")
    if transform.chExt.cx <= 0 or transform.chExt.cy <= 0:
        raise ValueError("shape.ungroup requires a group with positive child extents")
    scale_x = transform.ext.cx / transform.chExt.cx
    scale_y = transform.ext.cy / transform.chExt.cy
    for member in members:
        left = transform.off.x + (member.left - transform.chOff.x) * scale_x
        top = transform.off.y + (member.top - transform.chOff.y) * scale_y
        width = member.width * scale_x
        height = member.height * scale_y
        member.left = Emu(round(left))
        member.top = Emu(round(top))
        member.width = Emu(round(width))
        member.height = Emu(round(height))
    return members


class _WordPictureTarget:
    """One editable DrawingML picture in a Word package story."""

    def __init__(
        self,
        element: Any,
        *,
        part: Any,
        root: Any,
        stories: tuple[str, ...],
        section_indices: tuple[int, ...],
    ) -> None:
        self.element = element
        self.part = part
        self.root = root
        self.stories = stories
        self.section_indices = section_indices

    @property
    def layout(self) -> str:
        return "floating" if self.element.tag.endswith("}anchor") else "inline"

    @property
    def paragraph_element(self) -> Any:
        ancestor = self.element
        while ancestor is not None and not ancestor.tag.endswith("}p"):
            ancestor = ancestor.getparent()
        if ancestor is None:
            raise ValueError("Word picture is not contained in a paragraph")
        return ancestor

    @property
    def drawing_element(self) -> Any:
        drawing = self.element.getparent()
        if drawing is None or not drawing.tag.endswith("}drawing"):
            raise ValueError("Word picture is not contained in a drawing")
        return drawing

    @property
    def extent(self) -> Any:
        from docx.oxml.ns import qn

        extent = self.element.find(qn("wp:extent"))
        if extent is None:
            raise ValueError("Word picture has no editable extent")
        return extent

    @property
    def width(self) -> int:
        return int(self.extent.get("cx", "0"))

    @property
    def height(self) -> int:
        return int(self.extent.get("cy", "0"))

    @property
    def doc_properties(self) -> Any:
        from docx.oxml.ns import qn

        properties = self.element.find(qn("wp:docPr"))
        if properties is None:
            raise ValueError("Word picture has no non-visual properties")
        return properties

    @property
    def blip(self) -> Any:
        from docx.oxml.ns import qn

        blip = next(self.element.iter(qn("a:blip")), None)
        if blip is None or not blip.get(qn("r:embed")):
            raise ValueError("Word picture has no embedded image relationship")
        return blip

    @property
    def relationship_id(self) -> str:
        from docx.oxml.ns import qn

        return self.blip.get(qn("r:embed"))

    def set_relationship_id(self, relationship_id: str) -> None:
        from docx.oxml.ns import qn

        self.blip.set(qn("r:embed"), relationship_id)

    def set_size(self, width: int, height: int) -> None:
        from docx.oxml.ns import qn

        self.extent.set("cx", str(width))
        self.extent.set("cy", str(height))
        transform_extent = next(self.element.iter(qn("a:ext")), None)
        if transform_extent is not None:
            transform_extent.set("cx", str(width))
            transform_extent.set("cy", str(height))


_WORD_STORY_REFERENCES = (
    ("header", "headerReference", "default"),
    ("first_page_header", "headerReference", "first"),
    ("even_page_header", "headerReference", "even"),
    ("footer", "footerReference", "default"),
    ("first_page_footer", "footerReference", "first"),
    ("even_page_footer", "footerReference", "even"),
)


def _word_story_parts(document: Any) -> list[tuple[Any, Any, tuple[str, ...], tuple[int, ...]]]:
    """Return body plus explicitly defined/inherited header and footer parts."""
    from docx.oxml.ns import qn

    grouped: dict[Any, dict[str, Any]] = {}
    inherited: dict[str, Any] = {}
    for section_index, section in enumerate(document.sections):
        references = list(section._sectPr)
        for story, reference_name, reference_type in _WORD_STORY_REFERENCES:
            reference = next((
                item for item in references
                if item.tag == qn(f"w:{reference_name}") and item.get(qn("w:type"), "default") == reference_type
            ), None)
            if reference is not None:
                relationship_id = reference.get(qn("r:id"))
                relationship = document.part.rels.get(relationship_id)
                inherited[story] = relationship.target_part if relationship is not None else None
            part = inherited.get(story)
            if part is None:
                continue
            item = grouped.setdefault(part, {"stories": [], "sections": []})
            if story not in item["stories"]:
                item["stories"].append(story)
            if section_index not in item["sections"]:
                item["sections"].append(section_index)
    return [
        (document.part, document.element, ("body",), ()),
        *[
            (part, part.element, tuple(item["stories"]), tuple(item["sections"]))
            for part, item in grouped.items()
        ],
    ]


def _word_pictures(document: Any) -> list[_WordPictureTarget]:
    from docx.oxml.ns import qn

    container_tags = {qn("wp:inline"), qn("wp:anchor")}
    pictures = []
    for part, root, stories, section_indices in _word_story_parts(document):
        for element in root.iter():
            if element.tag not in container_tags:
                continue
            blip = next(element.iter(qn("a:blip")), None)
            if blip is not None and blip.get(qn("r:embed")):
                pictures.append(_WordPictureTarget(
                    element,
                    part=part,
                    root=root,
                    stories=stories,
                    section_indices=section_indices,
                ))
    return pictures


class _WordTextBoxTarget:
    """One native Word text box content container in a package story."""

    def __init__(
        self,
        content: Any,
        *,
        part: Any,
        root: Any,
        stories: tuple[str, ...],
        section_indices: tuple[int, ...],
    ) -> None:
        self.content = content
        self.part = part
        self.root = root
        self.stories = stories
        self.section_indices = section_indices

    @property
    def paragraphs(self) -> list[Any]:
        from docx.oxml.ns import qn
        from docx.text.paragraph import Paragraph

        return [Paragraph(element, self) for element in self.content.findall(qn("w:p"))]

    @property
    def tables(self) -> list[Any]:
        from docx.oxml.ns import qn
        from docx.table import Table

        return [Table(element, self) for element in self.content.findall(qn("w:tbl"))]

    @property
    def paragraph_element(self) -> Any:
        ancestor = self.content
        while ancestor is not None and not ancestor.tag.endswith("}p"):
            ancestor = ancestor.getparent()
        if ancestor is None:
            raise ValueError("Word text box is not contained in a paragraph")
        return ancestor

    @property
    def drawing_element(self) -> Any:
        ancestor = self.content
        while ancestor is not None and not ancestor.tag.endswith(("}drawing", "}pict")):
            ancestor = ancestor.getparent()
        if ancestor is None:
            raise ValueError("Word text box is not contained in DrawingML or VML")
        return ancestor

    @property
    def frame_element(self) -> Any:
        ancestor = self.content
        while ancestor is not None and not ancestor.tag.endswith(("}anchor", "}inline", "}shape", "}rect")):
            ancestor = ancestor.getparent()
        if ancestor is None:
            raise ValueError("Word text box has no editable frame")
        return ancestor

    @property
    def shape_element(self) -> Any:
        ancestor = self.content
        while ancestor is not None and not ancestor.tag.endswith(("}wsp", "}shape", "}rect")):
            ancestor = ancestor.getparent()
        if ancestor is None:
            raise ValueError("Word text box has no editable shape")
        return ancestor

    @property
    def group_elements(self) -> tuple[Any, ...]:
        """Return the containing Word/VML groups from nearest to outermost."""
        drawing = self.drawing_element
        ancestor = self.shape_element.getparent()
        groups = []
        while ancestor is not None and ancestor is not drawing:
            if ancestor.tag.rsplit("}", 1)[-1] in {"wgp", "grpSp", "group"}:
                groups.append(ancestor)
            ancestor = ancestor.getparent()
        return tuple(groups)

    @property
    def is_grouped(self) -> bool:
        from docx.oxml.ns import qn

        return bool(self.group_elements) or len(list(self.drawing_element.iter(qn("w:txbxContent")))) > 1

    @property
    def layout(self) -> str:
        return "inline" if self.frame_element.tag.endswith("}inline") else "floating"


def _word_text_boxes(document: Any) -> list[_WordTextBoxTarget]:
    from docx.oxml.ns import qn

    text_boxes = []
    for part, root, stories, section_indices in _word_story_parts(document):
        for content in root.iter(qn("w:txbxContent")):
            text_boxes.append(_WordTextBoxTarget(
                content,
                part=part,
                root=root,
                stories=stories,
                section_indices=section_indices,
            ))
    return text_boxes


def _word_group_has_visual_members(group: Any) -> bool:
    """Whether a WordprocessingShape or VML group still owns a drawable child."""
    namespace = group.tag.partition("}")[0].removeprefix("{")
    local_name = group.tag.rsplit("}", 1)[-1]
    if namespace == "urn:schemas-microsoft-com:vml" or local_name == "group":
        visual_names = {
            "arc", "curve", "group", "image", "line", "oval", "polyline", "rect",
            "roundrect", "shape",
        }
    else:
        visual_names = {"contentPart", "graphicFrame", "grpSp", "pic", "wsp"}
    return any(child.tag.rsplit("}", 1)[-1] in visual_names for child in group)


def _delete_word_text_box_shape(text_box: _WordTextBoxTarget) -> None:
    """Delete one native text-box member without removing grouped siblings."""
    from docx.oxml.ns import qn

    drawing = text_box.drawing_element
    shape = text_box.shape_element
    shape_contents = list(shape.iter(qn("w:txbxContent")))
    if len(shape_contents) != 1 or shape_contents[0] is not text_box.content:
        raise ValueError("Cannot safely identify one Word text-box shape member")

    groups = text_box.group_elements
    if not groups:
        if len(list(drawing.iter(qn("w:txbxContent")))) != 1:
            raise ValueError("Cannot delete one text box from a shared Word drawing without an explicit group")
        relationship_ids = _office_relationship_ids(drawing)
        run = drawing.getparent()
        run.remove(drawing)
        if run.tag.endswith("}r") and not any(
            child.tag.rsplit("}", 1)[-1] != "rPr" for child in run
        ):
            run.getparent().remove(run)
        _drop_unused_office_relationships(text_box.part, text_box.root, relationship_ids)
        return

    relationship_ids = _office_relationship_ids(shape)
    shape.getparent().remove(shape)
    for group in groups:
        parent = group.getparent()
        if parent is None or _word_group_has_visual_members(group):
            break
        parent.remove(group)

    outer_group = groups[-1]
    if outer_group.getparent() is None:
        relationship_ids.update(_office_relationship_ids(drawing))
        run = drawing.getparent()
        run.remove(drawing)
        if run.tag.endswith("}r") and not any(
            child.tag.rsplit("}", 1)[-1] != "rPr" for child in run
        ):
            run.getparent().remove(run)
    _drop_unused_office_relationships(text_box.part, text_box.root, relationship_ids)


def _word_content_target(document: Any, patch: dict[str, Any]) -> tuple[Any, dict[str, Any]]:
    if "text_box_index" in patch:
        if "story" in patch or "section_index" in patch:
            raise ValueError("text_box_index is package-wide and cannot be combined with story or section_index")
        text_boxes = _word_text_boxes(document)
        if not text_boxes:
            raise ValueError("Word document has no editable text boxes")
        index = _integer(patch.get("text_box_index"), "text_box_index", 0, len(text_boxes) - 1)
        return text_boxes[index], {"text_box_index": index}
    story_object, story, section_index = _word_story(document, patch.get("story"), patch.get("section_index"))
    selector: dict[str, Any] = {"story": story}
    if section_index is not None:
        selector["section_index"] = section_index
    return story_object, selector


def _word_target_root(target: Any) -> Any:
    return target.content if isinstance(target, _WordTextBoxTarget) else _word_story_element(target)


def _insert_word_target_paragraph(document: Any, target: Any, index: int, text: str, style: Any) -> None:
    if not isinstance(target, _WordTextBoxTarget):
        paragraphs = target.paragraphs
        if index == len(paragraphs):
            target.add_paragraph(text, style=style)
        else:
            paragraphs[index].insert_paragraph_before(text, style=style)
        return
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph

    if style is not None:
        try:
            document.styles[style]
        except (KeyError, TypeError):
            raise ValueError(f"Unknown Word paragraph style: {style}") from None
    paragraph_element = OxmlElement("w:p")
    paragraphs = target.paragraphs
    if index < len(paragraphs):
        paragraphs[index]._p.addprevious(paragraph_element)
    else:
        target.content.append(paragraph_element)
    paragraph = Paragraph(paragraph_element, target)
    if style is not None:
        paragraph.style = style
    _set_office_paragraph_text(paragraph, text, word=True)


def _word_text_box_transform(text_box: _WordTextBoxTarget, transform: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    allowed = {"x", "y", "width", "height", "rotation"}
    if not isinstance(transform, dict) or not transform or set(transform) - allowed:
        raise ValueError("Word text-box transform accepts x, y, width, height and rotation")
    values = {
        key: (
            _number(value, key, -360, 360) if key == "rotation"
            else _office_points(value, key, word=True, minimum=1 if key in {"width", "height"} else -100000)
        )
        for key, value in transform.items()
    }
    frame = text_box.frame_element
    shape = text_box.shape_element
    explicit_group = bool(text_box.group_elements)
    if explicit_group and frame.tag.endswith(("}anchor", "}inline")):
        shape_properties = next((node for node in shape if node.tag.endswith("}spPr")), None)
        shape_transform = (
            next((node for node in shape_properties if node.tag.endswith("}xfrm")), None)
            if shape_properties is not None else None
        )
        offset = (
            next((node for node in shape_transform if node.tag.endswith("}off")), None)
            if shape_transform is not None else None
        )
        extent = (
            next((node for node in shape_transform if node.tag.endswith("}ext")), None)
            if shape_transform is not None else None
        )
        if shape_transform is None or offset is None or extent is None:
            raise ValueError("Grouped Word text box has no editable member transform")
        for key, attribute in (("x", "x"), ("y", "y")):
            if key in values:
                offset.set(attribute, str(round(values[key] * 12700)))
        for key, attribute in (("width", "cx"), ("height", "cy")):
            if key in values:
                extent.set(attribute, str(round(values[key] * 12700)))
        if "rotation" in values:
            shape_transform.set("rot", str(round(values["rotation"] * 60000)))
        return
    if text_box.is_grouped and not explicit_group:
        raise ValueError("Word shared text-box geometry requires an explicit group-coordinate transform")
    if frame.tag.endswith(("}anchor", "}inline")):
        if frame.tag.endswith("}inline") and {"x", "y"} & values.keys():
            raise ValueError("inline Word text boxes do not have x/y positions")
        extent = frame.find(qn("wp:extent"))
        if extent is None:
            raise ValueError("Word text box has no editable extent")
        for key, attribute in (("width", "cx"), ("height", "cy")):
            if key in values:
                extent.set(attribute, str(round(values[key] * 12700)))
        if frame.tag.endswith("}anchor"):
            for key, axis in (("x", "H"), ("y", "V")):
                if key not in values:
                    continue
                position = frame.find(qn(f"wp:position{axis}"))
                if position is None:
                    position = OxmlElement(f"wp:position{axis}")
                    position.set("relativeFrom", "column" if axis == "H" else "paragraph")
                    first_extent = frame.find(qn("wp:extent"))
                    if first_extent is None:
                        frame.append(position)
                    else:
                        first_extent.addprevious(position)
                for child in list(position):
                    if child.tag in {qn("wp:align"), qn("wp:posOffset")}:
                        position.remove(child)
                offset = OxmlElement("wp:posOffset")
                offset.text = str(round(values[key] * 12700))
                position.append(offset)
        transform_element = next((node for node in shape.iter() if node.tag.endswith("}xfrm")), None)
        if transform_element is not None:
            transform_extent = next((node for node in transform_element if node.tag.endswith("}ext")), None)
            if transform_extent is not None:
                for key, attribute in (("width", "cx"), ("height", "cy")):
                    if key in values:
                        transform_extent.set(attribute, str(round(values[key] * 12700)))
            if "rotation" in values:
                transform_element.set("rot", str(round(values["rotation"] * 60000)))
        return
    style = {
        key.strip().lower(): value.strip()
        for item in frame.get("style", "").split(";") if ":" in item
        for key, value in [item.split(":", 1)]
    }
    position_names = {"x": "left", "y": "top"} if explicit_group else {
        "x": "margin-left", "y": "margin-top",
    }
    for key, property_name in (
        ("x", position_names["x"]), ("y", position_names["y"]),
        ("width", "width"), ("height", "height"),
    ):
        if key in values:
            style[property_name] = f"{values[key]:g}pt"
    if "rotation" in values:
        style["rotation"] = f"{values['rotation']:g}"
    frame.set("style", ";".join(f"{key}:{value}" for key, value in style.items()))


def _format_word_text_box(text_box: _WordTextBoxTarget, style: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    margins = {"margin_left", "margin_right", "margin_top", "margin_bottom"}
    allowed = margins | {
        "preset", "fill_color", "fill_opacity", "line_color", "line_opacity", "line_width",
        "vertical_alignment", "word_wrap", "alt_text", "wrap", "behind_text", "allow_overlap",
        "layout_in_cell", "distance_top", "distance_bottom", "distance_left", "distance_right",
    }
    if not isinstance(style, dict) or not style or set(style) - allowed:
        raise ValueError(f"Word text-box shape.format requires supported format keys: {', '.join(sorted(allowed))}")
    frame = text_box.frame_element
    shape = text_box.shape_element
    drawing = frame.tag.endswith(("}anchor", "}inline"))
    grouped = text_box.is_grouped
    shared_frame_keys = {
        "alt_text", "wrap", "behind_text", "allow_overlap", "layout_in_cell",
        "distance_top", "distance_bottom", "distance_left", "distance_right",
    }
    if grouped and shared_frame_keys & style.keys():
        raise ValueError("Grouped Word text boxes cannot change shared drawing-frame properties by member index")
    if "preset" in style:
        preset = style["preset"]
        if preset not in PresentationShapePreset.values() or preset == PresentationShapePreset.LINE.value:
            raise ValueError("Word text-box preset must be a supported non-line DrawingML preset")
        geometry = next((node for node in shape.iter() if node.tag.endswith("}prstGeom")), None)
        if geometry is None:
            raise ValueError("Word text box does not use editable preset geometry")
        geometry.set("prst", preset)
        for child in list(geometry):
            geometry.remove(child)
        geometry.append(OxmlElement("a:avLst"))

    def color(value: Any, name: str) -> str | None:
        if value is None:
            return None
        return _validated_cell_format({"fill_color": value}, spreadsheet=False)["fill_color"]

    if drawing:
        sp_pr = next((node for node in shape if node.tag.endswith("}spPr")), None)
        body_pr = next((node for node in shape if node.tag.endswith("}bodyPr")), None)
        if sp_pr is None or body_pr is None:
            raise ValueError("Word text box has no editable DrawingML shape/body properties")

        def replace_fill(owner: Any, value: Any, opacity: Any) -> None:
            fill_tags = {"noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill"}
            for child in list(owner):
                if child.tag.rsplit("}", 1)[-1] in fill_tags:
                    owner.remove(child)
            fill = OxmlElement("a:noFill")
            if value is not None:
                fill = OxmlElement("a:solidFill")
                rgb = OxmlElement("a:srgbClr")
                rgb.set("val", color(value, "color"))
                if opacity is not None:
                    alpha = OxmlElement("a:alpha")
                    alpha.set("val", str(round(_number(opacity, "opacity", 0, 1) * 100000)))
                    rgb.append(alpha)
                fill.append(rgb)
            if owner.tag.endswith("}spPr"):
                successor = next((
                    child for child in owner
                    if child.tag.rsplit("}", 1)[-1] in {
                        "ln", "effectLst", "effectDag", "scene3d", "sp3d", "extLst",
                    }
                ), None)
                if successor is not None:
                    successor.addprevious(fill)
                else:
                    owner.append(fill)
            else:
                owner.insert(0, fill)

        if "fill_color" in style:
            replace_fill(sp_pr, style["fill_color"], style.get("fill_opacity"))
        elif "fill_opacity" in style:
            raise ValueError("fill_opacity requires fill_color in the same Word text-box format operation")
        line = next((node for node in sp_pr if node.tag.endswith("}ln")), None)
        if {"line_color", "line_opacity", "line_width"} & style.keys():
            if line is None:
                line = OxmlElement("a:ln")
                sp_pr.append(line)
            if "line_color" in style:
                replace_fill(line, style["line_color"], style.get("line_opacity"))
            elif "line_opacity" in style:
                raise ValueError("line_opacity requires line_color in the same Word text-box format operation")
            if "line_width" in style:
                line.set("w", str(round(_office_points(style["line_width"], "line_width", word=True, maximum=72) * 12700)))
        margin_attributes = {
            "margin_left": "lIns", "margin_right": "rIns", "margin_top": "tIns", "margin_bottom": "bIns",
        }
        for key, attribute in margin_attributes.items():
            if key in style:
                body_pr.set(attribute, str(round(_office_points(style[key], key, word=True, maximum=10000) * 12700)))
        if "vertical_alignment" in style:
            alignment = style["vertical_alignment"]
            if alignment not in {"top", "middle", "bottom"}:
                raise ValueError("vertical_alignment must be top, middle or bottom")
            body_pr.set("anchor", {"top": "t", "middle": "ctr", "bottom": "b"}[alignment])
        if "word_wrap" in style:
            if type(style["word_wrap"]) is not bool:
                raise ValueError("word_wrap must be boolean")
            body_pr.set("wrap", "square" if style["word_wrap"] else "none")
        properties = frame.find(qn("wp:docPr"))
        if "alt_text" in style:
            alt_text = style["alt_text"]
            if properties is None or not isinstance(alt_text, str) or len(alt_text) > 1000 or "\x00" in alt_text:
                raise ValueError("alt_text must be a string of at most 1000 characters")
            properties.set("descr", alt_text)
        if frame.tag.endswith("}anchor"):
            if "wrap" in style:
                if style["wrap"] not in {"none", "square", "top_bottom"}:
                    raise ValueError("wrap must be none, square or top_bottom")
                _word_picture_set_wrap(frame, style["wrap"])
            for key, attribute in {
                "behind_text": "behindDoc", "allow_overlap": "allowOverlap", "layout_in_cell": "layoutInCell",
            }.items():
                if key in style:
                    if type(style[key]) is not bool:
                        raise ValueError(f"{key} must be boolean")
                    frame.set(attribute, "1" if style[key] else "0")
            for key, attribute in {
                "distance_top": "distT", "distance_bottom": "distB",
                "distance_left": "distL", "distance_right": "distR",
            }.items():
                if key in style:
                    points = _office_points(style[key], key, word=True, maximum=10000)
                    frame.set(attribute, str(round(points * 12700)))
        elif {"wrap", "behind_text", "allow_overlap", "layout_in_cell"} & style.keys():
            raise ValueError("floating layout properties require a floating Word text box")
        return
    if set(style) - {"fill_color", "line_color", "line_width", *margins, "alt_text"}:
        raise ValueError("this legacy VML text box supports fill/line/margins/alt text only")
    if "fill_color" in style:
        shape.set("filled", "f" if style["fill_color"] is None else "t")
        if style["fill_color"] is not None:
            shape.set("fillcolor", color(style["fill_color"], "fill_color"))
    if "line_color" in style:
        shape.set("stroked", "f" if style["line_color"] is None else "t")
        if style["line_color"] is not None:
            shape.set("strokecolor", color(style["line_color"], "line_color"))
    if "line_width" in style:
        shape.set("strokeweight", f"{_office_points(style['line_width'], 'line_width', word=True, maximum=72):g}pt")
    textbox = next((node for node in shape.iter() if node.tag.endswith("}textbox")), None)
    if textbox is not None and margins & style.keys():
        values = []
        for key in ("margin_left", "margin_top", "margin_right", "margin_bottom"):
            values.append(_office_points(style.get(key, 0), key, word=True, maximum=10000))
        textbox.set("inset", ",".join(f"{value:g}pt" for value in values))
    if "alt_text" in style:
        alt_text = style["alt_text"]
        if not isinstance(alt_text, str) or len(alt_text) > 1000 or "\x00" in alt_text:
            raise ValueError("alt_text must be a string of at most 1000 characters")
        shape.set("title", alt_text)


def _new_word_text_box_run(document: Any, text: str, transform: dict[str, Any], preset: str) -> Any:
    from lxml import etree
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    wps = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
    dimensions = {
        key: _office_points(transform[key], key, word=True, minimum=1 if key in {"width", "height"} else -100000)
        for key in ("x", "y", "width", "height")
    }
    rotation = _number(transform.get("rotation", 0), "rotation", -360, 360)
    doc_property_ids = [
        int(node.get("id"))
        for _part, root, _stories, _sections in _word_story_parts(document)
        for node in root.iter(qn("wp:docPr"))
        if str(node.get("id", "")).isdigit()
    ]
    doc_property_id = max(doc_property_ids, default=0) + 1
    run = OxmlElement("w:r")
    drawing = OxmlElement("w:drawing")
    anchor = OxmlElement("wp:anchor")
    for name, value in {
        "distT": "0", "distB": "0", "distL": "114300", "distR": "114300", "simplePos": "0",
        "relativeHeight": "251659264", "behindDoc": "0", "locked": "0", "layoutInCell": "1", "allowOverlap": "1",
    }.items():
        anchor.set(name, value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)
    for axis, relative, key in (("H", "column", "x"), ("V", "paragraph", "y")):
        position = OxmlElement(f"wp:position{axis}")
        position.set("relativeFrom", relative)
        offset = OxmlElement("wp:posOffset")
        offset.text = str(round(dimensions[key] * 12700))
        position.append(offset)
        anchor.append(position)
    extent = OxmlElement("wp:extent")
    extent.set("cx", str(round(dimensions["width"] * 12700)))
    extent.set("cy", str(round(dimensions["height"] * 12700)))
    anchor.append(extent)
    effect = OxmlElement("wp:effectExtent")
    for name in ("l", "t", "r", "b"):
        effect.set(name, "0")
    anchor.append(effect)
    wrap = OxmlElement("wp:wrapSquare")
    wrap.set("wrapText", "bothSides")
    anchor.append(wrap)
    properties = OxmlElement("wp:docPr")
    properties.set("id", str(doc_property_id))
    properties.set("name", f"Manor text box {doc_property_id}")
    anchor.append(properties)
    anchor.append(OxmlElement("wp:cNvGraphicFramePr"))
    graphic = OxmlElement("a:graphic")
    graphic_data = OxmlElement("a:graphicData")
    graphic_data.set("uri", wps)
    shape = etree.SubElement(graphic_data, f"{{{wps}}}wsp", nsmap={"wps": wps})
    nonvisual = etree.SubElement(shape, f"{{{wps}}}cNvSpPr")
    nonvisual.set("txBox", "1")
    shape_properties = etree.SubElement(shape, f"{{{wps}}}spPr")
    xfrm = OxmlElement("a:xfrm")
    if rotation:
        xfrm.set("rot", str(round(rotation * 60000)))
    offset = OxmlElement("a:off")
    offset.set("x", "0")
    offset.set("y", "0")
    shape_extent = OxmlElement("a:ext")
    shape_extent.set("cx", str(round(dimensions["width"] * 12700)))
    shape_extent.set("cy", str(round(dimensions["height"] * 12700)))
    xfrm.extend([offset, shape_extent])
    shape_properties.append(xfrm)
    geometry = OxmlElement("a:prstGeom")
    geometry.set("prst", preset)
    geometry.append(OxmlElement("a:avLst"))
    shape_properties.append(geometry)
    text_box = etree.SubElement(shape, f"{{{wps}}}txbx")
    content = OxmlElement("w:txbxContent")
    content.append(OxmlElement("w:p"))
    text_box.append(content)
    body = etree.SubElement(shape, f"{{{wps}}}bodyPr")
    body.set("wrap", "square")
    for attribute, value in {"lIns": 7.2, "rIns": 7.2, "tIns": 3.6, "bIns": 3.6}.items():
        body.set(attribute, str(round(value * 12700)))
    body.set("anchor", "t")
    graphic.append(graphic_data)
    anchor.append(graphic)
    drawing.append(anchor)
    run.append(drawing)
    target = _WordTextBoxTarget(content, part=document.part, root=document.element, stories=("body",), section_indices=())
    _set_office_paragraph_text(target.paragraphs[0], text, word=True)
    return run




def _word_picture_paragraph_element(picture: _WordPictureTarget) -> Any:
    return picture.paragraph_element


def _word_story(document: Any, story: Any, section_index: Any) -> tuple[Any, str, int | None]:
    story = "body" if story is None else story
    allowed = {"body", *[item[0] for item in _WORD_STORY_REFERENCES]}
    if story not in allowed:
        raise ValueError(f"Word picture story must be one of: {', '.join(sorted(allowed))}")
    if story == "body":
        if section_index is not None:
            raise ValueError("section_index is supported only for Word header/footer pictures")
        return document, story, None
    section_index = _integer(0 if section_index is None else section_index, "section_index", 0, len(document.sections) - 1)
    section = document.sections[section_index]
    story_object = getattr(section, story)
    if section_index and story_object.is_linked_to_previous:
        raise ValueError(
            f"Section {section_index} {story} is linked to the previous section; target the discovered shared picture "
            "or create an independent header/footer definition first"
        )
    return story_object, story, section_index


def _word_story_element(story_object: Any) -> Any:
    return getattr(story_object, "element", story_object._element)


def _office_relationship_ids(element: Any) -> set[str]:
    namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    return {
        value
        for node in element.iter()
        for name, value in node.attrib.items()
        if name.startswith(namespace) and isinstance(value, str) and value
    }


def _drop_unused_office_relationships(part: Any, root: Any, relationship_ids: set[str]) -> None:
    referenced = _office_relationship_ids(root)
    for relationship_id in relationship_ids - referenced:
        if relationship_id in part.rels:
            part.drop_rel(relationship_id)


def _insert_presentation_paragraph(shape: Any, index: int, text: Any, style: Any = None) -> Any:
    """Insert a paragraph while inheriting the nearest template paragraph style."""
    frame = shape.text_frame
    existing = list(frame.paragraphs)
    source = existing[index - 1] if index else (existing[0] if existing else None)
    paragraph = frame.add_paragraph()
    if index < len(existing):
        existing[index]._p.addprevious(paragraph._p)
    if source is not None:
        namespace = source._p.tag.rsplit("}", 1)[0] + "}"
        for child in list(paragraph._p):
            if child.tag in {namespace + "pPr", namespace + "endParaRPr"}:
                paragraph._p.remove(child)
        source_properties = source._p.find(namespace + "pPr")
        if source_properties is not None:
            paragraph._p.insert(0, copy.deepcopy(source_properties))
        source_typing_style = source._p.find(namespace + "endParaRPr")
        if source_typing_style is not None:
            paragraph._p.append(copy.deepcopy(source_typing_style))
        else:
            source_run_style = _first_office_run_properties(source._p)
            if source_run_style is not None:
                source_run_style = copy.deepcopy(source_run_style)
                source_run_style.tag = namespace + "endParaRPr"
                paragraph._p.append(source_run_style)
    _set_office_paragraph_text(paragraph, text, word=False)
    if style is not None:
        _format_office_paragraph(paragraph, {"format": style}, word=False)
    return paragraph


def _rewrite_office_relationship_ids(root: Any, replacements: dict[str, str]) -> None:
    namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    for node in root.iter():
        for name, value in list(node.attrib.items()):
            if name.startswith(namespace) and value in replacements:
                node.set(name, replacements[value])


def _next_cloned_partname(package: Any, partname: Any) -> Any:
    source = str(partname)
    match = re.fullmatch(r"(.*?)([0-9]+)(\.[^/]+)", source)
    if match:
        template = f"{match.group(1)}%d{match.group(3)}"
    else:
        stem, separator, suffix = source.rpartition(".")
        template = f"{stem}%d{separator}{suffix}" if separator else f"{source}%d"
    return package.next_partname(template)


def _clone_presentation_part(source_part: Any, package: Any, memo: dict[Any, Any]) -> Any:
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    if source_part in memo:
        return memo[source_part]
    target_part = source_part.__class__.load(
        _next_cloned_partname(package, source_part.partname),
        source_part.content_type,
        package,
        source_part.blob,
    )
    memo[source_part] = target_part
    shared_relationships = {
        RT.AUDIO, RT.IMAGE, RT.MEDIA, RT.NOTES_MASTER, RT.SLIDE_LAYOUT,
        RT.SLIDE_MASTER, RT.TABLE_STYLES, RT.THEME, RT.THEME_OVERRIDE, RT.VIDEO,
    }
    replacements = {}
    for relationship in source_part.rels.values():
        if relationship.is_external:
            target = relationship.target_ref
        elif relationship.target_part in memo:
            target = memo[relationship.target_part]
        elif relationship.reltype in shared_relationships:
            target = relationship.target_part
        else:
            target = _clone_presentation_part(relationship.target_part, package, memo)
        replacements[relationship.rId] = target_part.relate_to(
            target,
            relationship.reltype,
            is_external=relationship.is_external,
        )
    if hasattr(target_part, "_element"):
        _rewrite_office_relationship_ids(target_part._element, replacements)
    return target_part


def _duplicate_presentation_slide(document: Any, source_number: int, index: int) -> dict[str, Any]:
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    source = document.slides[source_number - 1]
    target = document.slides.add_slide(source.slide_layout)
    cloned_element = copy.deepcopy(source._element)
    target._element = cloned_element
    target.part._element = cloned_element
    memo = {source.part: target.part}
    replacements = {}
    shared_relationships = {
        RT.AUDIO, RT.IMAGE, RT.MEDIA, RT.SLIDE_LAYOUT, RT.VIDEO,
    }
    for relationship in source.part.rels.values():
        if relationship.is_external:
            related = relationship.target_ref
        elif relationship.reltype in shared_relationships:
            related = relationship.target_part
        elif relationship.target_part in memo:
            related = memo[relationship.target_part]
        else:
            related = _clone_presentation_part(relationship.target_part, document.part.package, memo)
        replacements[relationship.rId] = target.part.relate_to(
            related,
            relationship.reltype,
            is_external=relationship.is_external,
        )
    _rewrite_office_relationship_ids(cloned_element, replacements)
    slide_ids = document.slides._sldIdLst
    added = slide_ids[-1]
    slide_ids.remove(added)
    slide_ids.insert(index, added)
    return {"slide": index + 1, "index": index, "source_slide": source_number}


def _word_picture_position(anchor: Any, axis: str) -> Any:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    position = anchor.find(qn(f"wp:position{axis}"))
    if position is not None:
        return position
    position = OxmlElement(f"wp:position{axis}")
    position.set("relativeFrom", "column" if axis == "H" else "paragraph")
    offset = OxmlElement("wp:posOffset")
    offset.text = "0"
    position.append(offset)
    extent = anchor.find(qn("wp:extent"))
    if extent is None:
        raise ValueError("Word floating picture has no extent")
    extent.addprevious(position)
    return position


def _word_picture_set_position(position: Any, value: float) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for child in list(position):
        if child.tag in {qn("wp:align"), qn("wp:posOffset")}:
            position.remove(child)
    offset = OxmlElement("wp:posOffset")
    offset.text = str(round(value * 12700))
    position.append(offset)


def _word_picture_as_anchor(picture: _WordPictureTarget) -> None:
    if picture.layout == "floating":
        return
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    inline = picture.element
    anchor = OxmlElement("wp:anchor")
    for name, value in {
        "distT": "0", "distB": "0", "distL": "114300", "distR": "114300",
        "simplePos": "0", "relativeHeight": "251659264", "behindDoc": "0",
        "locked": "0", "layoutInCell": "1", "allowOverlap": "1",
    }.items():
        anchor.set(name, value)
    simple_position = OxmlElement("wp:simplePos")
    simple_position.set("x", "0")
    simple_position.set("y", "0")
    anchor.append(simple_position)
    for axis, relative_from in (("H", "column"), ("V", "paragraph")):
        position = OxmlElement(f"wp:position{axis}")
        position.set("relativeFrom", relative_from)
        offset = OxmlElement("wp:posOffset")
        offset.text = "0"
        position.append(offset)
        anchor.append(position)
    for tag in ("wp:extent", "wp:effectExtent"):
        child = inline.find(qn(tag))
        if child is not None:
            anchor.append(child)
    if anchor.find(qn("wp:effectExtent")) is None:
        effect_extent = OxmlElement("wp:effectExtent")
        for name in ("l", "t", "r", "b"):
            effect_extent.set(name, "0")
        anchor.append(effect_extent)
    wrap = OxmlElement("wp:wrapSquare")
    wrap.set("wrapText", "bothSides")
    anchor.append(wrap)
    for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        child = inline.find(qn(tag))
        if child is not None:
            anchor.append(child)
    parent = inline.getparent()
    if parent is None:
        raise ValueError("Word inline picture is detached")
    parent.replace(inline, anchor)
    picture.element = anchor


def _word_picture_as_inline(picture: _WordPictureTarget) -> None:
    if picture.layout == "inline":
        return
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    anchor = picture.element
    inline = OxmlElement("wp:inline")
    for name in ("distT", "distB", "distL", "distR"):
        inline.set(name, anchor.get(name, "0"))
    for tag in ("wp:extent", "wp:effectExtent", "wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        child = anchor.find(qn(tag))
        if child is not None:
            inline.append(child)
    parent = anchor.getparent()
    if parent is None:
        raise ValueError("Word floating picture is detached")
    parent.replace(anchor, inline)
    picture.element = inline


def _word_picture_wrap(anchor: Any) -> str | None:
    names = {
        "wrapNone": "none",
        "wrapSquare": "square",
        "wrapTopAndBottom": "top_bottom",
        "wrapTight": "tight",
        "wrapThrough": "through",
    }
    return next((names.get(child.tag.rsplit("}", 1)[-1]) for child in anchor if child.tag.rsplit("}", 1)[-1] in names), None)


def _word_picture_set_wrap(anchor: Any, wrap: str) -> None:
    from docx.oxml import OxmlElement

    wrap_names = {
        "none": "wp:wrapNone",
        "square": "wp:wrapSquare",
        "top_bottom": "wp:wrapTopAndBottom",
    }
    existing = next((child for child in anchor if child.tag.rsplit("}", 1)[-1].startswith("wrap")), None)
    replacement = OxmlElement(wrap_names[wrap])
    if wrap == "square":
        replacement.set("wrapText", "bothSides")
    if existing is not None:
        existing.addnext(replacement)
        anchor.remove(existing)
        return
    properties = next((child for child in anchor if child.tag.rsplit("}", 1)[-1] == "docPr"), None)
    if properties is None:
        anchor.append(replacement)
    else:
        properties.addprevious(replacement)


def _format_word_picture(picture: _WordPictureTarget, style: Any) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    floating_keys = {
        "position_x", "position_y", "relative_from_horizontal", "relative_from_vertical", "wrap",
        "behind_text", "allow_overlap", "layout_in_cell", "distance_top", "distance_bottom",
        "distance_left", "distance_right",
    }
    allowed = {"width", "height", "alt_text", "alignment", "layout"} | floating_keys
    if not isinstance(style, dict) or not style or set(style) - allowed:
        raise ValueError(f"Word picture.format requires supported format keys: {', '.join(sorted(allowed))}")
    layout = style.get("layout", picture.layout)
    if layout not in {"inline", "floating"}:
        raise ValueError("Word picture layout must be inline or floating")
    if set(style) & floating_keys and layout != "floating":
        raise ValueError("Word floating picture properties require layout floating")
    if "alignment" in style and layout != "inline":
        raise ValueError("Word picture alignment is supported only for inline pictures")
    horizontal_values = {
        "character", "column", "insideMargin", "leftMargin", "margin", "outsideMargin", "page", "rightMargin",
    }
    vertical_values = {
        "bottomMargin", "insideMargin", "line", "margin", "outsideMargin", "page", "paragraph", "topMargin",
    }
    if "relative_from_horizontal" in style and style["relative_from_horizontal"] not in horizontal_values:
        raise ValueError("relative_from_horizontal is not a supported Word anchor reference")
    if "relative_from_vertical" in style and style["relative_from_vertical"] not in vertical_values:
        raise ValueError("relative_from_vertical is not a supported Word anchor reference")
    if "wrap" in style and style["wrap"] not in {"none", "square", "top_bottom"}:
        raise ValueError("Word picture wrap must be none, square or top_bottom")
    for key in ("behind_text", "allow_overlap", "layout_in_cell"):
        if key in style and type(style[key]) is not bool:
            raise ValueError(f"{key} must be boolean")
    positions = {
        key: _office_points(style[key], key, word=True, minimum=-100000, maximum=100000)
        for key in ("position_x", "position_y") if key in style
    }
    distances = {
        key: _office_points(style[key], key, word=True, maximum=10000)
        for key in ("distance_top", "distance_bottom", "distance_left", "distance_right") if key in style
    }
    original_width, original_height = int(picture.width), int(picture.height)
    dimensions = {
        key: _office_points(style[key], key, word=True, minimum=1, maximum=10000)
        for key in ("width", "height") if key in style
    }
    if dimensions and (original_width <= 0 or original_height <= 0):
        raise ValueError("Word picture dimensions must be positive before proportional resizing")
    if "width" in dimensions and "height" not in dimensions:
        dimensions["height"] = dimensions["width"] * original_height / original_width
    elif "height" in dimensions and "width" not in dimensions:
        dimensions["width"] = dimensions["height"] * original_width / original_height
    if dimensions:
        picture.set_size(
            round(dimensions.get("width", original_width / 12700) * 12700),
            round(dimensions.get("height", original_height / 12700) * 12700),
        )
    if "alt_text" in style:
        alt_text = style["alt_text"]
        if not isinstance(alt_text, str) or len(alt_text) > 1000 or "\x00" in alt_text:
            raise ValueError("alt_text must be a string of at most 1000 characters")
        picture.doc_properties.set("descr", alt_text)
    if layout == "floating":
        _word_picture_as_anchor(picture)
    else:
        _word_picture_as_inline(picture)
    if picture.layout == "floating":
        horizontal = _word_picture_position(picture.element, "H")
        vertical = _word_picture_position(picture.element, "V")
        if "relative_from_horizontal" in style:
            horizontal.set("relativeFrom", style["relative_from_horizontal"])
        if "relative_from_vertical" in style:
            vertical.set("relativeFrom", style["relative_from_vertical"])
        if "position_x" in positions:
            _word_picture_set_position(horizontal, positions["position_x"])
        if "position_y" in positions:
            _word_picture_set_position(vertical, positions["position_y"])
        if "wrap" in style:
            _word_picture_set_wrap(picture.element, style["wrap"])
        for key, attribute in {
            "behind_text": "behindDoc", "allow_overlap": "allowOverlap", "layout_in_cell": "layoutInCell",
        }.items():
            if key in style:
                picture.element.set(attribute, "1" if style[key] else "0")
        for key, attribute in {
            "distance_top": "distT", "distance_bottom": "distB", "distance_left": "distL", "distance_right": "distR",
        }.items():
            if key in distances:
                picture.element.set(attribute, str(round(distances[key] * 12700)))
    if "alignment" in style:
        alignment = style["alignment"]
        if alignment not in {"left", "center", "right"}:
            raise ValueError("Word picture alignment must be left, center or right")
        paragraph_properties = _word_picture_paragraph_element(picture).get_or_add_pPr()
        paragraph_properties.get_or_add_jc().val = getattr(WD_ALIGN_PARAGRAPH, alignment.upper())


def _format_presentation_picture(picture: Any, style: Any) -> None:
    """Change requested native picture properties without replacing its media."""
    from pptx.oxml.xmlchemy import OxmlElement

    allowed = {"crop", "opacity", "alt_text"}
    if not isinstance(style, dict) or not style or set(style) - allowed:
        raise ValueError(f"picture.format requires supported format keys: {', '.join(sorted(allowed))}")
    if picture._element.tag.rsplit("}", 1)[-1] != "pic":
        raise ValueError("picture.format requires a top-level picture")
    if "crop" in style:
        crop = style["crop"]
        keys = {"left", "top", "right", "bottom"}
        if not isinstance(crop, dict) or not crop or set(crop) - keys:
            raise ValueError("picture crop requires left, top, right or bottom fractions")
        values = {
            key: _number(crop.get(key, getattr(picture, f"crop_{key}")), f"crop.{key}", 0, 0.95)
            for key in keys
        }
        if values["left"] + values["right"] >= 1 or values["top"] + values["bottom"] >= 1:
            raise ValueError("opposite picture crop fractions must total less than 1")
        for key, value in values.items():
            setattr(picture, f"crop_{key}", value)
    if "opacity" in style:
        opacity = _number(style["opacity"], "opacity", 0, 1)
        blip = picture._element.blipFill.blip
        namespace = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        for child in list(blip):
            if child.tag in {namespace + "alphaModFix", namespace + "alpha", namespace + "alphaMod", namespace + "alphaOff"}:
                blip.remove(child)
        if opacity < 1:
            alpha = OxmlElement("a:alphaModFix")
            alpha.set("amt", str(round(opacity * 100000)))
            blip.append(alpha)
    if "alt_text" in style:
        alt_text = style["alt_text"]
        if not isinstance(alt_text, str) or len(alt_text) > 1000 or "\x00" in alt_text:
            raise ValueError("alt_text must be a string of at most 1000 characters")
        picture._element.nvPicPr.cNvPr.set("descr", alt_text)


def _picture_fit_geometry(data: bytes, geometry: dict[str, float], fit: Any) -> tuple[dict[str, float], dict[str, float]]:
    from PIL import Image

    if fit not in {None, "stretch", "cover", "contain"}:
        raise ValueError("picture fit must be stretch, cover or contain")
    if fit in {None, "stretch"}:
        return geometry, {}
    with Image.open(io.BytesIO(data)) as image:
        source_ratio = image.width / image.height
    frame_ratio = geometry["width"] / geometry["height"]
    if fit == "cover":
        if source_ratio > frame_ratio:
            crop = (1 - frame_ratio / source_ratio) / 2
            return geometry, {"left": crop, "right": crop}
        crop = (1 - source_ratio / frame_ratio) / 2
        return geometry, {"top": crop, "bottom": crop}
    fitted = dict(geometry)
    if source_ratio > frame_ratio:
        height = geometry["width"] / source_ratio
        fitted["y"] += (geometry["height"] - height) / 2
        fitted["height"] = height
    else:
        width = geometry["height"] * source_ratio
        fitted["x"] += (geometry["width"] - width) / 2
        fitted["width"] = width
    return fitted, {}


def _presentation_chart_types() -> dict[str, Any]:
    from pptx.enum.chart import XL_CHART_TYPE

    return {
        OfficeChartType.COLUMN.value: XL_CHART_TYPE.COLUMN_CLUSTERED,
        OfficeChartType.COLUMN_STACKED.value: XL_CHART_TYPE.COLUMN_STACKED,
        OfficeChartType.COLUMN_STACKED_100.value: XL_CHART_TYPE.COLUMN_STACKED_100,
        OfficeChartType.BAR.value: XL_CHART_TYPE.BAR_CLUSTERED,
        OfficeChartType.BAR_STACKED.value: XL_CHART_TYPE.BAR_STACKED,
        OfficeChartType.BAR_STACKED_100.value: XL_CHART_TYPE.BAR_STACKED_100,
        OfficeChartType.LINE.value: XL_CHART_TYPE.LINE,
        OfficeChartType.LINE_MARKERS.value: XL_CHART_TYPE.LINE_MARKERS,
        OfficeChartType.AREA.value: XL_CHART_TYPE.AREA,
        OfficeChartType.AREA_STACKED.value: XL_CHART_TYPE.AREA_STACKED,
        OfficeChartType.AREA_STACKED_100.value: XL_CHART_TYPE.AREA_STACKED_100,
        OfficeChartType.PIE.value: XL_CHART_TYPE.PIE,
        OfficeChartType.DOUGHNUT.value: XL_CHART_TYPE.DOUGHNUT,
        OfficeChartType.SCATTER.value: XL_CHART_TYPE.XY_SCATTER,
        OfficeChartType.SCATTER_LINES.value: XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS,
        OfficeChartType.SCATTER_LINES_MARKERS.value: XL_CHART_TYPE.XY_SCATTER_LINES,
        OfficeChartType.SCATTER_SMOOTH.value: XL_CHART_TYPE.XY_SCATTER_SMOOTH_NO_MARKERS,
        OfficeChartType.SCATTER_SMOOTH_MARKERS.value: XL_CHART_TYPE.XY_SCATTER_SMOOTH,
    }


def _presentation_chart_type_name(chart_type: Any) -> str | None:
    return next((name for name, value in _presentation_chart_types().items() if value == chart_type), None)


def _is_combo_chart_type(chart_type: Any) -> bool:
    return chart_type == OfficeChartType.COMBO_COLUMN_LINE.value


def _is_stock_chart_type(chart_type: Any) -> bool:
    return chart_type in {
        OfficeChartType.STOCK_HLC.value,
        OfficeChartType.STOCK_OHLC.value,
        OfficeChartType.STOCK_VHLC.value,
        OfficeChartType.STOCK_VOHLC.value,
    }


def _is_volume_stock_chart_type(chart_type: Any) -> bool:
    return chart_type in {OfficeChartType.STOCK_VHLC.value, OfficeChartType.STOCK_VOHLC.value}


def _stock_chart_series_count(chart_type: Any) -> int | None:
    return {
        OfficeChartType.STOCK_HLC.value: 3,
        OfficeChartType.STOCK_OHLC.value: 4,
        OfficeChartType.STOCK_VHLC.value: 4,
        OfficeChartType.STOCK_VOHLC.value: 5,
    }.get(chart_type)


def _stock_price_series_count(chart_type: Any) -> int | None:
    return {
        OfficeChartType.STOCK_HLC.value: 3,
        OfficeChartType.STOCK_OHLC.value: 4,
        OfficeChartType.STOCK_VHLC.value: 3,
        OfficeChartType.STOCK_VOHLC.value: 4,
    }.get(chart_type)


def _presentation_category_value_axis_nodes(chart: Any) -> tuple[Any | None, Any | None, Any | None]:
    """Resolve category, primary-value, and secondary-value axes from plot axId references."""
    from pptx.oxml.ns import qn

    plot_area = chart._chartSpace.chart.plotArea
    category_axes = plot_area.findall(qn("c:catAx")) + plot_area.findall(qn("c:dateAx"))
    value_axes = plot_area.findall(qn("c:valAx"))

    def axis_id(node: Any) -> str | None:
        identifier = node.find(qn("c:axId"))
        return identifier.get("val") if identifier is not None else None

    def referenced_axis(plot: Any, axes_by_id: dict[str, Any]) -> Any | None:
        if plot is None:
            return None
        return next(
            (
                axes_by_id.get(identifier.get("val"))
                for identifier in plot.findall(qn("c:axId"))
                if identifier.get("val") in axes_by_id
            ),
            None,
        )

    categories_by_id = {
        identifier: node for node in category_axes if (identifier := axis_id(node)) is not None
    }
    values_by_id = {
        identifier: node for node in value_axes if (identifier := axis_id(node)) is not None
    }
    plots = list(plot_area.iter_xCharts())
    bar_plot = plot_area.find(qn("c:barChart"))
    primary_plot = bar_plot if bar_plot is not None else (plots[0] if plots else None)
    category_axis = referenced_axis(primary_plot, categories_by_id)
    if category_axis is None and category_axes:
        category_axis = category_axes[0]
    primary_value_axis = referenced_axis(primary_plot, values_by_id)
    if primary_value_axis is None and value_axes:
        primary_value_axis = value_axes[0]
    primary_value_id = axis_id(primary_value_axis) if primary_value_axis is not None else None
    secondary_value_axis = next(
        (
            candidate
            for plot in plots
            if plot is not primary_plot
            and (candidate := referenced_axis(plot, values_by_id)) is not None
            and axis_id(candidate) != primary_value_id
        ),
        None,
    )
    return category_axis, primary_value_axis, secondary_value_axis


def _presentation_chart_type_name_for_chart(chart: Any) -> str | None:
    from pptx.oxml.ns import qn

    plot_area = chart._chartSpace.chart.plotArea
    plot_tags = [node.tag for node in plot_area.iter_xCharts()]
    if len(plot_tags) == 2 and set(plot_tags) == {qn("c:barChart"), qn("c:stockChart")}:
        bar_chart = plot_area.find(qn("c:barChart"))
        stock_chart = plot_area.find(qn("c:stockChart"))
        bar_direction = bar_chart.find(qn("c:barDir")) if bar_chart is not None else None
        grouping = bar_chart.find(qn("c:grouping")) if bar_chart is not None else None
        bar_series_count = len(bar_chart.findall(qn("c:ser"))) if bar_chart is not None else 0
        stock_series_count = len(stock_chart.findall(qn("c:ser"))) if stock_chart is not None else 0
        if (
            bar_direction is not None and bar_direction.get("val") == "col"
            and grouping is not None and grouping.get("val") in {"clustered", "standard"}
            and bar_series_count == 1
        ):
            return {
                3: OfficeChartType.STOCK_VHLC.value,
                4: OfficeChartType.STOCK_VOHLC.value,
            }.get(stock_series_count)
        return None
    if plot_tags == [qn("c:stockChart")]:
        stock_chart = plot_area.find(qn("c:stockChart"))
        series_count = len(stock_chart.findall(qn("c:ser"))) if stock_chart is not None else 0
        if series_count == 3:
            return OfficeChartType.STOCK_HLC.value
        if series_count == 4:
            return OfficeChartType.STOCK_OHLC.value
        return None
    if len(plot_tags) == 2 and set(plot_tags) == {qn("c:barChart"), qn("c:lineChart")}:
        bar_chart = plot_area.find(qn("c:barChart"))
        bar_direction = bar_chart.find(qn("c:barDir")) if bar_chart is not None else None
        grouping = bar_chart.find(qn("c:grouping")) if bar_chart is not None else None
        if (
            bar_direction is not None and bar_direction.get("val") == "col"
            and grouping is not None and grouping.get("val") in {"clustered", "standard"}
        ):
            return OfficeChartType.COMBO_COLUMN_LINE.value
    if len(plot_tags) != 1:
        return None
    try:
        return _presentation_chart_type_name(chart.chart_type)
    except ValueError:
        return None


def _make_presentation_stock_chart(chart: Any, chart_type: str) -> None:
    """Build an editable stock plot, including a dual-axis volume variant."""
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement

    expected_series = _stock_chart_series_count(chart_type)
    if expected_series is None:
        raise ValueError("unsupported stock chart type")
    plot_area = chart._chartSpace.chart.plotArea
    volume_stock = _is_volume_stock_chart_type(chart_type)
    bar_plot = plot_area.find(qn("c:barChart")) if volume_stock else None
    plot = plot_area.find(qn("c:lineChart"))
    if plot is None:
        plot = plot_area.find(qn("c:stockChart"))
    if volume_stock:
        if bar_plot is None or len(list(plot_area.iter_xCharts())) not in {1, 2}:
            raise ValueError("volume stock charts require one column plot and one stock plot")
        series = [
            item
            for owner in (bar_plot, plot)
            if owner is not None
            for item in owner.findall(qn("c:ser"))
        ]
        if len(series) != expected_series:
            raise ValueError(f"{chart_type} charts require exactly {expected_series} series")

        def setting(parent: Any, tag: str, value: str) -> Any:
            element = parent.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                first_series = parent.find(qn("c:ser"))
                if first_series is None:
                    parent.append(element)
                else:
                    first_series.addprevious(element)
            element.set("val", value)
            return element

        def append_series(parent: Any, item: Any, boundary_names: set[str]) -> None:
            current_parent = item.getparent()
            if current_parent is not None:
                current_parent.remove(item)
            boundary = next(
                (child for child in parent if child.tag.rsplit("}", 1)[-1] in boundary_names),
                None,
            )
            if boundary is None:
                parent.append(item)
            else:
                boundary.addprevious(item)

        setting(bar_plot, "c:barDir", "col")
        setting(bar_plot, "c:grouping", "clustered")
        if plot is None:
            plot = OxmlElement("c:stockChart")
            bar_plot.addnext(plot)
        for item in series[:1]:
            append_series(bar_plot, item, {"dLbls", "gapWidth", "overlap", "serLines", "axId", "extLst"})
        for item in series[1:]:
            append_series(
                plot,
                item,
                {"dLbls", "dropLines", "hiLowLines", "upDownBars", "marker", "smooth", "axId", "extLst"},
            )

        category_axis, primary_value_axis, secondary_axis = _presentation_category_value_axis_nodes(chart)
        if category_axis is None or primary_value_axis is None:
            raise ValueError("volume stock charts require native category and value axes")
        category_id = category_axis.find(qn("c:axId")).get("val")
        primary_value_id = primary_value_axis.find(qn("c:axId")).get("val")
        if secondary_axis is None:
            used_ids = {
                int(item.get("val"))
                for item in plot_area.iter()
                if item.tag in {qn("c:axId"), qn("c:crossAx")}
                and str(item.get("val", "")).isdigit()
            }
            secondary_id = max(used_ids or {100}) + 1
            secondary_axis = copy.deepcopy(primary_value_axis)
            secondary_axis.find(qn("c:axId")).set("val", str(secondary_id))
            axis_position = secondary_axis.find(qn("c:axPos"))
            if axis_position is None:
                axis_position = OxmlElement("c:axPos")
                secondary_axis.insert(1, axis_position)
            axis_position.set("val", "r")
            cross_axis = secondary_axis.find(qn("c:crossAx"))
            if cross_axis is None:
                cross_axis = OxmlElement("c:crossAx")
                secondary_axis.append(cross_axis)
            cross_axis.set("val", category_id)
            crosses_at = secondary_axis.find(qn("c:crossesAt"))
            if crosses_at is not None:
                secondary_axis.remove(crosses_at)
            crosses = secondary_axis.find(qn("c:crosses"))
            if crosses is None:
                crosses = OxmlElement("c:crosses")
                cross_axis.addnext(crosses)
            crosses.set("val", "max")
            boundary = next(
                (child for child in plot_area if child.tag in {qn("c:dTable"), qn("c:spPr"), qn("c:extLst")}),
                None,
            )
            if boundary is None:
                plot_area.append(secondary_axis)
            else:
                boundary.addprevious(secondary_axis)
        else:
            secondary_id = int(secondary_axis.find(qn("c:axId")).get("val"))

        for axis_id in plot.findall(qn("c:axId")):
            plot.remove(axis_id)
        for value in (category_id, str(secondary_id)):
            axis_id = OxmlElement("c:axId")
            axis_id.set("val", value)
            plot.append(axis_id)
        bar_axis_ids = [item.get("val") for item in bar_plot.findall(qn("c:axId"))]
        if bar_axis_ids != [category_id, primary_value_id]:
            for axis_id in bar_plot.findall(qn("c:axId")):
                bar_plot.remove(axis_id)
            for value in (category_id, primary_value_id):
                axis_id = OxmlElement("c:axId")
                axis_id.set("val", value)
                bar_plot.append(axis_id)
    elif plot is None or len(list(plot_area.iter_xCharts())) != 1:
        raise ValueError("stock charts require one native stock plot")
    elif len(plot.findall(qn("c:ser"))) != expected_series:
        raise ValueError(f"{chart_type} charts require exactly {expected_series} series")
    plot.tag = qn("c:stockChart")
    for tag in ("c:grouping", "c:varyColors", "c:marker", "c:smooth", "c:hiLowLines", "c:upDownBars"):
        child = plot.find(qn(tag))
        if child is not None:
            plot.remove(child)
    first_axis = plot.find(qn("c:axId"))
    high_low_lines = OxmlElement("c:hiLowLines")
    if first_axis is None:
        plot.append(high_low_lines)
    else:
        first_axis.addprevious(high_low_lines)
    if chart_type in {OfficeChartType.STOCK_OHLC.value, OfficeChartType.STOCK_VOHLC.value}:
        up_down_bars = OxmlElement("c:upDownBars")
        gap_width = OxmlElement("c:gapWidth")
        gap_width.set("val", "150")
        up_down_bars.append(gap_width)
        high_low_lines.addnext(up_down_bars)


@contextmanager
def _presentation_stock_chart_as_line(chart: Any):
    """Temporarily expose a stock plot through python-pptx's supported line API."""
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement

    chart_type = _presentation_chart_type_name_for_chart(chart)
    if not _is_stock_chart_type(chart_type):
        yield chart
        return
    plot = chart._chartSpace.chart.plotArea.find(qn("c:stockChart"))
    if plot is None:
        raise ValueError("stock chart plot is missing")
    retained_stock_style = {
        tag: copy.deepcopy(plot.find(qn(tag)))
        for tag in ("c:hiLowLines", "c:upDownBars")
        if plot.find(qn(tag)) is not None
    }
    line_plot = OxmlElement("c:lineChart")
    for child in list(plot):
        line_plot.append(child)
    plot.addprevious(line_plot)
    plot.getparent().remove(plot)
    plot = line_plot
    if _is_volume_stock_chart_type(chart_type):
        bar_plot = chart._chartSpace.chart.plotArea.find(qn("c:barChart"))
        if bar_plot is not None and plot.getprevious() is not bar_plot:
            plot.getparent().remove(plot)
            bar_plot.addnext(plot)
    for tag in ("c:hiLowLines", "c:upDownBars"):
        child = plot.find(qn(tag))
        if child is not None:
            plot.remove(child)
    first_series = plot.find(qn("c:ser"))
    for tag, value in (("c:grouping", "standard"), ("c:varyColors", "0")):
        child = OxmlElement(tag)
        child.set("val", value)
        if first_series is None:
            plot.insert(0, child)
        else:
            first_series.addprevious(child)
    try:
        yield chart
    finally:
        _make_presentation_stock_chart(chart, chart_type)
        restored_plot = chart._chartSpace.chart.plotArea.find(qn("c:stockChart"))
        for tag, retained in retained_stock_style.items():
            generated = restored_plot.find(qn(tag)) if restored_plot is not None else None
            if generated is not None:
                generated.addprevious(retained)
                restored_plot.remove(generated)


def _make_presentation_combo_chart(chart: Any) -> None:
    """Keep every series except the last as columns and the last as a line."""
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement

    plot_area = chart._chartSpace.chart.plotArea
    bar_charts = plot_area.findall(qn("c:barChart"))
    line_charts = plot_area.findall(qn("c:lineChart"))
    if len(bar_charts) != 1 or len(line_charts) > 1:
        raise ValueError("combo_column_line requires one native column plot and at most one line plot")
    bar_chart = bar_charts[0]
    line_chart = line_charts[0] if line_charts else None
    series = [
        item
        for owner in (bar_chart, line_chart)
        if owner is not None
        for item in owner.findall(qn("c:ser"))
    ]
    if len(series) < 2:
        raise ValueError("combo_column_line charts require at least two series")

    def setting(parent: Any, tag: str, value: str) -> Any:
        element = parent.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            first_series = parent.find(qn("c:ser"))
            if first_series is None:
                parent.append(element)
            else:
                first_series.addprevious(element)
        element.set("val", value)
        return element

    setting(bar_chart, "c:barDir", "col")
    setting(bar_chart, "c:grouping", "clustered")
    if line_chart is None:
        line_chart = OxmlElement("c:lineChart")
        setting(line_chart, "c:grouping", "standard")
        setting(line_chart, "c:varyColors", "0")
        for axis_id in bar_chart.findall(qn("c:axId")):
            line_chart.append(copy.deepcopy(axis_id))
        bar_chart.addnext(line_chart)
    setting(line_chart, "c:grouping", "standard")
    setting(line_chart, "c:varyColors", "0")
    setting(line_chart, "c:marker", "1")
    setting(line_chart, "c:smooth", "0")

    def append_series(parent: Any, item: Any, boundary_names: set[str]) -> None:
        current_parent = item.getparent()
        if current_parent is not None:
            current_parent.remove(item)
        boundary = next(
            (child for child in parent if child.tag.rsplit("}", 1)[-1] in boundary_names),
            None,
        )
        if boundary is None:
            parent.append(item)
        else:
            boundary.addprevious(item)

    for item in series[:-1]:
        for tag in ("c:marker", "c:smooth"):
            child = item.find(qn(tag))
            if child is not None:
                item.remove(child)
        append_series(bar_chart, item, {"dLbls", "gapWidth", "overlap", "serLines", "axId", "extLst"})
    line_series = series[-1]
    inverted = line_series.find(qn("c:invertIfNegative"))
    if inverted is not None:
        line_series.remove(inverted)
    append_series(
        line_chart,
        line_series,
        {"dLbls", "dropLines", "hiLowLines", "upDownBars", "marker", "smooth", "axId", "extLst"},
    )


def _is_scatter_chart_type(chart_type: Any) -> bool:
    return chart_type in {
        OfficeChartType.SCATTER.value,
        OfficeChartType.SCATTER_LINES.value,
        OfficeChartType.SCATTER_LINES_MARKERS.value,
        OfficeChartType.SCATTER_SMOOTH.value,
        OfficeChartType.SCATTER_SMOOTH_MARKERS.value,
    }


def _normalized_chart_data(
    patch: dict[str, Any], *, single_series: bool, numeric_categories: bool = False,
    minimum_series: int = 1, exact_series: int | None = None,
) -> tuple[list[str | int | float], list[tuple[str, list[int | float | None]]]]:
    """Validate the common inline data contract used by native PPT/Excel charts."""
    categories = patch.get("categories")
    series = patch.get("series")
    if not isinstance(categories, list) or not 1 <= len(categories) <= 1000:
        raise ValueError("chart categories must contain 1 to 1000 labels")
    if not isinstance(series, list) or not 1 <= len(series) <= 50:
        raise ValueError("chart series must contain 1 to 50 series")
    if len(series) < minimum_series:
        raise ValueError(f"chart type requires at least {minimum_series} series")
    if exact_series is not None and len(series) != exact_series:
        raise ValueError(f"chart type requires exactly {exact_series} series")
    if len(categories) * len(series) > 10000:
        raise ValueError("chart data cannot exceed 10000 values")
    labels: list[str | int | float] = []
    for category in categories:
        if numeric_categories and (
            isinstance(category, bool) or type(category) not in {int, float} or not math.isfinite(category)
        ):
            raise ValueError("scatter chart categories must be finite numeric x values")
        if isinstance(category, bool) or type(category) not in {str, int, float}:
            raise ValueError("chart categories must be strings or finite numbers")
        if isinstance(category, float) and not math.isfinite(category):
            raise ValueError("chart categories must be finite")
        if isinstance(category, str) and (len(category) > 1000 or "\x00" in category):
            raise ValueError("chart category labels must be at most 1000 characters without NUL")
        labels.append(category)
    if single_series and len(series) != 1:
        raise ValueError("pie and doughnut charts require exactly one series")
    normalized_series = []
    for item in series:
        if not isinstance(item, dict) or set(item) != {"name", "values"}:
            raise ValueError("each chart series requires only name and values")
        name, values = item["name"], item["values"]
        if not isinstance(name, str) or not name.strip() or len(name) > 255 or "\x00" in name:
            raise ValueError("chart series name must be a non-empty string of at most 255 characters")
        if not isinstance(values, list) or len(values) != len(labels):
            raise ValueError("each chart series values array must match the category count")
        normalized_values = []
        for value in values:
            if value is None:
                normalized_values.append(None)
            elif isinstance(value, bool) or type(value) not in {int, float} or not math.isfinite(value):
                raise ValueError("chart values must be finite numbers or null")
            else:
                normalized_values.append(value)
        normalized_series.append((name, normalized_values))
    return labels, normalized_series


def _presentation_chart_data(patch: dict[str, Any], chart_type: Any):
    """Build bounded native chart data from the shared JSON operation contract."""
    from pptx.chart.data import CategoryChartData, XyChartData

    chart_type_name = (
        chart_type if chart_type in OfficeChartType.values()
        else _presentation_chart_type_name(chart_type)
    )
    is_scatter = _is_scatter_chart_type(chart_type_name)
    labels, series = _normalized_chart_data(
        patch,
        single_series=chart_type_name in {OfficeChartType.PIE.value, OfficeChartType.DOUGHNUT.value},
        numeric_categories=is_scatter,
        minimum_series=2 if _is_combo_chart_type(chart_type_name) else 1,
        exact_series=_stock_chart_series_count(chart_type_name),
    )
    if is_scatter:
        data = XyChartData()
        for name, values in series:
            xy_series = data.add_series(name)
            for x_value, y_value in zip(labels, values):
                xy_series.add_data_point(x_value, y_value)
        return data
    data = CategoryChartData()
    data.categories = labels
    for name, values in series:
        data.add_series(name, values)
    return data


def _presentation_plot_data_labels(plot: Any, *, create: bool = False) -> Any | None:
    from pptx.chart.datalabel import DataLabels
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement

    element = plot._element
    labels = element.find(qn("c:dLbls"))
    if labels is None and create:
        labels = OxmlElement("c:dLbls")
        first_axis_id = element.find(qn("c:axId"))
        if first_axis_id is None:
            element.append(labels)
        else:
            first_axis_id.addprevious(labels)
    return DataLabels(labels) if labels is not None else None


def _set_presentation_plot_data_labels(plot: Any, visible: bool) -> Any | None:
    from pptx.oxml.ns import qn

    labels = _presentation_plot_data_labels(plot, create=visible)
    if not visible:
        element = plot._element.find(qn("c:dLbls"))
        if element is not None:
            plot._element.remove(element)
        return None
    return labels


def _normalized_chart_trendlines(value: Any, series_count: int) -> list[tuple[int, dict[str, Any] | None]]:
    if not isinstance(value, list) or not 1 <= len(value) <= 50:
        raise ValueError("trendlines must contain 1 to 50 series configurations")
    if series_count < 1:
        raise ValueError("trendlines require at least one chart series")
    types = {"linear", "exponential", "logarithmic", "polynomial", "power", "moving_average"}
    allowed = {
        "series_index", "type", "name", "order", "period", "forward", "backward", "intercept",
        "display_equation", "display_r_squared",
    }
    normalized = []
    seen = set()
    for item in value:
        if not isinstance(item, dict) or set(item) - allowed or not {"series_index", "type"} <= set(item):
            raise ValueError("each trendline requires series_index and type plus supported trendline fields")
        index = _integer(item["series_index"], "trendline series_index", 0, series_count - 1)
        if index in seen:
            raise ValueError("trendline series_index values must be unique")
        seen.add(index)
        trendline_type = item["type"]
        if trendline_type is None:
            if set(item) != {"series_index", "type"}:
                raise ValueError("a null trendline type clears the series trendline and accepts no other fields")
            normalized.append((index, None))
            continue
        if trendline_type not in types:
            raise ValueError(f"trendline type must be one of: {', '.join(sorted(types))}, or null")
        config: dict[str, Any] = {"type": trendline_type}
        if "name" in item:
            name = item["name"]
            if not isinstance(name, str) or not name or len(name) > 255 or "\x00" in name:
                raise ValueError("trendline name must be a non-empty string of at most 255 characters")
            config["name"] = name
        if "order" in item:
            if trendline_type != "polynomial":
                raise ValueError("trendline order is supported only for polynomial trendlines")
            config["order"] = _integer(item["order"], "trendline order", 2, 6)
        elif trendline_type == "polynomial":
            config["order"] = 2
        if "period" in item:
            if trendline_type != "moving_average":
                raise ValueError("trendline period is supported only for moving_average trendlines")
            config["period"] = _integer(item["period"], "trendline period", 2, 255)
        elif trendline_type == "moving_average":
            config["period"] = 2
        for key in ("forward", "backward"):
            if key in item:
                config[key] = _number(item[key], f"trendline {key}", 0, 1e6)
        if "intercept" in item:
            config["intercept"] = _number(item["intercept"], "trendline intercept", -1e15, 1e15)
        for key in ("display_equation", "display_r_squared"):
            if key in item:
                if type(item[key]) is not bool:
                    raise ValueError(f"trendline {key} must be boolean")
                config[key] = item[key]
        normalized.append((index, config))
    return normalized


def _normalized_chart_error_bars(
    value: Any,
    series_count: int,
    *,
    scatter: bool,
) -> list[tuple[int, dict[str, Any] | None]]:
    if not isinstance(value, list) or not 1 <= len(value) <= 50:
        raise ValueError("error_bars must contain 1 to 50 series configurations")
    if series_count < 1:
        raise ValueError("error_bars require at least one chart series")
    allowed = {"series_index", "type", "direction", "side", "value", "end_style"}
    types = {"fixed", "percentage", "standard_deviation", "standard_error"}
    normalized = []
    seen = set()
    for item in value:
        if not isinstance(item, dict) or set(item) - allowed or not {"series_index", "type"} <= set(item):
            raise ValueError("each error bar requires series_index and type plus supported error-bar fields")
        index = _integer(item["series_index"], "error-bar series_index", 0, series_count - 1)
        if index in seen:
            raise ValueError("error-bar series_index values must be unique")
        seen.add(index)
        error_type = item["type"]
        if error_type is None:
            if set(item) != {"series_index", "type"}:
                raise ValueError("a null error-bar type clears the series error bars and accepts no other fields")
            normalized.append((index, None))
            continue
        if error_type not in types:
            raise ValueError(f"error-bar type must be one of: {', '.join(sorted(types))}, or null")
        direction = item.get("direction", "y")
        if direction not in {"x", "y"} or (direction == "x" and not scatter):
            raise ValueError("error-bar direction must be y, or x for scatter charts")
        side = item.get("side", "both")
        if side not in {"both", "plus", "minus"}:
            raise ValueError("error-bar side must be both, plus or minus")
        end_style = item.get("end_style", "cap")
        if end_style not in {"cap", "no_cap"}:
            raise ValueError("error-bar end_style must be cap or no_cap")
        config: dict[str, Any] = {
            "type": error_type,
            "direction": direction,
            "side": side,
            "end_style": end_style,
        }
        if error_type in {"fixed", "percentage", "standard_deviation"}:
            if "value" not in item:
                raise ValueError(f"{error_type} error bars require value")
            maximum = 100 if error_type == "percentage" else 1e15
            config["value"] = _number(item["value"], "error-bar value", 0, maximum)
        elif "value" in item:
            raise ValueError("standard_error error bars do not accept value")
        normalized.append((index, config))
    return normalized


def _presentation_chart_trendline_details(series: Any) -> dict[str, Any] | None:
    namespace = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
    trendline = series._element.find(namespace + "trendline")
    if trendline is None:
        return None
    native_types = {
        "linear": "linear", "exp": "exponential", "log": "logarithmic",
        "poly": "polynomial", "power": "power", "movingAvg": "moving_average",
    }
    type_element = trendline.find(namespace + "trendlineType")
    details: dict[str, Any] = {"type": native_types.get(type_element.get("val") if type_element is not None else "")}
    name = trendline.find(namespace + "name")
    if name is not None and name.text:
        details["name"] = name.text[:255]
    for native, public, converter in (
        ("order", "order", int), ("period", "period", int),
        ("forward", "forward", float), ("backward", "backward", float), ("intercept", "intercept", float),
    ):
        element = trendline.find(namespace + native)
        if element is not None and element.get("val") is not None:
            try:
                details[public] = converter(element.get("val"))
            except (TypeError, ValueError):
                pass
    for native, public in (("dispEq", "display_equation"), ("dispRSqr", "display_r_squared")):
        element = trendline.find(namespace + native)
        if element is not None:
            details[public] = element.get("val", "1") not in {"0", "false", "off"}
    return details


def _presentation_chart_error_bar_details(series: Any) -> dict[str, Any] | None:
    namespace = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
    error_bars = series._element.find(namespace + "errBars")
    if error_bars is None:
        return None
    types = {
        "fixedVal": "fixed", "percentage": "percentage",
        "stdDev": "standard_deviation", "stdErr": "standard_error",
    }

    def value(tag: str, default: Any = None) -> Any:
        element = error_bars.find(namespace + tag)
        return element.get("val", default) if element is not None else default

    details: dict[str, Any] = {
        "type": types.get(value("errValType")),
        "direction": value("errDir", "y"),
        "side": value("errBarType", "both"),
        "end_style": "no_cap" if value("noEndCap", "0") in {"1", "true", "on"} else "cap",
    }
    raw_value = value("val")
    if raw_value is not None:
        try:
            details["value"] = float(raw_value)
        except (TypeError, ValueError):
            pass
    return details


def _format_presentation_chart_trendlines(chart: Any, value: Any) -> None:
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement

    chart_type = _presentation_chart_type_name_for_chart(chart)
    supported = {
        OfficeChartType.COLUMN.value, OfficeChartType.BAR.value,
        OfficeChartType.LINE.value, OfficeChartType.LINE_MARKERS.value,
        OfficeChartType.SCATTER.value, OfficeChartType.SCATTER_LINES.value,
        OfficeChartType.SCATTER_LINES_MARKERS.value, OfficeChartType.SCATTER_SMOOTH.value,
        OfficeChartType.SCATTER_SMOOTH_MARKERS.value,
        OfficeChartType.COMBO_COLUMN_LINE.value,
    }
    if chart_type not in supported:
        raise ValueError("trendlines require a non-stacked column, bar, line or scatter chart")
    series_items = list(chart.series)
    native_types = {
        "linear": "linear", "exponential": "exp", "logarithmic": "log",
        "polynomial": "poly", "power": "power", "moving_average": "movingAvg",
    }
    for index, config in _normalized_chart_trendlines(value, len(series_items)):
        series = series_items[index]
        for existing in list(series._element.findall(qn("c:trendline"))):
            series._element.remove(existing)
        if config is None:
            continue
        trendline = OxmlElement("c:trendline")
        if "name" in config:
            name = OxmlElement("c:name")
            name.text = config["name"]
            trendline.append(name)
        for tag, raw_value in (
            ("trendlineType", native_types[config["type"]]),
            ("order", config.get("order")), ("period", config.get("period")),
            ("forward", config.get("forward")), ("backward", config.get("backward")),
            ("intercept", config.get("intercept")),
            ("dispRSqr", config.get("display_r_squared")),
            ("dispEq", config.get("display_equation")),
        ):
            if raw_value is None:
                continue
            element = OxmlElement(f"c:{tag}")
            element.set("val", "1" if raw_value is True else "0" if raw_value is False else str(raw_value))
            trendline.append(element)
        insertion = next((
            series._element.find(qn(f"c:{tag}")) for tag in ("errBars", "cat", "val", "xVal", "yVal")
            if series._element.find(qn(f"c:{tag}")) is not None
        ), None)
        if insertion is None:
            series._element.append(trendline)
        else:
            insertion.addprevious(trendline)


def _format_presentation_chart_error_bars(chart: Any, value: Any) -> None:
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement

    chart_type = _presentation_chart_type_name_for_chart(chart)
    scatter = _is_scatter_chart_type(chart_type)
    supported = {
        OfficeChartType.COLUMN.value, OfficeChartType.BAR.value,
        OfficeChartType.LINE.value, OfficeChartType.LINE_MARKERS.value,
        OfficeChartType.SCATTER.value, OfficeChartType.SCATTER_LINES.value,
        OfficeChartType.SCATTER_LINES_MARKERS.value, OfficeChartType.SCATTER_SMOOTH.value,
        OfficeChartType.SCATTER_SMOOTH_MARKERS.value,
        OfficeChartType.COMBO_COLUMN_LINE.value,
    }
    if chart_type not in supported:
        raise ValueError("error bars require a non-stacked column, bar, line or scatter chart")
    series_items = list(chart.series)
    native_types = {
        "fixed": "fixedVal", "percentage": "percentage",
        "standard_deviation": "stdDev", "standard_error": "stdErr",
    }
    for index, config in _normalized_chart_error_bars(value, len(series_items), scatter=scatter):
        series = series_items[index]
        for existing in list(series._element.findall(qn("c:errBars"))):
            series._element.remove(existing)
        if config is None:
            continue
        error_bars = OxmlElement("c:errBars")
        for tag, raw_value in (
            ("errDir", config["direction"]), ("errBarType", config["side"]),
            ("errValType", native_types[config["type"]]),
            ("noEndCap", config["end_style"] == "no_cap"), ("val", config.get("value")),
        ):
            if raw_value is None:
                continue
            element = OxmlElement(f"c:{tag}")
            element.set("val", "1" if raw_value is True else "0" if raw_value is False else str(raw_value))
            error_bars.append(element)
        insertion = next((
            series._element.find(qn(f"c:{tag}")) for tag in ("cat", "val", "xVal", "yVal")
            if series._element.find(qn(f"c:{tag}")) is not None
        ), None)
        if insertion is None:
            series._element.append(error_bars)
        else:
            insertion.addprevious(error_bars)


def _format_presentation_chart(chart: Any, style: Any) -> None:
    """Apply explicit chart presentation properties without flattening the chart."""
    from pptx.chart.axis import ValueAxis
    from pptx.enum.chart import XL_LABEL_POSITION, XL_LEGEND_POSITION

    allowed = {
        "title", "style", "has_legend", "legend_position", "legend_include_in_layout",
        "vary_colors", "show_data_labels", "show_category_name", "show_series_name",
        "show_value", "show_percentage", "data_label_position", "category_axis_title",
        "value_axis_title", "show_category_axis", "show_value_axis", "category_axis_min",
        "category_axis_max", "category_axis_major_unit", "value_axis_min", "value_axis_max",
        "value_axis_major_unit", "show_major_gridlines", "series_colors", "category_colors", "trendlines",
        "error_bars", "secondary_value_axis_title", "show_secondary_value_axis",
        "secondary_value_axis_min", "secondary_value_axis_max", "secondary_value_axis_major_unit",
        "show_secondary_major_gridlines",
    }
    if not isinstance(style, dict) or not style or set(style) - allowed:
        raise ValueError(f"chart.format requires supported format keys: {', '.join(sorted(allowed))}")

    def text(value: Any, name: str, maximum: int = 500) -> str | None:
        if value is not None and (not isinstance(value, str) or len(value) > maximum or "\x00" in value):
            raise ValueError(f"{name} must be null or a string of at most {maximum} characters")
        return value

    def boolean(value: Any, name: str) -> bool:
        if type(value) is not bool:
            raise ValueError(f"{name} must be boolean")
        return value

    if "title" in style:
        title = text(style["title"], "title")
        chart.has_title = title is not None
        if title is not None:
            chart.chart_title.text_frame.text = title
    if "style" in style:
        chart.chart_style = _integer(style["style"], "chart style", 1, 48)
    if "has_legend" in style:
        chart.has_legend = boolean(style["has_legend"], "has_legend")
    if "legend_position" in style:
        position = style["legend_position"]
        positions = {
            "top": XL_LEGEND_POSITION.TOP,
            "bottom": XL_LEGEND_POSITION.BOTTOM,
            "left": XL_LEGEND_POSITION.LEFT,
            "right": XL_LEGEND_POSITION.RIGHT,
        }
        if position not in positions:
            raise ValueError("legend_position must be top, bottom, left or right")
        chart.has_legend = True
        chart.legend.position = positions[position]
    if "legend_include_in_layout" in style:
        chart.has_legend = True
        chart.legend.include_in_layout = boolean(style["legend_include_in_layout"], "legend_include_in_layout")
    plots = list(chart.plots)
    if "vary_colors" in style:
        vary_colors = boolean(style["vary_colors"], "vary_colors")
        for target_plot in plots:
            target_plot.vary_by_categories = vary_colors
    label_keys = {
        "show_category_name", "show_series_name", "show_value", "show_percentage", "data_label_position",
    }
    if "show_data_labels" in style:
        show_data_labels = boolean(style["show_data_labels"], "show_data_labels")
        for target_plot in plots:
            _set_presentation_plot_data_labels(target_plot, show_data_labels)
    if label_keys & style.keys():
        labels_by_plot = [_set_presentation_plot_data_labels(target_plot, True) for target_plot in plots]
        for key in label_keys - {"data_label_position"}:
            if key in style:
                label_value = boolean(style[key], key)
                for labels in labels_by_plot:
                    setattr(labels, key, label_value)
        if "data_label_position" in style:
            positions = {
                "best_fit": XL_LABEL_POSITION.BEST_FIT,
                "center": XL_LABEL_POSITION.CENTER,
                "inside_end": XL_LABEL_POSITION.INSIDE_END,
                "outside_end": XL_LABEL_POSITION.OUTSIDE_END,
            }
            position = style["data_label_position"]
            if position not in positions:
                raise ValueError("data_label_position must be best_fit, center, inside_end or outside_end")
            for labels in labels_by_plot:
                labels.position = positions[position]

    axes = {}
    for name, attribute in (("category", "category_axis"), ("value", "value_axis")):
        try:
            axes[name] = getattr(chart, attribute)
        except ValueError:
            axes[name] = None
    _category_node, primary_value_node, secondary_value_node = _presentation_category_value_axis_nodes(chart)
    if _category_node is not None and primary_value_node is not None:
        axes["value"] = ValueAxis(primary_value_node)
    for name in ("category", "value"):
        title_key = f"{name}_axis_title"
        visible_key = f"show_{name}_axis"
        if title_key in style:
            if axes[name] is None:
                raise ValueError(f"this chart type has no {name} axis")
            title = text(style[title_key], title_key)
            axes[name].has_title = title is not None
            if title is not None:
                axes[name].axis_title.text_frame.text = title
        if visible_key in style:
            if axes[name] is None:
                raise ValueError(f"this chart type has no {name} axis")
            axes[name].visible = boolean(style[visible_key], visible_key)
    for name in ("category", "value"):
        axis = axes[name]
        scale_values = {}
        for suffix, attribute in (
            ("min", "minimum_scale"), ("max", "maximum_scale"), ("major_unit", "major_unit"),
        ):
            key = f"{name}_axis_{suffix}"
            if key not in style:
                continue
            if axis is None or (name == "category" and not _is_scatter_chart_type(
                _presentation_chart_type_name_for_chart(chart)
            )):
                raise ValueError(f"this chart type has no numeric {name} axis")
            value = style[key]
            if value is not None:
                value = _number(value, key, -1e15 if suffix != "major_unit" else 1e-12, 1e15)
            scale_values[attribute] = value
        prospective_min = scale_values.get("minimum_scale", getattr(axis, "minimum_scale", None) if axis else None)
        prospective_max = scale_values.get("maximum_scale", getattr(axis, "maximum_scale", None) if axis else None)
        if prospective_min is not None and prospective_max is not None and prospective_min >= prospective_max:
            raise ValueError(f"{name}_axis_min must be less than {name}_axis_max")
        for attribute, value in scale_values.items():
            setattr(axis, attribute, value)
    value_axis = axes["value"]
    if "show_major_gridlines" in style:
        if value_axis is None:
            raise ValueError("this chart type has no value axis")
        value_axis.has_major_gridlines = boolean(style["show_major_gridlines"], "show_major_gridlines")

    secondary_axis = ValueAxis(secondary_value_node) if secondary_value_node is not None else None
    secondary_keys = {
        "secondary_value_axis_title", "show_secondary_value_axis",
        "secondary_value_axis_min", "secondary_value_axis_max", "secondary_value_axis_major_unit",
        "show_secondary_major_gridlines",
    }
    if secondary_keys & style.keys() and secondary_axis is None:
        raise ValueError("this chart type has no secondary value axis")
    if secondary_axis is not None:
        if "secondary_value_axis_title" in style:
            title = text(style["secondary_value_axis_title"], "secondary_value_axis_title")
            secondary_axis.has_title = title is not None
            if title is not None:
                secondary_axis.axis_title.text_frame.text = title
        if "show_secondary_value_axis" in style:
            secondary_axis.visible = boolean(style["show_secondary_value_axis"], "show_secondary_value_axis")
        scale_values = {}
        for suffix, attribute in (
            ("min", "minimum_scale"), ("max", "maximum_scale"), ("major_unit", "major_unit"),
        ):
            key = f"secondary_value_axis_{suffix}"
            if key not in style:
                continue
            value = style[key]
            if value is not None:
                value = _number(value, key, -1e15 if suffix != "major_unit" else 1e-12, 1e15)
            scale_values[attribute] = value
        prospective_min = scale_values.get("minimum_scale", secondary_axis.minimum_scale)
        prospective_max = scale_values.get("maximum_scale", secondary_axis.maximum_scale)
        if prospective_min is not None and prospective_max is not None and prospective_min >= prospective_max:
            raise ValueError("secondary_value_axis_min must be less than secondary_value_axis_max")
        for attribute, value in scale_values.items():
            setattr(secondary_axis, attribute, value)
        if "show_secondary_major_gridlines" in style:
            secondary_axis.has_major_gridlines = boolean(
                style["show_secondary_major_gridlines"], "show_secondary_major_gridlines",
            )

    if "series_colors" in style:
        key, targets = "series_colors", list(chart.series)
        colors = style[key]
        if not isinstance(colors, list) or not colors or len(colors) > len(targets):
            raise ValueError(f"{key} must be a non-empty RGB array no longer than its chart target count")
        for color_value, target in zip(colors, targets):
            color = _validated_cell_format({"fill_color": color_value}, spreadsheet=False)["fill_color"]
            target.format.fill.solid()
            _set_office_rgb(target.format.fill.fore_color, color, word=False)
            target.format.line.fill.solid()
            _set_office_rgb(target.format.line.color, color, word=False)
    if "category_colors" in style:
        key = "category_colors"
        if _presentation_chart_type_name_for_chart(chart) not in {
            OfficeChartType.PIE.value, OfficeChartType.DOUGHNUT.value,
        }:
            raise ValueError("category_colors is supported only for pie and doughnut charts")
        targets = list(chart.series[0].points) if len(chart.series) else []
        colors = style[key]
        if not isinstance(colors, list) or not colors or len(colors) > len(targets):
            raise ValueError(f"{key} must be a non-empty RGB array no longer than its chart target count")
        for color_value, target in zip(colors, targets):
            color = _validated_cell_format({"fill_color": color_value}, spreadsheet=False)["fill_color"]
            target.format.fill.solid()
            _set_office_rgb(target.format.fill.fore_color, color, word=False)
            target.format.line.fill.solid()
            _set_office_rgb(target.format.line.color, color, word=False)
    if "trendlines" in style:
        _format_presentation_chart_trendlines(chart, style["trendlines"])
    if "error_bars" in style:
        _format_presentation_chart_error_bars(chart, style["error_bars"])


def apply_office_structure_patch(
    abs_path: str,
    patch: dict[str, Any],
    *,
    resources: dict[tuple[str, str], bytes] | None = None,
) -> dict[str, Any]:
    extension = file_type_from_path(abs_path)
    operation = patch["operation"]
    if extension == "docx":
        from docx import Document

        document = Document(abs_path)
        if operation in {"paragraph_style.insert", "paragraph_style.format", "paragraph_style.delete"}:
            allowed = (
                {"operation", "style"}
                if operation == "paragraph_style.delete"
                else {"operation", "style", "base_style", "next_style", "format"}
            )
            if set(patch) - allowed:
                raise ValueError(f"Word {operation} contains unsupported properties")
            style_name = _validated_word_paragraph_style_name(patch.get("style"))
            if operation == "paragraph_style.insert":
                from docx.enum.style import WD_STYLE_TYPE

                if any(candidate.name == style_name for candidate in document.styles):
                    raise ValueError(f"Word style already exists: {style_name}")
                existing_style_ids = {candidate.style_id for candidate in document.styles}
                try:
                    style = document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
                except ValueError:
                    raise ValueError("Word style name conflicts with an existing style name or ID") from None
                if style.style_id in existing_style_ids:
                    raise ValueError(
                        f"Word style name would duplicate the existing style ID: {style.style_id}",
                    )
                base_name = patch.get("base_style", "Normal")
                if base_name is not None:
                    base = _word_paragraph_style(document, base_name)
                    if base.style_id == style.style_id:
                        raise ValueError("base_style cannot reference the style itself")
                    style.base_style = base
            else:
                style = _word_paragraph_style(document, style_name)
            if operation == "paragraph_style.delete":
                if style.builtin:
                    raise ValueError("Built-in Word paragraph styles cannot be deleted")
                if _word_paragraph_style_is_used(document, style.style_id):
                    raise ValueError("Word paragraph style is still used by document content")
                if _word_paragraph_style_is_referenced(document, style):
                    raise ValueError("Word paragraph style is still referenced by another style")
                style_id = style.style_id
                style.delete()
                selector = {"style": style_name, "style_id": style_id}
            else:
                if operation == "paragraph_style.format" and not {
                    "base_style", "next_style", "format",
                } & patch.keys():
                    raise ValueError("paragraph_style.format requires base_style, next_style or format")
                if "base_style" in patch:
                    base_name = patch["base_style"]
                    base = None if base_name is None else _word_paragraph_style(document, base_name)
                    if base is not None and base.style_id == style.style_id:
                        raise ValueError("base_style cannot reference the style itself")
                    style.base_style = base
                if "next_style" in patch:
                    next_name = patch["next_style"]
                    style.next_paragraph_style = (
                        None if next_name is None else _word_paragraph_style(document, next_name)
                    )
                if "format" in patch:
                    _format_word_paragraph_style(style, patch["format"])
                selector = {"style": style.name, "style_id": style.style_id}
        elif operation == "section.insert":
            from docx.enum.section import WD_SECTION

            if set(patch) - {"operation", "index", "start_type", "inherit_headers_footers", "format"}:
                raise ValueError(
                    "Word section.insert accepts append index, start_type, inherit_headers_footers and format"
                )
            index = _integer(patch.get("index", len(document.sections)), "index", 0, len(document.sections))
            if index != len(document.sections):
                raise ValueError("Word section.insert currently appends; index must equal section_count")
            start_type = patch.get("start_type", "new_page")
            start_types = {
                "continuous": WD_SECTION.CONTINUOUS,
                "new_page": WD_SECTION.NEW_PAGE,
                "even_page": WD_SECTION.EVEN_PAGE,
                "odd_page": WD_SECTION.ODD_PAGE,
            }
            if start_type not in start_types:
                raise ValueError(f"start_type must be one of: {', '.join(sorted(start_types))}")
            inherit = patch.get("inherit_headers_footers", True)
            if type(inherit) is not bool:
                raise ValueError("inherit_headers_footers must be boolean")
            section = document.add_section(start_types[start_type])
            section_index = len(document.sections) - 1
            if not inherit:
                for story in [
                    section.header,
                    section.first_page_header,
                    section.even_page_header,
                    section.footer,
                    section.first_page_footer,
                    section.even_page_footer,
                ]:
                    story.is_linked_to_previous = False
            if "format" in patch:
                _setup_office_page(
                    document,
                    {"operation": "page.setup", "section_index": section_index, "format": patch["format"]},
                    word=True,
                )
            selector = {
                "index": section_index,
                "section_index": section_index,
                "start_type": start_type,
                "inherit_headers_footers": inherit,
            }
        elif operation == "page.setup":
            selector = _setup_office_page(document, patch, word=True)
        elif operation == "picture.insert":
            from docx.shared import Emu

            if set(patch) - {"operation", "index", "story", "section_index", "source", "transform", "format"}:
                raise ValueError("Word picture.insert accepts index, story, section_index, source, transform and format")
            story_object, story, section_index = _word_story(document, patch.get("story"), patch.get("section_index"))
            paragraphs = story_object.paragraphs
            index = _integer(patch.get("index", len(paragraphs)), "index", 0, len(paragraphs))
            anchor = paragraphs[index]._p if index < len(paragraphs) else None
            transform = patch.get("transform", {})
            if not isinstance(transform, dict) or set(transform) - {"width", "height"}:
                raise ValueError("Word picture transform accepts only width and height in points")
            dimensions = {
                key: _office_points(value, key, word=True, minimum=1, maximum=10000)
                for key, value in transform.items()
            }
            data = _office_picture_resource(patch, resources)
            paragraph = story_object.add_paragraph()
            if anchor is not None:
                anchor.addprevious(paragraph._p)
            inline_picture = paragraph.add_run().add_picture(
                io.BytesIO(data),
                **{key: Emu(round(value * 12700)) for key, value in dimensions.items()},
            )
            picture = next(
                item for item in _word_pictures(document)
                if item.element is inline_picture._inline
            )
            if "format" in patch:
                _format_word_picture(picture, patch["format"])
            picture_index = next(
                item_index for item_index, item in enumerate(_word_pictures(document))
                if item.element is picture.element
            )
            selector = {"index": index, "picture_index": picture_index, "story": story}
            if section_index is not None:
                selector["section_index"] = section_index
        elif operation in {"picture.delete", "picture.replace", "picture.format"}:
            allowed = (
                {"operation", "picture_index", "source"} if operation == "picture.replace"
                else {"operation", "picture_index", "format"} if operation == "picture.format"
                else {"operation", "picture_index"}
            )
            if set(patch) - allowed:
                raise ValueError(f"Word {operation} contains unsupported properties")
            pictures = _word_pictures(document)
            if not pictures:
                raise ValueError("Word document has no editable DrawingML pictures")
            picture_index = _integer(patch.get("picture_index"), "picture_index", 0, len(pictures) - 1)
            picture = pictures[picture_index]
            if operation == "picture.delete":
                drawing = picture.drawing_element
                relationship_ids = _office_relationship_ids(drawing)
                run = drawing.getparent()
                run.remove(drawing)
                if not any(child.tag.rsplit("}", 1)[-1] != "rPr" for child in run):
                    run.getparent().remove(run)
                _drop_unused_office_relationships(picture.part, picture.root, relationship_ids)
            elif operation == "picture.format":
                _format_word_picture(picture, patch.get("format"))
            else:
                data = _office_picture_resource(patch, resources)
                old_relationship = picture.relationship_id
                relationship, _image = picture.part.get_or_add_image(io.BytesIO(data))
                picture.set_relationship_id(relationship)
                if old_relationship != relationship:
                    _drop_unused_office_relationships(picture.part, picture.root, {old_relationship})
            selector = {"picture_index": picture_index}
        elif operation == "textbox.insert":
            if set(patch) - {
                "operation", "index", "story", "section_index", "text", "preset", "transform", "format",
            }:
                raise ValueError(
                    "Word textbox.insert accepts index, story, section_index, text, preset, transform and format",
                )
            story_object, story, section_index = _word_story(
                document, patch.get("story"), patch.get("section_index"),
            )
            paragraphs = story_object.paragraphs
            index = _integer(patch.get("index", len(paragraphs)), "index", 0, len(paragraphs))
            text = patch.get("text", "")
            if not isinstance(text, str):
                raise ValueError("text must be a string")
            transform = patch.get("transform")
            if not isinstance(transform, dict) or set(transform) - {"x", "y", "width", "height", "rotation"}:
                raise ValueError("Word textbox.insert transform accepts x, y, width, height and rotation")
            if not {"x", "y", "width", "height"} <= transform.keys():
                raise ValueError("Word textbox.insert requires x, y, width and height")
            preset = patch.get("preset", PresentationShapePreset.RECTANGLE.value)
            if preset not in PresentationShapePreset.values() or preset == PresentationShapePreset.LINE.value:
                raise ValueError("Word textbox.insert preset must be a supported non-line DrawingML preset")
            paragraph = story_object.add_paragraph()
            if index < len(paragraphs):
                paragraphs[index]._p.addprevious(paragraph._p)
            run = _new_word_text_box_run(document, text, transform, preset)
            paragraph._p.append(run)
            text_box = next(
                item for item in _word_text_boxes(document)
                if item.content is next(node for node in run.iter() if node.tag.endswith("}txbxContent"))
            )
            if "format" in patch:
                _format_word_text_box(text_box, patch["format"])
            text_box_index = next(
                item_index for item_index, item in enumerate(_word_text_boxes(document))
                if item.content is text_box.content
            )
            selector = {"index": index, "text_box_index": text_box_index, "story": story}
            if section_index is not None:
                selector["section_index"] = section_index
        elif operation in {"shape.transform", "shape.format", "shape.delete"}:
            allowed = (
                {"operation", "text_box_index", "transform"} if operation == "shape.transform"
                else {"operation", "text_box_index", "format"} if operation == "shape.format"
                else {"operation", "text_box_index"}
            )
            if set(patch) - allowed:
                raise ValueError(f"Word {operation} targets only text_box_index and its payload")
            text_boxes = _word_text_boxes(document)
            if not text_boxes:
                raise ValueError("Word document has no editable text boxes")
            text_box_index = _integer(
                patch.get("text_box_index"), "text_box_index", 0, len(text_boxes) - 1,
            )
            text_box = text_boxes[text_box_index]
            if operation == "shape.transform":
                _word_text_box_transform(text_box, patch.get("transform"))
            elif operation == "shape.format":
                _format_word_text_box(text_box, patch.get("format"))
            else:
                _delete_word_text_box_shape(text_box)
            selector = {"text_box_index": text_box_index}
        elif operation == "paragraph.delete":
            if set(patch) - {"operation", "index", "story", "section_index", "text_box_index"}:
                raise ValueError("Word paragraph.delete accepts index and a story or text_box_index target")
            target, selector = _word_content_target(document, patch)
            paragraphs = target.paragraphs
            index = _integer(patch.get("index"), "index", 0, len(paragraphs) - 1)
            if len(paragraphs) == 1 and (isinstance(target, _WordTextBoxTarget) or selector.get("story") != "body"):
                raise ValueError("A Word header/footer or text box must retain at least one paragraph; use text.set with empty text")
            paragraph = paragraphs[index]
            properties = paragraph._p.pPr
            if selector.get("story") == "body" and properties is not None and properties.sectPr is not None:
                raise ValueError("Cannot delete a Word paragraph that owns a section break")
            relationship_ids = _office_relationship_ids(paragraph._p)
            paragraph._p.getparent().remove(paragraph._p)
            _drop_unused_office_relationships(target.part, _word_target_root(target), relationship_ids)
            selector["index"] = index
        elif operation in {"text.set", "paragraph.format"}:
            disallowed = {"slide", "shape_id", "table_index", "cell", "sheet", "table"}
            if set(patch) & disallowed:
                raise ValueError(f"Word {operation} targets a story or text-box paragraph index")
            target, selector = _word_content_target(document, patch)
            paragraphs = target.paragraphs
            index = _integer(patch.get("index"), "index", 0, len(paragraphs) - 1)
            protected = word_complex_field_nodes(_word_target_root(target))
            if operation == "text.set":
                _set_office_paragraph_text(paragraphs[index], patch.get("text"), word=True, protected_nodes=protected)
            else:
                _format_office_paragraph(paragraphs[index], patch, word=True, protected_nodes=protected)
            selector["index"] = index
        elif operation in {"set_cell", "cell.format"}:
            if any(key in patch for key in ("sheet", "slide", "shape_id", "table", "inherit_from_row")):
                raise ValueError("Word cell edits use table_index and cell, not spreadsheet/slide selectors")
            target, selector = _word_content_target(document, patch)
            tables = target.tables
            table_index = _integer(patch.get("table_index"), "table_index", 0, len(tables) - 1)
            cell = _office_table_cell(tables[table_index], patch.get("cell"), word=True)
            protected = word_complex_field_nodes(_word_target_root(target))
            if operation == "cell.format":
                _format_office_table_cell(cell, patch.get("format"), word=True, protected_nodes=protected)
            else:
                if "value" not in patch:
                    raise ValueError("value is required for cell.set")
                value = spreadsheet_cell_value(patch["value"])
                _set_office_table_cell_text(cell, "" if value is None else str(value), word=True, protected_nodes=protected)
            selector.update({"table_index": table_index, "cell": patch["cell"].upper()})
        elif operation == "table.insert":
            target, selector = _word_content_target(document, patch)
            paragraphs = target.paragraphs
            index = _integer(patch.get("index"), "index", 0, len(paragraphs))
            values = _office_table_rows(patch.get("rows"))
            anchor = paragraphs[index]._p if index < len(paragraphs) else None
            if isinstance(target, _WordTextBoxTarget) or selector.get("story") == "body":
                table = document.add_table(len(values), len(values[0]), style=patch.get("style"))
            else:
                section = document.sections[selector["section_index"]]
                width = section.page_width - section.left_margin - section.right_margin
                table = target.add_table(len(values), len(values[0]), width)
                if patch.get("style") is not None:
                    table.style = patch["style"]
            for row, values_row in zip(table.rows, values):
                for cell, value in zip(row.cells, values_row):
                    cell.text = "" if value is None else str(value)
            if anchor is not None:
                anchor.addprevious(table._tbl)
            elif isinstance(target, _WordTextBoxTarget):
                target.content.append(table._tbl)
            selector.update({
                "index": index,
                "table_index": next(i for i, item in enumerate(target.tables) if item._tbl is table._tbl),
            })
        elif operation == "table.format":
            if set(patch) - {
                "operation", "table_index", "story", "section_index", "text_box_index", "format",
            }:
                raise ValueError(
                    "Word table.format accepts table_index, format and a story or text_box_index target",
                )
            target, selector = _word_content_target(document, patch)
            tables = target.tables
            table_index = _integer(patch.get("table_index"), "table_index", 0, len(tables) - 1)
            _format_word_table(tables[table_index], patch.get("format"))
            selector["table_index"] = table_index
        elif operation == "table.delete":
            if set(patch) - {"operation", "table_index", "story", "section_index", "text_box_index"}:
                raise ValueError("Word table.delete accepts table_index and a story or text_box_index target")
            target, selector = _word_content_target(document, patch)
            tables = target.tables
            table_index = _integer(patch.get("table_index"), "table_index", 0, len(tables) - 1)
            table = tables[table_index]
            relationship_ids = _office_relationship_ids(table._tbl)
            table._tbl.getparent().remove(table._tbl)
            _drop_unused_office_relationships(target.part, _word_target_root(target), relationship_ids)
            selector["table_index"] = table_index
        elif operation == "paragraph.insert":
            if set(patch) - {"operation", "index", "story", "section_index", "text_box_index", "text", "style"}:
                raise ValueError("Word paragraph.insert accepts index, text, style and a story or text_box_index target")
            target, selector = _word_content_target(document, patch)
            paragraphs = target.paragraphs
            index = _integer(patch.get("index"), "index", 0, len(paragraphs))
            text = patch.get("text")
            if not isinstance(text, str):
                raise ValueError("text must be a string")
            _insert_word_target_paragraph(document, target, index, text, patch.get("style"))
            selector["index"] = index
        else:
            raise ValueError("Unsupported Word structure operation")
    else:
        from pptx import Presentation
        from pptx.util import Emu

        document = Presentation(abs_path)
        if operation == "page.setup":
            selector = _setup_office_page(document, patch, word=False)
        elif operation == "slide.format":
            if set(patch) - {"operation", "slide", "format"}:
                raise ValueError("slide.format accepts only slide and format")
            slide_number = _integer(patch.get("slide"), "slide", 1, len(document.slides))
            selector = {
                "slide": slide_number,
                **_format_presentation_slide(
                    document.slides[slide_number - 1], patch.get("format"),
                ),
            }
        elif operation == "slide.insert":
            index = _integer(patch.get("index"), "index", 0, len(document.slides))
            layout = _integer(patch.get("layout_index"), "layout_index", 0, len(document.slide_layouts) - 1)
            document.slides.add_slide(document.slide_layouts[layout])
            slide_ids = document.slides._sldIdLst
            added = slide_ids[-1]
            slide_ids.remove(added)
            slide_ids.insert(index, added)
            selector = {"index": index, "slide": index + 1, "layout_index": layout}
        elif operation == "slide.duplicate":
            if set(patch) - {"operation", "slide", "index"}:
                raise ValueError("slide.duplicate accepts source slide and optional zero-based insertion index")
            source_number = _integer(patch.get("slide"), "slide", 1, len(document.slides))
            index = _integer(patch.get("index", source_number), "index", 0, len(document.slides))
            selector = _duplicate_presentation_slide(document, source_number, index)
        elif operation == "slide.reorder":
            if set(patch) - {"operation", "slide", "index"}:
                raise ValueError("slide.reorder accepts only slide and index")
            source_number = _integer(patch.get("slide"), "slide", 1, len(document.slides))
            index = _integer(patch.get("index"), "index", 0, len(document.slides) - 1)
            previous_index = source_number - 1
            if index != previous_index:
                slide_ids = document.slides._sldIdLst
                slide_id = slide_ids[previous_index]
                slide_ids.remove(slide_id)
                slide_ids.insert(index, slide_id)
            selector = {
                "source_slide": source_number,
                "slide": index + 1,
                "previous_index": previous_index,
                "index": index,
            }
        elif operation == "slide.delete":
            if set(patch) - {"operation", "index"}:
                raise ValueError("slide.delete accepts only zero-based index")
            index = _integer(patch.get("index"), "index", 0, len(document.slides) - 1)
            slide_id = document.slides._sldIdLst[index]
            relationship_id = slide_id.rId
            document.slides._sldIdLst.remove(slide_id)
            document.part.drop_rel(relationship_id)
            selector = {"index": index, "slide": index + 1}
        elif operation == "shape.group":
            if set(patch) - {"operation", "slide", "shape_ids"}:
                raise ValueError("shape.group accepts only slide and shape_ids")
            slide_number = _integer(patch.get("slide"), "slide", 1, len(document.slides))
            slide = document.slides[slide_number - 1]
            shape_ids = patch.get("shape_ids")
            if not isinstance(shape_ids, list) or not 2 <= len(shape_ids) <= 100:
                raise ValueError("shape.group requires 2–100 shape_ids")
            normalized_ids = [_integer(value, "shape_ids item", 1, 2**31 - 1) for value in shape_ids]
            if len(set(normalized_ids)) != len(normalized_ids):
                raise ValueError("shape.group shape_ids must be unique")
            targets = [_presentation_shape_target(slide, shape_id) for shape_id in normalized_ids]
            parent_paths = {parent_shape_ids for _shape, parent_shape_ids in targets}
            if len(parent_paths) != 1:
                raise ValueError("shape.group requires sibling objects in the same shape tree")
            parent_shape_ids = next(iter(parent_paths))
            parent_shapes = _presentation_shape_tree(slide, parent_shape_ids)
            sibling_order = {shape.shape_id: index for index, shape in enumerate(parent_shapes)}
            selected_positions = sorted(sibling_order[shape.shape_id] for shape, _parent in targets)
            if selected_positions != list(range(selected_positions[0], selected_positions[-1] + 1)):
                raise ValueError("shape.group requires contiguous sibling objects so unrelated z-order is preserved")
            ordered_shapes = sorted((shape for shape, _parent in targets), key=lambda shape: sibling_order[shape.shape_id])
            first_xml_index = ordered_shapes[0]._element.getparent().index(ordered_shapes[0]._element)
            next_shape_id = max(shape.shape_id for shape, _parent in _iter_presentation_shapes(slide.shapes)) + 1
            group = parent_shapes.add_group_shape(ordered_shapes)
            group._element.nvGrpSpPr.cNvPr.id = next_shape_id
            parent_xml = group._element.getparent()
            parent_xml.remove(group._element)
            parent_xml.insert(first_xml_index, group._element)
            selector = _presentation_shape_selector(
                slide_number,
                group,
                parent_shape_ids,
                grouped_shape_ids=[shape.shape_id for shape in ordered_shapes],
            )
        elif operation == "shape.ungroup":
            if set(patch) - {"operation", "slide", "shape_id"}:
                raise ValueError("shape.ungroup accepts only slide and shape_id")
            slide_number = _integer(patch.get("slide"), "slide", 1, len(document.slides))
            slide = document.slides[slide_number - 1]
            group, parent_shape_ids = _presentation_shape_target(slide, patch.get("shape_id"))
            if group._element.tag.rsplit("}", 1)[-1] != "grpSp":
                raise ValueError("shape_id must identify a PowerPoint group")
            members = _flatten_presentation_group_members(group)
            if not members:
                raise ValueError("Cannot ungroup an empty PowerPoint group")
            group_id = group.shape_id
            group_path = [*parent_shape_ids, group_id]
            for member in members:
                group._element.addprevious(member._element)
            group._element.getparent().remove(group._element)
            selector = {
                "slide": slide_number,
                "shape_id": group_id,
                "group_path": group_path,
                "ungrouped_shape_ids": [member.shape_id for member in members],
            }
        elif operation == "shape.reorder":
            if set(patch) - {"operation", "slide", "shape_id", "z_index"}:
                raise ValueError("shape.reorder accepts only slide, shape_id and z_index")
            slide_number = _integer(patch.get("slide"), "slide", 1, len(document.slides))
            slide = document.slides[slide_number - 1]
            shape, parent_shape_ids = _presentation_shape_target(slide, patch.get("shape_id"))
            siblings = list(_presentation_shape_tree(slide, parent_shape_ids))
            previous_z_index = next(
                index for index, sibling in enumerate(siblings) if sibling._element is shape._element
            )
            z_index = _integer(patch.get("z_index"), "z_index", 0, len(siblings) - 1)
            if z_index != previous_z_index:
                remaining = [sibling for sibling in siblings if sibling._element is not shape._element]
                parent = shape._element.getparent()
                parent.remove(shape._element)
                if z_index == len(remaining):
                    anchor = remaining[-1]._element
                    parent.insert(parent.index(anchor) + 1, shape._element)
                else:
                    anchor = remaining[z_index]._element
                    parent.insert(parent.index(anchor), shape._element)
            selector = _presentation_shape_selector(
                slide_number,
                shape,
                parent_shape_ids,
                previous_z_index=previous_z_index,
                z_index=z_index,
                sibling_count=len(siblings),
            )
        elif operation in {"shape.delete", "picture.delete", "chart.delete", "table.delete"}:
            if set(patch) - {"operation", "slide", "shape_id"}:
                raise ValueError(f"{operation} accepts only slide and shape_id")
            slide_number = _integer(patch.get("slide"), "slide", 1, len(document.slides))
            slide = document.slides[slide_number - 1]
            shape, parent_shape_ids = _presentation_shape_target(slide, patch.get("shape_id"))
            if operation == "picture.delete" and shape._element.tag.rsplit("}", 1)[-1] != "pic":
                raise ValueError("shape_id must identify a PowerPoint picture")
            if operation == "chart.delete" and not getattr(shape, "has_chart", False):
                raise ValueError("shape_id must identify a PowerPoint chart")
            if operation == "table.delete" and not getattr(shape, "has_table", False):
                raise ValueError("shape_id must identify a PowerPoint table")
            if parent_shape_ids:
                parent, _ = _presentation_shape_target(slide, parent_shape_ids[-1])
                if len(parent.shapes) <= 1:
                    raise ValueError("Cannot delete the only member of a PowerPoint group; delete the group instead")
            relationship_ids = _office_relationship_ids(shape._element)
            shape._element.getparent().remove(shape._element)
            _drop_unused_office_relationships(slide.part, slide._element, relationship_ids)
            selector = _presentation_shape_selector(slide_number, shape, parent_shape_ids)
        elif operation == "shape.format":
            if set(patch) - {"operation", "slide", "shape_id", "format"}:
                raise ValueError("shape.format targets only slide, shape_id and format")
            slide_number = _integer(patch.get("slide"), "slide", 1, len(document.slides))
            shape, parent_shape_ids = _presentation_shape_target(
                document.slides[slide_number - 1], patch.get("shape_id"),
            )
            _format_presentation_shape(shape, patch.get("format"))
            selector = _presentation_shape_selector(slide_number, shape, parent_shape_ids)
        elif operation == "table.format":
            if set(patch) - {"operation", "slide", "shape_id", "format"}:
                raise ValueError("PowerPoint table.format targets only slide, shape_id and format")
            slide_number = _integer(patch.get("slide"), "slide", 1, len(document.slides))
            shape, parent_shape_ids = _presentation_shape_target(
                document.slides[slide_number - 1], patch.get("shape_id"),
            )
            if not shape.has_table:
                raise ValueError("shape_id must identify a PowerPoint table")
            _format_presentation_table(shape.table, patch.get("format"))
            selector = _presentation_shape_selector(slide_number, shape, parent_shape_ids)
        elif operation in {"chart.data", "chart.format"}:
            allowed = (
                {"operation", "slide", "shape_id", "categories", "series"}
                if operation == "chart.data"
                else {"operation", "slide", "shape_id", "format"}
            )
            if set(patch) - allowed:
                raise ValueError(f"{operation} contains unsupported properties")
            slide_number = _integer(patch.get("slide"), "slide", 1, len(document.slides))
            slide = document.slides[slide_number - 1]
            chart_shape, parent_shape_ids = _presentation_shape_target(slide, patch.get("shape_id"))
            if not getattr(chart_shape, "has_chart", False):
                raise ValueError("shape_id must identify a PowerPoint chart")
            if operation == "chart.data":
                chart_type_name = _presentation_chart_type_name_for_chart(chart_shape.chart)
                if chart_type_name is None:
                    raise ValueError(
                        "chart.data supports only chart types listed by inspect_file_engine; recreate or preserve this chart",
                    )
                if _is_stock_chart_type(chart_type_name):
                    with _presentation_stock_chart_as_line(chart_shape.chart):
                        chart_shape.chart.replace_data(_presentation_chart_data(patch, chart_type_name))
                else:
                    chart_shape.chart.replace_data(_presentation_chart_data(patch, chart_type_name))
                    if _is_combo_chart_type(chart_type_name):
                        _make_presentation_combo_chart(chart_shape.chart)
            else:
                chart_type_name = _presentation_chart_type_name_for_chart(chart_shape.chart)
                if _is_stock_chart_type(chart_type_name):
                    style = patch.get("format")
                    if isinstance(style, dict) and {"vary_colors", "trendlines", "error_bars"} & style.keys():
                        raise ValueError("stock charts do not support vary_colors, trendlines or error bars")
                    with _presentation_stock_chart_as_line(chart_shape.chart):
                        _format_presentation_chart(chart_shape.chart, style)
                else:
                    _format_presentation_chart(chart_shape.chart, patch.get("format"))
            selector = _presentation_shape_selector(slide_number, chart_shape, parent_shape_ids)
        elif operation in {"picture.format", "picture.replace"}:
            allowed = {"operation", "slide", "shape_id", "format"} if operation == "picture.format" else {"operation", "slide", "shape_id", "source"}
            if set(patch) - allowed:
                raise ValueError(f"{operation} contains unsupported properties")
            slide_number = _integer(patch.get("slide"), "slide", 1, len(document.slides))
            slide = document.slides[slide_number - 1]
            picture, parent_shape_ids = _presentation_shape_target(slide, patch.get("shape_id"))
            if picture._element.tag.rsplit("}", 1)[-1] != "pic":
                raise ValueError("shape_id must identify a PowerPoint picture")
            if operation == "picture.format":
                _format_presentation_picture(picture, patch.get("format"))
            else:
                data = _office_picture_resource(patch, resources)
                old_relationship = picture._element.blipFill.blip.rEmbed
                _image_part, relationship = slide.part.get_or_add_image_part(io.BytesIO(data))
                picture._element.blipFill.blip.rEmbed = relationship
                if old_relationship != relationship:
                    _drop_unused_office_relationships(slide.part, slide._element, {old_relationship})
            selector = _presentation_shape_selector(slide_number, picture, parent_shape_ids)
        elif operation in {"paragraph.insert", "paragraph.delete"}:
            allowed = (
                {"operation", "slide", "shape_id", "index", "text", "format"}
                if operation == "paragraph.insert"
                else {"operation", "slide", "shape_id", "index"}
            )
            if set(patch) - allowed:
                raise ValueError(f"PowerPoint {operation} contains unsupported properties")
            slide_number = _integer(patch.get("slide"), "slide", 1, len(document.slides))
            slide = document.slides[slide_number - 1]
            shape, parent_shape_ids = _presentation_shape_target(slide, patch.get("shape_id"))
            if not shape.has_text_frame:
                raise ValueError("shape_id must identify a PowerPoint text frame")
            paragraphs = shape.text_frame.paragraphs
            if operation == "paragraph.insert":
                index = _integer(patch.get("index", len(paragraphs)), "index", 0, len(paragraphs))
                _insert_presentation_paragraph(shape, index, patch.get("text"), patch.get("format"))
                selector = _presentation_shape_selector(slide_number, shape, parent_shape_ids, index=index)
            elif len(paragraphs) <= 1:
                raise ValueError("Cannot delete the only paragraph in a PowerPoint text frame; use text.set with empty text")
            else:
                index = _integer(patch.get("index"), "index", 0, len(paragraphs) - 1)
                paragraph = paragraphs[index]
                relationship_ids = _office_relationship_ids(paragraph._p)
                paragraph._p.getparent().remove(paragraph._p)
                _drop_unused_office_relationships(slide.part, slide._element, relationship_ids)
                selector = _presentation_shape_selector(slide_number, shape, parent_shape_ids, index=index)
        elif operation in {"text.set", "paragraph.format"}:
            if any(key in patch for key in ("table_index", "cell", "sheet", "table", "section_index")):
                raise ValueError(f"PowerPoint {operation} targets slide, shape_id and paragraph index")
            slide_number = _integer(patch.get("slide"), "slide", 1, len(document.slides))
            shape, parent_shape_ids = _presentation_shape_target(
                document.slides[slide_number - 1], patch.get("shape_id"),
            )
            if not shape.has_text_frame:
                raise ValueError("shape_id must identify a PowerPoint text frame")
            index = _integer(patch.get("index"), "index", 0, len(shape.text_frame.paragraphs) - 1)
            if operation == "text.set":
                _set_office_paragraph_text(shape.text_frame.paragraphs[index], patch.get("text"), word=False)
            else:
                _format_office_paragraph(shape.text_frame.paragraphs[index], patch, word=False)
            selector = _presentation_shape_selector(slide_number, shape, parent_shape_ids, index=index)
        elif operation in {"set_cell", "cell.format"}:
            if any(key in patch for key in ("sheet", "table_index", "table", "inherit_from_row")):
                raise ValueError("PowerPoint cell edits use slide, shape_id and cell")
            slide_number = _integer(patch.get("slide"), "slide", 1, len(document.slides))
            shape, parent_shape_ids = _presentation_shape_target(
                document.slides[slide_number - 1], patch.get("shape_id"),
            )
            if not shape.has_table:
                raise ValueError("shape_id must identify a PowerPoint table")
            cell = _office_table_cell(shape.table, patch.get("cell"), word=False)
            if operation == "cell.format":
                _format_office_table_cell(cell, patch.get("format"), word=False)
            else:
                if "value" not in patch:
                    raise ValueError("value is required for cell.set")
                value = spreadsheet_cell_value(patch["value"])
                _set_office_table_cell_text(cell, "" if value is None else str(value), word=False)
            selector = _presentation_shape_selector(
                slide_number, shape, parent_shape_ids, cell=patch["cell"].upper(),
            )
        else:
            slide_number = _integer(patch.get("slide"), "slide", 1, len(document.slides))
            shapes = document.slides[slide_number - 1].shapes
            transform = patch.get("transform")
            allowed = {"x", "y", "width", "height", "rotation", "flip_horizontal", "flip_vertical"}
            if not isinstance(transform, dict) or not transform or set(transform) - allowed:
                raise ValueError("transform accepts x, y, width, height (points), rotation (degrees), and flip booleans")
            flips = {}
            for key in {"flip_horizontal", "flip_vertical"} & transform.keys():
                if type(transform[key]) is not bool:
                    raise ValueError(f"{key} must be boolean")
                flips[key] = transform[key]
            is_line = operation == "shape.insert" and patch.get("preset") == PresentationShapePreset.LINE.value
            geometry = {key: _number(value, key, (0 if is_line else 0.01) if key in {"width", "height"} else -100000, 100000)
                        for key, value in transform.items() if key not in flips}
            if operation in {"textbox.insert", "table.insert", "shape.insert", "picture.insert", "chart.insert"}:
                if not {"x", "y", "width", "height"} <= geometry.keys():
                    raise ValueError(f"{operation} requires x, y, width and height")
                if flips:
                    raise ValueError(f"{operation} does not accept flip booleans; insert first, then shape.transform")
                if operation == "picture.insert":
                    if set(patch) - {"operation", "slide", "source", "transform", "fit", "format"}:
                        raise ValueError("picture.insert accepts slide, source, transform, fit and format")
                    data = _office_picture_resource(patch, resources)
                    geometry, fit_crop = _picture_fit_geometry(data, geometry, patch.get("fit"))
                    dimensions = [Emu(round(geometry[key] * 12700)) for key in ("x", "y", "width", "height")]
                    shape = shapes.add_picture(io.BytesIO(data), *dimensions)
                    shape._element.nvPicPr.cNvPr.set(
                        "descr", str(patch["source"]["path"]).replace("\\", "/").rsplit("/", 1)[-1],
                    )
                    if fit_crop:
                        _format_presentation_picture(shape, {"crop": fit_crop})
                    if "format" in patch:
                        _format_presentation_picture(shape, patch["format"])
                elif operation == "chart.insert":
                    if set(patch) - {
                        "operation", "slide", "chart_type", "categories", "series", "transform", "format",
                    }:
                        raise ValueError(
                            "chart.insert accepts slide, chart_type, categories, series, transform and format",
                        )
                    chart_type = patch.get("chart_type")
                    chart_types = _presentation_chart_types()
                    if chart_type not in OfficeChartType.values():
                        raise ValueError(
                            f"chart_type must be one of {', '.join(OfficeChartType.values())}",
                        )
                    dimensions = [Emu(round(geometry[key] * 12700)) for key in ("x", "y", "width", "height")]
                    native_chart_type = (
                        chart_types[OfficeChartType.COLUMN.value]
                        if _is_combo_chart_type(chart_type) or _is_volume_stock_chart_type(chart_type)
                        else chart_types[OfficeChartType.LINE.value]
                        if _is_stock_chart_type(chart_type)
                        else chart_types[chart_type]
                    )
                    data = _presentation_chart_data(patch, chart_type)
                    shape = shapes.add_chart(native_chart_type, *dimensions, data)
                    if _is_combo_chart_type(chart_type):
                        _make_presentation_combo_chart(shape.chart)
                    if _is_stock_chart_type(chart_type):
                        _make_presentation_stock_chart(shape.chart, chart_type)
                        if "format" in patch:
                            style = patch["format"]
                            if isinstance(style, dict) and {"vary_colors", "trendlines", "error_bars"} & style.keys():
                                raise ValueError("stock charts do not support vary_colors, trendlines or error bars")
                            with _presentation_stock_chart_as_line(shape.chart):
                                _format_presentation_chart(shape.chart, style)
                    elif "format" in patch:
                        _format_presentation_chart(shape.chart, patch["format"])
                elif operation == "shape.insert":
                    dimensions = [Emu(round(geometry[key] * 12700)) for key in ("x", "y", "width", "height")]
                    if set(patch) - {"operation", "slide", "preset", "transform", "text", "format"}:
                        raise ValueError("shape.insert accepts slide, preset, transform, text and format")
                    preset = patch.get("preset")
                    if preset not in PresentationShapePreset.values():
                        raise ValueError(f"preset must be one of {', '.join(PresentationShapePreset.values())}")
                    if is_line:
                        from pptx.enum.shapes import MSO_CONNECTOR
                        if "text" in patch or (dimensions[2] == 0 and dimensions[3] == 0):
                            raise ValueError("line requires a nonzero extent and cannot contain text")
                        shape = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, dimensions[0], dimensions[1], dimensions[0] + dimensions[2], dimensions[1] + dimensions[3])
                    else:
                        from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
                        shape = shapes.add_shape(MSO_AUTO_SHAPE_TYPE.from_xml(preset), *dimensions)
                        if "text" in patch:
                            _set_office_paragraph_text(shape.text_frame.paragraphs[0], patch["text"], word=False)
                    if "format" in patch:
                        _format_presentation_shape(shape, patch["format"])
                elif operation == "table.insert":
                    dimensions = [Emu(round(geometry[key] * 12700)) for key in ("x", "y", "width", "height")]
                    if patch.get("style") is not None:
                        raise ValueError("PowerPoint table.insert does not accept Word style names")
                    values = _office_table_rows(patch.get("rows"))
                    shape = shapes.add_table(len(values), len(values[0]), *dimensions)
                    for row_index, row in enumerate(values):
                        for column_index, value in enumerate(row):
                            shape.table.cell(row_index, column_index).text = "" if value is None else str(value)
                else:
                    dimensions = [Emu(round(geometry[key] * 12700)) for key in ("x", "y", "width", "height")]
                    text = patch.get("text")
                    if not isinstance(text, str):
                        raise ValueError("text must be a string")
                    shape = shapes.add_textbox(*dimensions)
                    shape.text_frame.text = text
            elif operation == "shape.transform":
                shape, parent_shape_ids = _presentation_shape_target(
                    document.slides[slide_number - 1], patch.get("shape_id"),
                )
            else:
                raise ValueError("Unsupported PowerPoint structure operation")
            for key, number in geometry.items():
                if key == "rotation":
                    shape.rotation = number % 360
                else:
                    setattr(shape, {"x": "left", "y": "top"}.get(key, key), Emu(round(number * 12700)))
            for key, value in flips.items():
                setattr(shape._element.xfrm, {"flip_horizontal": "flipH", "flip_vertical": "flipV"}[key], value)
            selector = _presentation_shape_selector(
                slide_number, shape, parent_shape_ids if operation == "shape.transform" else (),
            )
    output = io.BytesIO()
    document.save(output)
    return {"operation": patch["operation"], **selector, "_persisted_bytes": output.getvalue()}


_SPREADSHEET_CHART_DATA_MARKER = "__MANOR_NATIVE_CHART_DATA_V1__"


def hydrate_spreadsheet_chart_styles(abs_path: str, workbook: Any) -> None:
    """Restore chart-space styles that openpyxl reads but does not map to ChartBase."""
    import zipfile

    from openpyxl.drawing.spreadsheet_drawing import SpreadsheetDrawing
    from openpyxl.packaging.relationship import get_dependents, get_rels_path
    from openpyxl.xml.functions import fromstring

    drawing_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing"
    chart_namespace = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
    with zipfile.ZipFile(abs_path) as archive:
        names = set(archive.namelist())
        for worksheet in workbook.worksheets:
            drawing_relationship = next((rel for rel in worksheet._rels if rel.Type == drawing_type), None)
            if drawing_relationship is None or not isinstance(drawing_relationship.target, str):
                continue
            drawing_path = drawing_relationship.target.lstrip("/")
            if drawing_path not in names:
                continue
            try:
                drawing = SpreadsheetDrawing.from_tree(fromstring(archive.read(drawing_path)))
                dependents = get_dependents(archive, get_rels_path(drawing_path))
            except (KeyError, TypeError, ValueError):
                # Style hydration is a fidelity supplement. A drawing that
                # openpyxl could not model must not block an unrelated cell edit.
                continue
            relationship_by_id = {relationship.Id: relationship for relationship in dependents}
            for chart, chart_relationship in zip(worksheet._charts, drawing._chart_rels):
                relationship = relationship_by_id.get(chart_relationship.id)
                chart_path = relationship.target.lstrip("/") if relationship is not None else ""
                if chart_path not in names:
                    continue
                try:
                    root = fromstring(archive.read(chart_path))
                    style = root.find(chart_namespace + "style")
                    if style is not None and style.get("val") is not None:
                        chart.style = _integer(int(style.get("val")), "chart style", 1, 48)
                except (KeyError, TypeError, ValueError):
                    continue


def _spreadsheet_chart_components(chart: Any) -> list[Any]:
    components = list(getattr(chart, "_charts", ()) or ())
    return components if components else [chart]


def _spreadsheet_chart_series(chart: Any) -> list[Any]:
    return [series for component in _spreadsheet_chart_components(chart) for series in component.ser]


def _spreadsheet_chart_type(chart: Any) -> str | None:
    from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, PieChart, ScatterChart, StockChart

    components = _spreadsheet_chart_components(chart)
    if (
        len(components) == 2
        and isinstance(components[0], BarChart)
        and components[0].type == "col"
        and components[0].grouping in {"clustered", "standard"}
        and len(components[0].ser) == 1
        and isinstance(components[1], StockChart)
    ):
        return {
            3: OfficeChartType.STOCK_VHLC.value,
            4: OfficeChartType.STOCK_VOHLC.value,
        }.get(len(components[1].ser))
    if (
        len(components) == 2
        and isinstance(components[0], BarChart)
        and components[0].type == "col"
        and components[0].grouping in {"clustered", "standard"}
        and isinstance(components[1], LineChart)
    ):
        return OfficeChartType.COMBO_COLUMN_LINE.value
    if isinstance(chart, StockChart):
        return {
            3: OfficeChartType.STOCK_HLC.value,
            4: OfficeChartType.STOCK_OHLC.value,
        }.get(len(chart.ser))
    if isinstance(chart, BarChart):
        prefix = "column" if chart.type == "col" else "bar"
        return {
            "clustered": prefix,
            "standard": prefix,
            "stacked": f"{prefix}_stacked",
            "percentStacked": f"{prefix}_stacked_100",
        }.get(chart.grouping)
    if isinstance(chart, LineChart):
        if chart.grouping == "stacked":
            return None
        if chart.grouping == "percentStacked":
            return None
        has_markers = bool(chart.ser) and all(getattr(item.marker, "symbol", None) not in {None, "none"} for item in chart.ser)
        return OfficeChartType.LINE_MARKERS.value if has_markers else OfficeChartType.LINE.value
    if isinstance(chart, AreaChart):
        return {
            "standard": OfficeChartType.AREA.value,
            "stacked": OfficeChartType.AREA_STACKED.value,
            "percentStacked": OfficeChartType.AREA_STACKED_100.value,
        }.get(chart.grouping)
    if isinstance(chart, DoughnutChart):
        return OfficeChartType.DOUGHNUT.value
    if isinstance(chart, PieChart):
        return OfficeChartType.PIE.value
    if isinstance(chart, ScatterChart):
        return {
            "marker": OfficeChartType.SCATTER.value,
            "line": OfficeChartType.SCATTER_LINES.value,
            "lineMarker": OfficeChartType.SCATTER_LINES_MARKERS.value,
            "smooth": OfficeChartType.SCATTER_SMOOTH.value,
            "smoothMarker": OfficeChartType.SCATTER_SMOOTH_MARKERS.value,
        }.get(chart.scatterStyle)
    return None


def _new_spreadsheet_chart(chart_type: Any) -> Any:
    from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, PieChart, ScatterChart, StockChart
    from openpyxl.chart.axis import ChartLines
    from openpyxl.chart.updown_bars import UpDownBars

    if chart_type not in OfficeChartType.values():
        raise ValueError(f"chart_type must be one of {', '.join(OfficeChartType.values())}")
    if _is_volume_stock_chart_type(chart_type):
        chart = BarChart()
        chart.type = "col"
        chart.grouping = "clustered"
        stock_chart = StockChart()
        stock_chart.hiLowLines = ChartLines()
        if chart_type == OfficeChartType.STOCK_VOHLC.value:
            stock_chart.upDownBars = UpDownBars()
        stock_chart.y_axis.axId = 200
        stock_chart.y_axis.axPos = "r"
        stock_chart.y_axis.crosses = "max"
        chart += stock_chart
        return chart
    if chart_type.startswith(("column", "bar")) or _is_combo_chart_type(chart_type):
        chart = BarChart()
        chart.type = "col" if chart_type.startswith("column") or _is_combo_chart_type(chart_type) else "bar"
        chart.grouping = (
            "percentStacked" if chart_type.endswith("_100")
            else "stacked" if chart_type.endswith("_stacked")
            else "clustered"
        )
        if chart.grouping != "clustered":
            chart.overlap = 100
        return chart
    if chart_type in {OfficeChartType.LINE.value, OfficeChartType.LINE_MARKERS.value}:
        return LineChart()
    if chart_type.startswith("area"):
        chart = AreaChart()
        chart.grouping = (
            "percentStacked" if chart_type.endswith("_100")
            else "stacked" if chart_type.endswith("_stacked")
            else "standard"
        )
        return chart
    if chart_type == OfficeChartType.PIE.value:
        return PieChart()
    if _is_scatter_chart_type(chart_type):
        chart = ScatterChart()
        chart.scatterStyle = {
            OfficeChartType.SCATTER.value: "marker",
            OfficeChartType.SCATTER_LINES.value: "line",
            OfficeChartType.SCATTER_LINES_MARKERS.value: "lineMarker",
            OfficeChartType.SCATTER_SMOOTH.value: "smooth",
            OfficeChartType.SCATTER_SMOOTH_MARKERS.value: "smoothMarker",
        }[chart_type]
        return chart
    if _is_stock_chart_type(chart_type):
        chart = StockChart()
        chart.hiLowLines = ChartLines()
        if chart_type == OfficeChartType.STOCK_OHLC.value:
            chart.upDownBars = UpDownBars()
        return chart
    return DoughnutChart()


def _spreadsheet_chart_data_sheet(workbook: Any) -> Any:
    for worksheet in workbook.worksheets:
        if worksheet.title.startswith("_manor_chart_data") and worksheet["A1"].value == _SPREADSHEET_CHART_DATA_MARKER:
            worksheet.sheet_state = "veryHidden"
            return worksheet
    base, name, suffix = "_manor_chart_data", "_manor_chart_data", 1
    while name in workbook.sheetnames:
        name = f"{base}_{suffix}"
        suffix += 1
    worksheet = workbook.create_sheet(name)
    worksheet["A1"] = _SPREADSHEET_CHART_DATA_MARKER
    worksheet.sheet_state = "veryHidden"
    return worksheet


def _spreadsheet_managed_chart_block(workbook: Any, chart: Any) -> tuple[Any, int, int, int] | None:
    from openpyxl.utils.cell import range_to_tuple

    chart_series = _spreadsheet_chart_series(chart)
    if not chart_series:
        return None

    def formula(source: Any, *names: str) -> str | None:
        for name in names:
            value = getattr(getattr(source, name, None), "f", None)
            if isinstance(value, str) and value:
                return value
        return None

    category_formula = formula(
        getattr(chart_series[0], "cat", None) or getattr(chart_series[0], "xVal", None),
        "strRef", "numRef",
    )
    try:
        sheet_name, (category_column, first_row, category_max_column, last_row) = range_to_tuple(category_formula)
        worksheet = workbook[sheet_name]
    except (KeyError, TypeError, ValueError):
        return None
    if (
        category_column != 1 or category_max_column != 1 or first_row < 2
        or not worksheet.title.startswith("_manor_chart_data")
        or worksheet["A1"].value != _SPREADSHEET_CHART_DATA_MARKER
    ):
        return None
    header_row = first_row - 1
    for series_index, series in enumerate(chart_series, 2):
        try:
            values_sheet, values_range = range_to_tuple(formula(
                getattr(series, "val", None) or getattr(series, "yVal", None), "numRef",
            ))
            name_sheet, name_range = range_to_tuple(formula(getattr(series, "tx", None), "strRef"))
        except (TypeError, ValueError):
            return None
        if values_sheet != sheet_name or name_sheet != sheet_name:
            return None
        if values_range != (series_index, first_row, series_index, last_row):
            return None
        if name_range != (series_index, header_row, series_index, header_row):
            return None
    return worksheet, header_row, last_row, len(chart_series) + 1


def _spreadsheet_chart_data_references(
    workbook: Any,
    patch: dict[str, Any],
    chart_type: str,
    *,
    replace_chart: Any | None = None,
):
    from openpyxl.chart import Reference

    labels, series = _normalized_chart_data(
        patch,
        single_series=chart_type in {OfficeChartType.PIE.value, OfficeChartType.DOUGHNUT.value},
        numeric_categories=_is_scatter_chart_type(chart_type),
        minimum_series=2 if _is_combo_chart_type(chart_type) else 1,
        exact_series=_stock_chart_series_count(chart_type),
    )
    worksheet = _spreadsheet_chart_data_sheet(workbook)
    old_block = _spreadsheet_managed_chart_block(workbook, replace_chart) if replace_chart is not None else None
    start_row = max(3, worksheet.max_row + 2)
    if old_block is not None:
        old_worksheet, old_start, old_end, old_end_column = old_block
        desired_end = old_start + len(labels)
        desired_end_column = len(series) + 1
        occupied = any(
            coordinate in old_worksheet._cells
            and not (old_start <= coordinate[0] <= old_end and 1 <= coordinate[1] <= old_end_column)
            for coordinate in (
                (row, column)
                for row in range(old_start, desired_end + 1)
                for column in range(1, desired_end_column + 1)
            )
        )
        if not occupied:
            worksheet = old_worksheet
            start_row = old_start
        for row in range(old_start, old_end + 1):
            for column in range(1, old_end_column + 1):
                old_worksheet._cells.pop((row, column), None)
    worksheet.cell(start_row, 1, "Category")
    for column, (name, _values) in enumerate(series, 2):
        worksheet.cell(start_row, column, name)
    for row, category in enumerate(labels, start_row + 1):
        worksheet.cell(row, 1, category)
    for column, (_name, values) in enumerate(series, 2):
        for row, value in enumerate(values, start_row + 1):
            worksheet.cell(row, column, value)
    end_row = start_row + len(labels)
    data = Reference(worksheet, min_col=2, max_col=1 + len(series), min_row=start_row, max_row=end_row)
    categories = Reference(worksheet, min_col=1, min_row=start_row + 1, max_row=end_row)
    return data, categories


def _spreadsheet_chart_selector(worksheet: Any, patch: dict[str, Any]) -> tuple[int, Any]:
    if not worksheet._charts:
        raise ValueError("selected worksheet has no native charts")
    index = _integer(patch.get("chart_index"), "chart_index", 0, len(worksheet._charts) - 1)
    return index, worksheet._charts[index]


def _spreadsheet_drawing_anchor(value: Any) -> str:
    from openpyxl.utils.cell import coordinate_to_tuple

    if not isinstance(value, str) or not re.fullmatch(r"[A-Za-z]{1,3}[1-9][0-9]{0,6}", value):
        raise ValueError("drawing anchor must be a single A1 cell")
    row, column = coordinate_to_tuple(value.upper())
    _integer(row, "drawing anchor row", 1, EXCEL_MAX_ROW)
    _integer(column, "drawing anchor column", 1, EXCEL_MAX_COLUMN)
    return value.upper()


def _spreadsheet_drawing_anchor_cell(drawing: Any) -> str | None:
    from openpyxl.utils.cell import get_column_letter

    anchor = drawing.anchor
    if isinstance(anchor, str):
        return _spreadsheet_drawing_anchor(anchor)
    marker = getattr(anchor, "_from", None)
    if marker is None:
        return None
    return f"{get_column_letter(marker.col + 1)}{marker.row + 1}"


def _move_spreadsheet_drawing(drawing: Any, value: Any) -> None:
    """Move a loaded chart or picture without replacing its native size anchor."""
    from openpyxl.utils.cell import coordinate_to_tuple

    anchor_cell = _spreadsheet_drawing_anchor(value)
    anchor = drawing.anchor
    if isinstance(anchor, str):
        drawing.anchor = anchor_cell
        return
    if getattr(anchor, "_from", None) is None:
        raise ValueError("absolute drawing anchors cannot be moved to a cell without explicit resizing")
    row, column = coordinate_to_tuple(anchor_cell)
    row_delta = row - 1 - anchor._from.row
    column_delta = column - 1 - anchor._from.col
    anchor._from.row += row_delta
    anchor._from.col += column_delta
    target = getattr(anchor, "to", None)
    if target is not None:
        target.row += row_delta
        target.col += column_delta
        if not 0 <= target.row < EXCEL_MAX_ROW or not 0 <= target.col < EXCEL_MAX_COLUMN:
            raise ValueError("moving this drawing would place its far edge outside the Excel grid")


def _spreadsheet_picture_selector(worksheet: Any, patch: dict[str, Any]) -> tuple[int, Any]:
    if not worksheet._images:
        raise ValueError("selected worksheet has no native pictures")
    index = _integer(patch.get("picture_index"), "picture_index", 0, len(worksheet._images) - 1)
    return index, worksheet._images[index]


def _spreadsheet_picture_frame(image: Any) -> Any:
    anchor = image.anchor
    child = getattr(anchor, "pic", None)
    if child is None:
        group = getattr(anchor, "groupShape", None)
        child = getattr(group, "pic", None)
    if child is None:
        raise ValueError("Excel picture has an unsupported drawing frame")
    return child


def _spreadsheet_picture_size(image: Any) -> tuple[float, float] | None:
    anchor = image.anchor
    extent = getattr(anchor, "ext", None)
    if extent is not None and int(extent.cx) > 0 and int(extent.cy) > 0:
        return int(extent.cx) / 12700, int(extent.cy) / 12700
    if isinstance(anchor, str) and image.width > 0 and image.height > 0:
        return float(image.width) * 72 / 96, float(image.height) * 72 / 96
    return None


def _spreadsheet_picture_one_cell_anchor(
    cell: str, width: float, height: float, *, picture: Any | None = None,
) -> Any:
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.utils.cell import coordinate_to_tuple

    row, column = coordinate_to_tuple(_spreadsheet_drawing_anchor(cell))
    return OneCellAnchor(
        _from=AnchorMarker(col=column - 1, row=row - 1),
        ext=XDRPositiveSize2D(cx=round(width * 12700), cy=round(height * 12700)),
        pic=picture,
    )


def _spreadsheet_new_picture_frame(worksheet: Any, *, name: str, alt_text: str) -> Any:
    from openpyxl.drawing.spreadsheet_drawing import SpreadsheetDrawing

    used_ids = []
    for image in worksheet._images:
        try:
            used_ids.append(int(_spreadsheet_picture_frame(image).nvPicPr.cNvPr.id))
        except (AttributeError, TypeError, ValueError):
            continue
    identifier = max([len(worksheet._charts) + len(worksheet._images), *used_ids], default=0) + 1
    frame = SpreadsheetDrawing()._picture_frame(identifier)
    frame.nvPicPr.cNvPr.name = name
    frame.nvPicPr.cNvPr.descr = alt_text
    return frame


def _format_spreadsheet_picture(image: Any, style: Any) -> None:
    from openpyxl.drawing.spreadsheet_drawing import AbsoluteAnchor, OneCellAnchor, TwoCellAnchor

    allowed = {"anchor", "width", "height", "alt_text"}
    if not isinstance(style, dict) or not style or set(style) - allowed:
        raise ValueError(f"Excel picture.format requires supported format keys: {', '.join(sorted(allowed))}")
    requested = {
        key: _number(style[key], f"picture {key}", 1, 4032)
        for key in ("width", "height") if key in style
    }
    original_anchor = image.anchor
    if "anchor" in style and isinstance(original_anchor, AbsoluteAnchor):
        if set(requested) != {"width", "height"}:
            raise ValueError("moving an absolute Excel picture to a cell requires both width and height")
        image.anchor = _spreadsheet_picture_one_cell_anchor(
            style["anchor"], requested["width"], requested["height"],
            picture=_spreadsheet_picture_frame(image),
        )
        requested = {}
    elif "anchor" in style:
        _move_spreadsheet_drawing(image, style["anchor"])
    if requested:
        original = _spreadsheet_picture_size(image)
        if len(requested) == 1:
            if original is None:
                raise ValueError("resizing a two-cell anchored Excel picture requires both width and height")
            if "width" in requested:
                requested["height"] = requested["width"] * original[1] / original[0]
            else:
                requested["width"] = requested["height"] * original[0] / original[1]
        anchor = image.anchor
        if isinstance(anchor, (OneCellAnchor, AbsoluteAnchor)):
            anchor.ext.width = round(requested["width"] * 12700)
            anchor.ext.height = round(requested["height"] * 12700)
        elif isinstance(anchor, TwoCellAnchor):
            image.anchor = _spreadsheet_picture_one_cell_anchor(
                _spreadsheet_drawing_anchor_cell(image) or "A1",
                requested["width"], requested["height"],
                picture=_spreadsheet_picture_frame(image),
            )
        elif isinstance(anchor, str):
            image.width = requested["width"] * 96 / 72
            image.height = requested["height"] * 96 / 72
        else:
            raise ValueError("Excel picture has an unsupported drawing anchor")
    if "alt_text" in style:
        alt_text = style["alt_text"]
        if not isinstance(alt_text, str) or len(alt_text) > 1000 or "\x00" in alt_text:
            raise ValueError("alt_text must be a string of at most 1000 characters")
        _spreadsheet_picture_frame(image).nvPicPr.cNvPr.descr = alt_text


def insert_spreadsheet_picture(
    worksheet: Any,
    patch: dict[str, Any],
    resources: dict[tuple[str, str], bytes] | None,
) -> dict[str, Any]:
    from openpyxl.drawing.image import Image

    if set(patch) - {"operation", "sheet", "source", "anchor", "transform", "format"}:
        raise ValueError("Excel picture.insert accepts sheet, source, anchor, transform and format")
    anchor = _spreadsheet_drawing_anchor(patch.get("anchor", "A1"))
    transform = patch.get("transform", {})
    if not isinstance(transform, dict) or set(transform) - {"width", "height"}:
        raise ValueError("Excel picture transform accepts only width and height in points")
    data = _office_picture_resource(patch, resources)
    image = Image(io.BytesIO(data))
    source_width, source_height = float(image.width) * 72 / 96, float(image.height) * 72 / 96
    dimensions = {
        key: _number(value, f"picture {key}", 1, 4032)
        for key, value in transform.items()
    }
    if "width" in dimensions and "height" not in dimensions:
        dimensions["height"] = dimensions["width"] * source_height / source_width
    elif "height" in dimensions and "width" not in dimensions:
        dimensions["width"] = dimensions["height"] * source_width / source_height
    width = dimensions.get("width", source_width)
    height = dimensions.get("height", source_height)
    filename = str(patch["source"]["path"]).replace("\\", "/").rsplit("/", 1)[-1]
    image.anchor = _spreadsheet_picture_one_cell_anchor(
        anchor, width, height,
        picture=_spreadsheet_new_picture_frame(worksheet, name=filename[:255] or "Picture", alt_text=filename),
    )
    worksheet.add_image(image)
    if "format" in patch:
        _format_spreadsheet_picture(image, patch["format"])
    return {
        "updated": True,
        "operation": "picture.insert",
        "sheet": worksheet.title,
        "picture_index": len(worksheet._images) - 1,
        "anchor": _spreadsheet_drawing_anchor_cell(image),
    }


def replace_spreadsheet_picture(
    worksheet: Any,
    patch: dict[str, Any],
    resources: dict[tuple[str, str], bytes] | None,
) -> dict[str, Any]:
    from openpyxl.drawing.image import Image

    if set(patch) - {"operation", "sheet", "picture_index", "source"}:
        raise ValueError("Excel picture.replace accepts sheet, picture_index and source")
    index, current = _spreadsheet_picture_selector(worksheet, patch)
    replacement = Image(io.BytesIO(_office_picture_resource(patch, resources)))
    replacement.anchor = copy.deepcopy(current.anchor)
    worksheet._images[index] = replacement
    return {
        "updated": True,
        "operation": "picture.replace",
        "sheet": worksheet.title,
        "picture_index": index,
        "anchor": _spreadsheet_drawing_anchor_cell(replacement),
    }


def format_spreadsheet_picture(worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    if set(patch) - {"operation", "sheet", "picture_index", "format"}:
        raise ValueError("Excel picture.format accepts sheet, picture_index and format")
    index, image = _spreadsheet_picture_selector(worksheet, patch)
    _format_spreadsheet_picture(image, patch.get("format"))
    return {
        "updated": True,
        "operation": "picture.format",
        "sheet": worksheet.title,
        "picture_index": index,
        "anchor": _spreadsheet_drawing_anchor_cell(image),
    }


def delete_spreadsheet_picture(worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    if set(patch) - {"operation", "sheet", "picture_index"}:
        raise ValueError("Excel picture.delete accepts sheet and picture_index")
    index, image = _spreadsheet_picture_selector(worksheet, patch)
    anchor = _spreadsheet_drawing_anchor_cell(image)
    worksheet._images.pop(index)
    return {
        "updated": True,
        "operation": "picture.delete",
        "sheet": worksheet.title,
        "picture_index": index,
        "anchor": anchor,
    }


def _spreadsheet_chart_point_count(series: Any) -> int:
    """Read the native value count whether the series uses a cache or worksheet reference."""
    from openpyxl.utils.cell import range_boundaries

    values = getattr(series, "val", None) or getattr(series, "yVal", None)
    reference = getattr(values, "numRef", None)
    literal = getattr(values, "numLit", None)
    cache = getattr(reference, "numCache", None) or literal
    if cache is not None:
        count = getattr(getattr(cache, "ptCount", None), "val", None)
        if isinstance(count, int):
            return count
        return len(getattr(cache, "pt", ()) or ())
    formula = getattr(reference, "f", None)
    if isinstance(formula, str) and formula:
        cell_range = formula.rsplit("!", 1)[-1].replace("$", "")
        min_column, min_row, max_column, max_row = range_boundaries(cell_range)
        return (max_column - min_column + 1) * (max_row - min_row + 1)
    return 0


def _resize_spreadsheet_chart(chart: Any, style: dict[str, Any]) -> None:
    from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.utils.units import cm_to_EMU

    requested = {
        key: _number(style[key], f"chart {key}", 20, 4032)
        for key in ("width", "height") if key in style
    }
    if not requested:
        return
    anchor = chart.anchor
    if isinstance(anchor, str):
        for key, points in requested.items():
            setattr(chart, key, points / 72 * 2.54)
        return
    if isinstance(anchor, OneCellAnchor):
        if "width" in requested:
            anchor.ext.width = cm_to_EMU(requested["width"] / 72 * 2.54)
        if "height" in requested:
            anchor.ext.height = cm_to_EMU(requested["height"] / 72 * 2.54)
        return
    if set(requested) != {"width", "height"}:
        raise ValueError("resizing a two-cell anchored template chart requires both width and height")
    chart.anchor = OneCellAnchor(
        _from=copy.deepcopy(anchor._from),
        ext=XDRPositiveSize2D(
            cx=cm_to_EMU(requested["width"] / 72 * 2.54),
            cy=cm_to_EMU(requested["height"] / 72 * 2.54),
        ),
    )


def _format_spreadsheet_chart(chart: Any, style: Any) -> None:
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.legend import Legend
    from openpyxl.chart.marker import DataPoint
    from openpyxl.chart.axis import ChartLines

    common = {
        "title", "style", "has_legend", "legend_position", "legend_include_in_layout",
        "vary_colors", "show_data_labels", "show_category_name", "show_series_name",
        "show_value", "show_percentage", "data_label_position", "category_axis_title",
        "value_axis_title", "show_category_axis", "show_value_axis", "category_axis_min",
        "category_axis_max", "category_axis_major_unit", "value_axis_min", "value_axis_max",
        "value_axis_major_unit", "show_major_gridlines", "series_colors", "category_colors",
        "anchor", "width", "height", "trendlines", "error_bars",
        "secondary_value_axis_title", "show_secondary_value_axis",
        "secondary_value_axis_min", "secondary_value_axis_max", "secondary_value_axis_major_unit",
        "show_secondary_major_gridlines",
    }
    if not isinstance(style, dict) or not style or set(style) - common:
        raise ValueError(f"chart.format requires supported format keys: {', '.join(sorted(common))}")
    if _is_stock_chart_type(_spreadsheet_chart_type(chart)) and "vary_colors" in style:
        raise ValueError("stock charts do not support vary_colors")
    chart_components = _spreadsheet_chart_components(chart)
    chart_series = _spreadsheet_chart_series(chart)

    def text(value: Any, name: str) -> str | None:
        if value is not None and (not isinstance(value, str) or len(value) > 500 or "\x00" in value):
            raise ValueError(f"{name} must be null or a string of at most 500 characters")
        return value

    def boolean(value: Any, name: str) -> bool:
        if type(value) is not bool:
            raise ValueError(f"{name} must be boolean")
        return value

    if "title" in style:
        chart.title = text(style["title"], "title")
    if "style" in style:
        chart.style = _integer(style["style"], "chart style", 1, 48)
    if "has_legend" in style:
        chart.legend = Legend() if boolean(style["has_legend"], "has_legend") and chart.legend is None else (
            chart.legend if style["has_legend"] else None
        )
    if "legend_position" in style:
        positions = {"top": "t", "bottom": "b", "left": "l", "right": "r"}
        if style["legend_position"] not in positions:
            raise ValueError("legend_position must be top, bottom, left or right")
        chart.legend = chart.legend or Legend()
        chart.legend.position = positions[style["legend_position"]]
    if "legend_include_in_layout" in style:
        chart.legend = chart.legend or Legend()
        chart.legend.overlay = not boolean(style["legend_include_in_layout"], "legend_include_in_layout")
    if "vary_colors" in style:
        vary_colors = boolean(style["vary_colors"], "vary_colors")
        for component in chart_components:
            component.varyColors = vary_colors

    label_keys = {"show_category_name", "show_series_name", "show_value", "show_percentage", "data_label_position"}
    if "show_data_labels" in style:
        show_data_labels = boolean(style["show_data_labels"], "show_data_labels")
        for component in chart_components:
            component.dLbls = DataLabelList() if show_data_labels else None
    if label_keys & style.keys():
        for component in chart_components:
            component.dLbls = component.dLbls or DataLabelList()
        for key, attribute in (
            ("show_category_name", "showCatName"),
            ("show_series_name", "showSerName"),
            ("show_value", "showVal"),
            ("show_percentage", "showPercent"),
        ):
            if key in style:
                label_value = boolean(style[key], key)
                for component in chart_components:
                    setattr(component.dLbls, attribute, label_value)
        if "data_label_position" in style:
            positions = {"best_fit": "bestFit", "center": "ctr", "inside_end": "inEnd", "outside_end": "outEnd"}
            if style["data_label_position"] not in positions:
                raise ValueError("data_label_position must be best_fit, center, inside_end or outside_end")
            for component in chart_components:
                component.dLbls.dLblPos = positions[style["data_label_position"]]

    axes = {"category": getattr(chart, "x_axis", None), "value": getattr(chart, "y_axis", None)}
    for name in ("category", "value"):
        title_key, visible_key = f"{name}_axis_title", f"show_{name}_axis"
        if title_key in style:
            if axes[name] is None:
                raise ValueError(f"this chart type has no {name} axis")
            axes[name].title = text(style[title_key], title_key)
        if visible_key in style:
            if axes[name] is None:
                raise ValueError(f"this chart type has no {name} axis")
            axes[name].delete = not boolean(style[visible_key], visible_key)
    for name in ("category", "value"):
        axis = axes[name]
        scale_values = {}
        for suffix, attribute in (("min", "min"), ("max", "max"), ("major_unit", "majorUnit")):
            key = f"{name}_axis_{suffix}"
            if key not in style:
                continue
            if axis is None or (name == "category" and not _is_scatter_chart_type(_spreadsheet_chart_type(chart))):
                raise ValueError(f"this chart type has no numeric {name} axis")
            value = style[key]
            if value is not None:
                value = _number(value, key, -1e15 if suffix != "major_unit" else 1e-12, 1e15)
            scale_values[attribute] = value
        current_min = axis.scaling.min if axis is not None else None
        current_max = axis.scaling.max if axis is not None else None
        prospective_min = scale_values.get("min", current_min)
        prospective_max = scale_values.get("max", current_max)
        if prospective_min is not None and prospective_max is not None and prospective_min >= prospective_max:
            raise ValueError(f"{name}_axis_min must be less than {name}_axis_max")
        for attribute, value in scale_values.items():
            if attribute == "majorUnit":
                axis.majorUnit = value
            else:
                setattr(axis.scaling, attribute, value)
    value_axis = axes["value"]
    if "show_major_gridlines" in style:
        if value_axis is None:
            raise ValueError("this chart type has no value axis")
        value_axis.majorGridlines = ChartLines() if boolean(style["show_major_gridlines"], "show_major_gridlines") else None

    components = _spreadsheet_chart_components(chart)
    secondary_axis = components[1].y_axis if (
        _is_volume_stock_chart_type(_spreadsheet_chart_type(chart)) and len(components) > 1
    ) else None
    secondary_keys = {
        "secondary_value_axis_title", "show_secondary_value_axis",
        "secondary_value_axis_min", "secondary_value_axis_max", "secondary_value_axis_major_unit",
        "show_secondary_major_gridlines",
    }
    if secondary_keys & style.keys() and secondary_axis is None:
        raise ValueError("this chart type has no secondary value axis")
    if secondary_axis is not None:
        if "secondary_value_axis_title" in style:
            secondary_axis.title = text(style["secondary_value_axis_title"], "secondary_value_axis_title")
        if "show_secondary_value_axis" in style:
            secondary_axis.delete = not boolean(style["show_secondary_value_axis"], "show_secondary_value_axis")
        scale_values = {}
        for suffix, attribute in (("min", "min"), ("max", "max"), ("major_unit", "majorUnit")):
            key = f"secondary_value_axis_{suffix}"
            if key not in style:
                continue
            value = style[key]
            if value is not None:
                value = _number(value, key, -1e15 if suffix != "major_unit" else 1e-12, 1e15)
            scale_values[attribute] = value
        prospective_min = scale_values.get("min", secondary_axis.scaling.min)
        prospective_max = scale_values.get("max", secondary_axis.scaling.max)
        if prospective_min is not None and prospective_max is not None and prospective_min >= prospective_max:
            raise ValueError("secondary_value_axis_min must be less than secondary_value_axis_max")
        for attribute, value in scale_values.items():
            if attribute == "majorUnit":
                secondary_axis.majorUnit = value
            else:
                setattr(secondary_axis.scaling, attribute, value)
        if "show_secondary_major_gridlines" in style:
            secondary_axis.majorGridlines = (
                ChartLines()
                if boolean(style["show_secondary_major_gridlines"], "show_secondary_major_gridlines")
                else None
            )

    if "series_colors" in style:
        colors = style["series_colors"]
        if not isinstance(colors, list) or not colors or len(colors) > min(50, len(chart_series)):
            raise ValueError("series_colors must be a non-empty RGB array no longer than the chart series count")
        for color_value, series in zip(colors, chart_series):
            color = _validated_cell_format({"fill_color": color_value}, spreadsheet=False)["fill_color"]
            series.graphicalProperties.solidFill = color
            series.graphicalProperties.line.solidFill = color
    if "category_colors" in style:
        if _spreadsheet_chart_type(chart) not in {OfficeChartType.PIE.value, OfficeChartType.DOUGHNUT.value}:
            raise ValueError("category_colors is supported only for pie and doughnut charts")
        colors = style["category_colors"]
        category_count = _spreadsheet_chart_point_count(chart.ser[0]) if chart.ser else 0
        if not isinstance(colors, list) or not colors or len(colors) > min(1000, category_count):
            raise ValueError("category_colors must be a non-empty RGB array no longer than the category count")
        points = []
        for index, color_value in enumerate(colors):
            color = _validated_cell_format({"fill_color": color_value}, spreadsheet=False)["fill_color"]
            point = DataPoint(idx=index)
            point.graphicalProperties.solidFill = color
            point.graphicalProperties.line.solidFill = color
            points.append(point)
        chart.ser[0].dPt = points
    if "trendlines" in style:
        chart_type = _spreadsheet_chart_type(chart)
        supported = {
            OfficeChartType.COLUMN.value, OfficeChartType.BAR.value,
            OfficeChartType.LINE.value, OfficeChartType.LINE_MARKERS.value,
            OfficeChartType.SCATTER.value, OfficeChartType.SCATTER_LINES.value,
            OfficeChartType.SCATTER_LINES_MARKERS.value, OfficeChartType.SCATTER_SMOOTH.value,
            OfficeChartType.SCATTER_SMOOTH_MARKERS.value,
            OfficeChartType.COMBO_COLUMN_LINE.value,
        }
        if chart_type not in supported:
            raise ValueError("trendlines require a non-stacked column, bar, line or scatter chart")
        from openpyxl.chart.trendline import Trendline

        native_types = {
            "linear": "linear", "exponential": "exp", "logarithmic": "log",
            "polynomial": "poly", "power": "power", "moving_average": "movingAvg",
        }
        for index, config in _normalized_chart_trendlines(style["trendlines"], len(chart_series)):
            if config is None:
                chart_series[index].trendline = None
                continue
            chart_series[index].trendline = Trendline(
                name=config.get("name"),
                trendlineType=native_types[config["type"]],
                order=config.get("order"),
                period=config.get("period"),
                forward=config.get("forward"),
                backward=config.get("backward"),
                intercept=config.get("intercept"),
                dispRSqr=config.get("display_r_squared"),
                dispEq=config.get("display_equation"),
            )
    if "error_bars" in style:
        chart_type = _spreadsheet_chart_type(chart)
        scatter = _is_scatter_chart_type(chart_type)
        supported = {
            OfficeChartType.COLUMN.value, OfficeChartType.BAR.value,
            OfficeChartType.LINE.value, OfficeChartType.LINE_MARKERS.value,
            OfficeChartType.SCATTER.value, OfficeChartType.SCATTER_LINES.value,
            OfficeChartType.SCATTER_LINES_MARKERS.value, OfficeChartType.SCATTER_SMOOTH.value,
            OfficeChartType.SCATTER_SMOOTH_MARKERS.value,
            OfficeChartType.COMBO_COLUMN_LINE.value,
        }
        if chart_type not in supported:
            raise ValueError("error bars require a non-stacked column, bar, line or scatter chart")
        from openpyxl.chart.error_bar import ErrorBars

        native_types = {
            "fixed": "fixedVal", "percentage": "percentage",
            "standard_deviation": "stdDev", "standard_error": "stdErr",
        }
        for index, config in _normalized_chart_error_bars(
            style["error_bars"], len(chart_series), scatter=scatter,
        ):
            if config is None:
                chart_series[index].errBars = None
                continue
            chart_series[index].errBars = ErrorBars(
                errDir=config["direction"],
                errBarType=config["side"],
                errValType=native_types[config["type"]],
                noEndCap=config["end_style"] == "no_cap",
                val=config.get("value"),
            )
    if "anchor" in style:
        _move_spreadsheet_drawing(chart, style["anchor"])
    _resize_spreadsheet_chart(chart, style)


def _populate_spreadsheet_chart_data(chart: Any, data: Any, categories: Any, chart_type: str) -> None:
    if _is_volume_stock_chart_type(chart_type):
        from openpyxl.chart import Reference, StockChart

        components = _spreadsheet_chart_components(chart)
        stock_chart = next((component for component in components[1:] if isinstance(component, StockChart)), None)
        if stock_chart is None:
            raise ValueError("volume stock chart is missing its native stock plot")
        volume_data = Reference(
            data.worksheet,
            min_col=data.min_col,
            max_col=data.min_col,
            min_row=data.min_row,
            max_row=data.max_row,
        )
        stock_data = Reference(
            data.worksheet,
            min_col=data.min_col + 1,
            max_col=data.max_col,
            min_row=data.min_row,
            max_row=data.max_row,
        )
        chart.add_data(volume_data, titles_from_data=True, from_rows=False)
        chart.set_categories(categories)
        stock_chart.add_data(stock_data, titles_from_data=True, from_rows=False)
        stock_chart.set_categories(categories)
        for series in stock_chart.ser:
            series.marker.symbol = "none"
        return
    if _is_combo_chart_type(chart_type):
        from openpyxl.chart import LineChart, Reference

        if data.max_col - data.min_col + 1 < 2:
            raise ValueError("combo_column_line charts require at least two series")
        components = _spreadsheet_chart_components(chart)
        line_chart = next((component for component in components[1:] if isinstance(component, LineChart)), None)
        if line_chart is None:
            line_chart = LineChart()
            chart += line_chart
        column_data = Reference(
            data.worksheet,
            min_col=data.min_col,
            max_col=data.max_col - 1,
            min_row=data.min_row,
            max_row=data.max_row,
        )
        line_data = Reference(
            data.worksheet,
            min_col=data.max_col,
            max_col=data.max_col,
            min_row=data.min_row,
            max_row=data.max_row,
        )
        chart.add_data(column_data, titles_from_data=True, from_rows=False)
        chart.set_categories(categories)
        line_chart.add_data(line_data, titles_from_data=True, from_rows=False)
        line_chart.set_categories(categories)
        for series in line_chart.ser:
            series.marker.symbol = "circle"
        return
    if not _is_scatter_chart_type(chart_type):
        chart.add_data(data, titles_from_data=True, from_rows=False)
        chart.set_categories(categories)
        if _is_stock_chart_type(chart_type):
            for series in chart.ser:
                series.marker.symbol = "none"
        return
    from openpyxl.chart import Reference, Series

    for column in range(data.min_col, data.max_col + 1):
        values = Reference(
            data.worksheet,
            min_col=column,
            min_row=data.min_row,
            max_row=data.max_row,
        )
        series = Series(values, categories, title_from_data=True)
        if chart_type in {
            OfficeChartType.SCATTER.value,
            OfficeChartType.SCATTER_LINES_MARKERS.value,
            OfficeChartType.SCATTER_SMOOTH_MARKERS.value,
        }:
            series.marker.symbol = "circle"
        else:
            series.marker.symbol = "none"
        series.smooth = chart_type in {
            OfficeChartType.SCATTER_SMOOTH.value,
            OfficeChartType.SCATTER_SMOOTH_MARKERS.value,
        }
        if chart_type == OfficeChartType.SCATTER.value:
            series.graphicalProperties.line.noFill = True
        chart.series.append(series)


def insert_spreadsheet_chart(workbook: Any, worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    allowed = {"operation", "sheet", "chart_type", "categories", "series", "anchor", "transform", "format"}
    if set(patch) - allowed:
        raise ValueError("Excel chart.insert accepts sheet, chart_type, categories, series, anchor, transform and format")
    chart_type = patch.get("chart_type")
    chart = _new_spreadsheet_chart(chart_type)
    data, categories = _spreadsheet_chart_data_references(workbook, patch, chart_type)
    _populate_spreadsheet_chart_data(chart, data, categories, chart_type)
    if chart_type == OfficeChartType.LINE_MARKERS.value:
        for series in chart.ser:
            series.marker.symbol = "circle"
    transform = patch.get("transform")
    if transform is not None:
        if not isinstance(transform, dict) or not transform or set(transform) - {"width", "height"}:
            raise ValueError("Excel chart transform accepts only width and height in points")
        _format_spreadsheet_chart(chart, transform)
    anchor = _spreadsheet_drawing_anchor(patch.get("anchor", "A1"))
    worksheet.add_chart(chart, anchor)
    if "format" in patch:
        _format_spreadsheet_chart(chart, patch["format"])
    return {"operation": patch["operation"], "sheet": worksheet.title, "chart_index": len(worksheet._charts) - 1}


def update_spreadsheet_chart_data(workbook: Any, worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    if set(patch) - {"operation", "sheet", "chart_index", "categories", "series"}:
        raise ValueError("Excel chart.data accepts sheet, chart_index, categories and series")
    index, chart = _spreadsheet_chart_selector(worksheet, patch)
    chart_type = _spreadsheet_chart_type(chart)
    if chart_type is None:
        raise ValueError("chart.data supports only chart types listed by inspect_file_engine")
    data, categories = _spreadsheet_chart_data_references(
        workbook, patch, chart_type, replace_chart=chart,
    )
    previous = _spreadsheet_chart_series(chart)
    for component in _spreadsheet_chart_components(chart):
        component.ser = []
    _populate_spreadsheet_chart_data(chart, data, categories, chart_type)
    updated_series = _spreadsheet_chart_series(chart)
    for old, new in zip(previous, updated_series):
        for attribute in ("graphicalProperties", "marker", "smooth", "dPt", "trendline", "errBars", "invertIfNegative"):
            if hasattr(old, attribute) and hasattr(new, attribute):
                try:
                    setattr(new, attribute, copy.deepcopy(getattr(old, attribute)))
                except (TypeError, ValueError):
                    pass
    if chart_type in {OfficeChartType.LINE_MARKERS.value, OfficeChartType.COMBO_COLUMN_LINE.value}:
        targets = updated_series if chart_type == OfficeChartType.LINE_MARKERS.value else updated_series[-1:]
        for series in targets:
            if series.marker.symbol in {None, "none"}:
                series.marker.symbol = "circle"
    return {"operation": patch["operation"], "sheet": worksheet.title, "chart_index": index}


def format_spreadsheet_chart(worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    if set(patch) - {"operation", "sheet", "chart_index", "format"}:
        raise ValueError("Excel chart.format accepts sheet, chart_index and format")
    index, chart = _spreadsheet_chart_selector(worksheet, patch)
    _format_spreadsheet_chart(chart, patch.get("format"))
    return {"operation": patch["operation"], "sheet": worksheet.title, "chart_index": index}


def delete_spreadsheet_chart(workbook: Any, worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    if set(patch) - {"operation", "sheet", "chart_index"}:
        raise ValueError("Excel chart.delete accepts sheet and chart_index")
    index, chart = _spreadsheet_chart_selector(worksheet, patch)
    managed_block = _spreadsheet_managed_chart_block(workbook, chart)
    worksheet._charts.pop(index)
    if managed_block is not None:
        _clear_spreadsheet_managed_chart_blocks(workbook, [managed_block])
    return {
        "updated": True,
        "operation": "chart.delete",
        "sheet": worksheet.title,
        "chart_index": index,
    }


def delete_spreadsheet_table(worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    if set(patch) - {"operation", "sheet", "table"}:
        raise ValueError("Excel table.delete accepts sheet and table")
    name = patch.get("table")
    if not isinstance(name, str) or not name:
        raise ValueError("table must be a non-empty name")
    table = next((item for item in worksheet.tables.values() if item.name.lower() == name.lower()), None)
    if table is None:
        raise ValueError(f"Table not found in worksheet {worksheet.title}: {name}")
    del worksheet.tables[table.name]
    return {
        "updated": True,
        "operation": "table.delete",
        "sheet": worksheet.title,
        "table": table.name,
        "table_ref": table.ref,
    }


def _spreadsheet_table_selector(worksheet: Any, name: Any) -> Any:
    if not isinstance(name, str) or not name:
        raise ValueError("table must be a non-empty name")
    table = next(
        (item for item in worksheet.tables.values() if item.name.casefold() == name.casefold()),
        None,
    )
    if table is None:
        raise ValueError(f"Table not found in worksheet {worksheet.title}: {name}")
    return table


def _spreadsheet_table_style(value: Any, *, default: bool = False) -> Any:
    from openpyxl.worksheet.table import TableStyleInfo

    if value is None:
        value = {"style_name": "TableStyleMedium2", "show_row_stripes": True} if default else None
    if not isinstance(value, dict) or not value:
        raise ValueError("table format must be a non-empty object")
    allowed = {
        "style_name", "show_first_column", "show_last_column",
        "show_row_stripes", "show_column_stripes",
    }
    if set(value) - allowed:
        raise ValueError("table format contains unsupported properties")
    style_name = value.get("style_name", "TableStyleMedium2" if default else None)
    if style_name is None:
        if set(value) != {"style_name"}:
            raise ValueError("style_name null must be the only table format property")
        return None
    if not isinstance(style_name, str) or not re.fullmatch(
        r"TableStyle(?:Light(?:[1-9]|1[0-9]|2[01])|Medium(?:[1-9]|1[0-9]|2[0-8])|Dark(?:[1-9]|1[01]))",
        style_name,
    ):
        raise ValueError("style_name must be a built-in Excel TableStyleLight1-21, Medium1-28 or Dark1-11")
    values = {}
    for key, attribute in (
        ("show_first_column", "showFirstColumn"),
        ("show_last_column", "showLastColumn"),
        ("show_row_stripes", "showRowStripes"),
        ("show_column_stripes", "showColumnStripes"),
    ):
        if key in value:
            if type(value[key]) is not bool:
                raise ValueError(f"{key} must be boolean")
            values[attribute] = value[key]
    return TableStyleInfo(name=style_name, **values)


def _spreadsheet_table_name(workbook: Any, value: Any) -> str:
    existing = {
        table.name.casefold()
        for worksheet in workbook.worksheets
        for table in worksheet.tables.values()
    }
    existing.update(name.casefold() for name in workbook.defined_names if isinstance(name, str))
    if value is None:
        index = 1
        while f"table{index}" in existing:
            index += 1
        return f"Table{index}"
    cell_reference = False
    if isinstance(value, str) and (match := re.fullmatch(r"([A-Za-z]{1,3})([1-9][0-9]{0,6})", value)):
        from openpyxl.utils.cell import column_index_from_string

        cell_reference = (
            column_index_from_string(match.group(1)) <= EXCEL_MAX_COLUMN
            and int(match.group(2)) <= EXCEL_MAX_ROW
        )
    if (
        not isinstance(value, str)
        or len(value) > 255
        or not re.fullmatch(r"[A-Za-z_\\][A-Za-z0-9_.\\]*", value)
        or cell_reference
        or re.fullmatch(r"[Rr][1-9][0-9]*[Cc][1-9][0-9]*", value)
    ):
        raise ValueError("table must be a valid Excel table name, not a cell reference")
    if value.casefold() in existing:
        raise ValueError(f"Excel table name already exists: {value}")
    return value


def _spreadsheet_table_bounds(value: Any) -> tuple[str, tuple[int, int, int, int]]:
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    if not isinstance(value, str) or not re.fullmatch(
        r"[A-Za-z]{1,3}[1-9][0-9]{0,6}:[A-Za-z]{1,3}[1-9][0-9]{0,6}", value,
    ):
        raise ValueError("range must be a rectangular A1 range such as A1:C10")
    min_column, min_row, max_column, max_row = range_boundaries(value.upper())
    if min_column > max_column or min_row > max_row:
        raise ValueError("table range must run from its top-left cell to its bottom-right cell")
    if max_column > EXCEL_MAX_COLUMN or max_row > EXCEL_MAX_ROW:
        raise ValueError("table range is outside Excel limits")
    normalized = (
        f"{get_column_letter(min_column)}{min_row}:"
        f"{get_column_letter(max_column)}{max_row}"
    )
    return normalized, (min_column, min_row, max_column, max_row)


def insert_spreadsheet_table(workbook: Any, worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    from openpyxl.utils.cell import range_boundaries
    from openpyxl.worksheet.table import Table

    if set(patch) - {"operation", "sheet", "range", "table", "format"}:
        raise ValueError("Excel table.insert accepts sheet, range, table and format")
    if worksheet["A1"].value == _SPREADSHEET_CHART_DATA_MARKER:
        raise ValueError("Cannot create a user table on a managed chart-data worksheet")
    reference, bounds = _spreadsheet_table_bounds(patch.get("range"))
    for area in worksheet.merged_cells.ranges:
        if _spreadsheet_ranges_overlap(bounds, range_boundaries(str(area))):
            raise ValueError(f"table range overlaps merged range {area}")
    for existing in worksheet.tables.values():
        if _spreadsheet_ranges_overlap(bounds, range_boundaries(existing.ref)):
            raise ValueError(f"table range overlaps Excel table {existing.name}")
    min_column, min_row, max_column, _max_row = bounds
    headers = [worksheet.cell(min_row, column).value for column in range(min_column, max_column + 1)]
    if any(not isinstance(header, str) or not header.strip() for header in headers):
        raise ValueError("Excel table header cells must contain non-empty strings")
    normalized_headers = [header.casefold() for header in headers]
    if len(set(normalized_headers)) != len(normalized_headers):
        raise ValueError("Excel table header cells must be unique, ignoring case")
    name = _spreadsheet_table_name(workbook, patch.get("table"))
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = _spreadsheet_table_style(patch.get("format"), default=True)
    table._initialise_columns()
    for column, header in zip(table.tableColumns, headers):
        column.name = header
    worksheet.add_table(table)
    return {
        "updated": True,
        "operation": "table.insert",
        "sheet": worksheet.title,
        "table": table.name,
        "table_ref": table.ref,
    }


def format_spreadsheet_table(worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    if set(patch) - {"operation", "sheet", "table", "format"}:
        raise ValueError("Excel table.format accepts sheet, table and format")
    table = _spreadsheet_table_selector(worksheet, patch.get("table"))
    style = patch.get("format")
    if not isinstance(style, dict) or not style:
        raise ValueError("table format must be a non-empty object")
    if style.get("style_name", object()) is None:
        table.tableStyleInfo = _spreadsheet_table_style(style)
    else:
        current = table.tableStyleInfo
        if current is not None and "style_name" not in style:
            allowed = {
                "show_first_column": "showFirstColumn",
                "show_last_column": "showLastColumn",
                "show_row_stripes": "showRowStripes",
                "show_column_stripes": "showColumnStripes",
            }
            if set(style) - set(allowed):
                raise ValueError("table format contains unsupported properties")
            updated = copy.copy(current)
            for key, value in style.items():
                if type(value) is not bool:
                    raise ValueError(f"{key} must be boolean")
                setattr(updated, allowed[key], value)
            table.tableStyleInfo = updated
        else:
            merged = {
                "style_name": "TableStyleMedium2",
                "show_first_column": bool(getattr(current, "showFirstColumn", False)),
                "show_last_column": bool(getattr(current, "showLastColumn", False)),
                "show_row_stripes": bool(getattr(current, "showRowStripes", False)),
                "show_column_stripes": bool(getattr(current, "showColumnStripes", False)),
                **style,
            }
            table.tableStyleInfo = _spreadsheet_table_style(merged)
    return {
        "updated": True,
        "operation": "table.format",
        "sheet": worksheet.title,
        "table": table.name,
        "table_ref": table.ref,
    }


def _spreadsheet_reference_targets_sheet(reference: Any, sheet_name: str) -> bool:
    if not isinstance(reference, str) or "!" not in reference:
        return False
    quoted = "'" + sheet_name.replace("'", "''") + "'!"
    if quoted.casefold() in reference.casefold():
        return True
    return re.search(
        rf"(?<![A-Za-z0-9_.]){re.escape(sheet_name)}!",
        reference,
        flags=re.IGNORECASE,
    ) is not None


def spreadsheet_sheet_title(value: Any) -> str:
    if (
        not isinstance(value, str)
        or not value.strip()
        or len(value) > 31
        or value.startswith("'")
        or value.endswith("'")
        or value.casefold() == "history"
        or re.search(r"[\x00-\x1f*?:/\\\[\]]", value)
    ):
        raise ValueError(
            "Sheet names must be nonblank, at most 31 characters, Excel-compatible, "
            "not History, and cannot start/end with an apostrophe",
        )
    return value


_SPREADSHEET_SHEET_REFERENCE_PATTERN = re.compile(
    r"(^|[^A-Za-z0-9_.$\]])"
    r"(?:(?:'((?:[^']|'')+)'|([^'!+\-*/^&=<>%,(){}\[\]\s:]+)):)?"
    r"(?:'((?:[^']|'')+)'|([^'!+\-*/^&=<>%,(){}\[\]\s:]+))!",
    re.IGNORECASE,
)


def _rename_spreadsheet_formula_references(formula: str, old_name: str, new_name: str) -> str:
    lookup = {old_name.casefold(): new_name}

    def quoted(value: str) -> str:
        return "'" + value.replace("'", "''") + "'"

    def references(code: str) -> str:
        def replace(match: re.Match[str]) -> str:
            prefix, quoted_start, unquoted_start, quoted_name, unquoted_name = match.groups()
            current_name = (quoted_name or unquoted_name or "").replace("''", "'")
            if quoted_name and ":" in current_name and not current_name.startswith("["):
                range_start, range_end = current_name.split(":", 1)
                next_start = lookup.get(range_start.casefold(), range_start)
                next_end = lookup.get(range_end.casefold(), range_end)
                return (
                    prefix + quoted(next_start + ":" + next_end) + "!"
                    if (next_start, next_end) != (range_start, range_end)
                    else match.group(0)
                )
            if current_name.startswith("["):
                return match.group(0)
            next_name = lookup.get(current_name.casefold())
            range_start = quoted_start or unquoted_start
            if range_start is not None:
                current_start = range_start.replace("''", "'")
                if current_start.startswith("["):
                    return match.group(0)
                next_start = lookup.get(current_start.casefold(), current_start)
                next_end = next_name or current_name
                return (
                    prefix + quoted(next_start + ":" + next_end) + "!"
                    if next_name or next_start != current_start
                    else match.group(0)
                )
            return prefix + quoted(next_name) + "!" if next_name else match.group(0)

        return _SPREADSHEET_SHEET_REFERENCE_PATTERN.sub(replace, code)

    output, code_start, index = [], 0, 0
    while index < len(formula):
        if formula[index] != '"':
            index += 1
            continue
        output.append(references(formula[code_start:index]))
        literal_start = index
        index += 1
        while index < len(formula):
            if formula[index] != '"':
                index += 1
            elif index + 1 < len(formula) and formula[index + 1] == '"':
                index += 2
            else:
                index += 1
                break
        output.append(formula[literal_start:index])
        code_start = index
    output.append(references(formula[code_start:]))
    return "".join(output)


def _rename_spreadsheet_xml_references(xml: str, old_name: str, new_name: str) -> str:
    names = (
        "calculatedColumnFormula|totalsRowFormula|definedName|formula1|formula2|formula|f"
    )
    formula_pattern = re.compile(
        rf"(<(?:[A-Za-z_][\w.-]*:)?(?:{names})\b[^>]*>)([\s\S]*?)"
        rf"(</(?:[A-Za-z_][\w.-]*:)?(?:{names})>)",
        re.IGNORECASE,
    )

    def formula(match: re.Match[str]) -> str:
        source = html.unescape(match.group(2))
        updated = _rename_spreadsheet_formula_references(source, old_name, new_name)
        if updated == source:
            return match.group(0)
        return match.group(1) + html.escape(updated, quote=True) + match.group(3)

    output = formula_pattern.sub(formula, xml)

    def attribute(tag: str, name: str, value: str) -> str:
        return re.sub(
            rf'(\b{name}=")[^"]*(")',
            lambda match: match.group(1) + html.escape(value, quote=True) + match.group(2),
            tag,
            count=1,
            flags=re.IGNORECASE,
        )

    hyperlink_pattern = re.compile(
        r"<(?:(?:[A-Za-z_][\w.-]*):)?hyperlink\b[^>]*\blocation=\"[^\"]*\"[^>]*?/?>",
        re.IGNORECASE,
    )

    def hyperlink(match: re.Match[str]) -> str:
        tag = match.group(0)
        if re.search(r'\b[A-Za-z_][\w.-]*:id="[^"]*"', tag, re.IGNORECASE):
            return tag
        location = re.search(r'\blocation="([^"]*)"', tag, re.IGNORECASE)
        if location is None:
            return tag
        source = html.unescape(location.group(1))
        updated = _rename_spreadsheet_formula_references(source, old_name, new_name)
        return attribute(tag, "location", updated) if updated != source else tag

    output = hyperlink_pattern.sub(hyperlink, output)
    worksheet_source_pattern = re.compile(
        r"<(?:(?:[A-Za-z_][\w.-]*):)?worksheetSource\b[^>]*\bsheet=\"[^\"]*\"[^>]*?/?>",
        re.IGNORECASE,
    )

    def worksheet_source(match: re.Match[str]) -> str:
        tag = match.group(0)
        if re.search(r'\b[A-Za-z_][\w.-]*:id="[^"]*"', tag, re.IGNORECASE):
            return tag
        current = re.search(r'\bsheet="([^"]*)"', tag, re.IGNORECASE)
        if current is None or html.unescape(current.group(1)).casefold() != old_name.casefold():
            return tag
        return attribute(tag, "sheet", new_name)

    return worksheet_source_pattern.sub(worksheet_source, output)


def rename_spreadsheet_package_references(data: bytes, old_name: str, new_name: str) -> bytes:
    source = io.BytesIO(data)
    target = io.BytesIO()
    with zipfile.ZipFile(source, "r") as archive, zipfile.ZipFile(
        target, "w", compression=zipfile.ZIP_DEFLATED,
    ) as output:
        for info in archive.infolist():
            payload = archive.read(info.filename)
            if info.filename.startswith("xl/") and info.filename.endswith(".xml"):
                xml = payload.decode("utf-8")
                payload = _rename_spreadsheet_xml_references(
                    xml, old_name, new_name,
                ).encode("utf-8")
            output.writestr(info, payload)
    return target.getvalue()


def _spreadsheet_tree_references_sheet(root: Any, sheet_name: str) -> bool:
    return any(
        _spreadsheet_reference_targets_sheet(value, sheet_name)
        for node in root.iter()
        for value in [node.text, *node.attrib.values()]
    )


def _spreadsheet_sheet_references(workbook: Any, target: Any) -> list[str]:
    references: list[str] = []
    for worksheet in workbook.worksheets:
        if worksheet is target:
            continue
        for cell in worksheet._cells.values():
            if cell.data_type == "f" and _spreadsheet_reference_targets_sheet(cell.value, target.title):
                references.append(f"formula {worksheet.title}!{cell.coordinate}")
                if len(references) >= 10:
                    return references
            hyperlink = cell.hyperlink
            hyperlink_location = getattr(hyperlink, "location", None)
            hyperlink_target = getattr(hyperlink, "target", None)
            if _spreadsheet_reference_targets_sheet(hyperlink_location, target.title) or (
                isinstance(hyperlink_target, str)
                and hyperlink_target.startswith("#")
                and _spreadsheet_reference_targets_sheet(hyperlink_target, target.title)
            ):
                references.append(f"hyperlink {worksheet.title}!{cell.coordinate}")
                if len(references) >= 10:
                    return references
        for validation in worksheet.data_validations.dataValidation:
            if any(
                _spreadsheet_reference_targets_sheet(value, target.title)
                for value in (validation.formula1, validation.formula2)
            ):
                references.append(f"data validation {worksheet.title}!{validation.sqref}")
                if len(references) >= 10:
                    return references
        for conditional_range, rules in worksheet.conditional_formatting._cf_rules.items():
            if any(
                _spreadsheet_reference_targets_sheet(formula, target.title)
                for rule in rules
                for formula in (rule.formula or ())
            ):
                references.append(f"conditional formatting {worksheet.title}!{conditional_range.sqref}")
                if len(references) >= 10:
                    return references
        for table in worksheet.tables.values():
            if _spreadsheet_tree_references_sheet(table.to_tree(), target.title):
                references.append(f"table {worksheet.title}!{table.name}")
                if len(references) >= 10:
                    return references
        for chart_index, chart in enumerate(worksheet._charts):
            if _spreadsheet_tree_references_sheet(chart.to_tree(), target.title):
                references.append(f"chart {worksheet.title}[{chart_index}]")
                if len(references) >= 10:
                    return references
    for name, defined_name in workbook.defined_names.items():
        if _spreadsheet_reference_targets_sheet(defined_name.attr_text, target.title):
            references.append(f"defined name {name}")
            if len(references) >= 10:
                break
    return references


def _clear_spreadsheet_managed_chart_blocks(workbook: Any, blocks: list[tuple[Any, int, int, int]]) -> None:
    data_sheets = set()
    for data_sheet, first_row, last_row, last_column in blocks:
        data_sheets.add(data_sheet)
        for row in range(first_row, last_row + 1):
            for column in range(1, last_column + 1):
                data_sheet._cells.pop((row, column), None)
    live_data_sheets = {
        block[0]
        for sheet in workbook.worksheets
        for chart in sheet._charts
        if (block := _spreadsheet_managed_chart_block(workbook, chart)) is not None
    }
    for data_sheet in data_sheets - live_data_sheets:
        if data_sheet in workbook.worksheets and len(workbook.worksheets) > 1:
            workbook.remove(data_sheet)


def delete_spreadsheet_sheet(workbook: Any, worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    if set(patch) - {"operation", "sheet"}:
        raise ValueError("Excel sheet.delete accepts only sheet")
    if not isinstance(patch.get("sheet"), str) or not patch["sheet"]:
        raise ValueError("sheet.delete requires an explicit sheet name")
    if worksheet.title.startswith("_manor_chart_data") and worksheet["A1"].value == _SPREADSHEET_CHART_DATA_MARKER:
        raise ValueError("Managed chart data sheets cannot be deleted directly; delete their charts")
    remaining_visible = [
        sheet for sheet in workbook.worksheets
        if sheet is not worksheet and sheet.sheet_state == "visible"
    ]
    if not remaining_visible:
        raise ValueError("Cannot delete the last visible Excel worksheet")
    references = _spreadsheet_sheet_references(workbook, worksheet)
    if references:
        raise ValueError(
            f"Cannot delete worksheet {worksheet.title}; it is referenced by {', '.join(references)}",
        )
    title = worksheet.title
    managed_blocks = [
        block for chart in worksheet._charts
        if (block := _spreadsheet_managed_chart_block(workbook, chart)) is not None
    ]
    workbook.remove(worksheet)
    _clear_spreadsheet_managed_chart_blocks(workbook, managed_blocks)
    return {
        "updated": True,
        "operation": "sheet.delete",
        "sheet": title,
        "sheet_names": list(workbook.sheetnames),
    }


def reorder_spreadsheet_sheet(workbook: Any, worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    if set(patch) - {"operation", "sheet", "index"}:
        raise ValueError("Excel sheet.reorder accepts only sheet and index")
    if not isinstance(patch.get("sheet"), str) or not patch["sheet"]:
        raise ValueError("sheet.reorder requires an explicit sheet name")
    if (
        worksheet.title.startswith("_manor_chart_data")
        and worksheet["A1"].value == _SPREADSHEET_CHART_DATA_MARKER
    ):
        raise ValueError("Managed chart data sheets cannot be reordered directly")
    previous_index = workbook.worksheets.index(worksheet)
    index = _integer(patch.get("index"), "index", 0, len(workbook.worksheets) - 1)
    active = workbook.active
    if index != previous_index:
        workbook._sheets.pop(previous_index)
        workbook._sheets.insert(index, worksheet)
        workbook._active_sheet_index = workbook._sheets.index(active)
    return {
        "updated": True,
        "operation": "sheet.reorder",
        "sheet": worksheet.title,
        "previous_index": previous_index,
        "index": index,
        "sheet_names": list(workbook.sheetnames),
    }


def rename_spreadsheet_sheet(workbook: Any, worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    if set(patch) - {"operation", "sheet", "new_sheet"}:
        raise ValueError("Excel sheet.rename accepts only sheet and new_sheet")
    if not isinstance(patch.get("sheet"), str) or not patch["sheet"]:
        raise ValueError("sheet.rename requires an explicit sheet name")
    if (
        worksheet.title.startswith("_manor_chart_data")
        and worksheet["A1"].value == _SPREADSHEET_CHART_DATA_MARKER
    ):
        raise ValueError("Managed chart data sheets cannot be renamed directly")
    new_name = spreadsheet_sheet_title(patch.get("new_sheet"))
    if new_name.casefold() != worksheet.title.casefold() and new_name.casefold() in {
        name.casefold() for name in workbook.sheetnames
    }:
        raise ValueError(f"Worksheet already exists: {new_name}")
    old_name = worksheet.title
    if new_name != old_name:
        worksheet.title = new_name
    return {
        "updated": new_name != old_name,
        "operation": "sheet.rename",
        "sheet": old_name,
        "new_sheet": new_name,
        "sheet_names": list(workbook.sheetnames),
        "_sheet_rename": (old_name, new_name) if new_name != old_name else None,
    }


def _spreadsheet_merge_bounds(value: Any) -> tuple[str, tuple[int, int, int, int]]:
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    if not isinstance(value, str) or not re.fullmatch(
        r"[A-Za-z]{1,3}[1-9][0-9]{0,6}:[A-Za-z]{1,3}[1-9][0-9]{0,6}", value,
    ):
        raise ValueError("range must be a rectangular A1 range such as A1:C2")
    min_column, min_row, max_column, max_row = range_boundaries(value.upper())
    if max_column > EXCEL_MAX_COLUMN or max_row > EXCEL_MAX_ROW:
        raise ValueError("merge range is outside Excel limits")
    if min_column == max_column and min_row == max_row:
        raise ValueError("merge range must contain at least two cells")
    normalized = (
        f"{get_column_letter(min_column)}{min_row}:"
        f"{get_column_letter(max_column)}{max_row}"
    )
    return normalized, (min_column, min_row, max_column, max_row)


def _spreadsheet_ranges_overlap(first: tuple[int, int, int, int], second: tuple[int, int, int, int]) -> bool:
    left, top, right, bottom = first
    other_left, other_top, other_right, other_bottom = second
    return left <= other_right and right >= other_left and top <= other_bottom and bottom >= other_top


def set_spreadsheet_merge(worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    from openpyxl.utils.cell import range_boundaries

    if set(patch) - {"operation", "sheet", "range"}:
        raise ValueError("Excel merge.set accepts sheet and range")
    reference, bounds = _spreadsheet_merge_bounds(patch.get("range"))
    existing = {str(area).upper(): area for area in worksheet.merged_cells.ranges}
    if reference in existing:
        return {"updated": False, "operation": "merge.set", "sheet": worksheet.title, "range": reference}
    for existing_reference, area in existing.items():
        if _spreadsheet_ranges_overlap(bounds, range_boundaries(existing_reference)):
            raise ValueError(f"merge range overlaps existing merged range {area}")
    for table in worksheet.tables.values():
        if _spreadsheet_ranges_overlap(bounds, range_boundaries(table.ref)):
            raise ValueError(f"merge range overlaps Excel table {table.name}")
    min_column, min_row, max_column, max_row = bounds
    anchor = (min_row, min_column)
    for row in range(min_row, max_row + 1):
        for column in range(min_column, max_column + 1):
            if (row, column) == anchor:
                continue
            cell = worksheet._cells.get((row, column))
            if cell is not None and (
                cell.value is not None or cell.comment is not None or cell.hyperlink is not None or cell.has_style
            ):
                raise ValueError(f"merge range would discard content or style in {cell.coordinate}")
    worksheet.merge_cells(reference)
    return {"updated": True, "operation": "merge.set", "sheet": worksheet.title, "range": reference}


def clear_spreadsheet_merge(worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    if set(patch) - {"operation", "sheet", "range"}:
        raise ValueError("Excel merge.clear accepts sheet and range")
    reference, _bounds = _spreadsheet_merge_bounds(patch.get("range"))
    existing = {str(area).upper() for area in worksheet.merged_cells.ranges}
    if reference not in existing:
        raise ValueError("merge.clear requires an exact existing merged range")
    worksheet.unmerge_cells(reference)
    return {"updated": True, "operation": "merge.clear", "sheet": worksheet.title, "range": reference}


def _spreadsheet_validation_range(value: Any) -> tuple[str, tuple[int, int, int, int]]:
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    if not isinstance(value, str) or not re.fullmatch(
        r"[A-Za-z]{1,3}[1-9][0-9]{0,6}(?::[A-Za-z]{1,3}[1-9][0-9]{0,6})?", value,
    ):
        raise ValueError("range must be one A1 cell or rectangular A1 range such as B2:B100")
    min_column, min_row, max_column, max_row = range_boundaries(value.upper())
    if min_column > max_column or min_row > max_row:
        raise ValueError("validation range must run from its top-left cell to its bottom-right cell")
    if max_column > EXCEL_MAX_COLUMN or max_row > EXCEL_MAX_ROW:
        raise ValueError("validation range is outside Excel limits")
    normalized = f"{get_column_letter(min_column)}{min_row}"
    if min_column != max_column or min_row != max_row:
        normalized += f":{get_column_letter(max_column)}{max_row}"
    return normalized, (min_column, min_row, max_column, max_row)


def _spreadsheet_validation_formula(value: Any, name: str) -> Any:
    if type(value) not in {str, int, float}:
        raise ValueError(f"{name} must be a string or finite number")
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError(f"{name} must be a string or finite number")
    if isinstance(value, str) and (not value or len(value) > 255 or "\x00" in value):
        raise ValueError(f"{name} must contain 1 to 255 valid characters")
    return value


def _spreadsheet_validation_config(value: Any, *, base: Any = None, inserting: bool = False) -> Any:
    from openpyxl.worksheet.datavalidation import DataValidation

    allowed = {
        "type", "operator", "formula1", "formula2", "values", "allow_blank",
        "show_dropdown", "show_error_message", "error_style", "error_title", "error",
        "show_input_message", "prompt_title", "prompt",
    }
    if not isinstance(value, dict) or not value or set(value) - allowed:
        raise ValueError("validation must be a non-empty object with supported properties")
    types = {"whole", "decimal", "list", "date", "time", "textLength", "custom"}
    previous_type = getattr(base, "type", None)
    kind = value.get("type", previous_type)
    if kind not in types:
        raise ValueError("validation type must be whole, decimal, list, date, time, textLength or custom")
    if inserting and "type" not in value:
        raise ValueError("validation.insert requires validation.type")
    changed_type = base is not None and "type" in value and kind != previous_type
    previous = None if changed_type else base

    formula1 = getattr(previous, "formula1", None)
    formula2 = getattr(previous, "formula2", None)
    if "values" in value:
        if kind != "list" or "formula1" in value:
            raise ValueError("validation.values is available only for list type and cannot accompany formula1")
        values = value["values"]
        if not isinstance(values, list) or not 1 <= len(values) <= 100:
            raise ValueError("validation.values must contain 1 to 100 strings")
        if any(
            not isinstance(item, str) or not item or any(char in item for char in ('"', ",", "\n", "\r", "\x00"))
            for item in values
        ):
            raise ValueError("inline validation values must be non-empty strings without quotes, commas or line breaks")
        formula1 = '"' + ",".join(values) + '"'
        if len(formula1) > 255:
            raise ValueError("inline validation values cannot exceed Excel's 255-character limit")
    elif "formula1" in value:
        formula1 = _spreadsheet_validation_formula(value["formula1"], "validation.formula1")
    if "formula2" in value:
        formula2 = _spreadsheet_validation_formula(value["formula2"], "validation.formula2")

    operators = {
        "between", "notBetween", "equal", "notEqual", "lessThan", "lessThanOrEqual",
        "greaterThan", "greaterThanOrEqual",
    }
    operator = value.get("operator", getattr(previous, "operator", None))
    if "operator" in value and operator not in {"between", "notBetween"} and "formula2" not in value:
        formula2 = None
    if kind in {"list", "custom"}:
        if operator is not None or "formula2" in value:
            raise ValueError(f"{kind} validation does not accept operator or formula2")
        operator = formula2 = None
    else:
        operator = operator or "between"
        if operator not in operators:
            raise ValueError("validation.operator is not supported")
        if operator in {"between", "notBetween"}:
            if formula2 is None:
                raise ValueError(f"{operator} validation requires formula2")
        elif formula2 is not None:
            raise ValueError(f"{operator} validation does not accept formula2")
    if formula1 is None:
        raise ValueError(f"{kind} validation requires formula1 or values")

    result = copy.copy(base) if base is not None else DataValidation()
    result.type = kind
    result.operator = operator
    result.formula1 = formula1
    result.formula2 = formula2
    for key, attribute in (
        ("allow_blank", "allowBlank"),
        ("show_error_message", "showErrorMessage"),
        ("show_input_message", "showInputMessage"),
    ):
        if key in value:
            if type(value[key]) is not bool:
                raise ValueError(f"validation.{key} must be boolean")
            setattr(result, attribute, value[key])
        elif base is None:
            setattr(result, attribute, False)
    if "show_dropdown" in value:
        if kind != "list" or type(value["show_dropdown"]) is not bool:
            raise ValueError("validation.show_dropdown must be boolean and is available only for list type")
        result.showDropDown = not value["show_dropdown"]
    elif base is None or changed_type:
        result.showDropDown = None
    if "error_style" in value:
        error_style = value["error_style"]
        if error_style is not None and error_style not in {"stop", "warning", "information"}:
            raise ValueError("validation.error_style must be stop, warning, information or null")
        result.errorStyle = error_style
    for key, attribute, limit in (
        ("error_title", "errorTitle", 32), ("prompt_title", "promptTitle", 32),
        ("error", "error", 255), ("prompt", "prompt", 255),
    ):
        if key not in value:
            continue
        message = value[key]
        if message is not None and (
            not isinstance(message, str) or len(message) > limit or "\x00" in message
        ):
            raise ValueError(f"validation.{key} must be null or a string of at most {limit} characters")
        setattr(result, attribute, message)
    return result


def _spreadsheet_validation_ranges(validation: Any) -> list[tuple[str, tuple[int, int, int, int]]]:
    from openpyxl.utils.cell import range_boundaries

    result = []
    for reference in sorted(str(area) for area in validation.ranges.ranges):
        min_column, min_row, max_column, max_row = range_boundaries(reference)
        result.append((reference, (
            min_column or 1,
            min_row or 1,
            max_column or EXCEL_MAX_COLUMN,
            max_row or EXCEL_MAX_ROW,
        )))
    return result


def _assert_spreadsheet_validation_range_available(
    worksheet: Any,
    bounds: tuple[int, int, int, int],
    *,
    excluding: Any = None,
) -> None:
    for validation in worksheet.data_validations.dataValidation:
        if validation is excluding:
            continue
        for reference, existing_bounds in _spreadsheet_validation_ranges(validation):
            if _spreadsheet_ranges_overlap(bounds, existing_bounds):
                raise ValueError(f"validation range overlaps existing data validation {reference}")


def insert_spreadsheet_validation(worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    if set(patch) - {"operation", "sheet", "range", "validation"}:
        raise ValueError("Excel validation.insert accepts sheet, range and validation")
    reference, bounds = _spreadsheet_validation_range(patch.get("range"))
    _assert_spreadsheet_validation_range_available(worksheet, bounds)
    validation = _spreadsheet_validation_config(patch.get("validation"), inserting=True)
    validation.add(reference)
    worksheet.add_data_validation(validation)
    index = len(worksheet.data_validations.dataValidation) - 1
    return {
        "updated": True,
        "operation": "validation.insert",
        "sheet": worksheet.title,
        "validation_index": index,
        "range": reference,
    }


def format_spreadsheet_validation(worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    if set(patch) - {"operation", "sheet", "validation_index", "range", "validation"}:
        raise ValueError("Excel validation.format accepts sheet, validation_index, optional range and validation")
    validations = worksheet.data_validations.dataValidation
    index = _integer(patch.get("validation_index"), "validation_index", 0, len(validations) - 1)
    current = validations[index]
    updated = _spreadsheet_validation_config(patch.get("validation"), base=current)
    if "range" in patch:
        reference, bounds = _spreadsheet_validation_range(patch["range"])
        _assert_spreadsheet_validation_range_available(worksheet, bounds, excluding=current)
        updated.sqref = reference
    validations[index] = updated
    return {
        "updated": True,
        "operation": "validation.format",
        "sheet": worksheet.title,
        "validation_index": index,
        "ranges": [reference for reference, _bounds in _spreadsheet_validation_ranges(updated)],
    }


def delete_spreadsheet_validation(worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    if set(patch) - {"operation", "sheet", "validation_index"}:
        raise ValueError("Excel validation.delete accepts sheet and validation_index")
    validations = worksheet.data_validations.dataValidation
    index = _integer(patch.get("validation_index"), "validation_index", 0, len(validations) - 1)
    removed = validations.pop(index)
    return {
        "updated": True,
        "operation": "validation.delete",
        "sheet": worksheet.title,
        "validation_index": index,
        "ranges": [reference for reference, _bounds in _spreadsheet_validation_ranges(removed)],
    }


def _spreadsheet_conditional_format_range(value: Any) -> str:
    if not isinstance(value, str) or not re.fullmatch(
        r"[A-Za-z]{1,3}[1-9][0-9]{0,6}(?::[A-Za-z]{1,3}[1-9][0-9]{0,6})?",
        value,
    ):
        raise ValueError("range must be one A1 cell or rectangular A1 range such as B2:B100")
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    min_column, min_row, max_column, max_row = range_boundaries(value.upper())
    if min_column > max_column or min_row > max_row:
        raise ValueError("conditional-format range must run from top-left to bottom-right")
    if max_column > EXCEL_MAX_COLUMN or max_row > EXCEL_MAX_ROW:
        raise ValueError("conditional-format range is outside Excel limits")
    normalized = f"{get_column_letter(min_column)}{min_row}"
    if min_column != max_column or min_row != max_row:
        normalized += f":{get_column_letter(max_column)}{max_row}"
    return normalized


def _spreadsheet_conditional_format_rules(worksheet: Any) -> list[tuple[Any, int, Any]]:
    return [
        (container, rule_index, rule)
        for container, rules in worksheet.conditional_formatting._cf_rules.items()
        for rule_index, rule in enumerate(rules)
    ]


def _spreadsheet_conditional_formula(value: Any, name: str) -> str:
    if type(value) not in {str, int, float}:
        raise ValueError(f"{name} must be a formula string or finite number")
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError(f"{name} must be a formula string or finite number")
    result = str(value)
    if isinstance(value, str):
        result = value[1:] if value.startswith("=") else value
        if not result or len(result) > 8192 or "\x00" in result:
            raise ValueError(f"{name} must contain 1 to 8192 valid characters")
    return result


def _spreadsheet_conditional_formulas(value: Any, *, count: int | tuple[int, int]) -> list[str]:
    if not isinstance(value, list):
        raise ValueError("conditional_format.formulas must be an array")
    minimum, maximum = (count, count) if isinstance(count, int) else count
    if not minimum <= len(value) <= maximum:
        expected = str(minimum) if minimum == maximum else f"{minimum} to {maximum}"
        raise ValueError(f"conditional_format.formulas must contain {expected} formula(s)")
    return [
        _spreadsheet_conditional_formula(formula, f"conditional_format.formulas[{index}]")
        for index, formula in enumerate(value)
    ]


def _spreadsheet_conditional_thresholds(value: Any, *, count: int | tuple[int, int]) -> list[Any]:
    from openpyxl.formatting.rule import FormatObject

    if not isinstance(value, list):
        raise ValueError("conditional_format.thresholds must be an array")
    minimum, maximum = (count, count) if isinstance(count, int) else count
    if not minimum <= len(value) <= maximum:
        expected = str(minimum) if minimum == maximum else f"{minimum} to {maximum}"
        raise ValueError(f"conditional_format.thresholds must contain {expected} threshold(s)")
    thresholds = []
    for index, threshold in enumerate(value):
        if not isinstance(threshold, dict) or not threshold or set(threshold) - {"type", "value", "gte"}:
            raise ValueError("each conditional-format threshold accepts type, value and optional gte")
        kind = threshold.get("type")
        if kind not in {"min", "max", "num", "percent", "percentile", "formula"}:
            raise ValueError("threshold type must be min, max, num, percent, percentile or formula")
        if "gte" in threshold and type(threshold["gte"]) is not bool:
            raise ValueError("conditional-format threshold gte must be boolean")
        raw_value = threshold.get("value")
        if kind in {"min", "max"}:
            if "value" in threshold and raw_value is not None:
                raise ValueError(f"{kind} threshold cannot have a value")
            parsed = None
        elif kind == "formula":
            parsed = _spreadsheet_conditional_formula(
                raw_value,
                f"conditional_format.thresholds[{index}].value",
            )
        else:
            parsed = _number(raw_value, f"conditional_format.thresholds[{index}].value", -1e15, 1e15)
            if kind in {"percent", "percentile"} and not 0 <= parsed <= 100:
                raise ValueError(f"{kind} threshold value must be between 0 and 100")
        thresholds.append(FormatObject(type=kind, val=parsed, gte=threshold.get("gte")))
    return thresholds


def _validate_spreadsheet_conditional_threshold_order(thresholds: list[Any]) -> None:
    if any(item.type == "min" for item in thresholds[1:]):
        raise ValueError("min threshold is allowed only in the first position")
    if any(item.type == "max" for item in thresholds[:-1]):
        raise ValueError("max threshold is allowed only in the final position")


def _spreadsheet_conditional_color(value: Any, name: str) -> str:
    if not isinstance(value, str) or not re.fullmatch(r"[0-9a-fA-F]{6}", value):
        raise ValueError(f"{name} must be a six-digit RGB hex string without #")
    return "FF" + value.upper()


def _spreadsheet_conditional_colors(value: Any, count: int) -> list[str]:
    if not isinstance(value, list) or len(value) != count:
        raise ValueError(f"conditional_format.colors must contain {count} RGB colors")
    return [
        _spreadsheet_conditional_color(color, f"conditional_format.colors[{index}]")
        for index, color in enumerate(value)
    ]


def _spreadsheet_conditional_dxf(base: Any, value: Any) -> Any:
    from openpyxl.styles import Border, Font, PatternFill
    from openpyxl.styles.differential import DifferentialStyle

    allowed = {
        "bold", "italic", "underline", "strike", "font_name", "font_size",
        "font_color", "fill_color", "borders",
    }
    if not isinstance(value, dict) or not value or set(value) - allowed:
        raise ValueError(
            "conditional_format.format must contain supported font, fill or border properties",
        )
    result = copy.deepcopy(base) if base is not None else DifferentialStyle()
    font = copy.deepcopy(result.font) if result.font is not None else Font()
    border = copy.deepcopy(result.border) if result.border is not None else Border()
    for key, item in value.items():
        if key in {"bold", "italic", "underline", "strike"}:
            if item is not None and type(item) is not bool:
                raise ValueError(f"conditional_format.format.{key} must be boolean or null")
            attribute = {"bold": "b", "italic": "i", "strike": "strike"}.get(key, "u")
            setattr(font, attribute, "single" if key == "underline" and item else (None if key == "underline" else item))
        elif key == "font_name":
            if item is not None and (not isinstance(item, str) or not item or len(item) > 255 or "\x00" in item):
                raise ValueError("conditional_format.format.font_name must be a valid string or null")
            font.name = item
        elif key == "font_size":
            font.sz = None if item is None else _number(item, key, 1, 409)
        elif key == "font_color":
            font.color = None if item is None else _spreadsheet_conditional_color(item, key)
        elif key == "fill_color":
            result.fill = None if item is None else PatternFill(
                "solid",
                fgColor=_spreadsheet_conditional_color(item, key),
            )
        else:
            if not isinstance(item, dict) or not item or set(item) - {"left", "right", "top", "bottom"}:
                raise ValueError("conditional_format.format.borders must specify supported sides")
            from openpyxl.styles import Side

            for edge, side in item.items():
                if side is None:
                    setattr(border, edge, Side())
                    continue
                if not isinstance(side, dict) or set(side) != {"style", "color"}:
                    raise ValueError("a conditional-format border requires style and RGB color; null clears it")
                if not isinstance(side["style"], str):
                    raise ValueError("conditional-format border style must be a string")
                setattr(border, edge, Side(
                    style=side["style"],
                    color=_spreadsheet_conditional_color(side["color"], f"borders.{edge}.color"),
                ))
    if set(value) & {"bold", "italic", "underline", "strike", "font_name", "font_size", "font_color"}:
        result.font = font
    if "borders" in value:
        result.border = border
    return result


def _spreadsheet_conditional_format_config(value: Any, *, base: Any = None, inserting: bool = False) -> Any:
    from openpyxl.formatting.rule import Rule
    from openpyxl.styles import Color

    if not isinstance(value, dict) or not value:
        raise ValueError("conditional_format must be a non-empty object with supported properties")
    native_to_public = {
        "cellIs": "cell",
        "expression": "formula",
        "colorScale": "color_scale",
        "dataBar": "data_bar",
        "iconSet": "icon_set",
    }
    public_to_native = {public: native for native, public in native_to_public.items()}
    previous_type = native_to_public.get(getattr(base, "type", None))
    kind = value.get("type", previous_type)
    if inserting and "type" not in value:
        raise ValueError("conditional_format.insert requires conditional_format.type")
    if kind not in public_to_native:
        if base is not None and set(value) == {"stop_if_true"}:
            result = copy.deepcopy(base)
            if type(value["stop_if_true"]) is not bool:
                raise ValueError("conditional_format.stop_if_true must be boolean")
            result.stopIfTrue = value["stop_if_true"]
            return result
        raise ValueError("conditional-format type must be cell, formula, color_scale, data_bar or icon_set")
    changed_type = base is not None and "type" in value and kind != previous_type
    result = Rule(type=public_to_native[kind]) if base is None or changed_type else copy.deepcopy(base)
    if base is not None and changed_type:
        result.priority = base.priority

    common = {"type"}
    allowed = {
        "cell": common | {"operator", "formulas", "stop_if_true", "format"},
        "formula": common | {"formulas", "stop_if_true", "format"},
        "color_scale": common | {"thresholds", "colors"},
        "data_bar": common | {"thresholds", "color", "show_value", "min_length", "max_length"},
        "icon_set": common | {"thresholds", "icon_style", "show_value", "reverse", "percent"},
    }[kind]
    if set(value) - allowed:
        raise ValueError(f"conditional_format contains properties unsupported for {kind}")

    if kind in {"cell", "formula"}:
        if "stop_if_true" in value:
            if type(value["stop_if_true"]) is not bool:
                raise ValueError("conditional_format.stop_if_true must be boolean")
            result.stopIfTrue = value["stop_if_true"]
        if "format" in value:
            result.dxf = _spreadsheet_conditional_dxf(result.dxf, value["format"])
        if kind == "formula":
            formulas = (
                _spreadsheet_conditional_formulas(value["formulas"], count=1)
                if "formulas" in value
                else list(result.formula or ())
            )
            if not formulas:
                raise ValueError("formula conditional formatting requires formulas")
            result.formula = formulas
            result.operator = None
        else:
            operators = {
                "between", "notBetween", "equal", "notEqual", "lessThan", "lessThanOrEqual",
                "greaterThan", "greaterThanOrEqual",
            }
            operator = value.get("operator", result.operator)
            if operator not in operators:
                raise ValueError("conditional_format.operator is not supported")
            expected_count = 2 if operator in {"between", "notBetween"} else 1
            if "formulas" in value:
                formulas = _spreadsheet_conditional_formulas(value["formulas"], count=expected_count)
            else:
                formulas = list(result.formula or ())
                if len(formulas) > expected_count and "operator" in value:
                    formulas = formulas[:expected_count]
                if len(formulas) != expected_count:
                    raise ValueError(f"{operator} conditional formatting requires {expected_count} formula(s)")
            result.operator = operator
            result.formula = formulas
        return result

    if base is None or changed_type:
        result.dxf = None
        result.stopIfTrue = None
        result.operator = None
        result.formula = []
    if kind == "color_scale":
        current = result.colorScale
        thresholds = (
            _spreadsheet_conditional_thresholds(value["thresholds"], count=(2, 3))
            if "thresholds" in value
            else copy.deepcopy(getattr(current, "cfvo", None))
        )
        if not thresholds or len(thresholds) not in {2, 3}:
            raise ValueError("color_scale conditional formatting requires two or three thresholds")
        _validate_spreadsheet_conditional_threshold_order(thresholds)
        colors = (
            [Color(rgb=color) for color in _spreadsheet_conditional_colors(value["colors"], len(thresholds))]
            if "colors" in value
            else copy.deepcopy(getattr(current, "color", None))
        )
        if not colors or len(colors) != len(thresholds):
            raise ValueError("color_scale requires one color for every threshold")
        from openpyxl.formatting.rule import ColorScale

        if current is None:
            result.colorScale = ColorScale(cfvo=thresholds, color=colors)
        else:
            result.colorScale = copy.deepcopy(current)
            result.colorScale.cfvo = thresholds
            result.colorScale.color = colors
        result.dataBar = result.iconSet = None
    elif kind == "data_bar":
        current = result.dataBar
        thresholds = (
            _spreadsheet_conditional_thresholds(value["thresholds"], count=2)
            if "thresholds" in value
            else copy.deepcopy(getattr(current, "cfvo", None))
        )
        if not thresholds or len(thresholds) != 2:
            raise ValueError("data_bar conditional formatting requires two thresholds")
        _validate_spreadsheet_conditional_threshold_order(thresholds)
        color = (
            Color(rgb=_spreadsheet_conditional_color(value["color"], "conditional_format.color"))
            if "color" in value
            else copy.deepcopy(getattr(current, "color", None))
        )
        if color is None:
            raise ValueError("data_bar conditional formatting requires color")
        from openpyxl.formatting.rule import DataBar

        minimum = value.get("min_length", getattr(current, "minLength", None))
        maximum = value.get("max_length", getattr(current, "maxLength", None))
        minimum = None if minimum is None else _integer(minimum, "min_length", 0, 100)
        maximum = None if maximum is None else _integer(maximum, "max_length", 0, 100)
        if minimum is not None and maximum is not None and minimum > maximum:
            raise ValueError("conditional_format.min_length cannot exceed max_length")
        show_value = value.get("show_value", getattr(current, "showValue", None))
        if show_value is not None and type(show_value) is not bool:
            raise ValueError("conditional_format.show_value must be boolean")
        if current is None:
            result.dataBar = DataBar(
                cfvo=thresholds,
                color=color,
                showValue=show_value,
                minLength=minimum,
                maxLength=maximum,
            )
        else:
            result.dataBar = copy.deepcopy(current)
            result.dataBar.cfvo = thresholds
            result.dataBar.color = color
            result.dataBar.showValue = show_value
            result.dataBar.minLength = minimum
            result.dataBar.maxLength = maximum
        result.colorScale = result.iconSet = None
    else:
        current = result.iconSet
        icon_style = value.get("icon_style", getattr(current, "iconSet", None))
        icon_styles = {
            "3Arrows", "3ArrowsGray", "3Flags", "3Signs", "3Symbols", "3Symbols2",
            "3TrafficLights1", "3TrafficLights2", "4Arrows", "4ArrowsGray", "4Rating",
            "4RedToBlack", "4TrafficLights", "5Arrows", "5ArrowsGray", "5Quarters", "5Rating",
        }
        if icon_style not in icon_styles:
            raise ValueError("conditional_format.icon_style is not a supported native Excel icon set")
        icon_count = int(icon_style[0])
        thresholds = (
            _spreadsheet_conditional_thresholds(value["thresholds"], count=icon_count)
            if "thresholds" in value
            else copy.deepcopy(getattr(current, "cfvo", None))
        )
        if not thresholds or len(thresholds) != icon_count:
            raise ValueError(f"{icon_style} conditional formatting requires {icon_count} thresholds")
        _validate_spreadsheet_conditional_threshold_order(thresholds)
        options = {}
        for key, attribute in (("show_value", "showValue"), ("reverse", "reverse"), ("percent", "percent")):
            option = value.get(key, getattr(current, attribute, None))
            if option is not None and type(option) is not bool:
                raise ValueError(f"conditional_format.{key} must be boolean")
            options[attribute] = option
        from openpyxl.formatting.rule import IconSet

        if current is None:
            result.iconSet = IconSet(iconSet=icon_style, cfvo=thresholds, **options)
        else:
            result.iconSet = copy.deepcopy(current)
            result.iconSet.iconSet = icon_style
            result.iconSet.cfvo = thresholds
            for attribute, option in options.items():
                setattr(result.iconSet, attribute, option)
        result.colorScale = result.dataBar = None
    return result


def insert_spreadsheet_conditional_format(worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    if set(patch) - {"operation", "sheet", "range", "conditional_format"}:
        raise ValueError("Excel conditional_format.insert accepts sheet, range and conditional_format")
    reference = _spreadsheet_conditional_format_range(patch.get("range"))
    rule = _spreadsheet_conditional_format_config(patch.get("conditional_format"), inserting=True)
    worksheet.conditional_formatting.add(reference, rule)
    rules = _spreadsheet_conditional_format_rules(worksheet)
    index = next(index for index, (_container, _position, item) in enumerate(rules) if item is rule)
    return {
        "updated": True,
        "operation": "conditional_format.insert",
        "sheet": worksheet.title,
        "conditional_format_index": index,
        "range": reference,
    }


def format_spreadsheet_conditional_format(worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    if set(patch) - {"operation", "sheet", "conditional_format_index", "range", "conditional_format"}:
        raise ValueError(
            "Excel conditional_format.format accepts sheet, conditional_format_index, optional range and conditional_format",
        )
    if "range" not in patch and "conditional_format" not in patch:
        raise ValueError("conditional_format.format requires range or conditional_format changes")
    entries = _spreadsheet_conditional_format_rules(worksheet)
    if not entries:
        raise ValueError("selected worksheet has no conditional-format rules")
    index = _integer(
        patch.get("conditional_format_index"),
        "conditional_format_index",
        0,
        len(entries) - 1,
    )
    container, rule_index, current = entries[index]
    updated = (
        _spreadsheet_conditional_format_config(patch["conditional_format"], base=current)
        if "conditional_format" in patch
        else current
    )
    rules = worksheet.conditional_formatting._cf_rules[container]
    if "range" not in patch:
        rules[rule_index] = updated
        new_index = index
        ranges = sorted(str(area) for area in container.sqref.ranges)
    else:
        reference = _spreadsheet_conditional_format_range(patch["range"])
        existing_ranges = sorted(str(area) for area in container.sqref.ranges)
        if existing_ranges == [reference]:
            rules[rule_index] = updated
            new_index = index
        else:
            rules.pop(rule_index)
            if not rules:
                del worksheet.conditional_formatting._cf_rules[container]
            worksheet.conditional_formatting.add(reference, updated)
            entries = _spreadsheet_conditional_format_rules(worksheet)
            new_index = next(
                item_index
                for item_index, (_container, _position, item) in enumerate(entries)
                if item is updated
            )
        ranges = [reference]
    return {
        "updated": True,
        "operation": "conditional_format.format",
        "sheet": worksheet.title,
        "conditional_format_index": new_index,
        "ranges": ranges,
    }


def delete_spreadsheet_conditional_format(worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    if set(patch) - {"operation", "sheet", "conditional_format_index"}:
        raise ValueError("Excel conditional_format.delete accepts sheet and conditional_format_index")
    entries = _spreadsheet_conditional_format_rules(worksheet)
    if not entries:
        raise ValueError("selected worksheet has no conditional-format rules")
    index = _integer(
        patch.get("conditional_format_index"),
        "conditional_format_index",
        0,
        len(entries) - 1,
    )
    container, rule_index, _rule = entries[index]
    rules = worksheet.conditional_formatting._cf_rules[container]
    rules.pop(rule_index)
    if not rules:
        del worksheet.conditional_formatting._cf_rules[container]
    return {
        "updated": True,
        "operation": "conditional_format.delete",
        "sheet": worksheet.title,
        "conditional_format_index": index,
        "ranges": sorted(str(area) for area in container.sqref.ranges),
    }


def spreadsheet_cell_value(value: Any) -> Any:
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    if value is not None and type(value) not in {str, int, float, bool}:
        raise ValueError("Spreadsheet cell values must be scalar")
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("Spreadsheet numbers must be finite")
    if isinstance(value, str):
        if len(value) > EXCEL_MAX_CELL_CHARS:
            raise ValueError(f"Cell text cannot exceed {EXCEL_MAX_CELL_CHARS} characters")
        if ILLEGAL_CHARACTERS_RE.search(value):
            raise ValueError("Cell text contains an invalid control character")
    return value


def spreadsheet_column_map(columns: Any) -> dict[str, list[int]]:
    result: dict[str, list[int]] = {}
    for index, value in columns:
        key = re.sub(r"\s+", " ", str(value if value is not None else "").strip()).lower()
        if key:
            result.setdefault(key, []).append(index)
    return result


def spreadsheet_header_map(worksheet: Any, header_row: Any) -> dict[str, list[int]]:
    _integer(header_row, "header_row", 1, EXCEL_MAX_ROW)
    result = spreadsheet_column_map((cell.column, cell.value) for cell in worksheet[header_row])
    if not result:
        raise ValueError(f"No headers found on row {header_row}")
    return result


def spreadsheet_column_index(columns: dict[str, list[int]], name: str) -> int:
    key = re.sub(r"\s+", " ", name.strip()).lower()
    if not key:
        raise ValueError("Column name is required")
    if key not in columns:
        raise ValueError(f"Column not found: {name}. Known columns: {', '.join(sorted(columns)[:20])}")
    if len(columns[key]) != 1:
        raise ValueError(f"Ambiguous column name: {name}; duplicate headers at columns {columns[key]}")
    return columns[key][0]


def _translate_spreadsheet_formula(formula: str, origin: str, target: str) -> str:
    from openpyxl.formula.tokenizer import Token, TokenizerError
    from openpyxl.formula.translate import Translator, TranslatorError
    from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries

    row, column = coordinate_to_tuple(target)
    source_row, source_column = coordinate_to_tuple(origin)

    def shift(reference: str) -> str:
        if reference.startswith("@"):
            return "@" + shift(reference[1:])
        is_range = Translator.ROW_RANGE_RE.fullmatch(reference) or Translator.COL_RANGE_RE.fullmatch(reference)
        if not is_range and ":" in reference:
            return ":".join(shift(part) for part in reference.split(":"))
        if not is_range and not Translator.CELL_REF_RE.fullmatch(reference):
            return reference  # Structured references and defined names.
        limits = (EXCEL_MAX_COLUMN, EXCEL_MAX_ROW, EXCEL_MAX_COLUMN, EXCEL_MAX_ROW)
        if any(value and value > limit for value, limit in zip(range_boundaries(reference), limits)):
            return reference  # Outside-grid identifiers can be defined names (e.g. XFE1).
        translated = Translator.translate_range(reference, row - source_row, column - source_column)
        if any(value and value > limit for value, limit in zip(range_boundaries(translated), limits)):
            raise ValueError("Inherited formula reference is outside Excel limits")
        return translated

    try:
        tokens = Translator(formula, origin).get_tokens()
        translated = []
        for token in tokens:
            value = token.value
            if token.type == Token.OPERAND and token.subtype == Token.RANGE:
                prefix, reference = Translator.strip_ws_name(value)
                if "!" in re.sub(r"'(?:[^']|'')*'", "", prefix[:-1]):
                    raise ValueError("Cannot safely inherit a multiply-qualified range; provide an explicit formula")
                # Preserve qualified names too: Translator.translate_formula
                # drops the sheet prefix for a named (rather than A1) range.
                value = prefix + shift(reference)
            translated.append(value)
        return "=" + "".join(translated)
    except (TokenizerError, TranslatorError) as exc:
        raise ValueError(f"Cannot safely inherit formula from {origin}: {exc}") from exc


def append_spreadsheet_row(worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    if ("row" in patch) == ("values" in patch):
        raise ValueError("Provide exactly one of row or values for append_row")
    row_object, values = patch.get("row"), patch.get("values")
    if not ((isinstance(row_object, dict) and row_object) or (isinstance(values, list) and values)):
        raise ValueError("row or values is required for append_row")
    table = None
    min_column, max_column = 1, worksheet.max_column
    next_row = worksheet.max_row + 1 if worksheet._cells else 1
    template_row = None
    columns = None
    if "table" in patch:
        name = patch["table"]
        if not isinstance(name, str) or not name:
            raise ValueError("table must be a non-empty name")
        if "header_row" in patch or "inherit_from_row" in patch:
            raise ValueError("Named tables own their headers and template row; omit header_row/inherit_from_row")
        table = next((item for item in worksheet.tables.values() if item.name.lower() == name.lower()), None)
        if table is None:
            raise ValueError(f"Table not found in worksheet {worksheet.title}: {name}")
        if table.totalsRowCount or table.totalsRowShown:
            raise ValueError("Appending a table with a totals row is not supported without moving references")
        if table.connectionId is not None or table.tableType not in {None, "worksheet"}:
            raise ValueError("Cannot append a table backed by an external connection")
        min_column, first_row, max_column, last_row = range_boundaries(table.ref)
        if table.headerRowCount not in {0, 1} or len(table.tableColumns) != max_column - min_column + 1:
            raise ValueError("Table headers/columns do not match its range")
        first_data_row = first_row + table.headerRowCount
        template_row = last_row if last_row >= first_data_row else None
        next_row = last_row + 1
        columns = spreadsheet_column_map((min_column + i, col.name) for i, col in enumerate(table.tableColumns))
    elif "inherit_from_row" in patch:
        template_row = _integer(patch["inherit_from_row"], "inherit_from_row", 1, worksheet.max_row)
        if not any(row == template_row for row, _ in worksheet._cells):
            raise ValueError("inherit_from_row must reference an existing row")

    if next_row > EXCEL_MAX_ROW or max_column > EXCEL_MAX_COLUMN:
        raise ValueError("No room to append within Excel limits")
    explicit = {}
    if isinstance(row_object, dict):
        if columns is None:
            columns = spreadsheet_header_map(worksheet, patch.get("header_row", 1))
        for name, value in row_object.items():
            column = spreadsheet_column_index(columns, str(name))
            if column in explicit:
                raise ValueError(f"Duplicate update for column: {name}")
            explicit[column] = spreadsheet_cell_value(value)
    else:
        if len(values) > (max_column - min_column + 1 if table is not None else EXCEL_MAX_COLUMN):
            raise ValueError("values exceed the target table/Excel column limit")
        explicit = {min_column + i: spreadsheet_cell_value(value) for i, value in enumerate(values)}

    if table is not None:
        for area in list(worksheet.merged_cells.ranges) + [item.ref for item in worksheet.tables.values() if item is not table]:
            left, top, right, bottom = range_boundaries(str(area))
            if top <= next_row <= bottom and left <= max_column and right >= min_column:
                raise ValueError("Table append overlaps another table or merged range")
        for column in range(min_column, max_column + 1):
            cell = worksheet._cells.get((next_row, column))
            if cell is not None and (cell.value is not None or cell.comment is not None or cell.hyperlink is not None):
                raise ValueError(f"Table append would overwrite occupied cell {cell.coordinate}")

    planned = dict(explicit)
    styles = {}
    if template_row is not None or table is not None:
        for column in range(min_column, max_column + 1):
            source = worksheet._cells.get((template_row, column)) if template_row else None
            if source is not None and source.has_style:
                styles[column] = copy.copy(source._style)
            if column in explicit:
                continue
            declared = table.tableColumns[column - min_column].calculatedColumnFormula if table is not None else None
            target = f"{get_column_letter(column)}{next_row}"
            if declared is not None:
                if declared.array or not declared.text:
                    raise ValueError("Cannot inherit an array or empty calculated-column formula")
                formula = declared.text if declared.text.startswith("=") else "=" + declared.text
                planned[column] = _translate_spreadsheet_formula(formula, f"{get_column_letter(column)}{first_data_row}", target)
            elif source is not None and source.data_type == "f":
                if not isinstance(source.value, str):
                    raise ValueError("Cannot inherit array/data-table formulas into a single cell")
                planned[column] = _translate_spreadsheet_formula(source.value, source.coordinate, target)

    updated_cells = {}
    for column in sorted(planned.keys() | styles.keys()):
        cell = spreadsheet_cell(worksheet, f"{get_column_letter(column)}{next_row}")
        if column in styles:
            cell._style = styles[column]
        if column in planned:
            value = spreadsheet_cell_value(planned[column])
            cell.value = value
            updated_cells[cell.coordinate] = {"old": None, "new": value}
    result = {"updated": True, "operation": "append_row", "sheet": worksheet.title,
              "row_number": next_row, "updated_cells": updated_cells}
    if template_row is not None:
        result["inherited_from_row"] = template_row
    if table is not None:
        def expand_sort_ref(reference: str) -> str:
            left, top, right, bottom = range_boundaries(reference)
            if min_column <= left <= right <= max_column and first_row <= top <= bottom == last_row:
                return f"{get_column_letter(left)}{top}:{get_column_letter(right)}{next_row}"
            return reference

        for state in (table.sortState, table.autoFilter.sortState if table.autoFilter is not None else None):
            if state is not None and not state.columnSort:
                state.ref = expand_sort_ref(state.ref)
                for condition in state.sortCondition:
                    condition.ref = expand_sort_ref(condition.ref)
        table.ref = f"{get_column_letter(min_column)}{first_row}:{get_column_letter(max_column)}{next_row}"
        if table.autoFilter is not None:
            table.autoFilter.ref = table.ref
        result.update(table=table.name, table_ref=table.ref)
    return result


def _spreadsheet_coordinate(reference: Any) -> tuple[int, int]:
    from openpyxl.utils.cell import coordinate_to_tuple

    if not isinstance(reference, str) or not re.fullmatch(r"[A-Za-z]{1,3}[1-9][0-9]{0,6}", reference):
        raise ValueError("cell must be an A1 coordinate")
    row, column = coordinate_to_tuple(reference)
    if column > EXCEL_MAX_COLUMN or row > EXCEL_MAX_ROW:
        raise ValueError("cell is outside Excel limits")
    return row, column


def spreadsheet_cell(worksheet: Any, reference: Any) -> Any:
    from openpyxl.cell.cell import MergedCell

    row, column = _spreadsheet_coordinate(reference)
    cell = worksheet.cell(row=row, column=column)
    if isinstance(cell, MergedCell):
        merged = next(area for area in worksheet.merged_cells.ranges if cell.coordinate in area)
        raise ValueError(f"Cell {cell.coordinate} is merged; target its anchor {merged.start_cell.coordinate}")
    return cell


def _validated_cell_format(style: Any, *, spreadsheet: bool) -> dict[str, Any]:
    allowed = {"bold", "italic", "font_size", "font_name", "font_color", "fill_color"}
    if spreadsheet:
        allowed |= {"number_format", "wrap_text", "underline", "strike", "horizontal", "vertical",
                    "indent", "rotation", "shrink_to_fit", "borders"}
    if not isinstance(style, dict) or not style or set(style) - allowed:
        raise ValueError(f"format must contain supported keys: {', '.join(sorted(allowed))}")
    result = dict(style)
    for key, value in style.items():
        if key in {"bold", "italic", "wrap_text", "underline", "strike", "shrink_to_fit"}:
            if type(value) is not bool:
                raise ValueError(f"{key} must be boolean")
        elif key in {"font_color", "fill_color"}:
            if not isinstance(value, str) or not re.fullmatch(r"[0-9a-fA-F]{6}", value):
                raise ValueError(f"{key} must be a six-digit RGB hex string without #")
            result[key] = value.upper()
        elif key == "font_size":
            result[key] = _number(value, key, 1, 409)
        elif key == "horizontal":
            if value not in ("general", "left", "center", "right", "fill", "justify", "centerContinuous", "distributed"):
                raise ValueError("Unsupported horizontal alignment")
        elif key == "vertical":
            if value not in ("top", "center", "bottom", "justify", "distributed"):
                raise ValueError("Unsupported vertical alignment")
        elif key in {"indent", "rotation"}:
            result[key] = _integer(value, key, 0, 250 if key == "indent" else 180)
        elif key == "borders":
            from openpyxl.styles import Side

            if not isinstance(value, dict) or not value or set(value) - {"left", "right", "top", "bottom"}:
                raise ValueError("borders must specify left/right/top/bottom sides")
            borders = {}
            for edge, side in value.items():
                if side is None:
                    borders[edge] = Side()
                    continue
                if not isinstance(side, dict) or set(side) != {"style", "color"}:
                    raise ValueError("A border requires style and RGB color; null clears it")
                color = _validated_cell_format({"font_color": side["color"]}, spreadsheet=False)["font_color"]
                if not isinstance(side["style"], str):
                    raise ValueError("Border style must be a string")
                borders[edge] = Side(style=side["style"], color="FF" + color)
            result[key] = borders
        elif not isinstance(value, str) or not value:
            raise ValueError(f"{key} must be a non-empty string")
        else:
            spreadsheet_cell_value(value)
    return result


def format_spreadsheet_cell(worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    from openpyxl.styles import PatternFill

    cell = spreadsheet_cell(worksheet, patch.get("cell"))
    style = _validated_cell_format(patch.get("format"), spreadsheet=True)
    font, alignment, border = copy.copy(cell.font), copy.copy(cell.alignment), copy.copy(cell.border)
    for key, value in style.items():
        if key in {"bold", "italic", "strike"}:
            setattr(font, key, value)
        elif key == "underline":
            font.underline = "single" if value else None
        elif key in {"wrap_text", "shrink_to_fit", "horizontal", "vertical", "indent", "rotation"}:
            setattr(alignment, "textRotation" if key == "rotation" else key, value)
        elif key == "borders":
            for edge, side in value.items():
                setattr(border, edge, side)
        elif key in {"font_color", "fill_color"}:
            if key == "font_color":
                font.color = "FF" + value
            else:
                cell.fill = PatternFill("solid", fgColor="FF" + value)
        elif key == "font_size":
            font.sz = value
        elif key == "font_name":
            font.name = value
        else:
            cell.number_format = value
    cell.font, cell.alignment, cell.border = font, alignment, border
    return {"updated": True, "operation": patch["operation"], "sheet": worksheet.title, "cell": cell.coordinate}


def _worksheet_layout_format(patch: dict[str, Any], allowed: set[str]) -> dict[str, Any]:
    if any(key in patch for key in ("cell", "table_index", "section_index", "index", "slide", "shape_id", "style")):
        raise ValueError("Worksheet layout uses only the sheet selector")
    style = patch.get("format")
    if not isinstance(style, dict) or not style or set(style) - allowed:
        raise ValueError(f"format must contain supported keys: {', '.join(sorted(allowed))}")
    return style


def format_spreadsheet_sheet(worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    from openpyxl.utils.cell import column_index_from_string, get_column_letter

    style = _worksheet_layout_format(patch, {"column_widths", "row_heights", "freeze_panes", "show_gridlines"})
    for key in ("column_widths", "row_heights"):
        if key not in style:
            continue
        dimensions = style[key]
        if not isinstance(dimensions, dict) or not dimensions or len(dimensions) > 200:
            raise ValueError(f"{key} must contain 1–200 explicit dimensions")
        seen = set()
        for label, raw_size in dimensions.items():
            is_column = key == "column_widths"
            pattern = r"[A-Za-z]{1,3}" if is_column else r"[1-9][0-9]{0,6}"
            if not isinstance(label, str) or not re.fullmatch(pattern, label):
                raise ValueError(f"Invalid {key} target: {label}")
            index = column_index_from_string(label) if is_column else int(label)
            _integer(index, key, 1, EXCEL_MAX_COLUMN if is_column else EXCEL_MAX_ROW)
            if index in seen:
                raise ValueError(f"Duplicate {key} target: {label}")
            seen.add(index)
            size = _number(raw_size, key, 0.1, 255 if is_column else 409)
            if not is_column:
                worksheet.row_dimensions[index].height = size
                continue
            # A single <col> can span a grouped range. Split it around the
            # target instead of creating overlapping dimensions or restyling peers.
            covering = [(name, dimension) for name, dimension in worksheet.column_dimensions.items()
                        if (dimension.min or column_index_from_string(name)) <= index
                        <= (dimension.max or column_index_from_string(name))]
            if len(covering) > 1:
                raise ValueError(f"Overlapping column dimensions at {label}")
            if covering:
                name, dimension = covering[0]
                first = dimension.min or column_index_from_string(name)
                last = dimension.max or first
                if first != last:
                    del worksheet.column_dimensions[name]
                    for start, end in ((first, index - 1), (index, index), (index + 1, last)):
                        if start > end:
                            continue
                        part = copy.copy(dimension)
                        part.index, part.min, part.max = get_column_letter(start), start, end
                        worksheet.column_dimensions[part.index] = part
            worksheet.column_dimensions[get_column_letter(index)].width = size
    if "freeze_panes" in style:
        from openpyxl.worksheet.views import Selection

        reference = style["freeze_panes"]
        if reference is not None:
            _spreadsheet_coordinate(reference)
            reference = reference.upper()
        view = worksheet.sheet_view
        active_pane = view.pane.activePane if view.pane is not None else None
        active = next((selection for selection in view.selection if selection.pane == active_pane), None)
        # The library adds pane selections on every assignment. Retain the
        # active selection, then rebuild pane-specific selections exactly once.
        selection = copy.copy(active) if active is not None else Selection()
        selection.pane = None
        view.selection = [selection]
        worksheet.freeze_panes = reference
    if "show_gridlines" in style:
        if type(style["show_gridlines"]) is not bool:
            raise ValueError("show_gridlines must be boolean")
        worksheet.sheet_view.showGridLines = style["show_gridlines"]
    return {"updated": True, "operation": patch["operation"], "sheet": worksheet.title}


def setup_spreadsheet_page(worksheet: Any, patch: dict[str, Any]) -> dict[str, Any]:
    from openpyxl.utils.cell import column_index_from_string
    from openpyxl.worksheet.properties import PageSetupProperties

    margins = {"margin_left": "left", "margin_right": "right", "margin_top": "top", "margin_bottom": "bottom",
               "header_distance": "header", "footer_distance": "footer"}
    style = _worksheet_layout_format(patch, set(margins) | {
        "orientation", "paper_size", "fit_width", "fit_height", "print_area", "repeat_rows", "repeat_columns",
    })
    papers = {"A3": worksheet.PAPERSIZE_A3, "A4": worksheet.PAPERSIZE_A4, "A5": worksheet.PAPERSIZE_A5,
              "letter": worksheet.PAPERSIZE_LETTER, "legal": worksheet.PAPERSIZE_LEGAL}
    for key, value in style.items():
        if key in margins:
            setattr(worksheet.page_margins, margins[key], _number(value, key, 0, 720) / 72)
        elif key == "orientation":
            if value not in ("portrait", "landscape"):
                raise ValueError("orientation must be portrait or landscape")
            worksheet.page_setup.orientation = value
        elif key == "paper_size":
            if not isinstance(value, str) or value not in papers:
                raise ValueError("paper_size must be A3/A4/A5/letter/legal")
            worksheet.page_setup.paperSize = papers[value]
        elif key in {"fit_width", "fit_height"}:
            setattr(worksheet.page_setup, "fitToWidth" if key == "fit_width" else "fitToHeight", _integer(value, key, 0, 32767))
            if worksheet.sheet_properties.pageSetUpPr is None:
                worksheet.sheet_properties.pageSetUpPr = PageSetupProperties()
            worksheet.sheet_properties.pageSetUpPr.fitToPage = True
            worksheet.page_setup.scale = None
        else:
            if value is not None:
                if not isinstance(value, str):
                    raise ValueError(f"{key} must be a local range or null")
                if key == "print_area":
                    endpoints = value.split(":")
                    if len(endpoints) not in (1, 2):
                        raise ValueError("print_area must be a single local A1 range")
                    first, last = _spreadsheet_coordinate(endpoints[0]), _spreadsheet_coordinate(endpoints[-1])
                    if first[0] > last[0] or first[1] > last[1]:
                        raise ValueError("print_area must have ordered endpoints")
                else:
                    rows = key == "repeat_rows"
                    pattern = r"([1-9][0-9]{0,6}):([1-9][0-9]{0,6})" if rows else r"([A-Za-z]{1,3}):([A-Za-z]{1,3})"
                    match = re.fullmatch(pattern, value)
                    if not match:
                        raise ValueError(f"Invalid {key} range")
                    convert = int if rows else column_index_from_string
                    first, last = map(convert, match.groups())
                    _integer(last, key, first, EXCEL_MAX_ROW if rows else EXCEL_MAX_COLUMN)
                value = value.upper()
            attribute = {"print_area": "print_area", "repeat_rows": "print_title_rows", "repeat_columns": "print_title_cols"}[key]
            if value is None and key != "print_area":
                # The public title setters ignore None rather than clearing it.
                setattr(worksheet, "_print_rows" if key == "repeat_rows" else "_print_cols", None)
            else:
                setattr(worksheet, attribute, value)
    return {"updated": True, "operation": patch["operation"], "sheet": worksheet.title}


def apply_pdf_patch_sequence(abs_path: str, patches: list[dict[str, Any]]) -> dict[str, Any]:
    from pypdf import PdfReader, PdfWriter

    with open(abs_path, "rb") as source:
        reader = PdfReader(source)
        if reader.is_encrypted:
            return {"error": "encrypted_pdf_not_editable"}
        fields = reader.get_fields() or {}
        if reader.trailer["/Root"].get("/Perms") or any(field.get("/FT") == "/Sig" for field in fields.values()):
            return {"error": "signed_pdf_not_editable", "hint": "Editing would invalidate digital signatures."}
        writer = PdfWriter(clone_from=reader)
        results = []
        try:
            for index, patch in enumerate(patches):
                try:
                    page = _integer(patch.get("page"), "page", 1, len(writer.pages))
                    degrees = _integer(patch.get("degrees"), "degrees", -360, 360)
                    if degrees % 90:
                        raise ValueError("degrees must be a multiple of 90")
                    writer.pages[page - 1].rotate(degrees)
                    results.append({"operation": "page.rotate", "page": page, "degrees": degrees})
                except ValueError as exc:
                    return {"error": str(exc), "operation_index": index}
            output = io.BytesIO()
            writer.write(output)
            return {
                "patched": True,
                "edited": True,
                "operation": "patch",
                "file_type": "pdf",
                "operations_applied": len(results),
                "operation_results": results,
                "_persisted_bytes": output.getvalue(),
            }
        finally:
            writer.close()


def _presentation_shape_details(shape: Any) -> dict[str, Any]:
    """Bounded explicit properties for choosing subsequent shape patches."""
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    properties = shape._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}spPr")
    if properties is None:
        return {}
    style: dict[str, Any] = {}

    def color(parent):
        if parent is None:
            return None
        rgb = parent.find(ns + "srgbClr")
        if rgb is None:
            return None
        alpha = rgb.find(ns + "alpha")
        return {"color": rgb.get("val"), "opacity": int(alpha.get("val")) / 100000 if alpha is not None else 1}

    for name, parent in (("fill", properties), ("line", properties.find(ns + "ln"))):
        if parent is None:
            continue
        if parent.find(ns + "noFill") is not None:
            style[name + "_color"] = None
        solid = color(parent.find(ns + "solidFill"))
        if solid:
            style.update({name + "_color": solid["color"], name + "_opacity": solid["opacity"]})
        if name == "line":
            if parent.get("w") is not None:
                style["line_width"] = int(parent.get("w")) / 12700
            dash = parent.find(ns + "prstDash")
            if dash is not None:
                style["line_dash"] = dash.get("val")
    gradient = properties.find(ns + "gradFill")
    if gradient is not None:
        linear = gradient.find(ns + "lin")
        stops = gradient.findall(ns + "gsLst/" + ns + "gs")
        if linear is not None and 2 <= len(stops) <= 16 and all(color(stop) is not None for stop in stops):
            style["gradient_fill"] = {"angle": int(linear.get("ang", "0")) / 60000,
                                      "stops": [{"position": int(stop.get("pos", "0")) / 100000, **color(stop)} for stop in stops]}
    preset = properties.find(ns + "prstGeom")
    preset_name = preset.get("prst") if preset is not None else None
    if shape._element.tag.endswith("}cxnSp") and preset_name is None:
        preset_name = "line"
    if preset_name == "roundRect":
        style["corner_radius"] = round(min(shape.width, shape.height) / 12700 * shape.adjustments[0], 2)
    if shape.has_text_frame:
        body = shape.text_frame._txBody.bodyPr
        for name, xml_name in (("left", "lIns"), ("right", "rIns"), ("top", "tIns"), ("bottom", "bIns")):
            if body.get(xml_name) is not None:
                style["margin_" + name] = int(body.get(xml_name)) / 12700
        if body.get("anchor") in {"t", "ctr", "b"}:
            style["vertical_alignment"] = {"t": "top", "ctr": "middle", "b": "bottom"}[body.get("anchor")]
        if body.get("wrap") in {"square", "none"}:
            style["word_wrap"] = body.get("wrap") == "square"
    effects = properties.find(ns + "effectLst")
    if effects is not None:
        shadow = effects.find(ns + "outerShdw")
        if shadow is None:
            style["shadow"] = None
        elif color(shadow):
            style["shadow"] = {**color(shadow), "blur": int(shadow.get("blurRad", "0")) / 12700,
                               "distance": int(shadow.get("dist", "0")) / 12700, "angle": int(shadow.get("dir", "0")) / 60000}
    return {"preset": preset_name, "format": style}


def _presentation_slide_details(slide: Any) -> dict[str, Any]:
    """Report only explicit editable slide background properties."""
    drawing = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    presentation = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    background = slide._element.cSld.find(presentation + "bg")
    if background is None:
        return {"follow_master_background": True, "format": {"background_color": None}}
    properties = background.find(presentation + "bgPr")
    if properties is None:
        return {"follow_master_background": False, "format": {}, "background_type": "theme_reference"}

    def color(parent: Any) -> dict[str, Any] | None:
        if parent is None:
            return None
        rgb = parent.find(drawing + "srgbClr")
        if rgb is None or not re.fullmatch(r"[0-9A-Fa-f]{6}", rgb.get("val", "")):
            return None
        alpha = rgb.find(drawing + "alpha")
        return {
            "color": rgb.get("val").upper(),
            "opacity": int(alpha.get("val")) / 100000 if alpha is not None else 1,
        }

    solid = color(properties.find(drawing + "solidFill"))
    if solid is not None:
        return {
            "follow_master_background": False,
            "background_type": "solid",
            "format": {"background_color": solid["color"]},
        }
    gradient = properties.find(drawing + "gradFill")
    if gradient is not None:
        linear = gradient.find(drawing + "lin")
        stops = gradient.findall(drawing + "gsLst/" + drawing + "gs")
        colors = [color(stop) for stop in stops]
        if linear is not None and 2 <= len(stops) <= 16 and all(colors):
            return {
                "follow_master_background": False,
                "background_type": "gradient",
                "format": {
                    "background_gradient": {
                        "angle": int(linear.get("ang", "0")) / 60000,
                        "stops": [
                            {
                                "position": int(stop.get("pos", "0")) / 100000,
                                **stop_color,
                            }
                            for stop, stop_color in zip(stops, colors)
                        ],
                    },
                },
            }
    return {"follow_master_background": False, "format": {}, "background_type": "unsupported_explicit"}


def _presentation_picture_details(shape: Any) -> dict[str, Any]:
    if shape._element.tag.rsplit("}", 1)[-1] != "pic":
        return {}
    blip = shape._element.blipFill.blip
    alpha = blip.find("{http://schemas.openxmlformats.org/drawingml/2006/main}alphaModFix")
    crop = {
        key: round(getattr(shape, f"crop_{key}"), 5)
        for key in ("left", "top", "right", "bottom")
    }
    return {
        "picture": {
            "crop": crop,
            "opacity": int(alpha.get("amt")) / 100000 if alpha is not None else 1,
            "alt_text": shape._element.nvPicPr.cNvPr.get("descr", ""),
        },
    }


def _presentation_chart_details(shape: Any, *, limit: int = 50) -> dict[str, Any]:
    if not getattr(shape, "has_chart", False):
        return {}
    from pptx.chart.axis import ValueAxis
    from pptx.enum.chart import XL_LABEL_POSITION, XL_LEGEND_POSITION

    chart = shape.chart
    chart_type = _presentation_chart_type_name_for_chart(chart)
    if _is_stock_chart_type(chart_type):
        with _presentation_stock_chart_as_line(chart):
            details = _presentation_chart_details(shape, limit=limit)
        details["chart"]["type"] = chart_type
        return details
    if chart_type is None:
        return {"chart": {
            "type": "unsupported",
            "categories": [],
            "series": [],
            "format": {
                "title": chart.chart_title.text_frame.text if chart.has_title else None,
                "style": chart.chart_style,
                "has_legend": chart.has_legend,
            },
            "truncated": False,
        }}
    plot = chart.plots[0]
    series = list(chart.series)
    is_scatter = _is_scatter_chart_type(chart_type)

    def xy_values(item: Any) -> list[int | float]:
        namespace = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
        source = item._element.find(namespace + "xVal")
        if source is None:
            return []
        cache = source.find(f"{namespace}numRef/{namespace}numCache")
        if cache is None:
            cache = source.find(namespace + "numLit")
        if cache is None:
            return []
        values = []
        for point in sorted(cache.findall(namespace + "pt"), key=lambda node: int(node.get("idx", "0"))):
            value = point.find(namespace + "v")
            try:
                number = float(value.text)
            except (AttributeError, TypeError, ValueError):
                continue
            values.append(int(number) if number.is_integer() else number)
        return values

    scatter_x_values = [xy_values(item) for item in series] if is_scatter else []
    categories = scatter_x_values[0] if scatter_x_values else list(plot.categories)

    def title_text(owner: Any) -> str | None:
        return owner.text_frame.text if owner is not None else None

    def explicit_color(target: Any) -> str | None:
        try:
            rgb = target.format.fill.fore_color.rgb
            return str(rgb) if rgb is not None else None
        except (AttributeError, TypeError, ValueError):
            return None

    legend_positions = {
        XL_LEGEND_POSITION.TOP: "top",
        XL_LEGEND_POSITION.BOTTOM: "bottom",
        XL_LEGEND_POSITION.LEFT: "left",
        XL_LEGEND_POSITION.RIGHT: "right",
    }
    label_positions = {
        XL_LABEL_POSITION.BEST_FIT: "best_fit",
        XL_LABEL_POSITION.CENTER: "center",
        XL_LABEL_POSITION.INSIDE_END: "inside_end",
        XL_LABEL_POSITION.OUTSIDE_END: "outside_end",
    }
    chart_format: dict[str, Any] = {
        "title": title_text(chart.chart_title) if chart.has_title else None,
        "style": chart.chart_style,
        "has_legend": chart.has_legend,
        "vary_colors": plot.vary_by_categories,
        "show_data_labels": _presentation_plot_data_labels(plot) is not None,
    }
    if chart.has_legend:
        chart_format.update({
            "legend_position": legend_positions.get(chart.legend.position),
            "legend_include_in_layout": chart.legend.include_in_layout,
        })
    labels = _presentation_plot_data_labels(plot)
    if labels is not None:
        chart_format.update({
            "show_category_name": labels.show_category_name,
            "show_series_name": labels.show_series_name,
            "show_value": labels.show_value,
            "show_percentage": labels.show_percentage,
            "data_label_position": label_positions.get(labels.position),
        })
    _category_node, primary_value_node, secondary_value_node = _presentation_category_value_axis_nodes(chart)
    for name, attribute in (("category", "category_axis"), ("value", "value_axis")):
        try:
            axis = (
                ValueAxis(primary_value_node)
                if name == "value" and _category_node is not None and primary_value_node is not None
                else getattr(chart, attribute)
            )
        except ValueError:
            continue
        chart_format[f"show_{name}_axis"] = axis.visible
        chart_format[f"{name}_axis_title"] = title_text(axis.axis_title) if axis.has_title else None
        if name == "value" or (name == "category" and is_scatter):
            chart_format.update({
                f"{name}_axis_min": axis.minimum_scale,
                f"{name}_axis_max": axis.maximum_scale,
                f"{name}_axis_major_unit": axis.major_unit,
            })
        if name == "value":
            chart_format["show_major_gridlines"] = axis.has_major_gridlines
    if secondary_value_node is not None:
        secondary_axis = ValueAxis(secondary_value_node)
        chart_format.update({
            "show_secondary_value_axis": secondary_axis.visible,
            "secondary_value_axis_title": (
                title_text(secondary_axis.axis_title) if secondary_axis.has_title else None
            ),
            "secondary_value_axis_min": secondary_axis.minimum_scale,
            "secondary_value_axis_max": secondary_axis.maximum_scale,
            "secondary_value_axis_major_unit": secondary_axis.major_unit,
            "show_secondary_major_gridlines": secondary_axis.has_major_gridlines,
        })
    chart_format["series_colors"] = [explicit_color(item) for item in series[:limit]]
    series_details = []
    for index, item in enumerate(series[:limit]):
        series_detail = {
            "name": str(item.name)[:255],
            "values": list(item.values)[:limit],
        }
        if is_scatter:
            series_detail["x_values"] = scatter_x_values[index][:limit]
        trendline = _presentation_chart_trendline_details(item)
        if trendline is not None:
            series_detail["trendline"] = trendline
        error_bars = _presentation_chart_error_bar_details(item)
        if error_bars is not None:
            series_detail["error_bars"] = error_bars
        series_details.append(series_detail)
    details = {
        "type": chart_type,
        "categories": (
            categories[:limit] if is_scatter
            else [str(category)[:1000] for category in categories[:limit]]
        ),
        "series": series_details,
        "format": chart_format,
        "truncated": (
            len(categories) > limit or len(series) > limit
            or any(len(item.values) > limit for item in series)
            or any(len(values) > limit for values in scatter_x_values)
        ),
    }
    return {"chart": details}


def _spreadsheet_chart_title(chart: Any) -> str | None:
    title = getattr(chart, "title", None)
    rich = getattr(getattr(title, "tx", None), "rich", None)
    if rich is None:
        return None
    parts = []
    for paragraph in getattr(rich, "p", ()) or ():
        parts.extend(str(getattr(run, "t", "") or "") for run in getattr(paragraph, "r", ()) or ())
        parts.extend(str(getattr(field, "t", "") or "") for field in getattr(paragraph, "fld", ()) or ())
    return "".join(parts)[:500] or None


def _spreadsheet_chart_reference_values(workbook: Any, formula: Any, limit: int) -> tuple[list[Any], bool]:
    from openpyxl.utils.cell import range_to_tuple

    if not isinstance(formula, str) or not formula:
        return [], False
    try:
        sheet_name, (min_column, min_row, max_column, max_row) = range_to_tuple(formula)
        worksheet = workbook[sheet_name]
    except (KeyError, ValueError):
        return [], False
    total = (max_column - min_column + 1) * (max_row - min_row + 1)
    values = []
    for row in range(min_row, max_row + 1):
        for column in range(min_column, max_column + 1):
            if len(values) == limit:
                return values, True
            cell = worksheet._cells.get((row, column))
            value = cell.value if cell is not None else None
            values.append(value if value is None or type(value) in {str, int, float, bool} else str(value))
    return values, total > limit


def _spreadsheet_chart_reference(source: Any, *names: str) -> str | None:
    for name in names:
        reference = getattr(source, name, None)
        formula = getattr(reference, "f", None)
        if isinstance(formula, str) and formula:
            return formula
    return None


def _spreadsheet_chart_color(graphical_properties: Any) -> str | None:
    choice = getattr(graphical_properties, "solidFill", None)
    value = getattr(choice, "srgbClr", None)
    return value if isinstance(value, str) else None


def _spreadsheet_picture_details(image: Any, index: int) -> dict[str, Any]:
    from openpyxl.utils.cell import get_column_letter

    size = _spreadsheet_picture_size(image)
    anchor = image.anchor
    target = getattr(anchor, "to", None)
    frame = _spreadsheet_picture_frame(image)
    properties = frame.nvPicPr.cNvPr
    return {
        "picture_index": index,
        "anchor": _spreadsheet_drawing_anchor_cell(image),
        "to_anchor": (
            f"{get_column_letter(target.col + 1)}{target.row + 1}"
            if target is not None else None
        ),
        "anchor_type": anchor.__class__.__name__,
        "width": size[0] if size is not None else None,
        "height": size[1] if size is not None else None,
        "intrinsic_width_px": image.width,
        "intrinsic_height_px": image.height,
        "name": properties.name,
        "alt_text": properties.descr or "",
        "format": image.format,
    }


def _spreadsheet_conditional_color_details(color: Any) -> Any:
    if color is None:
        return None
    kind = getattr(color, "type", None)
    value = getattr(color, kind, None) if isinstance(kind, str) else None
    if kind == "rgb" and isinstance(value, str):
        return value[-6:].upper()
    return {"type": kind, "value": value, "tint": getattr(color, "tint", 0)}


def _spreadsheet_conditional_threshold_details(threshold: Any) -> dict[str, Any]:
    result = {"type": threshold.type}
    if threshold.val is not None:
        result["value"] = threshold.val
    if threshold.gte is not None:
        result["gte"] = threshold.gte
    return result


def _spreadsheet_conditional_dxf_details(dxf: Any) -> dict[str, Any]:
    if dxf is None:
        return {}
    result: dict[str, Any] = {}
    font = dxf.font
    if font is not None:
        for key, attribute in (
            ("bold", "b"),
            ("italic", "i"),
            ("strike", "strike"),
        ):
            if getattr(font, attribute, None) is not None:
                result[key] = bool(getattr(font, attribute))
        if font.u is not None:
            result["underline"] = bool(font.u)
        if font.name is not None:
            result["font_name"] = font.name
        if font.sz is not None:
            result["font_size"] = font.sz
        if font.color is not None:
            result["font_color"] = _spreadsheet_conditional_color_details(font.color)
    fill = dxf.fill
    if fill is not None and getattr(fill, "patternType", None) == "solid":
        result["fill_color"] = _spreadsheet_conditional_color_details(fill.fgColor)
    border = dxf.border
    borders = {}
    if border is not None:
        for edge in ("left", "right", "top", "bottom"):
            side = getattr(border, edge, None)
            if side is not None and (side.style is not None or side.color is not None):
                borders[edge] = {
                    "style": side.style,
                    "color": _spreadsheet_conditional_color_details(side.color),
                }
    if borders:
        result["borders"] = borders
    return result


def _spreadsheet_conditional_format_details(
    container: Any,
    rule: Any,
    index: int,
    *,
    limit: int,
) -> tuple[dict[str, Any], bool]:
    native_to_public = {
        "cellIs": "cell",
        "expression": "formula",
        "colorScale": "color_scale",
        "dataBar": "data_bar",
        "iconSet": "icon_set",
    }
    ranges = sorted(str(area) for area in container.sqref.ranges)
    details: dict[str, Any] = {
        "index": index,
        "ranges": ranges[:limit],
        "type": native_to_public.get(rule.type, rule.type),
        "priority": rule.priority,
    }
    if rule.stopIfTrue is not None:
        details["stop_if_true"] = rule.stopIfTrue
    if rule.formula:
        details["formulas"] = list(rule.formula)[:limit]
    if rule.type == "cellIs":
        details["operator"] = rule.operator
    if rule.type in {"cellIs", "expression"}:
        style = _spreadsheet_conditional_dxf_details(rule.dxf)
        if style:
            details["format"] = style
    elif rule.type == "colorScale" and rule.colorScale is not None:
        details["thresholds"] = [
            _spreadsheet_conditional_threshold_details(item)
            for item in rule.colorScale.cfvo[:limit]
        ]
        details["colors"] = [
            _spreadsheet_conditional_color_details(item)
            for item in rule.colorScale.color[:limit]
        ]
    elif rule.type == "dataBar" and rule.dataBar is not None:
        details.update({
            "thresholds": [
                _spreadsheet_conditional_threshold_details(item)
                for item in rule.dataBar.cfvo[:limit]
            ],
            "color": _spreadsheet_conditional_color_details(rule.dataBar.color),
            "show_value": rule.dataBar.showValue,
            "min_length": rule.dataBar.minLength,
            "max_length": rule.dataBar.maxLength,
        })
    elif rule.type == "iconSet" and rule.iconSet is not None:
        details.update({
            "thresholds": [
                _spreadsheet_conditional_threshold_details(item)
                for item in rule.iconSet.cfvo[:limit]
            ],
            "icon_style": rule.iconSet.iconSet,
            "show_value": rule.iconSet.showValue,
            "reverse": rule.iconSet.reverse,
            "percent": rule.iconSet.percent,
        })
    item_count = max(
        len(ranges),
        len(rule.formula or ()),
        len(getattr(getattr(rule, "colorScale", None), "cfvo", ()) or ()),
        len(getattr(getattr(rule, "dataBar", None), "cfvo", ()) or ()),
        len(getattr(getattr(rule, "iconSet", None), "cfvo", ()) or ()),
    )
    return details, item_count > limit


def _spreadsheet_chart_details(workbook: Any, chart: Any, index: int, *, limit: int = 200) -> dict[str, Any]:
    from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor
    from openpyxl.utils.units import EMU_to_cm

    anchor = chart.anchor
    width = height = None
    if isinstance(anchor, OneCellAnchor):
        width = EMU_to_cm(anchor.ext.width) / 2.54 * 72
        height = EMU_to_cm(anchor.ext.height) / 2.54 * 72
    series_items, truncated = [], False
    categories = []
    categories_ref = None
    is_scatter = _is_scatter_chart_type(_spreadsheet_chart_type(chart))
    for series_index, series in enumerate(_spreadsheet_chart_series(chart)):
        if series_index >= limit:
            truncated = True
            break
        values_source = getattr(series, "val", None) or getattr(series, "yVal", None)
        values_ref = _spreadsheet_chart_reference(values_source, "numRef")
        values, values_truncated = _spreadsheet_chart_reference_values(workbook, values_ref, limit)
        name_ref = _spreadsheet_chart_reference(getattr(series, "tx", None), "strRef")
        name_values, name_truncated = _spreadsheet_chart_reference_values(workbook, name_ref, 1)
        name = getattr(getattr(series, "tx", None), "v", None)
        if name is None and name_values:
            name = name_values[0]
        if series_index == 0:
            category_source = getattr(series, "cat", None) or getattr(series, "xVal", None)
            categories_ref = _spreadsheet_chart_reference(category_source, "strRef", "numRef", "multiLvlStrRef")
            categories, categories_truncated = _spreadsheet_chart_reference_values(workbook, categories_ref, limit)
            truncated = truncated or categories_truncated
        series_item = {
            "index": series_index,
            "name": str(name)[:255] if name is not None else None,
            "name_ref": name_ref,
            "values_ref": values_ref,
            "values": values,
            "color": _spreadsheet_chart_color(getattr(series, "graphicalProperties", None)),
        }
        if is_scatter:
            x_values_ref = _spreadsheet_chart_reference(getattr(series, "xVal", None), "numRef")
            x_values, x_values_truncated = _spreadsheet_chart_reference_values(workbook, x_values_ref, limit)
            series_item.update({"x_values_ref": x_values_ref, "x_values": x_values})
            truncated = truncated or x_values_truncated
        trendline = getattr(series, "trendline", None)
        if trendline is not None:
            public_types = {
                "linear": "linear", "exp": "exponential", "log": "logarithmic",
                "poly": "polynomial", "power": "power", "movingAvg": "moving_average",
            }
            trendline_details = {"type": public_types.get(trendline.trendlineType)}
            for attribute, public in (
                ("name", "name"), ("order", "order"), ("period", "period"),
                ("forward", "forward"), ("backward", "backward"), ("intercept", "intercept"),
                ("dispEq", "display_equation"), ("dispRSqr", "display_r_squared"),
            ):
                value = getattr(trendline, attribute, None)
                if value is not None:
                    trendline_details[public] = value
            series_item["trendline"] = trendline_details
        error_bars = getattr(series, "errBars", None)
        if error_bars is not None:
            public_types = {
                "fixedVal": "fixed", "percentage": "percentage",
                "stdDev": "standard_deviation", "stdErr": "standard_error",
            }
            error_bar_details = {
                "type": public_types.get(error_bars.errValType),
                "direction": error_bars.errDir or "y",
                "side": error_bars.errBarType or "both",
                "end_style": "no_cap" if error_bars.noEndCap else "cap",
            }
            if error_bars.val is not None:
                error_bar_details["value"] = error_bars.val
            series_item["error_bars"] = error_bar_details
        series_items.append(series_item)
        truncated = truncated or values_truncated or name_truncated
    legend = getattr(chart, "legend", None)
    labels = getattr(chart, "dLbls", None)
    chart_format = {
        "title": _spreadsheet_chart_title(chart),
        "style": getattr(chart, "style", None),
        "has_legend": legend is not None,
        "legend_position": ({"t": "top", "b": "bottom", "l": "left", "r": "right"}.get(
            getattr(legend, "position", None),
        ) if legend is not None else None),
        "legend_include_in_layout": (not bool(legend.overlay) if legend is not None else None),
        "show_data_labels": labels is not None,
        "show_category_name": getattr(labels, "showCatName", None),
        "show_series_name": getattr(labels, "showSerName", None),
        "show_value": getattr(labels, "showVal", None),
        "show_percentage": getattr(labels, "showPercent", None),
        "series_colors": [item["color"] for item in series_items],
    }
    for name, axis in (("category", getattr(chart, "x_axis", None)), ("value", getattr(chart, "y_axis", None))):
        if axis is None:
            continue
        chart_format.update({
            f"show_{name}_axis": not bool(getattr(axis, "delete", False)),
            f"{name}_axis_title": _spreadsheet_chart_title(axis),
        })
        if name == "value" or (name == "category" and is_scatter):
            chart_format.update({
                f"{name}_axis_min": axis.scaling.min,
                f"{name}_axis_max": axis.scaling.max,
                f"{name}_axis_major_unit": axis.majorUnit,
            })
    value_axis = getattr(chart, "y_axis", None)
    if value_axis is not None:
        chart_format["show_major_gridlines"] = value_axis.majorGridlines is not None
    components = _spreadsheet_chart_components(chart)
    if _is_volume_stock_chart_type(_spreadsheet_chart_type(chart)) and len(components) > 1:
        secondary_axis = components[1].y_axis
        chart_format.update({
            "show_secondary_value_axis": not bool(getattr(secondary_axis, "delete", False)),
            "secondary_value_axis_title": _spreadsheet_chart_title(secondary_axis),
            "secondary_value_axis_min": secondary_axis.scaling.min,
            "secondary_value_axis_max": secondary_axis.scaling.max,
            "secondary_value_axis_major_unit": secondary_axis.majorUnit,
            "show_secondary_major_gridlines": secondary_axis.majorGridlines is not None,
        })
    return {
        "index": index,
        "type": _spreadsheet_chart_type(chart) or chart.__class__.__name__,
        "anchor": _spreadsheet_drawing_anchor_cell(chart),
        "width": width,
        "height": height,
        "categories_ref": categories_ref,
        "categories": categories,
        "series": series_items,
        "format": chart_format,
        "truncated": truncated,
    }


def _word_table_structure(table_element: Any, index: int, *, style_name: str | None = None) -> dict[str, Any]:
    from docx.oxml.ns import qn

    properties = table_element.find(qn("w:tblPr"))
    grid = table_element.find(qn("w:tblGrid"))
    rows = table_element.findall(qn("w:tr"))

    def child(parent: Any, name: str) -> Any:
        return parent.find(qn(f"w:{name}")) if parent is not None else None

    def points(element: Any, attribute: str = "w") -> float | None:
        if element is None:
            return None
        raw = element.get(qn(f"w:{attribute}"))
        try:
            return int(raw) / 20 if raw is not None else None
        except ValueError:
            return None

    width = child(properties, "tblW")
    indent = child(properties, "tblInd")
    row_heights = [points(child(child(row, "trPr"), "trHeight"), "val") for row in rows]
    public_height_rules = {"atLeast": "at_least", "exact": "exact", "auto": "auto"}
    row_height_rules = []
    for row in rows:
        height = child(child(row, "trPr"), "trHeight")
        native_rule = height.get(qn("w:hRule")) if height is not None else None
        row_height_rules.append(public_height_rules.get(native_rule, native_rule))
    repeat_header_rows = 0
    for row in rows:
        if child(child(row, "trPr"), "tblHeader") is None:
            break
        repeat_header_rows += 1
    result = {
        "index": index,
        "row_count": len(rows),
        "column_count": len(list(grid)) if grid is not None else 0,
        "style_id": child(properties, "tblStyle").get(qn("w:val")) if child(properties, "tblStyle") is not None else None,
        "alignment": child(properties, "jc").get(qn("w:val")) if child(properties, "jc") is not None else None,
        "autofit": child(properties, "tblLayout") is None or child(properties, "tblLayout").get(qn("w:type")) != "fixed",
        "width": points(width) if width is not None and width.get(qn("w:type")) == "dxa" else None,
        "indent": points(indent) if indent is not None and indent.get(qn("w:type")) == "dxa" else None,
        "column_widths": [points(column) for column in list(grid)] if grid is not None else [],
        "row_heights": row_heights,
        "row_height_rules": row_height_rules,
        "repeat_header_rows": repeat_header_rows,
    }
    if style_name is not None:
        result["style_name"] = style_name
    return result


def _presentation_table_structure(table: Any) -> dict[str, Any]:
    namespace = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    style_id = table._tbl.tblPr.find(namespace + "tableStyleId")
    return {
        "row_count": len(table.rows),
        "column_count": len(table.columns),
        "column_widths": [column.width.pt for column in table.columns],
        "row_heights": [row.height.pt for row in table.rows],
        "first_row": table.first_row,
        "last_row": table.last_row,
        "first_column": table.first_col,
        "last_column": table.last_col,
        "banded_rows": table.horz_banding,
        "banded_columns": table.vert_banding,
        "style_id": style_id.text if style_id is not None else None,
    }


def describe_file_structure(abs_path: str) -> dict[str, Any]:
    """Bounded selectors for structural edits; pair with read_file's byte hash."""
    extension = file_type_from_path(abs_path)
    limit = 200
    if extension == "docx":
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.enum.style import WD_STYLE_TYPE
        from docx.oxml.ns import qn

        document = Document(abs_path)
        paragraph_style_objects = [
            style for style in document.styles if style.type == WD_STYLE_TYPE.PARAGRAPH
        ]
        paragraph_styles = [style.name for style in paragraph_style_objects]
        word_pictures = _word_pictures(document)
        word_text_boxes = _word_text_boxes(document)
        paragraph_indices = {
            part: {
                paragraph: index
                for index, paragraph in enumerate(
                    [item._p for item in document.paragraphs] if stories == ("body",) else root.findall(qn("w:p"))
                )
            }
            for part, root, stories, _section_indices in _word_story_parts(document)
        }
        pictures = []
        for picture_index, picture in enumerate(word_pictures[:limit]):
            paragraph_element = _word_picture_paragraph_element(picture)
            alignment = getattr(getattr(paragraph_element.pPr, "jc", None), "val", None) if paragraph_element.pPr is not None else None
            item = {
                "picture_index": picture_index,
                "story": picture.stories[0],
                "stories": list(picture.stories),
                "section_indices": list(picture.section_indices),
                "layout": picture.layout,
                "paragraph_index": paragraph_indices.get(picture.part, {}).get(paragraph_element),
                "width": picture.width / 12700,
                "height": picture.height / 12700,
                "name": picture.doc_properties.get("name", ""),
                "alt_text": picture.doc_properties.get("descr", ""),
                "alignment": (
                    {0: "left", 1: "center", 2: "right"}.get(int(alignment))
                    if picture.layout == "inline" and alignment is not None else None
                ),
            }
            if picture.layout == "floating":
                horizontal = picture.element.find(qn("wp:positionH"))
                vertical = picture.element.find(qn("wp:positionV"))
                horizontal_offset = horizontal.find(qn("wp:posOffset")) if horizontal is not None else None
                vertical_offset = vertical.find(qn("wp:posOffset")) if vertical is not None else None
                item["floating"] = {
                    "position_x": (int(horizontal_offset.text or 0) / 12700 if horizontal_offset is not None else None),
                    "position_y": (int(vertical_offset.text or 0) / 12700 if vertical_offset is not None else None),
                    "relative_from_horizontal": horizontal.get("relativeFrom") if horizontal is not None else None,
                    "relative_from_vertical": vertical.get("relativeFrom") if vertical is not None else None,
                    "wrap": _word_picture_wrap(picture.element),
                    "behind_text": picture.element.get("behindDoc", "0") in {"1", "true", "on"},
                    "allow_overlap": picture.element.get("allowOverlap", "1") in {"1", "true", "on"},
                    "layout_in_cell": picture.element.get("layoutInCell", "1") in {"1", "true", "on"},
                    **{
                        key: int(picture.element.get(attribute, "0")) / 12700
                        for key, attribute in {
                            "distance_top": "distT", "distance_bottom": "distB",
                            "distance_left": "distL", "distance_right": "distR",
                        }.items()
                    },
                }
            pictures.append(item)
        text_boxes = []
        for text_box_index, text_box in enumerate(word_text_boxes[:limit]):
            frame = text_box.frame_element
            shape = text_box.shape_element
            group_elements = text_box.group_elements
            explicit_group = bool(group_elements)
            extent = frame.find(qn("wp:extent")) if frame.tag.endswith(("}anchor", "}inline")) else None
            horizontal = frame.find(qn("wp:positionH")) if frame.tag.endswith("}anchor") else None
            vertical = frame.find(qn("wp:positionV")) if frame.tag.endswith("}anchor") else None
            horizontal_offset = horizontal.find(qn("wp:posOffset")) if horizontal is not None else None
            vertical_offset = vertical.find(qn("wp:posOffset")) if vertical is not None else None
            preset = next((node.get("prst") for node in shape.iter() if node.tag.endswith("}prstGeom")), None)
            vml_style = frame.get("style", "") if frame.tag.endswith(("}shape", "}rect")) else ""

            def vml_points(name: str) -> float | None:
                match = re.search(rf"(?:^|;)\s*{re.escape(name)}\s*:\s*(-?[0-9.]+)pt\b", vml_style, re.I)
                return float(match.group(1)) if match else None

            paragraph_elements = text_box.content.findall(qn("w:p"))
            table_elements = text_box.content.findall(qn("w:tbl"))
            text_box_format: dict[str, Any] = {}
            rotation = None
            member_offset = None
            member_extent = None
            if frame.tag.endswith(("}anchor", "}inline")):
                properties = frame.find(qn("wp:docPr"))
                shape_properties = next((node for node in shape if node.tag.endswith("}spPr")), None)
                body_properties = next((node for node in shape if node.tag.endswith("}bodyPr")), None)
                transform = (
                    next((node for node in shape_properties if node.tag.endswith("}xfrm")), None)
                    if shape_properties is not None else None
                )
                if transform is not None and transform.get("rot") is not None:
                    rotation = int(transform.get("rot")) / 60000
                if explicit_group and transform is not None:
                    member_offset = next(
                        (node for node in transform if node.tag.endswith("}off")), None,
                    )
                    member_extent = next(
                        (node for node in transform if node.tag.endswith("}ext")), None,
                    )

                def explicit_fill(owner: Any) -> tuple[str | None, float | None] | None:
                    if owner is None:
                        return None
                    if next((node for node in owner if node.tag.endswith("}noFill")), None) is not None:
                        return None, None
                    fill = next((node for node in owner if node.tag.endswith("}solidFill")), None)
                    if fill is None:
                        return None
                    rgb = next((node for node in fill if node.tag.endswith("}srgbClr")), None)
                    if rgb is None:
                        return None
                    alpha = next((node for node in rgb if node.tag.endswith("}alpha")), None)
                    return rgb.get("val"), (int(alpha.get("val")) / 100000 if alpha is not None else 1)

                fill = explicit_fill(shape_properties)
                if fill is not None:
                    text_box_format.update({"fill_color": fill[0], "fill_opacity": fill[1]})
                line = (
                    next((node for node in shape_properties if node.tag.endswith("}ln")), None)
                    if shape_properties is not None else None
                )
                line_fill = explicit_fill(line)
                if line_fill is not None:
                    text_box_format.update({"line_color": line_fill[0], "line_opacity": line_fill[1]})
                if line is not None and line.get("w") is not None:
                    text_box_format["line_width"] = int(line.get("w")) / 12700
                if body_properties is not None:
                    for key, attribute in {
                        "margin_left": "lIns", "margin_right": "rIns",
                        "margin_top": "tIns", "margin_bottom": "bIns",
                    }.items():
                        if body_properties.get(attribute) is not None:
                            text_box_format[key] = int(body_properties.get(attribute)) / 12700
                    text_box_format["vertical_alignment"] = {
                        "t": "top", "ctr": "middle", "b": "bottom",
                    }.get(body_properties.get("anchor", "t"), body_properties.get("anchor"))
                    text_box_format["word_wrap"] = body_properties.get("wrap", "square") != "none"
                if properties is not None:
                    text_box_format["alt_text"] = properties.get("descr", "")
                if frame.tag.endswith("}anchor"):
                    text_box_format.update({
                        "wrap": _word_picture_wrap(frame),
                        "behind_text": frame.get("behindDoc", "0") in {"1", "true", "on"},
                        "allow_overlap": frame.get("allowOverlap", "1") in {"1", "true", "on"},
                        "layout_in_cell": frame.get("layoutInCell", "1") in {"1", "true", "on"},
                        **{
                            key: int(frame.get(attribute, "0")) / 12700
                            for key, attribute in {
                                "distance_top": "distT", "distance_bottom": "distB",
                                "distance_left": "distL", "distance_right": "distR",
                            }.items()
                        },
                    })
            else:
                rotation_match = re.search(
                    r"(?:^|;)\s*rotation\s*:\s*(-?[0-9.]+)(?:deg)?\b", vml_style, re.I,
                )
                rotation = float(rotation_match.group(1)) if rotation_match else None
                text_box_format = {
                    "fill_color": None if frame.get("filled") == "f" else frame.get("fillcolor"),
                    "line_color": None if frame.get("stroked") == "f" else frame.get("strokecolor"),
                    "line_width": vml_points("strokeweight") or (
                        float(frame.get("strokeweight", "0pt").removesuffix("pt"))
                        if frame.get("strokeweight", "").endswith("pt") else None
                    ),
                    "alt_text": frame.get("title", ""),
                }
            text_boxes.append({
                "text_box_index": text_box_index,
                "story": text_box.stories[0],
                "stories": list(text_box.stories),
                "section_indices": list(text_box.section_indices),
                "grouped": text_box.is_grouped,
                "group_depth": len(group_elements),
                "coordinate_space": "group_local" if explicit_group else "document",
                "layout": text_box.layout,
                "paragraph_index": paragraph_indices.get(text_box.part, {}).get(text_box.paragraph_element),
                "preset": preset or ("rect" if frame.tag.endswith("}rect") else "shape"),
                "width": (
                    int(member_extent.get("cx", "0")) / 12700 if member_extent is not None
                    else int(extent.get("cx", "0")) / 12700 if extent is not None
                    else vml_points("width")
                ),
                "height": (
                    int(member_extent.get("cy", "0")) / 12700 if member_extent is not None
                    else int(extent.get("cy", "0")) / 12700 if extent is not None
                    else vml_points("height")
                ),
                "position_x": (
                    int(member_offset.get("x", "0")) / 12700 if member_offset is not None
                    else int(horizontal_offset.text or 0) / 12700 if horizontal_offset is not None
                    else vml_points("left" if explicit_group else "margin-left")
                ),
                "position_y": (
                    int(member_offset.get("y", "0")) / 12700 if member_offset is not None
                    else int(vertical_offset.text or 0) / 12700 if vertical_offset is not None
                    else vml_points("top" if explicit_group else "margin-top")
                ),
                "rotation": rotation,
                "format": text_box_format,
                "paragraph_count": len(paragraph_elements),
                "paragraphs": [
                    {
                        "index": index,
                        "text": "".join((node.text or "") for node in paragraph.iter(qn("w:t")))[:200],
                    }
                    for index, paragraph in enumerate(paragraph_elements[:limit])
                ],
                "table_count": len(table_elements),
                "tables": [
                    _word_table_structure(table, index)
                    for index, table in enumerate(table_elements[:limit])
                ],
            })
        sections = []
        section_start_types = {
            WD_SECTION.CONTINUOUS: "continuous",
            WD_SECTION.NEW_COLUMN: "new_column",
            WD_SECTION.NEW_PAGE: "new_page",
            WD_SECTION.EVEN_PAGE: "even_page",
            WD_SECTION.ODD_PAGE: "odd_page",
        }
        for i, section in enumerate(document.sections):
            if i >= limit:
                break
            attributes = {"width": "page_width", "height": "page_height", "header_distance": "header_distance", "footer_distance": "footer_distance"}
            attributes.update({"margin_" + side: side + "_margin" for side in ("top", "bottom", "left", "right")})
            sections.append({
                "index": i,
                "start_type": section_start_types.get(section.start_type),
                **{
                    key: getattr(section, attr).pt if getattr(section, attr) is not None else None
                    for key, attr in attributes.items()
                },
            })
        story_structures = []
        story_content_total = 0
        for _part, root, stories, section_indices in _word_story_parts(document)[1:]:
            paragraph_elements = root.findall(qn("w:p"))
            table_elements = root.findall(qn("w:tbl"))
            story_content_total += len(paragraph_elements) + len(table_elements)
            story_structures.append({
                "story": stories[0],
                "stories": list(stories),
                "section_indices": list(section_indices),
                "paragraph_count": len(paragraph_elements),
                "paragraphs": [
                    {
                        "index": index,
                        "text": "".join((node.text or "") for node in paragraph.iter(qn("w:t")))[:200],
                    }
                    for index, paragraph in enumerate(paragraph_elements[:limit])
                ],
                "table_count": len(table_elements),
                "tables": [
                    _word_table_structure(table, index)
                    for index, table in enumerate(table_elements[:limit])
                ],
            })
        return {
            "paragraph_count": len(document.paragraphs),
            "paragraphs": [{"index": i, "text": p.text[:200]} for i, p in enumerate(document.paragraphs[:limit])],
            "table_count": len(document.tables),
            "tables": [
                _word_table_structure(
                    table._tbl,
                    i,
                    style_name=table.style.name if table.style is not None else None,
                )
                for i, table in enumerate(document.tables[:limit])
            ],
            "cell_addressing": "cell.set/table.format: table_index plus A1 cell or table layout; merged continuations and omitted grid positions are rejected",
            "picture_count": len(word_pictures),
            "pictures": pictures,
            "picture_addressing": (
                "picture_index is zero-based across body, header and footer DrawingML pictures; body pictures come first, then "
                "header/footer parts by first section reference. story/section_indices identify shared page furniture; paragraph_index "
                "is present only for a direct paragraph in that story."
            ),
            "text_box_count": len(word_text_boxes),
            "text_boxes": text_boxes,
            "text_box_addressing": (
                "text_box_index is package-wide across body, header and footer text-box contents. Use it instead of story/section_index "
                "with paragraph, table and cell operations; their paragraph/table indices are local to that text box."
            ),
            "truncated": (
                len(document.paragraphs) > limit or len(document.tables) > limit or len(document.sections) > limit
                or len(paragraph_styles) > limit or len(word_pictures) > limit or len(word_text_boxes) > limit
                or story_content_total > limit
            ),
            "index_scope": (
                "top-level paragraphs/tables describe the body only; story_structures provides independent zero-based "
                "header/footer paragraph and table indices"
            ),
            "section_count": len(document.sections),
            "sections": sections,
            "story_structures": story_structures,
            "paragraph_styles": paragraph_styles[:limit],
            "paragraph_style_definitions": [
                _word_paragraph_style_details(style) for style in paragraph_style_objects[:limit]
            ],
            "paragraph_style_addressing": (
                "paragraph_style insert/format/delete uses the exact style name. Definitions report direct formatting; "
                "unspecified properties inherit through base_style."
            ),
            "layout_units": "points",
            "story_addressing": (
                "Word paragraph/table/cell operations default to body; pass story plus section_index for a header/footer. "
                "Linked later sections must target the defining prior section."
            ),
        }
    if extension == "pptx":
        from pptx import Presentation

        document = Presentation(abs_path)
        shapes = []
        total = 0
        remaining_paragraphs = limit
        text_truncated = False
        for number, slide in enumerate(document.slides, 1):
            for shape, parent_shape_ids in _iter_presentation_shapes(slide.shapes):
                total += 1
                if len(shapes) < limit:
                    siblings = list(_presentation_shape_tree(slide, parent_shape_ids))
                    z_index = next(
                        index for index, sibling in enumerate(siblings)
                        if sibling._element is shape._element
                    )
                    text_frame = {}
                    if shape.has_text_frame:
                        paragraphs = shape.text_frame.paragraphs
                        shown = paragraphs[:remaining_paragraphs]
                        remaining_paragraphs -= len(shown)
                        truncated = len(shown) < len(paragraphs)
                        text_truncated = text_truncated or truncated
                        text_frame = {"text_frame": {
                            "paragraph_count": len(paragraphs),
                            "paragraphs": [{"index": i, "text": p.text[:200]} for i, p in enumerate(shown)],
                            "truncated": truncated,
                        }}
                    shapes.append(
                        {
                            "slide": number,
                            "shape_id": shape.shape_id,
                            "group_path": [*parent_shape_ids, shape.shape_id],
                            "parent_shape_id": parent_shape_ids[-1] if parent_shape_ids else None,
                            "z_index": z_index,
                            "sibling_count": len(siblings),
                            "coordinate_space": (
                                f"group:{parent_shape_ids[-1]}" if parent_shape_ids else "slide"
                            ),
                            "name": shape.name,
                            "x": shape.left.pt,
                            "y": shape.top.pt,
                            "width": shape.width.pt,
                            "height": shape.height.pt,
                            "rotation": shape.rotation,
                            "flip_horizontal": bool(shape._element.xfrm.flipH),
                            "flip_vertical": bool(shape._element.xfrm.flipV),
                            **_presentation_shape_details(shape),
                            **_presentation_picture_details(shape),
                            **_presentation_chart_details(shape),
                            "has_text_frame": shape.has_text_frame,
                            **text_frame,
                            "has_table": shape.has_table,
                            "is_group": shape._element.tag.rsplit("}", 1)[-1] == "grpSp",
                            **(
                                {"group_member_count": len(shape.shapes)}
                                if shape._element.tag.rsplit("}", 1)[-1] == "grpSp"
                                else {}
                            ),
                            **({"table": _presentation_table_structure(shape.table)} if shape.has_table else {}),
                        }
                    )
        return {"slide_count": len(document.slides), "shape_count": total, "shapes": shapes, "units": "points",
                "slides": [
                    {"slide": number, **_presentation_slide_details(slide)}
                    for number, slide in enumerate(islice(document.slides, limit), 1)
                ],
                "page_size": {"width": document.slide_width.pt, "height": document.slide_height.pt},
                "layouts": [{"index": i, "name": layout.name} for i, layout in enumerate(document.slide_layouts) if i < limit],
                "shape_addressing": (
                    "shape_id is unique across the whole slide, including nested group members; nested x/y/width/height "
                    "use the parent group's local coordinate space shown by coordinate_space; z_index is relative to sibling_count"
                ),
                "text_addressing": "text.set: slide + shape_id + zero-based text-frame paragraph index; grouped text shapes use the same selector; newlines are soft breaks",
                "shape_format_scope": "Explicit supported properties; inherited theme styles are not materialized here.",
                "truncated": total > limit or len(document.slides) > limit or len(document.slide_layouts) > limit or text_truncated}
    if extension in {"xlsx", "xlsm"}:
        from openpyxl import load_workbook
        from openpyxl.utils.cell import get_column_letter

        workbook = load_workbook(abs_path)
        try:
            hydrate_spreadsheet_chart_styles(abs_path, workbook)
            total_chart_count = sum(len(sheet._charts) for sheet in workbook.worksheets)
            total_picture_count = sum(len(sheet._images) for sheet in workbook.worksheets)
            total_merge_count = sum(len(sheet.merged_cells.ranges) for sheet in workbook.worksheets)
            total_validation_count = sum(
                len(sheet.data_validations.dataValidation) for sheet in workbook.worksheets
            )
            total_conditional_format_count = sum(
                len(_spreadsheet_conditional_format_rules(sheet)) for sheet in workbook.worksheets
            )
            sheets, table_count = [], 0
            shown_chart_count = shown_picture_count = shown_merge_count = shown_validation_count = 0
            shown_conditional_format_count = 0
            dimension_budget = limit
            truncated = (
                len(workbook.worksheets) > limit
                or total_chart_count > limit
                or total_picture_count > limit
                or total_merge_count > limit
                or total_validation_count > limit
                or total_conditional_format_count > limit
            )
            for sheet in workbook.worksheets[:limit]:
                tables = []
                for table in sheet.tables.values():
                    table_count += 1
                    if table_count > limit:
                        truncated = True
                        continue
                    style = table.tableStyleInfo
                    tables.append({"name": table.name, "ref": table.ref, "columns": table.column_names[:limit],
                                   "column_count": len(table.tableColumns), "has_totals": bool(table.totalsRowCount or table.totalsRowShown),
                                   "style": None if style is None else {
                                       "style_name": style.name,
                                       "show_first_column": bool(style.showFirstColumn),
                                       "show_last_column": bool(style.showLastColumn),
                                       "show_row_stripes": bool(style.showRowStripes),
                                       "show_column_stripes": bool(style.showColumnStripes),
                                   }})
                    truncated = truncated or len(table.tableColumns) > limit
                columns = {}
                for label, dimension in islice(sheet.column_dimensions.items(), dimension_budget):
                    if dimension.min and dimension.max and dimension.min != dimension.max:
                        label = f"{get_column_letter(dimension.min)}:{get_column_letter(dimension.max)}"
                    columns[label] = dimension.width
                dimension_budget -= len(columns)
                rows = {str(index): dimension.height for index, dimension in islice(sheet.row_dimensions.items(), dimension_budget)}
                dimension_budget -= len(rows)
                truncated = truncated or len(columns) < len(sheet.column_dimensions) or len(rows) < len(sheet.row_dimensions)
                charts = []
                for chart_index, chart in enumerate(sheet._charts):
                    if shown_chart_count == limit:
                        truncated = True
                        continue
                    shown_chart_count += 1
                    details = _spreadsheet_chart_details(workbook, chart, chart_index, limit=limit)
                    charts.append(details)
                    truncated = truncated or details["truncated"]
                pictures = []
                for picture_index, picture in enumerate(sheet._images):
                    if shown_picture_count == limit:
                        truncated = True
                        continue
                    shown_picture_count += 1
                    pictures.append(_spreadsheet_picture_details(picture, picture_index))
                merged_ranges = []
                for area in sheet.merged_cells.ranges:
                    if shown_merge_count == limit:
                        truncated = True
                        continue
                    shown_merge_count += 1
                    merged_ranges.append(str(area))
                validations = []
                for validation_index, validation in enumerate(sheet.data_validations.dataValidation):
                    if shown_validation_count == limit:
                        truncated = True
                        continue
                    shown_validation_count += 1
                    ranges = sorted(str(area) for area in validation.ranges.ranges)
                    if len(ranges) > limit:
                        truncated = True
                    validations.append({
                        "index": validation_index,
                        "ranges": ranges[:limit],
                        "type": validation.type,
                        "operator": validation.operator,
                        "formula1": validation.formula1,
                        "formula2": validation.formula2,
                        "allow_blank": bool(validation.allowBlank),
                        "show_dropdown": not bool(validation.showDropDown),
                        "show_error_message": bool(validation.showErrorMessage),
                        "error_style": validation.errorStyle,
                        "error_title": validation.errorTitle,
                        "error": validation.error,
                        "show_input_message": bool(validation.showInputMessage),
                        "prompt_title": validation.promptTitle,
                        "prompt": validation.prompt,
                    })
                conditional_formats = []
                for conditional_format_index, (container, _rule_index, rule) in enumerate(
                    _spreadsheet_conditional_format_rules(sheet),
                ):
                    if shown_conditional_format_count == limit:
                        truncated = True
                        continue
                    details, details_truncated = _spreadsheet_conditional_format_details(
                        container,
                        rule,
                        conditional_format_index,
                        limit=limit,
                    )
                    shown_conditional_format_count += 1
                    conditional_formats.append(details)
                    truncated = truncated or details_truncated
                sheets.append({"name": sheet.title, "state": sheet.sheet_state,
                    "managed_chart_data": sheet["A1"].value == _SPREADSHEET_CHART_DATA_MARKER,
                    "tables": tables, "charts": charts, "pictures": pictures,
                    "merged_ranges": merged_ranges, "validations": validations,
                    "conditional_formats": conditional_formats, "layout": {
                    "column_widths": columns, "row_heights": rows,
                    "freeze_panes": sheet.freeze_panes, "show_gridlines": sheet.sheet_view.showGridLines,
                }, "page_setup": {
                    "orientation": sheet.page_setup.orientation, "paper_size_code": sheet.page_setup.paperSize,
                    "print_area": sheet.print_area, "repeat_rows": sheet.print_title_rows, "repeat_columns": sheet.print_title_cols,
                    "fit_width": sheet.page_setup.fitToWidth, "fit_height": sheet.page_setup.fitToHeight,
                    **{f"margin_{side}": getattr(sheet.page_margins, side) * 72 for side in ("left", "right", "top", "bottom")},
                    "header_distance": sheet.page_margins.header * 72, "footer_distance": sheet.page_margins.footer * 72,
                }})
            return {"sheet_count": len(workbook.worksheets), "chart_count": total_chart_count,
                    "picture_count": total_picture_count, "merge_count": total_merge_count,
                    "validation_count": total_validation_count,
                    "conditional_format_count": total_conditional_format_count,
                    "sheets": sheets, "truncated": truncated}
        finally:
            workbook.close()
    return {}
